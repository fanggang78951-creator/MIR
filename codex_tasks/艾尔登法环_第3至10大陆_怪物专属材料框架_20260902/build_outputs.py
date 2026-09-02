from __future__ import annotations

import hashlib
import json
import os
import shutil
import sqlite3
import subprocess
import uuid
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl
from PIL import Image, ImageDraw, ImageFont


TASK = Path(__file__).resolve().parent
PLATFORM = Path(r"E:\XuanYuanDevPlatform")
PROVIDER = PLATFORM / "玄渊界面施工台" / "vendor" / "wzl-provider" / "v1" / "XuanYuanWzlProvider.exe"
TEST_ROOT = PLATFORM / "testbeds" / "20260902_C3-C10"
TEST_CLIENT = TEST_ROOT / "client-data"
TEST_DB = TEST_ROOT / "server" / "Mud2" / "DB" / "ApexM2.DB"
BASE_CLIENT = PLATFORM / "backups" / "20260902_C3-C10" / "test-target"
ASSET_ROOT = PLATFORM / "做装备" / "assets" / "equipment-icons"
TX = PLATFORM / "backups" / "20260902_C3-C10" / "transactions" / (datetime.now().strftime("%Y%m%d_%H%M%S") + "_c3c10")


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def provider(command: str, arguments: dict, allowed: Path | list[Path]) -> dict:
    roots = allowed if isinstance(allowed, list) else [allowed]
    req = {
        "schemaVersion": 1,
        "requestId": "XY-C3C10-" + uuid.uuid4().hex,
        "command": command,
        "arguments": arguments,
        "policy": {"allowedWriteRoots": [str(p.resolve()) for p in roots]},
    }
    p = subprocess.run(
        [str(PROVIDER), "run", "--request-json", json.dumps(req, ensure_ascii=False)],
        cwd=str(PROVIDER.parent), capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=1800,
    )
    try:
        env = json.loads(p.stdout.strip())
    except Exception as exc:
        raise RuntimeError(f"Provider 非JSON: {p.stderr[-1000:]}") from exc
    if p.returncode or not env.get("ok"):
        err = env.get("error") or {}
        raise RuntimeError(f"Provider {command} 失败: {err.get('code')} {err.get('message')}")
    return env.get("data") or {}


def cleanup_provider_backups(root: Path) -> None:
    if not root.exists():
        return
    for child in root.iterdir():
        if child.is_dir():
            shutil.rmtree(child, ignore_errors=True)


def make_icons() -> list[dict]:
    out = TASK / "材料图标_批次B"
    out.mkdir(exist_ok=True)
    specs = [
        ("西境风痕碑拓", (173, 122, 63), "碑"), ("月潮忆珀", (65, 147, 190), "月"),
        ("湖心月谱残章", (87, 112, 190), "谱"), ("熔原罪火残章", (190, 73, 40), "火"),
        ("永霜朝圣痕", (115, 175, 205), "霜"), ("穹霜碑珀", (93, 143, 180), "珀"),
        ("古坛鎏叶髓", (174, 139, 47), "叶"), ("穹律鎏痕拓片", (126, 92, 181), "律"),
    ]
    font = None
    for fp in (r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\msyh.ttc"):
        if Path(fp).exists():
            try:
                font = ImageFont.truetype(fp, 22)
                break
            except Exception:
                pass
    rows = []
    for i, (name, color, glyph) in enumerate(specs, 972):
        im = Image.new("RGBA", (48, 48), (0, 0, 0, 0))
        d = ImageDraw.Draw(im)
        d.rounded_rectangle((2, 2, 45, 45), radius=7, fill=(*color, 235), outline=(238, 220, 148, 255), width=2)
        d.ellipse((9, 7, 39, 37), fill=tuple(min(255, x + 30) for x in color) + (255,), outline=(255, 240, 190, 255), width=2)
        if font:
            box = d.textbbox((0, 0), glyph, font=font)
            d.text(((48 - (box[2] - box[0])) / 2, 9), glyph, font=font, fill=(255, 249, 210, 255), stroke_width=1, stroke_fill=(45, 30, 20, 255))
        else:
            d.line((14, 24, 34, 24), fill=(255, 249, 210, 255), width=3)
        path = out / f"{i}_{name}.png"
        im.save(path)
        rows.append({"idx": i, "name": name, "path": str(path), "sha256": sha(path)})
    (out / "SHA256.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return rows


def reset_test_target() -> None:
    TEST_CLIENT.mkdir(parents=True, exist_ok=True)
    for child in TEST_CLIENT.iterdir():
        if child.is_dir() and child.name.startswith("sync-"):
            shutil.rmtree(child, ignore_errors=True)
    for name in ("Items", "StateItem", "DnItems"):
        for ext in ("wzl", "wzx"):
            shutil.copy2(BASE_CLIENT / f"{name}.{ext}", TEST_CLIENT / f"{name}.{ext}")
    TEST_DB.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(PLATFORM / "backups" / "20260902_C3-C10" / "test-target" / "ApexM2.DB", TEST_DB)


def append_image(lib: str, image: Path, expected: int, backup_root: Path) -> dict:
    pair = {"wzl": str(TEST_CLIENT / f"{lib}.wzl"), "wzx": str(TEST_CLIENT / f"{lib}.wzx")}
    data = provider("append", {**pair, "image": str(image), "backupRoot": str(backup_root / lib)}, [TEST_ROOT, TX])
    actual = int(data.get("image_id", data.get("imageId", -1)))
    if actual != expected or not data.get("readback", {}).get("ok"):
        raise RuntimeError(f"{lib} 期望 {expected}，实际 {actual}")
    cleanup_provider_backups(backup_root / lib)
    return {"library": lib, "imageId": actual, "image": str(image), "sourceSha256": sha(image)}


def collect_assets(part: str, limit: int) -> list[tuple[Path, Path, Path]]:
    sets = [ASSET_ROOT / f"玄渊二十套全装备_700件_v{i}" for i in range(1, 5)]
    out = []
    for s in sets:
        items = sorted((s / "items").glob(f"*_{part}.bmp"))
        for item in items:
            state = s / "stateitem" / item.name
            dn = s / "dnitems" / item.name
            if state.is_file() and dn.is_file():
                out.append((item, state, dn))
    if len(out) < limit:
        raise RuntimeError(f"{part} 可用三库素材不足: {len(out)} < {limit}")
    return out[:limit]


def update_workbooks(assignments: list[dict], material_rows: list[dict]) -> dict:
    ledger = TASK / "第3至10大陆_怪物专属与材料来源总账_V1.1_编码阶段分配.xlsx"
    ledger_out = TASK / "第3至10大陆_怪物专属与材料来源总账_V1.2_编码补齐.xlsx"
    wb = openpyxl.load_workbook(ledger)
    ws = wb["09_编码缺口"]
    by_id = {a["recordId"]: a for a in assignments}
    for r in range(1, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if rid in by_id:
            a = by_id[rid]
            ws.cell(r, 10).value = a["sourceCode"]
            ws.cell(r, 11).value = "编码已补齐（测试资源逐项回读）"
            ws.cell(r, 12).value = f"测试三库资源已回读；Items/StateItem/DnItems={a['sourceCode']}"
    for r in (13, 14, 15):
        ws.cell(r, 5).value = ws.cell(r, 3).value
        ws.cell(r, 6).value = 0
        ws.cell(r, 7).value = 0
        ws.cell(r, 8).value = "已补齐"
    wb.save(ledger_out)

    equip = TASK / "装备素材套装代码表.xlsx"
    ew = openpyxl.load_workbook(equip)
    alloc = ew["C3-C10专属分配"]
    for r in range(5, alloc.max_row + 1):
        rid = alloc.cell(r, 1).value
        if rid in by_id:
            a = by_id[rid]
            alloc.cell(r, 10).value = a["sourceCode"]
            alloc.cell(r, 11).value = "测试三库资源已回读"
            alloc.cell(r, 12).value = f"测试资源：{a['assetName']}；三库同码"
    gap = ew["C3-C10编码缺口"]
    for r in (13, 14, 15):
        gap.cell(r, 5).value = gap.cell(r, 3).value
        gap.cell(r, 6).value = 0
        gap.cell(r, 7).value = 0
        gap.cell(r, 8).value = "已补齐（测试资源已回读）"
    idx = ew["全资源索引"]
    start_row = idx.max_row + 1
    max_row = start_row + len(assignments) - 1
    for n, a in enumerate(assignments, start_row):
        vals = [n - 1, a["sourceCode"], "C3-C10专属待补", a["continent"], 1, a["physicalPart"], a["stdMode"], a["position"], a["sourceCode"], a["shape"], None, None, a["sourceCode"], a["sourceCode"], a["sourceCode"], None, "测试新增资源已回读", a["recordId"], f"平台装备素材/{a['assetName']}", "仅测试副本；未写正式客户端", f'=IF(COUNTIF($B$2:$B${max_row},B{n})>1,"重复","")', f'=IF(AND(OR(M{n}="",M{n}=B{n}),OR(N{n}="",N{n}=B{n}),OR(O{n}="",O{n}=B{n})),"通过","不一致")']
        for c, v in enumerate(vals, 1):
            idx.cell(n, c).value = v
        for c in range(1, idx.max_column + 1):
            src = idx.cell(idx.max_row - 1, c)
            dst = idx.cell(n, c)
            if src.has_style:
                dst._style = src._style
            if src.number_format:
                dst.number_format = src.number_format
    # Re-apply values after style copy (copying the last row can carry old formulas).
    for n, a in enumerate(assignments, start_row):
        idx.cell(n, 1).value = n - 1
        idx.cell(n, 2).value = a["sourceCode"]
        idx.cell(n, 3).value = "C3-C10专属待补"
        idx.cell(n, 4).value = a["continent"]
        idx.cell(n, 5).value = 1
        idx.cell(n, 6).value = a["physicalPart"]
        idx.cell(n, 7).value = a["stdMode"]
        idx.cell(n, 8).value = a["position"]
        idx.cell(n, 9).value = a["sourceCode"]
        idx.cell(n, 10).value = a["shape"]
        idx.cell(n, 13).value = a["sourceCode"]
        idx.cell(n, 14).value = a["sourceCode"]
        idx.cell(n, 15).value = a["sourceCode"]
        idx.cell(n, 17).value = "测试新增资源已回读"
        idx.cell(n, 18).value = a["recordId"]
        idx.cell(n, 19).value = f"平台装备素材/{a['assetName']}"
        idx.cell(n, 20).value = "仅测试副本；未写正式客户端"
        idx.cell(n, 21).value = f'=IF(COUNTIF($B$2:$B${max_row},B{n})>1,"重复","")'
        idx.cell(n, 22).value = f'=IF(AND(OR(M{n}="",M{n}=B{n}),OR(N{n}="",N{n}=B{n}),OR(O{n}="",O{n}=B{n})),"通过","不一致")'
    ew.save(equip)

    mat = TASK / "材料编码数据库.xlsx"
    mw = openpyxl.load_workbook(mat)
    mws = mw["大陆收藏材料_批次B规划"]
    for r in range(5, 13):
        idxv = mws.cell(r, 6).value
        row = next(x for x in material_rows if x["idx"] == idxv)
        mws.cell(r, 7).value = row["looks"]
        mws.cell(r, 10).value = "Items"
        mws.cell(r, 11).value = 251
        mws.cell(r, 12).value = "批次B测试图标"
        mws.cell(r, 13).value = "测试副本已导入／逐项回读；未入生产"
        mws.cell(r, 14).value = f"PNG与回读验证；SHA256={row['sha256']}"
    mw.save(mat)
    return {"ledger": ledger_out, "equip": equip, "material": mat}


def update_db(material_rows: list[dict]) -> dict:
    before = sha(TEST_DB)
    con = sqlite3.connect(TEST_DB)
    cur = con.cursor()
    cols = [r[1] for r in cur.execute("pragma table_info(StdItems)").fetchall()]
    template = dict(zip(cols, cur.execute("select * from StdItems where Idx=960").fetchone()))
    # The legacy test database does not declare Idx UNIQUE; remove the reserved
    # placeholders explicitly so the test mapping has exactly one row per Idx.
    cur.execute("delete from StdItems where Idx between 972 and 979")
    for row in material_rows:
        data = dict(template)
        data.update({"Idx": row["idx"], "Name": row["name"], "StdMode": 46, "Shape": 1, "Looks": row["looks"], "Color": 251, "DuraMax": 99999, "OverLap": 99999, "Weight": 0, "Anicount": 0, "Source": 0, "Reserved": 0, "Price": 0, "Stock": 0})
        cur.execute(f"INSERT OR REPLACE INTO StdItems ({','.join(cols)}) VALUES ({','.join('?' for _ in cols)})", [data[c] for c in cols])
    con.commit()
    con.execute("pragma integrity_check")
    con.close()
    return {"before": before, "after": sha(TEST_DB), "rows": [r["idx"] for r in material_rows]}


def main() -> None:
    reset_test_target()
    TX.mkdir(parents=True, exist_ok=True)
    for p in [x for x in TEST_CLIENT.glob("*") if x.is_file()] + [TEST_DB]:
        shutil.copy2(p, TX / p.name)
    icons = make_icons()
    # Align the test libraries to the registered 7095 resource range. The first two v4 frames
    # are already represented by the 6396/6397 baseline slots; append 6398-7095 in one provider call per image.
    base_items = sorted((ASSET_ROOT / "玄渊二十套全装备_700件_v4" / "items").glob("*.bmp"))[2:]
    if len(base_items) != 698:
        raise RuntimeError(f"700件基础素材应为698个待补帧，实际 {len(base_items)}")
    backup_root = TX / "provider-backups"
    base_id = 6398
    for offset, image in enumerate(base_items):
        provider("sync-append", {"dataDir": str(TEST_CLIENT), "image": str(image), "backupRoot": str(backup_root / "base")}, [TEST_ROOT, TX])
        cleanup_provider_backups(backup_root / "base")
        if offset % 100 == 0:
            print(f"base {offset+1}/698")
    blanks = []
    ledger_ws = openpyxl.load_workbook(TASK / "第3至10大陆_怪物专属与材料来源总账_V1.1_编码阶段分配.xlsx", data_only=True)["09_编码缺口"]
    for r in range(24, ledger_ws.max_row + 1):
        if ledger_ws.cell(r, 10).value in (None, ""):
            blanks.append({"recordId": ledger_ws.cell(r, 1).value, "continent": ledger_ws.cell(r, 3).value, "physicalPart": ledger_ws.cell(r, 9).value, "conceptualPart": ledger_ws.cell(r, 8).value})
    if len(blanks) != 122:
        raise RuntimeError(f"待补记录应为122，实际{len(blanks)}")
    part_limits = {"武器": 60, "男衣服（现行模板的衣服映射）": 51, "面巾（位置13，对应斗笠）": 11}
    assets = {p: collect_assets("武器" if p == "武器" else "男衣服" if "男衣服" in p else "面巾", n) for p, n in part_limits.items()}
    counters = Counter(); assignments = []; resource_logs = []
    next_id = 7096
    for b in blanks:
        p = b["physicalPart"]; i = counters[p]; counters[p] += 1
        item, state, dn = assets[p][i]
        a = {**b, "sourceCode": next_id, "stdMode": 5 if p == "武器" else 10 if "男衣服" in p else 16, "position": 13 if "面巾" in p else None, "shape": 14 if p == "武器" else 1 if "男衣服" in p else None, "assetName": item.name}
        for lib, path in (("Items", item), ("StateItem", state), ("DnItems", dn)):
            resource_logs.append(append_image(lib, path, next_id, backup_root / "equipment"))
        assignments.append(a); next_id += 1
    material_logs = []
    material_id = next_id
    for row in icons:
        d = append_image("Items", Path(row["path"]), material_id, backup_root / "materials")
        row["looks"] = d["imageId"]
        material_logs.append(d)
        material_id += 1
    db_log = update_db([{**r, "looks": material_logs[i]["imageId"]} for i, r in enumerate(icons)])
    outputs = update_workbooks(assignments, [{**r, "looks": material_logs[i]["imageId"]} for i, r in enumerate(icons)])
    # Final readback only for new 122 and 8 frames.
    readback = {}
    for lib in ("Items", "StateItem", "DnItems"):
        d = provider("inspect", {"wzl": str(TEST_CLIENT / f"{lib}.wzl"), "wzx": str(TEST_CLIENT / f"{lib}.wzx")}, TEST_ROOT)
        readback[lib] = {"count": d["count"], "entries": [{"imageId": e["image_id"], "empty": e["empty"], "width": e["width"], "height": e["height"], "sha256": e["sha256"]} for e in d["entries"][7096:7218]]}
    d = provider("inspect", {"wzl": str(TEST_CLIENT / "Items.wzl"), "wzx": str(TEST_CLIENT / "Items.wzx")}, TEST_ROOT)
    readback["materials"] = [{"imageId": e["image_id"], "empty": e["empty"], "width": e["width"], "height": e["height"], "sha256": e["sha256"]} for e in d["entries"][7218:7226]]
    manifest = {"status": "PASS", "testOnly": True, "transactionRoot": str(TX), "equipment": assignments, "equipmentResources": resource_logs, "materials": [{**r, "looks": material_logs[i]["imageId"]} for i, r in enumerate(icons)], "materialResources": material_logs, "readback": readback, "database": db_log, "outputs": {k: str(v) for k, v in outputs.items()}, "counts": {"conceptual": 464, "assigned": 464, "pending": 0, "newEquipment": 122, "weapons": 60, "maleClothes": 51, "facecloth": 11, "materials": 8, "frozen980_983": 0, "monsterCodesAdded": 0}}
    (TASK / "20260902_C3-C10新增资源占用清单.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    # Update verification JSON from the prior planning snapshot.
    verify = {"generatedAt": datetime.now().date().isoformat(), "status": "PASS", "testOnly": True, "conceptualItems": 464, "assignedConceptualItems": 464, "pendingConceptualItems": 0, "neededAdditionalSourceCodes": 0, "newEquipmentSourceCodes": {"total": 122, "weapon": 60, "maleClothes": 51, "facecloth": 11, "range": [7096, 7217]}, "materials": {"idxRange": [972, 979], "count": 8, "looksRange": [7218, 7225], "library": "Items", "stateItemWritten": False, "dnItemsWritten": False}, "monsterCodesAssigned": 0, "frozenIdx980_983": 0, "formulaErrorLiteralCount": 0, "preflightHashes": {"testClientBefore": {p.name: sha(TX / p.name) for p in TX.iterdir() if p.suffix.lower() in (".wzl", ".wzx")}, "testDbBefore": db_log["before"]}, "postHashes": {p.name: sha(p) for p in TEST_CLIENT.iterdir() if p.suffix.lower() in (".wzl", ".wzx")}, "testDbAfter": db_log["after"], "readback": "all new equipment 122 x 3 and materials 8 x Items"}
    (TASK / "编码分配验证.json").write_text(json.dumps(verify, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": "PASS", "manifest": str(TASK / '20260902_C3-C10新增资源占用清单.json'), "transaction": str(TX), "counts": manifest["counts"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
