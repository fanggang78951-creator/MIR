from __future__ import annotations

import sqlite3
import shutil
import struct
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from xydp.initial_camp import (
    EXPECTED_FILES,
    InitialCampError,
    InitialCampService,
    _donate_anchor_contents,
    _enable_setup_flags,
    _equipment_chain_script,
    _cost_lines,
    _newbie_gift_script,
    _merge_named_description_lines,
    _rebirth_anchor_contents,
    _rebirth_power_steps,
    _rebirth_script,
    _single_tier_route_lines,
    _sync_sponsor_panel_values,
    _sponsor_anchor_contents,
    _sponsor_script,
    _title_chain_script,
    _upsert_title_definitions,
    read_workbook,
)
from xydp.encoding import TextDocument
from xydp.installer import InstallPlan
from xydp.textpatch import install_managed_anchor_hook


class InitialCampTests(unittest.TestCase):
    platform = Path(r"E:\XuanYuanDevPlatform")
    source_materials = platform / "所需材料表格汇总"
    materials = source_materials

    def test_setup_flag_update_preserves_target_crlf_without_normalizing_legacy_lf(self):
        document = TextDocument(
            "ACAttackSpeed=30\n"
            "SendItemDescList=0\r\n"
            "SendTzItemDescList=0\r\n"
            "Other=1\r\n",
            "gb18030",
            "\r\n",
        )

        updated = _enable_setup_flags(
            document, ("SendItemDescList", "SendTzItemDescList")
        ).decode("gb18030")

        self.assertEqual(
            updated,
            "ACAttackSpeed=30\n"
            "SendItemDescList=1\r\n"
            "SendTzItemDescList=1\r\n"
            "Other=1\r\n",
        )

    def test_named_description_update_keeps_existing_line_position(self):
        document = TextDocument(
            "第一项=说明一\r\n"
            "沙城捐献=旧说明\r\n"
            "最后一项=说明二\r\n",
            "gb18030",
            "\r\n",
        )

        updated = _merge_named_description_lines(
            document, ["沙城捐献=新说明"]
        ).decode("gb18030")

        self.assertEqual(
            updated,
            "第一项=说明一\r\n"
            "沙城捐献=新说明\r\n"
            "最后一项=说明二\r\n",
        )

    @classmethod
    def setUpClass(cls) -> None:
        # 用户会持续编辑正式母表；回归测试必须使用稳定副本，不能被业务中的临时值污染。
        cls._materials_temp = tempfile.TemporaryDirectory(prefix="xydp-camp-materials-")
        cls.materials = Path(cls._materials_temp.name)
        for name in EXPECTED_FILES:
            shutil.copy2(cls.source_materials / name, cls.materials / name)
        donate_path = cls.materials / "02_捐献.xlsx"
        workbook = load_workbook(donate_path)
        sheet = workbook.active
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if str(cell.value).strip() == "44":
                    cell.value = None
        workbook.save(donate_path)
        workbook.close()

    @classmethod
    def tearDownClass(cls) -> None:
        cls._materials_temp.cleanup()
        cls.materials = cls.source_materials

    def make_target(self, root: Path) -> Path:
        envir = root / "Mir200" / "Envir"
        (envir / "Market_Def" / "玄渊运营").mkdir(parents=True)
        (envir / "MapQuest_Def").mkdir(parents=True)
        (root / "Mir200" / "M2Server.exe").write_bytes(b"M2")
        (root / "Mir200" / "!setup.txt").write_text(
            "[Setup]\r\nSendItemDescList=0\r\nSendTzItemDescList=0\r\nRevivalTime=60000\r\n", encoding="gb18030"
        )
        (envir / "MapInfo.txt").write_text("[xycamp|vx79 初始营地]\r\n", encoding="gb18030")
        (envir / "Market_Def" / "QFunction-0.txt").write_bytes((
            "; 中文编码夹具\r\n"
            "; XY_EXECUTION_LAB_CHANCE_ANCHOR\r\n"
            "; XY_EXECUTION_LAB_TOUGHNESS_ANCHOR\r\n"
            "; XY_EQUIP_MAKER_DROP_ANCHOR\r\n"
            "; XY_EQUIP_MAKER_RUNTIME_DROP_ANCHOR\r\n"
            "; XY_EQUIP_MAKER_RUNTIME_DROP_MAX_ANCHOR\r\n"
            "; XY_EQUIP_MAKER_RUNTIME_DAMAGE_COEFFICIENT_ANCHOR\r\n"
            "; XY_EQUIP_MAKER_POWER_ANCHOR\r\n"
            "; XY_EQUIP_MAKER_RUNTIME_POWER_ANCHOR\r\n"
            "; XY_EQUIP_MAKER_ATTACK_ANCHOR\r\n"
        ).encode("gb18030"))
        (envir / "MapQuest_Def" / "QManage.txt").write_text("[@Login]\r\n#IF\r\n#ACT\r\n", encoding="gb18030")
        (envir / "ItemDescList.txt").write_text("旧物品=旧说明\r\n", encoding="gb18030")
        (envir / "MerChant.txt").write_text(
            "玄渊运营/狂暴之力\t0\t322\t272\t狂暴之力\t0\t15\t0\r\n"
            "玄渊运营/沙城捐献\t0\t324\t272\t沙城捐献\t0\t9\t0\r\n"
            "旧NPC/NPC_A\txycamp\t20\t20\tNPC_A\t0\t1\t0\r\n",
            encoding="gb18030",
        )
        width, height, cell_size = 55, 57, 14
        map_data = bytearray(52 + width * height * cell_size)
        struct.pack_into("<HH", map_data, 0, width, height)
        (root / "Mir200" / "Map").mkdir(parents=True)
        (root / "Mir200" / "Map" / "vx79.map").write_bytes(map_data)
        db = root / "Mud2" / "DB" / "ApexM2.DB"
        db.parent.mkdir(parents=True)
        connection = sqlite3.connect(db)
        try:
            connection.execute(
                "CREATE TABLE StdItems (Idx INTEGER, Name TEXT, StdMode INTEGER, Shape INTEGER, Weight INTEGER, "
                "Anicount INTEGER, Source INTEGER, Reserved INTEGER, Looks INTEGER, DuraMax INTEGER, Ac INTEGER, Ac2 INTEGER, "
                "Mac INTEGER, Mac2 INTEGER, Dc INTEGER, Dc2 INTEGER, Mc INTEGER, Mc2 INTEGER, Sc INTEGER, Sc2 INTEGER, "
                "Need INTEGER, NeedLevel INTEGER, Price INTEGER, Stock INTEGER, Color INTEGER, OverLap INTEGER, HP INTEGER, "
                "MP INTEGER, Light INTEGER, Horse INTEGER)"
            )
            index = 1
            for prefix, std_mode, shape_base in (("圣律之剑LV", 5, 100), ("黄金圣物LV", 30, 120)):
                for level in range(1, 11):
                    connection.execute(
                        "INSERT INTO StdItems (Idx, Name, StdMode, Shape) VALUES (?, ?, ?, ?)",
                        (index, f"{prefix}{level}", std_mode, shape_base + level),
                    )
                    index += 1
            connection.commit()
        finally:
            connection.close()
        return root

    def make_client(self, root: Path) -> Path:
        (root / "data").mkdir(parents=True)
        (root / "data" / "fenghao.dat").write_text("旧称号=旧说明\r\n", encoding="gb18030")
        return root

    def empty_dependency_plan(self, target: Path) -> InstallPlan:
        return InstallPlan(str(target), None, ["xy.optional.rage.core", "xy.optional.donate.core", "xy.ui.attr-overview"],
                           {"xy.optional.rage.core": "test", "xy.optional.donate.core": "test", "xy.ui.attr-overview": "test"}, {}, [],
                           candidate_packages=["xy.optional.rage.core", "xy.optional.donate.core", "xy.ui.attr-overview"])

    def database_with_items(self, *names: str) -> bytes:
        with tempfile.TemporaryDirectory(prefix="xydp-items-") as td:
            db = self.make_target(Path(td) / "server") / "Mud2" / "DB" / "ApexM2.DB"
            connection = sqlite3.connect(db)
            try:
                connection.execute("DELETE FROM StdItems")
                for index, name in enumerate(names, 1):
                    connection.execute(
                        "INSERT INTO StdItems (Idx, Name, StdMode, Shape) VALUES (?, ?, 30, ?)",
                        (index, name, index),
                    )
                connection.commit()
            finally:
                connection.close()
            return db.read_bytes()

    def test_all_seven_workbooks_are_readable(self):
        counts = []
        for name in EXPECTED_FILES:
            book = read_workbook(self.materials / name)
            self.assertTrue(book.sha256)
            self.assertTrue(book.rows)
            counts.append(len(book.rows))
        self.assertEqual(counts, [1, 1, 4, 10, 10, 9, 1])

    def test_user_named_equipment_and_newbie_gift_workbooks_are_recognized(self):
        sword = read_workbook(self.materials / "05_圣律之剑.xlsx").rows[0]
        relic = read_workbook(self.materials / "06_黄金圣物.xlsx").rows[0]
        gift = read_workbook(self.materials / "07_新手礼包.xlsx").rows[0]
        self.assertEqual(sword["NPC名称"], "圣律之剑")
        self.assertEqual(relic["NPC名称"], "黄金圣物")
        self.assertEqual(gift["NPC名称"], "新手礼包")
        self.assertEqual(gift["奖励装备A"], "圣律之剑LV1")
        self.assertEqual(gift["奖励装备B"], "黄金圣物LV1")
        self.assertEqual(str(gift["礼包码"]), "888888")

    def test_horse_exchange_is_xlsx_driven_and_does_not_write_attributes(self):
        rows = (
            {"状态": "可安装", "档位/等级": "1升2", "当前装备": "青铜马牌", "下一装备": "白银马牌", "材料A": "灵符", "材料A数量": 10},
            {"状态": "可安装", "档位/等级": "2升3", "当前装备": "白银马牌", "下一装备": "黄金马牌", "货币A": "元宝", "货币A数量": 20},
        )
        script = _equipment_chain_script(
            "马牌升级", "坐骑进阶大师", rows, "XY_HORSE",
            self.database_with_items("青铜马牌", "白银马牌", "黄金马牌"),
        )
        self.assertIn("坐骑进阶大师进阶", script)
        self.assertIn("CHECKITEM 黄金马牌 1\n#ACT\nGOTO @XY_HORSE_FULL", script)
        self.assertIn("CHECKITEM 青铜马牌 1\n#ACT\nGOTO @XY_HORSE_VIEW_1", script)
        self.assertNotIn("1升2：青铜马牌→白银马牌", script)
        self.assertIn("当前装备:/FCOLOR=161>{青铜马牌/SCOLOR=249}", script)
        self.assertIn("下一装备:/FCOLOR=161>{白银马牌/SCOLOR=253}", script)
        self.assertIn("<ItemShow:2:0:260:-90:1>", script)
        self.assertIn("CHECKITEM 青铜马牌 1", script)
        self.assertIn("TAKE 青铜马牌 1", script)
        self.assertIn("CHECKGAMEGIRD > 9", script)
        self.assertIn("GAMEGIRD - 10", script)
        self.assertNotIn("CHECKITEM 灵符", script)
        self.assertNotIn("TAKE 灵符", script)
        self.assertIn("GIVE 白银马牌 1", script)
        self.assertIn("CHECKGAMEGOLD > 19", script)
        self.assertIn("GAMEGOLD - 20", script)
        self.assertNotIn("CHECKITEM 元宝", script)
        self.assertNotIn("TAKE 元宝", script)
        self.assertIn("GOTO @Main", script)
        self.assertIn("当前装备已经达到最高档", script)
        self.assertNotIn("CHANGEITEM", script.upper())
        self.assertNotIn("SETCUSTOMITEM", script.upper())

    def test_equipment_chain_uses_native_gold_and_single_next_stage(self):
        rows = (
            {"状态": "可安装", "档位/等级": "1级", "当前装备": "装备LV1", "下一装备": "装备LV2", "货币A": "金币", "货币A数量": 50000},
            {"状态": "可安装", "档位/等级": "2级", "当前装备": "装备LV2", "下一装备": "装备LV3", "货币A": "金币", "货币A数量": 100000},
        )
        script = _equipment_chain_script(
            "装备进阶", "装备大师", rows, "XY_EQUIP",
            self.database_with_items("装备LV1", "装备LV2", "装备LV3"),
        )
        self.assertIn("CHECKGOLD 50000", script)
        self.assertIn("GOLDCOUNT - 50000", script)
        self.assertNotIn("CHECKITEM 金币", script)
        self.assertNotIn("TAKE 金币", script)
        self.assertEqual(script.count("<ItemShow:"), 2)
        self.assertEqual(script.count(":0:260:-90:1>"), 2)
        self.assertNotIn("1级：装备LV1→装备LV2", script)

    def test_single_tier_router_supports_explicit_progress_variable(self):
        lines = _single_tier_route_lines(
            "XY_TEST", ("0", "1", "2"), "variable", state_variable="N$XY_TEST_LEVEL"
        )
        text = "\n".join(lines)
        self.assertIn("EQUAL N$XY_TEST_LEVEL 2\n#ACT\nGOTO @XY_TEST_FULL", text)
        self.assertIn("EQUAL N$XY_TEST_LEVEL 1\n#ACT\nGOTO @XY_TEST_VIEW_2", text)
        self.assertIn("EQUAL N$XY_TEST_LEVEL 0\n#ACT\nGOTO @XY_TEST_VIEW_1", text)
        self.assertIn("GOTO @XY_TEST_INVALID", text)

    def test_newbie_gift_grants_two_bound_items_once(self):
        rows = ({
            "状态": "可安装", "奖励装备A": "圣律之剑LV1",
            "奖励装备B": "黄金圣物LV1", "礼包码": "888888",
        },)
        script = _newbie_gift_script(
            "新手礼包", rows, self.database_with_items("圣律之剑LV1", "黄金圣物LV1")
        )
        self.assertIn("<$NPCINPUT(1)>", script)
        self.assertIn("EQUAL <$NPCINPUT(1)> 888888", script)
        self.assertIn("<TEXT:提交礼包码:15:110:1/@XY_GIFT_CLAIM>", script)
        self.assertIn("CHECKBAGSIZE 2", script)
        self.assertIn("GiveStateItem 圣律之剑LV1 1 0 0 0 0 0 0 1", script)
        self.assertIn("GiveStateItem 黄金圣物LV1 1 0 0 0 0 0 0 1", script)
        self.assertEqual(script.count("ADDNAMELIST"), 1)

    def test_newbie_gift_blocks_unknown_or_duplicate_reward_names(self):
        with self.assertRaisesRegex(InitialCampError, "StdItems中唯一存在"):
            _newbie_gift_script(
                "新手礼包",
                ({"状态": "可安装", "奖励装备A": "圣律之剑LV1", "奖励装备B": "不存在", "礼包码": "888888"},),
                self.database_with_items("圣律之剑LV1"),
            )
        with self.assertRaisesRegex(InitialCampError, "不能相同"):
            _newbie_gift_script(
                "新手礼包",
                ({"状态": "可安装", "奖励装备A": "圣律之剑LV1", "奖励装备B": "圣律之剑LV1", "礼包码": "888888"},),
                self.database_with_items("圣律之剑LV1"),
            )

    def test_equipment_chain_blocks_unknown_or_non_unique_names(self):
        rows = ({"状态": "可安装", "当前装备": "不存在马牌", "下一装备": "白银马牌", "材料A": "灵符", "材料A数量": 1},)
        with self.assertRaisesRegex(InitialCampError, "StdItems中唯一存在"):
            _equipment_chain_script(
                "马牌升级", "马牌NPC", rows, "XY_HORSE", self.database_with_items("白银马牌")
            )

    def test_equipment_chain_blocks_invalid_or_shared_item_index(self):
        database = self.database_with_items("装备LV1", "装备LV2")
        with tempfile.TemporaryDirectory(prefix="xydp-invalid-idx-") as td:
            path = Path(td) / "items.db"
            path.write_bytes(database)
            connection = sqlite3.connect(path)
            try:
                connection.execute("UPDATE StdItems SET Idx = 0 WHERE Name = '装备LV2'")
                connection.commit()
            finally:
                connection.close()
            invalid = path.read_bytes()
        rows = ({"状态": "可安装", "当前装备": "装备LV1", "下一装备": "装备LV2", "货币A": "金币", "货币A数量": 1},)
        with self.assertRaisesRegex(InitialCampError, "Idx"):
            _equipment_chain_script("装备进阶", "装备NPC", rows, "XY_EQUIP", invalid)

    def test_preflight_moves_unique_npcs_and_generates_configured_scripts(self):
        with tempfile.TemporaryDirectory(prefix="xydp-camp-") as td:
            server = self.make_target(Path(td) / "server")
            client = self.make_client(Path(td) / "client")
            service = InitialCampService(self.platform)
            with patch.object(service.installer, "preflight", return_value=self.empty_dependency_plan(server)) as dependency_preflight:
                plan = service.preflight(server, self.materials, client)
            self.assertEqual(plan.blockers, [])
            self.assertEqual(len(plan.npcs), 7)
            self.assertEqual(len({(item.x, item.y) for item in plan.npcs}), 7)
            merchant = next(item for item in plan.changes if item.relative_path.endswith("MerChant.txt")).after.decode("gb18030")
            self.assertNotIn("玄渊运营/狂暴之力\t0\t", merchant)
            self.assertNotIn("玄渊运营/沙城捐献\t0\t", merchant)
            self.assertEqual(merchant.count("玄渊运营/狂暴之力\txycamp\t"), 1)
            self.assertEqual(merchant.count("玄渊运营/沙城捐献\txycamp\t"), 1)
            donate = next(item for item in plan.changes if item.relative_path.endswith("沙城捐献-xycamp.txt")).after.decode("gb18030")
            self.assertIn("@XY_DONATE_TRIGGER", donate)
            self.assertNotIn("GIVEFENGHAO 沙城捐献", donate)
            self.assertNotIn("SetNewFengHaoValue 沙城捐献", donate)
            donate_config = next(
                item for item in plan.changes if item.relative_path.endswith("沙城捐献配置.txt")
            ).after.decode("gb18030")
            self.assertIn("消耗类型=原生灵符", donate_config)
            self.assertIn("消耗名称=灵符", donate_config)
            self.assertNotIn("累计增加=", donate_config)
            self.assertNotIn("封榜阈值=", donate_config)
            qfunction = next(
                item for item in plan.changes if item.relative_path.endswith("QFunction-0.txt")
            ).after.decode("gb18030")
            self.assertIn("CHECKFENGHAO 沙城捐献", qfunction)
            self.assertIn("INC N$XY_最终爆率 100", qfunction)
            self.assertIn("INC N$XY_最大爆率 5", qfunction)
            self.assertIn("INC N$XY_RT_Drop 100", qfunction)
            self.assertIn("INC N$XY_RT_DropMax 5", qfunction)
            self.assertNotIn("CHECKFENGHAO 沙城捐献\n#ACT\nINC N$XY_PVE", qfunction.replace("\r\n", "\n"))
            rebirth = next(item for item in plan.changes if item.relative_path.endswith("转生-xycamp.txt")).after.decode("gb18030")
            self.assertIn("GIVEFENGHAO 转生1重 1", rebirth)
            self.assertIn("CHECKRENEWLEVEL = 0", rebirth)
            self.assertIn("RENEWLEVEL 1", rebirth)
            self.assertIn("TAKE 小卢恩 5", rebirth)
            self.assertIn("CHECKRENEWLEVEL = 1\n#ACT\nMOV N$XY_REBIRTH_PowerBonus 5", qfunction.replace("\r\n", "\n"))
            self.assertIn("CHECKRENEWLEVEL > 9\n#ACT\nMOV N$XY_REBIRTH_PowerBonus 50", qfunction.replace("\r\n", "\n"))
            self.assertIn("INC N$倍攻 <$STR(N$XY_REBIRTH_PowerBonus)>", qfunction)
            self.assertIn("INC N$XY_RT_Power <$STR(N$XY_REBIRTH_PowerBonus)>", qfunction)
            horse = next(item for item in plan.changes if item.relative_path.endswith("马牌升级-xycamp.txt")).after.decode("gb18030")
            relic = next(item for item in plan.changes if item.relative_path.endswith("军鼓升级-xycamp.txt")).after.decode("gb18030")
            self.assertIn("GIVE 圣律之剑LV2 1", horse)
            self.assertIn("GIVE 黄金圣物LV2 1", relic)
            gift = next(item for item in plan.changes if item.relative_path.endswith("新手礼包-xycamp.txt")).after.decode("gb18030")
            self.assertIn("EQUAL <$NPCINPUT(1)> 888888", gift)
            self.assertIn("GiveStateItem 圣律之剑LV1", gift)
            self.assertIn("GiveStateItem 黄金圣物LV1", gift)
            sponsor = next(item for item in plan.changes if item.relative_path.endswith("赞助称号-xycamp.txt")).after.decode("gb18030")
            self.assertIn("<查看当前赞助属性/@XY_SPONSOR_STATUS>", sponsor)
            self.assertIn("#CALL [\\玄渊功能\\非常驻\\属性总览\\玄渊三属性按钮.txt] @XY_UI_STATUS_REFRESH", sponsor)
            self.assertIn("当前拥有：赞助4档", sponsor)
            itemdesc = next(item for item in plan.changes if item.relative_path.endswith("ItemDescList.txt")).after.decode("gb18030")
            self.assertIn("沙城捐献=\\242/", itemdesc)
            self.assertIn("基础爆率+100%", itemdesc)
            self.assertIn("最大爆率+5%", itemdesc)
            self.assertIn("赞助4档=\\242/", itemdesc)
            setup = next(item for item in plan.changes if item.relative_path.endswith("!setup.txt")).after.decode("gb18030")
            self.assertIn("SendItemDescList=1", setup)
            self.assertIn("SendTzItemDescList=1", setup)
            fenghao = next(item for item in plan.changes if item.scope == "client" and item.relative_path.endswith("fenghao.dat"))
            self.assertIn("赞助4档=处决概率+10%", fenghao.after.decode("gb18030"))
            self.assertIn("沙城捐献=攻击+99-99", fenghao.after.decode("gb18030"))
            dependency_params = dependency_preflight.call_args.args[2]
            self.assertEqual(dependency_params["donate_title"], "沙城捐献")
            self.assertEqual(dependency_params["donate_critical_rate"], 10)
            panel_a_coefficient = next(
                item for item in dependency_params["panel_a"]["items"] if item["label"] == "伤害系数"
            )
            panel_c_toughness = next(
                item for item in dependency_params["panel_c"]["items"] if item["label"] == "韧性"
            )
            self.assertEqual([item["amount"] for item in panel_a_coefficient["title_additions"]], [5, 6, 10, 15])
            self.assertEqual([item["amount"] for item in panel_c_toughness["title_additions"]], [20, 30, 40, 50])

            database = next(
                item for item in plan.changes if item.relative_path.endswith("ApexM2.DB")
            ).after
            with tempfile.TemporaryDirectory(prefix="xydp-donate-db-") as db_dir:
                db_path = Path(db_dir) / "items.db"
                db_path.write_bytes(database)
                connection = sqlite3.connect(db_path)
                try:
                    row = connection.execute(
                        "SELECT Dc, Dc2, Mc, Mc2, Sc, Sc2 FROM StdItems WHERE Name='沙城捐献'"
                    ).fetchone()
                finally:
                    connection.close()
            self.assertEqual(row, (99, 99, 99, 99, 99, 99))

    def test_preflight_replaces_legacy_donate_title_hooks_with_current_xlsx_values(self):
        with tempfile.TemporaryDirectory(prefix="xydp-camp-donate-upgrade-") as td:
            server = self.make_target(Path(td) / "server")
            client = self.make_client(Path(td) / "client")
            qfunction_path = server / "Mir200/Envir/Market_Def/QFunction-0.txt"
            text = qfunction_path.read_text(encoding="gb18030")
            legacy = {
                "XY_EQUIP_MAKER_POWER_ANCHOR": "#IF\nCHECKFENGHAO 沙城捐献\n#ACT\nINC N$XY_PVE 50",
                "XY_EQUIP_MAKER_ATTACK_ANCHOR": "#IF\nCHECKFENGHAO 沙城捐献\n#ACT\nINC N$XY_PVE_Extra 50",
                "XY_EQUIP_MAKER_DROP_ANCHOR": "#IF\nCHECKFENGHAO 沙城捐献\n#ACT\nINC N$XY_最大爆率 20",
                "XY_EQUIP_MAKER_RUNTIME_DROP_MAX_ANCHOR": "#IF\nCHECKFENGHAO 沙城捐献\n#ACT\nINC N$XY_RT_DropMax 20",
            }
            for anchor, content in legacy.items():
                text = install_managed_anchor_hook(
                    text, "xy.ops.donate-title", anchor, content, "\r\n"
                ).text
            qfunction_path.write_text(text, encoding="gb18030", newline="")

            service = InitialCampService(self.platform)
            with patch.object(service.installer, "preflight", return_value=self.empty_dependency_plan(server)):
                plan = service.preflight(server, self.materials, client)
            qfunction = next(
                item for item in plan.changes if item.relative_path.endswith("QFunction-0.txt")
            ).after.decode("gb18030")
            self.assertNotIn("xy.ops.donate-title", qfunction)
            self.assertNotIn("INC N$XY_PVE 50", qfunction)
            self.assertNotIn("INC N$XY_PVE_Extra 50", qfunction)
            self.assertNotIn("INC N$XY_最大爆率 20", qfunction)
            self.assertIn("INC N$XY_最终爆率 100", qfunction)
            self.assertIn("INC N$XY_最大爆率 5", qfunction)

    def test_auto_position_skips_an_occupied_first_preference(self):
        with tempfile.TemporaryDirectory(prefix="xydp-camp-occupied-") as td:
            server = self.make_target(Path(td) / "server")
            client = self.make_client(Path(td) / "client")
            merchant = server / "Mir200" / "Envir" / "MerChant.txt"
            merchant.write_text(
                merchant.read_text(encoding="gb18030")
                + "其他NPC/占位\txycamp\t20\t25\t其他NPC\t0\t1\t0\r\n",
                encoding="gb18030",
            )
            service = InitialCampService(self.platform)
            with patch.object(service.installer, "preflight", return_value=self.empty_dependency_plan(server)):
                plan = service.preflight(server, self.materials, client)
            self.assertEqual(plan.blockers, [])
            self.assertEqual(len(plan.npcs), 7)
            self.assertNotIn((20, 25), {(item.x, item.y) for item in plan.npcs})

    def test_rebirth_uses_native_progress_and_title_only_as_display_mirror(self):
        rows = tuple(
            {
                "状态": "可安装", "档位/等级": f"{index}转", "称号名称": f"转生{index}重",
                "称号编号": 29 + index, "属性与数值": "神力倍攻+0.05", "神力增加": 0.05,
                "材料A": "灵符", "材料A数量": index,
            }
            for index in range(1, 11)
        )
        script = _rebirth_script(rows)
        self.assertIn("CHECKRENEWLEVEL = 0", script)
        self.assertIn("CHECKRENEWLEVEL = 1\n#ACT\nGOTO @XY_REBIRTH_VIEW_2", script)
        self.assertIn("下一转生:/FCOLOR=161>{2转（转生2重）/SCOLOR=253}", script)
        second = script[script.index("[@XY_REBIRTH_APPLY_2]"):script.index("[@XY_REBIRTH_VIEW_3]")]
        self.assertLess(second.index("GIVEFENGHAO 转生2重 1"), second.index("RENEWLEVEL 1"))
        self.assertLess(second.index("RENEWLEVEL 1"), second.index("GAMEGIRD - 2"))
        self.assertLess(second.index("GAMEGIRD - 2"), second.index("RECYCFENGHAO 转生1重"))
        self.assertIn("[@XY_REBIRTH_MIGRATE_VIEW_10]", script)
        self.assertIn("RENEWLEVEL 10", script)
        self.assertIn("本次未扣除物品和货币", script)
        self.assertIn("转生进阶：/FCOLOR=158> <突破桎梏，重塑根基", script)
        self.assertIn("当前转生:/FCOLOR=161>{转生1重/SCOLOR=249}", script)
        self.assertIn("转生传承：/FCOLOR=158> <找回您曾经达到的转生境界", script)
        self.assertIn("确认传承/@XY_REBIRTH_MIGRATE_10", script)
        self.assertIn("恭喜您晋升至转生10重", script)
        self.assertNotIn("引擎", script)
        self.assertNotIn("同步", script)
        self.assertNotIn("数据", script)
        self.assertNotIn("TAKE 灵符", script)
        self.assertNotIn("DELAYGOTO", script)
        self.assertEqual(_rebirth_power_steps(rows), list(range(5, 51, 5)))
        anchors = _rebirth_anchor_contents(rows)
        self.assertIn("INC N$倍攻 <$STR(N$XY_REBIRTH_PowerBonus)>", anchors["XY_EQUIP_MAKER_POWER_ANCHOR"])
        self.assertIn("INC N$XY_RT_Power <$STR(N$XY_REBIRTH_PowerBonus)>", anchors["XY_EQUIP_MAKER_RUNTIME_POWER_ANCHOR"])
        self.assertNotIn("CHECKFENGHAO", "\n".join(anchors.values()))

    def test_active_title_level_without_cost_is_blocked(self):
        rows = ({"状态": "可安装", "档位/等级": "1转", "称号名称": "转生1重", "称号编号": 30},)
        with self.assertRaisesRegex(InitialCampError, "没有填写材料或货币消耗"):
            _title_chain_script("转生", rows, "XY_REBIRTH")

    def test_active_title_level_without_attribute_description_is_blocked(self):
        rows = ({"状态": "可安装", "档位/等级": "1转", "称号名称": "转生1重", "称号编号": 30,
                 "材料A": "小卢恩", "材料A数量": 1},)
        with self.assertRaisesRegex(InitialCampError, "属性与数值说明"):
            _title_chain_script("转生", rows, "XY_REBIRTH")

    def test_title_definitions_use_xlsx_shape_and_block_conflicts(self):
        rows = (
            {"状态": "可安装", "称号名称": "转生1重", "称号编号": 30},
            {"状态": "可安装", "称号名称": "转生2重", "称号编号": 31},
        )
        database = _upsert_title_definitions(
            self.database_with_items("占位装备"), "转生", rows, 248
        )
        with tempfile.TemporaryDirectory(prefix="xydp-title-read-") as td:
            path = Path(td) / "db.sqlite"
            path.write_bytes(database)
            connection = sqlite3.connect(path)
            try:
                records = connection.execute(
                    "SELECT Name, StdMode, Shape, Looks, Anicount FROM StdItems "
                    "WHERE Name LIKE '转生%重' ORDER BY Shape"
                ).fetchall()
            finally:
                connection.close()
        self.assertEqual(records, [("转生1重", 70, 30, 150, 1), ("转生2重", 70, 31, 155, 1)])

        with tempfile.TemporaryDirectory(prefix="xydp-title-conflict-") as td:
            db = self.make_target(Path(td) / "server") / "Mud2" / "DB" / "ApexM2.DB"
            connection = sqlite3.connect(db)
            try:
                connection.execute(
                    "INSERT INTO StdItems (Idx, Name, StdMode, Shape) VALUES (99, '其他称号', 70, 30)"
                )
                connection.commit()
            finally:
                connection.close()
            with self.assertRaisesRegex(ValueError, "字段占用冲突|称号"):
                _upsert_title_definitions(db.read_bytes(), "转生", rows, 248)

    def test_sponsor_hooks_are_xlsx_driven(self):
        rows = ({
            "状态": "可安装", "称号名称": "赞助1档", "称号编号": 11,
            "属性与数值": "处决+5%；韧性+20；基础爆率+100%；最大爆率+3%；伤害系数+5%",
            "处决": 5, "韧性": 20, "基础爆率": 100, "最大爆率": 3, "伤害系数": 5,
        },)
        hooks = _sponsor_anchor_contents(rows)
        self.assertIn("INC N$XY_EXEC_ChanceBP 500", hooks["XY_EXECUTION_LAB_CHANCE_ANCHOR"])
        self.assertIn("INC N$XY_EXEC_Toughness 20", hooks["XY_EXECUTION_LAB_TOUGHNESS_ANCHOR"])
        self.assertIn("INC N$XY_最终爆率 100", hooks["XY_EQUIP_MAKER_DROP_ANCHOR"])
        self.assertIn("INC N$XY_最大爆率 3", hooks["XY_EQUIP_MAKER_DROP_ANCHOR"])
        self.assertIn("INC N$XY_RT_DamageCoeff 5", hooks["XY_EQUIP_MAKER_RUNTIME_DAMAGE_COEFFICIENT_ANCHOR"])

    def test_real_dependency_packages_compile_into_one_plan(self):
        with tempfile.TemporaryDirectory(prefix="xydp-camp-real-") as td:
            server = self.make_target(Path(td) / "server")
            client = self.make_client(Path(td) / "client")
            (server / "Mir200" / "Envir" / "Market_Def" / "QFunction-0.txt").write_bytes((
                "; XY_EXECUTION_LAB_CHANCE_ANCHOR\r\n"
                "; XY_EXECUTION_LAB_TOUGHNESS_ANCHOR\r\n"
                "; XY_EQUIP_MAKER_RUNTIME_DAMAGE_COEFFICIENT_ANCHOR\r\n"
                "; XY_EQUIP_MAKER_POWER_ANCHOR\r\n"
            ).encode("gb18030"))
            plan = InitialCampService(self.platform).preflight(server, self.materials, client)
            self.assertEqual(plan.blockers, [])
            self.assertEqual(len(plan.npcs), 7)
            self.assertIn("xy.optional.rage.core", plan.install_plan.package_ids)
            self.assertIn("xy.optional.donate.core", plan.install_plan.package_ids)
            self.assertIn("xy.ui.attr-overview", plan.install_plan.package_ids)
            self.assertIn("xy.initial-camp.seven-npcs", plan.install_plan.package_ids)
            self.assertTrue(any(item.relative_path.endswith("狂暴之力-xycamp.txt") for item in plan.changes))
            rage_core = next(
                item for item in plan.changes
                if item.relative_path.endswith("玄渊功能/狂暴/狂暴NPC接口.txt")
            ).after.decode("gb18030")
            self.assertIn("[@XY_RAGE_NPC_MAIN]\r\n{", rage_core)
            self.assertIn("CHECKGAMEGIRD > 99", rage_core)
            self.assertIn("GAMEGIRD - 100", rage_core)
            self.assertNotIn("CHECKGAMEDIAMOND", rage_core)

    def test_lingfu_aliases_always_use_native_gamegird(self):
        for field, alias in (("材料A", "灵符"), ("材料A", "账户灵符"), ("货币A", "原生灵符")):
            rows = ({"状态": "可安装", field: alias, field + "数量": 7,
                     "称号名称": "测试称号", "称号编号": 20, "属性与数值": "测试属性+1"},)
            script = _title_chain_script("测试", rows, "XY_TEST")
            self.assertIn("CHECKGAMEGIRD > 6", script)
            self.assertIn("GAMEGIRD - 7", script)
            self.assertNotIn("CHECKITEM 灵符", script)
            self.assertNotIn("TAKE 灵符", script)

    def test_currency_columns_only_accept_native_currency(self):
        cases = (
            ("金币", "CHECKGOLD 7", "GOLDCOUNT - 7"),
            ("原生金币", "CHECKGOLD 7", "GOLDCOUNT - 7"),
            ("元宝", "CHECKGAMEGOLD > 6", "GAMEGOLD - 7"),
            ("原生元宝", "CHECKGAMEGOLD > 6", "GAMEGOLD - 7"),
            ("灵符", "CHECKGAMEGIRD > 6", "GAMEGIRD - 7"),
            ("金刚石", "CHECKGAMEDIAMOND 7", "GAMEDIAMOND - 7"),
            ("账户金刚石", "CHECKGAMEDIAMOND 7", "GAMEDIAMOND - 7"),
        )
        for name, expected_check, expected_action in cases:
            checks, actions = _cost_lines({"货币A": name, "货币A数量": 7})
            self.assertEqual(checks, [expected_check])
            self.assertEqual(actions, [expected_action])
            self.assertFalse(any(line.startswith("CHECKITEM ") for line in checks))
            self.assertFalse(any(line.startswith("TAKE ") for line in actions))

        with self.assertRaisesRegex(InitialCampError, "货币A只允许填写原生货币"):
            _cost_lines({"货币A": "失色锻造石", "货币A数量": 7})

    def test_sponsor_script_is_xlsx_driven_and_exposes_status(self):
        rows = (
            {"状态": "可安装", "档位/等级": "赞助1档", "称号名称": "赞助1档", "称号编号": 11,
             "材料A": "灵符", "材料A数量": 1000,
             "属性与数值": "处决+5%；韧性+20；基础爆率+100%；最大爆率+3%；伤害系数+5%",
             "处决": 5, "韧性": 20, "基础爆率": 100, "最大爆率": 3, "伤害系数": 5},
            {"状态": "可安装", "档位/等级": "赞助2档", "称号名称": "赞助2档", "称号编号": 12,
             "材料A": "灵符", "材料A数量": 2000,
             "属性与数值": "处决+6%；韧性+30；基础爆率+200%；最大爆率+3%；伤害系数+6%",
             "处决": 6, "韧性": 30, "基础爆率": 200, "最大爆率": 3, "伤害系数": 6},
        )
        script = _sponsor_script(rows, "XY_SPONSOR", "赞助称号")
        self.assertIn("赞助1档：处决5%、韧性20、基础爆率100%、最大爆率3%、伤害系数5%", script)
        self.assertIn("<升级赞助1档（1000灵符）/@XY_SPONSOR_1>", script)
        self.assertIn("CHECKGAMEGIRD > 999", script)
        self.assertIn("[@XY_SPONSOR_STATUS]", script)
        self.assertLess(script.index("CHECKFENGHAO 赞助2档", script.index("[@XY_SPONSOR_STATUS]")),
                        script.index("CHECKFENGHAO 赞助1档", script.index("[@XY_SPONSOR_STATUS]")))

    def test_sponsor_panel_sync_preserves_extended_hooks_and_only_updates_title_values(self):
        rows = (
            {"状态": "可安装", "档位/等级": "赞助1档", "称号名称": "赞助1档", "称号编号": 11,
             "处决": 5, "韧性": 20, "基础爆率": 100, "最大爆率": 3, "伤害系数": 7,
             "属性与数值": "处决+5%；韧性+20；基础爆率+100%；最大爆率+3%；伤害系数+7%"},
        )
        original = (
            "#IF\r\nCHECKFENGHAO 赞助1档\r\n#ACT\r\nINC N$XY_UI_A_Value02 1\r\n"
            "; XY_EQUIP_MAKER_PANEL_EXEC_CHANCE_ANCHOR\r\n"
            "MOV N$XY_UI_C_Value01 <$STR(N$XY_EXEC_Toughness)>\r\n"
            "#IF\r\nCHECKFENGHAO 赞助1档\r\n#ACT\r\nINC N$XY_UI_C_Value02 1\r\n"
            "; KEEP-THIRD-PARTY-HOOK\r\n"
        )

        changed = _sync_sponsor_panel_values(original, rows, "\r\n")

        self.assertIn("INC N$XY_UI_A_Value02 7", changed)
        self.assertIn("INC N$XY_UI_C_Value02 5", changed)
        self.assertIn("; KEEP-THIRD-PARTY-HOOK", changed)
        self.assertIn("; XY_EQUIP_MAKER_PANEL_EXEC_CHANCE_ANCHOR", changed)

    def test_composite_install_and_rollback_are_byte_exact(self):
        with tempfile.TemporaryDirectory(prefix="xydp-camp-tx-") as td:
            server = self.make_target(Path(td) / "server")
            client = self.make_client(Path(td) / "client")
            service = InitialCampService(self.platform)
            original = {path: path.read_bytes() for path in server.rglob("*") if path.is_file()}
            client_original = {path: path.read_bytes() for path in client.rglob("*") if path.is_file()}
            with patch.object(service.installer, "preflight", return_value=self.empty_dependency_plan(server)):
                plan = service.preflight(server, self.materials, client)
            receipt = service.install(plan)
            self.assertTrue((server / "Mir200" / "Envir" / "Market_Def" / "玄渊成长" / "转生-xycamp.txt").is_file())
            with patch.object(service.installer, "preflight", return_value=self.empty_dependency_plan(server)):
                second = service.preflight(server, self.materials, client)
            self.assertEqual(second.blockers, [])
            self.assertEqual(second.changes, [])
            service.installer.rollback(server, receipt.transaction_id)
            for path, data in original.items():
                self.assertEqual(path.read_bytes(), data)
            for path, data in client_original.items():
                self.assertEqual(path.read_bytes(), data)
            self.assertFalse((server / "Mir200" / "Envir" / "Market_Def" / "玄渊成长" / "转生-xycamp.txt").exists())

    def test_selected_donate_does_not_regenerate_other_six_npcs(self):
        with tempfile.TemporaryDirectory(prefix="xydp-camp-selected-donate-") as td:
            server = self.make_target(Path(td) / "server")
            client = self.make_client(Path(td) / "client")
            service = InitialCampService(self.platform)
            empty = InstallPlan(str(server), str(client), ["xy.optional.donate.core"],
                                {"xy.optional.donate.core": "test"}, {}, [])
            with patch.object(service.installer, "preflight", return_value=empty):
                plan = service.preflight_selected(server, self.materials, {"donate"}, client)
            self.assertEqual(plan.blockers, [])
            self.assertEqual([item.feature_id for item in plan.npcs], ["donate"])
            paths = {item.relative_path for item in plan.changes}
            self.assertIn("Mir200/Envir/Market_Def/玄渊运营/沙城捐献-xycamp.txt", paths)
            self.assertFalse(any("赞助称号-xycamp.txt" in path for path in paths))
            self.assertFalse(any("转生-xycamp.txt" in path for path in paths))
            self.assertFalse(any("马牌升级-xycamp.txt" in path for path in paths))
            self.assertFalse(any("军鼓升级-xycamp.txt" in path for path in paths))
            self.assertFalse(any("新手礼包-xycamp.txt" in path for path in paths))
            merchant = next(item for item in plan.changes if item.relative_path.endswith("MerChant.txt")).after.decode("gb18030")
            self.assertIn("玄渊运营/狂暴之力\t0\t322\t272", merchant)
            self.assertEqual(merchant.count("玄渊运营/沙城捐献\txycamp\t"), 1)

    def test_selected_donate_keeps_unchanged_shared_hooks_in_place(self):
        with tempfile.TemporaryDirectory(prefix="xydp-camp-selected-hook-order-") as td:
            server = self.make_target(Path(td) / "server")
            client = self.make_client(Path(td) / "client")
            qfunction_path = server / "Mir200" / "Envir" / "Market_Def" / "QFunction-0.txt"
            qfunction = qfunction_path.read_bytes().decode("gb18030")
            donate_row = read_workbook(self.materials / "02_捐献.xlsx").rows[0]
            donate_title = str(donate_row["称号名称"]).strip()
            donate_contents = _donate_anchor_contents(donate_row, donate_title)
            sponsor_contents = _sponsor_anchor_contents(
                read_workbook(self.materials / "03_赞助称号.xlsx").rows
            )
            for anchor in sorted(set(donate_contents) | set(sponsor_contents)):
                content = "\n".join(
                    part for part in (
                        donate_contents.get(anchor, ""),
                        sponsor_contents.get(anchor, ""),
                    ) if part
                )
                qfunction = install_managed_anchor_hook(
                    qfunction, "xy.test.foreign", anchor,
                    "#IF\n#ACT\nMOV N$XY_FOREIGN 1", "\r\n",
                ).text
                qfunction = install_managed_anchor_hook(
                    qfunction, "xy.initial-camp.seven-npcs", anchor,
                    content, "\r\n",
                ).text
            qfunction_path.write_bytes(qfunction.encode("gb18030"))

            service = InitialCampService(self.platform)
            empty = InstallPlan(str(server), str(client), ["xy.optional.donate.core"],
                                {"xy.optional.donate.core": "test"}, {}, [])
            with patch.object(service.installer, "preflight", return_value=empty):
                plan = service.preflight_selected(server, self.materials, {"donate"}, client)

            self.assertEqual(plan.blockers, [])
            self.assertFalse(any(
                item.relative_path.endswith("QFunction-0.txt")
                for item in plan.changes
            ))

    def test_selected_rebirth_needs_no_client_and_is_transactional(self):
        with tempfile.TemporaryDirectory(prefix="xydp-camp-selected-rebirth-") as td:
            server = self.make_target(Path(td) / "server")
            service = InitialCampService(self.platform)
            before = {path: path.read_bytes() for path in server.rglob("*") if path.is_file()}
            plan = service.preflight_selected(server, self.materials, {"rebirth"})
            self.assertEqual(plan.blockers, [])
            self.assertIsNone(plan.client)
            self.assertEqual([item.feature_id for item in plan.npcs], ["rebirth"])
            self.assertIn("xy.platform.config-sync.initial-camp", plan.install_plan.package_ids)
            self.assertNotIn("xy.initial-camp.seven-npcs", plan.install_plan.package_ids)
            self.assertTrue(any(item.relative_path.endswith("转生-xycamp.txt") for item in plan.changes))
            rebirth = next(
                item for item in plan.changes if item.relative_path.endswith("转生-xycamp.txt")
            ).after.decode("gb18030")
            self.assertIn("CHECKRENEWLEVEL", rebirth)
            self.assertIn("RENEWLEVEL 1", rebirth)
            qfunction = next(
                item for item in plan.changes if item.relative_path.endswith("QFunction-0.txt")
            ).after.decode("gb18030")
            self.assertIn("INC N$倍攻 <$STR(N$XY_REBIRTH_PowerBonus)>", qfunction)
            self.assertIn("INC N$XY_RT_Power <$STR(N$XY_REBIRTH_PowerBonus)>", qfunction)
            self.assertFalse(any("赞助称号-xycamp.txt" in item.relative_path for item in plan.changes))
            receipt = service.install(plan)
            service.installer.rollback(server, receipt.transaction_id)
            for path, data in before.items():
                self.assertEqual(path.read_bytes(), data)
            self.assertFalse((server / "Mir200/Envir/Market_Def/玄渊成长/转生-xycamp.txt").exists())


if __name__ == "__main__":
    unittest.main()
