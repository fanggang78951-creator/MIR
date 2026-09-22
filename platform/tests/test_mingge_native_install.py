from __future__ import annotations

import importlib
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))


V2_DEFINITION_HEADERS = (
    "候选ID", "启用", "目标装备名称", "中文具体部位", "命格名称", "默认文字颜色", "备注", "使用说明",
)
V2_PROPERTY_HEADERS = ("候选ID", "顺序", "中文属性", "属性值", "启用", "备注", "使用说明")
V2_SEGMENT_HEADERS = (
    "候选ID", "顺序", "片段角色", "来源属性", "片段文字", "颜色", "启用", "备注", "使用说明",
)
V2_PART_HELP_HEADERS = ("中文具体部位", "支持状态", "说明")
V2_PROPERTY_HELP_HEADERS = ("中文属性", "实效状态", "单位", "说明")


class NativeTestInstallTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _workbook(self) -> Path:
        path = self.root / "native-p1.xlsx"
        wb = Workbook()
        definitions = wb.active
        definitions.title = "候选定义"
        definitions.append(V2_DEFINITION_HEADERS)
        definitions.append((
            "tiebi-defence-5-p1", True, "鞭尸灵玉", "灵玉", "【铁壁命格·测试】", 255,
            "P1隔离样本", "只用于管理员隔离游戏验收",
        ))
        properties = wb.create_sheet("自定义属性")
        properties.append(V2_PROPERTY_HEADERS)
        properties.append(("tiebi-defence-5-p1", 1, "防御", 5, True, "绑定1+行17黄金样本", "只填中文属性和值"))
        segments = wb.create_sheet("显示片段")
        segments.append(V2_SEGMENT_HEADERS)
        segments.append(("tiebi-defence-5-p1", 1, "命格名称", None, None, 249, True, "", "名称由候选定义派生"))
        segments.append(("tiebi-defence-5-p1", 2, "属性名称", "防御", None, 250, True, "", "一组必须同源"))
        segments.append(("tiebi-defence-5-p1", 3, "正负号", "防御", None, 69, True, "", "一组必须同源"))
        segments.append(("tiebi-defence-5-p1", 4, "属性值", "防御", None, 69, True, "", "一组必须同源"))
        part_help = wb.create_sheet("部位说明")
        part_help.append(V2_PART_HELP_HEADERS)
        part_help.append(("灵玉", "支持", "当前端实测位置17"))
        property_help = wb.create_sheet("属性说明")
        property_help.append(V2_PROPERTY_HELP_HEADERS)
        property_help.append(("防御", "引擎直接属性（已游戏实测）", "", "绑定1+行17"))
        wb.save(path)
        wb.close()
        return path

    def _server(self) -> Path:
        server = self.root / "MirServer"
        envir = server / "Mir200" / "Envir"
        (envir / "Market_Def").mkdir(parents=True)
        (envir / "QuestDiary" / "玄渊命格").mkdir(parents=True)
        (envir / "UserCmd.txt").write_bytes("命格显示验收\t96\r\n".encode("gb18030"))
        (envir / "Market_Def" / "QFunction-0.txt").write_bytes("[@Login]\r\n#ACT\r\nBREAK\r\n".encode("gb18030"))
        (envir / "QuestDiary" / "玄渊命格" / "命格核心.txt").write_bytes(b"OLD-ROUTE-SENTINEL\r\n")
        return server

    def test_preflight_plans_only_usercmd_qfunction_and_isolated_test_script(self) -> None:
        try:
            api = importlib.import_module("xydp.mingge_native_install")
        except ModuleNotFoundError:
            self.fail("缺少命格P1隔离安装模块")
        server = self._server()
        plan = api.plan_native_test_install(
            self._workbook(), "tiebi-defence-5-p1", server
        )
        self.assertEqual(plan.blockers, ())
        self.assertEqual(plan.candidate_id, "tiebi-defence-5-p1")
        relative_paths = {item.relative_path for item in plan.files}
        self.assertEqual(relative_paths, {
            "Mir200/Envir/UserCmd.txt",
            "Mir200/Envir/Market_Def/QFunction-0.txt",
            "Mir200/Envir/QuestDiary/玄渊验收/命格平台单件验收.txt",
        })
        script = next(item.after for item in plan.files if item.relative_path.endswith("命格平台单件验收.txt"))
        decoded = script.decode("gb18030")
        self.assertIn("EQUAL <$JADE> 鞭尸灵玉", decoded)
        self.assertIn("SetCustomItemText 17 {【铁壁命格·测试】|249}{防御|250}{+|69}{5|69}", decoded)
        script_lines = decoded.splitlines()
        self.assertEqual(script_lines.count("LockUpdateItem 17"), 1)
        self.assertEqual(script_lines.count("UpdateItem 17"), 1)
        self.assertNotIn("GAMEGOLD", decoded)
        self.assertNotIn("LINKGIVEITEM", decoded.upper())

    def test_install_and_rollback_are_byte_exact_and_leave_old_route_untouched(self) -> None:
        api = importlib.import_module("xydp.mingge_native_install")
        self.assertTrue(hasattr(api, "install_native_test_candidate"), "缺少P1事务安装入口")
        self.assertTrue(hasattr(api, "rollback_native_test_install"), "缺少P1逐字节回滚入口")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_test_install(
            self._workbook(), "tiebi-defence-5-p1", server
        )
        before = {
            item.relative_path: item.path.read_bytes() if item.path.exists() else None
            for item in plan.files
        }
        old_route = server / "Mir200" / "Envir" / "QuestDiary" / "玄渊命格" / "命格核心.txt"
        old_route_before = old_route.read_bytes()

        receipt = api.install_native_test_candidate(plan, platform)

        self.assertEqual(receipt.status, "installed-test-candidate")
        self.assertTrue(receipt.receipt_path.is_file())
        self.assertEqual(old_route.read_bytes(), old_route_before)
        for item in plan.files:
            self.assertEqual(item.path.read_bytes(), item.after)

        rollback = api.rollback_native_test_install(
            platform, receipt.transaction_id, server
        )

        self.assertEqual(rollback.status, "rolled-back")
        self.assertEqual(old_route.read_bytes(), old_route_before)
        for item in plan.files:
            expected = before[item.relative_path]
            if expected is None:
                self.assertFalse(item.path.exists())
            else:
                self.assertEqual(item.path.read_bytes(), expected)

    def test_cli_preflight_reports_exact_changes_without_writing_server(self) -> None:
        from xydp.cli import main as cli_main

        server = self._server()
        workbook = self._workbook()
        before = {
            path.relative_to(server).as_posix(): path.read_bytes()
            for path in server.rglob("*") if path.is_file()
        }
        stdout = io.StringIO()
        stderr = io.StringIO()
        try:
            with redirect_stdout(stdout), redirect_stderr(stderr):
                code = cli_main([
                    "--root", str(self.root / "platform"),
                    "mingge-native-test-preflight",
                    "--input", str(workbook),
                    "--candidate", "tiebi-defence-5-p1",
                    "--server", str(server),
                ])
        except SystemExit as exc:
            self.fail(f"缺少P1预检CLI：{exc}")
        self.assertEqual(code, 0)
        summary = json.loads(stdout.getvalue())
        self.assertEqual(summary["operation"], "mingge-native-p1-test")
        self.assertEqual(summary["blockers"], [])
        self.assertEqual(len(summary["changes"]), 3)
        after = {
            path.relative_to(server).as_posix(): path.read_bytes()
            for path in server.rglob("*") if path.is_file()
        }
        self.assertEqual(after, before)

    def test_preflight_blocks_any_second_property(self) -> None:
        api = importlib.import_module("xydp.mingge_native_install")
        workbook = self._workbook()
        book = load_workbook(workbook)
        book["自定义属性"].append((
            "tiebi-defence-5-p1", 2, "神力倍攻", 10, True, "不允许进入P1", "只用于阻断测试",
        ))
        book.save(workbook)
        book.close()

        plan = api.plan_native_test_install(workbook, "tiebi-defence-5-p1", self._server())

        self.assertTrue(any("只允许单条真实属性" in blocker for blocker in plan.blockers))

    def test_install_refuses_when_target_drifts_after_preflight(self) -> None:
        api = importlib.import_module("xydp.mingge_native_install")
        server = self._server()
        plan = api.plan_native_test_install(self._workbook(), "tiebi-defence-5-p1", server)
        qfunction = server / "Mir200" / "Envir" / "Market_Def" / "QFunction-0.txt"
        qfunction.write_bytes(qfunction.read_bytes() + b"; EXTERNAL-CHANGE\r\n")
        before = qfunction.read_bytes()

        with self.assertRaisesRegex(api.NativeTestInstallError, "预检后发生变化"):
            api.install_native_test_candidate(plan, self.root / "platform")

        self.assertEqual(qfunction.read_bytes(), before)
        self.assertFalse((server / "Mir200" / "Envir" / "QuestDiary" / "玄渊验收" / "命格平台单件验收.txt").exists())

    def test_rollback_refuses_to_overwrite_post_install_drift(self) -> None:
        api = importlib.import_module("xydp.mingge_native_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_test_install(self._workbook(), "tiebi-defence-5-p1", server)
        receipt = api.install_native_test_candidate(plan, platform)
        qfunction = server / "Mir200" / "Envir" / "Market_Def" / "QFunction-0.txt"
        qfunction.write_bytes(qfunction.read_bytes() + b"; EXTERNAL-CHANGE\r\n")
        drifted = qfunction.read_bytes()

        with self.assertRaisesRegex(api.NativeTestInstallError, "已漂移"):
            api.rollback_native_test_install(platform, receipt.transaction_id, server)

        self.assertEqual(qfunction.read_bytes(), drifted)

    def test_cli_install_requires_confirmation_and_cli_rollback_is_exact(self) -> None:
        from xydp.cli import main as cli_main

        server = self._server()
        workbook = self._workbook()
        platform = self.root / "platform"
        before = {
            path.relative_to(server).as_posix(): path.read_bytes()
            for path in server.rglob("*") if path.is_file()
        }
        with self.assertRaises(SystemExit):
            cli_main([
                "--root", str(platform), "mingge-native-test-install",
                "--input", str(workbook), "--candidate", "tiebi-defence-5-p1",
                "--server", str(server),
            ])

        stdout = io.StringIO()
        with redirect_stdout(stdout):
            code = cli_main([
                "--root", str(platform), "mingge-native-test-install",
                "--input", str(workbook), "--candidate", "tiebi-defence-5-p1",
                "--server", str(server), "--yes",
            ])
        self.assertEqual(code, 0)
        install_summary = json.loads(stdout.getvalue())
        transaction_id = install_summary["transaction_id"]

        stdout = io.StringIO()
        with redirect_stdout(stdout):
            code = cli_main([
                "--root", str(platform), "mingge-native-test-rollback",
                "--transaction", transaction_id, "--server", str(server), "--yes",
            ])
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(stdout.getvalue())["status"], "rolled-back")
        after = {
            path.relative_to(server).as_posix(): path.read_bytes()
            for path in server.rglob("*") if path.is_file()
        }
        self.assertEqual(after, before)


if __name__ == "__main__":
    unittest.main()
