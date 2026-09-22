from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from xydp.installer import InstallError, Installer  # noqa: E402
from xydp.repository import PackageRepository  # noqa: E402
from xydp.resident import ResidentService  # noqa: E402
from xydp.validator import validate_package  # noqa: E402


PACKAGE_ID = "xy.resident.equipment-level-bonus"
PACKAGE_ROOT = ROOT / "packages/candidate" / PACKAGE_ID
QFUNCTION = Path("Mir200/Envir/Market_Def/QFunction-0.txt")
TEXTVAR = Path("Mir200/Envir/CustomItemPropertyTextVarList.txt")
CORE = Path("Mir200/Envir/QuestDiary/玄渊等级加成/等级加成核心.txt")
TEXTVAR41 = "{人物等级∶|251}+$$2"


def make_server(root: Path, *, line41: str = "", extra_qfunction: str = "") -> Path:
    envir = root / "Mir200/Envir"
    envir.mkdir(parents=True)
    (root / "Mir200/M2Server.exe").write_bytes(b"M2")
    (envir / "MapInfo.txt").write_bytes(b"[0 test]\r\n")
    qfunction = root / QFUNCTION
    qfunction.parent.mkdir(parents=True)
    qfunction.write_bytes((
        "[@PlayLogin]\r\n#IF\r\n#ACT\r\nBREAK\r\n"
        "[@TakeOnEx]\r\n#IF\r\n#ACT\r\nBREAK\r\n"
        "[@TakeOffEx]\r\n#IF\r\n#ACT\r\nBREAK\r\n"
        + extra_qfunction
    ).encode("gb18030"))
    rows = [""] * 41
    rows[40] = line41
    textvar = root / TEXTVAR
    textvar.write_bytes(("\r\n".join(rows) + "\r\n").encode("gb18030"))
    return root


class EquipmentLevelBonusPackageTests(unittest.TestCase):
    def _require_package(self) -> None:
        self.assertTrue(PACKAGE_ROOT.is_dir(), f"缺少常驻包：{PACKAGE_ROOT}")

    def test_manifest_is_candidate_resident_and_collected_by_resident_base(self):
        self._require_package()
        package = validate_package(PACKAGE_ROOT)
        self.assertEqual((package.status, package.residency, package.install_route), (
            "candidate", "resident", "generic"
        ))
        repository = PackageRepository(ROOT / "packages")
        repository.refresh()
        self.assertIn(PACKAGE_ID, ResidentService(repository, ROOT / "backups").package_ids())

    def test_preflight_compiles_textvar_hooks_fixed_anchor_and_literal_delta_table(self):
        self._require_package()
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            repository = PackageRepository(ROOT / "packages")
            repository.refresh()
            plan = Installer(repository, base / "backups").preflight(server, [PACKAGE_ID], {})
        rendered = {Path(item.relative_path): item.after.decode("gb18030") for item in plan.changes}
        self.assertEqual(rendered[TEXTVAR].splitlines()[40], TEXTVAR41)
        qfunction = rendered[QFUNCTION]
        core = rendered[CORE]
        self.assertIn("; XY-LB-V1-BEGIN", qfunction)
        self.assertIn("; XY_EQUIP_MAKER_LEVEL_BONUS_ANCHOR", qfunction)
        self.assertIn("GetAllCustomItemValueByTextLine 60 -1 41 N$XY_LB_TMP N$XY_LB_CUSTOM N$XY_LB_TMP", core)
        self.assertIn("INC N$XY_LB_DESIRED <$STR(N$XY_LB_FIXED)>", core)
        self.assertIn("LARGE N$XY_LB_DESIRED 50", core)
        self.assertIn("MOV N$XY_LB_MAX_EFFECTIVE 500", core)
        self.assertIn("MOV U470 <$STR(N$XY_LB_DESIRED)>", core)
        for value in (1, 10, 50):
            self.assertIn(f"EQUAL N$XY_LB_DELTA {value}", core)
            self.assertIn(f"CHANGELEVEL + {value}", core)
            self.assertIn(f"EQUAL N$XY_LB_DELTA -{value}", core)
            self.assertIn(f"CHANGELEVEL - {value}", core)
        self.assertNotIn("CHANGELEVEL + <$STR", core)
        self.assertNotIn("CHANGELEVEL - <$STR", core)
        self.assertIn("DELAYGOTO 1200 @XY_LEVEL_BONUS_RECALC_DELAY", qfunction)
        self.assertEqual(qfunction.count("DELAYGOTO 150 @XY_LEVEL_BONUS_RECALC_DELAY"), 2)
        for forbidden in ("UserCmd103", "UserCmd104", "UserCmd105", "等级测1", "等级测查", "等级测清", "SENDMSG"):
            self.assertNotIn(forbidden, qfunction + core)

    def test_conflicting_textvar_u_variable_and_legacy_test_each_block_without_writes(self):
        self._require_package()
        repository = PackageRepository(ROOT / "packages")
        repository.refresh()
        cases = (
            ("textvar", {"line41": "其他系统占用"}, "第41行"),
            ("legacy", {"extra_qfunction": "; XY-LV-VERIFY-BEGIN\r\n"}, "禁止文本"),
        )
        for name, kwargs, expected in cases:
            with self.subTest(name=name), tempfile.TemporaryDirectory() as td:
                base = Path(td)
                server = make_server(base / "server", **kwargs)
                before = {p.relative_to(server): p.read_bytes() for p in server.rglob("*") if p.is_file()}
                with self.assertRaisesRegex(InstallError, expected):
                    Installer(repository, base / "backups").preflight(server, [PACKAGE_ID], {})
                after = {p.relative_to(server): p.read_bytes() for p in server.rglob("*") if p.is_file()}
                self.assertEqual(after, before)

        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            conflict = server / "Mir200/Envir/QuestDiary/其他系统/冲突.txt"
            conflict.parent.mkdir(parents=True)
            conflict.write_text("MOV U470 9\n", encoding="gb18030")
            with self.assertRaisesRegex(InstallError, "U470.*冲突.txt"):
                Installer(repository, base / "backups").preflight(server, [PACKAGE_ID], {})

    def test_install_is_idempotent_and_rollback_restores_target_bytes(self):
        self._require_package()
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            before = {p.relative_to(server): p.read_bytes() for p in server.rglob("*") if p.is_file()}
            repository = PackageRepository(ROOT / "packages")
            repository.refresh()
            installer = Installer(repository, base / "backups")
            receipt = installer.install(installer.preflight(server, [PACKAGE_ID], {}))
            self.assertEqual(installer.preflight(server, [PACKAGE_ID], {}).changes, [])
            installer.rollback(server, receipt.transaction_id)
            after = {
                p.relative_to(server): p.read_bytes()
                for p in server.rglob("*")
                if p.is_file() and ".xydp" not in p.parts
            }
        self.assertEqual(after, before)

    def test_preexisting_different_core_is_blocked_instead_of_overwritten(self):
        self._require_package()
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            core = server / CORE
            core.parent.mkdir(parents=True)
            original = "其他等级系统\r\n".encode("gb18030")
            core.write_bytes(original)
            repository = PackageRepository(ROOT / "packages")
            repository.refresh()
            with self.assertRaisesRegex(InstallError, "目标文件已存在且内容不同"):
                Installer(repository, base / "backups").preflight(server, [PACKAGE_ID], {})
            self.assertEqual(core.read_bytes(), original)

    def test_shared_workbooks_expose_level_bonus_before_trailing_metadata_columns(self):
        for path in (
            ROOT / "所需材料表格汇总/09_装备批量生成.xlsx",
            ROOT / "做装备/templates/XuanYuanItems.xlsx",
        ):
            with self.subTest(path=path):
                workbook = load_workbook(path, data_only=False, read_only=False)
                sheet = workbook["装备导入表"]
                headers = [cell.value for cell in sheet[1]]
                self.assertEqual(headers.count("等级增加"), 1)
                self.assertEqual(headers.index("等级增加") + 1, 67)
                self.assertEqual(headers[-2:], ["备注", "悬浮分类"])
                validations = []
                for validation in sheet.data_validations.dataValidation:
                    covered_rows: set[int] = set()
                    for cell_range in validation.sqref.ranges:
                        if cell_range.min_col <= 67 <= cell_range.max_col:
                            covered_rows.update(range(cell_range.min_row, cell_range.max_row + 1))
                    if covered_rows.issuperset(range(2, 1001)):
                        validations.append(validation)
                self.assertEqual(len(validations), 1)
                self.assertEqual(
                    (validations[0].type, str(validations[0].formula1), str(validations[0].formula2)),
                    ("whole", "0", "10"),
                )
                source = workbook["资料来源"]
                self.assertEqual((source["A12"].value, source["B12"].value), (
                    "等级增加", PACKAGE_ID
                ))
                workbook.close()


if __name__ == "__main__":
    unittest.main()
