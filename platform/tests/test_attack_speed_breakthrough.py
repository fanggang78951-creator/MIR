from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from xydp.attack_speed_breakthrough import (  # noqa: E402
    CORE_RELATIVE,
    PACKAGE_ID,
    QFUNCTION_RELATIVE,
    ROUTE,
    TEXTVAR_RELATIVE,
    AttackSpeedBreakthroughService,
    load_attack_speed_workbook,
)
from xydp.config_sync import ConfigSyncService  # noqa: E402
from xydp.validator import validate_package  # noqa: E402


TEXTVAR40 = "{攻速突破∶|251}+$$2"
QMANAGE_RELATIVE = Path("Mir200/Envir/MapQuest_Def/QManage.txt")


def make_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "基础设置"
    ws.append(("配置项", "配置值", "说明"))
    for row in (
        ("配置版本", 1, "固定"),
        ("系统编号", "attack_speed_breakthrough", "固定"),
        ("基础攻速阈值", 20, "固定"),
        ("单件突破上限", 30, "固定"),
        ("全身突破上限", 30, "固定"),
        ("绝对攻速上限", 50, "固定"),
        ("TextVar行", 40, "固定"),
        ("登录延迟毫秒", 1000, "固定"),
        ("穿脱延迟毫秒", 100, "固定"),
    ):
        ws.append(row)
    help_ws = wb.create_sheet("填写说明")
    help_ws.append(("项目", "说明"))
    help_ws.append(("攻速突破", "单件0至30，多件相加后全身封顶30。"))
    wb.save(path)
    wb.close()
    return path


def make_server(root: Path, textvar40: str = "") -> Path:
    qf = root / QFUNCTION_RELATIVE
    qf.parent.mkdir(parents=True, exist_ok=True)
    qf.write_bytes(
        (
            "[@PlayLogin]\r\n#IF\r\n#ACT\r\nBREAK\r\n"
            "[@TakeOnEx]\r\n#IF\r\n#ACT\r\nBREAK\r\n"
            "[@TakeOffEx]\r\n#IF\r\n#ACT\r\nBREAK\r\n"
        ).encode("gb18030")
    )
    textvar = root / TEXTVAR_RELATIVE
    textvar.parent.mkdir(parents=True, exist_ok=True)
    textvar.write_bytes((("\r\n" * 39) + textvar40 + "\r\n").encode("gb18030"))
    qmanage = root / QMANAGE_RELATIVE
    qmanage.parent.mkdir(parents=True, exist_ok=True)
    qmanage.write_bytes(
        (
            "[@Login]\r\n#IF\r\n#ACT\r\nGOTO @MAIN1\r\n"
            "[@MAIN1]\r\n#IF\r\n#ACT\r\nBREAK\r\n"
        ).encode("gb18030")
    )
    return root


class AttackSpeedBreakthroughTests(unittest.TestCase):
    def test_workbook_contract_and_boundaries(self):
        with tempfile.TemporaryDirectory() as td:
            path = make_workbook(Path(td) / "43.xlsx")
            book = load_attack_speed_workbook(path)
        self.assertEqual(book.base_cap, 20)
        self.assertEqual(book.item_maximum, 30)
        self.assertEqual(book.total_maximum, 30)
        self.assertEqual(book.absolute_maximum, 50)
        self.assertEqual(book.textvar_line, 40)

    def test_preflight_compiles_login_bridge_core_three_hooks_and_textvar_without_test_leakage(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            plan = AttackSpeedBreakthroughService(ROOT).preflight(
                make_server(base / "server"), make_workbook(base / "43.xlsx")
            )
        self.assertFalse(plan.blockers)
        rendered = {Path(change.relative_path): change.after.decode("gb18030") for change in plan.changes}
        core = rendered[CORE_RELATIVE]
        qfunction = rendered[QFUNCTION_RELATIVE]
        self.assertNotIn(QMANAGE_RELATIVE, rendered, "登录统一由属性总览分发PlayLogin，攻速包不得再生成第二条QManage桥")
        textvar = rendered[TEXTVAR_RELATIVE]
        self.assertIn("GetAllCustomItemValueByTextLine 60 -1 40", core)
        self.assertIn("GetAllCustomItemValueByTextLine 60 -1 40 N$XY_AS_TMP N$XY_AS_BREAK_RAW N$XY_AS_TMP", core)
        self.assertEqual(core.count("ChangeSpeed 2"), 2)
        self.assertIn("INC N$XY_AS_BREAK_RAW <$STR(N$XY_AS_FIXED_BREAK)>", core)
        self.assertIn("DELAYGOTO 1000 @XY_AS_FORMAL_RECALC_DELAY", qfunction)
        self.assertEqual(qfunction.count("DELAYGOTO 100 @XY_AS_FORMAL_RECALC_DELAY"), 2)
        self.assertIn("; XY_EQUIP_MAKER_ATTACK_SPEED_BREAK_ANCHOR", qfunction)
        self.assertIn("#CALL [\\玄渊攻速突破\\全身攻速阈值核心.txt] @XY_AS_CAP_RECALC", qfunction)
        self.assertEqual(textvar.splitlines()[39], TEXTVAR40)
        for forbidden in ("UserCmd100", "UserCmd101", "UserCmd102", "UserCmd103", "攻速测20", "攻速测40", "攻速测50", "测试木剑", "SENDMSG"):
            self.assertNotIn(forbidden, core + qfunction)

    def test_preflight_removes_the_retired_exact_qmanage_login_bridge(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            qmanage_path = server / QMANAGE_RELATIVE
            qmanage = qmanage_path.read_bytes().decode("gb18030")
            qmanage = qmanage.replace(
                "[@MAIN1]\r\n",
                "[@MAIN1]\r\n"
                "; XY-AS-LOGIN-BRIDGE-MAIN1-BEGIN\r\n"
                "#IF\r\n#ACT\r\nDELAYGOTO 1000 @XY_AS_LOGIN_RECALC_BRIDGE\r\n"
                "; XY-AS-LOGIN-BRIDGE-MAIN1-END\r\n",
                1,
            )
            qmanage += (
                "; XY-AS-LOGIN-BRIDGE-LABEL-BEGIN\r\n"
                "[@XY_AS_LOGIN_RECALC_BRIDGE]\r\n#IF\r\n#ACT\r\n"
                "GOTOLABEL 8 @XY_AS_FORMAL_RECALC_DELAY <$X> <$Y> 0 0\r\nBREAK\r\n"
                "; XY-AS-LOGIN-BRIDGE-LABEL-END\r\n"
            )
            qmanage_path.write_bytes(qmanage.encode("gb18030"))

            plan = AttackSpeedBreakthroughService(ROOT).preflight(server, make_workbook(base / "43.xlsx"))
        self.assertFalse(plan.blockers)
        rendered = {Path(change.relative_path): change.after.decode("gb18030") for change in plan.changes}
        self.assertIn(QMANAGE_RELATIVE, rendered)
        self.assertNotIn("XY-AS-LOGIN-BRIDGE", rendered[QMANAGE_RELATIVE])
        self.assertNotIn("[@XY_AS_LOGIN_RECALC_BRIDGE]", rendered[QMANAGE_RELATIVE])

    def test_other_active_changespeed2_blocks_with_zero_changes(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            conflict = server / "Mir200/Envir/QuestDiary/其他系统/冲突.txt"
            conflict.parent.mkdir(parents=True)
            conflict.write_bytes(b"[@X]\r\n#ACT\r\nChangeSpeed 2 5\r\n")
            plan = AttackSpeedBreakthroughService(ROOT).preflight(server, make_workbook(base / "43.xlsx"))
        self.assertFalse(plan.changes)
        self.assertTrue(any("ChangeSpeed 2" in blocker for blocker in plan.blockers))

    def test_utf8_business_script_is_scanned_for_changespeed_conflict(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            utf8_script = server / "Mir200/Envir/QuestDiary/无关系统/UTF8脚本.txt"
            utf8_script.parent.mkdir(parents=True)
            utf8_script.write_text("[@测试]\r\n#ACT\r\nChangeSpeed 2 5\r\n", encoding="utf-8-sig")
            plan = AttackSpeedBreakthroughService(ROOT).preflight(server, make_workbook(base / "43.xlsx"))
        self.assertFalse(plan.changes)
        self.assertTrue(any("UTF8脚本.txt" in blocker and "ChangeSpeed 2" in blocker for blocker in plan.blockers))

    def test_undecodable_envir_business_script_blocks_with_exact_path(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            undecodable = server / "Mir200/Envir/QuestDiary/无关系统/不可判定.txt"
            undecodable.parent.mkdir(parents=True)
            undecodable.write_bytes(b"\x81")
            plan = AttackSpeedBreakthroughService(ROOT).preflight(server, make_workbook(base / "43.xlsx"))
        self.assertFalse(plan.changes)
        self.assertTrue(any("不可判定.txt" in blocker and "编码" in blocker for blocker in plan.blockers))

    def test_textvar40_conflict_blocks_with_zero_changes(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            plan = AttackSpeedBreakthroughService(ROOT).preflight(
                make_server(base / "server", "其他系统占用"), make_workbook(base / "43.xlsx")
            )
        self.assertFalse(plan.changes)
        self.assertTrue(any("TextVar第40行" in blocker for blocker in plan.blockers))

    def test_install_is_idempotent_and_rollback_restores_bytes(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            workbook = make_workbook(base / "43.xlsx")
            before = {p.relative_to(server): p.read_bytes() for p in server.rglob("*") if p.is_file()}
            service = AttackSpeedBreakthroughService(base / "platform")
            receipt = service.install(service.preflight(server, workbook))
            self.assertFalse(service.preflight(server, workbook).changes)
            service.rollback(server, receipt.transaction_id)
            after = {p.relative_to(server): p.read_bytes() for p in server.rglob("*") if p.is_file() and ".xydp" not in p.parts}
        self.assertEqual(after, before)

    def test_preflight_preserves_compatible_wash_speed_and_fixed_equipment_extensions(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            server = make_server(base / "server")
            workbook = make_workbook(base / "43.xlsx")
            service = AttackSpeedBreakthroughService(base / "platform")
            service.install(service.preflight(server, workbook))

            core_path = server / CORE_RELATIVE
            core = core_path.read_bytes().decode("gb18030")
            core = core.replace(
                "MOV N$XY_AS_RAW <$HITSPD>\r\n",
                "MOV N$XY_AS_RAW <$HITSPD>\r\n"
                "; XY-EQUIPMENT-WASH-IMPORT-V1-NORMAL-SPEED-BEGIN\r\n"
                "MOV N$XY_WASH_AS_TMP 0\r\n"
                "MOV N$XY_WASH_AS_NORMAL 0\r\n"
                "GetAllCustomItemValueByTextLine 60 -1 46 N$XY_WASH_AS_TMP N$XY_WASH_AS_NORMAL N$XY_WASH_AS_TMP\r\n"
                "INC N$XY_AS_RAW <$STR(N$XY_WASH_AS_NORMAL)>\r\n"
                "; XY-EQUIPMENT-WASH-IMPORT-V1-NORMAL-SPEED-END\r\n",
                1,
            )
            core_path.write_bytes(core.encode("gb18030"))

            qfunction_path = server / QFUNCTION_RELATIVE
            qfunction = qfunction_path.read_bytes().decode("gb18030")
            qfunction = qfunction.replace(
                "; XY_EQUIP_MAKER_ATTACK_SPEED_BREAK_ANCHOR\r\n",
                "; XY_EQUIP_MAKER_ATTACK_SPEED_BREAK_ANCHOR\r\n"
                "; XY-EQUIP-MAKER 测试装备: 攻速突破+7\r\n"
                "#IF\r\nCHECKITEMW 测试装备 1\r\n#ACT\r\nINC N$XY_AS_FIXED_BREAK 7\r\n",
                1,
            )
            qfunction_path.write_bytes(qfunction.encode("gb18030"))

            plan = service.preflight(server, workbook)
        self.assertFalse(plan.blockers)
        self.assertFalse(plan.changes)

    def test_registered_document_and_candidate_package(self):
        self.assertEqual(ConfigSyncService._route([type("D", (), {"id": "attack_speed_breakthrough"})()]), ROUTE)
        package = validate_package(ROOT / "packages/candidate" / PACKAGE_ID)
        self.assertEqual(package.status, "candidate")
        self.assertEqual(package.residency, "optional")
        self.assertEqual(package.install_route, ROUTE)
        evidence = json.loads(
            (ROOT / "packages/candidate" / PACKAGE_ID / "manifest.json").read_text(encoding="utf-8")
        )["evidence"]
        self.assertIn("current-server:formula-and-equip-recalc-game-accepted", evidence)
        self.assertIn("current-server:unified-playlogin-dispatch-pending-validation", evidence)

    def test_formal_workbooks_and_document_registry_expose_attack_speed_breakthrough(self):
        registry = json.loads((ROOT / "所需材料表格汇总/00_填写文档注册表.json").read_text(encoding="utf-8"))
        entries = [item for item in registry["documents"] if item["id"] == "attack_speed_breakthrough"]
        self.assertEqual(entries, [{
            "id": "attack_speed_breakthrough",
            "file": "43_攻速突破.xlsx",
            "consumer": "attack-speed-breakthrough",
            "legacy": None,
        }])
        formal = load_attack_speed_workbook(ROOT / "所需材料表格汇总/43_攻速突破.xlsx")
        self.assertEqual((formal.base_cap, formal.item_maximum, formal.total_maximum, formal.absolute_maximum), (20, 30, 30, 50))

        for path in (ROOT / "所需材料表格汇总/09_装备批量生成.xlsx", ROOT / "做装备/templates/XuanYuanItems.xlsx"):
            wb = load_workbook(path, data_only=False, read_only=False)
            ws = wb["装备导入表"]
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(headers.count("攻速突破"), 1)
            self.assertEqual(headers.index("攻速突破") + 1, 66)
            self.assertEqual(headers[-2:], ["备注", "悬浮分类"])
            validations = []
            for validation in ws.data_validations.dataValidation:
                covered_rows: set[int] = set()
                for cell_range in validation.sqref.ranges:
                    if cell_range.min_col <= 66 <= cell_range.max_col:
                        covered_rows.update(range(cell_range.min_row, cell_range.max_row + 1))
                if covered_rows.issuperset(range(2, 1001)):
                    validations.append(validation)
            self.assertEqual(len(validations), 1)
            self.assertEqual((validations[0].type, str(validations[0].formula1), str(validations[0].formula2)), ("whole", "0", "30"))
            source = wb["资料来源"]
            self.assertEqual((source["A11"].value, source["B11"].value), ("攻速突破", PACKAGE_ID))
            wb.close()


if __name__ == "__main__":
    unittest.main()
