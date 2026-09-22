from __future__ import annotations

import hashlib
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from xydp.mingge_native import (  # noqa: E402
    NativeCandidateError,
    TRUSTED_EQUIP_SLOTS,
    compile_native_candidate,
    generate_native_candidates,
    read_native_candidate_workbook,
)
from xydp import mingge_native as mingge_native_module  # noqa: E402
from xydp.cli import main as cli_main  # noqa: E402


DEFINITION_HEADERS = (
    "候选ID",
    "启用",
    "目标装备名称",
    "装备位置",
    "写真实属性",
    "属性行",
    "属性颜色",
    "属性绑定",
    "属性位置",
    "属性模式",
    "属性百分比类型",
    "属性值1",
    "属性值2",
    "属性值3",
    "默认文字颜色",
    "备注",
)
SEGMENT_HEADERS = (
    "候选ID",
    "顺序",
    "片段角色",
    "片段文字",
    "颜色",
    "启用",
    "备注",
)
V2_DEFINITION_HEADERS = (
    "候选ID",
    "启用",
    "目标装备名称",
    "中文具体部位",
    "命格名称",
    "默认文字颜色",
    "备注",
    "使用说明",
)
V2_PROPERTY_HEADERS = ("候选ID", "顺序", "中文属性", "属性值", "启用", "备注", "使用说明")
V2_SEGMENT_HEADERS = (
    "候选ID",
    "顺序",
    "片段角色",
    "来源属性",
    "片段文字",
    "颜色",
    "启用",
    "备注",
    "使用说明",
)
V2_PART_HELP_HEADERS = ("中文具体部位", "支持状态", "说明")
V2_PROPERTY_HELP_HEADERS = ("中文属性", "实效状态", "单位", "说明")
SIMPLE_CONFIG_HEADERS = (
    "启用", "目标装备", "部位", "命格名称", "颜色",
    "属性1", "数值1", "属性2", "数值2", "属性3", "数值3",
)
SIMPLE_DICTIONARY_HEADERS = ("中文属性", "单位", "中文部位", "颜色")


class NativeCandidateTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _workbook(
        self,
        *,
        definitions: list[tuple[object, ...]] | None = None,
        segments: list[tuple[object, ...]] | None = None,
    ) -> Path:
        path = self.root / "native.xlsx"
        wb = Workbook()
        definitions_sheet = wb.active
        definitions_sheet.title = "候选定义"
        definitions_sheet.append(DEFINITION_HEADERS)
        rows = definitions or [
            (
                "four-colors",
                True,
                "鞭尸灵玉",
                17,
                True,
                17,
                250,
                1,
                17,
                0,
                9,
                5,
                0,
                0,
                255,
                "游戏黄金样本",
            )
        ]
        for row in rows:
            definitions_sheet.append(row)
        segments_sheet = wb.create_sheet("显示片段")
        segments_sheet.append(SEGMENT_HEADERS)
        segment_rows = segments or [
            ("four-colors", 1, "固定文字", "红", 249, True, ""),
            ("four-colors", 2, "固定文字", "橙", 69, True, ""),
            ("four-colors", 3, "固定文字", "绿", 250, True, ""),
            ("four-colors", 4, "固定文字", "彩", 251, True, ""),
        ]
        for row in segment_rows:
            segments_sheet.append(row)
        wb.save(path)
        wb.close()
        return path

    def _v2_workbook(
        self,
        *,
        part: str = "灵玉",
        properties_rows: list[tuple[object, ...]] | None = None,
        segments_rows: list[tuple[object, ...]] | None = None,
    ) -> Path:
        path = self.root / "native-v2.xlsx"
        wb = Workbook()
        definitions = wb.active
        definitions.title = "候选定义"
        definitions.append(V2_DEFINITION_HEADERS)
        definitions.append(("v2-two-properties", True, "鞭尸灵玉", part, "铁壁命格", 255, "", "填写唯一英文候选ID与具体中文部位"))
        properties = wb.create_sheet("自定义属性")
        properties.append(V2_PROPERTY_HEADERS)
        for row in properties_rows or [
            ("v2-two-properties", 1, "防御", 5, True, "直接属性", "只填中文属性和值"),
            ("v2-two-properties", 2, "神力倍攻", 10, True, "依赖09号脚本", "只填中文属性和值"),
        ]:
            properties.append(row)
        segments = wb.create_sheet("显示片段")
        segments.append(V2_SEGMENT_HEADERS)
        for row in segments_rows or [
            ("v2-two-properties", 1, "命格名称", None, None, 249, True, "", "名称由候选定义派生"),
            ("v2-two-properties", 2, "属性名称", "防御", None, 250, True, "", "一组必须同源"),
            ("v2-two-properties", 3, "正负号", "防御", None, 69, True, "", "一组必须同源"),
            ("v2-two-properties", 4, "属性值", "防御", None, 69, True, "", "一组必须同源"),
            ("v2-two-properties", 5, "属性名称", "神力倍攻", None, 251, True, "", "一组必须同源"),
            ("v2-two-properties", 6, "正负号", "神力倍攻", None, 251, True, "", "一组必须同源"),
            ("v2-two-properties", 7, "属性值", "神力倍攻", None, 251, True, "", "一组必须同源"),
            ("v2-two-properties", 8, "单位", "神力倍攻", None, 251, True, "", "一组必须同源"),
        ]:
            segments.append(row)
        part_help = wb.create_sheet("部位说明")
        part_help.append(V2_PART_HELP_HEADERS)
        part_help.append(("灵玉", "支持", "固定映射为灵玉单槽"))
        property_help = wb.create_sheet("属性说明")
        property_help.append(V2_PROPERTY_HELP_HEADERS)
        property_help.append(("防御", "引擎直接属性（已游戏实测）", "", "绑定1"))
        property_help.append(("神力倍攻", "依赖09号QFunction/QManage实效出口", "%", "绑定40仅负责显示"))
        wb.save(path)
        wb.close()
        return path

    def _simple_workbook(self) -> Path:
        path = self.root / "命格简表.xlsx"
        wb = Workbook()
        config = wb.active
        config.title = "命格配置"
        config.append(SIMPLE_CONFIG_HEADERS)
        config.append((
            "是", "鞭尸灵玉", "灵玉", "【赤焰命格·测试】", "整条红色",
            "防御", 5, None, None, "魔法值", 100,
        ))
        dictionary = wb.create_sheet("数据字典")
        dictionary.append(SIMPLE_DICTIONARY_HEADERS)
        dictionary.append(("防御", "", "灵玉", "整条红色"))
        dictionary.append(("魔法值", "", None, "原生彩色"))
        wb.save(path)
        wb.close()
        return path

    def _v3_workbook_with_stale_display_segments(self) -> Path:
        path = self.root / "旧六页表.xlsx"
        wb = Workbook()
        definitions = wb.active
        definitions.title = "候选定义"
        definitions.append(mingge_native_module.V3_DEFINITION_HEADERS)
        definitions.append((
            "uniform-red", "是", "鞭尸灵玉", "灵玉",
            "【赤焰命格·测试】", "red", "", "",
        ))
        properties = wb.create_sheet("自定义属性")
        properties.append(mingge_native_module.V3_PROPERTY_HEADERS)
        properties.append(("uniform-red", 1, "对怪伤害吸收", 5, "是", "", ""))
        segments = wb.create_sheet("显示片段")
        segments.append(mingge_native_module.V3_SEGMENT_HEADERS)
        segments.append(("uniform-red", 1, "命格名称", None, None, "是", "", ""))
        segments.append(("uniform-red", 2, "属性名称", "防御", None, "是", "", ""))
        segments.append(("uniform-red", 3, "正负号", "防御", None, "是", "", ""))
        segments.append(("uniform-red", 4, "属性值", "防御", None, "是", "", ""))
        colors = wb.create_sheet("颜色方案")
        colors.append(mingge_native_module.V3_COLOR_HEADERS)
        colors.append(("red", "整条红色", "uniform", 249, None, None, None, None, None, None, "", ""))
        parts = wb.create_sheet("部位说明")
        parts.append(mingge_native_module.V3_PART_HELP_HEADERS)
        parts.append(("灵玉", "支持", ""))
        property_help = wb.create_sheet("属性说明")
        property_help.append(mingge_native_module.V3_PROPERTY_HELP_HEADERS)
        property_help.append(("对怪伤害吸收", "已验证", "%", ""))
        wb.save(path)
        wb.close()
        return path

    def test_compiles_verified_four_color_literal_and_real_property_separately(self) -> None:
        book = read_native_candidate_workbook(self._workbook())
        self.assertEqual(len(book.candidates), 1)
        compiled = compile_native_candidate(book.candidates[0])
        expected = (
            "LockUpdateItem 17\r\n"
            "SetCustomItemAbil 17 17 0 250\r\n"
            "SetCustomItemAbil 17 17 1 1\r\n"
            "SetCustomItemAbil 17 17 2 17\r\n"
            "SetCustomItemAbil 17 17 3 0\r\n"
            "SetCustomItemAbil 17 17 4 9\r\n"
            "SetCustomItemValueEx 17 17 = 5 0 0\r\n"
            "SetCustomItemText 17 {红|249}{橙|69}{绿|250}{彩|251}\r\n"
            "SetCustomItemTextColor 17 255\r\n"
            "UpdateItem 17\r\n"
        )
        self.assertEqual(compiled.script, expected)
        self.assertEqual(compiled.payload, expected.encode("gb18030"))
        self.assertEqual(compiled.display_text, "红橙绿彩")
        self.assertNotIn("防御+5", compiled.display_text)

    def test_v2_accepts_two_complete_same_source_display_groups_without_user_supplied_codes(self) -> None:
        book = read_native_candidate_workbook(self._v2_workbook())
        candidate = book.candidates[0]
        self.assertEqual(candidate.equip_slot, 17)
        self.assertEqual(
            [(item.display_name, item.binding, item.property_row) for item in candidate.properties],
            [("防御", 1, 17), ("神力倍攻", 40, 18)],
        )
        compiled = compile_native_candidate(candidate)
        self.assertIn("SetCustomItemAbil 17 17 1 1", compiled.script)
        self.assertIn("SetCustomItemAbil 17 18 1 40", compiled.script)
        self.assertIn("SetCustomItemValue 17 18 = 10", compiled.script)
        self.assertNotIn("SetCustomItemValueEx 17 18", compiled.script)
        self.assertEqual(compiled.display_text, "铁壁命格防御+5神力倍攻+10%")
        receipt = generate_native_candidates(
            self._v2_workbook(), self.root / "candidate", allowed_output_root=self.root
        )
        evidence = json.loads((Path(receipt.output) / "v2-two-properties.json").read_text(encoding="utf-8"))
        self.assertEqual(evidence["schema_version"], 2)
        self.assertEqual(evidence["v2_properties"][0]["实效状态"], "引擎直接属性（已游戏实测）")
        self.assertEqual(evidence["v2_properties"][1]["依赖"], "09号QFunction/QManage实效出口")

    def test_v2_editable_sheet_headers_end_with_exact_usage_instruction_column(self) -> None:
        self.assertEqual(mingge_native_module.V2_DEFINITION_HEADERS[-1], "使用说明")
        self.assertEqual(mingge_native_module.V2_PROPERTY_HEADERS[-1], "使用说明")
        self.assertEqual(mingge_native_module.V2_SEGMENT_HEADERS[-1], "使用说明")

    def test_v2_blocks_cross_property_display_group(self) -> None:
        path = self._v2_workbook(segments_rows=[
            ("v2-two-properties", 1, "命格名称", None, None, 249, True, "", ""),
            ("v2-two-properties", 2, "属性名称", "防御", None, 250, True, "", ""),
            ("v2-two-properties", 3, "正负号", "防御", None, 69, True, "", ""),
            ("v2-two-properties", 4, "属性值", "防御", None, 69, True, "", ""),
            ("v2-two-properties", 5, "单位", "神力倍攻", None, 251, True, "", ""),
        ])
        with self.assertRaisesRegex(NativeCandidateError, "显示组.*同一来源"):
            read_native_candidate_workbook(path)

    def test_v2_allocates_only_rows_17_to_19_and_blocks_a_fourth_property(self) -> None:
        properties_rows = [
            ("v2-two-properties", 1, "防御", 5, True, "", ""),
            ("v2-two-properties", 2, "攻击", 6, True, "", ""),
            ("v2-two-properties", 3, "魔法", 7, True, "", ""),
            ("v2-two-properties", 4, "道术", 8, True, "", ""),
        ]
        segments_rows = [
            ("v2-two-properties", 1, "命格名称", None, None, 249, True, "", ""),
            ("v2-two-properties", 2, "属性名称", "防御", None, 250, True, "", ""),
            ("v2-two-properties", 3, "正负号", "防御", None, 69, True, "", ""),
            ("v2-two-properties", 4, "属性值", "防御", None, 69, True, "", ""),
        ]
        three = read_native_candidate_workbook(
            self._v2_workbook(properties_rows=properties_rows[:3], segments_rows=segments_rows)
        ).candidates[0]
        self.assertEqual([item.property_row for item in three.properties], [17, 18, 19])
        self.assertEqual(three.properties[0].effect_status, "引擎直接属性（已游戏实测）")
        self.assertEqual(three.properties[1].effect_status, "官方扩展口径静态候选待实测")
        self.assertEqual(three.properties[2].effect_status, "官方扩展口径静态候选待实测")
        with self.assertRaisesRegex(NativeCandidateError, "最多3条|17至19"):
            read_native_candidate_workbook(self._v2_workbook(properties_rows=properties_rows, segments_rows=segments_rows))

    def test_v2_marks_defence_at_row18_static_when_attack_precedes_it(self) -> None:
        candidate = read_native_candidate_workbook(self._v2_workbook(
            properties_rows=[
                ("v2-two-properties", 1, "攻击", 6, True, "", ""),
                ("v2-two-properties", 2, "防御", 5, True, "", ""),
            ],
            segments_rows=[
                ("v2-two-properties", 1, "命格名称", None, None, 249, True, "", ""),
                ("v2-two-properties", 2, "属性名称", "攻击", None, 250, True, "", ""),
                ("v2-two-properties", 3, "正负号", "攻击", None, 69, True, "", ""),
                ("v2-two-properties", 4, "属性值", "攻击", None, 69, True, "", ""),
                ("v2-two-properties", 5, "属性名称", "防御", None, 250, True, "", ""),
                ("v2-two-properties", 6, "正负号", "防御", None, 69, True, "", ""),
                ("v2-two-properties", 7, "属性值", "防御", None, 69, True, "", ""),
            ],
        )).candidates[0]
        self.assertEqual(
            [(item.display_name, item.property_row, item.effect_status) for item in candidate.properties],
            [
                ("攻击", 17, "官方扩展口径静态候选待实测"),
                ("防御", 18, "官方扩展口径静态候选待实测"),
            ],
        )

    def test_v2_marks_defence_at_row19_static_when_attack_and_magic_precede_it(self) -> None:
        candidate = read_native_candidate_workbook(self._v2_workbook(
            properties_rows=[
                ("v2-two-properties", 1, "攻击", 6, True, "", ""),
                ("v2-two-properties", 2, "魔法", 7, True, "", ""),
                ("v2-two-properties", 3, "防御", 5, True, "", ""),
            ],
            segments_rows=[
                ("v2-two-properties", 1, "命格名称", None, None, 249, True, "", ""),
                ("v2-two-properties", 2, "属性名称", "攻击", None, 250, True, "", ""),
                ("v2-two-properties", 3, "正负号", "攻击", None, 69, True, "", ""),
                ("v2-two-properties", 4, "属性值", "攻击", None, 69, True, "", ""),
                ("v2-two-properties", 5, "属性名称", "魔法", None, 250, True, "", ""),
                ("v2-two-properties", 6, "正负号", "魔法", None, 69, True, "", ""),
                ("v2-two-properties", 7, "属性值", "魔法", None, 69, True, "", ""),
                ("v2-two-properties", 8, "属性名称", "防御", None, 250, True, "", ""),
                ("v2-two-properties", 9, "正负号", "防御", None, 69, True, "", ""),
                ("v2-two-properties", 10, "属性值", "防御", None, 69, True, "", ""),
            ],
        )).candidates[0]
        self.assertEqual(
            [(item.display_name, item.property_row, item.effect_status) for item in candidate.properties],
            [
                ("攻击", 17, "官方扩展口径静态候选待实测"),
                ("魔法", 18, "官方扩展口径静态候选待实测"),
                ("防御", 19, "官方扩展口径静态候选待实测"),
            ],
        )

    def test_v2_blocks_all_same_binding_alias_pairs(self) -> None:
        pairs = (
            ("防御", "自身防御"), ("攻击", "自身攻击"), ("魔法", "自身魔法"),
            ("道术", "自身道术"), ("HP", "生命值"), ("MP", "魔法值"),
            ("鞭尸", "鞭尸概率"),
        )
        for first, second in pairs:
            with self.subTest(first=first, second=second):
                properties_rows = [
                    ("v2-two-properties", 1, first, 5, True, "", ""),
                    ("v2-two-properties", 2, second, 6, True, "", ""),
                ]
                segments_rows = [
                    ("v2-two-properties", 1, "命格名称", None, None, 249, True, "", ""),
                    ("v2-two-properties", 2, "属性名称", first, None, 250, True, "", ""),
                    ("v2-two-properties", 3, "正负号", first, None, 69, True, "", ""),
                    ("v2-two-properties", 4, "属性值", first, None, 69, True, "", ""),
                ]
                with self.assertRaisesRegex(NativeCandidateError, "规范属性|绑定"):
                    read_native_candidate_workbook(
                        self._v2_workbook(properties_rows=properties_rows, segments_rows=segments_rows)
                    )

    def test_v2_covers_every_09_single_value_part_alias_and_excludes_ranges(self) -> None:
        expected = {
            "武器": 1, "衣服": 0, "男衣服": 0, "女衣服": 0, "盔甲": 0,
            "照明物": 2, "勋章": 2, "项链": 3, "头盔": 4, "左手镯": 5,
            "右手镯": 6, "左戒指": 7, "右戒指": 8, "护身符": 9, "护符": 9,
            "腰带": 10, "鞋子": 11, "靴子": 11, "宝石": 12, "斗笠": 13,
            "面巾": 13, "军鼓": 14, "马牌": 15, "盾牌": 16, "灵玉": 17,
            "时装衣服": 18, "时装男衣服": 18, "时装女衣服": 18, "时装盔甲": 18,
            "时装武器": 19, "时装项链": 20, "时装头盔": 21, "时装左手镯": 22,
            "时装右手镯": 23, "时装左戒指": 24, "时装右戒指": 25, "时装勋章": 26,
            "时装腰带": 27, "时装鞋子": 28, "时装靴子": 28, "时装宝石": 29,
        }
        for number in range(1, 7):
            expected[f"首饰盒{number}"] = 29 + number
            expected[f"普通首饰盒{number}"] = 29 + number
            expected[f"时装首饰盒{number}"] = 69 + number
        for number in range(1, 13):
            expected[f"生肖{number}"] = 39 + number
            expected[f"生肖盒{number}"] = 39 + number
            expected[f"普通生肖{number}"] = 39 + number
            expected[f"普通生肖盒{number}"] = 39 + number
            expected[f"时装生肖{number}"] = 79 + number
            expected[f"时装生肖盒{number}"] = 79 + number
        self.assertEqual(TRUSTED_EQUIP_SLOTS, expected)
        for forbidden in ("手镯", "戒指", "时装手镯", "时装戒指", "首饰盒", "生肖", "背包神器", "称号卷"):
            self.assertNotIn(forbidden, TRUSTED_EQUIP_SLOTS)

    def test_v2_generation_receipt_summarizes_evidence_dependencies_and_zero_writes(self) -> None:
        receipt = generate_native_candidates(
            self._v2_workbook(), self.root / "candidate", allowed_output_root=self.root
        )
        data = json.loads((Path(receipt.output) / "generation-receipt.json").read_text(encoding="utf-8"))
        self.assertEqual(data["schema_version"], 2)
        self.assertEqual(data["candidate_evidence_level"], "candidate")
        self.assertEqual(data["dependencies"], ["09号QFunction/QManage实效出口"])
        self.assertFalse(data["writes_server"])
        self.assertFalse(data["writes_client"])
        self.assertFalse(data["writes_database"])
        self.assertFalse(data["image_gradient_pipeline_called"])

    def test_v1_keeps_nondefault_position_and_all_three_values(self) -> None:
        definitions = [
            (
                "v1-nondefault", True, "鞭尸灵玉", 17, True, 17, 250, 1, 19,
                0, 9, 5, 6, 7, 255, "V1兼容回归",
            )
        ]
        segments = [("v1-nondefault", 1, "固定文字", "甲", 249, True, "")]
        candidate = read_native_candidate_workbook(
            self._workbook(definitions=definitions, segments=segments)
        ).candidates[0]
        compiled = compile_native_candidate(candidate)
        self.assertIn("SetCustomItemAbil 17 17 2 19", compiled.script)
        self.assertIn("SetCustomItemValueEx 17 17 = 5 6 7", compiled.script)

    def test_v2_reuses_all_16_frozen_09_display_contracts(self) -> None:
        expected = (
            ("神力倍攻", 40, True, "%"), ("打怪伤害", 41, True, "%"), ("暴击伤害", 42, True, "%"),
            ("固定切割", 43, False, ""), ("爆率", 44, True, "%"), ("最大爆率", 45, True, "%"),
            ("首刀斩杀", 46, True, "%"), ("尾刀斩杀", 47, True, "%"), ("鞭尸概率", 48, True, "%"),
            ("处决概率", 49, True, "%"), ("韧性", 50, False, ""), ("处决倍率", 51, True, "%"),
            ("处决时间", 52, False, "秒"), ("伤害系数", 53, True, "%"), ("吸血", 54, True, "%"),
            ("每秒回血", 55, False, ""),
        )
        for name, binding, percent, unit in expected:
            with self.subTest(name=name):
                properties_rows = [("v2-two-properties", 1, name, 10, True, "", "")]
                segments_rows = [
                    ("v2-two-properties", 1, "命格名称", None, None, 249, True, "", ""),
                    ("v2-two-properties", 2, "属性名称", name, None, 251, True, "", ""),
                    ("v2-two-properties", 3, "正负号", name, None, 251, True, "", ""),
                    ("v2-two-properties", 4, "属性值", name, None, 251, True, "", ""),
                ]
                if unit:
                    segments_rows.append(("v2-two-properties", 5, "单位", name, None, 251, True, "", ""))
                compiled = compile_native_candidate(read_native_candidate_workbook(
                    self._v2_workbook(properties_rows=properties_rows, segments_rows=segments_rows)
                ).candidates[0])
                self.assertIn(f"SetCustomItemAbil 17 17 0 251", compiled.script)
                self.assertIn(f"SetCustomItemAbil 17 17 1 {binding}", compiled.script)
                self.assertIn(f"SetCustomItemAbil 17 17 3 {int(percent)}", compiled.script)
                self.assertIn("SetCustomItemValue 17 17 = 10", compiled.script)
                self.assertNotIn("SetCustomItemAbil 17 17 4", compiled.script)
                self.assertNotIn("SetCustomItemValueEx 17 17", compiled.script)

    def test_v2_blocks_ambiguous_parts_and_unsupported_properties(self) -> None:
        with self.assertRaisesRegex(NativeCandidateError, "多槽歧义"):
            read_native_candidate_workbook(self._v2_workbook(part="手镯"))

    def test_all_39_catalog_properties_have_explicit_route_metadata(self) -> None:
        from openpyxl import load_workbook

        source = Path(r"E:\XuanYuanDevPlatform\所需材料表格汇总\39_命格P3颜色导入测试.xlsx")
        wb = load_workbook(source, read_only=True, data_only=True)
        try:
            dictionary_sheet = "数据字典" if "数据字典" in wb.sheetnames else "属性说明"
            names = [row[0].value for row in wb[dictionary_sheet].iter_rows(min_row=2) if row[0].value]
        finally:
            wb.close()
        self.assertEqual(len(names), 61)
        self.assertIn("攻速突破", names)
        specs = [mingge_native_module._v2_property_spec(name, index + 2) for index, name in enumerate(names)]
        self.assertEqual(
            {spec.route for spec in specs},
            {"direct", "script_bind", "textline_runtime"},
        )
        self.assertTrue(all(spec.route in {"direct", "script_bind", "textline_runtime"} for spec in specs))

    def test_attack_speed_breakthrough_reuses_existing_textline40_interface_without_local_threshold(self) -> None:
        spec = mingge_native_module._v2_property_spec("攻速突破", 2, instance_routes=True)

        self.assertEqual(spec.route, "textline_runtime")
        self.assertEqual(spec.binding, 60)
        self.assertEqual(spec.text_line, 40)
        self.assertEqual(spec.value_command, "value_ex")

        path = self._simple_workbook()
        from openpyxl import load_workbook

        workbook = load_workbook(path)
        config = workbook["命格配置"]
        config.cell(2, 4).value = "攻速命格"
        config.cell(2, 6).value = "攻速突破"
        config.cell(2, 7).value = 99
        config.cell(2, 10).value = None
        config.cell(2, 11).value = None
        workbook["数据字典"].append(("攻速突破", "", None, None))
        workbook.save(path)
        workbook.close()
        candidate = read_native_candidate_workbook(path).candidates[0]
        self.assertEqual(candidate.properties[0].value, 99)
        script = compile_native_candidate(candidate).script
        self.assertIn("SetCustomItemAbil 17 17 1 60", script)
        self.assertIn("SetCustomItemAbil 17 17 2 40", script)
        self.assertIn("SetCustomItemValueEx 17 17 = 40 99 0", script)
        self.assertEqual(
            script.count("#CALL [\\玄渊攻速突破\\全身攻速阈值核心.txt] @XY_AS_CAP_RECALC"),
            1,
        )
        self.assertNotIn("ChangeSpeed 2", script)

    def test_current_39_six_candidates_derive_segments_and_runtime_routes(self) -> None:
        source = Path(r"E:\XuanYuanDevPlatform\所需材料表格汇总\39_命格P3颜色导入测试.xlsx")
        book = read_native_candidate_workbook(source, derive_v3_segments=True)
        self.assertEqual(len(book.candidates), 6)
        self.assertEqual(
            [(item.display_name, item.route, item.text_line) for item in book.candidates[0].properties],
            [("对怪伤害吸收", "textline_runtime", 33), ("道术加成", "textline_runtime", 34), ("伤害系数", "script_bind", None)],
        )
        self.assertEqual(
            [(item.display_name, item.route, item.text_line) for item in book.candidates[3].properties],
            [("致命伤害", "textline_runtime", 35), ("处决概率", "textline_runtime", 36), ("鞭尸", "textline_runtime", 37)],
        )
        for candidate in book.candidates:
            derived_names = [segment.source_property for segment in candidate.segments if segment.role == "属性名称"]
            self.assertEqual(derived_names, [item.display_name for item in candidate.properties])

    def test_simple_chinese_workbook_derives_internal_fields_and_compacts_property_gaps(self) -> None:
        book = read_native_candidate_workbook(self._simple_workbook())
        self.assertEqual(len(book.candidates), 1)
        candidate = book.candidates[0]
        self.assertEqual(candidate.candidate_id, "candidate-0002")
        self.assertEqual(candidate.target_item_name, "鞭尸灵玉")
        self.assertEqual(candidate.equip_slot, 17)
        self.assertEqual(
            [(item.order, item.display_name, item.value) for item in candidate.properties],
            [(1, "防御", 5), (2, "魔法值", 100)],
        )
        self.assertEqual(
            [segment.source_property for segment in candidate.segments if segment.role == "属性名称"],
            ["防御", "魔法值"],
        )
        self.assertEqual({segment.color for segment in candidate.segments}, {249})

    def test_offline_generation_auto_derives_v3_segments_when_legacy_sheet_is_stale(self) -> None:
        allowed = self.root / "outputs"
        receipt = generate_native_candidates(
            self._v3_workbook_with_stale_display_segments(),
            allowed / "命格候选",
            allowed_output_root=allowed,
        )
        self.assertEqual(receipt.candidate_ids, ("uniform-red",))
        published = Path(receipt.output)
        script = (published / "uniform-red.txt").read_text(encoding="gb18030")
        self.assertIn("对怪伤害吸收", script)
        self.assertNotIn("防御", script)

    def test_disabled_candidate_and_disabled_segment_are_not_emitted(self) -> None:
        definitions = [
            (
                "enabled",
                True,
                "鞭尸灵玉",
                17,
                False,
                17,
                250,
                1,
                17,
                0,
                9,
                5,
                0,
                0,
                255,
                "",
            ),
            (
                "disabled",
                False,
                "鞭尸灵玉",
                17,
                True,
                17,
                250,
                1,
                17,
                0,
                9,
                5,
                0,
                0,
                255,
                "",
            ),
        ]
        segments = [
            ("enabled", 1, "命格名称", "【铁壁命格·测试】", 249, True, ""),
            ("enabled", 2, "固定文字", "+5", 69, False, ""),
            ("disabled", 1, "命格名称", "不应生成", 249, True, ""),
        ]
        book = read_native_candidate_workbook(
            self._workbook(definitions=definitions, segments=segments)
        )
        self.assertEqual([item.candidate_id for item in book.candidates], ["enabled"])
        compiled = compile_native_candidate(book.candidates[0])
        self.assertIn("{【铁壁命格·测试】|249}", compiled.script)
        self.assertNotIn("+5", compiled.script)
        self.assertNotIn("SetCustomItemAbil", compiled.script)

    def test_multiple_candidates_are_published_with_evidence(self) -> None:
        definitions = []
        segments = []
        for candidate_id, value, color in (("defence-5", 5, 249), ("attack-8", 8, 69)):
            definitions.append(
                (
                    candidate_id,
                    True,
                    "鞭尸灵玉",
                    17,
                    True,
                    17,
                    250,
                    1,
                    17,
                    0,
                    9,
                    value,
                    0,
                    0,
                    255,
                    "",
                )
            )
            segments.append((candidate_id, 1, "数值", None, color, True, ""))
        workbook = self._workbook(definitions=definitions, segments=segments)
        output = self.root / "candidate"
        receipt = generate_native_candidates(
            workbook, output, allowed_output_root=self.root
        )
        published = Path(receipt.output)
        self.assertEqual(receipt.backend, "native_segments")
        self.assertEqual(receipt.status, "candidate")
        self.assertEqual(receipt.candidate_ids, ("defence-5", "attack-8"))
        for candidate_id in receipt.candidate_ids:
            text_path = published / f"{candidate_id}.txt"
            json_path = published / f"{candidate_id}.json"
            self.assertTrue(text_path.is_file())
            self.assertTrue(json_path.is_file())
            raw = text_path.read_bytes()
            self.assertNotRegex(raw, rb"(?<!\r)\n|\r(?!\n)")
            evidence = json.loads(json_path.read_text(encoding="utf-8"))
            self.assertEqual(evidence["backend"], "native_segments")
            self.assertEqual(evidence["status"], "candidate")
            self.assertEqual(evidence["script_sha256"], hashlib.sha256(raw).hexdigest())
            self.assertFalse(evidence["writes_server"])
            self.assertFalse(evidence["writes_client"])
            self.assertFalse(evidence["writes_database"])

    def test_publish_uses_independent_atomic_batches_and_preserves_unrelated_files(self) -> None:
        workbook = self._workbook()
        output = self.root / "candidate"
        output.mkdir()
        unrelated = output / "keep.me"
        unrelated.write_text("user", encoding="utf-8")
        first = generate_native_candidates(
            workbook, output, allowed_output_root=self.root
        )
        first_dir = Path(first.output)
        first_script = (first_dir / "four-colors.txt").read_bytes()
        second = generate_native_candidates(
            workbook, output, allowed_output_root=self.root
        )
        second_dir = Path(second.output)
        self.assertNotEqual(first_dir, second_dir)
        self.assertEqual(first_script, (second_dir / "four-colors.txt").read_bytes())
        self.assertEqual(first.workbook_sha256, second.workbook_sha256)
        self.assertEqual(unrelated.read_text(encoding="utf-8"), "user")

    def test_second_batch_cannot_mix_removed_candidates_from_first_batch(self) -> None:
        definitions = []
        segments = []
        for candidate_id in ("alpha", "beta"):
            definitions.append(
                (
                    candidate_id, True, "鞭尸灵玉", 17, False, 17, 250, 1,
                    17, 0, 9, 5, 0, 0, 255, "",
                )
            )
            segments.append((candidate_id, 1, "固定文字", candidate_id, 249, True, ""))
        first = generate_native_candidates(
            self._workbook(definitions=definitions, segments=segments),
            self.root / "candidate",
            allowed_output_root=self.root,
        )
        second = generate_native_candidates(
            self._workbook(definitions=definitions[:1], segments=segments[:1]),
            self.root / "candidate",
            allowed_output_root=self.root,
        )
        self.assertTrue((Path(first.output) / "beta.txt").is_file())
        self.assertFalse((Path(second.output) / "beta.txt").exists())
        self.assertEqual(second.candidate_ids, ("alpha",))

    def test_numeric_segment_is_derived_from_real_property_value(self) -> None:
        definitions = [
            (
                "derived-value",
                True,
                "鞭尸灵玉",
                17,
                True,
                17,
                250,
                1,
                17,
                0,
                9,
                8,
                0,
                0,
                255,
                "",
            )
        ]
        segments = [
            ("derived-value", 1, "命格名称", "【铁壁命格】", 249, True, ""),
            ("derived-value", 2, "数值", None, 69, True, ""),
        ]
        book = read_native_candidate_workbook(
            self._workbook(definitions=definitions, segments=segments)
        )
        compiled = compile_native_candidate(book.candidates[0])
        self.assertIn("{【铁壁命格】|249}{+8|69}", compiled.script)
        self.assertIn("SetCustomItemValueEx 17 17 = 8 0 0", compiled.script)
        self.assertEqual(compiled.display_text, "【铁壁命格】+8")

    def test_numeric_segment_rejects_a_second_manually_typed_value(self) -> None:
        segments = [("four-colors", 1, "数值", "+99", 249, True, "")]
        with self.assertRaisesRegex(NativeCandidateError, "数值片段.*留空"):
            read_native_candidate_workbook(self._workbook(segments=segments))

    def test_cli_generates_candidate_without_server_or_client_arguments(self) -> None:
        workbook = self._workbook()
        runtime_root = self.root / "runtime-platform"
        output = runtime_root / "outputs" / "cli-candidate"
        stdout = io.StringIO()
        with mock.patch("xydp.cli.platform_root", return_value=runtime_root), redirect_stdout(stdout):
            code = cli_main(
                [
                    "--root",
                    str(self.root / "unused-platform-root"),
                    "mingge-native-candidate",
                    "--input",
                    str(workbook),
                    "--output",
                    str(output),
                ]
            )
        self.assertEqual(code, 0)
        summary = json.loads(stdout.getvalue())
        self.assertEqual(summary["backend"], "native_segments")
        self.assertEqual(summary["status"], "candidate")
        self.assertNotIn("server", summary)
        self.assertNotIn("client", summary)
        self.assertTrue((Path(summary["output"]) / "four-colors.txt").is_file())

    def test_rejects_output_outside_platform_output_root_before_writing(self) -> None:
        workbook = self._workbook()
        allowed = self.root / "platform" / "outputs"
        forbidden = self.root / "pretend-server"
        with self.assertRaisesRegex(NativeCandidateError, "平台outputs"):
            generate_native_candidates(
                workbook, forbidden, allowed_output_root=allowed
            )
        self.assertFalse(forbidden.exists())

    def test_rejects_resolved_directory_link_escape(self) -> None:
        workbook = self._workbook()
        allowed = self.root / "platform" / "outputs"
        outside = self.root / "outside"
        allowed.mkdir(parents=True)
        outside.mkdir()
        link = allowed / "linked"
        try:
            link.symlink_to(outside, target_is_directory=True)
        except OSError as exc:
            self.skipTest(f"当前环境不能创建目录链接：{exc}")
        with self.assertRaisesRegex(NativeCandidateError, "平台outputs"):
            generate_native_candidates(
                workbook, link / "candidate", allowed_output_root=allowed
            )
        self.assertEqual(list(outside.iterdir()), [])

    def test_rejects_allowed_outputs_root_when_it_is_a_directory_link(self) -> None:
        workbook = self._workbook()
        platform = self.root / "platform"
        outside = self.root / "outside"
        platform.mkdir()
        outside.mkdir()
        allowed_link = platform / "outputs"
        try:
            allowed_link.symlink_to(outside, target_is_directory=True)
        except OSError as exc:
            self.skipTest(f"当前环境不能创建目录链接：{exc}")
        with self.assertRaisesRegex(NativeCandidateError, "链接|联接"):
            generate_native_candidates(
                workbook,
                allowed_link / "candidate",
                allowed_output_root=allowed_link,
            )
        self.assertEqual(list(outside.iterdir()), [])

    def test_rejects_non_contiguous_order(self) -> None:
        segments = [
            ("four-colors", 1, "固定文字", "红", 249, True, ""),
            ("four-colors", 3, "固定文字", "绿", 250, True, ""),
        ]
        with self.assertRaisesRegex(NativeCandidateError, "连续"):
            read_native_candidate_workbook(self._workbook(segments=segments))

    def test_rejects_bad_colors_bool_numbers_and_default_not_255(self) -> None:
        bad_definitions = [
            (
                "four-colors",
                True,
                "鞭尸灵玉",
                17,
                True,
                17,
                250,
                1,
                17,
                0,
                9,
                5,
                0,
                0,
                250,
                "",
            )
        ]
        with self.assertRaisesRegex(NativeCandidateError, "默认文字颜色.*255"):
            read_native_candidate_workbook(self._workbook(definitions=bad_definitions))
        for bad in (True, -1, 256, 1.5):
            with self.subTest(bad=bad):
                segments = [("four-colors", 1, "固定文字", "红", bad, True, "")]
                with self.assertRaisesRegex(NativeCandidateError, "颜色"):
                    read_native_candidate_workbook(self._workbook(segments=segments))

    def test_rejects_control_characters_and_script_variables(self) -> None:
        for text in ("坏|字", "坏{字", "坏}字", "坏\\字", "坏^字", "坏<字", "坏>字", "坏\t字", "坏\n字", "$$1", " 坏", "坏 "):
            with self.subTest(text=text):
                segments = [("four-colors", 1, "固定文字", text, 249, True, "")]
                with self.assertRaisesRegex(NativeCandidateError, "片段文字"):
                    read_native_candidate_workbook(self._workbook(segments=segments))

    def test_rejects_casefold_collisions_reserved_names_and_trailing_dot(self) -> None:
        base = (
            True, "鞭尸灵玉", 17, False, 17, 250, 1, 17, 0, 9,
            5, 0, 0, 255, "",
        )
        definitions = [("same", *base), ("SAME", *base)]
        segments = [
            ("same", 1, "固定文字", "甲", 249, True, ""),
            ("SAME", 1, "固定文字", "乙", 250, True, ""),
        ]
        with self.assertRaisesRegex(NativeCandidateError, "大小写"):
            read_native_candidate_workbook(
                self._workbook(definitions=definitions, segments=segments)
            )
        for candidate_id in ("CON", "nul.txt", "name."):
            with self.subTest(candidate_id=candidate_id):
                definitions = [(candidate_id, *base)]
                segments = [(candidate_id, 1, "固定文字", "甲", 249, True, "")]
                with self.assertRaisesRegex(NativeCandidateError, "Windows"):
                    read_native_candidate_workbook(
                        self._workbook(definitions=definitions, segments=segments)
                    )

    def test_rejects_unknown_sheet_or_header_drift(self) -> None:
        path = self._workbook()
        from openpyxl import load_workbook

        wb = load_workbook(path)
        wb.create_sheet("意外工作表")
        wb.save(path)
        wb.close()
        with self.assertRaisesRegex(NativeCandidateError, "中文简表两页结构"):
            read_native_candidate_workbook(path)


if __name__ == "__main__":
    unittest.main()
