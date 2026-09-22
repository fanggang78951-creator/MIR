from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

import xydp.mingge_dual as mingge_dual_module

from xydp.mingge_dual import (
    BRIDGE_RELATIVE,
    CONTENT_PROVIDER_RELATIVE,
    CONTENT_QFUNCTION_MARKER,
    CONTENT_TEXTVAR_START,
    NPC_RULES_RELATIVE,
    TEXTVAR_RELATIVE,
    MinggeContentService,
    MinggeNpcService,
    load_mingge_content_workbook,
    load_mingge_npc_workbook,
)


def _save_npc_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "系统设置"
    ws.append(["配置项", "配置值", "说明"])
    for row in (
        ("系统编号", "MGNPC001", "稳定系统编号"),
        ("目标装备", "鞭尸灵玉", "只允许准确名称"),
        ("装备部位", "灵玉", "中文部位"),
        ("装备位置", 17, "平台可由部位校验"),
        ("槽位数量", 8, "1至8"),
        ("NPC脚本相对路径", "Mir200/Envir/Market_Def/玄渊命格/龙魂觉醒-3.txt", "只接入现有NPC"),
        ("NPC接入标签", "main", "不带方括号和@"),
        ("默认候选池ID", 1, "内容包候选池"),
        ("内容包缺失提示", "命格内容包尚未安装，本次不会扣费。", "安全提示"),
    ):
        ws.append(row)

    ws = wb.create_sheet("槽位开放")
    ws.append(["槽位", "启用", "显示名称", "条件类型", "条件运算", "条件值", "消耗类型", "消耗名称", "消耗数量", "失败提示"])
    for slot, level in enumerate((1, 10, 20, 30, 40, 50, 60, 70), start=1):
        ws.append([slot, "是", f"命格槽位{slot}", "等级", ">=", level, "无", "", 0, f"需要等级{level}"])

    ws = wb.create_sheet("洗练规则")
    ws.append(["规则ID", "启用", "名称", "次数", "消耗类型", "消耗名称", "消耗数量", "前置条件类型", "前置条件值", "候选池ID", "失败提示"])
    ws.append([1, "是", "单次洗练", 1, "元宝", "元宝", 100, "已选择槽位", 1, 1, "请先选择槽位"])
    ws.append([2, "是", "十连洗练", 10, "元宝", "元宝", 1000, "已选择槽位", 1, 1, "请先选择槽位"])

    ws = wb.create_sheet("品质保底")
    ws.append(["品质ID", "品质名称", "基础权重", "连续未出保底次数", "固定周期次数", "命中后清零", "说明"])
    ws.append([1, "绿色", 50, 0, 0, "是", "普通品质"])
    ws.append([2, "蓝色", 30, 0, 0, "是", "普通品质"])
    ws.append([3, "红色", 15, 20, 0, "是", "20次红保底"])
    ws.append([4, "彩色", 5, 100, 10, "是", "每10次参与彩判定，100次保底"])

    ws = wb.create_sheet("命格羁绊")
    ws.append(["羁绊ID", "启用", "羁绊名称", "成员类型", "成员ID列表", "需要数量", "奖励属性", "奖励数值", "显示文字"])
    ws.append([1, "否", "八荒齐聚", "命格ID", "1001,1002,1003,1004", 4, "防御", 10, "集齐四种命格后激活"])

    ws = wb.create_sheet("NPC文案")
    ws.append(["文案ID", "启用", "显示文字", "说明"])
    ws.append(["标题", "是", "龙魂觉醒·命格洗练", "NPC标题"])
    ws.append(["保底说明", "是", "20次必出红色命格，100次必出彩色命格", "由规则自动校验"])
    wb.save(path)


def _save_content_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "命格定义"
    ws.append(["命格ID", "启用", "命格名称", "品质ID", "候选池ID", "适用部位", "颜色方案ID", "备注"])
    ws.append([1001, "是", "铁壁命格", 1, 1, "灵玉", 101, "绿色示例"])
    ws.append([1002, "是", "沧海命格", 2, 1, "灵玉", 102, "蓝色示例"])
    ws.append([1003, "是", "赤霄命格", 3, 1, "灵玉", 103, "红色示例"])
    ws.append([1004, "是", "绮彩命格", 4, 1, "灵玉", 104, "彩色示例"])

    ws = wb.create_sheet("命格属性")
    ws.append(["命格ID", "顺序", "中文属性", "数值", "单位", "属性颜色方案ID", "备注"])
    ws.append([1001, 1, "防御", 5, "", "", "direct"])
    ws.append([1002, 1, "生命值", 100, "", "", "direct"])
    ws.append([1003, 1, "魔法值", 100, "", "", "direct"])
    ws.append([1004, 1, "防御", 5, "", "", "彩色属性"])

    ws = wb.create_sheet("颜色方案")
    ws.append(["方案ID", "中文名称", "类型", "名称颜色序列", "属性颜色序列", "分隔符颜色", "说明"])
    ws.append([101, "整条绿色", "统一色", "250", "250", 255, "绿色"])
    ws.append([102, "整条蓝色", "统一色", "252", "252", 255, "蓝色"])
    ws.append([103, "整条红色", "统一色", "249", "249", 255, "红色"])
    ws.append([104, "彩色浅色属性", "原生多色", "249,69,250,251", "31,147,239", 255, "已验收配色"])

    ws = wb.create_sheet("数据字典")
    ws.append(["类型", "中文名称", "单位", "内部标识", "路由", "绑定", "支持状态", "说明"])
    ws.append(["属性", "防御", "", "defence", "direct", 1, "已游戏实测", "引擎直接属性"])
    ws.append(["属性", "生命值", "", "hp", "direct", 6, "已游戏实测", "引擎直接属性"])
    ws.append(["属性", "魔法值", "", "mp", "direct", 7, "已游戏实测", "引擎直接属性"])
    ws.append(["部位", "灵玉", "", "17", "equip_slot", 17, "已确认", "装备位置"])
    wb.save(path)


def _save_legacy_p3_content_workbook(path: Path) -> None:
    _save_content_workbook(path)
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    definitions = workbook["命格定义"]
    definitions.delete_rows(2, definitions.max_row - 1)
    for candidate_id, name, quality_id, color_id in (
        (1001, "【赤焰命格·测试1】", 3, 103),
        (1002, "【碧落命格·测试2】", 1, 101),
        (1003, "【沧澜命格·测试3】", 2, 102),
        (1004, "【绮彩命格·彩】", 4, 104),
        (1005, "【沧澜命格·红】", 3, 103),
        (1006, "【绮彩命格·测试】", 4, 104),
    ):
        definitions.append([candidate_id, "退役", name, quality_id, 1, "灵玉", color_id, "旧P3候选"])
    attributes = workbook["命格属性"]
    attributes.delete_rows(2, attributes.max_row - 1)
    for candidate_id in range(1001, 1007):
        attributes.append([candidate_id, 1, "防御", 5, "", "", "迁移测试"])
    workbook.save(path)
    workbook.close()


LEGACY_P3_TEXTVAR_LINES = (
    "{【赤|249}{焰命|249}{格·|249}{测试|249}{1】|249}\\{对怪伤害吸收+5%|249}{·|249}{道术加成+100%|249}{·|249}{伤害系数+5%|249}",
    "{【碧|250}{落命|250}{格·|250}{测试|250}{2】|250}\\{防御+5|250}{·|250}{生命值+100|250}{·|250}{魔法值+100|250}",
    "{【沧|252}{澜命|252}{格·|252}{测试|252}{3】|252}\\{防御+5|252}{·|252}{生命值+100|252}{·|252}{魔法值+100|252}",
    "{【|249}{绮|69}{彩|250}{命|251}{格|249}{·|69}{彩|250}{】|251}\\{致命伤害+5%|31}{·|255}{处决概率+2%|147}{·|255}{鞭尸+3%|239}",
    "{【沧|249}{澜命|249}{格·|249}{红】|249}\\{鞭尸+3%|249}{·|249}{暴击伤害+3%|249}{·|249}{处决倍率+10%|249}",
    "{【|249}{绮|69}{彩|250}{命|251}{格|249}{·|69}{测|250}{试|251}{】|249}\\{尾刀斩杀+4%|31}{·|255}{鞭尸+5%|147}{·|255}{韧性+6|239}",
)


def _replace_with_many_registered_properties(path: Path) -> tuple[str, ...]:
    """Create a legal workbook whose candidate catalog spans more than 11 routes."""
    from openpyxl import load_workbook

    names = (
        "防御", "自定魔防", "攻击", "魔法", "道术", "生命值", "魔法值",
        "神力倍攻", "打怪伤害", "固定切割", "爆率", "最大爆率", "首刀斩杀",
        "韧性", "处决倍率", "处决时间", "伤害系数", "吸血", "每秒回血",
        "攻击加成", "魔法加成", "伤害吸收上限", "回收增加", "暴击几率",
        "攻击伤害", "魔御", "幸运", "准确",
    )
    workbook = load_workbook(path)
    definitions = workbook["命格定义"]
    required_candidates = (len(names) + 2) // 3
    for offset in range(4, required_candidates):
        candidate_id = 1001 + offset
        definitions.append([candidate_id, "是", f"容量命格{candidate_id}", 1, 1, "灵玉", 101, "容量回归"])
    sheet = workbook["命格属性"]
    sheet.delete_rows(2, sheet.max_row - 1)
    for index, name in enumerate(names):
        candidate_id = 1001 + index // 3
        order = index % 3 + 1
        sheet.append([candidate_id, order, name, index + 1, "", "", "容量回归"])
    workbook.save(path)
    workbook.close()
    return names


def _make_server(root: Path) -> Path:
    envir = root / "Mir200" / "Envir"
    npc = envir / "Market_Def" / "玄渊命格" / "龙魂觉醒-3.txt"
    npc.parent.mkdir(parents=True)
    npc.write_bytes("[@main]\r\n#IF\r\n#ACT\r\n#CALL [\\玄渊命格\\命格核心.txt] @XY_MG_MAIN\r\nBREAK\r\n".encode("gb18030"))
    core = envir / "QuestDiary" / "玄渊命格" / "命格核心.txt"
    core.parent.mkdir(parents=True)
    labels = ["XY_MG_MAIN", *(f"XY_MG_SLOT{i}" for i in range(1, 9)), "XY_MG_WASH_ONE", "XY_MG_WASH_TEN", "XY_MG_BOND_INFO", "XY_MG_PANEL_ROUTE"]
    core.write_bytes(("\r\n".join(f"[@{label}]\r\n#IF\r\n#ACT\r\n#SAY\r\n原界面-{label}" for label in labels) + "\r\n").encode("gb18030"))
    qfunction = envir / "Market_Def" / "QFunction-0.txt"
    qfunction.parent.mkdir(parents=True, exist_ok=True)
    qfunction.write_bytes(
        (
            "[@Login]\r\n#IF\r\n#ACT\r\n"
            "[@PlayLogin]\r\n#IF\r\n#ACT\r\n"
            "; XY_EXECUTION_LAB_TOUGHNESS_ANCHOR\r\n"
            "; XY_EXECUTION_LAB_PVE_BONUS_ANCHOR\r\n"
            "; XY_EQUIP_MAKER_RUNTIME_DAMAGE_COEFFICIENT_ANCHOR\r\n"
            "[@TakeOnEx]\r\n#IF\r\n#ACT\r\n"
            "[@TakeOffEx]\r\n#IF\r\n#ACT\r\n"
        ).encode("gb18030")
    )
    textvar = envir / "CustomItemPropertyTextVarList.txt"
    textvar.write_bytes(("\r\n".join(f"占位{index}" for index in range(1, 33)) + "\r\n").encode("gb18030"))
    return root


class MinggeDualWorkbookTests(unittest.TestCase):
    def test_new_public_bridge_has_lfm2_callable_envelope(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-new-bridge-") as td:
            before, bridge = mingge_dual_module._bridge_before(Path(td) / "missing.txt")
            lines = bridge.splitlines()

            self.assertIsNone(before)
            self.assertEqual(lines[0], "{")
            self.assertEqual(lines[-1], "}")
            self.assertEqual(lines.count("{"), 1)
            self.assertEqual(lines.count("}"), 1)

    def test_existing_bare_public_bridge_is_wrapped_without_body_changes(self) -> None:
        legacy = """; XY-MG-API-V1 PLATFORM SHARED BRIDGE
; XY-MG-CONTENT-STATUS-ANCHOR
; XY-MG-CONTENT-DRAW-ANCHOR
; XY-MG-CONTENT-READ-ANCHOR
; XY-MG-CONTENT-RECALC-ANCHOR
BREAK
"""
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-existing-bridge-") as td:
            path = Path(td) / "bridge.txt"
            path.write_bytes(mingge_dual_module._script_bytes(legacy))

            before, bridge = mingge_dual_module._bridge_before(path)
            lines = bridge.splitlines()

            self.assertEqual(before, path.read_bytes())
            self.assertEqual(lines[0], "{")
            self.assertEqual(lines[-1], "}")
            self.assertEqual(lines[1:-1], legacy.splitlines())

    def test_content_provider_has_lfm2_callable_envelope(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-callable-") as td:
            path = Path(td) / "41.xlsx"
            _save_content_workbook(path)
            book = load_mingge_content_workbook(path)

            provider = mingge_dual_module._render_content_provider(
                book, {candidate.candidate_id: 33 + index for index, candidate in enumerate(book.candidates)}
            )
            lines = provider.splitlines()

            self.assertEqual(lines[0], "{")
            self.assertEqual(lines[1], "; XY-MG-CONTENT-PROVIDER-V1")
            self.assertEqual(lines[-1], "}")
            self.assertEqual(lines.count("{"), 1)
            self.assertEqual(lines.count("}"), 1)

    def test_retired_only_catalog_is_valid_and_has_empty_draw_pool(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-retired-only-") as td:
            path = Path(td) / "41.xlsx"
            _save_legacy_p3_content_workbook(path)

            book = load_mingge_content_workbook(path)
            provider = mingge_dual_module._render_content_provider(
                book, {candidate_id: line_no for candidate_id, line_no in zip(range(1001, 1007), range(33, 39))}
            )

            self.assertEqual([candidate.state for candidate in book.candidates], ["retired"] * 6)
            self.assertNotIn("GOTO @XY_MG_CONTENT_CHOOSE_", provider.split("[@XY_MG_CONTENT_CHOOSE_1001]", 1)[0])
            self.assertIn("当前候选池或品质没有启用命格", provider)

    def test_two_workbooks_parse_stable_ids_and_normalized_attributes(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-dual-books-") as td:
            root = Path(td)
            npc_path = root / "40.xlsx"
            content_path = root / "41.xlsx"
            _save_npc_workbook(npc_path)
            _save_content_workbook(content_path)

            npc = load_mingge_npc_workbook(npc_path)
            content = load_mingge_content_workbook(content_path)

            self.assertEqual(npc.settings.target_item, "鞭尸灵玉")
            self.assertEqual(npc.settings.slot_count, 8)
            self.assertEqual([rule.batch_count for rule in npc.wash_rules], [1, 10])
            self.assertEqual([item.candidate_id for item in content.candidates], [1001, 1002, 1003, 1004])
            self.assertEqual(content.attributes[1004][0].property_name, "防御")
            self.assertEqual(content.color_schemes[104].attribute_colors, (31, 147, 239))

    def test_attack_speed_breakthrough_provider_reuses_textline40_and_existing_recalc(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-dual-speed-") as td:
            path = Path(td) / "41.xlsx"
            _save_content_workbook(path)
            from openpyxl import load_workbook

            workbook = load_workbook(path)
            sheet = workbook["命格属性"]
            sheet.delete_rows(2, sheet.max_row - 1)
            for candidate_id in (1001, 1002, 1003, 1004):
                sheet.append([candidate_id, 1, "攻速突破", 7, "", "", "复用TextLine40"])
            workbook["数据字典"].append(
                ["属性", "攻速突破", "", "attack_speed_breakthrough", "textline_runtime", 60, "复用既有接口", "TextLine40"]
            )
            workbook.save(path)
            workbook.close()

            book = load_mingge_content_workbook(path)
            provider = mingge_dual_module._render_content_provider(
                book, {1001: 33, 1002: 34, 1003: 35, 1004: 36}
            )

            self.assertIn("SetCustomItemAbil <$STR(N$XY_MG_API_EQUIP_SLOT)> <$STR(N$XY_MG_CONTENT_PROP_ROW)> 1 60", provider)
            self.assertIn("SetCustomItemAbil <$STR(N$XY_MG_API_EQUIP_SLOT)> <$STR(N$XY_MG_CONTENT_PROP_ROW)> 2 40", provider)
            self.assertIn(
                "SetCustomItemValueEx <$STR(N$XY_MG_API_EQUIP_SLOT)> <$STR(N$XY_MG_CONTENT_PROP_ROW)> = 40 <$STR(N$XY_MG_CONTENT_SUM1A)> 0",
                provider,
            )
            self.assertEqual(
                provider.count("#CALL [\\玄渊攻速突破\\全身攻速阈值核心.txt] @XY_AS_CAP_RECALC"),
                1,
            )
            self.assertNotIn("ChangeSpeed 2", provider)

    def test_retired_candidate_keeps_identity_but_leaves_draw_pool(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-dual-retired-") as td:
            path = Path(td) / "41.xlsx"
            _save_content_workbook(path)
            from openpyxl import load_workbook

            workbook = load_workbook(path)
            sheet = workbook["命格定义"]
            sheet.cell(3, 2).value = "退役"
            sheet.cell(4, 2).value = "否"
            workbook.save(path)
            workbook.close()

            book = load_mingge_content_workbook(path)
            self.assertEqual([item.candidate_id for item in book.candidates], [1001, 1002, 1004])
            self.assertEqual([item.state for item in book.candidates], ["active", "retired", "active"])

            provider = mingge_dual_module._render_content_provider(
                book, {1001: 33, 1002: 34, 1004: 36}
            )
            self.assertIn("; CANDIDATE 1002 TEXTVAR 34", provider)
            self.assertIn("[@XY_MG_CONTENT_CHOOSE_1002]", provider)
            self.assertNotIn("GOTO @XY_MG_CONTENT_CHOOSE_1002", provider)
            self.assertNotIn("CANDIDATE 1003", provider)
            self.assertNotIn("XY_MG_CONTENT_CHOOSE_1003", provider)

    def test_catalog_over_eleven_properties_is_compiled_without_union_gate_or_truncation(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-dual-many-properties-") as td:
            path = Path(td) / "41.xlsx"
            _save_content_workbook(path)
            names = _replace_with_many_registered_properties(path)

            book = load_mingge_content_workbook(path)
            aggregates = mingge_dual_module._content_aggregates(book)
            line_by_id = {item.candidate_id: 33 + index for index, item in enumerate(book.candidates)}
            provider = mingge_dual_module._render_content_provider(book, line_by_id)

            self.assertEqual(len(aggregates), len(names))
            self.assertIn("N$XY_MG_CONTENT_SUM28A", provider)
            self.assertIn("N$XY_MG_CONTENT_PROP_ROW", provider)
            self.assertIn(
                "MESSAGEBOX 当前灵玉实际同时存在的原生属性超过装备记录行容量；本次未静默截断，请调整候选组合。\n"
                "#IF\n#ACT\nGOTO @XY_MG_CONTENT_RECALC_ALL\nBREAK",
                provider,
            )
            for spec, _position in aggregates:
                self.assertIn(
                    f"<$STR(N$XY_MG_CONTENT_PROP_ROW)> 1 {spec.binding}", provider
                )

    def test_long_multicolor_text_is_warning_not_install_blocker(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-dual-long-text-") as td:
            root = Path(td)
            workbook_path = root / "41.xlsx"
            _save_content_workbook(workbook_path)
            from openpyxl import load_workbook

            workbook = load_workbook(workbook_path)
            sheet = workbook["命格定义"]
            sheet.cell(2, 3).value = "超长命格名称" * 12
            workbook.save(workbook_path)
            workbook.close()

            plan = MinggeContentService(root / "platform").preflight(
                _make_server(root / "server"), workbook_path
            )

            self.assertEqual(plan.blockers, ())
            self.assertTrue(any("TextVar" in warning and "字节" in warning for warning in plan.warnings))


class MinggeDualInstallTests(unittest.TestCase):
    def test_exact_legacy_p3_textvar_lines_are_adopted_by_stable_ids(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-legacy-adopt-") as td:
            root = Path(td)
            server = _make_server(root / "server")
            textvar = server / "Mir200/Envir/CustomItemPropertyTextVarList.txt"
            original = textvar.read_text(encoding="gb18030")
            textvar.write_bytes((original + "\r\n".join(LEGACY_P3_TEXTVAR_LINES) + "\r\n").encode("gb18030"))
            workbook = root / "41.xlsx"
            _save_legacy_p3_content_workbook(workbook)

            plan = MinggeContentService(root / "platform").preflight(server, workbook)

            self.assertEqual(plan.blockers, ())
            provider = next(
                change.after.decode("gb18030")
                for change in plan.changes
                if change.relative_path == str(CONTENT_PROVIDER_RELATIVE).replace("\\", "/")
            )
            for candidate_id, line_no in zip(range(1001, 1007), range(33, 39)):
                self.assertIn(f"; CANDIDATE {candidate_id} TEXTVAR {line_no}", provider)
            textvar_change = next(
                change
                for change in plan.changes
                if change.relative_path == "Mir200/Envir/CustomItemPropertyTextVarList.txt"
            )
            self.assertEqual(textvar_change.metadata["adopted_legacy_lines"], [33, 34, 35, 36, 37, 38])
            self.assertTrue(any("旧P3命格TextVar第33至38行" in warning for warning in plan.warnings))

    def test_preflight_accepts_twenty_eight_registered_catalog_properties(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-content-28-") as td:
            root = Path(td)
            workbook = root / "41.xlsx"
            _save_content_workbook(workbook)
            _replace_with_many_registered_properties(workbook)

            plan = MinggeContentService(root / "platform").preflight(
                _make_server(root / "server"), workbook
            )

            self.assertEqual(plan.blockers, ())
            provider = next(
                change.after.decode("gb18030")
                for change in plan.changes
                if change.relative_path == str(CONTENT_PROVIDER_RELATIVE).replace("\\", "/")
            )
            self.assertIn("N$XY_MG_CONTENT_SUM28A", provider)

    def test_content_preflight_blocks_attack_speed_when_existing_interface_is_missing(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-speed-dependency-") as td:
            root = Path(td)
            workbook_path = root / "41.xlsx"
            _save_content_workbook(workbook_path)
            from openpyxl import load_workbook

            workbook = load_workbook(workbook_path)
            sheet = workbook["命格属性"]
            sheet.delete_rows(2, sheet.max_row - 1)
            for candidate_id in (1001, 1002, 1003, 1004):
                sheet.append([candidate_id, 1, "攻速突破", 7, "", "", "复用TextLine40"])
            workbook.save(workbook_path)
            workbook.close()

            server = _make_server(root / "server")
            plan = MinggeContentService(root / "platform").preflight(server, workbook_path)

            self.assertTrue(any("攻速突破依赖" in blocker for blocker in plan.blockers))

    def test_npc_only_install_never_deducts_before_content_ready_check(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-npc-only-") as td:
            root = Path(td)
            platform = root / "platform"
            server = _make_server(root / "server")
            workbook = root / "40.xlsx"
            _save_npc_workbook(workbook)
            service = MinggeNpcService(platform)

            plan = service.preflight(server, workbook)
            self.assertEqual(plan.blockers, ())
            receipt = service.install(plan)

            rules = (server / NPC_RULES_RELATIVE).read_text(encoding="gb18030")
            self.assertLess(rules.index("@XY_MG_API_STATUS"), rules.index("GAMEGOLD -"))
            self.assertIn("命格内容包尚未安装，本次不会扣费", rules)
            self.assertIn("MOV N$XY_MG_NPC_BATCH_LEFT 10", rules)
            self.assertIn("RANDOMEX 50 100", rules)
            self.assertIn("LARGE U", rules)
            self.assertIn("DEC N$XY_MG_NPC_BATCH_LEFT 1", rules)
            core = (server / "Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt").read_text(encoding="gb18030")
            self.assertIn("XY-MG-NPC-ACTION-SLOT-1-V1-BEGIN", core)
            self.assertIn("@XY_MG_NPC_WASH_2", core)
            self.assertTrue((server / BRIDGE_RELATIVE).is_file())
            self.assertFalse((server / CONTENT_PROVIDER_RELATIVE).exists())
            self.assertEqual(receipt.status, "installed-pending-game-verification")

    def test_content_only_install_and_rollback_leave_safe_shared_bridge(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-content-only-") as td:
            root = Path(td)
            platform = root / "platform"
            server = _make_server(root / "server")
            workbook = root / "41.xlsx"
            _save_content_workbook(workbook)
            service = MinggeContentService(platform)

            plan = service.preflight(server, workbook)
            self.assertEqual(plan.blockers, ())
            receipt = service.install(plan)

            bridge = (server / BRIDGE_RELATIVE).read_text(encoding="gb18030")
            self.assertIn("XY-MG-CONTENT-V1-BEGIN", bridge)
            self.assertTrue((server / CONTENT_PROVIDER_RELATIVE).is_file())
            provider = (server / CONTENT_PROVIDER_RELATIVE).read_text(encoding="gb18030")
            self.assertIn("@XY_MG_CONTENT_REBUILD", provider)
            self.assertRegex(provider, r"SetCustomItemAbil .*PROP_ROW.* 1 1")
            self.assertIn("N$XY_MG_CONTENT_SUM1A", provider)
            textvar_lines = (server / "Mir200/Envir/CustomItemPropertyTextVarList.txt").read_text(encoding="gb18030").splitlines()
            self.assertIn("{铁壁|250}", textvar_lines[CONTENT_TEXTVAR_START - 1])
            self.assertIn("{命格|250}", textvar_lines[CONTENT_TEXTVAR_START - 1])

            service.rollback(server, receipt.transaction_id)

            bridge_after = (server / BRIDGE_RELATIVE).read_text(encoding="gb18030")
            self.assertNotIn("XY-MG-CONTENT-V1-BEGIN", bridge_after)
            self.assertIn("@XY_MG_API_STATUS", bridge_after)
            qfunction = (server / "Mir200/Envir/Market_Def/QFunction-0.txt").read_text(encoding="gb18030")
            self.assertNotIn(CONTENT_QFUNCTION_MARKER, qfunction)

    def test_install_order_and_independent_rollback_preserve_other_package(self) -> None:
        for order in (("npc", "content"), ("content", "npc")):
            with self.subTest(order=order), tempfile.TemporaryDirectory(prefix="xydp-mingge-order-") as td:
                root = Path(td)
                platform = root / "platform"
                server = _make_server(root / "server")
                npc_book = root / "40.xlsx"
                content_book = root / "41.xlsx"
                _save_npc_workbook(npc_book)
                _save_content_workbook(content_book)
                npc = MinggeNpcService(platform)
                content = MinggeContentService(platform)
                receipts = {}
                for package in order:
                    service, workbook = (npc, npc_book) if package == "npc" else (content, content_book)
                    receipts[package] = service.install(service.preflight(server, workbook))

                npc.rollback(server, receipts["npc"].transaction_id)

                self.assertFalse((server / NPC_RULES_RELATIVE).exists())
                self.assertTrue((server / CONTENT_PROVIDER_RELATIVE).is_file())
                bridge = (server / BRIDGE_RELATIVE).read_text(encoding="gb18030")
                self.assertIn("XY-MG-CONTENT-V1-BEGIN", bridge)
                npc_script = (server / "Mir200/Envir/Market_Def/玄渊命格/龙魂觉醒-3.txt").read_text(encoding="gb18030")
                self.assertNotIn("XY-MG-NPC-V1-BEGIN", npc_script)

    def test_content_rollback_preserves_installed_npc_package(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-content-rollback-") as td:
            root = Path(td)
            platform = root / "platform"
            server = _make_server(root / "server")
            npc_book = root / "40.xlsx"
            content_book = root / "41.xlsx"
            _save_npc_workbook(npc_book)
            _save_content_workbook(content_book)
            npc = MinggeNpcService(platform)
            content = MinggeContentService(platform)
            npc.install(npc.preflight(server, npc_book))
            receipt = content.install(content.preflight(server, content_book))

            content.rollback(server, receipt.transaction_id)

            self.assertTrue((server / NPC_RULES_RELATIVE).is_file())
            npc_script = (server / "Mir200/Envir/Market_Def/玄渊命格/龙魂觉醒-3.txt").read_text(encoding="gb18030")
            self.assertIn("XY-MG-NPC-V1-BEGIN", npc_script)
            bridge = (server / BRIDGE_RELATIVE).read_text(encoding="gb18030")
            self.assertNotIn("XY-MG-CONTENT-V1-BEGIN", bridge)

    def test_unmanaged_textvar_line_is_skipped_without_overwrite(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-textvar-conflict-") as td:
            root = Path(td)
            server = _make_server(root / "server")
            textvar = server / "Mir200/Envir/CustomItemPropertyTextVarList.txt"
            textvar.write_bytes((textvar.read_text(encoding="gb18030") + "其他系统占用\r\n").encode("gb18030"))
            workbook = root / "41.xlsx"
            _save_content_workbook(workbook)

            plan = MinggeContentService(root / "platform").preflight(server, workbook)

            self.assertEqual(plan.blockers, ())
            self.assertTrue(any("自动跳过" in warning and "33" in warning for warning in plan.warnings))
            textvar_change = next(
                change for change in plan.changes
                if change.relative_path == str(TEXTVAR_RELATIVE).replace("\\", "/")
            )
            self.assertEqual(
                textvar_change.before.decode("gb18030").splitlines()[32],
                textvar_change.after.decode("gb18030").splitlines()[32],
            )

    def test_new_candidates_skip_occupied_textvar_lines(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-mingge-textvar-skip-") as td:
            root = Path(td)
            server = _make_server(root / "server")
            textvar = server / "Mir200/Envir/CustomItemPropertyTextVarList.txt"
            lines = textvar.read_text(encoding="gb18030").splitlines()
            lines.extend([*(f"其他系统占用{line}" for line in range(33, 42)), "", "其他系统占用43"])
            textvar.write_bytes(("\r\n".join(lines) + "\r\n").encode("gb18030"))
            workbook = root / "41.xlsx"
            _save_content_workbook(workbook)

            plan = MinggeContentService(root / "platform").preflight(server, workbook)

            self.assertEqual(plan.blockers, ())
            provider_change = next(change for change in plan.changes if change.relative_path == str(CONTENT_PROVIDER_RELATIVE).replace("\\", "/"))
            provider_text = provider_change.after.decode("gb18030")
            self.assertIn("; CANDIDATE 1001 TEXTVAR 42", provider_text)
            self.assertIn("; CANDIDATE 1002 TEXTVAR 44", provider_text)
            self.assertIn("; CANDIDATE 1003 TEXTVAR 45", provider_text)
            self.assertIn("; CANDIDATE 1004 TEXTVAR 46", provider_text)


if __name__ == "__main__":
    unittest.main()
