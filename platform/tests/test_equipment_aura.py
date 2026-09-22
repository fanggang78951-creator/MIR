from __future__ import annotations

import json
import sqlite3
import sys
import tempfile
import unittest
from dataclasses import dataclass
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

PAYLOAD_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PAYLOAD_ROOT / "src"))

from xydp.equipment_aura import (  # noqa: E402
    AuraBinding,
    EquipmentAuraError,
    EquipmentAuraService,
    compile_runtime_core,
    sha256,
)


@dataclass(frozen=True)
class AuraTransactionFixture:
    platform: Path
    server: Path
    client: Path
    login: Path
    workbook: Path


def _write_script(path: Path, labels: tuple[str, ...]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    sections = []
    for label in labels:
        sections.extend((f"[@{label}]", "{{", "#IF", "#ACT", "BREAK", "}}", ""))
    path.write_bytes(("\r\n".join(sections) + "\r\n").encode("gb18030"))


def rewrite_binding(
    workbook_path: Path,
    equipment_name: str,
    style_id: str,
    multiplier: int,
    interval: int,
    priority: int,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "装备绑定"
    sheet.append((
        "启用状态", "装备名称", "特效样式ID", "伤害倍率",
        "攻击间隔秒", "优先级", "备注",
    ))
    sheet.append(("是", equipment_name, style_id, multiplier, interval, priority, "测试"))
    workbook.create_sheet("样式说明")
    workbook.save(workbook_path)


def create_transaction_fixture(root: Path) -> AuraTransactionFixture:
    platform = root / "platform"
    server = root / "server"
    client = root / "client"
    login = root / "login"
    workbook = root / "36_装备范围光环.xlsx"
    (platform / "assets" / "equipment_aura" / "library").mkdir(parents=True)
    (platform / "assets" / "equipment_aura" / "library" / "XY_EquipmentAura_10.wzl").write_bytes(b"WZL-80-FRAMES")
    (platform / "assets" / "equipment_aura" / "library" / "XY_EquipmentAura_10.wzx").write_bytes(b"WZX-80-INDEX")
    (client / "data").mkdir(parents=True)
    login.mkdir(parents=True)
    envir = server / "Mir200" / "Envir"
    _write_script(
        envir / "Market_Def" / "QFunction-0.txt",
        ("PlayLogin", "TakeOnEx", "TakeOffEx", "PlayDie", "NpcRevival", "EnterMap"),
    )
    _write_script(envir / "MapQuest_Def" / "QManage.txt", ("Startup",))
    (envir / "EffectImageList.txt").write_bytes("Base.wzl\r\n".encode("gb18030"))
    database = server / "Mud2" / "DB" / "ApexM2.DB"
    database.parent.mkdir(parents=True)
    connection = sqlite3.connect(database)
    try:
        connection.execute("CREATE TABLE StdItems (Idx INTEGER, Name TEXT)")
        connection.executemany(
            "INSERT INTO StdItems (Idx, Name) VALUES (?, ?)",
            ((1, "灵玉"), (2, "灵玉1"), (3, "灵玉2")),
        )
        connection.commit()
    finally:
        connection.close()
    rewrite_binding(workbook, "灵玉", "style02", 2, 1, 300)
    return AuraTransactionFixture(platform, server, client, login, workbook)


def snapshot_bytes(paths: tuple[Path, ...]) -> dict[str, bytes | None]:
    return {str(path): path.read_bytes() if path.is_file() else None for path in paths}


class AuraPlatformTransactionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.fixture = create_transaction_fixture(Path(self.temporary.name))
        self.service = EquipmentAuraService(self.fixture.platform)

    def test_public_compiler_keeps_one_damage_path_and_ten_styles(self) -> None:
        script = compile_runtime_core(
            (
                AuraBinding(True, "灵玉", "style01", 2, 1, 300, ""),
                AuraBinding(True, "灵玉1", "style10", 3, 2, 200, ""),
            ),
            27,
        )
        self.assertEqual(script.count("RangeHarm"), 1)
        self.assertEqual(script.count("[@XY_EQUIP_AURA_TICK]"), 1)
        self.assertEqual(len(self.service.list_styles()), 10)
        self.assertTrue(all(style["range"] == 3 for style in self.service.list_styles()))

    def test_search_target_equipment_uses_real_database_names(self) -> None:
        partial = self.service.search_target_equipment(self.fixture.server, "灵玉")
        exact = self.service.search_target_equipment(self.fixture.server, "灵玉1")
        self.assertEqual([item["name"] for item in partial], ["灵玉", "灵玉1", "灵玉2"])
        self.assertEqual(exact, ({"idx": 2, "name": "灵玉1"},))

    def test_upsert_binding_preserves_workbook_and_updates_or_appends(self) -> None:
        row = self.service.upsert_binding(
            self.fixture.workbook, "灵玉", "style10", 5, 2, 500
        )
        self.assertEqual(row, 2)
        appended = self.service.upsert_binding(
            self.fixture.workbook, "灵玉1", "style01", 3, 1, 400
        )
        self.assertEqual(appended, 3)
        workbook = load_workbook(self.fixture.workbook, data_only=True)
        try:
            sheet = workbook["装备绑定"]
            self.assertEqual(tuple(cell.value for cell in sheet[2])[:6], ("是", "灵玉", "style10", 5, 2, 500))
            self.assertEqual(tuple(cell.value for cell in sheet[3])[:6], ("是", "灵玉1", "style01", 3, 1, 400))
            self.assertIn("样式说明", workbook.sheetnames)
        finally:
            workbook.close()
        backups = tuple((self.fixture.platform / "backups" / "equipment-aura-workbook").glob("*.xlsx"))
        self.assertEqual(len(backups), 2)

    def _preflight(self):
        return self.service.preflight(
            self.fixture.server,
            self.fixture.client,
            self.fixture.login,
            self.fixture.workbook,
        )

    def test_preflight_is_read_only_and_plans_all_required_targets(self) -> None:
        roots = (self.fixture.platform, self.fixture.server, self.fixture.client, self.fixture.login)
        before = {
            str(path): (path.stat().st_mtime_ns, path.stat().st_size)
            for root in roots
            for path in root.rglob("*")
            if path.is_file()
        }
        plan = self._preflight()
        after = {
            str(path): (path.stat().st_mtime_ns, path.stat().st_size)
            for root in roots
            for path in root.rglob("*")
            if path.is_file()
        }
        self.assertFalse(plan.blockers)
        self.assertEqual(plan.static_resource_action, "install")
        self.assertEqual(before, after)
        self.assertEqual(plan.resource_index, 1)
        self.assertIn("M2运行时仍可预检与部署", "".join(plan.warnings))

    def test_install_is_idempotent_and_binding_update_reuses_static_library(self) -> None:
        first_plan = self._preflight()
        first = self.service.install(first_plan)
        client_wzl = self.fixture.client / "data" / "XY_EquipmentAura_10.wzl"
        before_wzl = sha256(client_wzl)
        no_op = self._preflight()
        self.assertFalse(no_op.files)
        self.assertEqual(no_op.static_resource_action, "reuse")
        repeated = self.service.install(no_op)
        self.assertEqual(repeated.status, "already-current")
        rewrite_binding(self.fixture.workbook, "灵玉1", "style03", 3, 2, 200)
        update = self._preflight()
        self.assertFalse(update.blockers)
        self.assertEqual(update.static_resource_action, "reuse")
        self.assertNotIn(client_wzl, update.affected_paths)
        self.service.install(update)
        self.assertEqual(sha256(client_wzl), before_wzl)
        self.assertNotEqual(first.transaction_id, repeated.transaction_id)

    def test_failure_restores_every_changed_file(self) -> None:
        plan = self._preflight()
        originals = snapshot_bytes(plan.affected_paths)
        original_replace = self.service._atomic_replace
        calls = 0

        def fail_on_fourth(path: Path, data: bytes) -> None:
            nonlocal calls
            calls += 1
            if calls == 4:
                raise OSError("模拟第四个文件锁定")
            original_replace(path, data)

        with patch.object(self.service, "_atomic_replace", side_effect=fail_on_fourth):
            with self.assertRaisesRegex(OSError, "模拟第四个文件锁定"):
                self.service.install(plan)
        self.assertEqual(snapshot_bytes(plan.affected_paths), originals)

    def test_preflight_hash_change_blocks_apply(self) -> None:
        plan = self._preflight()
        plan.paths.qfunction.write_bytes(plan.paths.qfunction.read_bytes() + b"; changed\r\n")
        with self.assertRaisesRegex(EquipmentAuraError, "预检后文件已变化"):
            self.service.install(plan)

    def test_timer_resource_event_and_manual_core_conflicts_are_blocked(self) -> None:
        qmanage = self.fixture.server / "Mir200" / "Envir" / "MapQuest_Def" / "QManage.txt"
        qmanage.write_bytes(qmanage.read_bytes() + b"[@OnTimer97]\r\n{\r\nBREAK\r\n}\r\n")
        plan = self._preflight()
        self.assertTrue(any("定时器事件已被占用" in item for item in plan.blockers))

        self.fixture = create_transaction_fixture(Path(self.temporary.name) / "second")
        self.service = EquipmentAuraService(self.fixture.platform)
        target_wzl = self.fixture.client / "data" / "XY_EquipmentAura_10.wzl"
        target_wzx = self.fixture.client / "data" / "XY_EquipmentAura_10.wzx"
        target_wzl.write_bytes(b"foreign")
        target_wzx.write_bytes(b"foreign-index")
        plan = self._preflight()
        self.assertTrue(any("静态资源冲突" in item for item in plan.blockers))

    def test_rollback_is_byte_exact_and_refuses_hand_modified_install(self) -> None:
        plan = self._preflight()
        originals = snapshot_bytes(plan.affected_paths)
        receipt = self.service.install(plan)
        result = self.service.rollback(self.fixture.server, receipt.transaction_id)
        self.assertEqual(result, "rolled-back")
        self.assertEqual(snapshot_bytes(plan.affected_paths), originals)

        plan = self._preflight()
        receipt = self.service.install(plan)
        plan.paths.core_script.write_bytes(plan.paths.core_script.read_bytes() + b"manual")
        with self.assertRaisesRegex(EquipmentAuraError, "安装后文件已被修改"):
            self.service.rollback(self.fixture.server, receipt.transaction_id)

    def test_receipt_records_operation_bindings_hashes_and_backups(self) -> None:
        receipt = self.service.install(self._preflight())
        payload = json.loads(receipt.receipt_path.read_text(encoding="utf-8"))
        self.assertEqual(payload["operation"], "equipment-aura")
        self.assertEqual(payload["status"], "installed-pending-game-verification")
        self.assertEqual(payload["bindings"][0]["equipment_name"], "灵玉")
        self.assertTrue(all("before_hash" in entry and "after_hash" in entry for entry in payload["files"]))


class AuraPlatformSurfaceTests(unittest.TestCase):
    def test_gui_exposes_equipment_aura_tab_and_shared_service(self) -> None:
        source = (PAYLOAD_ROOT / "src" / "xydp" / "gui.py").read_text(encoding="utf-8")
        self.assertIn('"装备光环"', source)
        self.assertIn("self.equipment_aura = EquipmentAuraService", source)
        self.assertIn("def _build_equipment_aura", source)

    def test_cli_registers_all_equipment_aura_commands(self) -> None:
        source = (PAYLOAD_ROOT / "src" / "xydp" / "cli.py").read_text(encoding="utf-8")
        for command in (
            "equipment-aura-list-styles",
            "equipment-aura-preflight",
            "equipment-aura-install",
            "equipment-aura-rollback",
        ):
            self.assertIn(command, source)


if __name__ == "__main__":
    unittest.main()
