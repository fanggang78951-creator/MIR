from __future__ import annotations

import hashlib
import inspect
import shutil
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from xydp.equipment_collection import (
    REWARD_FIELDS,
    CollectionCategory,
    CollectionGroupReward,
    CollectionItem,
    CollectionReward,
    CollectionWorkbook,
    EquipmentCollectionService,
    ResolvedCollectionItem,
    _compile_core,
    _group_flag_map,
    read_collection_workbook,
)
from xydp.npc_bundle import NpcBundleService


class CollectionPersistentStateTests(unittest.TestCase):
    def test_large_collection_uses_human_variables_instead_of_invalid_four_digit_flags(self):
        from xydp.equipment_collection import (
            _state_check,
            _state_set_lines,
            _state_variable,
        )

        variable = _state_variable("item", 622)

        self.assertEqual(variable, "XY_COL_ITEM_0622")
        self.assertEqual(_state_check(variable, 1), "CHECKVAR HUMAN XY_COL_ITEM_0622 = 1")
        self.assertEqual(
            _state_set_lines(variable),
            [
                "CALCVAR HUMAN XY_COL_ITEM_0622 = 1",
                "SAVEVAR HUMAN XY_COL_ITEM_0622 ..\\QuestDiary\\XY_System\\XuanYuanCollectionVar.txt",
            ],
        )

    def test_collection_login_hook_declares_and_loads_every_active_state_once(self):
        from xydp.equipment_collection import _compile_persistent_state_login

        hook = _compile_persistent_state_login(
            ["XY_COL_ITEM_0004", "XY_COL_ITEM_0005"],
            ["XY_COL_GROUP_C01_G01"],
        )

        self.assertEqual(hook.count("VAR Integer HUMAN XY_COL_ITEM_0004"), 1)
        self.assertEqual(hook.count("LOADVAR HUMAN XY_COL_ITEM_0004"), 1)
        self.assertIn("VAR Integer HUMAN XY_COL_GROUP_C01_G01", hook)
        self.assertNotIn("SET [1000]", hook)
from xydp.textpatch import scan_labels


FIXTURE = Path(__file__).parent / "fixtures" / "equipment_collection_sample.xlsx"
BUNDLE_ZIP = Path(r"C:\Users\Administrator\Downloads\艾尔登法环_全NPC平台填写_V2.1_20260920.zip")


def _server(tmp_path: Path, *, omit: str | None = None) -> Path:
    root = tmp_path / "MirServer"
    market = root / "Mir200" / "Envir" / "Market_Def"
    qmanage = root / "Mir200" / "Envir" / "MapQuest_Def"
    database = root / "Mud2" / "DB"
    market.mkdir(parents=True)
    qmanage.mkdir(parents=True)
    database.mkdir(parents=True)
    attr_overview = root / "Mir200" / "Envir" / "QuestDiary" / "玄渊功能" / "非常驻" / "属性总览"
    attr_overview.mkdir(parents=True)
    (root / "Mir200" / "M2Server.exe").write_bytes(b"")
    (root / "Mir200" / "Envir" / "MapInfo.txt").write_bytes("[0 测试地图 0]\r\n".encode("gb18030"))
    (root / "Mir200" / "Envir" / "UserCmd.txt").write_bytes(b"")
    (root / "Mir200" / "Envir" / "EffectImageList.txt").write_bytes(
        "Icon.wil\r\nItems.wil\r\nPrguse.wil\r\n".encode("gb18030")
    )
    (qmanage / "QManage.txt").write_bytes("[@MAIN1]\r\n#IF\r\n#ACT\r\nBREAK\r\n".encode("gb18030"))
    (market / "QFunction-0.txt").write_bytes(
        "\r\n".join((
            "[@XYDP_RecalcPower]", "#IF", "#ACT", "MOV N$倍攻 100",
            "; XY_EQUIP_MAKER_POWER_ANCHOR", "BREAK", "",
            "[@AttackDamage]", "#IF", "#ACT", "MOV N$XY_RT_Power 100",
            "; XY_EQUIP_MAKER_RUNTIME_POWER_ANCHOR",
            "; XY_EQUIP_MAKER_ATTACK_ANCHOR",
            "; XY_EQUIP_MAKER_RUNTIME_DROP_ANCHOR",
            "; XY_EQUIP_MAKER_RUNTIME_DAMAGE_COEFFICIENT_ANCHOR", "BREAK", "",
            "[@XYDP_RecalcBlast]", "#IF", "#ACT", "BREAK", "",
            "[@XYDP_RecalcDrop]", "#IF", "#ACT", "; XY_EQUIP_MAKER_DROP_ANCHOR", "BREAK", "",
            "[@XYDP_RecalcMonsterAbsorb]", "#IF", "#ACT", "BREAK", "",
            "[@XYDP_RecalcSustain]", "#IF", "#ACT", "BREAK", "",
            "[@XYDP_ShowAttrBuff]", "#IF", "#ACT", "BREAK", "",
        )).encode("gb18030"),
    )
    (attr_overview / "玄渊三属性按钮.txt").write_bytes(
        "\r\n".join((
            "[@XY_UI_STATUS_REFRESH]", "#IF", "#ACT",
            "MOV N$XY_UI_A_Value02 0",
            "#IF", "SMALL N$XY_UI_A_Value02 0", "#ACT", "MOV N$XY_UI_A_Value02 0",
            "ADDBUTTON 17 91 1772 1773 1774 246 17 0 -1 250#攻击属性A\\254#伤害系数:+<$STR(N$XY_UI_A_Value02)>%",
            "BREAK", "",
        )).encode("gb18030")
    )
    connection = sqlite3.connect(database / "ApexM2.DB")
    try:
        connection.execute("CREATE TABLE StdItems (Idx INTEGER, Name TEXT, StdMode INTEGER, Shape INTEGER, Looks INTEGER)")
        rows = [
            (701, "木剑", 5, 11, 6101),
            (702, "铁剑", 5, 12, 6102),
            (703, "青铜剑", 5, 13, 6103),
        ]
        connection.executemany(
            "INSERT INTO StdItems (Idx, Name, StdMode, Shape, Looks) VALUES (?, ?, ?, ?, ?)",
            [row for row in rows if row[1] != omit],
        )
        connection.commit()
    finally:
        connection.close()
    return root


def _platform(tmp_path: Path) -> Path:
    root = tmp_path / "platform"
    build = root / "labs" / "equipment_collection" / "graphical_candidate" / "build"
    build.mkdir(parents=True)
    source = Path(r"E:\XuanYuanDevPlatform\labs\equipment_collection\graphical_candidate\build")
    shutil.copy2(source / "XY_EquipmentCollection.wzl", build / "XY_EquipmentCollection.wzl")
    shutil.copy2(source / "XY_EquipmentCollection.wzx", build / "XY_EquipmentCollection.wzx")
    return root


def _client(tmp_path: Path) -> Path:
    root = tmp_path / "client"
    (root / "data").mkdir(parents=True)
    (root / "data" / "Prguse.wzl").write_bytes(b"test-prguse-wzl")
    (root / "data" / "Prguse.wzx").write_bytes(b"test-prguse-wzx")
    return root


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class EquipmentCollectionTests(unittest.TestCase):
    def test_collection_anchor_hooks_have_stable_canonical_order(self) -> None:
        source = inspect.getsource(EquipmentCollectionService.preflight)

        self.assertGreaterEqual(source.count("canonical_before_existing_hooks=True"), 2)

    def test_v21_bundle_supports_622_ids_and_46_group_pages(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-collection-v21-") as temp:
            bundle_root = NpcBundleService(Path(r"E:\XuanYuanDevPlatform")).materialize_bundle(
                BUNDLE_ZIP, Path(temp)
            )
            source = next(bundle_root.rglob("33_装备收集图鉴.xlsx"))
            workbook = read_collection_workbook(source)

            self.assertEqual(len(workbook.items), 619)
            self.assertEqual(max(item.collection_id for item in workbook.items), 622)
            self.assertEqual(max(item.flag for item in workbook.items), 1021)
            self.assertEqual(len(workbook.categories), 46)
            self.assertEqual(len(workbook.groups), 46)
            group_flags = _group_flag_map(workbook)
            self.assertEqual(min(group_flags.values()), 1100)
            self.assertEqual(max(group_flags.values()), 1145)

    def test_large_collection_initializes_missing_state_file_without_overwriting_existing_save(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xydp-collection-state-") as temp:
            tmp_path = Path(temp)
            workbook_path = tmp_path / "persistent-collection.xlsx"
            shutil.copy2(FIXTURE, workbook_path)
            editable = load_workbook(workbook_path)
            editable["收集明细"]["B2"] = 622
            for row in editable["界面设置"].iter_rows():
                if row[0].value == "最大收集ID":
                    row[1].value = 622
                    break
            editable.save(workbook_path)
            server = _server(tmp_path)

            qmanage = server / "Mir200" / "Envir" / "MapQuest_Def" / "QManage.txt"
            qmanage.write_bytes(
                "[@Login]\r\n#IF\r\n#ACT\r\nBREAK\r\n".encode("gb18030")
            )

            service = EquipmentCollectionService(_platform(tmp_path))
            client = _client(tmp_path)
            plan = service.preflight(workbook_path, server, client)
            self.assertEqual(plan.blockers, [])
            state_changes = [
                change
                for change in plan.changes
                if change.relative_path.endswith("XuanYuanCollectionVar.txt")
            ]
            self.assertEqual(len(state_changes), 1)
            self.assertIsNone(state_changes[0].before)
            self.assertEqual(state_changes[0].after, b"\xef\xbb\xbf")
            self.assertEqual(state_changes[0].operation, "initialize-collection-state")

            state_path = (
                server
                / "Mir200"
                / "Envir"
                / "QuestDiary"
                / "XY_System"
                / "XuanYuanCollectionVar.txt"
            )
            state_path.parent.mkdir(parents=True, exist_ok=True)
            saved = b"\xef\xbb\xbf" + "[大胆]\r\nXY_COL_ITEM_0004=1\r\n".encode("utf-8")
            state_path.write_bytes(saved)

            second = service.preflight(workbook_path, server, client)
            self.assertFalse(
                any(
                    change.relative_path.endswith("XuanYuanCollectionVar.txt")
                    for change in second.changes
                )
            )
            self.assertEqual(state_path.read_bytes(), saved)

    def test_official_workbook_contains_first_game_test_group(self) -> None:
        workbook = read_collection_workbook(
            Path(r"E:\XuanYuanDevPlatform\所需材料表格汇总\33_装备收集图鉴.xlsx")
        )
        self.assertEqual(workbook.title, "装备收集图鉴")
        self.assertEqual([item.item_name for item in workbook.items], ["木剑", "铁剑", "青铜剑"])
        self.assertEqual([item.flag for item in workbook.items], [400, 401, 402])
        self.assertTrue(all(item.required_count == 1 and item.consume for item in workbook.items))
        self.assertEqual(workbook.items[0].reward.values["神力倍攻%"], 1)
        self.assertEqual(workbook.items[1].reward.values["打怪伤害%"], 1)
        self.assertEqual(workbook.items[2].reward.values["爆率%"], 1)
        self.assertEqual(len(workbook.groups), 1)
        self.assertEqual(workbook.groups[0].category_name, "漂流群岛")
        self.assertEqual(workbook.groups[0].reward.values["伤害系数%"], 1)
        self.assertEqual(
            [category.category_name for category in workbook.categories],
            [
                "漂流群岛", "宁姆格福", "宁姆格福西域", "立业尼亚海域", "湖之利耶尼亚",
                "盖利德", "化圣雪原", "巨人山顶", "亚坛高原", "王城罗德尔",
                "圣树分支", "法姆·亚兹拉", "永恒之城",
            ],
        )
        self.assertEqual([category.category_order for category in workbook.categories], list(range(1, 14)))

    def test_group_damage_coefficient_updates_effect_and_left_top_display(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp_path = Path(td)
            server = _server(tmp_path)
            client = _client(tmp_path)
            service = EquipmentCollectionService(_platform(tmp_path))
            workbook = Path(r"E:\XuanYuanDevPlatform\所需材料表格汇总\33_装备收集图鉴.xlsx")

            plan = service.preflight(workbook, server, client)
            self.assertEqual(plan.blockers, [])
            qfunction_change = next(
                change for change in plan.changes if change.relative_path.endswith("QFunction-0.txt")
            )
            qfunction = qfunction_change.after.decode("gb18030")
            self.assertIn("CHECK [600] 1\r\n#ACT\r\nINC N$XY_RT_DamageCoeff 1", qfunction)
            self.assertIn("[@XY_COLLECTION_GROUP_AUDIT_001]", qfunction)
            self.assertIn(
                "CHECK [400] 1\r\nCHECK [401] 1\r\nCHECK [402] 1\r\n#ACT\r\nSET [600] 1",
                qfunction,
            )

            overview_change = next(
                change for change in plan.changes if change.relative_path.endswith("玄渊三属性按钮.txt")
            )
            overview = overview_change.after.decode("gb18030")
            self.assertIn("; XY_EQUIPMENT_COLLECTION_DAMAGE_COEFFICIENT_DISPLAY_ANCHOR", overview)
            self.assertIn("CHECK [600] 1\r\n#ACT\r\nINC N$XY_UI_A_Value02 1", overview)

            paths = [
                server / "Mir200" / "Envir" / "Market_Def" / "QFunction-0.txt",
                server / "Mir200" / "Envir" / "MapQuest_Def" / "QManage.txt",
                server / "Mir200" / "Envir" / "UserCmd.txt",
                server / "Mir200" / "Envir" / "QuestDiary" / "玄渊功能" / "非常驻" / "属性总览" / "玄渊三属性按钮.txt",
            ]
            before = {path: _hash(path) for path in paths}
            receipt = service.install(plan)
            second = service.preflight(workbook, server, client)
            self.assertEqual(second.blockers, [])
            self.assertEqual(second.changes, [])
            service.rollback(server, receipt.transaction_id)
            self.assertEqual({path: _hash(path) for path in paths}, before)

    def test_target_idx_controls_itemshow_and_transaction_is_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp_path = Path(td)
            server = _server(tmp_path)
            client = _client(tmp_path)
            service = EquipmentCollectionService(_platform(tmp_path))
            plan = service.preflight(FIXTURE, server, client)
            self.assertEqual(plan.blockers, [])
            self.assertEqual([item["idx"] for item in plan.items], [701, 702, 703])
            self.assertEqual([item["itemshow"] for item in plan.items], ["ItemShow:701", "ItemShow:702", "ItemShow:703"])
            self.assertEqual([item["shape"] for item in plan.items], [11, 12, 13])
            qfunction_change = next(change for change in plan.changes if change.relative_path.endswith("QFunction-0.txt"))
            generated = qfunction_change.after.decode("gb18030")
            self.assertIn("<&ItemShow:701:", generated)
            self.assertIn("<&ItemShow:702:", generated)
            self.assertIn("<&ItemShow:703:", generated)
            self.assertNotIn("ItemShow:6", generated)
            self.assertIn("ItemShow:701:0:225:118:1:0:1/@XY_COLLECTION_APPLY_001", generated)
            self.assertIn("ItemShow:701:0:225:118:1:0:0", generated)
            self.assertIn("OPENMERCHANTBIGDLG 3 0 0 4 0 0 1 746 6 1", generated)
            self.assertIn("<&Text:鼠标移此可显示点亮属性|", generated)
            self.assertIn("<&Text:第1页 共1页:70:27{FCOLOR=146}>", generated)
            self.assertNotIn("第1/1页", generated)
            self.assertIn("<&Text:已点亮<$STR(N$XY_COL_DONE)>/3:70:49{FCOLOR=250}>", generated)
            self.assertNotIn("已点亮<$STR(N$XY_COL_DONE)>/3:405:", generated)
            self.assertIn("^250#武器收藏点亮属性", generated)
            self.assertNotIn("^146#已点亮:<$STR(N$XY_COL_DONE)>/3", generated)
            self.assertIn("MOV S$XY_COL_HOVER01 木剑[未点亮]增加", generated)
            self.assertIn("MOV S$XY_COL_HOVER01 木剑[已点亮]增加", generated)
            self.assertIn("^251#<$STR(S$XY_COL_HOVER01)>", generated)
            self.assertIn("^254#<$STR(S$XY_COL_HOVER02)>", generated)
            self.assertIn("^250#<$STR(S$XY_COL_HOVER03)>", generated)
            self.assertIn("MOV S$XY_COL_GROUP_HOVER 全部点亮[未激活]增加攻击下限5、攻击上限5", generated)
            self.assertIn("MOV S$XY_COL_GROUP_HOVER 全部点亮[已激活]增加攻击下限5、攻击上限5", generated)
            self.assertIn("^253#<$STR(S$XY_COL_GROUP_HOVER)>", generated)
            self.assertIn(":525:27{FCOLOR=250}>\\", generated)
            self.assertIn(
                "<&Text:一键收集本页:525:49{FCOLOR=251}/@XY_COLLECTION_BATCH_001_PAGE_01>",
                generated,
            )
            batch_one = generated.split("[@XY_COLLECTION_BATCH_001_PAGE_01]", 1)[1].split(
                "[@XY_COLLECTION_APPLY_001]", 1
            )[0]
            self.assertIn("CHECK [400] 0\r\nCHECKITEM 木剑 1\r\n#ACT\r\nTAKE 木剑 1", batch_one)
            self.assertIn("CHECK [401] 0\r\nCHECKITEM 铁剑 1", batch_one)
            self.assertIn("CHECK [402] 0\r\nCHECKITEM 青铜剑 1", batch_one)
            self.assertEqual(batch_one.count("INC N$XY_COL_BATCH_COUNT 1"), 3)
            self.assertIn(
                "CHECK [600] 0\r\nCHECK [400] 1\r\nCHECK [401] 1\r\nCHECK [402] 1",
                batch_one,
            )
            self.assertIn("EQUAL N$XY_COL_BATCH_COUNT 0", batch_one)
            self.assertIn("本页一键收集完成，共点亮<$STR(N$XY_COL_BATCH_COUNT)>件装备。", batch_one)
            self.assertNotIn("MESSAGEBOX 背包中缺少", batch_one)
            self.assertNotIn("<&ImgEx:", generated)
            self.assertNotIn("[@XY_COLLECTION_REWARD_", generated)
            self.assertNotIn("返回图鉴", generated)
            apply_one = generated.split("[@XY_COLLECTION_APPLY_001]", 1)[1].split(
                "[@XY_COLLECTION_APPLY_002]", 1
            )[0]
            self.assertIn("NOT CHECKITEM 木剑 1", apply_one)
            self.assertLess(apply_one.index("NOT CHECKITEM 木剑 1"), apply_one.index("TAKE 木剑 1"))
            self.assertNotIn("#ELSEACT", apply_one)
            self.assertIn("CHECK [400] 1", generated)
            self.assertIn("ChangeHumAbilityEX 11 + 5", generated)

            paths = [
                server / "Mir200" / "Envir" / "Market_Def" / "QFunction-0.txt",
                server / "Mir200" / "Envir" / "MapQuest_Def" / "QManage.txt",
                server / "Mir200" / "Envir" / "UserCmd.txt",
                server / "Mir200" / "Envir" / "EffectImageList.txt",
            ]
            before = {path: _hash(path) for path in paths}
            receipt = service.install(plan)
            self.assertTrue((client / "data" / "XY_EquipmentCollection.wzl").is_file())
            self.assertTrue((client / "data" / "XY_EquipmentCollection.wzx").is_file())
            second = service.preflight(FIXTURE, server, client)
            self.assertEqual(second.blockers, [])
            self.assertEqual(second.changes, [])
            service.rollback(server, receipt.transaction_id)
            self.assertEqual({path: _hash(path) for path in paths}, before)
            self.assertFalse((client / "data" / "XY_EquipmentCollection.wzl").exists())
            self.assertFalse((client / "data" / "XY_EquipmentCollection.wzx").exists())

    def test_missing_target_item_blocks_without_changes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp_path = Path(td)
            server = _server(tmp_path, omit="铁剑")
            plan = EquipmentCollectionService(_platform(tmp_path)).preflight(
                FIXTURE, server, _client(tmp_path)
            )
            self.assertTrue(any("目标服装备不存在：铁剑" in blocker for blocker in plan.blockers))
            self.assertEqual(plan.changes, [])

    def test_graphic_collection_requires_client_and_blocks_resource_collision(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp_path = Path(td)
            server = _server(tmp_path)
            service = EquipmentCollectionService(_platform(tmp_path))
            no_client = service.preflight(FIXTURE, server)
            self.assertTrue(any("必须选择目标客户端" in blocker for blocker in no_client.blockers))

            client = _client(tmp_path)
            (client / "data" / "Prguse.wzl").unlink()
            no_icon_library = service.preflight(FIXTURE, server, client)
            self.assertEqual(no_icon_library.blockers, [])
            (client / "data" / "XY_EquipmentCollection.wzl").write_bytes(b"foreign")
            collision = service.preflight(FIXTURE, server, client)
            self.assertTrue(any("同名但内容不同" in blocker for blocker in collision.blockers))
            self.assertEqual(collision.changes, [])

    def test_graphic_layout_generates_item_and_category_pagination(self) -> None:
        zero = CollectionReward({name: 0 for name in REWARD_FIELDS})
        items: list[CollectionItem] = []
        resolved: list[ResolvedCollectionItem] = []
        collection_id = 1
        for category_number in range(1, 8):
            count = 22 if category_number == 1 else 1
            for item_number in range(1, count + 1):
                item = CollectionItem(
                    collection_id=collection_id,
                    category_id=f"cat{category_number}",
                    category_name=f"分类{category_number}",
                    category_order=category_number,
                    item_order=item_number,
                    item_name=f"装备{collection_id}",
                    required_count=1,
                    consume=True,
                    reward=zero,
                )
                items.append(item)
                resolved.append(ResolvedCollectionItem(item, 700 + collection_id, 5, 0, 0))
                collection_id += 1
        groups = tuple(
            CollectionGroupReward(
                category_id=f"cat{number}",
                category_name=f"分类{number}",
                category_order=number,
                reward=zero,
                message=f"分类{number}完成",
            )
            for number in range(1, 8)
        )
        workbook = CollectionWorkbook(
            "synthetic.xlsx", "0" * 64, "装备收集图鉴", 21,
            "装备收集", 92, tuple(items), groups,
        )
        group_flags = {f"cat{number}": 599 + number for number in range(1, 8)}
        generated = _compile_core(workbook, tuple(resolved), group_flags, 14)
        self.assertIn("[@XY_COLLECTION_CAT_001_PAGE_02]", generated)
        self.assertIn("[@XY_COLLECTION_CAT_007_PAGE_01]", generated)
        self.assertIn("分类下页", generated)
        self.assertIn("分类上页", generated)
        self.assertIn("ItemShow:721:0:676:305:1:0:1", generated)
        self.assertIn("ItemShow:722:0:225:118:1:0:1", generated)
        batch_page_one = generated.split("[@XY_COLLECTION_BATCH_001_PAGE_01]", 1)[1].split(
            "[@XY_COLLECTION_CAT_001_PAGE_02]", 1
        )[0]
        batch_page_two = generated.split("[@XY_COLLECTION_BATCH_001_PAGE_02]", 1)[1].split(
            "[@XY_COLLECTION_CAT_002_PAGE_01]", 1
        )[0]
        self.assertEqual(batch_page_one.count("INC N$XY_COL_BATCH_COUNT 1"), 21)
        self.assertIn("CHECK [420] 0", batch_page_one)
        self.assertNotIn("CHECK [421] 0", batch_page_one)
        self.assertEqual(batch_page_two.count("INC N$XY_COL_BATCH_COUNT 1"), 1)
        self.assertIn("CHECK [421] 0", batch_page_two)
        page_two = generated.split("[@XY_COLLECTION_CAT_001_PAGE_02]", 1)[1].split(
            "[@XY_COLLECTION_CAT_002_PAGE_01]", 1
        )[0]
        self.assertIn("<&Text:第2页 共2页:70:27{FCOLOR=146}>", page_two)
        self.assertNotIn("第2/2页", page_two)
        self.assertEqual(page_two.count("INC N$XY_COL_DONE 1"), 22)
        self.assertNotIn("[@XY_COLLECTION_REWARD_", generated)
        self.assertEqual(generated.count("<&Text:鼠标移此可显示点亮属性|"), 8)
        self.assertFalse(scan_labels(generated).duplicates)

    def test_large_group_open_path_renders_selected_page_without_pre_render_audit(self) -> None:
        zero = CollectionReward({name: 0 for name in REWARD_FIELDS})
        items = tuple(
            CollectionItem(
                collection_id=index,
                category_id=f"cat{index:02d}",
                category_name=f"分类{index}",
                category_order=index,
                item_order=1,
                item_name=f"装备{index}",
                required_count=1,
                consume=True,
                reward=zero,
            )
            for index in range(1, 47)
        )
        resolved = tuple(
            ResolvedCollectionItem(item, 700 + item.collection_id, 5, 0, 0)
            for item in items
        )
        groups = tuple(
            CollectionGroupReward(
                category_id=f"cat{index:02d}",
                category_name=f"分类{index}",
                category_order=index,
                reward=zero,
                message=f"分类{index}完成",
            )
            for index in range(1, 47)
        )
        workbook = CollectionWorkbook(
            "synthetic.xlsx", "0" * 64, "装备收集图鉴", 21,
            "装备收集", 92, items, groups,
        )
        group_flags = {f"cat{index:02d}": 599 + index for index in range(1, 47)}

        generated = _compile_core(workbook, resolved, group_flags, 14)

        command = generated.split("[@XY_COLLECTION_COMMAND]", 1)[1].split(
            "[@XY_COLLECTION_MAIN]", 1
        )[0]
        main = generated.split("[@XY_COLLECTION_MAIN]", 1)[1].split(
            "[@XY_COLLECTION_CAT_001_PAGE_01]", 1
        )[0]
        self.assertIn("GOTO @XY_COLLECTION_CAT_001_PAGE_01", command)
        self.assertIn("GOTO @XY_COLLECTION_CAT_001_PAGE_01", main)
        self.assertNotIn("XY_COLLECTION_GROUP_AUDIT", generated)
        self.assertIn("/@XY_COLLECTION_CAT_021_PAGE_01>", generated)

    def test_empty_continent_pages_are_visible_and_do_not_allocate_reward_flags(self) -> None:
        zero = CollectionReward({name: 0 for name in REWARD_FIELDS})
        reward = CollectionReward({name: (1 if name == "神力倍攻%" else 0) for name in REWARD_FIELDS})
        items = tuple(
            CollectionItem(
                collection_id=index,
                category_id="starter",
                category_name="漂流群岛",
                category_order=1,
                item_order=index,
                item_name=name,
                required_count=1,
                consume=True,
                reward=reward,
            )
            for index, name in enumerate(("木剑", "铁剑", "青铜剑"), start=1)
        )
        resolved = tuple(
            ResolvedCollectionItem(item, 700 + item.collection_id, 5, 0, 0)
            for item in items
        )
        names = (
            "漂流群岛", "宁姆格福", "宁姆格福西域", "立业尼亚海域", "湖之利耶尼亚",
            "盖利德", "化圣雪原", "巨人山顶", "亚坛高原", "王城罗德尔",
            "圣树分支", "法姆·亚兹拉", "永恒之城",
        )
        categories = tuple(
            CollectionCategory("starter" if index == 1 else f"continent_{index:02d}", name, index)
            for index, name in enumerate(names, start=1)
        )
        group = CollectionGroupReward("starter", "漂流群岛", 1, zero, "漂流群岛完成")
        workbook = CollectionWorkbook(
            "synthetic.xlsx", "0" * 64, "装备收集图鉴", 21,
            "装备收集", 92, items, (group,), categories,
        )
        generated = _compile_core(workbook, resolved, {"starter": 600}, 14)
        self.assertIn("[@XY_COLLECTION_CAT_013_PAGE_01]", generated)
        self.assertIn("<&Text:宁姆格福:53:154", generated)
        self.assertIn("<&Text:永恒之城:53:98", generated)
        empty_page = generated.split("[@XY_COLLECTION_CAT_002_PAGE_01]", 1)[1].split(
            "[@XY_COLLECTION_CAT_003_PAGE_01]", 1
        )[0]
        self.assertIn("MOV S$XY_COL_HOVER01 本大陆暂无可点亮装备", empty_page)
        self.assertIn("MOV S$XY_COL_GROUP_HOVER 全部点亮[未配置]无额外属性奖励", empty_page)
        self.assertIn("已点亮<$STR(N$XY_COL_DONE)>/0", empty_page)
        self.assertNotIn("ItemShow:", empty_page)
        self.assertNotIn("一键收集本页", empty_page)
        self.assertNotIn("@XY_COLLECTION_BATCH_002_PAGE_01", generated)
        self.assertIn("分类下页", generated)
        self.assertIn("分类上页", generated)
        self.assertEqual(generated.count("; XY-COLLECTION-GROUP"), 1)
        self.assertFalse(scan_labels(generated).duplicates)
