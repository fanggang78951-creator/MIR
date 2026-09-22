from __future__ import annotations

import hashlib
import json
import os
import sys
import re
import shutil
import sqlite3
import subprocess
import tempfile
import time
import unittest
from copy import deepcopy
from dataclasses import dataclass, replace
from pathlib import Path, PurePosixPath
from unittest import mock

from openpyxl import load_workbook
from PIL import Image, ImageColor

PAYLOAD_ROOT = Path(__file__).resolve().parents[1]
FIXTURES = PAYLOAD_ROOT / "tests" / "fixtures" / "mingge"
sys.path.insert(0, str(PAYLOAD_ROOT / "src"))

import xydp.mingge as mingge_module  # noqa: E402
from xydp.mingge import (  # noqa: E402
    QUALITIES,
    MingGeError,
    QualityProgress,
    compile_mingge_payload,
    load_property_registry,
    read_mingge_workbook,
    resolve_quality,
)
from xydp.mingge_image import MingGeImageError, render_mingge_labels  # noqa: E402


REGISTRY_PATH = (
    PAYLOAD_ROOT
    / "packages"
    / "candidate"
    / "xy.optional.mingge-system"
    / "payload"
    / "property_registry.json"
)
REGISTRY = load_property_registry(REGISTRY_PATH)
VENDOR_ROOT = PAYLOAD_ROOT / "vendor"
PROVIDER_PATH = (
    PAYLOAD_ROOT
    / "wzl编辑"
    / "bin"
    / "XuanYuanWzlProvider"
    / "XuanYuanWzlProvider.exe"
)
TEMPLATE_DIR = PAYLOAD_ROOT / "wzl编辑" / "workspace" / "mingge-template"
PACKAGE = (
    PAYLOAD_ROOT
    / "packages"
    / "candidate"
    / "xy.optional.mingge-system"
)
SERVER_TEMPLATE_DIR = PACKAGE / "payload" / "server" / "templates"
INTERFACE_ROOT = PAYLOAD_ROOT / "接口"


class MingGeWorkbookTests(unittest.TestCase):
    @staticmethod
    def _set_cell(workbook, sheet_name: str, cell: str, value: object) -> None:
        workbook[sheet_name][cell] = value

    def _mutated_fixture(self, mutation):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        path = Path(temporary.name) / "invalid.xlsx"
        shutil.copy2(FIXTURES / "valid.xlsx", path)
        workbook = load_workbook(path)
        try:
            mutation(workbook)
            workbook.save(path)
        finally:
            workbook.close()
        return path

    @staticmethod
    def _rule_row(workbook, rule_id: str) -> int:
        sheet = workbook["品质规则"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, 1).value == rule_id:
                return row
        raise AssertionError(f"测试夹具缺少规则: {rule_id}")

    def test_valid_fixture_has_eight_affixes_four_qualities_and_six_ordered_rules(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        self.assertEqual(book.settings.target_item_name, "鞭尸灵玉")
        self.assertEqual(book.settings.equip_slot, 17)
        self.assertEqual(len(book.affixes), 8)
        self.assertEqual({row.quality for row in book.attributes}, {"绿", "橙", "红", "彩"})
        self.assertEqual(
            tuple((rule.rule_id, rule.priority) for rule in book.quality_rules),
            (
                ("MG_RULE_COLOR_PITY", 1),
                ("MG_RULE_COLOR_CYCLE", 2),
                ("MG_RULE_RED_PITY", 3),
                ("MG_RULE_RED_NATURAL", 4),
                ("MG_RULE_ORANGE_NATURAL", 5),
                ("MG_RULE_GREEN_FALLBACK", 6),
            ),
        )

    def test_each_non_green_core_rule_is_required(self) -> None:
        rule_ids = (
            "MG_RULE_COLOR_PITY",
            "MG_RULE_COLOR_CYCLE",
            "MG_RULE_RED_PITY",
            "MG_RULE_RED_NATURAL",
            "MG_RULE_ORANGE_NATURAL",
        )
        for rule_id in rule_ids:
            with self.subTest(rule_id=rule_id):
                def mutate(workbook, target=rule_id):
                    workbook["品质规则"].delete_rows(self._rule_row(workbook, target))

                path = self._mutated_fixture(mutate)
                with self.assertRaisesRegex(MingGeError, "启用品质规则必须恰好包含六条"):
                    read_mingge_workbook(path)

    def test_enabled_rule_priorities_must_be_unique(self) -> None:
        def mutate(workbook):
            row = self._rule_row(workbook, "MG_RULE_RED_PITY")
            workbook["品质规则"].cell(row, 2).value = 2

        path = self._mutated_fixture(mutate)
        with self.assertRaisesRegex(MingGeError, "启用品质规则优先级必须唯一且连续为1～6"):
            read_mingge_workbook(path)

    def test_enabled_rule_priorities_must_not_have_gaps(self) -> None:
        def mutate(workbook):
            row = self._rule_row(workbook, "MG_RULE_ORANGE_NATURAL")
            workbook["品质规则"].cell(row, 2).value = 7

        path = self._mutated_fixture(mutate)
        with self.assertRaisesRegex(MingGeError, "启用品质规则优先级必须唯一且连续为1～6"):
            read_mingge_workbook(path)

    def test_enabled_rule_priority_semantics_must_not_be_reordered(self) -> None:
        def mutate(workbook):
            sheet = workbook["品质规则"]
            red_row = self._rule_row(workbook, "MG_RULE_RED_NATURAL")
            orange_row = self._rule_row(workbook, "MG_RULE_ORANGE_NATURAL")
            sheet.cell(red_row, 2).value = 5
            sheet.cell(orange_row, 2).value = 4

        path = self._mutated_fixture(mutate)
        with self.assertRaisesRegex(MingGeError, "启用品质规则优先级语义必须依次为"):
            read_mingge_workbook(path)

    def test_missing_required_sheet_is_rejected(self) -> None:
        with self.assertRaisesRegex(MingGeError, "缺少工作表：品质规则"):
            read_mingge_workbook(FIXTURES / "missing-sheet.xlsx")

    def test_bad_header_is_rejected(self) -> None:
        path = self._mutated_fixture(
            lambda workbook: self._set_cell(workbook, "系统设置", "A1", "错误列名")
        )
        with self.assertRaisesRegex(MingGeError, "系统设置.*表头"):
            read_mingge_workbook(path)

    def test_slot_outside_the_eight_slot_range_is_rejected(self) -> None:
        path = self._mutated_fixture(
            lambda workbook: self._set_cell(workbook, "词条定义", "B2", 9)
        )
        with self.assertRaisesRegex(MingGeError, "槽位号.*1～8"):
            read_mingge_workbook(path)

    def test_incomplete_quality_coverage_is_rejected(self) -> None:
        def mutate(workbook):
            workbook["品质属性"].delete_rows(2)

        path = self._mutated_fixture(mutate)
        with self.assertRaisesRegex(MingGeError, "品质.*完整覆盖"):
            read_mingge_workbook(path)

    def test_duplicate_affix_id_is_rejected(self) -> None:
        path = self._mutated_fixture(
            lambda workbook: self._set_cell(workbook, "词条定义", "A3", "MG01")
        )
        with self.assertRaisesRegex(MingGeError, "词条ID重复"):
            read_mingge_workbook(path)

    def test_non_positive_probability_denominator_is_rejected(self) -> None:
        def mutate(workbook):
            sheet = workbook["品质规则"]
            sheet["G5"] = 0

        path = self._mutated_fixture(mutate)
        with self.assertRaisesRegex(MingGeError, "概率规则.*概率分母必须是正整数"):
            read_mingge_workbook(path)

    def test_non_positive_cumulative_threshold_is_rejected(self) -> None:
        path = self._mutated_fixture(
            lambda workbook: self._set_cell(workbook, "品质规则", "E2", 0)
        )
        with self.assertRaisesRegex(MingGeError, "累计阈值规则.*累计阈值必须是正整数"):
            read_mingge_workbook(path)

    def test_non_positive_color_cycle_fields_are_rejected(self) -> None:
        for cell in ("F3", "G3"):
            with self.subTest(cell=cell):
                path = self._mutated_fixture(
                    lambda workbook, target=cell: self._set_cell(
                        workbook, "品质规则", target, 0
                    )
                )
                with self.assertRaisesRegex(
                    MingGeError, "周期彩色判定规则.*周期和概率分母必须是正整数"
                ):
                    read_mingge_workbook(path)

    def test_missing_green_fallback_is_rejected(self) -> None:
        def mutate(workbook):
            sheet = workbook["品质规则"]
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == "MG_RULE_GREEN_FALLBACK":
                    sheet.delete_rows(row)
                    break

        path = self._mutated_fixture(mutate)
        with self.assertRaisesRegex(MingGeError, "绿色兜底规则必须恰好有一条"):
            read_mingge_workbook(path)

    def test_duplicate_green_fallback_is_rejected(self) -> None:
        def mutate(workbook):
            sheet = workbook["品质规则"]
            sheet.append(("MG_RULE_GREEN_2", 7, "绿", "绿色兜底", 0, None, None, "", "是"))

        path = self._mutated_fixture(mutate)
        with self.assertRaisesRegex(MingGeError, "绿色兜底规则必须恰好有一条"):
            read_mingge_workbook(path)

    def test_non_positive_cost_is_rejected(self) -> None:
        path = self._mutated_fixture(
            lambda workbook: self._set_cell(workbook, "洗练消耗", "D2", 0)
        )
        with self.assertRaisesRegex(MingGeError, "消耗数量必须是正整数"):
            read_mingge_workbook(path)

    def test_batch_quality_boundaries_match_current_server(self) -> None:
        rules = read_mingge_workbook(FIXTURES / "valid.xlsx").quality_rules
        miss = lambda denominator: False
        green = resolve_quality(QualityProgress(0, 0, 0), 10, rules, miss)
        red_batch = resolve_quality(QualityProgress(10, 0, 0), 10, rules, miss)
        below_red = resolve_quality(QualityProgress(9, 0, 0), 10, rules, miss)
        red_single = resolve_quality(QualityProgress(19, 0, 0), 1, rules, miss)
        color_single = resolve_quality(QualityProgress(0, 99, 0), 1, rules, miss)
        self.assertEqual((green.quality, green.matched_rule_id), ("绿", "MG_RULE_GREEN_FALLBACK"))
        self.assertEqual((red_batch.quality, red_batch.matched_rule_id), ("红", "MG_RULE_RED_PITY"))
        self.assertEqual((below_red.quality, below_red.matched_rule_id), ("绿", "MG_RULE_GREEN_FALLBACK"))
        self.assertEqual((red_single.quality, red_single.matched_rule_id), ("红", "MG_RULE_RED_PITY"))
        self.assertEqual((color_single.quality, color_single.matched_rule_id), ("彩", "MG_RULE_COLOR_PITY"))

    def test_ten_draws_add_ten_and_settle_quality_once(self) -> None:
        rules = read_mingge_workbook(FIXTURES / "valid.xlsx").quality_rules
        denominators = []

        def miss(denominator):
            denominators.append(denominator)
            return False

        result = resolve_quality(QualityProgress(0, 0, 0), 10, rules, miss)
        self.assertEqual(result.progress, QualityProgress(10, 10, 10))
        self.assertEqual(denominators, [10, 50, 10])

    def test_color_cycle_at_ten_uses_table_rule_and_resets_all_progress(self) -> None:
        rules = read_mingge_workbook(FIXTURES / "valid.xlsx").quality_rules
        result = resolve_quality(
            QualityProgress(0, 0, 0), 10, rules, lambda denominator: denominator == 10
        )
        self.assertEqual((result.quality, result.matched_rule_id), ("彩", "MG_RULE_COLOR_CYCLE"))
        self.assertEqual(result.progress, QualityProgress(0, 0, 0))

    def test_mutated_color_cycle_parameters_drive_resolution(self) -> None:
        def mutate(workbook):
            sheet = workbook["品质规则"]
            row = self._rule_row(workbook, "MG_RULE_COLOR_CYCLE")
            sheet.cell(row, 1).value = "ALT_COLOR_CYCLE"
            sheet.cell(row, 7).value = 7
            sheet.cell(row, 8).value = "彩色进度"

        rules = read_mingge_workbook(self._mutated_fixture(mutate)).quality_rules
        denominators = []

        def hit(denominator):
            denominators.append(denominator)
            return denominator == 7

        result = resolve_quality(QualityProgress(3, 4, 0), 10, rules, hit)
        self.assertEqual(denominators, [7])
        self.assertEqual((result.quality, result.matched_rule_id), ("彩", "ALT_COLOR_CYCLE"))
        self.assertEqual(result.progress, QualityProgress(13, 0, 10))

    def test_natural_red_uses_table_probability_rule(self) -> None:
        rules = read_mingge_workbook(FIXTURES / "valid.xlsx").quality_rules
        result = resolve_quality(
            QualityProgress(0, 0, 0),
            1,
            rules,
            lambda denominator: denominator in {50, 10},
        )
        self.assertEqual((result.quality, result.matched_rule_id), ("红", "MG_RULE_RED_NATURAL"))

    def test_natural_orange_uses_table_probability_rule(self) -> None:
        rules = read_mingge_workbook(FIXTURES / "valid.xlsx").quality_rules
        result = resolve_quality(
            QualityProgress(0, 0, 0), 1, rules, lambda denominator: denominator == 10
        )
        self.assertEqual((result.quality, result.matched_rule_id), ("橙", "MG_RULE_ORANGE_NATURAL"))

    def test_mutated_natural_rule_parameters_drive_resolution(self) -> None:
        def mutate(workbook):
            sheet = workbook["品质规则"]
            red_row = self._rule_row(workbook, "MG_RULE_RED_NATURAL")
            orange_row = self._rule_row(workbook, "MG_RULE_ORANGE_NATURAL")
            sheet.cell(red_row, 1).value = "ALT_RED_NATURAL"
            sheet.cell(red_row, 7).value = 37
            sheet.cell(red_row, 8).value = "彩色进度"
            sheet.cell(orange_row, 1).value = "ALT_ORANGE_NATURAL"
            sheet.cell(orange_row, 7).value = 8
            sheet.cell(orange_row, 8).value = "十连进度"

        rules = read_mingge_workbook(self._mutated_fixture(mutate)).quality_rules
        red_calls = []

        def hit_red(denominator):
            red_calls.append(denominator)
            return denominator == 37

        red = resolve_quality(QualityProgress(2, 3, 4), 1, rules, hit_red)
        self.assertEqual(red_calls, [37])
        self.assertEqual((red.quality, red.matched_rule_id), ("红", "ALT_RED_NATURAL"))
        self.assertEqual(red.progress, QualityProgress(3, 0, 5))

        orange_calls = []

        def hit_orange(denominator):
            orange_calls.append(denominator)
            return denominator == 8

        orange = resolve_quality(QualityProgress(2, 3, 4), 1, rules, hit_orange)
        self.assertEqual(orange_calls, [37, 8])
        self.assertEqual((orange.quality, orange.matched_rule_id), ("橙", "ALT_ORANGE_NATURAL"))
        self.assertEqual(orange.progress, QualityProgress(3, 4, 0))

    def test_all_probability_misses_use_table_green_fallback(self) -> None:
        rules = read_mingge_workbook(FIXTURES / "valid.xlsx").quality_rules
        result = resolve_quality(QualityProgress(0, 0, 0), 1, rules, lambda denominator: False)
        self.assertEqual((result.quality, result.matched_rule_id), ("绿", "MG_RULE_GREEN_FALLBACK"))

    def test_resolve_quality_rejects_missing_green_fallback(self) -> None:
        rules = tuple(
            rule
            for rule in read_mingge_workbook(FIXTURES / "valid.xlsx").quality_rules
            if rule.trigger_type != "绿色兜底"
        )
        with self.assertRaisesRegex(MingGeError, "绿色兜底规则必须恰好有一条"):
            resolve_quality(QualityProgress(0, 0, 0), 1, rules, lambda denominator: False)


class MingGeImageRendererTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.output = Path(temporary.name)
        self.book = read_mingge_workbook(FIXTURES / "valid.xlsx")

    def test_renderer_outputs_quality_major_frames_and_real_rainbow_pixels(self) -> None:
        frames = render_mingge_labels(self.book, self.output / "frames")
        expected_labels = tuple(
            f"{affix.name}·{quality}"
            for quality in QUALITIES
            for affix in self.book.affixes
            if affix.enabled
        )
        self.assertEqual(len(frames), 32)
        self.assertEqual(tuple(frame.label for frame in frames), expected_labels)
        self.assertEqual(tuple(frame.line_no for frame in frames), tuple(range(1, 33)))
        self.assertEqual(tuple(frame.frame_id for frame in frames), tuple(range(32)))
        self.assertEqual(
            tuple(frame.png_path.name for frame in frames),
            tuple(f"{frame_id:03}.png" for frame_id in range(32)),
        )

        for frame, expected_label in zip(frames, expected_labels, strict=True):
            with self.subTest(frame_id=frame.frame_id):
                self.assertEqual(frame.label, expected_label)
                self.assertNotRegex(frame.label, r"\d")
                image = Image.open(frame.png_path).convert("RGB")
                self.assertEqual(image.mode, "RGB")
                style = next(
                    style for style in self.book.styles if style.quality == expected_label[-1]
                )
                self.assertEqual(image.height, style.canvas_height)
                self.assertLessEqual(image.width, 512)
                self.assertEqual(image.getpixel((0, 0)), (0, 0, 0))
                self.assertTrue(
                    any(pixel != (0, 0, 0) for pixel in image.get_flattened_data())
                )

        for quality_index, quality in enumerate(("绿", "橙", "红")):
            with self.subTest(quality=quality):
                frame = frames[quality_index * len(self.book.affixes)]
                style = next(style for style in self.book.styles if style.quality == quality)
                colors = {
                    pixel
                    for pixel in Image.open(frame.png_path)
                    .convert("RGB")
                    .get_flattened_data()
                    if pixel != (0, 0, 0)
                }
                self.assertIn(ImageColor.getrgb(style.colors[0]), colors)

        rainbow = Image.open(frames[3 * len(self.book.affixes)].png_path).convert("RGB")
        rainbow_colors = {
            pixel for pixel in rainbow.get_flattened_data() if pixel != (0, 0, 0)
        }
        self.assertGreater(len(rainbow_colors), 8)

    def test_renderer_uses_enabled_affix_quality_combinations_not_a_fixed_frame_count(self) -> None:
        affixes = tuple(
            replace(affix, enabled=False) if affix.slot == 8 else affix
            for affix in self.book.affixes
        )
        frames = render_mingge_labels(
            replace(self.book, affixes=affixes), self.output / "enabled-only"
        )
        self.assertEqual(len(frames), 28)
        self.assertEqual(tuple(frame.frame_id for frame in frames), tuple(range(28)))
        self.assertEqual(tuple(frame.line_no for frame in frames), tuple(range(1, 29)))

    def test_renderer_builds_hover_text_from_the_matching_table_rows_in_order(self) -> None:
        frames = render_mingge_labels(self.book, self.output / "hover")
        attributes = tuple(self.book.attributes)
        for frame in frames:
            affix_name, quality = frame.label.rsplit("·", 1)
            affix = next(affix for affix in self.book.affixes if affix.name == affix_name)
            expected = "\n".join(
                attribute.hover_text
                for attribute in sorted(
                    (
                        attribute
                        for attribute in attributes
                        if attribute.affix_id == affix.affix_id
                        and attribute.quality == quality
                    ),
                    key=lambda attribute: attribute.order,
                )
            )
            with self.subTest(frame_id=frame.frame_id):
                self.assertEqual(frame.hover_text, expected)
                self.assertNotIn(frame.hover_text, frame.label)

    def test_renderer_is_deterministic_for_the_same_workbook(self) -> None:
        first = render_mingge_labels(self.book, self.output / "first")
        second = render_mingge_labels(self.book, self.output / "second")
        self.assertEqual(
            tuple(frame.sha256 for frame in first),
            tuple(frame.sha256 for frame in second),
        )
        self.assertEqual(
            tuple(frame.sha256 for frame in first),
            tuple(hashlib.sha256(frame.png_path.read_bytes()).hexdigest() for frame in first),
        )

    def test_renderer_rejects_missing_font_illegal_color_small_canvas_and_text_overflow(self) -> None:
        cases = (
            (
                "font",
                replace(self.book, styles=(replace(self.book.styles[0], font_name="Missing Font"),) + self.book.styles[1:]),
                "字体",
            ),
            (
                "color",
                replace(self.book, styles=(replace(self.book.styles[0], colors=("#not-a-color",)),) + self.book.styles[1:]),
                "颜色",
            ),
            (
                "canvas",
                replace(self.book, styles=(replace(self.book.styles[0], canvas_height=1),) + self.book.styles[1:]),
                "画布",
            ),
            (
                "overflow",
                replace(self.book, styles=(replace(self.book.styles[0], horizontal_padding=10_000),) + self.book.styles[1:]),
                "文字",
            ),
        )
        for name, book, message in cases:
            with self.subTest(case=name):
                with self.assertRaisesRegex(MingGeImageError, message):
                    render_mingge_labels(book, self.output / name)

    def test_renderer_rejects_negative_horizontal_padding_before_writing_pngs(self) -> None:
        book = replace(
            self.book,
            styles=(
                replace(self.book.styles[0], horizontal_padding=-10),
            )
            + self.book.styles[1:],
        )
        output = self.output / "negative-padding"
        with self.assertRaisesRegex(MingGeImageError, "左右留白"):
            render_mingge_labels(book, output)
        self.assertEqual(tuple(output.glob("*.png")), ())

    def test_renderer_keeps_the_outlined_glyph_inside_canvas_bounds(self) -> None:
        book = replace(
            self.book,
            styles=(
                replace(self.book.styles[0], horizontal_padding=1),
            )
            + self.book.styles[1:],
        )
        frame = render_mingge_labels(book, self.output / "outline-bounds")[0]
        image = Image.open(frame.png_path).convert("RGB")
        glyph_pixels = [
            (x, y)
            for y in range(image.height)
            for x in range(image.width)
            if image.getpixel((x, y)) != (0, 0, 0)
        ]
        self.assertTrue(glyph_pixels)
        xs, ys = zip(*glyph_pixels, strict=True)
        self.assertGreater(min(xs), 0)
        self.assertLess(max(xs), image.width - 1)
        self.assertGreater(min(ys), 0)
        self.assertLess(max(ys), image.height - 1)

    def test_vendor_pillow_is_importable_from_platform_vendor_root(self) -> None:
        environment = os.environ.copy()
        environment["PYTHONPATH"] = str(VENDOR_ROOT)
        environment["PYTHONDONTWRITEBYTECODE"] = "1"
        code = (
            "from io import BytesIO\n"
            "from pathlib import Path\n"
            "import PIL\n"
            "from PIL import Image, ImageDraw, ImageFont\n"
            "font = ImageFont.truetype(r'C:\\Windows\\Fonts\\msyhbd.ttc', 12)\n"
            "image = Image.new('RGB', (32, 20), (0, 0, 0))\n"
            "ImageDraw.Draw(image).text((1, 1), 'A', font=font, fill=(255, 255, 255))\n"
            "buffer = BytesIO()\n"
            "image.save(buffer, format='PNG')\n"
            "print(PIL.__version__)\n"
            "print(Path(PIL.__file__).resolve())\n"
            "print(image.mode)\n"
            "print(sum(pixel != (0, 0, 0) for pixel in image.get_flattened_data()))\n"
            "print(len(buffer.getvalue()))\n"
        )
        result = subprocess.run(
            (
                sys.executable,
                "-S",
                "-c",
                code,
            ),
            capture_output=True,
            check=False,
            encoding="utf-8",
            env=environment,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        version, package_path, mode, nonblack_pixels, png_bytes = result.stdout.splitlines()
        self.assertEqual(version, "12.3.0")
        self.assertTrue(Path(package_path).is_relative_to(VENDOR_ROOT.resolve()))
        self.assertEqual(mode, "RGB")
        self.assertGreater(int(nonblack_pixels), 0)
        self.assertGreater(int(png_bytes), 8)

    def test_vendor_readme_and_both_build_targets_register_pillow(self) -> None:
        readme = (VENDOR_ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn("Pillow 12.3.0", readme)
        self.assertIn("codex-primary-runtime", readme)
        wrapper = (PAYLOAD_ROOT / "build.ps1").read_text(encoding="utf-8")
        self.assertIn("tools\\build_platform_release.py", wrapper)
        builder = (PAYLOAD_ROOT / "tools/build_platform_release.py").read_text(encoding="utf-8")
        self.assertIn("'gui','run_gui.py'", builder)
        self.assertIn("'cli','run_cli.py'", builder)
        self.assertIn("'--collect-all','PIL'", builder)
        self.assertIn("gui_result.get('status') != 'passed'", builder)



class MingGeResourceBuilderTests(unittest.TestCase):
    @staticmethod
    def _provider_request(request: dict) -> dict:
        result = subprocess.run(
            (str(PROVIDER_PATH), "run", "--request", "-"),
            input=json.dumps(request, ensure_ascii=False, separators=(",", ":")),
            capture_output=True,
            check=False,
            encoding="utf-8",
            timeout=120,
        )
        if result.returncode != 0:
            raise AssertionError(result.stderr or result.stdout)
        payload = json.loads(result.stdout)
        if payload.get("ok") is not True:
            raise AssertionError(payload)
        return payload

    @staticmethod
    def _render(root: Path):
        return render_mingge_labels(
            read_mingge_workbook(FIXTURES / "valid.xlsx"), root / "frames"
        )

    @staticmethod
    def _cross_process_build_code() -> str:
        return (
            "import json,sys\n"
            "from pathlib import Path\n"
            "root=Path(sys.argv[1]); output=Path(sys.argv[2]); worker=sys.argv[3]\n"
            "sys.path.insert(0,str(root/'src'))\n"
            "from xydp.mingge import build_mingge_library,read_mingge_workbook\n"
            "from xydp.mingge_image import render_mingge_labels\n"
            "book=root/'tests'/'fixtures'/'mingge'/'valid.xlsx'\n"
            "provider=root/'wzl编辑'/'bin'/'XuanYuanWzlProvider'/'XuanYuanWzlProvider.exe'\n"
            "template=root/'wzl编辑'/'workspace'/'mingge-template'\n"
            "frames=render_mingge_labels(read_mingge_workbook(book),output.parent/f'frames-{worker}')\n"
            "library=build_mingge_library(frames,output,provider,workbook_path=book,template_dir=template)\n"
            "print(json.dumps({'config_hash':library.config_hash,'wzl_hash':library.wzl_hash,'wzx_hash':library.wzx_hash}))\n"
        )

    def test_real_provider_builds_all_32_frames_with_contiguous_mapping(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)

            library = mingge_module.build_mingge_library(
                frames,
                root / "library",
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=TEMPLATE_DIR,
            )

            self.assertEqual(len(library.mapping), 32)
            self.assertEqual(
                tuple(item["source_id"] for item in library.mapping), tuple(range(32))
            )
            self.assertEqual(
                tuple(item["target_id"] for item in library.mapping), tuple(range(32))
            )
            self.assertTrue(library.wzl_path.is_file())
            self.assertTrue(library.wzx_path.is_file())

    def test_real_provider_is_deterministic_and_exports_four_qualities_pixel_exactly(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            template_paths = (TEMPLATE_DIR / "Template.wzl", TEMPLATE_DIR / "Template.wzx")
            template_before = tuple(
                hashlib.sha256(path.read_bytes()).hexdigest() for path in template_paths
            )
            png_before = tuple(
                hashlib.sha256(frame.png_path.read_bytes()).hexdigest() for frame in frames
            )

            first = mingge_module.build_mingge_library(
                frames,
                root / "first",
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=TEMPLATE_DIR,
            )
            second = mingge_module.build_mingge_library(
                frames,
                root / "second",
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=TEMPLATE_DIR,
            )

            self.assertEqual(first.config_hash, second.config_hash)
            self.assertEqual(first.wzl_path.read_bytes(), second.wzl_path.read_bytes())
            self.assertEqual(first.wzx_path.read_bytes(), second.wzx_path.read_bytes())
            self.assertEqual(
                first.wzl_hash, hashlib.sha256(first.wzl_path.read_bytes()).hexdigest()
            )
            self.assertEqual(
                first.wzx_hash, hashlib.sha256(first.wzx_path.read_bytes()).hexdigest()
            )
            inspect_request = {
                "schemaVersion": 1,
                "requestId": "task5-test-inspect-32",
                "command": "inspect",
                "arguments": {
                    "wzl": str(first.wzl_path),
                    "wzx": str(first.wzx_path),
                    "selection": {"ids": list(range(32))},
                },
                "policy": {"allowedWriteRoots": []},
            }
            inspected = self._provider_request(inspect_request)["data"]
            self.assertEqual(inspected["count"], 32)
            self.assertEqual(
                tuple(entry["image_id"] for entry in inspected["entries"]), tuple(range(32))
            )
            self.assertTrue(all(entry["frame_type"] == 6 for entry in inspected["entries"]))

            for quality, frame_id in zip(QUALITIES, (0, 8, 16, 24), strict=True):
                with self.subTest(quality=quality):
                    export_dir = root / "exports" / quality
                    exported = self._provider_request(
                        {
                            "schemaVersion": 1,
                            "requestId": f"task5-test-export-{frame_id}",
                            "command": "export",
                            "arguments": {
                                "wzl": str(first.wzl_path),
                                "wzx": str(first.wzx_path),
                                "id": frame_id,
                                "output": str(export_dir),
                            },
                            "policy": {"allowedWriteRoots": [str(export_dir)]},
                        }
                    )["data"]
                    with Image.open(frames[frame_id].png_path) as source, Image.open(
                        exported["png_path"]
                    ) as readback:
                        source_rgb = source.convert("RGB")
                        readback_rgb = readback.convert("RGB")
                        self.assertEqual(source_rgb.size, readback_rgb.size)
                        self.assertEqual(source_rgb.tobytes(), readback_rgb.tobytes())

            self.assertEqual(
                template_before,
                tuple(hashlib.sha256(path.read_bytes()).hexdigest() for path in template_paths),
            )
            self.assertEqual(
                png_before,
                tuple(hashlib.sha256(frame.png_path.read_bytes()).hexdigest() for frame in frames),
            )

    def test_config_hash_covers_workbook_png_template_and_provider_version(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            baseline = mingge_module.build_mingge_library(
                frames,
                root / "baseline",
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=TEMPLATE_DIR,
            )

            workbook = root / "changed.xlsx"
            workbook.write_bytes((FIXTURES / "valid.xlsx").read_bytes() + b"task5")
            workbook_changed = mingge_module.build_mingge_library(
                frames,
                root / "workbook",
                PROVIDER_PATH,
                workbook_path=workbook,
                template_dir=TEMPLATE_DIR,
            )

            changed_png = root / "changed.png"
            with Image.open(frames[0].png_path) as image:
                image = image.convert("RGB")
                image.putpixel((0, 0), (1, 2, 3))
                image.save(changed_png, format="PNG")
            changed_frames = (
                replace(
                    frames[0],
                    png_path=changed_png,
                    sha256=hashlib.sha256(changed_png.read_bytes()).hexdigest(),
                ),
            ) + frames[1:]
            png_changed = mingge_module.build_mingge_library(
                changed_frames,
                root / "png",
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=TEMPLATE_DIR,
            )

            changed_template = root / "template"
            changed_template.mkdir()
            shutil.copy2(TEMPLATE_DIR / "Template.wzl", changed_template / "Template.wzl")
            shutil.copy2(TEMPLATE_DIR / "Template.wzx", changed_template / "Template.wzx")
            os.chmod(changed_template / "Template.wzl", 0o666)
            os.chmod(changed_template / "Template.wzx", 0o666)
            template_bytes = bytearray((changed_template / "Template.wzl").read_bytes())
            template_bytes[100] ^= 1
            (changed_template / "Template.wzl").write_bytes(template_bytes)
            template_changed = mingge_module.build_mingge_library(
                frames,
                root / "template-output",
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=changed_template,
            )

            real_process = mingge_module._provider_process

            def versioned_provider(provider_path, arguments, request=None):
                payload = deepcopy(real_process(provider_path, arguments, request))
                payload["providerVersion"] = "1.1.1-task5-test"
                return payload

            with mock.patch.object(
                mingge_module, "_provider_process", side_effect=versioned_provider
            ):
                provider_changed = mingge_module.build_mingge_library(
                    frames,
                    root / "provider-version",
                    PROVIDER_PATH,
                    workbook_path=FIXTURES / "valid.xlsx",
                    template_dir=TEMPLATE_DIR,
                )

            self.assertEqual(
                len(
                    {
                        baseline.config_hash,
                        workbook_changed.config_hash,
                        png_changed.config_hash,
                        template_changed.config_hash,
                        provider_changed.config_hash,
                    }
                ),
                5,
            )

    def test_provider_consumes_snapshot_when_original_inputs_are_replaced_then_restored(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            source_template = root / "source-template"
            source_template.mkdir()
            shutil.copyfile(TEMPLATE_DIR / "Template.wzl", source_template / "Template.wzl")
            shutil.copyfile(TEMPLATE_DIR / "Template.wzx", source_template / "Template.wzx")
            baseline = mingge_module.build_mingge_library(
                frames,
                root / "baseline",
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=source_template,
            )

            original_template = (source_template / "Template.wzl").read_bytes()
            original_png = frames[0].png_path.read_bytes()
            changed_template = bytearray(original_template)
            changed_template[0] ^= 1
            changed_png_path = root / "replacement.png"
            with Image.open(frames[0].png_path) as image:
                replacement = image.convert("RGB")
                replacement.putpixel((0, 0), (7, 8, 9))
                replacement.save(changed_png_path, format="PNG")
            changed_png = changed_png_path.read_bytes()
            attacked_output = root / "attacked"
            captured_paths = {}
            real_process = mingge_module._provider_process

            def replace_originals_during_provider(provider_path, arguments, request=None):
                if request and request.get("command") == "build-image-library":
                    captured_paths["template"] = Path(request["arguments"]["templateWzl"])
                    captured_paths["png"] = Path(request["arguments"]["images"][0]["path"])
                    (source_template / "Template.wzl").write_bytes(changed_template)
                    frames[0].png_path.write_bytes(changed_png)
                    try:
                        return real_process(provider_path, arguments, request)
                    finally:
                        (source_template / "Template.wzl").write_bytes(original_template)
                        frames[0].png_path.write_bytes(original_png)
                return real_process(provider_path, arguments, request)

            with mock.patch.object(
                mingge_module,
                "_provider_process",
                side_effect=replace_originals_during_provider,
            ):
                attacked = mingge_module.build_mingge_library(
                    frames,
                    attacked_output,
                    PROVIDER_PATH,
                    workbook_path=FIXTURES / "valid.xlsx",
                    template_dir=source_template,
                )

            self.assertEqual(attacked.config_hash, baseline.config_hash)
            self.assertEqual(attacked.wzl_hash, baseline.wzl_hash)
            self.assertEqual(attacked.wzx_hash, baseline.wzx_hash)
            self.assertTrue(
                captured_paths["template"].is_relative_to(attacked.wzl_path.parent)
            )
            self.assertTrue(captured_paths["png"].is_relative_to(attacked.wzl_path.parent))
            self.assertNotEqual(captured_paths["template"], source_template / "Template.wzl")
            self.assertNotEqual(captured_paths["png"], frames[0].png_path)
            self.assertEqual(
                Path(attacked.mapping[0]["source_path"]), frames[0].png_path.resolve()
            )
            self.assertEqual((source_template / "Template.wzl").read_bytes(), original_template)
            self.assertEqual(frames[0].png_path.read_bytes(), original_png)
            self.assertEqual(tuple(attacked_output.glob(".mingge-snapshot-*")), ())

    def test_protocol_mismatches_are_blocked_without_partial_resource_pair(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            real_process = mingge_module._provider_process
            mutations = {
                "schema": lambda payload: payload.__setitem__("schemaVersion", 2),
                "provider": lambda payload: payload.__setitem__("provider", "wrong-provider"),
                "protocol": lambda payload: payload.__setitem__("protocol", "wrong/1"),
                "version": lambda payload: payload.__setitem__("providerVersion", "0.0.0"),
                "request": lambda payload: payload.__setitem__("requestId", "wrong"),
                "command": lambda payload: payload.__setitem__("command", "inspect"),
                "ok": lambda payload: payload.__setitem__("ok", False),
                "count": lambda payload: payload["data"].__setitem__("count", 31),
                "readback": lambda payload: payload["data"].__setitem__("readback_ok", False),
                "mapping": lambda payload: payload["data"]["mapping"][0].__setitem__("target_id", 1),
                "source-path": lambda payload: payload["data"]["mapping"][0].__setitem__("source_path", r"C:\wrong.png"),
                "mapping-size": lambda payload: payload["data"]["mapping"][0].__setitem__("width", 999),
                "source-hash": lambda payload: payload["data"]["mapping"][0].__setitem__("source_sha256", "0" * 64),
                "output-path": lambda payload: payload["data"]["output_pair"].__setitem__("wzl", r"C:\wrong.wzl"),
                "output-hash": lambda payload: payload["data"]["output_pair"].__setitem__("wzlSha256", "0" * 64),
            }
            for name, mutate in mutations.items():
                with self.subTest(case=name):
                    output = root / name

                    def changed_provider(provider_path, arguments, request=None, change=mutate):
                        payload = deepcopy(real_process(provider_path, arguments, request))
                        if request and request.get("command") == "build-image-library":
                            change(payload)
                        return payload

                    with mock.patch.object(
                        mingge_module, "_provider_process", side_effect=changed_provider
                    ):
                        with self.assertRaises(mingge_module.MingGeResourceError):
                            mingge_module.build_mingge_library(
                                frames,
                                output,
                                PROVIDER_PATH,
                                workbook_path=FIXTURES / "valid.xlsx",
                                template_dir=TEMPLATE_DIR,
                            )
                    self.assertEqual(tuple(output.glob("*.wzl")), ())
                    self.assertEqual(tuple(output.glob("*.wzx")), ())

    def test_nonzero_timeout_and_non_json_are_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            cases = (
                subprocess.CompletedProcess(("provider",), 4, "", "failed"),
                subprocess.TimeoutExpired(("provider",), 120),
                subprocess.CompletedProcess(("provider",), 0, "not-json", ""),
                UnicodeDecodeError("utf-8", b"\xff", 0, 1, "invalid start byte"),
            )
            for index, result in enumerate(cases):
                with self.subTest(case=index), mock.patch.object(
                    mingge_module.subprocess,
                    "run",
                    side_effect=result if isinstance(result, Exception) else None,
                    return_value=None if isinstance(result, Exception) else result,
                ):
                    output = root / f"failure-{index}"
                    with self.assertRaises(mingge_module.MingGeResourceError):
                        mingge_module.build_mingge_library(
                            frames,
                            output,
                            PROVIDER_PATH,
                            workbook_path=FIXTURES / "valid.xlsx",
                            template_dir=TEMPLATE_DIR,
                        )
                    self.assertEqual(tuple(output.glob("*.wzl")), ())
                    self.assertEqual(tuple(output.glob("*.wzx")), ())

    def test_second_publish_failure_removes_the_first_resource_file(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            output = root / "publish-failure"
            real_replace = mingge_module.os.replace
            calls = 0

            def fail_second_publish(source, destination):
                nonlocal calls
                calls += 1
                if calls == 2:
                    raise OSError("task5 simulated second publish failure")
                return real_replace(source, destination)

            with mock.patch.object(
                mingge_module.os, "replace", side_effect=fail_second_publish
            ):
                with self.assertRaises(mingge_module.MingGeResourceError):
                    mingge_module.build_mingge_library(
                        frames,
                        output,
                        PROVIDER_PATH,
                        workbook_path=FIXTURES / "valid.xlsx",
                        template_dir=TEMPLATE_DIR,
                    )
            self.assertEqual(tuple(output.glob("*.wzl")), ())
            self.assertEqual(tuple(output.glob("*.wzx")), ())

    def test_publish_failure_does_not_delete_pair_replaced_by_another_builder(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            changed_png = root / "other.png"
            with Image.open(frames[0].png_path) as image:
                changed = image.convert("RGB")
                changed.putpixel((0, 0), (11, 12, 13))
                changed.save(changed_png, format="PNG")
            other_frames = (
                replace(
                    frames[0],
                    png_path=changed_png,
                    sha256=hashlib.sha256(changed_png.read_bytes()).hexdigest(),
                ),
            ) + frames[1:]
            other = mingge_module.build_mingge_library(
                other_frames,
                root / "other",
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=TEMPLATE_DIR,
            )
            other_wzl = other.wzl_path.read_bytes()
            other_wzx = other.wzx_path.read_bytes()
            output = root / "target"
            real_replace = mingge_module.os.replace
            calls = 0

            def replace_pair_then_fail(source, destination):
                nonlocal calls
                calls += 1
                destination = Path(destination)
                if calls == 2:
                    destination.with_suffix(".wzl").write_bytes(other_wzl)
                    destination.write_bytes(other_wzx)
                    raise OSError("task5 simulated replacement by another builder")
                return real_replace(source, destination)

            with mock.patch.object(
                mingge_module.os, "replace", side_effect=replace_pair_then_fail
            ):
                with self.assertRaises(mingge_module.MingGeResourceError):
                    mingge_module.build_mingge_library(
                        frames,
                        output,
                        PROVIDER_PATH,
                        workbook_path=FIXTURES / "valid.xlsx",
                        template_dir=TEMPLATE_DIR,
                    )

            self.assertEqual(len(tuple(output.glob("*.wzl"))), 1)
            self.assertEqual(len(tuple(output.glob("*.wzx"))), 1)
            self.assertEqual(next(output.glob("*.wzl")).read_bytes(), other_wzl)
            self.assertEqual(next(output.glob("*.wzx")).read_bytes(), other_wzx)

    def test_windows_lock_blocks_publication_and_is_released_when_holder_exits(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            baseline = mingge_module.build_mingge_library(
                frames,
                root / "baseline",
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=TEMPLATE_DIR,
            )
            output = root / "locked-output"
            output.mkdir()
            stem = baseline.wzl_path.stem
            lock_path = output / f".{stem}.lock"
            holder_code = (
                "import sys,time\n"
                "from pathlib import Path\n"
                f"sys.path.insert(0,{str(PAYLOAD_ROOT / 'src')!r})\n"
                "from xydp.mingge import _mingge_output_lock\n"
                "with _mingge_output_lock(Path(sys.argv[1])):\n"
                " print('locked',flush=True)\n"
                " time.sleep(60)\n"
            )
            holder = subprocess.Popen(
                (sys.executable, "-c", holder_code, str(lock_path)),
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                encoding="utf-8",
            )
            builder = None
            try:
                self.assertEqual(holder.stdout.readline().strip(), "locked")
                builder = subprocess.Popen(
                    (
                        sys.executable,
                        "-c",
                        self._cross_process_build_code(),
                        str(PAYLOAD_ROOT),
                        str(output),
                        "locked",
                    ),
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                    encoding="utf-8",
                )
                deadline = time.monotonic() + 20
                while (
                    not tuple(output.glob(".*.staged.wzx"))
                    and builder.poll() is None
                    and time.monotonic() < deadline
                ):
                    time.sleep(0.05)
                self.assertTrue(tuple(output.glob(".*.staged.wzx")))
                self.assertIsNone(builder.poll())
                self.assertFalse((output / f"{stem}.wzl").exists())
                self.assertFalse((output / f"{stem}.wzx").exists())

                holder.kill()
                holder.wait(timeout=5)
                stdout, stderr = builder.communicate(timeout=30)
                self.assertEqual(builder.returncode, 0, stderr)
                result = json.loads(stdout)
                final_wzl = output / f"{stem}.wzl"
                final_wzx = output / f"{stem}.wzx"
                self.assertEqual(
                    hashlib.sha256(final_wzl.read_bytes()).hexdigest(), result["wzl_hash"]
                )
                self.assertEqual(
                    hashlib.sha256(final_wzx.read_bytes()).hexdigest(), result["wzx_hash"]
                )
            finally:
                if holder.poll() is None:
                    holder.kill()
                    holder.wait(timeout=5)
                if builder is not None and builder.poll() is None:
                    builder.kill()
                    builder.wait(timeout=5)
                if holder.stdout is not None:
                    holder.stdout.close()
                if holder.stderr is not None:
                    holder.stderr.close()
                if builder is not None and builder.stdout is not None:
                    builder.stdout.close()
                if builder is not None and builder.stderr is not None:
                    builder.stderr.close()

    def test_two_real_processes_same_config_leave_one_complete_consistent_pair(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            output = root / "concurrent"
            workers = [
                subprocess.Popen(
                    (
                        sys.executable,
                        "-c",
                        self._cross_process_build_code(),
                        str(PAYLOAD_ROOT),
                        str(output),
                        str(worker),
                    ),
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                    encoding="utf-8",
                )
                for worker in range(2)
            ]
            for worker in workers:
                self.addCleanup(lambda process=worker: process.poll() is None and process.kill())
            results = []
            for worker in workers:
                stdout, stderr = worker.communicate(timeout=45)
                self.assertEqual(worker.returncode, 0, stderr)
                results.append(json.loads(stdout))

            self.assertEqual(results[0], results[1])
            wzl_files = tuple(output.glob("*.wzl"))
            wzx_files = tuple(output.glob("*.wzx"))
            self.assertEqual(len(wzl_files), 1)
            self.assertEqual(len(wzx_files), 1)
            self.assertEqual(
                hashlib.sha256(wzl_files[0].read_bytes()).hexdigest(),
                results[0]["wzl_hash"],
            )
            self.assertEqual(
                hashlib.sha256(wzx_files[0].read_bytes()).hexdigest(),
                results[0]["wzx_hash"],
            )
            self.assertEqual(tuple(output.glob(".mingge-snapshot-*")), ())
            self.assertEqual(tuple(output.glob(".*.staged.*")), ())

    def test_content_addressed_pair_is_reused_but_single_or_changed_pair_is_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            output = root / "idempotent"
            first = mingge_module.build_mingge_library(
                frames,
                output,
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=TEMPLATE_DIR,
            )
            second = mingge_module.build_mingge_library(
                frames,
                output,
                PROVIDER_PATH,
                workbook_path=FIXTURES / "valid.xlsx",
                template_dir=TEMPLATE_DIR,
            )
            self.assertEqual(first, second)
            self.assertEqual(len(tuple(output.glob("*.wzl"))), 1)
            self.assertEqual(len(tuple(output.glob("*.wzx"))), 1)

            original_wzl = first.wzl_path.read_bytes()
            first.wzx_path.unlink()
            with self.assertRaisesRegex(mingge_module.MingGeResourceError, "单边资源|成对存在"):
                mingge_module.build_mingge_library(
                    frames,
                    output,
                    PROVIDER_PATH,
                    workbook_path=FIXTURES / "valid.xlsx",
                    template_dir=TEMPLATE_DIR,
                )
            self.assertEqual(first.wzl_path.read_bytes(), original_wzl)

            first.wzx_path.write_bytes(second.wzx_path.read_bytes() if second.wzx_path.exists() else b"wrong")
            first.wzl_path.write_bytes(original_wzl + b"changed")
            with self.assertRaisesRegex(mingge_module.MingGeResourceError, "禁止覆盖"):
                mingge_module.build_mingge_library(
                    frames,
                    output,
                    PROVIDER_PATH,
                    workbook_path=FIXTURES / "valid.xlsx",
                    template_dir=TEMPLATE_DIR,
                )
            self.assertEqual(first.wzl_path.read_bytes(), original_wzl + b"changed")

    def test_workbook_path_is_required_keyword_only(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frames = self._render(root)
            with self.assertRaises(TypeError):
                mingge_module.build_mingge_library(
                    frames,
                    root / "library",
                    PROVIDER_PATH,
                    FIXTURES / "valid.xlsx",
                )


class MingGeCompilerTests(unittest.TestCase):
    @staticmethod
    def _compiled():
        return compile_mingge_payload(
            read_mingge_workbook(FIXTURES / "valid.xlsx"),
            resource_index=20,
            custom_text_base=1,
            registry=REGISTRY,
        )

    @staticmethod
    def _label_block(text: str, label: str) -> str:
        normalized = text.replace("\r\n", "\n")
        match = re.search(
            rf"^\[{re.escape(label)}\]\n(.*?)(?=^\[@|\Z)", normalized, re.MULTILINE | re.DOTALL
        )
        if match is None:
            raise AssertionError(f"缺少标签段: {label}")
        return match.group(1)

    @staticmethod
    def _core(compiled) -> str:
        return compiled.files[
            PurePosixPath("Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt")
        ].decode("gb18030")

    def _assert_cost_payment_success_branch(self, cost: str, mode: int, amount: int) -> None:
        check = f"CHECKGAMEGOLD > {amount - 1}"
        pay_jump = f"DELAYGOTO 50 @XY_MG_PAY_{mode}"
        check_index = cost.index(check)
        after_check = cost[check_index + len(check):]
        self.assertTrue(after_check.startswith("\n#ACT\n"))
        else_index = cost.index("#ELSEACT", check_index)
        success_branch = cost[check_index:else_index]
        insufficient_branch = cost[else_index + len("#ELSEACT"):]
        self.assertIn(pay_jump, success_branch)
        self.assertLess(cost.index(pay_jump), else_index)
        self.assertNotIn(pay_jump, insufficient_branch)
        self.assertNotIn("@XY_MG_PROGRESS_ROUTE", insufficient_branch)
        self.assertNotIn("@XY_MG_BATCH_APPLY", insufficient_branch)
        self.assertTrue(insufficient_branch.strip().endswith("BREAK"))

    def test_compiler_keeps_native_effect_rows_and_separate_display_rows(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        compiled = compile_mingge_payload(
            book, resource_index=20, custom_text_base=1, registry=REGISTRY
        )
        core = compiled.files[
            PurePosixPath("Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt")
        ].decode("gb18030")
        self.assertIn("SetCustomItemAbil 17 1 1 1", core)
        self.assertIn("SetCustomItemAbil 17 9 1 60", core)
        self.assertIn("SetCustomItemValueEX 17 9 =", core)
        self.assertNotIn("ITEMBOX", core)
        self.assertEqual(core.count("[@XY_MG_BATCH_APPLY]"), 1)

    def test_compiler_routes_every_quality_through_the_selected_slot_write_and_tail(self) -> None:
        core = self._core(self._compiled())
        dispatch = self._label_block(core, "@XY_MG_BATCH_APPLY").replace("\r\n", "\n")
        for slot in range(1, 9):
            with self.subTest(slot=slot):
                self.assertIn(
                    f"EQUAL N$XY_MG_SLOT {slot}\n#ACT\nGOTO @XY_MG_BATCH_APPLY_SLOT_{slot}\nBREAK",
                    dispatch,
                )
                apply = self._label_block(
                    core, f"@XY_MG_BATCH_APPLY_SLOT_{slot}"
                ).replace("\r\n", "\n")
                self.assertEqual(apply.count("LockUpdateItem 17"), 1)
                self.assertIn("#IF\n#ACT\nLockUpdateItem 17", apply)
                self.assertIn(
                    "#IF\n#ACT\nUpdateItem 17\nDELAYGOTO 50 @XY_MG_DRAW_TAIL\nBREAK",
                    apply,
                )
                for quality in range(2, 5):
                    self.assertIn(f"EQUAL N$XY_MG_QUALITY {quality}", apply)

    def test_compiler_rejects_unknown_property_key(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        attributes = (replace(book.attributes[0], property_key="not_registered"),) + book.attributes[1:]
        with self.assertRaisesRegex(MingGeError, "blocker.*未知属性键"):
            compile_mingge_payload(
                replace(book, attributes=attributes), 20, 1, REGISTRY
            )

    def test_compiler_rejects_value_outside_registry_range(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        attributes = (replace(book.attributes[0], value=1_000_001),) + book.attributes[1:]
        with self.assertRaisesRegex(MingGeError, "blocker.*属性数值越界"):
            compile_mingge_payload(
                replace(book, attributes=attributes), 20, 1, REGISTRY
            )

    def test_compiler_rejects_registry_unit_conflict(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        attributes = (replace(book.attributes[0], unit="%"),) + book.attributes[1:]
        with self.assertRaisesRegex(MingGeError, "blocker.*属性单位冲突"):
            compile_mingge_payload(
                replace(book, attributes=attributes), 20, 1, REGISTRY
            )

    def test_compiler_rejects_cross_quality_property_key_drift(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        self.assertIn("custom_magic_defense", REGISTRY)
        attributes = tuple(
            replace(attribute, property_key="custom_magic_defense")
            if attribute.affix_id == "MG01" and attribute.quality == "橙"
            else attribute
            for attribute in book.attributes
        )
        with self.assertRaisesRegex(MingGeError, "blocker.*四品质.*property_key.*unit"):
            compile_mingge_payload(
                replace(book, attributes=attributes), 20, 1, REGISTRY
            )

    def test_generated_files_use_gb18030_crlf_and_safe_batch_order(self) -> None:
        compiled = self._compiled()
        for relative_path, content in compiled.files.items():
            with self.subTest(relative_path=relative_path):
                self.assertEqual(content.decode("gb18030").encode("gb18030"), content)
                self.assertNotIn(b"\n", content.replace(b"\r\n", b""))
                self.assertNotIn("????", content.decode("gb18030"))
        core = compiled.files[
            PurePosixPath("Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt")
        ].decode("gb18030")
        self.assertLess(core.index("余额或材料不足，不扣除"), core.index("GAMEGOLD - 100"))
        self.assertLess(core.index("[@XY_MG_QUALITY]"), core.index("GAMEGOLD - 1000"))
        self.assertLess(core.index("GAMEGOLD - 1000"), core.index("[@XY_MG_BATCH_APPLY]"))
        self.assertEqual(core.count("[@XY_MG_BATCH_APPLY]"), 1)
        self.assertNotIn("ITEMBOX", core)

    def test_resource_index_only_drives_hover_rows_and_default_dialog_is_usable(self) -> None:
        compiled = compile_mingge_payload(
            read_mingge_workbook(FIXTURES / "valid.xlsx"), 21, 1, REGISTRY
        )
        core = self._core(compiled)
        self.assertNotIn("OPENMERCHANTBIGDLG", core)
        panel = self._label_block(core, "@XY_MG_PANEL_8")
        self.assertIn("#SAY", panel)
        self.assertIn("/@XY_MG_SLOT8>", panel)
        self.assertEqual(len(compiled.custom_property_lines), 32)
        self.assertTrue(compiled.custom_property_lines[0].startswith("<PlayImg:21:0:"))

    def test_compiler_uses_each_label_frame_once_without_cross_panel_offset(self) -> None:
        compiled = compile_mingge_payload(
            read_mingge_workbook(FIXTURES / "valid.xlsx"), 21, 1, REGISTRY
        )
        lines = compiled.custom_property_lines
        self.assertEqual(len(lines), 32)
        self.assertEqual(
            tuple(int(line.split(":", 3)[2]) for line in lines),
            tuple(range(32)),
        )
        self.assertTrue(all(":100:0:0:1:" in line for line in lines))
        self.assertTrue(all(":-92:" not in line for line in lines))

    def test_dialog_resource_index_is_separate_from_hover_resource_index(self) -> None:
        try:
            compiled = compile_mingge_payload(
                read_mingge_workbook(FIXTURES / "valid.xlsx"),
                21,
                1,
                REGISTRY,
                dialog_resource_index=37,
            )
        except TypeError as exc:
            self.fail(f"编译器缺少独立对话框资源参数: {exc}")
        core = self._core(compiled)
        self.assertIn("OPENMERCHANTBIGDLG 37", core)
        self.assertNotIn("OPENMERCHANTBIGDLG 21", core)
        self.assertEqual(len(compiled.custom_property_lines), 32)
        self.assertTrue(compiled.custom_property_lines[0].startswith("<PlayImg:21:0:"))

    def test_large_dialog_uses_fixed_frame_layout_without_say_text_leaks(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        compiled = compile_mingge_payload(
            book,
            21,
            1,
            REGISTRY,
            dialog_resource_index=37,
        )
        core = self._core(compiled)
        panel = self._label_block(core, "@XY_MG_PANEL_8")
        self.assertIn(f"<&USERITEM:{book.settings.equip_slot}:79:102:1>", panel)
        self.assertIn("<&Text:龙魂觉醒·命格洗练:240:18{FCOLOR=253}>", panel)
        self.assertIn(
            f"<&Text:仅支持身上准确名称为{book.settings.target_item_name}的装备，命格属性写入后请悬浮查看:34:43{{FCOLOR=161}}>",
            panel,
        )
        self.assertIn("<&Text:当前选择:476:47{FCOLOR=251}>", panel)
        self.assertIn(
            "<&Text:当前槽位：<$STR(N$XY_MG_SLOT)>:450:98{FCOLOR=250}>",
            panel,
        )
        for slot, rule in enumerate(book.slots, 1):
            y = 66 + (slot - 1) * 31
            expected = (
                f"<&Text:{rule.display_name}（{rule.condition_value}级开放）:182:{y}"
                f"{{FCOLOR=250}}/@XY_MG_SLOT{slot}>"
            )
            self.assertIn(expected, panel)
        costs = {cost.mode: cost.amount for cost in book.costs if cost.enabled}
        self.assertIn(
            f"<&Text:单次洗练|249#需要{costs['单次']}元宝:270:372"
            "{FCOLOR=69}/@XY_MG_WASH_ONE>",
            panel,
        )
        self.assertIn(
            f"<&Text:十连洗练|249#需要{costs['十连']}元宝:447:372"
            "{FCOLOR=249}/@XY_MG_WASH_TEN>",
            panel,
        )
        self.assertIn("<&Text:羁绊说明:105:372{FCOLOR=250}/@XY_MG_BOND_INFO>", panel)
        self.assertNotIn("命格洗练\\", panel)
        self.assertNotIn("<关闭/@exit>", panel)
        self.assertNotIn("\nBREAK", panel)

    def test_dialog_resource_index_rejects_bool_negative_and_non_integer(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        for value in (True, -1, 1.5, "37"):
            with self.subTest(value=value):
                try:
                    with self.assertRaisesRegex(MingGeError, "对话框资源索引"):
                        compile_mingge_payload(
                            book,
                            21,
                            1,
                            REGISTRY,
                            dialog_resource_index=value,
                        )
                except TypeError as exc:
                    self.fail(f"编译器缺少独立对话框资源参数: {exc}")

    def test_compiler_commits_table_selected_progress_resets_after_payment(self) -> None:
        core = self._core(self._compiled())
        apply = self._label_block(core, "@XY_MG_BATCH_APPLY")
        self.assertIn("EQUAL N$XY_MG_RESET_RED 1\r\n#ACT\r\nMOV U201 0", core)
        self.assertIn("EQUAL N$XY_MG_RESET_COLOR 1\r\n#ACT\r\nMOV U471 0", core)
        self.assertIn("EQUAL N$XY_MG_RESET_TENTH 1\r\n#ACT\r\nMOV U491 0", core)
        self.assertLess(core.index("GAMEGOLD - 1000"), core.index("MOV U201 0"))
        self.assertNotIn("INC U", apply)

    def test_npc_text_uses_the_registered_npc_and_map_specific_entry_file(self) -> None:
        compiled = self._compiled()
        npc_path = PurePosixPath(
            "Mir200/Envir/Market_Def/玄渊命格/龙魂觉醒-3.txt"
        )
        self.assertIn(npc_path, compiled.files)
        self.assertNotIn(
            PurePosixPath("Mir200/Envir/Market_Def/玄渊命格/龙魂觉醒.txt"),
            compiled.files,
        )
        npc = compiled.files[npc_path].decode("gb18030")
        self.assertEqual(compiled.npc_text, npc)
        self.assertIn("#CALL [\\玄渊命格\\命格核心.txt] @XY_MG_MAIN", compiled.npc_text)

    def test_level_range_panels_only_expose_opened_slots_and_direct_calls_stay_gated(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        core = self._core(compile_mingge_payload(book, 20, 1, REGISTRY))
        route = self._label_block(core, "@XY_MG_PANEL_ROUTE")
        thresholds = tuple(int(slot.condition_value) for slot in book.slots)
        for opened in range(0, 9):
            with self.subTest(opened=opened):
                panel = self._label_block(core, f"@XY_MG_PANEL_{opened}")
                for slot in range(1, 9):
                    link = f"/@XY_MG_SLOT{slot}>"
                    if slot <= opened:
                        self.assertIn(link, panel)
                    else:
                        self.assertNotIn(link, panel)
        for slot, threshold in enumerate(thresholds, 1):
            with self.subTest(slot=slot):
                label = self._label_block(core, f"@XY_MG_SLOT{slot}")
                self.assertIn(f"SMALL <$LEVEL> {threshold}", label)
                self.assertLess(
                    label.index(f"SMALL <$LEVEL> {threshold}"),
                    label.index(f"MOV N$XY_MG_SLOT {slot}"),
                )
                self.assertIn("DELAYGOTO 50 @XY_MG_PANEL_ROUTE", label)
        for threshold in thresholds:
            self.assertIn(f"SMALL <$LEVEL> {threshold}", route)

    def test_unsupported_or_non_monotonic_slot_conditions_block_compilation(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        mutations = (
            replace(book.slots[0], condition_type="任务"),
            replace(book.slots[0], condition_key="转生等级"),
            replace(book.slots[0], condition_value="一级"),
            replace(book.slots[1], condition_value=book.slots[0].condition_value),
        )
        for changed in mutations:
            with self.subTest(changed=changed):
                slots = tuple(
                    changed if slot.slot == changed.slot else slot for slot in book.slots
                )
                with self.assertRaisesRegex(MingGeError, "blocker.*槽位.*条件"):
                    compile_mingge_payload(replace(book, slots=slots), 20, 1, REGISTRY)

    def test_wash_precheck_blocks_unselected_or_wrong_equipment_before_cost(self) -> None:
        core = self._core(self._compiled())
        for label in ("@XY_MG_WASH_ONE", "@XY_MG_WASH_TEN"):
            with self.subTest(label=label):
                wash = self._label_block(core, label)
                self.assertIn("DELAYGOTO 50 @XY_MG_PRECHECK", wash)
                self.assertNotIn("@XY_MG_QUALITY", wash)
        precheck = self._label_block(core, "@XY_MG_PRECHECK")
        slot_check = precheck.index("EQUAL N$XY_MG_SLOT 0")
        occupied_check = precheck.index("NOT CHECKUSEITEM 17")
        exact_item_check = precheck.index("EQUAL <$JADE> 鞭尸灵玉")
        cost_route = precheck.index("DELAYGOTO 50 @XY_MG_COST_ROUTE")
        self.assertLess(slot_check, occupied_check)
        self.assertLess(occupied_check, exact_item_check)
        self.assertLess(exact_item_check, cost_route)
        self.assertNotIn("GAMEGOLD -", precheck)
        self.assertNotIn("@XY_MG_QUALITY", precheck)

    def test_paid_wash_increments_and_loads_selected_slot_before_quality(self) -> None:
        core = self._core(self._compiled())
        progress = self._label_block(core, "@XY_MG_PROGRESS_ROUTE")
        for mode, amount in ((1, 100), (10, 1000)):
            with self.subTest(mode=mode):
                pay = self._label_block(core, f"@XY_MG_PAY_{mode}")
                self.assertLess(
                    pay.index(f"GAMEGOLD - {amount}"),
                    pay.index("DELAYGOTO 50 @XY_MG_PROGRESS_ROUTE"),
                )
                self.assertNotIn("@XY_MG_BATCH_APPLY", pay)
        for slot in range(1, 9):
            with self.subTest(slot=slot):
                sequence = (
                    f"INC U{200 + slot} <$STR(N$XY_MG_PROGRESS_ADD)>\n"
                    f"INC U{470 + slot} <$STR(N$XY_MG_PROGRESS_ADD)>\n"
                    f"INC U{490 + slot} <$STR(N$XY_MG_PROGRESS_ADD)>\n"
                    f"MOV N$XY_MG_RED_PROGRESS <$STR(U{200 + slot})>\n"
                    f"MOV N$XY_MG_COLOR_PROGRESS <$STR(U{470 + slot})>\n"
                    f"MOV N$XY_MG_TENTH_PROGRESS <$STR(U{490 + slot})>\n"
                    "DELAYGOTO 50 @XY_MG_QUALITY"
                )
                self.assertIn(sequence, progress)
                self.assertEqual(progress.count(f"INC U{200 + slot} "), 1)
                self.assertEqual(progress.count(f"INC U{470 + slot} "), 1)
                self.assertEqual(progress.count(f"INC U{490 + slot} "), 1)

    def test_ten_pull_uses_workbook_progress_and_settles_once(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        core = self._core(compile_mingge_payload(book, 20, 1, REGISTRY))
        wash = self._label_block(core, "@XY_MG_WASH_TEN")
        pay = self._label_block(core, "@XY_MG_PAY_10")
        self.assertEqual(book.settings.batch_progress, 10)
        self.assertEqual(wash.count("MOV N$XY_MG_PROGRESS_ADD 10"), 1)
        self.assertEqual(wash.count("DELAYGOTO 50 @XY_MG_PRECHECK"), 1)
        self.assertEqual(core.count("GAMEGOLD - 1000"), 1)
        self.assertEqual(pay.count("DELAYGOTO 50 @XY_MG_PROGRESS_ROUTE"), 1)
        self.assertEqual(core.count("[@XY_MG_BATCH_APPLY]"), 1)

    def test_display_refresh_is_quiet_idempotent_and_guarded_by_slot_and_item(self) -> None:
        compiled = self._compiled()
        display = compiled.files[
            PurePosixPath("Mir200/Envir/QuestDiary/玄渊命格/命格显示.txt")
        ].decode("gb18030")
        refresh = self._label_block(display, "@XY_MG_DISPLAY")
        self.assertIn("CHECKUSEITEM 17", refresh)
        self.assertIn("EQUAL <$JADE> 鞭尸灵玉", refresh)
        self.assertIn("UpdateItem 17", refresh)
        occupied_check = refresh.index("CHECKUSEITEM 17")
        exact_item_check = refresh.index("EQUAL <$JADE> 鞭尸灵玉")
        update = refresh.index("UpdateItem 17")
        self.assertLess(occupied_check, exact_item_check)
        self.assertLess(exact_item_check, update)
        self.assertEqual(refresh.count("UpdateItem 17"), 1)
        self.assertNotIn("SENDMSG", display)
        self.assertNotIn("MESSAGEBOX", display)

    def test_called_questdiary_outputs_use_engine_required_outer_braces(self) -> None:
        compiled = self._compiled()
        for relative_path, entry_label in (
            ("命格核心.txt", "@XY_MG_MAIN"),
            ("命格显示.txt", "@XY_MG_DISPLAY"),
        ):
            text = compiled.files[
                PurePosixPath(f"Mir200/Envir/QuestDiary/玄渊命格/{relative_path}")
            ].decode("gb18030")
            lines = [line for line in text.splitlines() if line.strip()]
            self.assertEqual(lines[0], f"[{entry_label}]", relative_path)
            self.assertEqual(lines[1], "{", relative_path)
            self.assertEqual(lines[-1], "}", relative_path)
            self.assertEqual(sum(line == "{" for line in lines), 1, relative_path)
            self.assertEqual(sum(line == "}" for line in lines), 1, relative_path)

    def test_compiler_maps_all_native_and_display_rows_by_slot(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        core = compile_mingge_payload(book, 20, 1, REGISTRY).files[
            PurePosixPath("Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt")
        ].decode("gb18030")
        affix_slots = {affix.affix_id: affix.slot for affix in book.affixes}
        bindings_by_slot = {
            affix_slots[attribute.affix_id]: REGISTRY[attribute.property_key].binding
            for attribute in book.attributes
            if attribute.quality == "绿"
        }
        self.assertEqual(set(bindings_by_slot), set(range(1, 9)))
        for slot in range(1, 9):
            with self.subTest(slot=slot):
                display_row = 8 + slot
                self.assertIn(
                    f"SetCustomItemAbil 17 {slot} 1 {bindings_by_slot[slot]}", core
                )
                self.assertIn(
                    f"SetCustomItemAbil 17 {slot} 0 <$STR(N$XY_MG_COLOR)>", core
                )
                self.assertIn(f"SetCustomItemAbil 17 {slot} 4 9", core)
                self.assertNotIn(f"SetCustomItemAbil 17 {slot} 0 0", core)
                self.assertIn(f"SetCustomItemAbil 17 {display_row} 0 255", core)
                self.assertIn(f"SetCustomItemAbil 17 {display_row} 1 60", core)
                self.assertIn(f"SetCustomItemAbil 17 {display_row} 2 {slot}", core)
                self.assertIn(f"SetCustomItemAbil 17 {display_row} 3 0", core)
                self.assertIn(f"SetCustomItemAbil 17 {display_row} 4 9", core)
                self.assertIn(f"SetCustomItemValueEX 17 {display_row} =", core)

    def test_payment_control_flow_follows_precheck_to_progress_and_apply(self) -> None:
        core = self._compiled().files[
            PurePosixPath("Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt")
        ].decode("gb18030")
        quality = self._label_block(core, "@XY_MG_QUALITY")
        cost_route = self._label_block(core, "@XY_MG_COST_ROUTE")
        self.assertIn("DELAYGOTO 50 @XY_MG_BATCH_APPLY", quality)
        for mode, amount in ((1, 100), (10, 1000)):
            with self.subTest(mode=mode):
                cost_label = f"@XY_MG_COST_{mode}"
                pay_label = f"@XY_MG_PAY_{mode}"
                self.assertIn(f"DELAYGOTO 50 {cost_label}", cost_route)
                cost = self._label_block(core, cost_label)
                pay = self._label_block(core, pay_label)
                self._assert_cost_payment_success_branch(cost, mode, amount)
                self.assertIn(f"GAMEGOLD - {amount}", pay)
                self.assertIn("DELAYGOTO 50 @XY_MG_PROGRESS_ROUTE", pay)
        self.assertEqual(core.count("[@XY_MG_BATCH_APPLY]"), 1)

    def test_payment_control_flow_rejects_pay_jump_in_insufficient_branch(self) -> None:
        core = self._compiled().files[
            PurePosixPath("Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt")
        ].decode("gb18030")
        for mode, amount in ((1, 100), (10, 1000)):
            with self.subTest(mode=mode):
                cost = self._label_block(core, f"@XY_MG_COST_{mode}")
                pay_jump = f"DELAYGOTO 50 @XY_MG_PAY_{mode}"
                unsafe = cost.replace(
                    f"#ACT\n{pay_jump}\nBREAK\n#ELSEACT\n",
                    f"#ACT\nBREAK\n#ELSEACT\n{pay_jump}\n",
                    1,
                )
                self.assertNotEqual(unsafe, cost)
                with self.assertRaises(AssertionError):
                    self._assert_cost_payment_success_branch(unsafe, mode, amount)

    def test_custom_text_base_offsets_every_slot_and_quality_label_line(self) -> None:
        base = 101
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        core = compile_mingge_payload(book, 21, base, REGISTRY).files[
            PurePosixPath("Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt")
        ].decode("gb18030")
        for quality_index, quality in enumerate(("绿", "橙", "红", "彩")):
            for slot in range(1, 9):
                with self.subTest(quality=quality, slot=slot):
                    expected = base + quality_index * 8 + slot - 1
                    self.assertIn(f"MOV N$XY_MG_LABEL_LINE {expected}", core)
        self.assertNotIn("OPENMERCHANTBIGDLG", core)

    def test_reset_fields_only_emit_the_selected_progress_clear(self) -> None:
        book = read_mingge_workbook(FIXTURES / "valid.xlsx")
        rules = tuple(
            replace(rule, reset_fields=("彩色进度",))
            for rule in book.quality_rules
        )
        core = compile_mingge_payload(
            replace(book, quality_rules=rules), 20, 1, REGISTRY
        ).files[PurePosixPath("Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt")].decode("gb18030")
        self.assertIn("MOV U471 0", core)
        self.assertNotIn("MOV U201 0", core)
        self.assertNotIn("MOV U491 0", core)


@dataclass(frozen=True)
class MingGeTransactionFixture:
    platform: Path
    server: Path
    client: Path
    workbook: Path


def _write_mingge_script(path: Path, labels: tuple[str, ...]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines: list[str] = []
    for label in labels:
        lines.extend((f"[@{label}]", "#IF", "#ACT", "BREAK", ""))
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("gb18030"))


def _create_mingge_transaction_fixture(root: Path) -> MingGeTransactionFixture:
    root = Path(root).resolve()
    platform, server, client = root / "platform", root / "server", root / "client"
    workbook = root / "37_命格系统.xlsx"
    shutil.copyfile(FIXTURES / "valid.xlsx", workbook)
    registry_target = platform / "packages/candidate/xy.optional.mingge-system/payload/property_registry.json"
    registry_target.parent.mkdir(parents=True)
    shutil.copyfile(REGISTRY_PATH, registry_target)
    shutil.copytree(
        SERVER_TEMPLATE_DIR,
        platform / "packages/candidate/xy.optional.mingge-system/payload/server/templates",
    )
    envir = server / "Mir200/Envir"
    (envir / "Market_Def").mkdir(parents=True)
    (envir / "MapQuest_Def").mkdir(parents=True)
    (envir / "QuestDiary/玄渊命格").mkdir(parents=True)
    (client / "data").mkdir(parents=True)
    (envir / "MapInfo.txt").write_bytes("[3 测试地图 0]\r\n".encode("gb18030"))
    (envir / "MerChant.txt").write_bytes(b"")
    _write_mingge_script(envir / "MapQuest_Def/QManage.txt", ("Login",))
    _write_mingge_script(envir / "Market_Def/QFunction-0.txt", ("TakeOnEx", "TakeOffEx"))
    (server / "Mir200/!Setup.txt").write_bytes(
        ("UseSqliteDB=1\r\nCustomItemPropertyCheck60=1\r\n"
         "CustomItemPropertyBindName60=<TEXT:$$1>\r\n").encode("gb18030")
    )
    (envir / "CustomItemPropertyTextVarList.txt").write_bytes(b"")
    (envir / "EffectImageList.txt").write_bytes("Base00.wzl\r\nBase01.wzl\r\n".encode("gb18030"))
    database = server / "Mud2/DB/ApexM2.DB"
    database.parent.mkdir(parents=True)
    connection = sqlite3.connect(database)
    try:
        connection.execute("CREATE TABLE StdItems (Idx INTEGER, Name TEXT, StdMode INTEGER)")
        connection.execute("INSERT INTO StdItems VALUES (?, ?, ?)", (1001, "鞭尸灵玉", 90))
        connection.commit()
    finally:
        connection.close()
    return MingGeTransactionFixture(platform, server, client, workbook)


def _fake_mingge_library(frames, output_dir: Path, provider_path: Path, **kwargs):
    del provider_path, kwargs
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    wzl, wzx = output_dir / "XY_MingGeLabels_fixture.wzl", output_dir / "XY_MingGeLabels_fixture.wzx"
    wzl.write_bytes(b"MINGGE-WZL-32")
    wzx.write_bytes(b"MINGGE-WZX-32")
    return mingge_module.MingGeLibrary(
        wzl, wzx, hashlib.sha256(wzl.read_bytes()).hexdigest(),
        hashlib.sha256(wzx.read_bytes()).hexdigest(), "fixture-config-hash",
        tuple({"target_id": frame.frame_id} for frame in frames),
    )


def _snapshot_mingge_targets(*roots: Path) -> dict[str, tuple[int, int, str]]:
    return {str(path): (path.stat().st_size, path.stat().st_mtime_ns,
                        hashlib.sha256(path.read_bytes()).hexdigest())
            for root in roots for path in root.rglob("*") if path.is_file()}


def _snapshot_mingge_changes(changes) -> dict[str, bytes | None]:
    return {str(c.path): c.path.read_bytes() if c.path.is_file() else None for c in changes}


class MingGePlatformTransactionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.fixture = _create_mingge_transaction_fixture(Path(self.temporary.name))
        self.service = mingge_module.MingGeService(self.fixture.platform)
        builder = mock.patch.object(mingge_module, "build_mingge_library", side_effect=_fake_mingge_library)
        builder.start()
        self.addCleanup(builder.stop)

    def _preflight(self):
        return self.service.preflight(self.fixture.server, self.fixture.workbook, self.fixture.client)

    def test_service_compiles_server_templates_from_its_platform_root(self) -> None:
        missing_module_templates = self.fixture.platform / "module-temp" / "missing-templates"
        with mock.patch.object(
            mingge_module, "_template_root", return_value=missing_module_templates
        ):
            plan = self._preflight()

        self.assertFalse(plan.blockers)
        self.assertTrue(any(change.path.name == "命格核心.txt" for change in plan.changes))

    def _set_workbook_settings(self, **values: object) -> None:
        shutil.copyfile(FIXTURES / "valid.xlsx", self.fixture.workbook)
        workbook = load_workbook(self.fixture.workbook)
        try:
            sheet = workbook["系统设置"]
            rows = {
                str(sheet.cell(row, 1).value): row
                for row in range(2, sheet.max_row + 1)
            }
            for key, value in values.items():
                sheet.cell(rows[key], 2).value = value
            workbook.save(self.fixture.workbook)
        finally:
            workbook.close()

    def _install_exact_legacy_candidate(self) -> dict[str, str]:
        envir = self.fixture.server / "Mir200/Envir"
        custom = "".join(f"-\\<Img:{i}:20:0:0>\r\n" for i in range(8)).encode("gb18030")
        (envir / "CustomItemPropertyTextVarList.txt").write_bytes(custom)
        lines = [f"Base{i:02}.wzl" for i in range(19)] + ["XY_MingGeDialog.wzl", "XY_MingGeRainbow.wzl"]
        (envir / "EffectImageList.txt").write_bytes(("\r\n".join(lines) + "\r\n").encode("gb18030"))
        (envir / "MerChant.txt").write_bytes("玄渊命格\\龙魂觉醒 3 320 339 龙魂觉醒 0 220 0\r\n".encode("gb18030"))
        scripts = {
            "npc": envir / "Market_Def/玄渊命格/龙魂觉醒-3.txt",
            "rainbow": envir / "QuestDiary/玄渊命格/彩色显示.txt",
            "config": envir / "QuestDiary/玄渊命格/命格测试配置.txt",
        }
        scripts["npc"].parent.mkdir(parents=True, exist_ok=True)
        for name, path in scripts.items():
            path.write_bytes(f"legacy-{name}\r\n".encode("gb18030"))
        label_wzl, label_wzx = self.fixture.client / "data/XY_MingGeRainbow.wzl", self.fixture.client / "data/XY_MingGeRainbow.wzx"
        label_wzl.write_bytes(b"legacy-label-wzl")
        label_wzx.write_bytes(b"legacy-label-wzx")
        dialog_wzl, dialog_wzx = self.fixture.client / "data/XY_MingGeDialog.wzl", self.fixture.client / "data/XY_MingGeDialog.wzx"
        dialog_wzl.write_bytes(b"legacy-dialog-wzl")
        dialog_wzx.write_bytes(b"legacy-dialog-wzx")
        return {
            "custom": hashlib.sha256(custom).hexdigest(),
            "npc": hashlib.sha256(scripts["npc"].read_bytes()).hexdigest(),
            "rainbow": hashlib.sha256(scripts["rainbow"].read_bytes()).hexdigest(),
            "config": hashlib.sha256(scripts["config"].read_bytes()).hexdigest(),
            "label_wzl": hashlib.sha256(label_wzl.read_bytes()).hexdigest(),
            "label_wzx": hashlib.sha256(label_wzx.read_bytes()).hexdigest(),
            "dialog_wzl": hashlib.sha256(dialog_wzl.read_bytes()).hexdigest(),
            "dialog_wzx": hashlib.sha256(dialog_wzx.read_bytes()).hexdigest(),
        }

    def _legacy_fingerprint_patch(self, hashes: dict[str, str]):
        return mock.patch.multiple(
            mingge_module, LEGACY_CUSTOM_TEXT_HASH=hashes["custom"],
            LEGACY_LABEL_WZL_HASH=hashes["label_wzl"], LEGACY_LABEL_WZX_HASH=hashes["label_wzx"],
            LEGACY_DIALOG_WZL_HASH=hashes["dialog_wzl"], LEGACY_DIALOG_WZX_HASH=hashes["dialog_wzx"],
            LEGACY_SCRIPT_HASHES={
                "Mir200/Envir/Market_Def/玄渊命格/龙魂觉醒-3.txt": hashes["npc"],
                "Mir200/Envir/QuestDiary/玄渊命格/彩色显示.txt": hashes["rainbow"],
                "Mir200/Envir/QuestDiary/玄渊命格/命格测试配置.txt": hashes["config"],
            },
        )

    def test_path_traversal_and_containment_block_unsafe_npc_destinations(self) -> None:
        cases = (
            {"脚本目录": r"..\..\outside"},
            {"脚本目录": r"C:\outside"},
            {"脚本目录": "/outside"},
            {"脚本目录": r"玄渊命格\..\outside"},
            {"NPC": "bad/name"},
            {"NPC": r"bad\name"},
            {"地图": "../3"},
            {"地图": r"3\outside"},
        )
        for mutation in cases:
            with self.subTest(mutation=mutation):
                self._set_workbook_settings(**mutation)
                plan = self._preflight()
                self.assertTrue(any("路径" in item or "组件" in item for item in plan.blockers))
                self.assertFalse(plan.changes)

        self._set_workbook_settings()
        safe = self._preflight()
        original = next(c for c in safe.changes if c.scope == "server")
        escaped = self.fixture.server.parent / "escaped-by-forged-plan.txt"
        forged = replace(
            safe,
            changes=(replace(original, path=escaped, before_hash=None),),
        )
        with self.assertRaisesRegex(mingge_module.MingGeError, "目标根目录"):
            self.service.install(forged)
        self.assertFalse(escaped.exists())

    def test_fixed_target_and_slot_reject_workbook_drift(self) -> None:
        database = self.fixture.server / "Mud2/DB/ApexM2.DB"
        connection = sqlite3.connect(database)
        try:
            connection.execute("INSERT INTO StdItems VALUES (?, ?, ?)", (1002, "另一灵玉", 90))
            connection.commit()
        finally:
            connection.close()
        for mutation in ({"装备": "另一灵玉"}, {"装备位": 18}):
            with self.subTest(mutation=mutation):
                self._set_workbook_settings(**mutation)
                plan = self._preflight()
                self.assertTrue(any("固定" in item for item in plan.blockers))
                self.assertFalse(plan.changes)

    def test_dialog_missing_or_mutated_blocks_legacy_adopt(self) -> None:
        hashes = self._install_exact_legacy_candidate()
        dialog_wzl = self.fixture.client / "data/XY_MingGeDialog.wzl"
        dialog_wzx = self.fixture.client / "data/XY_MingGeDialog.wzx"
        dialog_wzl.unlink()
        with self._legacy_fingerprint_patch(hashes):
            missing = self._preflight()
        self.assertTrue(missing.blockers)
        self.assertFalse(missing.changes)
        self.assertNotIn("legacy-adopt", missing.warnings)

        dialog_wzl.write_bytes(b"legacy-dialog-wzl")
        mutated = bytearray(dialog_wzx.read_bytes())
        mutated[-1] ^= 1
        dialog_wzx.write_bytes(mutated)
        with self._legacy_fingerprint_patch(hashes):
            changed = self._preflight()
        self.assertTrue(changed.blockers)
        self.assertFalse(changed.changes)
        self.assertNotIn("legacy-adopt", changed.warnings)

    def test_receipt_operation_and_status_remain_immutable_after_rollback(self) -> None:
        receipt = self.service.install(self._preflight())
        receipt_before = receipt.receipt_path.read_bytes()
        payload_before = json.loads(receipt_before.decode("utf-8"))
        self.assertEqual(payload_before["operation"], "mingge-system")
        self.assertEqual(payload_before["status"], "installed-pending-game-verification")
        self.assertEqual(
            self.service.rollback(self.fixture.server, receipt.transaction_id),
            "rolled-back",
        )
        self.assertEqual(receipt.receipt_path.read_bytes(), receipt_before)
        payload_after = json.loads(receipt.receipt_path.read_text(encoding="utf-8"))
        self.assertEqual(payload_after["operation"], "mingge-system")
        self.assertEqual(payload_after["status"], "installed-pending-game-verification")
        rollback_path = receipt.receipt_path.with_name("rollback.json")
        rollback = json.loads(rollback_path.read_text(encoding="utf-8"))
        self.assertEqual(rollback["operation"], "mingge-system-rollback")
        self.assertEqual(rollback["status"], "rolled-back")

    def test_preflight_is_read_only_and_plans_complete_install(self) -> None:
        before = _snapshot_mingge_targets(self.fixture.server, self.fixture.client)
        plan = self._preflight()
        self.assertEqual(_snapshot_mingge_targets(self.fixture.server, self.fixture.client), before)
        self.assertFalse(plan.blockers)
        self.assertEqual((plan.resource_index, plan.custom_text_base), (2, 1))
        paths = {c.path for c in plan.changes}
        self.assertIn(self.fixture.server / "Mir200/Envir/Market_Def/玄渊命格/龙魂觉醒-3.txt", paths)
        self.assertIn(self.fixture.server / "Mir200/Envir/QuestDiary/玄渊命格/platform-state.json", paths)
        self.assertIn(self.fixture.client / "data/XY_MingGeLabels_fixture.wzl", paths)
        self.assertIn(self.fixture.client / "data/XY_MingGeLabels_fixture.wzx", paths)
        custom = next(c for c in plan.changes if c.path.name == "CustomItemPropertyTextVarList.txt")
        self.assertEqual(len(custom.after_bytes.decode("gb18030").splitlines()), 32)
        setup = next(c for c in plan.changes if c.path.name == "!Setup.txt")
        setup_text = setup.after_bytes.decode("gb18030")
        for binding in range(1, 8):
            self.assertEqual(
                setup_text.count(f"CustomItemPropertyCheck{binding}=0"), 1
            )
            self.assertNotIn(f"CustomItemPropertyCheck{binding}=1", setup_text)
        self.assertEqual(setup_text.count("CustomItemPropertyCheck60=1"), 1)
        self.assertEqual(
            setup_text.count("CustomItemPropertyBindName60=<TEXT:$$1>"), 1
        )
        core = next(c for c in plan.changes if c.path.name == "命格核心.txt")
        self.assertIn("鞭尸灵玉", core.after_bytes.decode("gb18030"))
        client_wzl = next(c for c in plan.changes if c.path.suffix == ".wzl")
        client_wzl.path.write_bytes(client_wzl.after_bytes)
        incomplete = self._preflight()
        self.assertTrue(any("资源不完整" in item for item in incomplete.blockers))
        self.assertFalse(incomplete.changes)

    def test_exact_legacy_candidate_is_adopted_without_overwriting_rainbow(self) -> None:
        hashes = self._install_exact_legacy_candidate()
        with self._legacy_fingerprint_patch(hashes):
            plan = self._preflight()
        self.assertFalse(plan.blockers)
        self.assertIn("legacy-adopt", plan.warnings)
        self.assertEqual(plan.resource_index, 21)
        custom = next(c for c in plan.changes if c.path.name == "CustomItemPropertyTextVarList.txt")
        self.assertEqual(len(custom.after_bytes.decode("gb18030").splitlines()), 32)
        self.assertTrue(
            all(
                line.startswith("<PlayImg:21:")
                for line in custom.after_bytes.decode("gb18030").splitlines()
            )
        )
        effect = next(c for c in plan.changes if c.path.name == "EffectImageList.txt")
        effect_lines = effect.after_bytes.decode("gb18030").splitlines()
        self.assertEqual(effect_lines[20], "XY_MingGeRainbow.wzl")
        self.assertEqual(effect_lines[21], "XY_MingGeLabels_fixture.wzl")
        state = next(c for c in plan.changes if c.path.name == "platform-state.json")
        self.assertEqual(
            json.loads(state.after_bytes.decode("utf-8"))["dialog_resource_index"],
            19,
        )
        core = next(c for c in plan.changes if c.path.name == "命格核心.txt")
        self.assertIn("OPENMERCHANTBIGDLG 19", core.after_bytes.decode("gb18030"))

    def test_managed_legacy_resource_collision_is_relocated_from_index_20(self) -> None:
        hashes = self._install_exact_legacy_candidate()
        with self._legacy_fingerprint_patch(hashes):
            self.service.install(self._preflight())

        envir = self.fixture.server / "Mir200/Envir"
        effect_path = envir / "EffectImageList.txt"
        effect_lines = effect_path.read_bytes().decode("gb18030").splitlines()
        managed_resource = effect_lines.pop(21)
        effect_lines[20] = managed_resource
        effect_path.write_bytes(("\r\n".join(effect_lines) + "\r\n").encode("gb18030"))

        custom_path = envir / "CustomItemPropertyTextVarList.txt"
        custom_bytes = custom_path.read_bytes().replace(b"<PlayImg:21:", b"<PlayImg:20:")
        custom_path.write_bytes(custom_bytes)

        state_path = envir / "QuestDiary/玄渊命格/platform-state.json"
        state = json.loads(state_path.read_text(encoding="utf-8"))
        state["custom_text"]["line_hash"] = hashlib.sha256(custom_bytes).hexdigest()
        state["effect_image_list"]["line_number"] = 20
        state_path.write_text(
            json.dumps(state, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
            encoding="utf-8",
        )

        with self._legacy_fingerprint_patch(hashes):
            plan = self._preflight()

        self.assertFalse(plan.blockers)
        self.assertIn("managed-legacy-resource-relocated", plan.warnings)
        self.assertEqual(plan.resource_index, 21)
        custom = next(c for c in plan.changes if c.path.name == "CustomItemPropertyTextVarList.txt")
        self.assertTrue(
            all(
                line.startswith("<PlayImg:21:")
                for line in custom.after_bytes.decode("gb18030").splitlines()
            )
        )
        effect = next(c for c in plan.changes if c.path.name == "EffectImageList.txt")
        repaired_lines = effect.after_bytes.decode("gb18030").splitlines()
        self.assertEqual(repaired_lines[20], "XY_MingGeRainbow.wzl")
        self.assertEqual(repaired_lines[21], managed_resource)

    def test_one_byte_legacy_mutation_blocks_without_fuzzy_adoption(self) -> None:
        hashes = self._install_exact_legacy_candidate()
        npc = self.fixture.server / "Mir200/Envir/Market_Def/玄渊命格/龙魂觉醒-3.txt"
        data = bytearray(npc.read_bytes())
        data[-3] ^= 1
        npc.write_bytes(data)
        with self._legacy_fingerprint_patch(hashes):
            plan = self._preflight()
        self.assertTrue(plan.blockers)
        self.assertFalse(plan.changes)
        self.assertNotIn("legacy-adopt", plan.warnings)

    def test_install_failure_on_fifth_write_restores_all_targets(self) -> None:
        plan = self._preflight()
        before = _snapshot_mingge_changes(plan.changes)
        original_replace, calls = self.service._atomic_replace, 0

        def fail_on_fifth(path: Path, data: bytes) -> None:
            nonlocal calls
            calls += 1
            if calls == 5:
                raise OSError("模拟第五次写入失败")
            original_replace(path, data)

        with mock.patch.object(self.service, "_atomic_replace", side_effect=fail_on_fifth):
            with self.assertRaisesRegex(OSError, "模拟第五次写入失败"):
                self.service.install(plan)
        self.assertEqual(_snapshot_mingge_changes(plan.changes), before)

    def test_repeated_preflight_and_install_are_idempotent(self) -> None:
        first = self.service.install(self._preflight())
        self.assertEqual(first.status, "installed-pending-game-verification")
        plan = self._preflight()
        self.assertFalse(plan.blockers)
        self.assertFalse(plan.changes)
        self.assertEqual(self.service.install(plan).status, "already-current")
        qmanage = self.fixture.server / "Mir200/Envir/MapQuest_Def/QManage.txt"
        managed = (
            "; XY-MINGGE-LOGIN-BEGIN\r\n"
            "#CALL [\\玄渊命格\\命格显示.txt] @XY_MG_DISPLAY\r\n"
            "; XY-MINGGE-LOGIN-END\r\n"
        ).encode("gb18030")
        qmanage.write_bytes(qmanage.read_bytes().replace(managed, b""))
        missing_hook = self._preflight()
        self.assertTrue(any("受管钩子缺失" in item for item in missing_hook.blockers))
        self.assertFalse(missing_hook.changes)

    def test_manual_edit_blocks_entire_rollback_before_first_write(self) -> None:
        plan = self._preflight()
        receipt = self.service.install(plan)
        edited = next(c for c in plan.changes if c.path.name == "命格核心.txt")
        edited.path.write_bytes(edited.path.read_bytes() + b"manual")
        before = _snapshot_mingge_changes(plan.changes)
        with self.assertRaisesRegex(mingge_module.MingGeError, "安装后文件已被修改"):
            self.service.rollback(self.fixture.server, receipt.transaction_id)
        self.assertEqual(_snapshot_mingge_changes(plan.changes), before)

    def test_rollback_is_byte_exact_and_deletes_new_client_pair(self) -> None:
        plan = self._preflight()
        before = _snapshot_mingge_changes(plan.changes)
        receipt = self.service.install(plan)
        client_changes = [c for c in plan.changes if c.scope == "client"]
        self.assertEqual(len(client_changes), 2)
        self.assertTrue(all(c.path.is_file() for c in client_changes))
        self.assertEqual(self.service.rollback(self.fixture.server, receipt.transaction_id), "rolled-back")
        self.assertEqual(_snapshot_mingge_changes(plan.changes), before)
        self.assertTrue(all(not c.path.exists() for c in client_changes))


class MingGeCandidatePackageTests(unittest.TestCase):
    def test_candidate_package_is_config_driven_and_executable_free(self) -> None:
        manifest_path = PACKAGE / "manifest.json"
        self.assertTrue(manifest_path.is_file(), "candidate manifest 尚未建立")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        self.assertEqual(manifest["id"], "xy.optional.mingge-system")
        self.assertEqual(manifest["status"], "candidate")
        self.assertEqual(manifest["operations"], [])
        self.assertEqual(manifest["preflight_checks"], [])
        self.assertEqual(manifest["post_checks"], [])
        forbidden = {".exe", ".dll", ".py", ".ps1", ".bat", ".cmd"}
        self.assertFalse(
            [
                path
                for path in PACKAGE.rglob("*")
                if path.is_file() and path.suffix.lower() in forbidden
            ]
        )

    def test_candidate_readme_and_static_evidence_keep_game_acceptance_pending(self) -> None:
        readme_path = PACKAGE / "README_使用说明.md"
        evidence_path = PACKAGE / "evidence" / "static_verification.json"
        self.assertTrue(readme_path.is_file(), "candidate 使用说明尚未建立")
        self.assertTrue(evidence_path.is_file(), "candidate 静态验证证据尚未建立")

        readme = readme_path.read_text(encoding="utf-8")
        for required in (
            "37_命格系统.xlsx",
            "config-sync-preflight",
            "config-sync-install",
            "config-sync-rollback",
            "--route mingge-system --yes",
            "待游戏验证",
        ):
            self.assertIn(required, readme)

        evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
        self.assertEqual(evidence["package_id"], "xy.optional.mingge-system")
        self.assertEqual(evidence["status"], "static_pass_game_pending")
        self.assertEqual(evidence["install_route"]["workbook"], "37_命格系统.xlsx")
        self.assertTrue(evidence["executable_free"])
        self.assertTrue(evidence["game_acceptance_required"])

    def test_interface_documents_expose_only_the_workbook_37_config_sync_route(self) -> None:
        interface_path = INTERFACE_ROOT / "39_命格系统表格与一键植入接口.txt"
        self.assertTrue(interface_path.is_file(), "39 号命格接口尚未建立")
        interface = interface_path.read_text(encoding="utf-8")

        for sheet in (
            "系统设置",
            "槽位开放",
            "词条定义",
            "品质属性",
            "品质规则",
            "洗练消耗",
            "品质样式",
        ):
            self.assertIn(sheet, interface)
        for managed in (
            r"Mir200\!Setup.txt",
            r"Mir200\Envir\MerChant.txt",
            r"Mir200\Envir\MapQuest_Def\QManage.txt",
            r"Mir200\Envir\Market_Def\QFunction-0.txt",
            r"Mir200\Envir\CustomItemPropertyTextVarList.txt",
            r"Mir200\Envir\EffectImageList.txt",
            "platform-state.json",
            "命格核心.txt",
            "命格显示.txt",
            "命格属性.txt",
            "XY_MingGeLabels_",
        ):
            self.assertIn(managed, interface)
        for required in (
            "37_命格系统.xlsx",
            "鞭尸灵玉",
            "StdMode=90",
            "装备位=17",
            "config-sync-preflight",
            "config-sync-install",
            "config-sync-rollback",
            "--route mingge-system --yes",
            "重载 M2 脚本",
            "待游戏验证",
        ):
            self.assertIn(required, interface)
        command_lines = [
            line for line in interface.splitlines() if line.startswith(r"bin\xydp-cli.exe ")
        ]
        self.assertEqual(
            {line.split()[1] for line in command_lines},
            {"config-sync-preflight", "config-sync-install", "config-sync-rollback"},
        )

        index = (INTERFACE_ROOT / "00_先看这里_接口总索引.txt").read_text(encoding="utf-8")
        config_sync = (INTERFACE_ROOT / "10_主动选择配置文件同步接口.txt").read_text(encoding="utf-8")
        platform_readme = (PAYLOAD_ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn("39_命格系统表格与一键植入接口.txt", index)
        self.assertIn("37_命格系统.xlsx", config_sync)
        self.assertIn("mingge-system", config_sync)
        self.assertIn("xy.optional.mingge-system", platform_readme)
        self.assertIn("37_命格系统.xlsx", platform_readme)


if __name__ == "__main__":
    unittest.main()
