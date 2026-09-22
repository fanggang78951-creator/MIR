from __future__ import annotations

import hashlib
import importlib
import importlib.util
import io
import json
import re
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from xydp.mingge_native import (  # noqa: E402
    NativeCandidateError,
    compile_native_candidate,
    read_native_candidate_workbook,
)


V3_DEFINITION_HEADERS = (
    "候选ID", "启用", "目标装备名称", "中文具体部位", "命格名称", "颜色方案ID", "备注", "使用说明",
)
V3_PROPERTY_HEADERS = ("候选ID", "顺序", "中文属性", "属性值", "启用", "备注", "使用说明")
V3_SEGMENT_HEADERS = (
    "候选ID", "顺序", "片段角色", "来源属性", "片段文字", "启用", "备注", "使用说明",
)
V3_COLOR_HEADERS = (
    "方案ID", "中文名称", "类型", "统一色号", "命格名称色", "属性名称色", "正负号色", "属性值色", "单位色", "固定文字色", "备注", "使用说明",
)
V3_PART_HELP_HEADERS = ("中文具体部位", "支持状态", "说明")
V3_PROPERTY_HELP_HEADERS = ("中文属性", "实效状态", "单位", "说明")


def _sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _visible_native_text(literal: str) -> str:
    return "".join(match.group(1) for match in re.finditer(r"\{([^{}|]*)\|\d+\}", literal))


class NativeV3ColorWorkbookTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _workbook(
        self,
        *,
        definition_rows: list[tuple[object, ...]] | None = None,
        property_rows: list[tuple[object, ...]] | None = None,
        segment_rows: list[tuple[object, ...]] | None = None,
        color_rows: list[tuple[object, ...]] | None = None,
    ) -> Path:
        path = self.root / "native-v3.xlsx"
        workbook = Workbook()
        definitions = workbook.active
        definitions.title = "候选定义"
        definitions.append(V3_DEFINITION_HEADERS)
        for row in definition_rows or [
            ("uniform", True, "鞭尸灵玉", "灵玉", "【统一命格】", "red", "", "V3颜色只来自方案"),
        ]:
            definitions.append(row)
        properties = workbook.create_sheet("自定义属性")
        properties.append(V3_PROPERTY_HEADERS)
        for row in property_rows or [
            ("uniform", 1, "神力倍攻", 10, True, "", "绑定40候选"),
        ]:
            properties.append(row)
        segments = workbook.create_sheet("显示片段")
        segments.append(V3_SEGMENT_HEADERS)
        for row in segment_rows or [
            ("uniform", 1, "命格名称", None, None, True, "", "由候选定义派生"),
            ("uniform", 2, "属性名称", "神力倍攻", None, True, "", "由属性派生"),
            ("uniform", 3, "正负号", "神力倍攻", None, True, "", "由属性值派生"),
            ("uniform", 4, "属性值", "神力倍攻", None, True, "", "由属性值派生"),
            ("uniform", 5, "单位", "神力倍攻", None, True, "", "由属性目录派生"),
            ("uniform", 6, "固定文字", None, "·", True, "", "固定分隔"),
        ]:
            segments.append(row)
        colors = workbook.create_sheet("颜色方案")
        colors.append(V3_COLOR_HEADERS)
        for row in color_rows or [
            ("red", "统一红", "uniform", 249, None, None, None, None, None, None, "", "全部角色统一色"),
        ]:
            colors.append(row)
        part_help = workbook.create_sheet("部位说明")
        part_help.append(V3_PART_HELP_HEADERS)
        part_help.append(("灵玉", "支持", "固定映射位置17"))
        property_help = workbook.create_sheet("属性说明")
        property_help.append(V3_PROPERTY_HELP_HEADERS)
        property_help.append(("神力倍攻", "依赖09号QFunction/QManage实效出口", "%", "绑定40"))
        workbook.save(path)
        workbook.close()
        return path

    def test_uniform_scheme_is_the_only_source_and_colors_every_role_identically(self) -> None:
        candidate = read_native_candidate_workbook(self._workbook()).candidates[0]
        compiled = compile_native_candidate(candidate)
        self.assertEqual(candidate.color_scheme.scheme_id, "red")
        self.assertEqual(candidate.color_scheme.kind, "uniform")
        self.assertEqual({segment.color for segment in candidate.segments}, {249})
        self.assertEqual(
            compiled.segment_literal,
            "{【统一命格】|249}{神力倍攻|249}{+|249}{10|249}{%|249}{·|249}",
        )

    def test_native_rainbow_maps_each_segment_role_to_the_scheme_column(self) -> None:
        path = self._workbook(
            definition_rows=[("rainbow", True, "鞭尸灵玉", "灵玉", "【彩色命格】", "rainbow", "", "角色映射")],
            property_rows=[("rainbow", 1, "神力倍攻", 10, True, "", "")],
            segment_rows=[
                ("rainbow", 1, "命格名称", None, None, True, "", ""),
                ("rainbow", 2, "属性名称", "神力倍攻", None, True, "", ""),
                ("rainbow", 3, "正负号", "神力倍攻", None, True, "", ""),
                ("rainbow", 4, "属性值", "神力倍攻", None, True, "", ""),
                ("rainbow", 5, "单位", "神力倍攻", None, True, "", ""),
                ("rainbow", 6, "固定文字", None, "·", True, "", ""),
            ],
            color_rows=[("rainbow", "原生分段彩色", "native-rainbow", None, 249, 250, 69, 251, 252, 255, "", "角色色")],
        )
        compiled = compile_native_candidate(read_native_candidate_workbook(path).candidates[0])
        self.assertEqual(
            compiled.segment_literal,
            "{【彩色命格】|249}{神力倍攻|250}{+|69}{10|251}{%|252}{·|255}",
        )

    def test_v3_allows_middle_dot_but_rejects_spaces_and_native_control_characters(self) -> None:
        compile_native_candidate(read_native_candidate_workbook(self._workbook()).candidates[0])
        for text in ("坏 字", "坏|字", "坏{字", "坏}字", "坏\\字", "$$1"):
            with self.subTest(text=text):
                rows = [
                    ("uniform", 1, "命格名称", None, None, True, "", ""),
                    ("uniform", 2, "属性名称", "神力倍攻", None, True, "", ""),
                    ("uniform", 3, "正负号", "神力倍攻", None, True, "", ""),
                    ("uniform", 4, "属性值", "神力倍攻", None, True, "", ""),
                    ("uniform", 5, "单位", "神力倍攻", None, True, "", ""),
                    ("uniform", 6, "固定文字", None, text, True, "", ""),
                ]
                with self.assertRaisesRegex(NativeCandidateError, "片段文字"):
                    read_native_candidate_workbook(self._workbook(segment_rows=rows))

    def test_v3_rejects_sheet_count_name_order_and_header_drift(self) -> None:
        for mutation in ("extra", "rename", "reorder", "header"):
            with self.subTest(mutation=mutation):
                path = self._workbook()
                workbook = load_workbook(path)
                if mutation == "extra":
                    workbook.create_sheet("意外工作表")
                elif mutation == "rename":
                    workbook["颜色方案"].title = "配色"
                elif mutation == "reorder":
                    workbook._sheets[2], workbook._sheets[3] = workbook._sheets[3], workbook._sheets[2]
                else:
                    workbook["颜色方案"]["D1"] = "直接色号"
                workbook.save(path)
                workbook.close()
                with self.assertRaisesRegex(
                    NativeCandidateError,
                    "恰好包含|标题必须逐字匹配|必须为中文简表两页结构",
                ):
                    read_native_candidate_workbook(path)

    def test_v3_rejects_missing_and_case_insensitive_duplicate_color_schemes(self) -> None:
        missing = self._workbook(
            definition_rows=[("uniform", True, "鞭尸灵玉", "灵玉", "【统一命格】", "missing", "", "")],
        )
        with self.assertRaisesRegex(NativeCandidateError, "未知颜色方案|颜色方案.*不存在"):
            read_native_candidate_workbook(missing)
        duplicate = self._workbook(color_rows=[
            ("red", "统一红", "uniform", 249, None, None, None, None, None, None, "", ""),
            ("RED", "重复红", "uniform", 249, None, None, None, None, None, None, "", ""),
        ])
        with self.assertRaisesRegex(NativeCandidateError, "颜色方案.*重复|大小写"):
            read_native_candidate_workbook(duplicate)

    def test_v3_rejects_illegal_scheme_type_non_applicable_fields_and_bad_colors(self) -> None:
        invalid_rows = (
            ("red", "坏类型", "gradient", 249, None, None, None, None, None, None, "", ""),
            ("red", "统一夹带角色色", "uniform", 249, 250, None, None, None, None, None, "", ""),
            ("rainbow", "彩色缺字段", "native-rainbow", None, 249, 250, None, 251, 252, 255, "", ""),
        )
        for row in invalid_rows:
            with self.subTest(row=row):
                with self.assertRaisesRegex(NativeCandidateError, "类型|留空|色号|颜色|整数"):
                    read_native_candidate_workbook(self._workbook(color_rows=[row]))
        for bad in (True, -1, 256, 1.5):
            with self.subTest(bad=bad):
                row = ("red", "统一红", "uniform", bad, None, None, None, None, None, None, "", "")
                with self.assertRaisesRegex(NativeCandidateError, "统一色号.*0.*255|统一色号.*整数"):
                    read_native_candidate_workbook(self._workbook(color_rows=[row]))


class NativeP3ColorInstallTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _api(self):
        spec = importlib.util.find_spec("xydp.mingge_native_p3_color_install")
        self.assertIsNotNone(spec, "P3颜色安装模块尚未实现")
        return importlib.import_module("xydp.mingge_native_p3_color_install")

    def _workbook(self, *, mutate=None) -> Path:
        path = self.root / "native-p3.xlsx"
        workbook = Workbook()
        definitions = workbook.active
        definitions.title = "候选定义"
        definitions.append(V3_DEFINITION_HEADERS)
        definitions_rows = (
            ("uniform-red", True, "鞭尸灵玉", "灵玉", "【统一红命格·测试】", "red", "", "统一红249"),
            ("uniform-green", True, "鞭尸灵玉", "灵玉", "【统一绿命格·测试】", "green", "", "统一绿250"),
            ("uniform-blue", True, "鞭尸灵玉", "灵玉", "【统一蓝命格·测试】", "blue", "", "统一蓝252"),
            ("native-rainbow", True, "鞭尸灵玉", "灵玉", "【原生彩色命格·测试】", "rainbow", "", "原生角色分段彩色"),
        )
        for row in definitions_rows:
            definitions.append(row)
        properties = workbook.create_sheet("自定义属性")
        properties.append(V3_PROPERTY_HEADERS)
        for candidate_id, *_ in definitions_rows:
            properties.append((candidate_id, 1, "防御", 5, True, "已实测确认", "绑定1、行17"))
            properties.append((candidate_id, 2, "生命值", 100, True, "已实测确认", "绑定6、行18"))
            properties.append((candidate_id, 3, "魔法值", 100, True, "已实测确认", "绑定7、行19"))
        segments = workbook.create_sheet("显示片段")
        segments.append(V3_SEGMENT_HEADERS)
        segment_specs = (
            (1, "命格名称", None, None),
            (2, "属性名称", "防御", None), (3, "正负号", "防御", None), (4, "属性值", "防御", None),
            (5, "固定文字", None, "·"),
            (6, "属性名称", "生命值", None), (7, "正负号", "生命值", None), (8, "属性值", "生命值", None),
            (9, "固定文字", None, "·"),
            (10, "属性名称", "魔法值", None), (11, "正负号", "魔法值", None), (12, "属性值", "魔法值", None),
        )
        for candidate_id, *_ in definitions_rows:
            for order, role, source, text in segment_specs:
                segments.append((candidate_id, order, role, source, text, True, "", "P3受限显示片段"))
        colors = workbook.create_sheet("颜色方案")
        colors.append(V3_COLOR_HEADERS)
        for row in (
            ("red", "统一红", "uniform", 249, None, None, None, None, None, None, "", "整条统一色"),
            ("green", "统一绿", "uniform", 250, None, None, None, None, None, None, "", "整条统一色"),
            ("blue", "统一蓝", "uniform", 252, None, None, None, None, None, None, "", "整条统一色"),
            ("rainbow", "原生分段彩色", "native-rainbow", None, 249, 250, 69, 251, 252, 255, "", "角色映射"),
        ):
            colors.append(row)
        part_help = workbook.create_sheet("部位说明")
        part_help.append(V3_PART_HELP_HEADERS)
        part_help.append(("灵玉", "支持", "当前端位置17"))
        property_help = workbook.create_sheet("属性说明")
        property_help.append(V3_PROPERTY_HELP_HEADERS)
        property_help.append(("防御", "已实测确认", "", "绑定1、行17"))
        property_help.append(("生命值", "已实测确认", "", "绑定6、行18"))
        property_help.append(("魔法值", "已实测确认", "", "绑定7、行19"))
        if mutate is not None:
            mutate(workbook)
        workbook.save(path)
        workbook.close()
        return path

    def _server(self) -> tuple[Path, dict[str, bytes]]:
        server = self.root / "MirServer"
        envir = server / "Mir200" / "Envir"
        (envir / "Market_Def").mkdir(parents=True, exist_ok=True)
        (envir / "QuestDiary" / "玄渊验收").mkdir(parents=True, exist_ok=True)
        core_lines = [
            "[@XY_MG_MAIN]", "{", "#IF", "CHECKUSEITEM 17", "EQUAL <$JADE> 鞭尸灵玉",
            "#ACT", "MOV N$XY_MG_SLOT 0", "DELAYGOTO 50 @XY_MG_PANEL_ROUTE", "BREAK",
        ]
        for panel in range(9):
            core_lines.extend((
                f"[@XY_MG_PANEL_{panel}]", "#IF", "#ACT", "#SAY",
                "<&Text:点击中间已开放槽位进行选择:420:126{FCOLOR=161}>",
            ))
        core_lines.extend((
            "[@XY_MG_QUALITY]", "#IF", "#ACT", "MOV N$XY_MG_QUALITY 1",
            "#IF", "LARGE N$XY_MG_TENTH_PROGRESS 9", "RANDOM 10", "#ACT",
            "MOV N$XY_MG_QUALITY 4", "BREAK",
            "#IF", "RANDOM 50", "#ACT", "MOV N$XY_MG_QUALITY 3", "BREAK",
            "#IF", "RANDOM 10", "#ACT", "MOV N$XY_MG_QUALITY 2", "BREAK",
            "[@XY_MG_DRAW_TAIL]", "#IF", "#ACT", "DELAYGOTO 50 @XY_MG_PANEL_ROUTE",
            "BREAK", "}",
        ))
        payloads = {
            "Mir200/Envir/UserCmd.txt": "P2-USERCMD-98\r\n".encode("gb18030"),
            "Mir200/Envir/MerChant.txt": "NPC-REGISTRY\r\n".encode("gb18030"),
            "Mir200/Envir/Market_Def/玄渊命格/龙魂觉醒-3.txt": (
                "[@main]\r\n#IF\r\n#ACT\r\n#CALL [\\玄渊命格\\命格核心.txt] @XY_MG_MAIN\r\nBREAK\r\n"
            ).encode("gb18030"),
            "Mir200/Envir/QuestDiary/玄渊命格/命格核心.txt": (
                "\r\n".join(core_lines) + "\r\n"
            ).encode("gb18030"),
            "Mir200/Envir/CustomItemPropertyTextVarList.txt": ("\r\n".join(f"旧行{i}" for i in range(1, 33)) + "\r\n").encode("gb18030"),
            "Mir200/Envir/Market_Def/QFunction-0.txt": (
                "[@PlayLogin]\r\n[@TakeOnEx]\r\n[@TakeOffEx]\r\n"
                "; XY_EQUIP_MAKER_MONSTER_ABSORB_ANCHOR\r\n"
                "; XY_EXECUTION_LAB_CHANCE_ANCHOR\r\n"
                "; XY_EXECUTION_LAB_TOUGHNESS_ANCHOR\r\n"
                "; XY_EXECUTION_LAB_PVE_BONUS_ANCHOR\r\n"
                "; XY_EQUIP_MAKER_CORPSE_ANCHOR\r\n"
                "; XY_EQUIP_MAKER_RUNTIME_BLAST_ANCHOR\r\n"
                "; XY_EQUIP_MAKER_RUNTIME_DAMAGE_COEFFICIENT_ANCHOR\r\n"
                "; XY_EQUIP_MAKER_KILL_RESET\r\n"
            ).encode("gb18030"),
            "Mir200/Envir/QuestDiary/玄渊验收/命格平台三属性验收.txt": "P2-BASELINE-SCRIPT\r\n".encode("gb18030"),
        }
        for relative, payload in payloads.items():
            path = server / Path(relative)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(payload)
        return server, payloads

    def _baseline_patch(self, api, payloads: dict[str, bytes]):
        return patch.object(api, "EXPECTED_P2_HASHES", {key: _sha256(value) for key, value in payloads.items()})

    def test_preflight_is_zero_write_and_plans_backend_plus_original_core_without_replacing_npc(self) -> None:
        api = self._api()
        workbook = self._workbook()
        server, payloads = self._server()
        platform = self.root / "platform"
        before = {relative: (server / Path(relative)).read_bytes() for relative in payloads}
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(workbook, server, platform)
        self.assertEqual(plan.blockers, ())
        self.assertFalse(platform.exists())
        self.assertEqual(tuple(item.relative_path for item in plan.files), (
            api.TEST_SCRIPT_RELATIVE,
            api.TEXTVAR_RELATIVE,
            api.QFUNCTION_RELATIVE,
            api.CORE_RELATIVE,
        ))
        self.assertNotIn(api.NPC_WRAPPER_RELATIVE, {item.relative_path for item in plan.files})
        self.assertNotIn(api.MERCHANT_RELATIVE, {item.relative_path for item in plan.files})
        self.assertEqual(
            (server / Path(api.NPC_WRAPPER_RELATIVE)).read_bytes(),
            payloads[api.NPC_WRAPPER_RELATIVE],
        )
        self.assertEqual({relative: (server / Path(relative)).read_bytes() for relative in payloads}, before)

    def test_original_ui_draw_tail_uses_same_file_eight_slot_pool_adapter(self) -> None:
        api = self._api()
        workbook = self._workbook()
        server, payloads = self._server()
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(workbook, server, self.root / "platform-original-ui")
        self.assertEqual(plan.blockers, ())
        by_relative = {item.relative_path: item for item in plan.files}
        core = by_relative[api.CORE_RELATIVE].after.decode("gb18030")
        script = by_relative[api.TEST_SCRIPT_RELATIVE].after.decode("gb18030")
        self.assertIn("[@XY_MG_DRAW_TAIL]", core)
        self.assertEqual(core.count("; XY-MG-P3-BEGIN ORIGINAL_UI_APPLY"), 1)
        self.assertEqual(core.count("; XY-MG-P3-END ORIGINAL_UI_APPLY"), 1)
        self.assertNotIn("@XY_MG_P3_APPLY_CURRENT_SLOT", core)
        self.assertIn("GOTO @XY_MG_P4_SELECT_POOL", core)
        self.assertIn("[@XY_MG_P4_WRITE_SLOT_1]", core)
        self.assertIn("[@XY_MG_P4_WRITE_SLOT_8]", core)
        for slot, row in enumerate(range(9, 17), start=1):
            self.assertIn(f"SetCustomItemAbil 17 {row} 1 60", core)
        self.assertIn(
            "GetCustomItemValueEx 17 9 N$XY_MG_P4_TMPP N$XY_MG_P4_SLOT1 N$XY_MG_P4_TMPB N$XY_MG_P4_TMPC",
            core,
        )
        self.assertIn(
            "GetCustomItemValueEx 17 9 N$XY_MG_P3_TMPP N$XY_MG_P3_SLOT1 N$XY_MG_P3_TMPB N$XY_MG_P3_TMPC",
            script,
        )
        self.assertNotRegex(core, r"(?:Abil|ValueEx) 17 (?:2[0-9]|[3-9][0-9])(?:\s|$)")
        self.assertIn("[@XY_MG_P3_RECALC_ALL]", script)
        self.assertNotIn("[@XY_MG_P3_APPLY_CURRENT_SLOT]", script)
        script_lines = [line for line in script.splitlines() if line.strip()]
        self.assertEqual(script_lines[-1], "}")
        self.assertNotIn("}", script_lines[1:-1])

    def test_attack_speed_breakthrough_is_aggregated_through_existing_textline40_interface(self) -> None:
        api = self._api()

        def attack_speed_only(workbook) -> None:
            properties = workbook["自定义属性"]
            properties.delete_rows(2, properties.max_row - 1)
            for candidate_id in ("uniform-red", "uniform-green", "uniform-blue", "native-rainbow"):
                properties.append((candidate_id, 1, "攻速突破", 7, True, "", "复用TextLine40"))
            workbook["属性说明"].append(("攻速突破", "复用已安装攻速突破接口", "", "TextLine40"))

        candidates = read_native_candidate_workbook(
            self._workbook(mutate=attack_speed_only), derive_v3_segments=True
        ).candidates
        self.assertEqual(api._candidate_blockers(candidates), [])

        core = "\r\n".join(api._render_core_block(candidates))
        self.assertIn("SetCustomItemAbil 17 1 1 60", core)
        self.assertIn("SetCustomItemAbil 17 1 2 40", core)
        self.assertIn(
            "SetCustomItemValueEx 17 1 = 40 <$STR(N$XY_MG_P4_SUM1A)> 0",
            core,
        )
        self.assertEqual(
            core.count("#CALL [\\玄渊攻速突破\\全身攻速阈值核心.txt] @XY_AS_CAP_RECALC"),
            1,
        )
        self.assertNotIn("ChangeSpeed 2", core)

    def test_attack_speed_breakthrough_blocks_when_existing_interface_is_missing(self) -> None:
        api = self._api()

        def attack_speed_only(workbook) -> None:
            properties = workbook["自定义属性"]
            properties.delete_rows(2, properties.max_row - 1)
            for candidate_id in ("uniform-red", "uniform-green", "uniform-blue", "native-rainbow"):
                properties.append((candidate_id, 1, "攻速突破", 7, True, "", "复用TextLine40"))

        workbook = self._workbook(mutate=attack_speed_only)
        server, payloads = self._server()
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(
                workbook, server, self.root / "platform-speed-dependency"
            )

        self.assertTrue(any("攻速突破依赖" in blocker for blocker in plan.blockers))

    def test_rainbow_text_var_keeps_name_colors_and_uses_clear_pastels_for_whole_properties(self) -> None:
        api = self._api()
        candidates = read_native_candidate_workbook(self._workbook(), derive_v3_segments=True).candidates
        rainbow = next(item for item in candidates if item.candidate_id == "native-rainbow")
        literal = api._candidate_display_literal(rainbow)
        self.assertEqual(
            literal,
            "{【|249}{原|69}{生|250}{彩|251}{色|249}{命|69}{格|250}{·|251}"
            "{测|249}{试|69}{】|250}\\{防御+5|31}{·|255}{生命值+100|147}"
            "{·|255}{魔法值+100|239}",
        )
        self.assertLessEqual(len(literal), 128)

    def test_uniform_text_var_uses_short_name_chunks_and_whole_properties(self) -> None:
        api = self._api()
        candidate = read_native_candidate_workbook(
            NativeV3ColorWorkbookTests._workbook(self),
            derive_v3_segments=True,
        ).candidates[0]
        self.assertEqual(
            api._candidate_display_literal(candidate),
            "{【统|249}{一命|249}{格】|249}\\{神力倍攻+10%|249}",
        )

    def test_current_39_text_var_payloads_fit_the_observed_128_byte_safe_target(self) -> None:
        api = self._api()
        workbook = Path(r"E:\XuanYuanDevPlatform\所需材料表格汇总\39_命格P3颜色导入测试.xlsx")
        candidates = read_native_candidate_workbook(workbook, derive_v3_segments=True).candidates
        for line_no, literal in api._candidate_text_var_entries(candidates):
            with self.subTest(line_no=line_no):
                self.assertLessEqual(len(literal.encode("gb18030")), 128)

    def test_multi_candidate_pool_falls_through_without_random_elseact(self) -> None:
        api = self._api()
        workbook = Path(r"E:\XuanYuanDevPlatform\所需材料表格汇总\39_命格P3颜色导入测试.xlsx")
        candidates = read_native_candidate_workbook(workbook, derive_v3_segments=True).candidates
        red = api._classify_candidates(candidates)["RED"]
        self.assertEqual(
            api._pool_lines("RED", red),
            (
                "[@XY_MG_P4_POOL_RED]", "#IF", "#ACT",
                "GOTO @XY_MG_P4_POOL_RED_ROLL_1", "BREAK", "",
                "[@XY_MG_P4_POOL_RED_ROLL_1]", "#IF", "RANDOMEX 1 2", "#ACT",
                "GOTO @XY_MG_P4_CHOOSE_CANDIDATE_0002", "BREAK",
                "#IF", "#ACT", "GOTO @XY_MG_P4_POOL_RED_ROLL_2", "BREAK", "",
                "[@XY_MG_P4_POOL_RED_ROLL_2]", "#IF", "#ACT",
                "GOTO @XY_MG_P4_CHOOSE_CANDIDATE_0006", "BREAK", "",
            ),
        )

    def test_original_core_normalizes_only_owned_progress_once(self) -> None:
        api = self._api()
        server, payloads = self._server()
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(
                self._workbook(), server, self.root / "platform-versioned-progress"
            )
        self.assertEqual(plan.blockers, ())
        core = {item.relative_path: item for item in plan.files}[api.CORE_RELATIVE].after.decode("gb18030")
        marker = "; XY-MG-P3-BEGIN ORIGINAL_UI_STATE_INIT"
        self.assertIn(marker, core)
        block = core.split(marker, 1)[1].split(
            "; XY-MG-P3-END ORIGINAL_UI_STATE_INIT", 1
        )[0]
        cleared = {
            line.split()[1]
            for line in block.splitlines()
            if line.startswith("MOV U") and line.endswith(" 0")
        }
        self.assertEqual(cleared, {
            *(f"U{value}" for value in range(201, 209)),
            *(f"U{value}" for value in range(471, 479)),
            *(f"U{value}" for value in range(491, 499)),
        })
        self.assertIn("NOT EQUAL U499 82802", block)
        self.assertIn("MOV U499 82802", block)

    def test_legacy_batch_apply_is_upgraded_to_unconditional_slot_write_and_tail(self) -> None:
        api = self._api()
        legacy = "\r\n".join((
            "[@XY_MG_BATCH_APPLY]",
            "#IF", "EQUAL N$XY_MG_SLOT 1", "#ACT",
            "MOV N$XY_MG_VALUE 5", "MOV N$XY_MG_LABEL_LINE 1",
            "#IF", "EQUAL N$XY_MG_QUALITY 2", "#ACT",
            "MOV N$XY_MG_VALUE 10", "MOV N$XY_MG_LABEL_LINE 9",
            "#IF", "EQUAL N$XY_MG_QUALITY 3", "#ACT",
            "MOV N$XY_MG_VALUE 20", "MOV N$XY_MG_LABEL_LINE 17",
            "#IF", "EQUAL N$XY_MG_QUALITY 4", "#ACT",
            "MOV N$XY_MG_VALUE 40", "MOV N$XY_MG_LABEL_LINE 25",
            "LockUpdateItem 17", "SetCustomItemValue 17 1 = <$STR(N$XY_MG_VALUE)>",
            "#IF", "EQUAL N$XY_MG_RESET_RED 1", "#ACT", "MOV U201 0",
            "#IF", "EQUAL N$XY_MG_RESET_COLOR 1", "#ACT", "MOV U471 0",
            "#IF", "EQUAL N$XY_MG_RESET_TENTH 1", "#ACT", "MOV U491 0",
            "UpdateItem 17", "DELAYGOTO 50 @XY_MG_DRAW_TAIL", "BREAK",
            "[@XY_MG_DRAW_TAIL]", "#IF", "#ACT", "BREAK",
        ))
        repaired = api._repair_legacy_batch_apply_control(legacy)
        self.assertIn(
            "EQUAL N$XY_MG_SLOT 1\r\n#ACT\r\nGOTO @XY_MG_BATCH_APPLY_SLOT_1\r\nBREAK",
            repaired,
        )
        slot = repaired.split("[@XY_MG_BATCH_APPLY_SLOT_1]", 1)[1].split(
            "[@XY_MG_DRAW_TAIL]", 1
        )[0]
        self.assertIn("#IF\r\n#ACT\r\nLockUpdateItem 17", slot)
        self.assertIn(
            "#IF\r\n#ACT\r\nUpdateItem 17\r\nDELAYGOTO 50 @XY_MG_DRAW_TAIL\r\nBREAK",
            slot,
        )

    def test_original_core_uses_precise_randomex_quality_odds(self) -> None:
        api = self._api()
        server, payloads = self._server()
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(
                self._workbook(), server, self.root / "platform-precise-quality"
            )
        self.assertEqual(plan.blockers, ())
        core = {item.relative_path: item for item in plan.files}[api.CORE_RELATIVE].after.decode("gb18030")
        lines = core.splitlines()
        self.assertEqual(lines.count("RANDOMEX 1 50"), 1)
        self.assertEqual(lines.count("RANDOMEX 1 10"), 2)
        self.assertNotIn("RANDOM 50", lines)
        self.assertNotIn("RANDOM 10", lines)

    def test_original_panels_reuse_live_candidate_result_with_uniform_or_rainbow_color(self) -> None:
        api = self._api()
        server, payloads = self._server()
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(
                self._workbook(), server, self.root / "platform-live-result"
            )
        self.assertEqual(plan.blockers, ())
        core = {item.relative_path: item for item in plan.files}[api.CORE_RELATIVE].after.decode("gb18030")
        self.assertEqual(core.count("<$STR(S$XY_MG_P4_PANEL_RESULT)>"), 9)
        self.assertNotIn("<&Text:点击中间已开放槽位进行选择:420:126{FCOLOR=161}>", core)
        self.assertIn(
            "MOV S$XY_MG_P4_PANEL_RESULT <&Text:本次洗出：【统一红命格·测试】（槽位<$STR(N$XY_MG_SLOT)>）:405:126{FCOLOR=249;FSIZE=9}>",
            core,
        )
        self.assertIn(
            "MOV S$XY_MG_P4_PANEL_RESULT <&Text:本次洗出：【原生彩色命格·测试】（槽位<$STR(N$XY_MG_SLOT)>）:405:126{AUTOCOLOR=254,251,168,191,250,70,245,249,253;FSIZE=9}>",
            core,
        )

    def test_color_classifier_keeps_multiple_candidates_per_pool_and_requires_all_pools(self) -> None:
        api = self._api()

        def duplicate_red_and_rainbow(workbook) -> None:
            for source_row, new_id, new_name in (
                (2, "uniform-red-2", "【统一红命格·备2】"),
                (2, "uniform-red-3", "【统一红命格·备3】"),
                (5, "native-rainbow-2", "【原生彩色命格·备2】"),
            ):
                source = [cell.value for cell in workbook["候选定义"][source_row]]
                source[0], source[4] = new_id, new_name
                workbook["候选定义"].append(source)
                old_id = workbook["候选定义"].cell(source_row, 1).value
                for sheet_name in ("自定义属性", "显示片段"):
                    sheet = workbook[sheet_name]
                    for row in list(sheet.iter_rows(min_row=2, values_only=True)):
                        if row[0] == old_id:
                            sheet.append((new_id, *row[1:]))

        workbook = self._workbook(mutate=duplicate_red_and_rainbow)
        candidates = read_native_candidate_workbook(workbook, derive_v3_segments=True).candidates
        pools = api._classify_candidates(candidates)
        self.assertEqual(tuple(pools), ("GREEN", "BLUE", "RED", "RAINBOW"))
        self.assertEqual(len(pools["RED"]), 3)
        self.assertEqual(len(pools["RAINBOW"]), 2)
        self.assertEqual(api._candidate_blockers(candidates), [])
        with self.assertRaisesRegex(api.NativeP3ColorTestInstallError, "GREEN|\u7eff"):
            api._classify_candidates(tuple(item for item in candidates if item.color_scheme.uniform_color != 250))

    def test_candidate_blockers_reject_unknown_script_bind_but_accept_current_50_51_53(self) -> None:
        api = self._api()
        current = read_native_candidate_workbook(
            Path(r"E:\XuanYuanDevPlatform\所需材料表格汇总\39_命格P3颜色导入测试.xlsx"),
            derive_v3_segments=True,
        ).candidates
        self.assertEqual(api._candidate_blockers(current), [])
        source = current[0]
        bad_property = replace(source.properties[0], route="script_bind", binding=40, text_line=None)
        bad_candidate = replace(source, properties=(bad_property, *source.properties[1:]))
        blockers = api._candidate_blockers((bad_candidate, *current[1:]))
        self.assertTrue(any("binding40" in item and "script_bind" in item for item in blockers))

    def test_candidate_blockers_reject_branch_key_collision(self) -> None:
        api = self._api()
        candidates = read_native_candidate_workbook(self._workbook(), derive_v3_segments=True).candidates
        colliding = (
            replace(candidates[0], candidate_id="a-b"),
            replace(candidates[1], candidate_id="a_b"),
            *candidates[2:],
        )
        blockers = api._candidate_blockers(colliding)
        self.assertTrue(any("A_B" in item and "冲突" in item for item in blockers))

    def test_preflight_rejects_unsafe_candidate_shape_and_target(self) -> None:
        api = self._api()
        server, payloads = self._server()
        mutations = {
            "target-item": lambda wb: setattr(wb["候选定义"]["C2"], "value", "其他灵玉"),
            "part": lambda wb: setattr(wb["候选定义"]["D2"], "value", "盾牌"),
            "too-many-properties": lambda wb: wb["自定义属性"].append(
                ("uniform-red", 4, "攻击", 6, True, "", "")
            ),
        }
        for name, mutation in mutations.items():
            with self.subTest(name=name), self._baseline_patch(api, payloads):
                try:
                    plan = api.plan_native_p3_color_test_install(
                        self._workbook(mutate=mutation), server, self.root / f"platform-{name}"
                    )
                except NativeCandidateError:
                    continue
                self.assertTrue(plan.blockers, name)

    def test_generated_backend_only_recalculates_and_core_owns_random_pool_write(self) -> None:
        api = self._api()
        workbook = self._workbook()
        server, payloads = self._server()
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(workbook, server, self.root / "platform")
        self.assertEqual(plan.blockers, ())
        by_relative = {item.relative_path: item for item in plan.files}
        script = by_relative[api.TEST_SCRIPT_RELATIVE].after.decode("gb18030")
        core = by_relative[api.CORE_RELATIVE].after.decode("gb18030")
        self.assertNotIn("LockUpdateItem", script)
        self.assertNotIn("SetCustomItemValue", script)
        self.assertNotIn("SetCustomItemText", script)
        self.assertIn("LockUpdateItem 17", core)
        self.assertIn("SetCustomItemTextColor 17 255", core)
        self.assertIn("MOV S$XY_MG_P4_EMPTY", core)
        self.assertIn("SetCustomItemText 17 <$STR(S$XY_MG_P4_EMPTY)>", core)
        self.assertIn("#CALL [\\玄渊命格\\命格P3实例快速版.txt] @XY_MG_P3_RECALC_ALL", core)
        for row in range(1, 9):
            self.assertIn(f"SetCustomItemValueEx 17 {row} = 0 0 0", core)
        for row in (17, 18, 19):
            self.assertEqual(core.count(f"SetCustomItemValueEx 17 {row} = 0 0 0"), 1)
        for slot, row in enumerate(range(9, 17), start=1):
            self.assertIn(f"MOV N$XY_MG_P4_VALID{slot} 0", core)
            self.assertIn(f"MOV N$XY_MG_P4_VALID{slot} 1", core)
            self.assertIn(f"EQUAL N$XY_MG_P4_VALID{slot} 0", core)
            self.assertIn(f"SetCustomItemValueEx 17 {row} = 0 0 0", core)
        for forbidden in ("GAMEGOLD", "GAMEPOINT", "SQL ", "MIR.DB", "UPDATE STDITEMS"):
            self.assertNotIn(forbidden, core.upper())

    def test_current_39_builds_six_candidate_text_lines_and_eight_slot_aggregates(self) -> None:
        api = self._api()
        workbook = Path(r"E:\XuanYuanDevPlatform\所需材料表格汇总\39_命格P3颜色导入测试.xlsx")
        server, payloads = self._server()
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(workbook, server, self.root / "platform-current39")
        self.assertEqual(plan.blockers, ())
        self.assertEqual(plan.candidate_ids, (
            "candidate-0002", "candidate-0003", "candidate-0004",
            "candidate-0005", "candidate-0006", "candidate-0007",
        ))
        self.assertEqual(tuple(item.relative_path for item in plan.files), (
            api.TEST_SCRIPT_RELATIVE, api.TEXTVAR_RELATIVE, api.QFUNCTION_RELATIVE, api.CORE_RELATIVE,
        ))
        by_relative = {item.relative_path: item for item in plan.files}
        script = by_relative[api.TEST_SCRIPT_RELATIVE].after.decode("gb18030")
        self.assertIn("[@XY_MG_P3_RECALC_ALL]", script)
        self.assertNotIn("LockUpdateItem", script)
        core = by_relative[api.CORE_RELATIVE].after.decode("gb18030")
        self.assertIn("RANDOMEX 1 2", core)
        self.assertIn("EQUAL N$XY_MG_QUALITY 1", core)
        self.assertIn("GOTO @XY_MG_P4_POOL_GREEN", core)
        self.assertIn("EQUAL N$XY_MG_QUALITY 2", core)
        self.assertIn("GOTO @XY_MG_P4_POOL_BLUE", core)
        self.assertIn("EQUAL N$XY_MG_QUALITY 3", core)
        self.assertIn("GOTO @XY_MG_P4_POOL_RED", core)
        self.assertIn("EQUAL N$XY_MG_QUALITY 4", core)
        self.assertIn("GOTO @XY_MG_P4_POOL_RAINBOW", core)
        self.assertEqual(core.count("GetCustomItemValueEx 17 9 "), 1)
        self.assertEqual(core.count("GetCustomItemValueEx 17 16 "), 1)
        self.assertNotRegex(core, r"(?:Abil|ValueEx) 17 (?:2[0-9]|[3-9][0-9])(?:\s|$)")
        self.assertIn("SetCustomItemAbil 17 6 1 50", core)
        upper = script.upper()
        for forbidden in ("UPDATE STDITEMS", "SQL ", "MIR.DB", "GAMEGOLD", "RANDOM"):
            self.assertNotIn(forbidden, upper)
        self.assertIn("ChangeHumAbility 9 = 0", script)
        self.assertIn("ChangeHumAbility 10 = 0", script)
        self.assertNotIn("ChangeHumAbility 5", script)
        self.assertNotIn("ChangeHumAbility 6", script)
        text_var = by_relative[api.TEXTVAR_RELATIVE].after.decode("gb18030").splitlines()
        self.assertEqual(len(plan.text_var_entries), 6)
        for index, candidate in enumerate(read_native_candidate_workbook(workbook, derive_v3_segments=True).candidates, start=33):
            value = text_var[index - 1]
            self.assertIn(candidate.mingge_name, _visible_native_text(value))
            self.assertIn("\\", value)
            for prop in candidate.properties:
                self.assertIn(prop.display_name, value)
        qfunction = by_relative[api.QFUNCTION_RELATIVE].after.decode("gb18030")
        self.assertNotIn("GetAllCustomItemValueByTextLine", qfunction)
        self.assertEqual(qfunction.count("; XY-MG-P3-BEGIN TEXT38_BLAST"), 1)
        self.assertEqual(qfunction.count("; XY-MG-P3-BEGIN BIND50_TOUGHNESS"), 1)
        self.assertEqual(qfunction.count("; XY-MG-P3-BEGIN BIND51_EXECUTION_BONUS"), 1)
        self.assertEqual(qfunction.count("; XY-MG-P3-BEGIN BIND53_DAMAGE_COEFFICIENT"), 1)
        self.assertIn("INC N$XY_EXEC_Toughness <$STR(N$XY_MG_P3_TOUGHNESS)>", qfunction)
        self.assertIn("INC N$XY_EXEC_PVEEquipBonusPercent <$STR(N$XY_MG_P3_EXECUTION_BONUS)>", qfunction)
        self.assertIn("INC N$XY_RT_DamageCoeff <$STR(N$XY_MG_P3_DAMAGE_COEFF)>", qfunction)
        for variable, value in (
            ("N$XY_MG_P3_TOUGHNESS", 6),
            ("N$XY_MG_P3_EXECUTION_BONUS", 10),
            ("N$XY_MG_P3_DAMAGE_COEFF", 5),
        ):
            self.assertIn(f"MOV {variable} 0", script)
            self.assertIn(f"INC {variable} {value}", script)
        self.assertNotIn("@XY_MG_P3_APPLY_CURRENT_SLOT", core)
        self.assertEqual((server / Path(api.NPC_WRAPPER_RELATIVE)).read_bytes(), payloads[api.NPC_WRAPPER_RELATIVE])
        for planned in plan.files:
            planned.after.decode("gb18030")
            self.assertNotIn(b"\n", planned.after.replace(b"\r\n", b""))

    def test_install_and_rollback_are_byte_exact_for_all_four_files(self) -> None:
        api = self._api()
        workbook = self._workbook()
        server, payloads = self._server()
        platform = self.root / "platform"
        p2_sentinel = platform / "backups" / "mingge-native-p2-test" / "active" / "formal-p2.json"
        p2_sentinel.parent.mkdir(parents=True)
        p2_sentinel.write_bytes(b"FORMAL-P2-ACTIVE")
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(workbook, server, platform)
            receipt = api.install_native_p3_color_test(plan, platform)
            self.assertEqual(receipt.files, tuple(item.relative_path for item in plan.files))
            for planned in plan.files:
                self.assertEqual((server / Path(planned.relative_path)).read_bytes(), planned.after)
            transaction = platform / "backups" / "mingge-native-p3-color-test" / receipt.transaction_id
            manifest = json.loads((transaction / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["schema_version"], 2)
            self.assertEqual(len(manifest["files"]), 4)
            self.assertFalse(manifest["files"][0]["existed_before"])
            self.assertIsNone(manifest["files"][0]["backup_relative_path"])
            for index, relative in enumerate((api.TEXTVAR_RELATIVE, api.QFUNCTION_RELATIVE, api.CORE_RELATIVE), start=1):
                self.assertEqual((transaction / "files" / f"{index:02d}.before.bin").read_bytes(), payloads[relative])
            self.assertEqual((server / Path(api.NPC_WRAPPER_RELATIVE)).read_bytes(), payloads[api.NPC_WRAPPER_RELATIVE])
            self.assertEqual(p2_sentinel.read_bytes(), b"FORMAL-P2-ACTIVE")
            with self.assertRaisesRegex(api.NativeP3ColorTestInstallError, "活动|安装后|预检后"):
                api.install_native_p3_color_test(plan, platform)
            active_path = api._active_path(platform, server)
            coordination_path = api._coordination_path(platform, server)
            self.assertTrue(active_path.is_file())
            self.assertTrue(coordination_path.is_file())
            rollback = api.rollback_native_p3_color_test(platform, receipt.transaction_id, server)
        self.assertEqual(rollback.status, "rolled-back")
        self.assertFalse((server / Path(api.TEST_SCRIPT_RELATIVE)).exists())
        for relative in (api.TEXTVAR_RELATIVE, api.QFUNCTION_RELATIVE, api.CORE_RELATIVE):
            self.assertEqual((server / Path(relative)).read_bytes(), payloads[relative])
        self.assertEqual((server / Path(api.NPC_WRAPPER_RELATIVE)).read_bytes(), payloads[api.NPC_WRAPPER_RELATIVE])
        self.assertEqual((server / Path(api.MERCHANT_RELATIVE)).read_bytes(), payloads[api.MERCHANT_RELATIVE])
        self.assertFalse(active_path.exists())
        self.assertFalse(coordination_path.exists())
        self.assertEqual(p2_sentinel.read_bytes(), b"FORMAL-P2-ACTIVE")

    def test_install_and_rollback_reject_hash_drift_without_overwrite(self) -> None:
        api = self._api()
        workbook = self._workbook()
        server, payloads = self._server()
        platform = self.root / "platform"
        target = server / Path(api.TEST_SCRIPT_RELATIVE)
        with self._baseline_patch(api, payloads):
            plan = api.plan_native_p3_color_test_install(workbook, server, platform)
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(b"DRIFT-BEFORE-INSTALL")
            with self.assertRaisesRegex(api.NativeP3ColorTestInstallError, "预检后|哈希"):
                api.install_native_p3_color_test(plan, platform)
            self.assertEqual(target.read_bytes(), b"DRIFT-BEFORE-INSTALL")
            target.unlink()
            plan = api.plan_native_p3_color_test_install(workbook, server, platform)
            receipt = api.install_native_p3_color_test(plan, platform)
            target.write_bytes(b"DRIFT-BEFORE-ROLLBACK")
            with self.assertRaisesRegex(api.NativeP3ColorTestInstallError, "漂移|哈希"):
                api.rollback_native_p3_color_test(platform, receipt.transaction_id, server)
        self.assertEqual(target.read_bytes(), b"DRIFT-BEFORE-ROLLBACK")
        self.assertTrue(api._active_path(platform, server).is_file())

    def test_preflight_blocks_textvar_occupation_and_partial_qfunction_marker(self) -> None:
        api = self._api()
        workbook = self._workbook()
        for case in ("textvar", "qfunction"):
            with self.subTest(case=case):
                server, payloads = self._server()
                if case == "textvar":
                    path = server / Path(api.TEXTVAR_RELATIVE)
                    path.write_bytes(path.read_bytes() + "占用33\r\n".encode("gb18030"))
                    payloads[api.TEXTVAR_RELATIVE] = path.read_bytes()
                else:
                    path = server / Path(api.QFUNCTION_RELATIVE)
                    path.write_bytes(path.read_bytes() + b"; XY-MG-P3-BEGIN TEXT38_BLAST\r\n")
                    payloads[api.QFUNCTION_RELATIVE] = path.read_bytes()
                with self._baseline_patch(api, payloads):
                    plan = api.plan_native_p3_color_test_install(workbook, server, self.root / f"platform-{case}")
                self.assertTrue(plan.blockers)

    def test_preflight_blocks_when_any_script_bind_runtime_anchor_is_missing(self) -> None:
        api = self._api()
        workbook = self._workbook()
        anchors = (
            "; XY_EXECUTION_LAB_TOUGHNESS_ANCHOR",
            "; XY_EXECUTION_LAB_PVE_BONUS_ANCHOR",
            "; XY_EQUIP_MAKER_RUNTIME_DAMAGE_COEFFICIENT_ANCHOR",
        )
        for index, anchor in enumerate(anchors):
            with self.subTest(anchor=anchor):
                server, payloads = self._server()
                path = server / Path(api.QFUNCTION_RELATIVE)
                text = path.read_bytes().decode("gb18030")
                path.write_bytes(text.replace(anchor + "\r\n", "", 1).encode("gb18030"))
                payloads[api.QFUNCTION_RELATIVE] = path.read_bytes()
                with self._baseline_patch(api, payloads):
                    plan = api.plan_native_p3_color_test_install(
                        workbook, server, self.root / f"platform-missing-script-bind-anchor-{index}"
                    )
                self.assertTrue(plan.blockers)
                self.assertTrue(any(anchor in blocker for blocker in plan.blockers))

    def test_cli_install_and_rollback_require_explicit_yes(self) -> None:
        api = self._api()
        from xydp.cli import main as cli_main

        workbook = self._workbook()
        server, payloads = self._server()
        platform = self.root / "platform"
        install_args = [
            "--root", str(platform), "mingge-native-p3-color-test-install",
            "--input", str(workbook), "--server", str(server),
        ]
        with self._baseline_patch(api, payloads):
            with self.assertRaisesRegex(SystemExit, "必须显式提供 --yes"):
                cli_main(install_args)
            stdout = io.StringIO()
            with redirect_stdout(stdout):
                self.assertEqual(cli_main([*install_args, "--yes"]), 0)
            transaction_id = json.loads(stdout.getvalue())["transaction_id"]
            rollback_args = [
                "--root", str(platform), "mingge-native-p3-color-test-rollback",
                "--transaction", transaction_id, "--server", str(server),
            ]
            with self.assertRaisesRegex(SystemExit, "必须显式提供 --yes"):
                cli_main(rollback_args)
            stdout = io.StringIO()
            with redirect_stdout(stdout):
                self.assertEqual(cli_main([*rollback_args, "--yes"]), 0)
        self.assertEqual(json.loads(stdout.getvalue())["status"], "rolled-back")


if __name__ == "__main__":
    unittest.main()
