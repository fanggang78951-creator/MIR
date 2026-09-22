from __future__ import annotations

import importlib
import io
import json
import os
import sys
import tempfile
import unittest
import uuid
from contextlib import redirect_stdout
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))


V2_DEFINITION_HEADERS = (
    "候选ID", "启用", "目标装备名称", "中文具体部位", "命格名称", "默认文字颜色", "备注", "使用说明",
)
V2_PROPERTY_HEADERS = ("候选ID", "顺序", "中文属性", "属性值", "启用", "备注", "使用说明")
V2_SEGMENT_HEADERS = (
    "候选ID", "顺序", "片段角色", "来源属性", "片段文字", "颜色", "启用", "备注", "使用说明",
)
V2_PART_HELP_HEADERS = ("中文具体部位", "支持状态", "说明")
V2_PROPERTY_HELP_HEADERS = ("中文属性", "实效状态", "单位", "说明")


class _SimulatedProcessCrash(BaseException):
    """Bypass in-process recovery to model a process terminating at a fault boundary."""


class NativeP2InstallTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.api = importlib.import_module("xydp.mingge_native_p2_install")
        self.real_restrict_secret_acl = self.api._restrict_secret_acl
        self.real_verify_secret_acl = self.api._verify_secret_acl
        self.patches = (
            patch.object(self.api, "_engine_identity", return_value=(
                self.api.EXPECTED_M2SERVER_VERSION, self.api.EXPECTED_M2SERVER_SHA256, [],
            )),
            patch.object(self.api, "_secret_directory", return_value=self.root / "secrets"),
            patch.object(self.api, "_restrict_secret_acl", return_value=None),
            patch.object(self.api, "_verify_secret_acl", return_value=None),
            patch.object(self.api, "_assert_secret_path_no_reparse", return_value=None),
            patch.object(self.api, "_dpapi_protect", side_effect=lambda value: value),
            patch.object(self.api, "_dpapi_unprotect", side_effect=lambda value: value),
        )
        for item in self.patches:
            item.start()

    def tearDown(self) -> None:
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def _workbook(self) -> Path:
        path = self.root / "native-p2.xlsx"
        wb = Workbook()
        definitions = wb.active
        definitions.title = "候选定义"
        definitions.append(V2_DEFINITION_HEADERS)
        definitions.append((
            "tiebi-vitality-p2", True, "鞭尸灵玉", "灵玉", "【三元命格·测试】", 255,
            "P2三属性隔离样本", "防御已实测；生命值、魔法值依据官网说明待游戏实测",
        ))
        properties = wb.create_sheet("自定义属性")
        properties.append(V2_PROPERTY_HEADERS)
        properties.append(("tiebi-vitality-p2", 1, "防御", 5, True, "已实测确认", "绑定1、行17"))
        properties.append(("tiebi-vitality-p2", 2, "生命值", 100, True, "网上资料待实测", "绑定6、行18"))
        properties.append(("tiebi-vitality-p2", 3, "魔法值", 100, True, "网上资料待实测", "绑定7、行19"))
        segments = wb.create_sheet("显示片段")
        segments.append(V2_SEGMENT_HEADERS)
        rows = (
            (1, "命格名称", None, None, 249),
            (2, "属性名称", "防御", None, 250), (3, "正负号", "防御", None, 69), (4, "属性值", "防御", None, 69),
            (5, "固定文字", None, "·", 255),
            (6, "属性名称", "生命值", None, 251), (7, "正负号", "生命值", None, 69), (8, "属性值", "生命值", None, 69),
            (9, "固定文字", None, "·", 255),
            (10, "属性名称", "魔法值", None, 250), (11, "正负号", "魔法值", None, 69), (12, "属性值", "魔法值", None, 69),
        )
        for order, role, source, text, color in rows:
            segments.append(("tiebi-vitality-p2", order, role, source, text, color, True, "", "P2受限显示片段"))
        part_help = wb.create_sheet("部位说明")
        part_help.append(V2_PART_HELP_HEADERS)
        part_help.append(("灵玉", "支持", "当前端位置17"))
        property_help = wb.create_sheet("属性说明")
        property_help.append(V2_PROPERTY_HELP_HEADERS)
        property_help.append(("防御", "已实测确认", "", "绑定1、行17"))
        property_help.append(("生命值", "网上资料待实测", "", "官网绑定6；P2待验收"))
        property_help.append(("魔法值", "网上资料待实测", "", "官网绑定7；P2待验收"))
        wb.save(path)
        wb.close()
        return path

    def _server(self, name: str = "MirServer") -> Path:
        server = self.root / name
        envir = server / "Mir200" / "Envir"
        (envir / "Market_Def").mkdir(parents=True)
        (envir / "QuestDiary" / "玄渊验收").mkdir(parents=True)
        (envir / "UserCmd.txt").write_bytes("命格平台验收\t97\r\n".encode("gb18030"))
        (envir / "Market_Def" / "QFunction-0.txt").write_bytes(
            ("[@Login]\r\n#ACT\r\nBREAK\r\n\r\n"
             "; XY-MG-NATIVE-P1-BEGIN\r\n[@UserCmd97]\r\n#CALL [\\玄渊验收\\命格平台单件验收.txt] @XY_MG_NATIVE_P1_MAIN\r\n; XY-MG-NATIVE-P1-END\r\n").encode("gb18030")
        )
        p1 = envir / "QuestDiary" / "玄渊验收" / "命格平台单件验收.txt"
        p1.write_bytes(b"P1-SENTINEL\r\n")
        return server

    def _business_bytes(self, plan) -> dict[Path, bytes | None]:
        return {
            item.path: item.path.read_bytes() if item.path.exists() else None
            for item in plan.files
        }

    def _restore_business_bytes(self, snapshot: dict[Path, bytes | None]) -> None:
        for path, payload in snapshot.items():
            if payload is None:
                if path.exists():
                    path.unlink()
            else:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(payload)

    def _transaction_ids(self, platform: Path) -> list[str]:
        base = platform / "backups" / "mingge-native-p2-test"
        if not base.is_dir():
            return []
        return sorted(item.name for item in base.iterdir() if item.is_dir() and item.name != "active")

    def _reseal_journal(self, api, platform: Path, transaction_id: str, server: Path, mutate) -> None:
        """Rewrite a sealed test journal so semantic guards, not the HMAC, reject it."""
        transaction, manifest, key, _ = api._load_sealed_manifest(
            platform, transaction_id, server,
        )
        paths = sorted((transaction / "journal").glob("*.json"))
        events = [json.loads(path.read_text(encoding="utf-8")) for path in paths]
        mutate(events)
        previous = ""
        for sequence, (path, event) in enumerate(zip(paths, events), start=1):
            event.pop("hmac", None)
            event["sequence"] = sequence
            event["previous_event_hmac"] = previous
            event["hmac"] = api._seal(event, key, api.JOURNAL_DOMAIN)
            path.write_bytes(api._canonical_json_bytes(event))
            previous = str(event["hmac"])
        generation = str(manifest["generation"])
        coordination = api._require_coordination(server, transaction_id, generation, key)
        coordination["journal_sequence"] = len(events)
        coordination["journal_head_hmac"] = previous
        coordination["phase"] = str(events[-1]["state"])
        coordination["pending_event"] = None
        api._write_coordination(coordination, key)

    def _recover_without_exception(self, api, platform: Path, transaction_id: str, server: Path):
        try:
            return api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
        except BaseException as exc:
            self.fail(f"preparation reservation should be explicitly recoverable: {type(exc).__name__}: {exc}")

    def test_preflight_accepts_exact_three_direct_properties_and_preserves_p1(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        p1 = server / "Mir200" / "Envir" / "QuestDiary" / "玄渊验收" / "命格平台单件验收.txt"
        p1_before = p1.read_bytes()
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        self.assertEqual(plan.blockers, ())
        self.assertEqual(plan.evidence_status, ("已实测确认", "网上资料待实测", "网上资料待实测"))
        self.assertEqual({item.relative_path for item in plan.files}, {
            "Mir200/Envir/UserCmd.txt",
            "Mir200/Envir/Market_Def/QFunction-0.txt",
            "Mir200/Envir/QuestDiary/玄渊验收/命格平台三属性验收.txt",
        })
        self.assertEqual(p1.read_bytes(), p1_before)
        script = next(item.after for item in plan.files if item.relative_path.endswith("命格平台三属性验收.txt")).decode("gb18030")
        self.assertIn("SetCustomItemValueEx 17 17 = 5 0 0", script)
        self.assertIn("SetCustomItemValueEx 17 18 = 100 0 0", script)
        self.assertIn("SetCustomItemValueEx 17 19 = 100 0 0", script)
        self.assertIn("SetCustomItemText 17 {【三元命格·测试】|249}{防御|250}{+|69}{5|69}{·|255}{生命值|251}{+|69}{100|69}{·|255}{魔法值|250}{+|69}{100|69}", script)
        self.assertEqual(script.splitlines().count("LockUpdateItem 17"), 1)
        self.assertEqual(script.splitlines().count("UpdateItem 17"), 1)
        self.assertNotIn("@RecalcAbilitys", script)

    def test_script_has_baseline_idempotent_and_three_stage_all_row_gates(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", self._server())
        script = next(item.after for item in plan.files if item.relative_path.endswith("命格平台三属性验收.txt")).decode("gb18030")
        for row in (17, 18, 19):
            self.assertGreaterEqual(script.count(f"GetCustomItemAbil 17 {row} 0"), 3)
            self.assertGreaterEqual(script.count(f"GetCustomItemValueEx 17 {row}"), 3)
        self.assertIn("@XY_MG_NATIVE_P2_ALREADY", script)
        self.assertIn("属性行18、19必须为空", script)
        self.assertIn("写后读回", script)

    def test_preflight_rejects_wrong_or_script_dependent_candidate(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        workbook = self._workbook()
        wb = load_workbook(workbook)
        wb["自定义属性"]["C4"] = "神力倍攻"
        for row in wb["显示片段"].iter_rows(min_row=2):
            if row[3].value == "魔法值":
                row[3].value = "神力倍攻"
        wb.save(workbook)
        wb.close()
        plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", self._server())
        self.assertTrue(any("精确三属性" in blocker or "脚本" in blocker for blocker in plan.blockers))

    def test_preflight_rejects_non_golden_id_even_when_structure_matches(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        workbook = self._workbook()
        wb = load_workbook(workbook)
        for sheet_name in ("候选定义", "自定义属性", "显示片段"):
            for row in wb[sheet_name].iter_rows(min_row=2):
                if row[0].value == "tiebi-vitality-p2":
                    row[0].value = "lookalike-not-golden"
        wb.save(workbook)
        wb.close()
        plan = api.plan_native_p2_test_install(workbook, "lookalike-not-golden", self._server())
        self.assertTrue(any("固定候选ID" in blocker for blocker in plan.blockers))

    def test_preflight_is_zero_write(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        before = {path.relative_to(server).as_posix(): path.read_bytes() for path in server.rglob("*") if path.is_file()}
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        self.assertEqual(plan.blockers, ())
        after = {path.relative_to(server).as_posix(): path.read_bytes() for path in server.rglob("*") if path.is_file()}
        self.assertEqual(after, before)

    def test_install_and_rollback_are_byte_exact_and_receipt_warns_runtime_instance(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = {item.relative_path: item.path.read_bytes() if item.path.exists() else None for item in plan.files}
        receipt = api.install_native_p2_test_candidate(plan, platform)
        data = json.loads(receipt.receipt_path.read_text(encoding="utf-8"))
        self.assertTrue(data["runtime_item_instance_may_change"])
        self.assertFalse(data["writes_database_directly"])
        self.assertEqual(data["game_validation_status"], "pending")
        api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        for item in plan.files:
            expected = before[item.relative_path]
            self.assertEqual(item.path.read_bytes() if item.path.exists() else None, expected)

    def test_install_refuses_post_preflight_drift(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        qfunction = server / "Mir200" / "Envir" / "Market_Def" / "QFunction-0.txt"
        qfunction.write_bytes(qfunction.read_bytes() + b"; DRIFT\r\n")
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "预检后发生变化"):
            api.install_native_p2_test_candidate(plan, self.root / "platform")

    def test_receipt_tampering_is_ignored_but_manifest_tampering_blocks_without_writes(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        unrelated = server / "Mir200" / "Envir" / "unrelated.txt"
        unrelated.write_bytes(b"KEEP")
        data = json.loads(receipt.receipt_path.read_text(encoding="utf-8"))
        data["files"].append({
            "relative_path": "Mir200/Envir/unrelated.txt", "existed_before": False,
            "before_sha256": None, "after_sha256": "bad", "backup_index": None,
        })
        receipt.receipt_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        installed = {item.path: item.path.read_bytes() for item in plan.files}
        manifest_path = receipt.receipt_path.parent / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["candidate"]["payment"] = True
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "manifest.*封印"):
            api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        self.assertEqual(unrelated.read_bytes(), b"KEEP")
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes(), payload)

    def test_rollback_validates_all_backups_before_restoring_first_file(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        installed = {item.path: item.path.read_bytes() for item in plan.files}
        transaction = receipt.receipt_path.parent
        (transaction / "files" / "01.bin").unlink()
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "备份缺失"):
            api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes(), payload)

    def test_preflight_blocks_false_game_verified_claims_and_duplicate_usercmd98(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        workbook = self._workbook()
        wb = load_workbook(workbook)
        wb["自定义属性"]["F3"] = "已实测确认"
        wb["属性说明"]["B3"] = "已实测确认"
        wb.save(workbook)
        wb.close()
        server = self._server()
        qfunction = server / "Mir200" / "Envir" / "Market_Def" / "QFunction-0.txt"
        qfunction.write_bytes(qfunction.read_bytes() + b"\r\n[@UserCmd98]\r\n#ACT\r\nBREAK\r\n")
        plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        self.assertTrue(any("证据状态" in item or "不得把生命值" in item for item in plan.blockers))
        self.assertTrue(any("@UserCmd98" in item for item in plan.blockers))

    def test_preflight_blocks_exact_duplicate_usercmd_mapping(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        usercmd = server / "Mir200" / "Envir" / "UserCmd.txt"
        usercmd.write_bytes(usercmd.read_bytes() + "命格平台三属性验收\t98\r\n命格平台三属性验收\t98\r\n".encode("gb18030"))
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        self.assertTrue(any("UserCmd" in item and "唯一" in item for item in plan.blockers))

    def test_evidence_drift_after_preflight_blocks_with_zero_server_writes(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        source = self.root / "official.htm"
        source.write_bytes(b"official-v1")
        digest = api._sha256_bytes(source.read_bytes()).upper()
        patched_evidence = tuple(
            api.NativeP2Evidence(
                item.property_name,
                item.evidence_status,
                item.source_id,
                item.source_type,
                item.source_locator,
                str(source) if item.source_type == "official_web" else item.snapshot_path,
                item.captured_at,
                digest if item.source_type == "official_web" else item.source_sha256,
            )
            for item in api.P2_EVIDENCE
        )
        server = self._server()
        with patch.object(api, "P2_EVIDENCE", patched_evidence):
            plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
            self.assertEqual(plan.blockers, ())
            before = {item.path: item.path.read_bytes() if item.path.exists() else None for item in plan.files}
            source.write_bytes(b"official-v2")
            with self.assertRaisesRegex(api.NativeP2TestInstallError, "证据在预检后发生变化"):
                api.install_native_p2_test_candidate(plan, self.root / "platform")
            for path, payload in before.items():
                self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    def test_final_presentation_receipt_failure_keeps_sealed_installed_state_and_rollback(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        real_atomic_write = api._atomic_write

        def fail_final_receipt(path: Path, payload: bytes) -> None:
            if path.name == "receipt.json" and b'"status": "INSTALLED"' in payload:
                raise OSError("fault-injected final presentation receipt failure")
            real_atomic_write(path, payload)

        with patch.object(api, "_atomic_write", side_effect=fail_final_receipt):
            receipt = api.install_native_p2_test_candidate(plan, platform)
        state = api.inspect_native_p2_test_transaction(platform, receipt.transaction_id, server)
        self.assertEqual(state["state"], "INSTALLED")
        api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)

    def test_recovery_rollback_handles_interrupted_mixed_before_after_state(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = {item.path: item.path.read_bytes() if item.path.exists() else None for item in plan.files}
        receipt = api.install_native_p2_test_candidate(plan, platform)
        transaction = receipt.receipt_path.parent
        manifest = json.loads((transaction / "manifest.json").read_text(encoding="utf-8"))
        first = plan.files[0]
        first.path.write_bytes((transaction / "files" / "00.bin").read_bytes())
        recovered = api.recover_native_p2_test_transaction(platform, receipt.transaction_id, server, "rollback")
        self.assertEqual(recovered.status, "rolled-back")
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    def test_journal_tampering_blocks_rollback_without_touching_installed_files(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        installed = {item.path: item.path.read_bytes() for item in plan.files}
        latest = sorted((receipt.receipt_path.parent / "journal").glob("*.json"))[-1]
        event = json.loads(latest.read_text(encoding="utf-8"))
        event["state"] = "ROLLED_BACK"
        latest.write_text(json.dumps(event, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "journal.*封印"):
            api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes(), payload)

    def test_install_complete_journal_failure_auto_restores_all_business_files(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = {item.path: item.path.read_bytes() if item.path.exists() else None for item in plan.files}
        real_append = api._append_journal

        def fail_installed(*args, **kwargs):
            state = args[3] if len(args) > 3 else kwargs.get("state")
            if state == "INSTALLED":
                raise OSError("fault-injected INSTALL_COMPLETE journal failure")
            return real_append(*args, **kwargs)

        with patch.object(api, "_append_journal", side_effect=fail_installed):
            with self.assertRaisesRegex(OSError, "INSTALL_COMPLETE"):
                api.install_native_p2_test_candidate(plan, platform)
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    def test_preparation_crash_immediately_after_coordination_reservation_is_abortable(self) -> None:
        """Leaving a current coordination reservation without a manifest must not deadlock installs."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        workbook = self._workbook()
        plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)

        def crash(name: str) -> None:
            if name == "after_coordination_reservation":
                raise _SimulatedProcessCrash("fault after coordination reservation")

        with patch.object(api, "_transaction_fault", side_effect=crash):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "coordination reservation"):
                api.install_native_p2_test_candidate(plan, platform)
        transaction_id, = self._transaction_ids(platform)
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)
        first = self._recover_without_exception(api, platform, transaction_id, server)
        second = self._recover_without_exception(api, platform, transaction_id, server)
        self.assertEqual((first.status, second.status), ("rolled-back", "rolled-back"))
        next_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        self.assertEqual(api.install_native_p2_test_candidate(next_plan, platform).status, "installed-test-candidate")

    def test_preparation_backup_write_failure_is_abortable_and_next_install_can_start(self) -> None:
        """A failed backup write after reservation must have a protected zero-write abort path."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        workbook = self._workbook()
        plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)
        real_atomic_write = api._atomic_write

        def fail_first_backup(path: Path, payload: bytes) -> None:
            if path.parent.name == "files" and path.name == "00.bin":
                raise OSError("fault-injected backup write failure")
            real_atomic_write(path, payload)

        with patch.object(api, "_atomic_write", side_effect=fail_first_backup):
            with self.assertRaisesRegex(OSError, "backup write failure"):
                api.install_native_p2_test_candidate(plan, platform)
        transaction_id, = self._transaction_ids(platform)
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)
        recovered = self._recover_without_exception(api, platform, transaction_id, server)
        self.assertEqual(recovered.status, "rolled-back")
        next_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        self.assertEqual(api.install_native_p2_test_candidate(next_plan, platform).status, "installed-test-candidate")

    def test_preparation_crash_after_manifest_json_before_hmac_is_abortable(self) -> None:
        """A manifest.json orphan must be abortable only through the protected zero-write reservation."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        workbook = self._workbook()
        plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)

        def crash(name: str) -> None:
            if name == "after_manifest_json":
                raise _SimulatedProcessCrash("fault after manifest.json")

        with patch.object(api, "_transaction_fault", side_effect=crash):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "manifest.json"):
                api.install_native_p2_test_candidate(plan, platform)
        transaction_id, = self._transaction_ids(platform)
        transaction = platform / "backups" / "mingge-native-p2-test" / transaction_id
        self.assertTrue((transaction / "manifest.json").is_file())
        self.assertFalse((transaction / "manifest.hmac").exists())
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)
        recovered = self._recover_without_exception(api, platform, transaction_id, server)
        self.assertEqual(recovered.status, "rolled-back")
        next_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        self.assertEqual(api.install_native_p2_test_candidate(next_plan, platform).status, "installed-test-candidate")

    def test_preparation_crash_after_manifest_hmac_before_verify_is_abortable(self) -> None:
        """A complete but unverified manifest pair must not strand the protected reservation."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        workbook = self._workbook()
        plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)

        def crash(name: str) -> None:
            if name == "after_manifest_hmac":
                raise _SimulatedProcessCrash("fault after manifest.hmac")

        with patch.object(api, "_transaction_fault", side_effect=crash):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "manifest.hmac"):
                api.install_native_p2_test_candidate(plan, platform)
        transaction_id, = self._transaction_ids(platform)
        transaction = platform / "backups" / "mingge-native-p2-test" / transaction_id
        self.assertTrue((transaction / "manifest.json").is_file())
        self.assertTrue((transaction / "manifest.hmac").is_file())
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)
        first = self._recover_without_exception(api, platform, transaction_id, server)
        second = self._recover_without_exception(api, platform, transaction_id, server)
        self.assertEqual((first.status, second.status), ("rolled-back", "rolled-back"))
        next_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        self.assertEqual(api.install_native_p2_test_candidate(next_plan, platform).status, "installed-test-candidate")

    def test_acl_failure_before_key_creation_leaves_no_key_and_cannot_be_bypassed_next_run(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        base = self.root / "platform" / "backups" / "mingge-native-p2-test"
        base.mkdir(parents=True)
        with patch.object(api, "_restrict_secret_acl", side_effect=api.NativeP2TestInstallError("ACL failure")):
            with self.assertRaisesRegex(api.NativeP2TestInstallError, "ACL failure"):
                api._load_or_create_integrity_key(base)
            self.assertFalse(api._integrity_key_path().exists())
            with self.assertRaisesRegex(api.NativeP2TestInstallError, "ACL failure"):
                api._load_or_create_integrity_key(base)
            self.assertFalse(api._integrity_key_path().exists())

    def test_raw_server_root_reparse_entry_is_blocked(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        original = api._is_reparse_point
        raw_server_entry_seen = False

        def fake_reparse(path: Path) -> bool:
            nonlocal raw_server_entry_seen
            if not raw_server_entry_seen and Path(path).absolute() == server.absolute():
                raw_server_entry_seen = True
                return True
            return original(Path(path))

        with patch.object(api, "_is_reparse_point", side_effect=fake_reparse):
            plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        self.assertTrue(any("服务端入口" in item for item in plan.blockers))

    def test_raw_platform_root_symlink_entry_is_blocked_before_resolve(self) -> None:
        """Resolving the raw platform entry before checking it must make this test fail."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        real_platform = self.root / "platform-real"
        raw_platform = self.root / "platform-link"
        real_platform.mkdir()
        try:
            os.symlink(real_platform, raw_platform, target_is_directory=True)
        except OSError as exc:
            self.skipTest(f"directory symlink unavailable: {exc}")
        before = {item.path: item.path.read_bytes() if item.path.exists() else None for item in plan.files}
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "平台.*重解析|平台.*符号链接"):
            api.install_native_p2_test_candidate(plan, raw_platform)
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    def test_mutated_plan_files_are_bound_to_fixed_targets_and_after_bytes_with_zero_writes(self) -> None:
        """Removing count/order/path/after binding from the install gate must make a subcase fail."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        outside = self.root / "outside.txt"
        outside.write_bytes(b"OUTSIDE-SENTINEL")
        for name in ("count", "order", "relative", "absolute", "after_sha256", "after_bytes"):
            with self.subTest(name=name):
                server = self._server(f"MirServer-{name}")
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
                original = {path: path.read_bytes() for path in server.rglob("*") if path.is_file()}
                first = plan.files[0]
                mutated_files = {
                    "count": plan.files[:2],
                    "order": (plan.files[1], plan.files[0], plan.files[2]),
                    "relative": (replace(first, relative_path="Mir200/Envir/Other.txt"), *plan.files[1:]),
                    "absolute": (replace(first, path=outside), *plan.files[1:]),
                    "after_sha256": (replace(first, after_sha256="0" * 64), *plan.files[1:]),
                    "after_bytes": (replace(first, after=first.after + b"DRIFT"), *plan.files[1:]),
                }[name]
                mutated = replace(plan, files=mutated_files)
                with self.assertRaisesRegex(api.NativeP2TestInstallError, "固定三文件|计划文件|安装后内容"):
                    api.install_native_p2_test_candidate(mutated, self.root / f"platform-{name}")
                self.assertEqual(outside.read_bytes(), b"OUTSIDE-SENTINEL")
                for path, payload in original.items():
                    self.assertEqual(path.read_bytes(), payload)

    def test_truncated_old_journal_and_missing_owner_cannot_revive_equal_after_new_generation(self) -> None:
        """Accepting a valid old HMAC prefix must let this test reproduce the generation rollback attack."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        workbook = self._workbook()
        first_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        first = api.install_native_p2_test_candidate(first_plan, platform)
        api.rollback_native_p2_test_install(platform, first.transaction_id, server)
        second_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        second = api.install_native_p2_test_candidate(second_plan, platform)
        installed = {item.path: item.path.read_bytes() for item in second_plan.files}
        base = platform / "backups" / "mingge-native-p2-test"
        owner = api._active_owner_path(base, server)
        owner.unlink()
        first_journal = first.receipt_path.parent / "journal"
        for event_path in sorted(first_journal.glob("*.json"))[4:]:
            event_path.unlink()
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "协调|尾锚|安装代|日志.*截断|归属"):
            api.recover_native_p2_test_transaction(platform, first.transaction_id, server, "finalize-install")
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "协调|尾锚|安装代|日志.*截断|归属|恢复预检"):
            api.rollback_native_p2_test_install(platform, first.transaction_id, server)
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes(), payload)
        self.assertEqual(
            api.inspect_native_p2_test_transaction(platform, second.transaction_id, server)["state"],
            "INSTALLED",
        )

    def test_empty_advanced_journal_is_not_manifest_only_and_business_files_stay_unchanged(self) -> None:
        """Treating an empty deleted journal as MANIFEST_ONLY must make this test fail."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        installed = {item.path: item.path.read_bytes() for item in plan.files}
        api._active_owner_path(platform / "backups" / "mingge-native-p2-test", server).unlink()
        for event_path in (receipt.receipt_path.parent / "journal").glob("*.json"):
            event_path.unlink()
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "协调|尾锚|journal.*截|日志.*截断|归属"):
            api.recover_native_p2_test_transaction(platform, receipt.transaction_id, server, "finalize-install")
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes(), payload)

    def test_crash_after_installed_journal_before_owner_recovers_deterministically(self) -> None:
        """Removing the recoverable INSTALLED-to-owner transition must make this test fail."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        workbook = self._workbook()
        first_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        first = api.install_native_p2_test_candidate(first_plan, platform)
        api.rollback_native_p2_test_install(platform, first.transaction_id, server)
        plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)

        def crash_at_boundary(name: str) -> None:
            if name == "after_installed_journal":
                raise RuntimeError("fault-injected crash after durable INSTALLED journal")

        with patch.object(api, "_transaction_fault", side_effect=crash_at_boundary, create=True):
            with self.assertRaisesRegex(RuntimeError, "durable INSTALLED"):
                api.install_native_p2_test_candidate(plan, platform)
        base = platform / "backups" / "mingge-native-p2-test"
        transactions = [item for item in base.iterdir() if item.is_dir() and item.name != "active"]
        self.assertEqual(len(transactions), 2)
        transaction_id = next(item.name for item in transactions if item.name != first.transaction_id)
        recovered = api.recover_native_p2_test_transaction(platform, transaction_id, server, "finalize-install")
        self.assertEqual(recovered.status, "installed-finalized")
        self.assertTrue(api.inspect_native_p2_test_transaction(platform, transaction_id, server)["owns_active_generation"])
        api.rollback_native_p2_test_install(platform, transaction_id, server)

    def test_crash_after_rolled_back_journal_before_owner_close_recovers_and_allows_next_install(self) -> None:
        """Removing the recoverable ROLLED_BACK-to-owner-close transition must make this test fail."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        workbook = self._workbook()
        plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        before = {item.path: (receipt.receipt_path.parent / "files" / f"{index:02d}.bin").read_bytes()
                  if item.existed_before else None for index, item in enumerate(plan.files)}

        def crash_at_boundary(name: str) -> None:
            if name == "after_rolled_back_journal":
                raise RuntimeError("fault-injected crash after durable ROLLED_BACK journal")

        with patch.object(api, "_transaction_fault", side_effect=crash_at_boundary, create=True):
            with self.assertRaisesRegex(RuntimeError, "durable ROLLED_BACK"):
                api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)
        recovered = api.recover_native_p2_test_transaction(platform, receipt.transaction_id, server, "rollback")
        self.assertEqual(recovered.status, "rolled-back")
        next_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        next_receipt = api.install_native_p2_test_candidate(next_plan, platform)
        self.assertEqual(next_receipt.status, "installed-test-candidate")

    @unittest.skipUnless(os.name == "nt", "Windows directory-handle integration test")
    def test_install_holds_target_parent_handle_without_delete_share_against_replacement(self) -> None:
        """Removing the held no-delete-share directory handles must redirect this write externally."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        target = plan.files[1].path
        parent = server / "Mir200" / "Envir"
        moved = parent.with_name(parent.name + "-moved")
        external = self.root / "external-install"
        external_target = external / "Market_Def" / "QFunction-0.txt"
        external_target.parent.mkdir(parents=True)
        external_target.write_bytes(b"EXTERNAL-INSTALL-SENTINEL")
        attack = {"attempted": False, "prevented": False}

        def attempt_replacement(name: str) -> None:
            if name == "after_target_leaves_validated_for_install" and not attack["attempted"]:
                attack["attempted"] = True
                try:
                    os.replace(parent, moved)
                except OSError:
                    attack["prevented"] = True
                else:
                    os.symlink(external, parent, target_is_directory=True)

        with patch.object(api, "_transaction_fault", side_effect=attempt_replacement):
            api.install_native_p2_test_candidate(plan, platform)
        self.assertTrue(attack["attempted"])
        self.assertTrue(attack["prevented"])
        self.assertEqual(external_target.read_bytes(), b"EXTERNAL-INSTALL-SENTINEL")

    @unittest.skipUnless(os.name == "nt", "Windows directory-handle integration test")
    def test_rollback_holds_target_parent_handle_without_delete_share_against_replacement(self) -> None:
        """Removing the rollback directory handles must redirect the byte-exact restore externally."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        target = plan.files[1].path
        parent = server / "Mir200" / "Envir"
        moved = parent.with_name(parent.name + "-moved")
        external = self.root / "external-rollback"
        external_target = external / "Market_Def" / "QFunction-0.txt"
        external_target.parent.mkdir(parents=True)
        installed_payload = target.read_bytes()
        external_target.write_bytes(installed_payload)
        attack = {"attempted": False, "prevented": False}

        def attempt_replacement(name: str) -> None:
            if name == "after_target_leaves_validated_for_rollback" and not attack["attempted"]:
                attack["attempted"] = True
                try:
                    os.replace(parent, moved)
                except OSError:
                    attack["prevented"] = True
                else:
                    os.symlink(external, parent, target_is_directory=True)

        with patch.object(api, "_transaction_fault", side_effect=attempt_replacement):
            api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        self.assertTrue(attack["attempted"])
        self.assertTrue(attack["prevented"])
        self.assertEqual(external_target.read_bytes(), installed_payload)

    @unittest.skipUnless(os.name == "nt", "Windows backup-leaf handle integration test")
    def test_validated_backup_leaves_reject_write_and_delete_before_rollback_uses_them(self) -> None:
        """A backup leaf mutation after full validation must stop before any business overwrite."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        installed = self._business_bytes(plan)
        transaction = receipt.receipt_path.parent
        backup_write = transaction / "files" / "00.bin"
        backup_delete = transaction / "files" / "01.bin"
        moved = transaction / "files" / "01.bin.moved"
        real_append = api._append_journal
        attack = {"attempted": False, "write_blocked": False, "delete_blocked": False}

        def attack_after_validation(*args, **kwargs):
            event = real_append(*args, **kwargs)
            action = args[4] if len(args) > 4 else kwargs.get("action")
            if action == "ROLLBACK_START" and not attack["attempted"]:
                attack["attempted"] = True
                try:
                    backup_write.write_bytes(b"ATTACK-BACKUP-WRITE")
                except OSError:
                    attack["write_blocked"] = True
                try:
                    os.replace(backup_delete, moved)
                except OSError:
                    attack["delete_blocked"] = True
                raise RuntimeError("stop after backup leaf race probe")
            return event

        with patch.object(api, "_append_journal", side_effect=attack_after_validation):
            with self.assertRaisesRegex(RuntimeError, "backup leaf race probe"):
                api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        self.assertTrue(attack["attempted"])
        self.assertTrue(attack["write_blocked"])
        self.assertTrue(attack["delete_blocked"])
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows backup-leaf handle integration test")
    def test_automatic_recovery_uses_held_validated_backup_bytes_not_a_reopened_path(self) -> None:
        """Mutating 00.bin during install failure recovery must never inject bytes into UserCmd."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)
        real_append = api._append_journal
        attack = {"attempted": False, "blocked": False}

        def fail_second_write_after_backup_attack(*args, **kwargs):
            event = real_append(*args, **kwargs)
            action = args[4] if len(args) > 4 else kwargs.get("action")
            index = args[5] if len(args) > 5 else kwargs.get("file_index")
            if action == "WRITE_INTENT" and index == 1 and not attack["attempted"]:
                attack["attempted"] = True
                backup = Path(args[0]) / "files" / "00.bin"
                try:
                    backup.write_bytes(b"ATTACK-THROUGH-TRUSTED-BACKUP")
                except OSError:
                    attack["blocked"] = True
                    raise
                raise RuntimeError("fault after backup mutation")
            return event

        with patch.object(api, "_append_journal", side_effect=fail_second_write_after_backup_attack):
            with self.assertRaises((OSError, RuntimeError)):
                api.install_native_p2_test_candidate(plan, platform)
        self.assertTrue(attack["attempted"])
        self.assertTrue(attack["blocked"])
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows target-leaf handle integration test")
    def test_install_target_leaf_drift_after_last_classification_is_not_overwritten(self) -> None:
        """A writer entering after target classification must be denied, not overwritten by install."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)
        target = plan.files[0].path
        real_atomic_write = api._atomic_write
        attack = {"attempted": False}

        def fault(name: str) -> None:
            if name == "after_target_leaves_validated_for_install":
                attack["attempted"] = True
                target.write_bytes(b"UNKNOWN-INSTALL-LEAF")

        def reproduce_old_window(path: Path, payload: bytes) -> None:
            if Path(path) == target and not attack["attempted"]:
                attack["attempted"] = True
                target.write_bytes(b"UNKNOWN-INSTALL-LEAF")
            real_atomic_write(Path(path), payload)

        with (
            patch.object(api, "_transaction_fault", side_effect=fault),
            patch.object(api, "_atomic_write", side_effect=reproduce_old_window),
        ):
            with self.assertRaises(OSError):
                api.install_native_p2_test_candidate(plan, platform)
        self.assertTrue(attack["attempted"])
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows target-leaf handle integration test")
    def test_rollback_target_leaf_drift_after_last_classification_is_not_overwritten(self) -> None:
        """A writer entering after rollback classification must be denied before any target changes."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        installed = self._business_bytes(plan)
        target = plan.files[1].path
        real_atomic_write = api._atomic_write
        attack = {"attempted": False}

        def fault(name: str) -> None:
            if name == "after_target_leaves_validated_for_rollback":
                attack["attempted"] = True
                target.write_bytes(b"UNKNOWN-ROLLBACK-LEAF")

        def reproduce_old_window(path: Path, payload: bytes) -> None:
            if Path(path) == target and not attack["attempted"]:
                attack["attempted"] = True
                target.write_bytes(b"UNKNOWN-ROLLBACK-LEAF")
            real_atomic_write(Path(path), payload)

        with (
            patch.object(api, "_transaction_fault", side_effect=fault),
            patch.object(api, "_atomic_write", side_effect=reproduce_old_window),
        ):
            with self.assertRaises(OSError):
                api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        self.assertTrue(attack["attempted"])
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows CREATE_NEW target integration test")
    def test_install_missing_target_uses_create_new_and_never_overwrites_a_racing_creator(self) -> None:
        """The absent test-script leaf must be reserved with CREATE_NEW at its write transition."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)
        target = plan.files[2].path
        attacker = b"RACING-CREATOR-SENTINEL"
        real_atomic_write = api._atomic_write
        attack = {"introduced": False}

        def fault(name: str) -> None:
            if name == "before_create_new_target:2" and not attack["introduced"]:
                target.write_bytes(attacker)
                attack["introduced"] = True

        def reproduce_old_window(path: Path, payload: bytes) -> None:
            if Path(path) == target and not attack["introduced"]:
                target.write_bytes(attacker)
                attack["introduced"] = True
            real_atomic_write(Path(path), payload)

        with (
            patch.object(api, "_transaction_fault", side_effect=fault),
            patch.object(api, "_atomic_write", side_effect=reproduce_old_window),
        ):
            with self.assertRaises((OSError, api.NativeP2TestInstallError)):
                api.install_native_p2_test_candidate(plan, platform)
        self.assertTrue(attack["introduced"])
        self.assertEqual(target.read_bytes(), attacker)
        for item in plan.files[:2]:
            self.assertEqual(item.path.read_bytes(), before[item.path])

    @unittest.skipUnless(os.name == "nt", "Windows same-handle delete integration test")
    def test_rollback_new_target_uses_same_exclusive_handle_for_delete(self) -> None:
        """Closing a target handle before unlink would reopen a write/delete race."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        installed = self._business_bytes(plan)
        target = plan.files[2].path
        moved = target.with_name(target.name + ".moved")
        real_unlink = os.unlink
        attack = {"attempted": False, "write_blocked": False, "delete_blocked": False}

        def fault(name: str) -> None:
            if name == "before_delete_target_handle:2":
                attack["attempted"] = True
                try:
                    target.write_bytes(b"UNKNOWN-BEFORE-DELETE")
                except OSError:
                    attack["write_blocked"] = True
                try:
                    os.replace(target, moved)
                except OSError:
                    attack["delete_blocked"] = True
                raise RuntimeError("stop before same-handle delete")

        def reproduce_old_window(path, *, dir_fd=None):
            candidate = Path(path)
            if candidate == target and not attack["attempted"]:
                attack["attempted"] = True
                target.write_bytes(b"UNKNOWN-BEFORE-DELETE")
            if dir_fd is None:
                return real_unlink(path)
            return real_unlink(path, dir_fd=dir_fd)

        with (
            patch.object(api, "_transaction_fault", side_effect=fault),
            patch.object(os, "unlink", side_effect=reproduce_old_window),
        ):
            with self.assertRaisesRegex(RuntimeError, "same-handle delete"):
                api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        self.assertTrue(attack["attempted"])
        self.assertTrue(attack["write_blocked"])
        self.assertTrue(attack["delete_blocked"])
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows same-handle crash recovery test")
    def test_install_same_handle_prefix_crash_recovers_only_current_write_intent(self) -> None:
        """A crash after truncate may recover an exact prefix of the current install payload."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)

        def crash(name: str) -> None:
            if name == "install_after_target_handle_truncate:0":
                raise _SimulatedProcessCrash("install crashed after target truncate")

        with patch.object(api, "_transaction_fault", side_effect=crash):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "install crashed"):
                api.install_native_p2_test_candidate(plan, platform)
        transaction_id, = self._transaction_ids(platform)
        self.assertEqual(plan.files[0].path.read_bytes(), b"")
        recovered = api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
        self.assertEqual(recovered.status, "rolled-back")
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows same-handle crash recovery test")
    def test_rollback_same_handle_prefix_crash_is_idempotently_recoverable(self) -> None:
        """A crash after rollback truncate may resume only the anchored rollback intent payload."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)
        receipt = api.install_native_p2_test_candidate(plan, platform)

        def crash(name: str) -> None:
            if name == "rollback_after_target_handle_truncate:1":
                raise _SimulatedProcessCrash("rollback crashed after target truncate")

        with patch.object(api, "_transaction_fault", side_effect=crash):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "rollback crashed"):
                api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        self.assertEqual(plan.files[1].path.read_bytes(), b"")
        recovered = api.recover_native_p2_test_transaction(platform, receipt.transaction_id, server, "rollback")
        self.assertEqual(recovered.status, "rolled-back")
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows two-stage crash recovery test")
    def test_install_prefix_recovery_survives_crash_after_anchored_rollback_start(self) -> None:
        """A second crash after ROLLBACK_START must not discard an accepted install-prefix intent."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)

        def crash_after_install_truncate(name: str) -> None:
            if name == "install_after_target_handle_truncate:0":
                raise _SimulatedProcessCrash("install prefix crash")

        with patch.object(api, "_transaction_fault", side_effect=crash_after_install_truncate):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "install prefix"):
                api.install_native_p2_test_candidate(plan, platform)
        transaction_id, = self._transaction_ids(platform)
        self.assertEqual(plan.files[0].path.read_bytes(), b"")

        real_append_journal = api._append_journal
        crash = {"unsafe_start": False, "safe_intent": False}

        def append_then_crash(*args, **kwargs):
            event = real_append_journal(*args, **kwargs)
            if (
                event["action"] == "ROLLBACK_START"
                and not crash["safe_intent"]
                and not crash["unsafe_start"]
            ):
                crash["unsafe_start"] = True
                raise _SimulatedProcessCrash("crash at rollback recovery boundary")
            return event

        def crash_after_safe_intent(name: str) -> None:
            if name == "rollback_prefix_reauthorized_before_write:0":
                crash["safe_intent"] = True
                raise _SimulatedProcessCrash("crash at rollback recovery boundary")

        with (
            patch.object(api, "_append_journal", side_effect=append_then_crash),
            patch.object(api, "_transaction_fault", side_effect=crash_after_safe_intent),
        ):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "rollback recovery boundary"):
                api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
        journal = platform / "backups" / "mingge-native-p2-test" / transaction_id / "journal"
        last_event = json.loads(sorted(journal.glob("*.json"))[-1].read_text(encoding="utf-8"))
        self.assertTrue(crash["safe_intent"])
        self.assertEqual(last_event["action"], "ROLLBACK_INTENT")
        self.assertEqual(last_event["file_index"], 0)
        self.assertEqual(last_event["observed_after_sha256"], plan.files[0].before_sha256)
        self.assertEqual(plan.files[0].path.read_bytes(), b"")

        try:
            recovered = api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
        except Exception as exc:
            self.fail(f"second recovery lost the protected install-prefix intent: {type(exc).__name__}: {exc}")
        self.assertEqual(recovered.status, "rolled-back")
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows two-stage crash recovery test")
    def test_rollback_prefix_recovery_survives_crash_after_anchored_rollback_start(self) -> None:
        """A second crash after ROLLBACK_START must not discard an accepted rollback-prefix intent."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)
        receipt = api.install_native_p2_test_candidate(plan, platform)

        def crash_after_rollback_truncate(name: str) -> None:
            if name == "rollback_after_target_handle_truncate:1":
                raise _SimulatedProcessCrash("rollback prefix crash")

        with patch.object(api, "_transaction_fault", side_effect=crash_after_rollback_truncate):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "rollback prefix"):
                api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
        self.assertEqual(plan.files[1].path.read_bytes(), b"")

        real_append_journal = api._append_journal
        crash = {"unsafe_start": False, "safe_intent": False}

        def append_then_crash(*args, **kwargs):
            event = real_append_journal(*args, **kwargs)
            if (
                event["action"] == "ROLLBACK_START"
                and not crash["safe_intent"]
                and not crash["unsafe_start"]
            ):
                crash["unsafe_start"] = True
                raise _SimulatedProcessCrash("crash at rollback recovery boundary")
            return event

        def crash_after_safe_intent(name: str) -> None:
            if name == "rollback_prefix_reauthorized_before_write:1":
                crash["safe_intent"] = True
                raise _SimulatedProcessCrash("crash at rollback recovery boundary")

        with (
            patch.object(api, "_append_journal", side_effect=append_then_crash),
            patch.object(api, "_transaction_fault", side_effect=crash_after_safe_intent),
        ):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "rollback recovery boundary"):
                api.recover_native_p2_test_transaction(
                    platform, receipt.transaction_id, server, "rollback",
                )
        journal = receipt.receipt_path.parent / "journal"
        last_event = json.loads(sorted(journal.glob("*.json"))[-1].read_text(encoding="utf-8"))
        self.assertTrue(crash["safe_intent"])
        self.assertEqual(last_event["action"], "ROLLBACK_INTENT")
        self.assertEqual(last_event["file_index"], 1)
        self.assertEqual(last_event["observed_after_sha256"], plan.files[1].before_sha256)
        self.assertEqual(plan.files[1].path.read_bytes(), b"")

        try:
            recovered = api.recover_native_p2_test_transaction(
                platform, receipt.transaction_id, server, "rollback",
            )
        except Exception as exc:
            self.fail(f"second recovery lost the protected rollback-prefix intent: {type(exc).__name__}: {exc}")
        self.assertEqual(recovered.status, "rolled-back")
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    def test_installed_transaction_with_deleted_active_owner_cannot_use_nonactive_prefix_resume(self) -> None:
        """Deleting an installed transaction's ordinary owner must not downgrade active rollback authority."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        installed = self._business_bytes(plan)
        base = platform / "backups" / "mingge-native-p2-test"
        api._active_owner_path(base, server).unlink()

        with self.assertRaises(api.NativeP2TestInstallError):
            api.recover_native_p2_test_transaction(
                platform, receipt.transaction_id, server, "rollback",
            )
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows protected rollback-done crash recovery test")
    def test_nonactive_prefix_recovery_done_boundary_is_replayable_but_cannot_authorize_drift(self) -> None:
        """Only a matching per-file ROLLBACK_DONE may resume the non-active current generation."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        for case in ("valid", "wrong-before-hash", "other-index", "target-not-before"):
            with self.subTest(case=case):
                server = self._server(f"MirServer-done-{case}")
                platform = self.root / f"platform-done-{case}"
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
                before = self._business_bytes(plan)

                def crash_after_install_truncate(name: str) -> None:
                    if name == "install_after_target_handle_truncate:0":
                        raise _SimulatedProcessCrash("install prefix for rollback-done boundary")

                with patch.object(api, "_transaction_fault", side_effect=crash_after_install_truncate):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "rollback-done boundary"):
                        api.install_native_p2_test_candidate(plan, platform)
                transaction_id, = self._transaction_ids(platform)
                self.assertEqual(plan.files[0].path.read_bytes(), b"")

                real_append_journal = api._append_journal

                def append_done_then_crash(*args, **kwargs):
                    values = list(args)
                    if len(values) > 4 and values[4] == "ROLLBACK_DONE":
                        if case == "wrong-before-hash":
                            values[7] = "0" * 64
                        elif case == "other-index":
                            values[5] = 1
                            values[6] = plan.files[1].after_sha256
                            values[7] = plan.files[1].before_sha256
                        event = real_append_journal(*values, **kwargs)
                        raise _SimulatedProcessCrash("crash after protected rollback done")
                    return real_append_journal(*args, **kwargs)

                with patch.object(api, "_append_journal", side_effect=append_done_then_crash):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "protected rollback done"):
                        api.recover_native_p2_test_transaction(
                            platform, transaction_id, server, "rollback",
                        )
                journal = platform / "backups" / "mingge-native-p2-test" / transaction_id / "journal"
                last_event = json.loads(sorted(journal.glob("*.json"))[-1].read_text(encoding="utf-8"))
                self.assertEqual(last_event["action"], "ROLLBACK_DONE")
                if case == "valid":
                    self.assertEqual(last_event["file_index"], 0)
                    self.assertEqual(last_event["observed_after_sha256"], plan.files[0].before_sha256)
                elif case == "target-not-before":
                    plan.files[0].path.write_bytes(b"NOT-THE-SEALED-BEFORE-PAYLOAD")

                if case == "valid":
                    try:
                        recovered = api.recover_native_p2_test_transaction(
                            platform, transaction_id, server, "rollback",
                        )
                    except Exception as exc:
                        self.fail(f"protected rollback-done boundary was not replayable: {type(exc).__name__}: {exc}")
                    self.assertEqual(recovered.status, "rolled-back")
                    for path, payload in before.items():
                        self.assertEqual(path.read_bytes() if path.exists() else None, payload)
                else:
                    drift = self._business_bytes(plan)
                    with self.assertRaises(api.NativeP2TestInstallError):
                        api.recover_native_p2_test_transaction(
                            platform, transaction_id, server, "rollback",
                        )
                    self.assertEqual(self._business_bytes(plan), drift)

    @unittest.skipUnless(os.name == "nt", "Windows pre-truncate rollback intent recovery test")
    def test_nonactive_existing_target_exact_after_intent_replays_before_truncate(self) -> None:
        """A current exact-after target may replay its matching rollback intent before truncate."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)

        def crash_install(name: str) -> None:
            if name == "install_after_target_handle_truncate:1":
                raise _SimulatedProcessCrash("install crash before exact-after rollback")

        with patch.object(api, "_transaction_fault", side_effect=crash_install):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "exact-after rollback"):
                api.install_native_p2_test_candidate(plan, platform)
        transaction_id, = self._transaction_ids(platform)
        self.assertEqual(plan.files[0].path.read_bytes(), plan.files[0].after)

        real_append_journal = api._append_journal

        def append_intent_then_crash(*args, **kwargs):
            event = real_append_journal(*args, **kwargs)
            if event["action"] == "ROLLBACK_INTENT" and event["file_index"] == 0:
                raise _SimulatedProcessCrash("crash after existing-target rollback intent")
            return event

        with patch.object(api, "_append_journal", side_effect=append_intent_then_crash):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "existing-target rollback intent"):
                api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
        transaction = platform / "backups" / "mingge-native-p2-test" / transaction_id
        tail = json.loads(sorted((transaction / "journal").glob("*.json"))[-1].read_text(encoding="utf-8"))
        self.assertEqual(tail["action"], "ROLLBACK_INTENT")
        self.assertEqual(tail["file_index"], 0)
        self.assertEqual(tail["observed_before_sha256"], plan.files[0].after_sha256)
        self.assertEqual(tail["observed_after_sha256"], plan.files[0].before_sha256)
        self.assertEqual(plan.files[0].path.read_bytes(), plan.files[0].after)

        try:
            recovered = api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
        except Exception as exc:
            self.fail(f"matching existing-target exact-after intent was not replayable: {type(exc).__name__}: {exc}")
        self.assertEqual(recovered.status, "rolled-back")
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows pre-delete rollback intent recovery test")
    def test_nonactive_created_target_exact_after_intent_replays_before_same_handle_delete(self) -> None:
        """A newly created exact-after target may replay its matching intent before same-handle delete."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)
        real_append_journal = api._append_journal

        def crash_after_final_install_write(*args, **kwargs):
            event = real_append_journal(*args, **kwargs)
            if event["action"] == "WRITE_DONE" and event["file_index"] == 2:
                raise _SimulatedProcessCrash("install crash with created exact-after target")
            return event

        with patch.object(api, "_append_journal", side_effect=crash_after_final_install_write):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "created exact-after"):
                api.install_native_p2_test_candidate(plan, platform)
        transaction_id, = self._transaction_ids(platform)
        self.assertEqual(plan.files[2].path.read_bytes(), plan.files[2].after)

        def append_delete_intent_then_crash(*args, **kwargs):
            event = real_append_journal(*args, **kwargs)
            if event["action"] == "ROLLBACK_INTENT" and event["file_index"] == 2:
                raise _SimulatedProcessCrash("crash before same-handle delete")
            return event

        with patch.object(api, "_append_journal", side_effect=append_delete_intent_then_crash):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "same-handle delete"):
                api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
        transaction = platform / "backups" / "mingge-native-p2-test" / transaction_id
        tail = json.loads(sorted((transaction / "journal").glob("*.json"))[-1].read_text(encoding="utf-8"))
        self.assertEqual(tail["action"], "ROLLBACK_INTENT")
        self.assertEqual(tail["file_index"], 2)
        self.assertEqual(tail["observed_before_sha256"], plan.files[2].after_sha256)
        self.assertIsNone(tail["observed_after_sha256"])
        self.assertEqual(plan.files[2].path.read_bytes(), plan.files[2].after)

        try:
            recovered = api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
        except Exception as exc:
            self.fail(f"matching created-target exact-after intent was not replayable: {type(exc).__name__}: {exc}")
        self.assertEqual(recovered.status, "rolled-back")
        for path, payload in before.items():
            self.assertEqual(path.read_bytes() if path.exists() else None, payload)

    @unittest.skipUnless(os.name == "nt", "Windows protected exact-after intent rejection test")
    def test_nonactive_exact_after_intent_rejects_wrong_hash_index_drift_and_old_generation(self) -> None:
        """Exact-after replay remains bound to both hashes, deterministic index, current generation and bytes."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        cases = ("wrong-observed-before", "wrong-intended-before", "other-index", "target-drift", "old-generation")
        for case in cases:
            with self.subTest(case=case):
                server = self._server(f"MirServer-intent-guard-{case}")
                platform = self.root / f"platform-intent-guard-{case}"
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
                real_append_journal = api._append_journal

                def crash_after_final_install_write(*args, **kwargs):
                    event = real_append_journal(*args, **kwargs)
                    if event["action"] == "WRITE_DONE" and event["file_index"] == 2:
                        raise _SimulatedProcessCrash("install crash for exact-after guard")
                    return event

                with patch.object(api, "_append_journal", side_effect=crash_after_final_install_write):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "exact-after guard"):
                        api.install_native_p2_test_candidate(plan, platform)
                transaction_id, = self._transaction_ids(platform)

                def append_guard_intent_then_crash(*args, **kwargs):
                    values = list(args)
                    if len(values) > 5 and values[4] == "ROLLBACK_INTENT" and values[5] == 2:
                        if case == "wrong-observed-before":
                            values[6] = "0" * 64
                        elif case == "wrong-intended-before":
                            values[7] = "0" * 64
                        elif case == "other-index":
                            values[5] = 1
                            values[6] = plan.files[1].after_sha256
                            values[7] = plan.files[1].before_sha256
                        event = real_append_journal(*values, **kwargs)
                        raise _SimulatedProcessCrash("crash after protected guard intent")
                    return real_append_journal(*args, **kwargs)

                with patch.object(api, "_append_journal", side_effect=append_guard_intent_then_crash):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "protected guard intent"):
                        api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
                if case == "target-drift":
                    plan.files[2].path.write_bytes(b"NOT-EXACT-AFTER")
                elif case == "old-generation":
                    transaction = platform / "backups" / "mingge-native-p2-test" / transaction_id
                    _, manifest, key, _ = api._load_sealed_manifest(platform, transaction_id, server)
                    coordination = api._require_coordination(server, transaction_id, str(manifest["generation"]), key)
                    coordination["generation"] = uuid.uuid4().hex
                    api._write_coordination(coordination, key)

                drift = self._business_bytes(plan)
                with self.assertRaises(api.NativeP2TestInstallError):
                    api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
                self.assertEqual(self._business_bytes(plan), drift)

    @unittest.skipUnless(os.name == "nt", "Windows generic rollback-start reentry test")
    def test_nonactive_rollback_start_survives_third_recovery_and_rejects_drift(self) -> None:
        """A current no-index START may replay only when every held target is exact before or after."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        for case in ("valid", "target-drift"):
            with self.subTest(case=case):
                server = self._server(f"MirServer-start-{case}")
                platform = self.root / f"platform-start-{case}"
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
                before = self._business_bytes(plan)

                def crash_install(name: str) -> None:
                    if name == "install_after_target_handle_truncate:1":
                        raise _SimulatedProcessCrash("install prefix for three-stage recovery")

                with patch.object(api, "_transaction_fault", side_effect=crash_install):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "three-stage recovery"):
                        api.install_native_p2_test_candidate(plan, platform)
                transaction_id, = self._transaction_ids(platform)
                real_append_journal = api._append_journal

                def append_done_then_crash(*args, **kwargs):
                    event = real_append_journal(*args, **kwargs)
                    if event["action"] == "ROLLBACK_DONE" and event["file_index"] == 1:
                        raise _SimulatedProcessCrash("first recovery crash after done")
                    return event

                with patch.object(api, "_append_journal", side_effect=append_done_then_crash):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "after done"):
                        api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")

                def append_start_then_crash(*args, **kwargs):
                    event = real_append_journal(*args, **kwargs)
                    if event["action"] == "ROLLBACK_START":
                        raise _SimulatedProcessCrash("second recovery crash after start")
                    return event

                with patch.object(api, "_append_journal", side_effect=append_start_then_crash):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "after start"):
                        api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
                transaction = platform / "backups" / "mingge-native-p2-test" / transaction_id
                tail = json.loads(sorted((transaction / "journal").glob("*.json"))[-1].read_text(encoding="utf-8"))
                self.assertEqual(tail["action"], "ROLLBACK_START")
                self.assertIsNone(tail["file_index"])
                self.assertIsNone(tail["observed_before_sha256"])
                self.assertIsNone(tail["observed_after_sha256"])
                if case == "target-drift":
                    plan.files[1].path.write_bytes(b"UNKNOWN-AFTER-GENERIC-START")

                if case == "valid":
                    try:
                        recovered = api.recover_native_p2_test_transaction(
                            platform, transaction_id, server, "rollback",
                        )
                    except Exception as exc:
                        self.fail(f"third recovery rejected current protected rollback start: {type(exc).__name__}: {exc}")
                    self.assertEqual(recovered.status, "rolled-back")
                    for path, payload in before.items():
                        self.assertEqual(path.read_bytes() if path.exists() else None, payload)
                else:
                    drift = self._business_bytes(plan)
                    with self.assertRaises(api.NativeP2TestInstallError):
                        api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
                    self.assertEqual(self._business_bytes(plan), drift)

    @unittest.skipUnless(os.name == "nt", "Windows pending-closed terminal recovery test")
    def test_pending_closed_recovery_only_finishes_terminal_state_and_rejects_invalid_state(self) -> None:
        """Pending-closed may add only ROLLBACK_COMPLETE and close the exact current owner."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        cases = ("valid-active", "valid-missing", "valid-closed", "target-drift", "other-owner", "old-generation")
        for case in cases:
            with self.subTest(case=case):
                server = self._server(f"MirServer-pending-{case}")
                platform = self.root / f"platform-pending-{case}"
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
                before = self._business_bytes(plan)
                receipt = api.install_native_p2_test_candidate(plan, platform)
                real_set_owner_state = api._set_coordination_owner_state

                def set_pending_then_crash(*args, **kwargs):
                    result = real_set_owner_state(*args, **kwargs)
                    owner_state = args[4] if len(args) > 4 else kwargs.get("owner_state")
                    if owner_state == "pending_closed":
                        raise _SimulatedProcessCrash("crash after protected pending-closed")
                    return result

                with patch.object(api, "_set_coordination_owner_state", side_effect=set_pending_then_crash):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "pending-closed"):
                        api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
                for path, payload in before.items():
                    self.assertEqual(path.read_bytes() if path.exists() else None, payload)

                base = platform / "backups" / "mingge-native-p2-test"
                transaction, manifest, key, events = api._load_sealed_manifest(
                    platform, receipt.transaction_id, server,
                )
                generation = str(manifest["generation"])
                coordination = api._require_coordination(server, receipt.transaction_id, generation, key)
                self.assertEqual(coordination["owner_state"], "pending_closed")
                self.assertEqual(coordination["phase"], "ROLLING_BACK")
                event_count = len(events)
                owner_path = api._active_owner_path(base, server)

                if case == "valid-missing":
                    owner_path.unlink()
                elif case == "valid-closed":
                    api._write_active_owner(base, server, receipt.transaction_id, generation, key, "closed")
                elif case == "target-drift":
                    plan.files[0].path.write_bytes(b"UNKNOWN-AFTER-PENDING-CLOSED")
                elif case == "other-owner":
                    api._write_active_owner(base, server, "other-transaction", uuid.uuid4().hex, key, "closed")
                elif case == "old-generation":
                    coordination["generation"] = uuid.uuid4().hex
                    api._write_coordination(coordination, key)

                if case.startswith("valid-"):
                    try:
                        recovered = api.recover_native_p2_test_transaction(
                            platform, receipt.transaction_id, server, "rollback",
                        )
                    except Exception as exc:
                        self.fail(f"pending-closed terminal recovery failed: {type(exc).__name__}: {exc}")
                    self.assertEqual(recovered.status, "rolled-back")
                    _, _, _, completed_events = api._load_sealed_manifest(
                        platform, receipt.transaction_id, server,
                    )
                    self.assertEqual(len(completed_events), event_count + 1)
                    self.assertEqual(completed_events[-1]["action"], "ROLLBACK_COMPLETE")
                    owner = api._read_active_owner(base, server, key)
                    self.assertEqual(owner["transaction_id"], receipt.transaction_id)
                    self.assertEqual(owner["generation"], generation)
                    self.assertEqual(owner["state"], "closed")
                    closed = api._require_coordination(server, receipt.transaction_id, generation, key)
                    self.assertEqual(closed["owner_state"], "closed")
                else:
                    drift = self._business_bytes(plan)
                    with self.assertRaises(api.NativeP2TestInstallError):
                        api.recover_native_p2_test_transaction(
                            platform, receipt.transaction_id, server, "rollback",
                        )
                    self.assertEqual(self._business_bytes(plan), drift)

    @unittest.skipUnless(os.name == "nt", "Windows predecessor-owner generation handoff test")
    def test_bound_predecessor_closed_owner_recovers_both_terminal_boundaries_and_allows_g3(self) -> None:
        """A naturally retained G1/closed owner must finish only the protected G2 terminal handoff."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        for boundary in ("pending-closed", "rolled-back"):
            with self.subTest(boundary=boundary):
                server = self._server(f"MirServer-predecessor-{boundary}")
                platform = self.root / f"platform-predecessor-{boundary}"
                workbook = self._workbook()
                g1_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
                g1 = api.install_native_p2_test_candidate(g1_plan, platform)
                api.rollback_native_p2_test_install(platform, g1.transaction_id, server)
                base = platform / "backups" / "mingge-native-p2-test"
                _, g1_manifest, key, _ = api._load_sealed_manifest(
                    platform, g1.transaction_id, server,
                )
                g1_generation = str(g1_manifest["generation"])
                retained = api._read_active_owner(base, server, key)
                self.assertEqual(
                    (retained["transaction_id"], retained["generation"], retained["state"]),
                    (g1.transaction_id, g1_generation, "closed"),
                )

                g2_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)

                def crash_g2_install(name: str) -> None:
                    if name == "install_after_target_handle_truncate:0":
                        raise _SimulatedProcessCrash("G2 interrupted after truncate")

                with patch.object(api, "_transaction_fault", side_effect=crash_g2_install):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "G2 interrupted"):
                        api.install_native_p2_test_candidate(g2_plan, platform)
                g2_id = next(item for item in self._transaction_ids(platform) if item != g1.transaction_id)
                _, g2_manifest, _, _ = api._load_sealed_manifest(platform, g2_id, server)
                g2_generation = str(g2_manifest["generation"])
                coordination = api._require_coordination(server, g2_id, g2_generation, key)
                self.assertEqual(coordination["schema_version"], 3)
                self.assertEqual(coordination["predecessor_owner"], {
                    "transaction_id": g1.transaction_id,
                    "generation": g1_generation,
                    "state": "closed",
                })

                if boundary == "pending-closed":
                    real_set_owner_state = api._set_coordination_owner_state

                    def crash_after_pending(*args, **kwargs):
                        result = real_set_owner_state(*args, **kwargs)
                        state = args[4] if len(args) > 4 else kwargs.get("owner_state")
                        if state == "pending_closed":
                            raise _SimulatedProcessCrash("G2 protected pending-closed")
                        return result

                    with patch.object(api, "_set_coordination_owner_state", side_effect=crash_after_pending):
                        with self.assertRaisesRegex(_SimulatedProcessCrash, "pending-closed"):
                            api.recover_native_p2_test_transaction(platform, g2_id, server, "rollback")
                else:
                    real_append = api._append_journal

                    def crash_after_terminal(*args, **kwargs):
                        event = real_append(*args, **kwargs)
                        if event["action"] == "ROLLBACK_COMPLETE":
                            raise _SimulatedProcessCrash("G2 durable rolled-back")
                        return event

                    with patch.object(api, "_append_journal", side_effect=crash_after_terminal):
                        with self.assertRaisesRegex(_SimulatedProcessCrash, "durable rolled-back"):
                            api.recover_native_p2_test_transaction(platform, g2_id, server, "rollback")

                retained_after_crash = api._read_active_owner(base, server, key)
                self.assertEqual(
                    (retained_after_crash["transaction_id"], retained_after_crash["generation"], retained_after_crash["state"]),
                    (g1.transaction_id, g1_generation, "closed"),
                )
                recovered = api.recover_native_p2_test_transaction(platform, g2_id, server, "rollback")
                self.assertEqual(recovered.status, "rolled-back")
                current = api._read_active_owner(base, server, key)
                self.assertEqual(
                    (current["transaction_id"], current["generation"], current["state"]),
                    (g2_id, g2_generation, "closed"),
                )
                g3_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
                g3 = api.install_native_p2_test_candidate(g3_plan, platform)
                self.assertEqual(g3.status, "installed-test-candidate")

    @unittest.skipUnless(os.name == "nt", "Windows predecessor-owner fail-closed test")
    def test_pending_closed_rejects_unbound_or_replaced_predecessor_owner_without_writes(self) -> None:
        """Only the exact predecessor sealed before G2 reservation may cross the terminal boundary."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        cases = ("unbound", "wrong-bound-generation", "replacement-closed", "replacement-active")
        for case in cases:
            with self.subTest(case=case):
                server = self._server(f"MirServer-predecessor-guard-{case}")
                platform = self.root / f"platform-predecessor-guard-{case}"
                workbook = self._workbook()
                g1_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
                g1 = api.install_native_p2_test_candidate(g1_plan, platform)
                api.rollback_native_p2_test_install(platform, g1.transaction_id, server)
                g2_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)

                def crash_g2_install(name: str) -> None:
                    if name == "install_after_target_handle_truncate:0":
                        raise _SimulatedProcessCrash("G2 predecessor guard setup")

                with patch.object(api, "_transaction_fault", side_effect=crash_g2_install):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "guard setup"):
                        api.install_native_p2_test_candidate(g2_plan, platform)
                g2_id = next(item for item in self._transaction_ids(platform) if item != g1.transaction_id)
                real_set_owner_state = api._set_coordination_owner_state

                def crash_after_pending(*args, **kwargs):
                    result = real_set_owner_state(*args, **kwargs)
                    state = args[4] if len(args) > 4 else kwargs.get("owner_state")
                    if state == "pending_closed":
                        raise _SimulatedProcessCrash("protected pending-closed guard")
                    return result

                with patch.object(api, "_set_coordination_owner_state", side_effect=crash_after_pending):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "pending-closed guard"):
                        api.recover_native_p2_test_transaction(platform, g2_id, server, "rollback")

                base = platform / "backups" / "mingge-native-p2-test"
                transaction, manifest, key, events = api._load_sealed_manifest(platform, g2_id, server)
                generation = str(manifest["generation"])
                coordination = api._require_coordination(server, g2_id, generation, key)
                if case == "unbound":
                    coordination["predecessor_owner"] = None
                    api._write_coordination(coordination, key)
                elif case == "wrong-bound-generation":
                    predecessor = dict(coordination.get("predecessor_owner") or {})
                    predecessor["generation"] = uuid.uuid4().hex
                    coordination["predecessor_owner"] = predecessor
                    api._write_coordination(coordination, key)
                elif case == "replacement-closed":
                    api._write_active_owner(base, server, "replacement", uuid.uuid4().hex, key, "closed")
                else:
                    api._write_active_owner(base, server, "replacement", uuid.uuid4().hex, key, "active")
                before_attempt = self._business_bytes(g2_plan)
                event_count = len(events)
                with self.assertRaises(api.NativeP2TestInstallError):
                    api.recover_native_p2_test_transaction(platform, g2_id, server, "rollback")
                self.assertEqual(self._business_bytes(g2_plan), before_attempt)
                _, _, _, after_events = api._load_sealed_manifest(platform, g2_id, server)
                self.assertEqual(len(after_events), event_count)

    @unittest.skipUnless(os.name == "nt", "Windows automatic prefix terminal recovery test")
    def test_auto_recovery_prefix_source_pair_survives_pending_closed_crash(self) -> None:
        """Auto-recovery must seal the real prefix source hash+length before terminal-only recovery."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server("MirServer-auto-prefix-terminal")
        platform = self.root / "platform-auto-prefix-terminal"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)
        real_set_owner_state = api._set_coordination_owner_state

        def faults(name: str) -> None:
            if name == "install_after_target_handle_truncate:0":
                raise OSError("capturable install truncate failure")

        def crash_after_pending(*args, **kwargs):
            result = real_set_owner_state(*args, **kwargs)
            state = args[4] if len(args) > 4 else kwargs.get("owner_state")
            if state == "pending_closed":
                raise _SimulatedProcessCrash("auto recovery pending-closed")
            return result

        with (
            patch.object(api, "_transaction_fault", side_effect=faults),
            patch.object(api, "_set_coordination_owner_state", side_effect=crash_after_pending),
        ):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "auto recovery pending-closed"):
                api.install_native_p2_test_candidate(plan, platform)
        transaction_id, = self._transaction_ids(platform)
        _, manifest, key, events = api._load_sealed_manifest(platform, transaction_id, server)
        generation = str(manifest["generation"])
        coordination = api._require_coordination(server, transaction_id, generation, key)
        self.assertEqual(coordination["owner_state"], "pending_closed")
        intent, done = events[-2:]
        self.assertEqual((intent["action"], done["action"]), ("ROLLBACK_INTENT", "ROLLBACK_DONE"))
        empty_hash = api._sha256_bytes(b"")
        for event in (intent, done):
            self.assertEqual(event.get("generation"), generation)
            self.assertEqual(event["file_index"], 0)
            self.assertEqual(event.get("relative_path"), api.TARGET_RELATIVES[0])
            self.assertEqual(event["observed_before_sha256"], empty_hash)
            self.assertEqual(event.get("observed_before_length"), 0)
            self.assertEqual(event["observed_after_sha256"], plan.files[0].before_sha256)
        recovered = api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
        self.assertEqual(recovered.status, "rolled-back")
        self.assertEqual(self._business_bytes(plan), before)

    @unittest.skipUnless(os.name == "nt", "Windows terminal-only auto recovery boundary test")
    def test_auto_recovery_zero_written_and_recovery_required_are_terminal_only_or_zero_cover(self) -> None:
        """All-before may only close; exact-after after either boundary must never be overwritten."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        for boundary, drift in (("zero-written", False), ("zero-written", True), ("recovery-required", False), ("recovery-required", True)):
            with self.subTest(boundary=boundary, drift=drift):
                server = self._server(f"MirServer-auto-boundary-{boundary}-{drift}")
                platform = self.root / f"platform-auto-boundary-{boundary}-{drift}"
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
                before = self._business_bytes(plan)
                real_append = api._append_journal
                real_set_owner_state = api._set_coordination_owner_state
                failed_done = {"value": False}

                def faults(name: str) -> None:
                    if boundary == "zero-written" and name == "after_target_leaves_validated_for_install":
                        raise OSError("capturable zero-written failure")
                    if boundary == "recovery-required" and name == "install_after_target_handle_truncate:0":
                        raise OSError("capturable install prefix failure")

                def fail_first_done(*args, **kwargs):
                    action = args[4] if len(args) > 4 else kwargs.get("action")
                    if boundary == "recovery-required" and action == "ROLLBACK_DONE" and not failed_done["value"]:
                        failed_done["value"] = True
                        raise OSError("fault before auto recovery DONE anchor")
                    return real_append(*args, **kwargs)

                def crash_after_pending(*args, **kwargs):
                    result = real_set_owner_state(*args, **kwargs)
                    state = args[4] if len(args) > 4 else kwargs.get("owner_state")
                    if boundary == "zero-written" and state == "pending_closed":
                        raise _SimulatedProcessCrash("zero-written pending-closed")
                    return result

                if boundary == "zero-written":
                    with (
                        patch.object(api, "_transaction_fault", side_effect=faults),
                        patch.object(api, "_set_coordination_owner_state", side_effect=crash_after_pending),
                    ):
                        with self.assertRaisesRegex(_SimulatedProcessCrash, "zero-written pending-closed"):
                            api.install_native_p2_test_candidate(plan, platform)
                else:
                    with (
                        patch.object(api, "_transaction_fault", side_effect=faults),
                        patch.object(api, "_append_journal", side_effect=fail_first_done),
                    ):
                        with self.assertRaises(api.NativeP2TestInstallError):
                            api.install_native_p2_test_candidate(plan, platform)
                transaction_id, = self._transaction_ids(platform)
                transaction, manifest, key, events = api._load_sealed_manifest(platform, transaction_id, server)
                generation = str(manifest["generation"])
                coordination = api._require_coordination(server, transaction_id, generation, key)
                if boundary == "zero-written":
                    self.assertEqual(coordination["owner_state"], "pending_closed")
                    self.assertEqual(events[-1]["action"], "TRANSACTION_PREPARED")
                else:
                    self.assertEqual(coordination["owner_state"], "none")
                    self.assertEqual(events[-1]["state"], "RECOVERY_REQUIRED")
                    self.assertEqual(events[-1]["action"], "AUTO_RECOVERY_FAILED")
                if drift:
                    plan.files[0].path.write_bytes(plan.files[0].after)
                    frozen = self._business_bytes(plan)
                    with self.assertRaises(api.NativeP2TestInstallError):
                        api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
                    self.assertEqual(self._business_bytes(plan), frozen)
                else:
                    event_count = len(events)
                    recovered = api.recover_native_p2_test_transaction(
                        platform, transaction_id, server, "rollback",
                    )
                    self.assertEqual(recovered.status, "rolled-back")
                    self.assertEqual(self._business_bytes(plan), before)
                    _, _, _, completed = api._load_sealed_manifest(platform, transaction_id, server)
                    self.assertEqual(len(completed), event_count + 1)
                    self.assertEqual(completed[-1]["action"], "ROLLBACK_COMPLETE")

    @unittest.skipUnless(os.name == "nt", "Windows automatic recovery source contract test")
    def test_auto_recovery_terminal_pair_rejects_wrong_source_contract_without_writes(self) -> None:
        """HMAC-valid wrong source/hash/length/index/relative/generation/pairing must fail closed."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        cases = ("source-hash", "source-length", "index", "relative", "generation", "non-pair")
        for case in cases:
            with self.subTest(case=case):
                server = self._server(f"MirServer-auto-source-{case}")
                platform = self.root / f"platform-auto-source-{case}"
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
                real_set_owner_state = api._set_coordination_owner_state

                def faults(name: str) -> None:
                    if name == "install_after_target_handle_truncate:0":
                        raise OSError("capturable source-contract prefix")

                def crash_after_pending(*args, **kwargs):
                    result = real_set_owner_state(*args, **kwargs)
                    state = args[4] if len(args) > 4 else kwargs.get("owner_state")
                    if state == "pending_closed":
                        raise _SimulatedProcessCrash("source-contract pending-closed")
                    return result

                with (
                    patch.object(api, "_transaction_fault", side_effect=faults),
                    patch.object(api, "_set_coordination_owner_state", side_effect=crash_after_pending),
                ):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "source-contract pending-closed"):
                        api.install_native_p2_test_candidate(plan, platform)
                transaction_id, = self._transaction_ids(platform)

                def mutate(events):
                    intent, done = events[-2:]
                    self.assertEqual((intent["action"], done["action"]), ("ROLLBACK_INTENT", "ROLLBACK_DONE"))
                    if case == "source-hash":
                        intent["observed_before_sha256"] = "f" * 64
                        done["observed_before_sha256"] = "f" * 64
                    elif case == "source-length":
                        intent["observed_before_length"] = 1
                        done["observed_before_length"] = 1
                    elif case == "index":
                        done["file_index"] = 1
                    elif case == "relative":
                        done["relative_path"] = api.TARGET_RELATIVES[1]
                    elif case == "generation":
                        done["generation"] = uuid.uuid4().hex
                    else:
                        done["observed_before_sha256"] = plan.files[0].after_sha256
                        done["observed_before_length"] = len(plan.files[0].after)

                self._reseal_journal(api, platform, transaction_id, server, mutate)
                frozen = self._business_bytes(plan)
                with self.assertRaises(api.NativeP2TestInstallError):
                    api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
                self.assertEqual(self._business_bytes(plan), frozen)

    @unittest.skipUnless(os.name == "nt", "Windows active rollback source contract test")
    def test_active_rollback_rejects_bad_source_length_or_pair_before_replay(self) -> None:
        """An active owner cannot bypass the same sealed rollback source contract."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        for case in ("intent-length", "done-pair"):
            with self.subTest(case=case):
                server = self._server(f"MirServer-active-source-{case}")
                platform = self.root / f"platform-active-source-{case}"
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
                receipt = api.install_native_p2_test_candidate(plan, platform)
                real_append = api._append_journal

                def corrupt_source_then_crash(*args, **kwargs):
                    values = list(args)
                    action = values[4] if len(values) > 4 else kwargs.get("action")
                    if action == "ROLLBACK_INTENT" and case == "intent-length":
                        values[8] = 1
                        event = real_append(*values, **kwargs)
                        raise _SimulatedProcessCrash("active bad intent source length")
                    if action == "ROLLBACK_DONE" and case == "done-pair":
                        values[6] = api._sha256_bytes(b"")
                        values[8] = 0
                        event = real_append(*values, **kwargs)
                        raise _SimulatedProcessCrash("active non-paired done source")
                    return real_append(*args, **kwargs)

                with patch.object(api, "_append_journal", side_effect=corrupt_source_then_crash):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "active .* source"):
                        api.rollback_native_p2_test_install(platform, receipt.transaction_id, server)
                frozen = self._business_bytes(plan)
                with self.assertRaises(api.NativeP2TestInstallError):
                    api.recover_native_p2_test_transaction(
                        platform, receipt.transaction_id, server, "rollback",
                    )
                self.assertEqual(self._business_bytes(plan), frozen)

    @unittest.skipUnless(os.name == "nt", "Windows active RECOVERY_REQUIRED terminal boundary test")
    def test_active_recovery_required_all_before_reenters_three_terminal_boundaries(self) -> None:
        """Active UNKNOWN_TARGET_DRIFT may only terminal-close after all targets are exact before."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server("MirServer-recovery-required-active-terminal")
        platform = self.root / "platform-recovery-required-active-terminal"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        before = self._business_bytes(plan)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        plan.files[0].path.write_bytes(b"ACTIVE-UNKNOWN-TARGET")
        with self.assertRaises(api.NativeP2TestInstallError):
            api.recover_native_p2_test_transaction(
                platform, receipt.transaction_id, server, "rollback",
            )
        self._restore_business_bytes(before)
        transaction, manifest, key, events = api._load_sealed_manifest(
            platform, receipt.transaction_id, server,
        )
        generation = str(manifest["generation"])
        self.assertEqual((events[-1]["state"], events[-1]["action"]), (
            "RECOVERY_REQUIRED", "UNKNOWN_TARGET_DRIFT",
        ))
        original_event_count = len(events)
        real_append = api._append_journal

        def crash_after_terminal(*args, **kwargs):
            event = real_append(*args, **kwargs)
            if event["action"] == "ROLLBACK_COMPLETE":
                raise _SimulatedProcessCrash("crash after recovery-required terminal journal")
            return event

        with patch.object(api, "_append_journal", side_effect=crash_after_terminal):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "terminal journal"):
                api.recover_native_p2_test_transaction(
                    platform, receipt.transaction_id, server, "rollback",
                )
        _, _, _, after_terminal = api._load_sealed_manifest(
            platform, receipt.transaction_id, server,
        )
        self.assertEqual(len(after_terminal), original_event_count + 1)
        self.assertEqual(after_terminal[-1]["action"], "ROLLBACK_COMPLETE")
        coordination = api._require_coordination(
            server, receipt.transaction_id, generation, key,
        )
        self.assertEqual(coordination["owner_state"], "pending_closed")

        real_set_owner_state = api._set_coordination_owner_state

        def crash_after_ordinary_owner(*args, **kwargs):
            state = args[4] if len(args) > 4 else kwargs.get("owner_state")
            if state == "closed":
                raise _SimulatedProcessCrash("crash after current ordinary owner")
            return real_set_owner_state(*args, **kwargs)

        with patch.object(api, "_set_coordination_owner_state", side_effect=crash_after_ordinary_owner):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "ordinary owner"):
                api.recover_native_p2_test_transaction(
                    platform, receipt.transaction_id, server, "rollback",
                )
        base = platform / "backups" / "mingge-native-p2-test"
        owner = api._read_active_owner(base, server, key)
        self.assertEqual(
            (owner["transaction_id"], owner["generation"], owner["state"]),
            (receipt.transaction_id, generation, "closed"),
        )
        coordination = api._require_coordination(
            server, receipt.transaction_id, generation, key,
        )
        self.assertEqual(coordination["owner_state"], "pending_closed")

        def crash_after_coordination_closed(*args, **kwargs):
            result = real_set_owner_state(*args, **kwargs)
            state = args[4] if len(args) > 4 else kwargs.get("owner_state")
            if state == "closed":
                raise _SimulatedProcessCrash("crash after protected coordination closed")
            return result

        with patch.object(api, "_set_coordination_owner_state", side_effect=crash_after_coordination_closed):
            with self.assertRaisesRegex(_SimulatedProcessCrash, "coordination closed"):
                api.recover_native_p2_test_transaction(
                    platform, receipt.transaction_id, server, "rollback",
                )
        coordination = api._require_coordination(
            server, receipt.transaction_id, generation, key,
        )
        self.assertEqual(coordination["owner_state"], "closed")
        final = api.recover_native_p2_test_transaction(
            platform, receipt.transaction_id, server, "rollback",
        )
        self.assertEqual(final.status, "rolled-back")
        self.assertEqual(self._business_bytes(plan), before)

    @unittest.skipUnless(os.name == "nt", "Windows active RECOVERY_REQUIRED zero-cover test")
    def test_active_recovery_required_exact_after_rollback_is_zero_cover(self) -> None:
        """A current active owner cannot turn a generic recovery-required tail into write authority."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server("MirServer-recovery-required-active-after")
        platform = self.root / "platform-recovery-required-active-after"
        plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
        receipt = api.install_native_p2_test_candidate(plan, platform)
        plan.files[0].path.write_bytes(b"ACTIVE-UNKNOWN-BEFORE-AFTER")
        with self.assertRaises(api.NativeP2TestInstallError):
            api.recover_native_p2_test_transaction(
                platform, receipt.transaction_id, server, "rollback",
            )
        for item in plan.files:
            item.path.parent.mkdir(parents=True, exist_ok=True)
            item.path.write_bytes(item.after)
        frozen = self._business_bytes(plan)
        with self.assertRaises(api.NativeP2TestInstallError):
            api.recover_native_p2_test_transaction(
                platform, receipt.transaction_id, server, "rollback",
            )
        self.assertEqual(self._business_bytes(plan), frozen)

    @unittest.skipUnless(os.name == "nt", "Windows nonactive RECOVERY_REQUIRED state-table test")
    def test_nonactive_recovery_required_actions_terminal_close_and_reject_other_states(self) -> None:
        """Every generated action shares one all-before terminal-only and fail-closed state table."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        positive = (
            ("UNKNOWN_TARGET_DRIFT", "missing"),
            ("AUTO_RECOVERY_FAILED", "missing"),
            ("ROLLBACK_DRIFT", "missing"),
            ("ROLLBACK_VERIFY_FAILED", "missing"),
            ("UNKNOWN_TARGET_DRIFT", "current-closed"),
            ("UNKNOWN_TARGET_DRIFT", "predecessor"),
        )
        for action, owner_source in positive:
            with self.subTest(kind="positive", action=action, owner_source=owner_source):
                server = self._server(f"MirServer-recovery-action-{action}-{owner_source}")
                platform = self.root / f"platform-recovery-action-{action}-{owner_source}"
                workbook = self._workbook()
                predecessor_id = None
                if owner_source == "predecessor":
                    g1_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
                    g1 = api.install_native_p2_test_candidate(g1_plan, platform)
                    api.rollback_native_p2_test_install(platform, g1.transaction_id, server)
                    predecessor_id = g1.transaction_id
                plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
                before = self._business_bytes(plan)

                def crash_install(name: str) -> None:
                    if name == "install_after_target_handle_truncate:0":
                        raise _SimulatedProcessCrash("nonactive recovery-required setup")

                with patch.object(api, "_transaction_fault", side_effect=crash_install):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "setup"):
                        api.install_native_p2_test_candidate(plan, platform)
                transaction_id = next(
                    item for item in self._transaction_ids(platform) if item != predecessor_id
                )
                plan.files[0].path.write_bytes(b"NONACTIVE-UNKNOWN-TARGET")
                with self.assertRaises(api.NativeP2TestInstallError):
                    api.recover_native_p2_test_transaction(
                        platform, transaction_id, server, "rollback",
                    )
                self._restore_business_bytes(before)
                transaction, manifest, key, _ = api._load_sealed_manifest(
                    platform, transaction_id, server,
                )
                generation = str(manifest["generation"])
                base = platform / "backups" / "mingge-native-p2-test"
                if owner_source == "current-closed":
                    api._write_active_owner(
                        base, server, transaction_id, generation, key, "closed",
                    )

                def set_action(events):
                    tail = events[-1]
                    tail["action"] = action
                    tail["file_index"] = None
                    tail["relative_path"] = None
                    tail["observed_before_sha256"] = None
                    tail["observed_after_sha256"] = None
                    tail["observed_before_length"] = None
                    if action == "ROLLBACK_DRIFT":
                        tail["file_index"] = 0
                        tail["relative_path"] = api.TARGET_RELATIVES[0]
                        tail["observed_before_sha256"] = api._sha256_bytes(b"NONACTIVE-UNKNOWN-TARGET")
                        tail["observed_before_length"] = len(b"NONACTIVE-UNKNOWN-TARGET")
                    elif action == "ROLLBACK_VERIFY_FAILED":
                        tail["file_index"] = 0
                        tail["relative_path"] = api.TARGET_RELATIVES[0]

                self._reseal_journal(api, platform, transaction_id, server, set_action)
                _, _, _, before_events = api._load_sealed_manifest(
                    platform, transaction_id, server,
                )
                recovered = api.recover_native_p2_test_transaction(
                    platform, transaction_id, server, "rollback",
                )
                self.assertEqual(recovered.status, "rolled-back")
                self.assertEqual(self._business_bytes(plan), before)
                _, _, _, after_events = api._load_sealed_manifest(
                    platform, transaction_id, server,
                )
                self.assertEqual(len(after_events), len(before_events) + 1)
                self.assertEqual(after_events[-1]["action"], "ROLLBACK_COMPLETE")

        negative = (
            "after", "prefix", "unknown", "unknown-action",
            "bad-unknown-shape", "bad-drift-shape", "bad-verify-shape", "bad-auto-shape",
        )
        for case in negative:
            with self.subTest(kind="negative", case=case):
                server = self._server(f"MirServer-recovery-negative-{case}")
                platform = self.root / f"platform-recovery-negative-{case}"
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)

                def crash_install(name: str) -> None:
                    if name == "install_after_target_handle_truncate:0":
                        raise _SimulatedProcessCrash("negative recovery-required setup")

                with patch.object(api, "_transaction_fault", side_effect=crash_install):
                    with self.assertRaisesRegex(_SimulatedProcessCrash, "negative.*setup"):
                        api.install_native_p2_test_candidate(plan, platform)
                transaction_id, = self._transaction_ids(platform)
                plan.files[0].path.write_bytes(b"NONACTIVE-NEGATIVE-UNKNOWN")
                with self.assertRaises(api.NativeP2TestInstallError):
                    api.recover_native_p2_test_transaction(
                        platform, transaction_id, server, "rollback",
                    )
                before = {item.path: None for item in plan.files}
                transaction = platform / "backups" / "mingge-native-p2-test" / transaction_id
                for index, item in enumerate(plan.files):
                    backup = transaction / "files" / f"{index:02d}.bin"
                    before[item.path] = backup.read_bytes() if backup.exists() else None
                self._restore_business_bytes(before)

                def mutate_tail(events):
                    tail = events[-1]
                    if case == "unknown-action":
                        tail["action"] = "NOT_A_PRODUCTION_ACTION"
                    elif case == "bad-unknown-shape":
                        tail["file_index"] = 0
                        tail["relative_path"] = api.TARGET_RELATIVES[0]
                    elif case == "bad-drift-shape":
                        tail["action"] = "ROLLBACK_DRIFT"
                    elif case == "bad-verify-shape":
                        tail["action"] = "ROLLBACK_VERIFY_FAILED"
                        tail["file_index"] = 0
                        tail["relative_path"] = api.TARGET_RELATIVES[0]
                        tail["observed_before_length"] = 0
                    elif case == "bad-auto-shape":
                        tail["action"] = "AUTO_RECOVERY_FAILED"
                        tail["file_index"] = 0
                        tail["relative_path"] = api.TARGET_RELATIVES[0]

                if case not in {"after", "prefix", "unknown"}:
                    self._reseal_journal(api, platform, transaction_id, server, mutate_tail)
                if case == "after":
                    plan.files[0].path.write_bytes(plan.files[0].after)
                elif case == "prefix":
                    plan.files[0].path.write_bytes(plan.files[0].after[:7])
                elif case == "unknown":
                    plan.files[0].path.write_bytes(b"UNKNOWN-AFTER-RECOVERY-REQUIRED")
                frozen = self._business_bytes(plan)
                with self.assertRaises(api.NativeP2TestInstallError):
                    api.recover_native_p2_test_transaction(
                        platform, transaction_id, server, "rollback",
                    )
                self.assertEqual(self._business_bytes(plan), frozen)

    @unittest.skipUnless(os.name == "nt", "Windows pending journal prevalidation test")
    def test_invalid_pending_event_never_mutates_journal_or_coordination_on_reconcile(self) -> None:
        """An HMAC-valid invalid pending event must be rejected before either durable side changes."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        for placement in ("missing", "present"):
            for defect in ("generation", "index-relative", "negative-length", "bool-length"):
                with self.subTest(placement=placement, defect=defect):
                    server = self._server(f"MirServer-pending-contract-{placement}-{defect}")
                    platform = self.root / f"platform-pending-contract-{placement}-{defect}"
                    plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)
                    receipt = api.install_native_p2_test_candidate(plan, platform)
                    transaction, manifest, key, _ = api._load_sealed_manifest(
                        platform, receipt.transaction_id, server,
                    )
                    generation = str(manifest["generation"])
                    coordination = api._require_coordination(
                        server, receipt.transaction_id, generation, key,
                    )
                    event = {
                        "sequence": int(coordination["journal_sequence"]) + 1,
                        "transaction_id": receipt.transaction_id,
                        "generation": generation,
                        "state": "RECOVERY_REQUIRED",
                        "action": "UNKNOWN_TARGET_DRIFT",
                        "file_index": None,
                        "relative_path": None,
                        "observed_before_sha256": None,
                        "observed_after_sha256": None,
                        "observed_before_length": None,
                        "previous_event_hmac": str(coordination["journal_head_hmac"]),
                        "created_at_utc": "2026-08-26T00:00:00+00:00",
                    }
                    if defect == "generation":
                        event["generation"] = uuid.uuid4().hex
                    elif defect == "index-relative":
                        event["file_index"] = 0
                        event["relative_path"] = api.TARGET_RELATIVES[1]
                    elif defect == "negative-length":
                        event["observed_before_length"] = -1
                    else:
                        event["observed_before_length"] = True
                    event["hmac"] = api._seal(event, key, api.JOURNAL_DOMAIN)
                    coordination["pending_event"] = event
                    api._write_coordination(coordination, key)
                    next_path = transaction / "journal" / f"{event['sequence']:06d}.json"
                    if placement == "present":
                        api._atomic_write(next_path, api._canonical_json_bytes(event))
                    protected_before = api._read_coordination(server, key, required=True)
                    journal_before = {
                        path.name: path.read_bytes()
                        for path in sorted((transaction / "journal").glob("*.json"))
                    }
                    with self.assertRaises(api.NativeP2TestInstallError):
                        api._load_sealed_manifest(
                            platform, receipt.transaction_id, server, reconcile_pending=True,
                        )
                    protected_after = api._read_coordination(server, key, required=True)
                    for field in ("journal_sequence", "journal_head_hmac", "phase", "pending_event"):
                        self.assertEqual(protected_after[field], protected_before[field])
                    journal_after = {
                        path.name: path.read_bytes()
                        for path in sorted((transaction / "journal").glob("*.json"))
                    }
                    self.assertEqual(journal_after, journal_before)

    @unittest.skipUnless(os.name == "nt", "Windows same-handle crash recovery test")
    def test_current_write_intent_rejects_non_prefix_and_other_index_drift(self) -> None:
        """A current intent never authorizes arbitrary bytes or a prefix on another file index."""
        api = importlib.import_module("xydp.mingge_native_p2_install")
        for drift_case in ("non-prefix", "other-index"):
            with self.subTest(drift_case=drift_case):
                server = self._server(f"MirServer-{drift_case}")
                platform = self.root / f"platform-{drift_case}"
                plan = api.plan_native_p2_test_install(self._workbook(), "tiebi-vitality-p2", server)

                def crash(name: str) -> None:
                    if name == "install_after_target_handle_truncate:0":
                        raise _SimulatedProcessCrash("install crashed for drift rejection")

                with patch.object(api, "_transaction_fault", side_effect=crash):
                    with self.assertRaises(_SimulatedProcessCrash):
                        api.install_native_p2_test_candidate(plan, platform)
                transaction_id, = self._transaction_ids(platform)
                if drift_case == "non-prefix":
                    drift_path = plan.files[0].path
                    drift_payload = b"NOT-A-PREFIX-OF-INTENDED-PAYLOAD"
                else:
                    drift_path = plan.files[1].path
                    drift_payload = plan.files[1].after[:8]
                drift_path.write_bytes(drift_payload)
                with self.assertRaisesRegex(api.NativeP2TestInstallError, "unknown|未知|前缀|意图|漂移"):
                    api.recover_native_p2_test_transaction(platform, transaction_id, server, "rollback")
                self.assertEqual(drift_path.read_bytes(), drift_payload)

    def test_rolled_back_transaction_cannot_finalize_and_old_transaction_cannot_rollback_new_generation(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        workbook = self._workbook()
        first_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        first = api.install_native_p2_test_candidate(first_plan, platform)
        api.rollback_native_p2_test_install(platform, first.transaction_id, server)
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "已回滚|不允许"):
            api.recover_native_p2_test_transaction(platform, first.transaction_id, server, "finalize-install")
        second_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        second = api.install_native_p2_test_candidate(second_plan, platform)
        installed = {item.path: item.path.read_bytes() for item in second_plan.files}
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "已回滚|安装代|事务恢复预检"):
            api.rollback_native_p2_test_install(platform, first.transaction_id, server)
        for path, payload in installed.items():
            self.assertEqual(path.read_bytes(), payload)

    def test_repeated_install_is_blocked_without_creating_a_second_transaction(self) -> None:
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = self._server()
        platform = self.root / "platform"
        workbook = self._workbook()
        first_plan = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        api.install_native_p2_test_candidate(first_plan, platform)
        base = platform / "backups" / "mingge-native-p2-test"
        before_transactions = {item.name for item in base.iterdir() if item.is_dir() and item.name != "active"}
        repeated = api.plan_native_p2_test_install(workbook, "tiebi-vitality-p2", server)
        self.assertTrue(any("无变化的新事务" in item for item in repeated.blockers))
        with self.assertRaisesRegex(api.NativeP2TestInstallError, "无变化的新事务|活动P2事务"):
            api.install_native_p2_test_candidate(repeated, platform)
        after_transactions = {item.name for item in base.iterdir() if item.is_dir() and item.name != "active"}
        self.assertEqual(after_transactions, before_transactions)

    def test_cli_requires_yes_and_roundtrips_transaction(self) -> None:
        from xydp.cli import main as cli_main

        server = self._server()
        workbook = self._workbook()
        platform = self.root / "platform"
        with self.assertRaises(SystemExit):
            cli_main(["--root", str(platform), "mingge-native-p2-test-install", "--input", str(workbook),
                      "--candidate", "tiebi-vitality-p2", "--server", str(server)])
        stdout = io.StringIO()
        with redirect_stdout(stdout):
            code = cli_main(["--root", str(platform), "mingge-native-p2-test-install", "--input", str(workbook),
                             "--candidate", "tiebi-vitality-p2", "--server", str(server), "--yes"])
        self.assertEqual(code, 0)
        transaction = json.loads(stdout.getvalue())["transaction_id"]
        stdout = io.StringIO()
        with redirect_stdout(stdout):
            code = cli_main(["--root", str(platform), "mingge-native-p2-test-rollback",
                             "--transaction", transaction, "--server", str(server), "--yes"])
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(stdout.getvalue())["status"], "rolled-back")


class NativeP2SecretAclIntegrationTests(unittest.TestCase):
    def _probe_directory(self) -> Path:
        program_data = os.environ.get("PROGRAMDATA")
        if not program_data:
            self.skipTest("PROGRAMDATA is unavailable")
        return Path(program_data) / "XuanYuanDevPlatform" / f"secrets-p2-{uuid.uuid4().hex}"

    def test_restricted_acl_uses_only_expected_sids_when_parent_module_path_is_poisoned(self) -> None:
        """Removing the child PSModulePath isolation must make this integration test fail."""
        if os.name != "nt":
            self.skipTest("Windows ACL integration test")
        api = importlib.import_module("xydp.mingge_native_p2_install")
        probe = self._probe_directory()
        try:
            try:
                probe.mkdir(parents=True)
            except PermissionError as exc:
                self.skipTest(f"cannot create ProgramData ACL probe: {exc}")
            api._restrict_secret_acl(probe)
            with patch.dict(os.environ, {"PSModulePath": r"C:\Program Files\PowerShell\Modules"}, clear=False):
                snapshot = api._acl_snapshot(probe)
                api._verify_secret_acl(probe)
            self.assertIs(snapshot.get("Protected"), True)
            rules = snapshot.get("Rules")
            self.assertIsInstance(rules, list)
            self.assertEqual({rule["Identity"] for rule in rules}, {"S-1-5-18", "S-1-5-32-544"})
        finally:
            if probe.exists():
                probe.rmdir()

    def test_coordination_file_uses_exact_protected_acl_and_is_removed(self) -> None:
        """Removing coordination-file ACL restriction must fail this real Windows integration."""
        if os.name != "nt":
            self.skipTest("Windows ACL integration test")
        api = importlib.import_module("xydp.mingge_native_p2_install")
        server = Path(tempfile.gettempdir()) / f"xydp-p2-coordination-server-{uuid.uuid4().hex}"
        key = os.urandom(32)
        path = api._coordination_path(server)
        payload = {
            "schema_version": 3,
            "server_root": str(server.resolve()),
            "transaction_id": f"acl-probe-{uuid.uuid4().hex}",
            "generation": uuid.uuid4().hex,
            "phase": "MANIFEST_ONLY",
            "setup_state": "RESERVED",
            "business_write_started": False,
            "reservation_files": [
                {
                    "index": index,
                    "relative_path": relative,
                    "before_sha256": None,
                    "after_sha256": str(index) * 64,
                }
                for index, relative in enumerate(api.TARGET_RELATIVES)
            ],
            "journal_sequence": 0,
            "journal_head_hmac": "",
            "pending_event": None,
            "owner_state": "none",
            "predecessor_owner": None,
            "target_relatives": list(api.TARGET_RELATIVES),
        }
        try:
            api._write_coordination(payload, key)
            snapshot = api._acl_snapshot(path)
            api._verify_secret_acl(path)
            self.assertIs(snapshot.get("Protected"), True)
            rules = snapshot.get("Rules")
            self.assertIsInstance(rules, list)
            self.assertEqual({rule["Identity"] for rule in rules}, {"S-1-5-18", "S-1-5-32-544"})
        finally:
            if path.exists():
                path.unlink()
            self.assertFalse(path.exists())


if __name__ == "__main__":
    unittest.main()
