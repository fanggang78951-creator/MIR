from __future__ import annotations

import hashlib
import sqlite3
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from xydp.item_synthesis import (
    ItemSynthesisError,
    ItemSynthesisService,
    SynthesisInput,
    SynthesisRecipe,
    SynthesisSettings,
    SynthesisWorkbook,
    _compile_script,
    read_item_synthesis_workbook,
    read_synthesis_npc_config,
)


FIXTURES = Path(__file__).parent / "fixtures"
SAMPLE = FIXTURES / "item_synthesis_sample.xlsx"
INVALID_CURRENCY = FIXTURES / "item_synthesis_invalid_currency.xlsx"


def _server(tmp_path: Path, *, omit: str | None = None) -> Path:
    root = tmp_path / "MirServer"
    envir = root / "Mir200" / "Envir"
    database = root / "Mud2" / "DB"
    (envir / "Market_Def").mkdir(parents=True)
    database.mkdir(parents=True)
    (root / "Mir200" / "M2Server.exe").write_bytes(b"")
    (envir / "MapInfo.txt").write_bytes(
        "[XY_NMGF_MAIN 宁姆格福 0]\r\n[XY_NMGF_WEST 宁姆格福西域 0]\r\n".encode("gb18030")
    )
    (envir / "MerChant.txt").write_bytes(b"")
    connection = sqlite3.connect(database / "ApexM2.DB")
    try:
        connection.execute("CREATE TABLE StdItems (Idx INTEGER, Name TEXT)")
        rows = [(6, "木剑"), (101, "铁矿"), (102, "木材")]
        connection.executemany(
            "INSERT INTO StdItems (Idx, Name) VALUES (?, ?)",
            [row for row in rows if row[1] != omit],
        )
        connection.commit()
    finally:
        connection.close()
    return root


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _multi_workbook(path: Path) -> SynthesisWorkbook:
    settings = SynthesisSettings("旧单NPC", "XY_NMGF_MAIN", 104, 75, 220, "旧界面", 8)
    inputs = (SynthesisInput("物品", "铁矿", 1, 1),)
    recipes = tuple(
        SynthesisRecipe(recipe_id, "测试", display, "木剑", 1, order, inputs)
        for order, (recipe_id, display) in enumerate(
            (("A1", "合成A1"), ("B2", "合成B2"), ("C1", "合成C1")), start=1
        )
    )
    digest = hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else "0" * 64
    return SynthesisWorkbook(str(path), digest, settings, recipes)


class ItemSynthesisTests(unittest.TestCase):
    def test_bundle_mode_accepts_blank_coordinates_without_changing_default_contract(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "auto-coordinate.xlsx"
            workbook = load_workbook(SAMPLE)
            sheet = workbook["界面设置"]
            sheet["B6"] = None
            sheet["B7"] = None
            workbook.save(path)

            with self.assertRaisesRegex(ItemSynthesisError, "NPC坐标X"):
                read_item_synthesis_workbook(path)
            parsed = read_item_synthesis_workbook(path, allow_auto_coordinates=True)
            self.assertEqual((parsed.settings.x, parsed.settings.y), (0, 0))

    def test_official_workbook_schema_is_readable_with_or_without_enabled_recipes(self) -> None:
        path = Path(r"E:\XuanYuanDevPlatform\所需材料表格汇总\34_通用物品合成.xlsx")
        try:
            workbook = read_item_synthesis_workbook(path)
        except ItemSynthesisError as exc:
            self.assertIn("没有启用配方", str(exc))
        else:
            self.assertTrue(workbook.recipes)

    def test_default_output_one_and_unlimited_cost_rows_compile_in_safe_order(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            server = _server(tmp)
            service = ItemSynthesisService(tmp / "platform")
            plan = service.preflight(SAMPLE, server)
            self.assertEqual(plan.blockers, [])
            self.assertEqual(plan.recipes[0]["output_amount"], 1)
            script_change = next(
                change for change in plan.changes if change.relative_path.endswith("通用合成-XY_NMGF_MAIN.txt")
            )
            script = script_change.after.decode("gb18030")
            main_page = script[script.index("[@XY_SYN_PAGE_1]"):script.index("[@XY_SYN_VIEW_SYN001]")]
            self.assertIn("<ItemShow:6:0:260:-30:1>", main_page)
            self.assertEqual(script.count("<ItemShow:6:0:260:-30:1>"), 2)
            self.assertNotIn("<ItemShow:6:0:260:-90:1>", script)
            self.assertIn("CHECKBAGSIZE 1", script)
            self.assertIn("CHECKITEM 铁矿 2", script)
            self.assertIn("CHECKITEM 木材 3", script)
            self.assertIn("CHECKGOLD 1000", script)
            self.assertIn("GIVE 木剑 1", script)
            checks_end = script.index("#ACT", script.index("[@XY_SYN_APPLY_SYN001]"))
            self.assertLess(script.index("CHECKITEM 铁矿 2"), checks_end)
            self.assertLess(script.index("CHECKITEM 木材 3"), checks_end)
            self.assertLess(script.index("CHECKGOLD 1000"), checks_end)
            self.assertGreater(script.index("TAKE 铁矿 2"), checks_end)
            self.assertGreater(script.index("GOLDCOUNT - 1000"), checks_end)
            self.assertLess(script.index("GOLDCOUNT - 1000"), script.index("GIVE 木剑 1"))

    def test_main_page_output_itemshows_use_real_idx_and_unique_grid_positions(self) -> None:
        workbook = _multi_workbook(Path("multi.xlsx"))
        compiled = _compile_script(workbook, {"木剑": 706})
        main_page = compiled[compiled.index("[@XY_SYN_PAGE_1]"):compiled.index("[@XY_SYN_VIEW_A1]")]
        self.assertIn("<ItemShow:706:0:260:-30:1>", main_page)
        self.assertIn("<ItemShow:706:0:318:-30:1>", main_page)
        self.assertIn("<ItemShow:706:0:376:-30:1>", main_page)
        self.assertNotIn("/@XY_SYN_APPLY_", main_page)

    def test_explicit_output_amount_is_used_for_bag_and_give(self) -> None:
        workbook = read_item_synthesis_workbook(SAMPLE)
        recipe = replace(workbook.recipes[0], output_amount=3)
        compiled = _compile_script(replace(workbook, recipes=(recipe,)), {"木剑": 6})
        self.assertIn("CHECKBAGSIZE 3", compiled)
        self.assertIn("GIVE 木剑 3", compiled)

    def test_invalid_native_currency_is_blocked(self) -> None:
        with self.assertRaisesRegex(ItemSynthesisError, "不支持的原生货币"):
            read_item_synthesis_workbook(INVALID_CURRENCY)

    def test_missing_target_item_blocks_without_changes(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            plan = ItemSynthesisService(tmp / "platform").preflight(
                SAMPLE, _server(tmp, omit="木材")
            )
            self.assertTrue(any("目标服物品不存在：木材" in item for item in plan.blockers))
            self.assertEqual(plan.changes, [])

    def test_foreign_dedicated_script_is_not_overwritten(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            server = _server(tmp)
            script = server / "Mir200" / "Envir" / "Market_Def" / "玄渊功能" / "通用合成-XY_NMGF_MAIN.txt"
            script.parent.mkdir(parents=True)
            script.write_bytes("[@Main]\r\n#SAY\r\n旧脚本".encode("gb18030"))
            plan = ItemSynthesisService(tmp / "platform").preflight(SAMPLE, server)
            self.assertTrue(any("不是平台受管文件" in item for item in plan.blockers))
            self.assertEqual(script.read_bytes(), "[@Main]\r\n#SAY\r\n旧脚本".encode("gb18030"))

    def test_transaction_is_idempotent_and_rolls_back_byte_exactly(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            server = _server(tmp)
            platform = tmp / "platform"
            service = ItemSynthesisService(platform)
            merchant = server / "Mir200" / "Envir" / "MerChant.txt"
            original_merchant = _hash(merchant)
            plan = service.preflight(SAMPLE, server)
            self.assertEqual(plan.blockers, [])
            receipt = service.install(plan)
            script = server / "Mir200" / "Envir" / "Market_Def" / "玄渊功能" / "通用合成-XY_NMGF_MAIN.txt"
            self.assertTrue(script.is_file())
            second = service.preflight(SAMPLE, server)
            self.assertEqual(second.blockers, [])
            self.assertEqual(second.changes, [])
            service.rollback(server, receipt.transaction_id)
            self.assertFalse(script.exists())
            self.assertEqual(_hash(merchant), original_merchant)

    def test_multi_map_multi_npc_generates_isolated_scripts(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            server = _server(tmp)
            workbook_path = tmp / "34.xlsx"
            workbook_path.write_bytes(b"fixture")
            config = tmp / "35.txt"
            config.write_text(
                "[NPC_A]\n状态=是\nNPC名称=地图1合成师A\n地图=XY_NMGF_MAIN\nX=104\nY=75\n外观=220\n界面标题=A合成\n顺序=1\n配方ID=A1\n\n"
                "[NPC_B]\n状态=是\nNPC名称=地图1合成师B\n地图=XY_NMGF_MAIN\nX=106\nY=75\n外观=221\n界面标题=B合成\n顺序=2\n配方ID=B2\n\n"
                "[NPC_C]\n状态=是\nNPC名称=地图2合成师C\n地图=XY_NMGF_WEST\nX=100\nY=75\n外观=222\n界面标题=C合成\n顺序=3\n配方ID=C1\n",
                encoding="utf-8",
            )
            workbook = _multi_workbook(workbook_path)
            service = ItemSynthesisService(tmp / "platform")
            merchant_path = server / "Mir200" / "Envir" / "MerChant.txt"
            original_merchant = _hash(merchant_path)
            with patch("xydp.item_synthesis.read_item_synthesis_workbook", return_value=workbook):
                plan = service.preflight(workbook_path, server, config)
            self.assertEqual(plan.blockers, [])
            self.assertEqual(len(plan.npcs), 3)
            script_changes = {
                change.relative_path: change.after.decode("gb18030")
                for change in plan.changes if change.operation == "item-synthesis-script"
            }
            self.assertEqual(len(script_changes), 3)
            a_script = next(value for key, value in script_changes.items() if "NPC_A" in key)
            b_script = next(value for key, value in script_changes.items() if "NPC_B" in key)
            c_script = next(value for key, value in script_changes.items() if "NPC_C" in key)
            self.assertIn("合成A1", a_script)
            self.assertNotIn("合成B2", a_script)
            self.assertIn("合成B2", b_script)
            self.assertNotIn("合成C1", b_script)
            self.assertIn("合成C1", c_script)
            merchant = next(change for change in plan.changes if change.operation == "item-synthesis-npc-register")
            merchant_text = merchant.after.decode("utf-8")
            self.assertIn("玄渊功能/通用合成/NPC_A\tXY_NMGF_MAIN\t104\t75", merchant_text)
            self.assertIn("玄渊功能/通用合成/NPC_B\tXY_NMGF_MAIN\t106\t75", merchant_text)
            self.assertIn("玄渊功能/通用合成/NPC_C\tXY_NMGF_WEST\t100\t75", merchant_text)
            receipt = service.install(plan)
            for npc_id, map_code in (("NPC_A", "XY_NMGF_MAIN"), ("NPC_B", "XY_NMGF_MAIN"), ("NPC_C", "XY_NMGF_WEST")):
                path = server / "Mir200" / "Envir" / "Market_Def" / "玄渊功能" / "通用合成" / f"{npc_id}-{map_code}.txt"
                self.assertTrue(path.is_file())
            with patch("xydp.item_synthesis.read_item_synthesis_workbook", return_value=workbook):
                second = service.preflight(workbook_path, server, config)
            self.assertEqual(second.blockers, [])
            self.assertEqual(second.changes, [])
            service.rollback(server, receipt.transaction_id)
            self.assertEqual(_hash(merchant_path), original_merchant)
            self.assertFalse((server / "Mir200" / "Envir" / "Market_Def" / "玄渊功能" / "通用合成" / "NPC_A-XY_NMGF_MAIN.txt").exists())

    def test_recipe_cannot_be_assigned_to_two_npcs(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            config = tmp / "35.txt"
            config.write_text(
                "[NPC_A]\n状态=是\nNPC名称=A\n地图=XY_NMGF_MAIN\nX=1\nY=1\n外观=1\n界面标题=A\n顺序=1\n配方ID=A1,B2,C1\n\n"
                "[NPC_B]\n状态=是\nNPC名称=B\n地图=XY_NMGF_MAIN\nX=2\nY=2\n外观=1\n界面标题=B\n顺序=2\n配方ID=A1\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ItemSynthesisError, "同时分配"):
                read_synthesis_npc_config(config, _multi_workbook(tmp / "34.xlsx"))


if __name__ == "__main__":
    unittest.main()
