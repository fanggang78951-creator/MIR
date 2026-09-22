from __future__ import annotations

import hashlib
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from xydp.cli import _parser
from xydp.monster_library import MONSTER_COLUMNS, MonsterLibraryError, MonsterLibraryService
from xydp.monster_workbook import HEADERS, LEGACY_HEADERS, MonsterWorkbookService
from tests.monster_v3_fixtures import (
    write_effect_image_list,
    write_smartmonster_ini,
    write_wzl_wzx_pair,
)


def monster_values(name: str, appr: int, **updates) -> dict[str, object]:
    values: dict[str, object] = {column: 0 for column in MONSTER_COLUMNS}
    values.update({
        "Name": name, "Race": 81, "RaceImg": 19, "Appr": appr, "Lvl": 80,
        "Exp": 50000, "HP": 100000, "AC": 100, "MAC": 100, "DC": 200,
        "DCMAX": 300, "HIT": 200, "WALK_SPD": 300, "WalkStep": 1,
        "WalkWait": 1, "ATTACK_SPD": 800,
    })
    values.update(updates)
    return values


def make_monster_db(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    try:
        definitions = ", ".join(
            f'"{column}" TEXT' if column == "Name" else f'"{column}" INTEGER'
            for column in MONSTER_COLUMNS
        )
        connection.execute(f'CREATE TABLE "Monster" ({definitions})')
        columns = ", ".join(f'"{item}"' for item in MONSTER_COLUMNS)
        placeholders = ", ".join("?" for _ in MONSTER_COLUMNS)
        connection.executemany(
            f'INSERT INTO "Monster" ({columns}) VALUES ({placeholders})',
            [tuple(row[item] for item in MONSTER_COLUMNS) for row in rows],
        )
        connection.commit()
    finally:
        connection.close()


class MonsterWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.platform = self.root / "platform"
        self.server = self.root / "server"
        self.client_data = self.root / "client" / "data"
        self.client_data.mkdir(parents=True)
        (self.client_data.parent / "Game.exe").write_bytes(b"CLIENT")
        (self.server / "登录器").mkdir(parents=True)
        (self.server / "登录器" / "pak.txt").write_bytes(b"")
        donor_db = self.root / "donor" / "Mud2" / "DB" / "ApexM2.DB"
        donor_wzl = self.root / "donor-client" / "data"
        donor_pak = self.root / "donor-client" / "mycm" / "data"
        donor_rules = self.root / "donor-login" / "pak.txt"
        self.donor_db = donor_db
        self.donor_wzl = donor_wzl
        self.donor_pak = donor_pak
        self.donor_rules = donor_rules
        self.donor_effect_list = self.root / "donor" / "Mir200" / "Envir" / "EffectImageList.txt"
        self.donor_smartmonster = self.donor_effect_list.parent / "SmartMonster"
        donor_wzl.mkdir(parents=True)
        donor_pak.mkdir(parents=True)
        make_monster_db(donor_db, [
            monster_values("模型甲", 1230, Race=81),
            monster_values("模型乙", 1800, Race=82),
            monster_values("模型丙", 215, Race=83),
        ])
        make_monster_db(
            self.server / "Mud2" / "DB" / "ApexM2.DB",
            [
                monster_values("本端已有", 10),
                monster_values(
                    "稻草人", 3,
                    Race=7,
                    RaceImg=8,
                    SPEED=311,
                    WALK_SPD=333,
                    ATTACK_SPD=777,
                    CoolEye=61,
                    ExploreItem=901,
                ),
            ],
        )
        rules = []
        for library in (124, 181, 22):
            pak = donor_pak / f"Mon{library}.Pak"
            pak.write_bytes(f"MON-{library}-FULL-ACTIONS".encode("ascii"))
            rules.append(f"{pak}|测试密码")
        donor_rules.parent.mkdir(parents=True)
        donor_rules.write_bytes(("\r\n".join(rules) + "\r\n").encode("gb18030"))
        MonsterLibraryService(self.platform).scan(
            donor_db, donor_wzl, donor_pak, donor_rules, self.server, self.client_data,
        )
        self.service = MonsterWorkbookService(self.platform)
        self.workbook = self.root / "20_怪物批量生成.xlsx"

    def tearDown(self):
        self.temp.cleanup()

    def write_workbook(self, rows: list[list[object]], headers: tuple[str, ...] = HEADERS) -> None:
        book = Workbook()
        sheet = book.active
        sheet.title = "怪物生成"
        sheet.append(list(headers))
        for row in rows:
            sheet.append(row)
        book.save(self.workbook)
        book.close()

    @staticmethod
    def row(
        name: str,
        *,
        model_id: object = "",
        model_appr: object = "",
        color: object = "",
        status: str = "可安装",
        headers: tuple[str, ...] = HEADERS,
    ) -> list[object]:
        values = {
            "状态": status,
            "怪物名称": name,
            "等级": 120,
            "经验": 888888,
            "血量": 2000000,
            "防御": 600,
            "魔防": 500,
            "最小攻击": 1000,
            "最大攻击": 1800,
            "命中": 999,
            "颜色": color,
            "模型库编号": model_id,
            "模型Appr": model_appr,
            "备注": "测试行",
        }
        return [values.get(header, "") for header in headers]

    def test_blank_models_are_stable_and_recorded(self):
        self.write_workbook([self.row("随机怪一"), self.row("随机怪二")])
        first = self.service.preflight(self.workbook, self.server, self.client_data)
        second = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertEqual(first.blockers, [])
        self.assertEqual(first.model_assignments, second.model_assignments)
        self.assertTrue(all(item["mode"] == "自动随机" for item in first.model_assignments))
        self.assertEqual(set(first.monster_names), {"随机怪一", "随机怪二"})

    def test_legacy_twelve_column_workbook_preserves_no_color_behavior(self):
        self.write_workbook(
            [self.row("旧表兼容怪", headers=LEGACY_HEADERS)],
            headers=LEGACY_HEADERS,
        )
        compiled = self.service.compile(self.workbook)
        self.assertEqual(compiled.blockers, ())
        self.assertIsNone(compiled.specs[0].name_color)

    def test_invalid_name_color_blocks(self):
        self.write_workbook([self.row("颜色错误怪", color="彩虹色")])
        compiled = self.service.compile(self.workbook)
        self.assertTrue(any("颜色只能填写" in item for item in compiled.blockers))

    def test_additional_name_color_candidates_are_supported(self):
        self.write_workbook([
            self.row("橙色怪", color="橙色"),
            self.row("绿色怪", color="绿色"),
            self.row("青色怪", color="青色"),
            self.row("粉色怪", color="粉色"),
        ])
        compiled = self.service.compile(self.workbook)
        self.assertEqual(compiled.blockers, ())
        self.assertEqual([spec.name_color for spec in compiled.specs], [70, 250, 254, 245])

    def test_name_colors_update_mongen_preserve_concentration_and_roll_back(self):
        mon_gen = self.server / "Mir200" / "Envir" / "MonGen.txt"
        mon_gen.parent.mkdir(parents=True)
        before = (
            "T001\t10\t10\t彩色怪\t20\t1\t5\r\n"
            "T001\t20\t20\t本端已有\t30\t2\t10\t37\r\n"
        ).encode("gb18030")
        mon_gen.write_bytes(before)
        self.write_workbook([
            self.row("彩色怪", color="红色"),
            self.row("本端已有", color="蓝色"),
        ])

        plan = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertEqual(plan.blockers, [])
        self.assertEqual([item["color"] for item in plan.name_color_assignments], [249, 154])
        self.assertEqual([item["matched_rows"] for item in plan.name_color_assignments], [1, 1])
        self.assertTrue(any(change.kind == "monster-name-colors" for change in plan.changes))

        receipt = self.service.install(plan)
        lines = mon_gen.read_bytes().decode("gb18030").splitlines()
        self.assertEqual(lines[0].split(), ["T001", "10", "10", "彩色怪", "20", "1", "5", "0", "249"])
        self.assertEqual(lines[1].split(), ["T001", "20", "20", "本端已有", "30", "2", "10", "37", "154"])

        repeat = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertEqual(repeat.blockers, [])
        self.assertFalse(any(change.kind == "monster-name-colors" for change in repeat.changes))
        self.service.rollback(receipt.transaction_id)
        self.assertEqual(mon_gen.read_bytes(), before)

    def test_explicit_model_overrides_stats_and_rolls_back(self):
        model = MonsterLibraryService(self.platform).list_monsters("ready")[0]
        self.write_workbook([self.row("自定义Boss", model_id=model.monster_id)])
        database = self.server / "Mud2" / "DB" / "ApexM2.DB"
        before = database.read_bytes()
        plan = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertEqual(plan.blockers, [])
        self.assertEqual(plan.model_assignments[0]["mode"], "指定库编号")
        receipt = self.service.install(plan)
        connection = sqlite3.connect(database)
        try:
            row = connection.execute(
                'SELECT "HP", "DC", "DCMAX", "HIT", "Appr", "Race" FROM "Monster" WHERE "Name"=?',
                ("自定义Boss",),
            ).fetchone()
        finally:
            connection.close()
        self.assertEqual(row, (2000000, 1000, 1800, 999, model.appearance_id, model.race))
        self.assertTrue((self.client_data / model.resource_name).is_file())
        self.service.rollback(receipt.transaction_id)
        self.assertEqual(database.read_bytes(), before)
        self.assertFalse((self.client_data / model.resource_name).exists())

    def test_new_monster_uses_target_neutral_defaults_not_donor_stats(self):
        model = MonsterLibraryService(self.platform).list_monsters("ready")[0]
        self.write_workbook([self.row("属性隔离怪", model_id=model.monster_id)])

        plan = self.service.preflight(self.workbook, self.server, self.client_data)

        self.assertEqual(plan.blockers, [])
        connection = sqlite3.connect(":memory:")
        try:
            connection.deserialize(plan.database_after)
            row = connection.execute(
                'SELECT "Name", "Race", "RaceImg", "Appr", "HP", "DC", "DCMAX", '
                '"SPEED", "WALK_SPD", "ATTACK_SPD", "CoolEye", "ExploreItem" '
                'FROM "Monster" WHERE "Name"=?',
                ("属性隔离怪",),
            ).fetchone()
        finally:
            connection.close()
        self.assertEqual(
            row,
            (
                "属性隔离怪", model.race, model.race_img, model.appearance_id,
                2000000, 1000, 1800, model.monster_values["SPEED"], 333, 777, 61, 901,
            ),
        )

    def test_blank_model_uses_only_ready_verified_pool(self):
        self.write_workbook([self.row("稳定池怪")])
        compiled = self.service.compile(self.workbook)
        ready = MonsterLibraryService(self.platform).list_monsters("all")
        expected_before, _ = self.service._resolve_model(
            compiled.specs[0], ready, compiled.workbook_hash,
        )
        catalog = self.platform / "怪物库" / "catalog.sqlite"
        connection = sqlite3.connect(catalog)
        try:
            connection.execute(
                "UPDATE monsters SET closure_status='ready_opaque' WHERE monster_id=?",
                (expected_before.monster_id,),
            )
            connection.commit()
        finally:
            connection.close()
        changed_ready = MonsterLibraryService(self.platform).list_monsters("all")

        first, _ = self.service._resolve_model(
            compiled.specs[0], changed_ready, compiled.workbook_hash,
        )
        second, _ = self.service._resolve_model(
            compiled.specs[0], list(reversed(changed_ready)), compiled.workbook_hash,
        )

        self.assertNotEqual(first.monster_id, expected_before.monster_id)
        self.assertEqual(first.closure_status, "ready_verified")
        self.assertEqual(first.monster_id, second.monster_id)

    def test_one_workbook_four_rows_preserves_attributes_models_and_preflight_targets(self):
        """同一临时XLSX必须同时覆盖新怪、保模、真实SmartMonster与稳定默认模型。"""
        connection = sqlite3.connect(self.donor_db)
        try:
            connection.execute(
                'UPDATE "Monster" SET "Race"=156, "RaceImg"=156, "Appr"=1123, '
                '"CoolEye"=7777, "ExploreItem"=8888 WHERE "Name"=?',
                ("模型丙",),
            )
            connection.commit()
        finally:
            connection.close()
        self.donor_smartmonster.mkdir(parents=True)
        write_wzl_wzx_pair(self.donor_wzl, "Mon7")
        write_effect_image_list(
            self.donor_effect_list,
            [f"Existing{index}.wzl" for index in range(78)] + ["Mon7.wzl"],
        )
        write_smartmonster_ini(self.donor_smartmonster / "模型丙.ini")
        target_envir = self.server / "Mir200" / "Envir"
        (target_envir / "SmartMonster").mkdir(parents=True)
        write_effect_image_list(target_envir / "EffectImageList.txt", ["History0.wzl"])
        (target_envir / "MonGen.txt").write_bytes(b"BASELINE-MONGEN\r\n")
        monitems = target_envir / "MonItems"
        monitems.mkdir()
        (monitems / "keep.txt").write_bytes(b"MONITEMS-BASELINE")
        (self.client_data / "History0.wzl").write_bytes(b"HISTORICAL-WZL")
        (self.client_data / "History0.wzx").write_bytes(b"HISTORICAL-WZX")
        MonsterLibraryService(self.platform).scan(
            self.donor_db, self.donor_wzl, self.donor_pak, self.donor_rules,
            self.server, self.client_data,
        )
        records = MonsterLibraryService(self.platform).list_monsters("all")
        smart = next(record for record in records if record.monster_name == "模型丙")
        self.assertEqual((smart.engine_mode, smart.closure_status), ("smartmonster", "ready_verified"))
        opaque = next(record for record in records if record.monster_name == "模型甲")
        catalog = self.platform / "怪物库" / "catalog.sqlite"
        connection = sqlite3.connect(catalog)
        try:
            connection.execute(
                "UPDATE monsters SET closure_status='ready_opaque' WHERE monster_id=?",
                (opaque.monster_id,),
            )
            connection.commit()
        finally:
            connection.close()

        def file_hashes(root: Path) -> dict[str, str]:
            if not root.exists():
                return {}
            return {
                str(path.relative_to(root)): hashlib.sha256(path.read_bytes()).hexdigest()
                for path in sorted(root.rglob("*")) if path.is_file()
            }

        database = self.server / "Mud2" / "DB" / "ApexM2.DB"
        mon_gen = target_envir / "MonGen.txt"
        before = {
            "database": hashlib.sha256(database.read_bytes()).hexdigest(),
            "mongen": hashlib.sha256(mon_gen.read_bytes()).hexdigest(),
            "monitems": file_hashes(monitems),
            "smart_ini": file_hashes(self.donor_smartmonster),
            "smart_resources": file_hashes(self.donor_wzl),
        }
        self.write_workbook([
            self.row("四行新怪"),
            self.row("本端已有"),
            self.row("显式Smart怪", model_id=smart.monster_id),
            self.row("稳定默认怪"),
        ])

        compiled = self.service.compile(self.workbook)
        catalog_records = MonsterLibraryService(self.platform).list_monsters("all")
        ready_verified = [
            record for record in catalog_records
            if record.status == "ready" and record.closure_status == "ready_verified"
        ]
        compatible_ready, compatibility_conflicts = self.service.library.compatible_models(
            ready_verified, self.server, self.client_data,
        )
        compatible_ids = {record.monster_id for record in compatible_ready}
        initial_candidates = {
            spec.name: self.service._resolve_model(spec, catalog_records, compiled.workbook_hash)[0].monster_id
            for spec in compiled.specs if spec.name in {"四行新怪", "稳定默认怪"}
        }

        plan = self.service.preflight(self.workbook, self.server, self.client_data)

        self.assertEqual(plan.blockers, [])
        after = {
            "database": hashlib.sha256(database.read_bytes()).hexdigest(),
            "mongen": hashlib.sha256(mon_gen.read_bytes()).hexdigest(),
            "monitems": file_hashes(monitems),
            "smart_ini": file_hashes(self.donor_smartmonster),
            "smart_resources": file_hashes(self.donor_wzl),
        }
        self.assertEqual(after, before)
        assignments = {item["monster_name"]: item for item in plan.model_assignments}
        self.assertEqual(assignments["本端已有"]["mode"], "保留现有模型")
        self.assertEqual(assignments["显式Smart怪"]["model_id"], smart.monster_id)
        self.assertEqual(assignments["显式Smart怪"]["mode"], "指定库编号")
        selected = {record.monster_id: record for record in MonsterLibraryService(self.platform).list_monsters("all")}
        repeated = self.service.preflight(self.workbook, self.server, self.client_data)
        repeated_assignments = {item["monster_name"]: item for item in repeated.model_assignments}
        reselected = {item["monster_name"]: item for item in plan.auto_reselected}
        self.assertEqual(repeated.auto_reselected, plan.auto_reselected)
        for name in ("四行新怪", "稳定默认怪"):
            assignment = assignments[name]
            self.assertEqual(selected[assignment["model_id"]].closure_status, "ready_verified")
            self.assertEqual(repeated_assignments[name]["model_id"], assignment["model_id"])
            if assignment["mode"] == "自动随机":
                self.assertEqual(assignment["model_id"], initial_candidates[name])
                self.assertNotIn(name, reselected)
            elif assignment["mode"] == "自动兼容改选":
                change = reselected[name]
                self.assertEqual(change["original_model_id"], initial_candidates[name])
                self.assertEqual(change["selected_model_id"], assignment["model_id"])
                self.assertNotEqual(change["selected_model_id"], change["original_model_id"])
                self.assertIn(change["selected_model_id"], compatible_ids)
                self.assertEqual(change["reason"], compatibility_conflicts[change["original_model_id"]])
            else:
                self.fail(f"{name}出现未定义的空模型选择模式：{assignment['mode']}")
        self.assertEqual(selected[assignments["显式Smart怪"]["model_id"]].engine_mode, "smartmonster")

        target = sqlite3.connect(":memory:")
        try:
            target.deserialize(plan.database_after)
            for name in ("四行新怪", "本端已有", "显式Smart怪", "稳定默认怪"):
                row = target.execute(
                    'SELECT "Lvl", "Exp", "HP", "AC", "MAC", "DC", "DCMAX", "HIT" '
                    'FROM "Monster" WHERE "Name"=?', (name,),
                ).fetchone()
                self.assertEqual(row, (120, 888888, 2000000, 600, 500, 1000, 1800, 999))
            for name in ("四行新怪", "显式Smart怪", "稳定默认怪"):
                row = target.execute(
                    'SELECT "CoolEye", "ExploreItem" FROM "Monster" WHERE "Name"=?', (name,),
                ).fetchone()
                self.assertEqual(row, (61, 901))
        finally:
            target.close()
        print(json.dumps({
            "task10_four_row_xlsx": {
                "rows": ["四行新怪", "本端已有", "显式Smart怪", "稳定默认怪"],
                "preflight_before_sha256": before,
                "preflight_after_sha256": after,
                "preflight_unchanged": after == before,
                "explicit_smartmonster": {
                    "model_id": smart.monster_id,
                    "engine_mode": smart.engine_mode,
                    "closure_status": smart.closure_status,
                },
                "blank_models": [
                    assignments["四行新怪"]["model_id"],
                    assignments["稳定默认怪"]["model_id"],
                ],
                "donor_attribute_leakage": 0,
            },
        }, ensure_ascii=False, sort_keys=True))

    def test_ready_opaque_requires_explicit_model_id(self):
        model = MonsterLibraryService(self.platform).list_monsters("ready")[0]
        catalog = self.platform / "怪物库" / "catalog.sqlite"
        connection = sqlite3.connect(catalog)
        try:
            connection.execute(
                "UPDATE monsters SET closure_status='ready_opaque' WHERE monster_id=?",
                (model.monster_id,),
            )
            connection.commit()
        finally:
            connection.close()
        self.write_workbook([self.row("不透明模型怪", model_id=model.monster_id)])

        plan = self.service.preflight(self.workbook, self.server, self.client_data)

        self.assertEqual(plan.blockers, [])
        self.assertEqual(plan.model_assignments[0]["model_id"], model.monster_id)
        self.assertTrue(plan.model_assignments[0]["single_monster_acceptance_required"])
        summary = self.service.plan_summary(plan)
        self.assertTrue(summary["model_assignments"][0]["single_monster_acceptance_required"])
        self.assertTrue(any("ready_opaque" in warning and "单怪验收" in warning for warning in plan.warnings))

    def test_neutral_template_cardinality_blocks_new_monsters_only(self):
        model = MonsterLibraryService(self.platform).list_monsters("ready")[0]
        database = self.server / "Mud2" / "DB" / "ApexM2.DB"
        connection = sqlite3.connect(database)
        try:
            connection.execute('DELETE FROM "Monster" WHERE "Name"=?', ("稻草人",))
            connection.commit()
        finally:
            connection.close()

        self.write_workbook([self.row("缺模板新怪", model_id=model.monster_id)])
        missing = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertTrue(any("缺少唯一中性模板：稻草人" in item for item in missing.blockers))

        self.write_workbook([self.row("本端已有")])
        existing = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertEqual(existing.blockers, [])

        connection = sqlite3.connect(database)
        try:
            connection.executemany(
                'INSERT INTO "Monster" (' + ", ".join(f'"{item}"' for item in MONSTER_COLUMNS) + ') '
                'VALUES (' + ", ".join("?" for _ in MONSTER_COLUMNS) + ')',
                [
                    tuple(monster_values("稻草人", 3)[item] for item in MONSTER_COLUMNS),
                    tuple(monster_values("稻草人", 4)[item] for item in MONSTER_COLUMNS),
                ],
            )
            connection.commit()
        finally:
            connection.close()
        self.write_workbook([self.row("重复模板新怪", model_id=model.monster_id)])
        duplicate = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertTrue(any("中性模板重复：稻草人" in item for item in duplicate.blockers))

    def test_explicit_incomplete_and_complex_models_are_rejected(self):
        model = MonsterLibraryService(self.platform).list_monsters("ready")[0]
        catalog = self.platform / "怪物库" / "catalog.sqlite"
        for closure_status in ("incomplete", "complex_ability"):
            with self.subTest(closure_status=closure_status):
                connection = sqlite3.connect(catalog)
                try:
                    connection.execute(
                        "UPDATE monsters SET closure_status=? WHERE monster_id=?",
                        (closure_status, model.monster_id),
                    )
                    connection.commit()
                finally:
                    connection.close()
                self.write_workbook([self.row(f"{closure_status}模型怪", model_id=model.monster_id)])

                plan = self.service.preflight(self.workbook, self.server, self.client_data)

                self.assertTrue(any(
                    f"资料未闭环（{closure_status}）" in item for item in plan.blockers
                ))

    def test_explicit_model_error_distinguishes_missing_id_from_unclosed_id(self):
        model = MonsterLibraryService(self.platform).list_monsters("ready")[0]
        catalog = self.platform / "怪物库" / "catalog.sqlite"
        connection = sqlite3.connect(catalog)
        try:
            connection.execute(
                "UPDATE monsters SET closure_status='incomplete' WHERE monster_id=?",
                (model.monster_id,),
            )
            connection.commit()
        finally:
            connection.close()

        self.write_workbook([self.row("未闭环编号怪", model_id=model.monster_id)])
        unclosed = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertTrue(any(
            f"模型库编号{model.monster_id}已收录，但资料未闭环（incomplete）" in item
            for item in unclosed.blockers
        ))

        self.write_workbook([self.row("不存在编号怪", model_id=999999)])
        missing = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertTrue(any("模型库编号999999不存在" in item for item in missing.blockers))

    def test_workbook_speed_overrides_win_over_neutral_defaults(self):
        model = MonsterLibraryService(self.platform).list_monsters("ready")[0]
        headers = HEADERS + ("行走速度", "攻击速度")
        overridden = self.row("速度覆盖怪", model_id=model.monster_id, headers=headers)
        overridden[headers.index("行走速度")] = 456
        overridden[headers.index("攻击速度")] = 789
        self.write_workbook([
            overridden,
            self.row("速度中性怪", model_id=model.monster_id, headers=headers),
        ], headers=headers)

        plan = self.service.preflight(self.workbook, self.server, self.client_data)

        self.assertEqual(plan.blockers, [])
        connection = sqlite3.connect(":memory:")
        try:
            connection.deserialize(plan.database_after)
            rows = {
                name: (walk_spd, attack_spd)
                for name, walk_spd, attack_spd in connection.execute(
                'SELECT "Name", "WALK_SPD", "ATTACK_SPD" FROM "Monster" '
                'WHERE "Name" IN (?, ?)',
                ("速度覆盖怪", "速度中性怪"),
                ).fetchall()
            }
        finally:
            connection.close()
        self.assertEqual(rows["速度覆盖怪"], (456, 789))
        self.assertEqual(rows["速度中性怪"], (333, 777))

    def test_existing_monster_blank_model_preserves_current_model(self):
        database = self.server / "Mud2" / "DB" / "ApexM2.DB"
        connection = sqlite3.connect(database)
        try:
            before = connection.execute(
                'SELECT "Race", "RaceImg", "Appr", "SPEED", "WalkStep", "WalkWait", '
                '"AttackState", "AttackSource", "DisableSimpleActor" '
                'FROM "Monster" WHERE "Name"=?',
                ("本端已有",),
            ).fetchone()
            connection.execute('DELETE FROM "Monster" WHERE "Name"=?', ("稻草人",))
            connection.commit()
        finally:
            connection.close()
        self.write_workbook([self.row("本端已有")])

        plan = self.service.preflight(self.workbook, self.server, self.client_data)

        self.assertEqual(plan.blockers, [])
        connection = sqlite3.connect(":memory:")
        try:
            connection.deserialize(plan.database_after)
            after = connection.execute(
                'SELECT "Race", "RaceImg", "Appr", "SPEED", "WalkStep", "WalkWait", '
                '"AttackState", "AttackSource", "DisableSimpleActor" '
                'FROM "Monster" WHERE "Name"=?',
                ("本端已有",),
            ).fetchone()
        finally:
            connection.close()
        self.assertEqual(after, before)

    def test_existing_name_updates_stats_and_preserves_visual_then_rolls_back(self):
        database = self.server / "Mud2" / "DB" / "ApexM2.DB"
        before_bytes = database.read_bytes()
        connection = sqlite3.connect(database)
        try:
            visual_before = connection.execute(
                'SELECT "Race", "RaceImg", "Appr", "CoolEye", "SPEED", "WalkStep", '
                '"WalkWait", "AttackState", "AttackSource", "ExploreItem", "DisableSimpleActor" '
                'FROM "Monster" WHERE "Name"=?',
                ("本端已有",),
            ).fetchone()
        finally:
            connection.close()

        self.write_workbook([self.row("本端已有")])
        plan = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertEqual(plan.blockers, [])
        self.assertEqual(plan.inserted_names, ())
        self.assertEqual(plan.updated_names, ("本端已有",))
        self.assertEqual(plan.appearance_changed_names, ())
        self.assertEqual(plan.model_assignments[0]["mode"], "保留现有模型")
        self.assertEqual(plan.library_numbers, ())

        receipt = self.service.install(plan)
        connection = sqlite3.connect(database)
        try:
            stats = connection.execute(
                'SELECT "Lvl", "Exp", "HP", "AC", "MAC", "DC", "DCMAX", "HIT" '
                'FROM "Monster" WHERE "Name"=?',
                ("本端已有",),
            ).fetchone()
            visual_after = connection.execute(
                'SELECT "Race", "RaceImg", "Appr", "CoolEye", "SPEED", "WalkStep", '
                '"WalkWait", "AttackState", "AttackSource", "ExploreItem", "DisableSimpleActor" '
                'FROM "Monster" WHERE "Name"=?',
                ("本端已有",),
            ).fetchone()
            count = connection.execute(
                'SELECT COUNT(*) FROM "Monster" WHERE "Name"=?', ("本端已有",)
            ).fetchone()[0]
        finally:
            connection.close()
        self.assertEqual(stats, (120, 888888, 2000000, 600, 500, 1000, 1800, 999))
        self.assertEqual(visual_after, visual_before)
        self.assertEqual(count, 1)
        self.service.rollback(receipt.transaction_id)
        self.assertEqual(database.read_bytes(), before_bytes)

    def test_existing_name_with_explicit_model_changes_visual(self):
        model = MonsterLibraryService(self.platform).list_monsters("ready")[1]
        self.write_workbook([self.row("本端已有", model_id=model.monster_id)])
        plan = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertEqual(plan.blockers, [])
        self.assertEqual(plan.inserted_names, ())
        self.assertEqual(plan.updated_names, ("本端已有",))
        self.assertEqual(plan.appearance_changed_names, ("本端已有",))
        self.assertEqual(plan.model_assignments[0]["action"], "更新属性并换模")
        receipt = self.service.install(plan)
        connection = sqlite3.connect(self.server / "Mud2" / "DB" / "ApexM2.DB")
        try:
            row = connection.execute(
                'SELECT "Appr", "Race", "HP" FROM "Monster" WHERE "Name"=?',
                ("本端已有",),
            ).fetchone()
        finally:
            connection.close()
        self.assertEqual(row, (model.appearance_id, model.race, 2000000))
        self.service.rollback(receipt.transaction_id)

    def test_auto_model_reselects_away_from_target_patch_conflict(self):
        self.write_workbook([self.row("自动避冲突怪")])
        compiled = self.service.compile(self.workbook)
        ready = MonsterLibraryService(self.platform).list_monsters("ready")
        original, _ = self.service._resolve_model(compiled.specs[0], ready, compiled.workbook_hash)
        (self.client_data / f"Mon{original.library_no}.wzl").write_bytes(b"TARGET-WZL")
        (self.client_data / f"Mon{original.library_no}.wzx").write_bytes(b"TARGET-WZX")

        first = self.service.preflight(self.workbook, self.server, self.client_data)
        second = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertEqual(first.blockers, [])
        self.assertEqual(first.skipped, [])
        self.assertEqual(first.auto_reselected, second.auto_reselected)
        self.assertEqual(len(first.auto_reselected), 1)
        self.assertEqual(first.model_assignments[0]["mode"], "自动兼容改选")
        self.assertNotEqual(first.model_assignments[0]["model_id"], original.monster_id)
        self.assertEqual(first.inserted_names, ("自动避冲突怪",))

    def test_repeat_sync_is_idempotent_and_first_transaction_still_rolls_back(self):
        database = self.server / "Mud2" / "DB" / "ApexM2.DB"
        before = database.read_bytes()
        self.write_workbook([self.row("幂等怪")])
        first = self.service.preflight(self.workbook, self.server, self.client_data)
        receipt1 = self.service.install(first)
        after_first = database.read_bytes()

        second = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertEqual(second.blockers, [])
        self.assertEqual(second.inserted_names, ())
        self.assertEqual(second.updated_names, ())
        self.assertEqual(second.unchanged_names, ("幂等怪",))
        self.assertEqual(second.changes, [])
        receipt2 = self.service.install(second)
        self.assertEqual(database.read_bytes(), after_first)
        self.service.rollback(receipt2.transaction_id)
        self.assertEqual(database.read_bytes(), after_first)
        self.service.rollback(receipt1.transaction_id)
        self.assertEqual(database.read_bytes(), before)

    def test_duplicate_target_name_blocks_sync(self):
        database = self.server / "Mud2" / "DB" / "ApexM2.DB"
        connection = sqlite3.connect(database)
        try:
            row = monster_values("本端已有", 11)
            columns = ", ".join(f'"{item}"' for item in MONSTER_COLUMNS)
            placeholders = ", ".join("?" for _ in MONSTER_COLUMNS)
            connection.execute(
                f'INSERT INTO "Monster" ({columns}) VALUES ({placeholders})',
                tuple(row[item] for item in MONSTER_COLUMNS),
            )
            connection.commit()
        finally:
            connection.close()
        self.write_workbook([self.row("本端已有")])
        plan = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertTrue(any("多个同名怪物" in item for item in plan.blockers))
        with self.assertRaises(MonsterLibraryError):
            self.service.install(plan)

    def test_pending_and_mismatched_model_block_install(self):
        model = MonsterLibraryService(self.platform).list_monsters("ready")[0]
        headers = HEADERS + ("模型Appr",)
        self.write_workbook([
            self.row("待填怪", status="待配置", headers=headers),
            self.row("冲突怪", model_id=model.monster_id, model_appr=model.appearance_id + 1, headers=headers),
            self.row("忽略怪", status="忽略", headers=headers),
        ], headers=headers)
        plan = self.service.preflight(self.workbook, self.server, self.client_data)
        self.assertTrue(any("待配置" in item for item in plan.blockers))
        self.assertTrue(any("不一致" in item for item in plan.blockers))
        self.assertTrue(any("状态为忽略" in item for item in plan.skipped))
        with self.assertRaises(MonsterLibraryError):
            self.service.install(plan)

    def test_workbook_change_after_preflight_blocks_install(self):
        self.write_workbook([self.row("哈希门禁怪")])
        plan = self.service.preflight(self.workbook, self.server, self.client_data)
        book = load_workbook(self.workbook)
        book["怪物生成"]["L2"] = "预检后改动"
        book.save(self.workbook)
        book.close()
        with self.assertRaisesRegex(MonsterLibraryError, "生成表发生变化"):
            self.service.install(plan)

    def test_cli_exposes_workbook_commands(self):
        parser = _parser()
        args = parser.parse_args([
            "monster-workbook-preflight", "--input", str(self.workbook),
            "--server", str(self.server), "--client-data", str(self.client_data),
        ])
        self.assertEqual(args.command, "monster-workbook-preflight")
        self.assertEqual(args.input, self.workbook)


if __name__ == "__main__":
    unittest.main()
