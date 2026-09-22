from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sqlite3
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from xydp.monster_library import _monster_rows
from xydp.monster_workbook import MonsterWorkbookService


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def run(platform: Path, source_server: Path, workbook: Path, evidence: Path) -> dict[str, object]:
    catalog = platform / "怪物库" / "catalog.sqlite"
    source_database = source_server / "Mud2" / "DB" / "ApexM2.DB"
    source_pak_rules = source_server / "登录器" / "pak.txt"
    for required in (catalog, source_database, source_pak_rules, workbook):
        if not required.is_file():
            raise RuntimeError(f"验证来源不存在：{required}")

    evidence.mkdir(parents=True, exist_ok=True)
    verification_name = f"玄渊批量生成验证怪_{datetime.now():%Y%m%d_%H%M%S}"
    with tempfile.TemporaryDirectory(prefix="monster-workbook-release-", dir=platform / "build") as temp_dir:
        temp = Path(temp_dir)
        temp_platform = temp / "platform"
        temp_server = temp / "server"
        temp_client = temp / "client" / "data"
        temp_workbook = temp / "20_怪物批量生成_验证.xlsx"
        (temp_platform / "怪物库").mkdir(parents=True)
        (temp_server / "Mud2" / "DB").mkdir(parents=True)
        (temp_server / "登录器").mkdir(parents=True)
        (temp_server / "Mir200" / "Envir").mkdir(parents=True)
        temp_client.mkdir(parents=True)
        shutil.copy2(catalog, temp_platform / "怪物库" / "catalog.sqlite")
        shutil.copy2(source_database, temp_server / "Mud2" / "DB" / "ApexM2.DB")
        shutil.copy2(source_pak_rules, temp_server / "登录器" / "pak.txt")

        source_rows = _monster_rows(source_database)
        if not source_rows:
            raise RuntimeError("来源Monster数据库没有可用于更新验证的现有怪物")
        existing_before = source_rows[0][1]
        existing_name = str(existing_before["Name"])
        existing_model = tuple(int(existing_before[field]) for field in ("Race", "RaceImg", "Appr"))

        def changed(value: object, amount: int = 1) -> int:
            current = int(value)
            return current + amount if current < 2_000_000_000 - amount else current - amount

        book = load_workbook(workbook)
        sheet = book["怪物生成"]
        headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value}
        new_values = {
            "状态": "可安装",
            "怪物名称": verification_name,
            "等级": 321,
            "经验": 7654321,
            "血量": 9876543,
            "防御": 432,
            "魔防": 321,
            "最小攻击": 2345,
            "最大攻击": 3456,
            "命中": 888,
            "颜色": "绿色",
            "模型库编号": None,
            "备注": "正式构建临时验证；自动随机模型",
        }
        existing_values = {
            "状态": "可安装",
            "怪物名称": existing_name,
            "等级": changed(existing_before["Lvl"], 3),
            "经验": changed(existing_before["Exp"], 7),
            "血量": changed(existing_before["HP"], 11),
            "防御": changed(existing_before["AC"], 13),
            "魔防": changed(existing_before["MAC"], 17),
            "最小攻击": changed(existing_before["DC"], 19),
            "最大攻击": max(changed(existing_before["DCMAX"], 23), changed(existing_before["DC"], 19)),
            "命中": changed(existing_before["HIT"], 29),
            "颜色": "青色",
            "模型库编号": None,
            "备注": "正式构建临时验证；同名更新且保留目标外观",
        }
        for header, value in new_values.items():
            sheet.cell(2, headers[header], value)
        for header, value in existing_values.items():
            sheet.cell(3, headers[header], value)
        book.save(temp_workbook)
        book.close()

        database = temp_server / "Mud2" / "DB" / "ApexM2.DB"
        pak_rules = temp_server / "登录器" / "pak.txt"
        mon_gen = temp_server / "Mir200" / "Envir" / "MonGen.txt"
        mon_gen.write_bytes((
            f"T001\t10\t10\t{verification_name}\t20\t1\t5\r\n"
            f"T001\t20\t20\t{existing_name}\t30\t2\t10\t37\r\n"
        ).encode("gb18030"))
        database_before = database.read_bytes()
        pak_rules_before = pak_rules.read_bytes()
        mon_gen_before = mon_gen.read_bytes()
        client_before = sorted(path.name for path in temp_client.iterdir())

        service = MonsterWorkbookService(temp_platform)
        plan = service.preflight(temp_workbook, temp_server, temp_client)
        if plan.blockers:
            raise RuntimeError("临时目标预检被阻止：" + "；".join(plan.blockers))
        if set(plan.monster_names) != {verification_name, existing_name}:
            raise RuntimeError(f"预检怪物名异常：{plan.monster_names}")
        if len(plan.model_assignments) != 2 or plan.model_assignments[0]["mode"] != "自动随机":
            raise RuntimeError(f"自动随机模型未生效：{plan.model_assignments}")
        if plan.model_assignments[1]["mode"] != "保留现有模型":
            raise RuntimeError(f"同名怪物未保留目标模型：{plan.model_assignments}")
        if plan.inserted_names != (verification_name,) or plan.updated_names != (existing_name,):
            raise RuntimeError(f"新增/更新分类异常：{plan.inserted_names} / {plan.updated_names}")
        if [item["color"] for item in plan.name_color_assignments] != [250, 254]:
            raise RuntimeError(f"名字颜色映射异常：{plan.name_color_assignments}")
        if [item["matched_rows"] for item in plan.name_color_assignments] != [1, 1]:
            raise RuntimeError(f"MonGen颜色匹配异常：{plan.name_color_assignments}")

        receipt = service.install(plan)
        connection = sqlite3.connect(database)
        try:
            installed_row = connection.execute(
                'SELECT "Lvl", "Exp", "HP", "AC", "MAC", "DC", "DCMAX", "HIT", "Appr", "Race", "RaceImg" '
                'FROM "Monster" WHERE "Name"=?',
                (verification_name,),
            ).fetchone()
            updated_row = connection.execute(
                'SELECT "Lvl", "Exp", "HP", "AC", "MAC", "DC", "DCMAX", "HIT", "Race", "RaceImg", "Appr" '
                'FROM "Monster" WHERE "Name"=?',
                (existing_name,),
            ).fetchone()
        finally:
            connection.close()
        expected_stats = (321, 7654321, 9876543, 432, 321, 2345, 3456, 888)
        if installed_row is None or tuple(installed_row[:8]) != expected_stats:
            raise RuntimeError(f"临时Monster基础属性不一致：{installed_row}")
        expected_updated_stats = tuple(existing_values[header] for header in (
            "等级", "经验", "血量", "防御", "魔防", "最小攻击", "最大攻击", "命中",
        ))
        if updated_row is None or tuple(updated_row[:8]) != expected_updated_stats:
            raise RuntimeError(f"同名Monster属性更新不一致：{updated_row}")
        if tuple(updated_row[8:]) != existing_model:
            raise RuntimeError(f"同名Monster模型没有保留：{existing_model} -> {updated_row[8:]}")
        client_after_install = sorted(path.name for path in temp_client.iterdir())
        if not client_after_install:
            raise RuntimeError("临时安装没有生成客户端怪物补丁")
        if database.read_bytes() == database_before:
            raise RuntimeError("临时安装没有修改Monster数据库")
        mon_gen_rows = [line.split() for line in mon_gen.read_bytes().decode("gb18030").splitlines()]
        if mon_gen_rows[0][-2:] != ["0", "250"]:
            raise RuntimeError(f"新怪MonGen颜色或集中概率异常：{mon_gen_rows[0]}")
        if mon_gen_rows[1][-2:] != ["37", "254"]:
            raise RuntimeError(f"已有MonGen集中概率未保留或颜色异常：{mon_gen_rows[1]}")

        repeat_plan = service.preflight(temp_workbook, temp_server, temp_client)
        if repeat_plan.blockers or repeat_plan.changes:
            raise RuntimeError(f"重复同步不幂等：{repeat_plan.blockers} / {repeat_plan.changes}")
        if set(repeat_plan.unchanged_names) != {verification_name, existing_name}:
            raise RuntimeError(f"重复同步未归类为无需改动：{repeat_plan.unchanged_names}")

        service.rollback(receipt.transaction_id)
        rollback_checks = {
            "database_byte_exact": database.read_bytes() == database_before,
            "pak_txt_byte_exact": pak_rules.read_bytes() == pak_rules_before,
            "mon_gen_byte_exact": mon_gen.read_bytes() == mon_gen_before,
            "client_files_exact": sorted(path.name for path in temp_client.iterdir()) == client_before,
        }
        if not all(rollback_checks.values()):
            raise RuntimeError(f"临时逐字节回滚失败：{rollback_checks}")

        catalog_connection = sqlite3.connect(catalog)
        try:
            catalog_ready_count = catalog_connection.execute(
                "SELECT count(1) FROM monsters WHERE status='ready'"
            ).fetchone()[0]
        finally:
            catalog_connection.close()
        result = {
            "status": "passed",
            "source_workbook": str(workbook),
            "source_workbook_sha256": sha256(workbook.read_bytes()),
            "catalog_ready_count": catalog_ready_count,
            "verification_monster": verification_name,
            "model_assignment": plan.model_assignments[0],
            "installed_stats": list(installed_row),
            "updated_existing_monster": existing_name,
            "updated_existing_stats": list(updated_row),
            "existing_model_preserved": tuple(updated_row[8:]) == existing_model,
            "repeat_sync_no_changes": not repeat_plan.changes,
            "name_color_assignments": list(plan.name_color_assignments),
            "mon_gen_rows_after_install": mon_gen_rows,
            "client_files_after_install": client_after_install,
            "transaction_id": receipt.transaction_id,
            "rollback_checks": rollback_checks,
            "real_target_written": False,
            "source_database_sha256": sha256(source_database.read_bytes()),
            "source_pak_txt_sha256": sha256(source_pak_rules.read_bytes()),
        }

    evidence_file = evidence / "monster-workbook-release-verification.json"
    evidence_file.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="用真实怪物库和目标只读副本验证20号怪物批量生成事务")
    parser.add_argument("--platform", type=Path, required=True)
    parser.add_argument("--source-server", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--evidence", type=Path, required=True)
    args = parser.parse_args()
    result = run(
        args.platform.resolve(),
        args.source_server.resolve(),
        args.workbook.resolve(),
        args.evidence.resolve(),
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
