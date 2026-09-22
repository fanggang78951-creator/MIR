from __future__ import annotations

import argparse
import sqlite3
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = [
    "状态",
    "怪物名称",
    "等级",
    "经验",
    "血量",
    "防御",
    "魔防",
    "最小攻击",
    "最大攻击",
    "命中",
    "颜色",
    "模型库编号",
    "备注",
]

REQUIRED = {
    "怪物名称",
    "等级",
    "经验",
    "血量",
    "防御",
    "魔防",
    "最小攻击",
    "最大攻击",
    "命中",
}

NAVY = "1F4E78"
BLUE = "D9EAF7"
LIGHT_BLUE = "EAF3F8"
GOLD = "FFF2CC"
GREEN = "E2F0D9"
RED = "FCE4D6"
PURPLE = "E4DFEC"
ORANGE = "F4B183"
CYAN = "CCFFFF"
PINK = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
TEXT = "1F2937"
MUTED = "5B6573"
LIGHT_BORDER = Side(style="thin", color="D9E2F3")
MEDIUM_BORDER = Side(style="medium", color=NAVY)


def _catalog_rows(catalog: Path) -> list[tuple[object, ...]]:
    with sqlite3.connect(catalog) as connection:
        return connection.execute(
            """
            SELECT monster_id, monster_name
            FROM monsters
            WHERE status = 'ready'
            ORDER BY monster_id
            """
        ).fetchall()


def _base_font(size: int = 10, bold: bool = False, color: str = TEXT) -> Font:
    return Font(name="Microsoft YaHei", size=size, bold=bold, color=color)


def _set_header(cell, *, required: bool = False, model: bool = False) -> None:
    if required:
        fill = PatternFill("solid", fgColor="17365D")
    elif model:
        fill = PatternFill("solid", fgColor="806000")
    else:
        fill = PatternFill("solid", fgColor=NAVY)
    cell.fill = fill
    cell.font = _base_font(10, True, WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=MEDIUM_BORDER)
    cell.number_format = "@"


def _add_generation_sheet(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "怪物生成"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = NAVY
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = "A1:M500"
    sheet.row_dimensions[1].height = 34

    widths = {
        "A": 12,
        "B": 26,
        "C": 10,
        "D": 14,
        "E": 16,
        "F": 11,
        "G": 11,
        "H": 13,
        "I": 13,
        "J": 11,
        "K": 12,
        "L": 14,
        "M": 42,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, column, header)
        _set_header(cell, required=header in REQUIRED, model=header == "模型库编号")
        requirement = "必填" if header in REQUIRED else "可选"
        if header == "状态":
            message = "必填。可安装=参与预检与生成；待配置=阻止提交；忽略=整行跳过。"
        elif header == "模型库编号":
            message = "可选。留空时，平台从完整怪物库中稳定随机选择外观和动作模型。"
        elif header == "颜色":
            message = "可选。基础四级为黄、蓝、紫、红，另提供橙、绿、青、粉候选；留空默认为黄色。"
        elif header == "备注":
            message = "可选，仅供填写人说明，不写入Monster数据库。"
        else:
            message = f"{requirement}。数值必须是整数。"
        cell.comment = Comment(message, "玄渊平台")

    example = [
        "忽略",
        "示例怪物_请修改",
        100,
        100000,
        1000000,
        500,
        500,
        1000,
        1500,
        500,
        "黄色",
        None,
        "示例行：修改名称和基础属性，再把状态改成“可安装”；模型库编号留空即稳定随机。",
    ]
    for column, value in enumerate(example, start=1):
        sheet.cell(2, column, value)

    for row in range(2, 501):
        sheet.row_dimensions[row].height = 24
        for column in range(1, 14):
            cell = sheet.cell(row, column)
            cell.font = _base_font(10)
            cell.alignment = Alignment(
                horizontal="left" if column in {2, 13} else "center",
                vertical="center",
                wrap_text=column == 13,
            )
            cell.border = Border(bottom=LIGHT_BORDER)
            if column == 12:
                cell.fill = PatternFill("solid", fgColor=GOLD)
            elif column == 11:
                cell.fill = PatternFill("solid", fgColor=PURPLE)
            elif HEADERS[column - 1] in REQUIRED:
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            cell.number_format = "@" if column in {1, 2, 12} else "#,##0"

    status_validation = DataValidation(
        type="list",
        formula1='"可安装,待配置,忽略"',
        allow_blank=False,
    )
    status_validation.promptTitle = "选择处理状态"
    status_validation.prompt = "可安装、待配置或忽略"
    status_validation.errorTitle = "状态不正确"
    status_validation.error = "只能选择：可安装、待配置、忽略"
    status_validation.errorStyle = "stop"
    status_validation.showErrorMessage = True
    status_validation.showInputMessage = True
    sheet.add_data_validation(status_validation)
    status_validation.add("A2:A500")

    integer_columns = ["C", "D", "E", "F", "G", "H", "I", "J", "L"]
    integer_validation = DataValidation(
        type="whole",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
    )
    integer_validation.errorTitle = "必须填写整数"
    integer_validation.error = "该列只允许填写0或更大的整数；需要继承模型时请留空。"
    integer_validation.errorStyle = "stop"
    integer_validation.showErrorMessage = True
    sheet.add_data_validation(integer_validation)
    for column in integer_columns:
        integer_validation.add(f"{column}2:{column}500")

    color_validation = DataValidation(
        type="list",
        formula1='"黄色,蓝色,紫色,红色,橙色,绿色,青色,粉色"',
        allow_blank=True,
    )
    color_validation.promptTitle = "选择怪物名字颜色"
    color_validation.prompt = "黄/蓝/紫/红为基础四级；另可选橙/绿/青/粉；留空默认黄色"
    color_validation.errorTitle = "颜色不正确"
    color_validation.error = "只能选择：黄色、蓝色、紫色、红色、橙色、绿色、青色、粉色"
    color_validation.errorStyle = "stop"
    color_validation.showErrorMessage = True
    color_validation.showInputMessage = True
    sheet.add_data_validation(color_validation)
    color_validation.add("K2:K500")

    sheet.conditional_formatting.add(
        "A2:A500",
        FormulaRule(formula=['$A2="可安装"'], fill=PatternFill("solid", fgColor=GREEN)),
    )
    sheet.conditional_formatting.add(
        "A2:A500",
        FormulaRule(formula=['$A2="待配置"'], fill=PatternFill("solid", fgColor=RED)),
    )
    sheet.conditional_formatting.add(
        "A2:A500",
        FormulaRule(formula=['$A2="忽略"'], fill=PatternFill("solid", fgColor=GRAY)),
    )
    for value, fill in (
        ("黄色", GOLD), ("蓝色", BLUE), ("紫色", PURPLE), ("红色", RED),
        ("橙色", ORANGE), ("绿色", GREEN), ("青色", CYAN), ("粉色", PINK),
    ):
        sheet.conditional_formatting.add(
            "K2:K500",
            FormulaRule(formula=[f'$K2="{value}"'], fill=PatternFill("solid", fgColor=fill)),
        )

    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _add_instruction_sheet(workbook: Workbook, catalog_count: int) -> None:
    sheet = workbook.create_sheet("填写说明")
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "5B9BD5"
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 92
    sheet.column_dimensions["C"].width = 16
    sheet.freeze_panes = "A4"

    sheet.merge_cells("A1:C1")
    title = sheet["A1"]
    title.value = "玄渊怪物批量生成母表"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = _base_font(16, True, WHITE)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 36

    sheet.merge_cells("A2:C2")
    subtitle = sheet["A2"]
    subtitle.value = f"当前怪物库含 {catalog_count} 条完整可用模型；本表生成怪物资料、动作补丁，并给已有MonGen刷新行同步名字颜色。"
    subtitle.fill = PatternFill("solid", fgColor=BLUE)
    subtitle.font = _base_font(10, False, NAVY)
    subtitle.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 30

    rows = [
        ("操作顺序", "1. 在“怪物生成”页填写；2. 保存并关闭Excel/WPS；3. 平台“怪物库”页点击“预检生成表”；4. blockers清零后点击“确认批量生成”；5. 记录事务号并进游戏验收。", "必须先预检"),
        ("状态", "可安装=参与生成；待配置=整份表阻止提交；忽略=该行跳过。模板示例行默认是忽略，防止误生成。", "三选一"),
        ("必填列", "怪物名称、等级、经验、血量、防御、魔防、最小攻击、最大攻击、命中。", "不能为空"),
        ("名字颜色", "基础四级可选黄色、蓝色、紫色、红色；另提供橙色、绿色、青色、粉色候选。留空自动按黄色。", "只多这一格"),
        ("随机模型", "模型库编号留空时，平台从完整怪物库稳定随机选择。相同未修改文件重复预检，外观模型不会变化。", "推荐入门"),
        ("指定模型", "需要指定外观时，只填写“模型库编号”；编号可在“模型参考”页按外观名称筛选。", "完全可选"),
        ("自动补丁", "PAK、WZL/WZX、图库号、槽位、密码和登录器规则全部由怪物库与平台事务自动处理，不在表格填写。", "平台自动"),
        ("同名同步", "目标Monster数据库已有唯一同名怪物时更新表格属性并保留外观；同名不唯一、攻击上下限错误或非法整数会阻止。", "可反复修改"),
        ("资源事务", "完整PAK或成对WZL/WZX、目标Monster数据库、登录器pak.txt和MonGen名字颜色在同一事务中备份；失败不提交，可逐字节回滚。", "可回滚"),
        ("M2门禁", "预检和确认写入前必须由用户关闭M2。平台不负责打开、关闭或重启游戏引擎。", "高风险门禁"),
        ("范围边界", "颜色只更新已有MonGen行的第9列；本表不新增刷怪行，不改地图、坐标、范围、数量、间隔、爆率、技能或AI。", "不代替刷怪表"),
    ]

    start = 4
    for offset, values in enumerate(rows):
        row = start + offset
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row, column, value)
            cell.font = _base_font(10, bold=column == 1, color=NAVY if column == 1 else TEXT)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(bottom=LIGHT_BORDER)
            if column == 1:
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            elif column == 3:
                cell.fill = PatternFill("solid", fgColor=GOLD)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 42

    table = Table(displayName="MonsterWorkbookHelp", ref=f"A3:C{start + len(rows) - 1}")
    sheet["A3"] = "项目"
    sheet["B3"] = "填写与执行说明"
    sheet["C3"] = "提示"
    for cell in sheet[3]:
        _set_header(cell)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _add_catalog_sheet(workbook: Workbook, rows: list[tuple[object, ...]]) -> None:
    sheet = workbook.create_sheet("模型参考")
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "70AD47"
    sheet.freeze_panes = "A2"

    headers = ["模型库编号", "外观名称"]
    widths = [16, 34]
    for column, (header, width) in enumerate(zip(headers, widths), start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
        cell = sheet.cell(1, column, header)
        _set_header(cell, model=column == 1)
    sheet.row_dimensions[1].height = 30

    for row_number, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            cell.font = _base_font(10)
            cell.alignment = Alignment(horizontal="left" if column == 2 else "center", vertical="center")
            cell.border = Border(bottom=LIGHT_BORDER)
            cell.number_format = "@" if column == 2 else "#,##0"
            if column == 1:
                cell.fill = PatternFill("solid", fgColor=GOLD)
        sheet.row_dimensions[row_number].height = 22

    if rows:
        table = Table(displayName="MonsterModelCatalog", ref=f"A1:B{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        sheet.auto_filter.ref = f"A1:B{len(rows) + 1}"

    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def build(catalog: Path, output: Path) -> None:
    rows = _catalog_rows(catalog)
    if not rows:
        raise RuntimeError("怪物库没有status=ready的完整模型，不能生成母表")

    workbook = Workbook()
    workbook.properties.creator = "玄渊成果平台"
    workbook.properties.title = "20_怪物批量生成"
    workbook.properties.subject = "翎风/LFM2怪物资料与动作补丁批量生成母表"
    workbook.properties.description = "由玄渊成果平台怪物库生成；支持已有MonGen行名字颜色，不新增地图刷怪。"

    _add_generation_sheet(workbook)
    _add_instruction_sheet(workbook, len(rows))
    _add_catalog_sheet(workbook, rows)

    workbook.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    check = load_workbook(output, data_only=False)
    if check.sheetnames != ["怪物生成", "填写说明", "模型参考"]:
        raise RuntimeError(f"工作表结构错误：{check.sheetnames}")
    headers = [check["怪物生成"].cell(1, column).value for column in range(1, len(HEADERS) + 1)]
    if headers != HEADERS:
        raise RuntimeError("怪物生成表头与平台合同不一致")
    if check["模型参考"].max_row != len(rows) + 1:
        raise RuntimeError("模型参考行数与怪物库不一致")
    for sheet in check.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    raise RuntimeError(f"母表不应包含公式：{sheet.title}!{cell.coordinate}")


def main() -> int:
    parser = argparse.ArgumentParser(description="生成玄渊怪物批量生成母表")
    parser.add_argument("--catalog", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.catalog.resolve(), args.output.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
