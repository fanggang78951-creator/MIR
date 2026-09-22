from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
from pathlib import Path


EQUIPMENT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(EQUIPMENT_ROOT / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from test_equipment_preflight import make_target, tree_hash


def make_hint_contract(target: Path, client_data: Path) -> None:
    from xyequip.item_hint import EFFECT_DEFINITIONS, RESOURCE_LIBRARY, RESOURCE_SLOT

    envir = target / "Mir200" / "Envir"
    (envir / "EffectImageList.txt").write_bytes(
        ("\r\n".join([f"Reserved{i}.wzl" for i in range(RESOURCE_SLOT)] + [RESOURCE_LIBRARY]) + "\r\n").encode("gb18030")
    )
    (envir / "EffectHintBG.txt").write_bytes(
        ("\r\n".join(EFFECT_DEFINITIONS[index] for index in sorted(EFFECT_DEFINITIONS)) + "\r\n").encode("gb18030")
    )
    (envir / "EffectHintBGItems.txt").write_bytes(b"")
    client_data.mkdir(parents=True, exist_ok=True)


class ItemHintContractTests(unittest.TestCase):
    def test_invalid_category_is_rejected_at_compile_time(self):
        from xyequip.batch import BatchEquipmentService, BatchSourceError

        with self.assertRaisesRegex(BatchSourceError, "悬浮分类"):
            BatchEquipmentService(EQUIPMENT_ROOT.parent).compile_hint_rows(
                [{"名称": "错误分类装备", "悬浮分类": "普通"}], Path("hint.xlsx")
            )

    def test_existing_equipment_hint_sync_installs_and_rolls_back_byte_exactly(self):
        from xyequip.batch import BatchEquipmentService
        from xyequip.installer import EquipmentInstaller
        from xyequip.planner import EquipmentPlanner

        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "Server"
            client_data = Path(directory) / "Client" / "data"
            make_target(target)
            make_hint_contract(target, client_data)
            desc = target / "Mir200" / "Envir" / "ItemDescList.txt"
            desc.write_bytes("青铜头盔=基础说明\r\n".encode("gb18030"))
            before_server = tree_hash(target)
            before_client = tree_hash(client_data)
            compiled = BatchEquipmentService(EQUIPMENT_ROOT.parent).compile_hint_rows(
                [{"名称": "青铜头盔", "悬浮分类": "稀有专属"}], Path("hint.xlsx"), "hint-hash"
            )
            plan = EquipmentPlanner(EQUIPMENT_ROOT.parent).preflight(
                target, compiled, client_data, mode="item_hint_sync"
            )

            self.assertEqual(plan.blockers, [])
            receipt = EquipmentInstaller(EQUIPMENT_ROOT.parent).install(plan)
            bindings = (target / "Mir200/Envir/EffectHintBGItems.txt").read_bytes().decode("gb18030")
            description = desc.read_bytes().decode("gb18030")
            self.assertIn("青铜头盔\t1\t2\t4", bindings)
            self.assertIn("\\<Img:27:10:0:0>", description)
            self.assertIn("\\<PlayImg:10:19:8:120:0:-82:1>", description)
            self.assertTrue((client_data / "XY_ItemHint.wzl").is_file())
            self.assertTrue((client_data / "XY_ItemHint.wzx").is_file())

            EquipmentInstaller(EQUIPMENT_ROOT.parent).rollback(target, receipt.transaction_id)
            restored = tree_hash(target)
            restored.pop(".xydp/equipment-installed.json", None)
            self.assertEqual(restored, before_server)
            self.assertEqual(tree_hash(client_data), before_client)

    def test_new_equipment_classification_is_applied_in_same_transaction(self):
        from xyequip.batch import BatchEquipmentService
        from xyequip.installer import EquipmentInstaller
        from xyequip.planner import EquipmentPlanner

        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "Server"
            client_data = Path(directory) / "Client" / "data"
            make_target(target)
            make_hint_contract(target, client_data)
            compiled = BatchEquipmentService(EQUIPMENT_ROOT.parent).compile_rows(
                [{"名称": "同事务悬浮头盔", "部位": "头盔", "防御": "0-1", "悬浮分类": "追梦神器"}],
                Path("items.xlsx"),
                "items-hash",
            )
            plan = EquipmentPlanner(EQUIPMENT_ROOT.parent).preflight(target, compiled, client_data)

            self.assertEqual(plan.blockers, [])
            receipt = EquipmentInstaller(EQUIPMENT_ROOT.parent).install(plan)
            bindings = (target / "Mir200/Envir/EffectHintBGItems.txt").read_bytes().decode("gb18030")
            description = (target / "Mir200/Envir/ItemDescList.txt").read_bytes().decode("gb18030")
            self.assertIn("同事务悬浮头盔\t1\t2\t5", bindings)
            self.assertIn("同事务悬浮头盔=", description)
            self.assertIn("\\<PlayImg:10:19:8:120:0:-82:1>", description)
            self.assertEqual(receipt.equipment_names, ("同事务悬浮头盔",))

    def test_duplicate_binding_blocks_without_writing_target(self):
        from xyequip.batch import BatchEquipmentService
        from xyequip.planner import EquipmentPlanner

        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "Server"
            client_data = Path(directory) / "Client" / "data"
            make_target(target)
            make_hint_contract(target, client_data)
            (target / "Mir200/Envir/ItemDescList.txt").write_bytes("青铜头盔=说明\r\n".encode("gb18030"))
            (target / "Mir200/Envir/EffectHintBGItems.txt").write_bytes(
                "青铜头盔 1 2 3\r\n青铜头盔 1 2 4\r\n".encode("gb18030")
            )
            before = tree_hash(target)
            compiled = BatchEquipmentService(EQUIPMENT_ROOT.parent).compile_hint_rows(
                [{"名称": "青铜头盔", "悬浮分类": "制式装备"}], Path("hint.xlsx")
            )
            plan = EquipmentPlanner(EQUIPMENT_ROOT.parent).preflight(
                target, compiled, client_data, mode="item_hint_sync"
            )

            self.assertTrue(any("重复悬浮绑定" in item for item in plan.blockers))
            self.assertEqual(tree_hash(target), before)

    def test_candidate_resource_hashes_are_stable(self):
        from xyequip.item_hint import item_hint_asset_paths

        wzl, wzx = item_hint_asset_paths(EQUIPMENT_ROOT.parent)
        self.assertEqual(hashlib.sha256(wzl.read_bytes()).hexdigest().upper(), "68A33C07949CF81562D1DE91D17F9F449414FB73D4878050C74340944F2CE3A7")
        self.assertEqual(hashlib.sha256(wzx.read_bytes()).hexdigest().upper(), "EB9075711E346FDACEDCEE667A4CEAEDE16F0B233DBE765142A9E96498295B63")

    def test_export_current_bindings_creates_reusable_workbook(self):
        from openpyxl import load_workbook
        from xyequip.item_hint import export_hint_workbook
        from xyequip.paths import EquipmentPaths

        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "Server"
            make_target(target)
            (target / "Mir200/Envir/EffectHintBGItems.txt").write_bytes(
                "青铜头盔 1 2 4\r\n".encode("gb18030")
            )
            output = Path(directory) / "export.xlsx"
            export_hint_workbook(EquipmentPaths.for_target(EQUIPMENT_ROOT.parent, target), output)

            workbook = load_workbook(output, read_only=True)
            sheet = workbook["装备悬浮分类"]
            values = list(sheet.iter_rows(min_row=2, values_only=True))
            workbook.close()
            self.assertIn(("青铜头盔", "稀有专属", None), values)


if __name__ == "__main__":
    unittest.main()
