from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


EQUIPMENT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(EQUIPMENT_ROOT / "src"))


class XlsxContractTests(unittest.TestCase):
    def test_service_rebinds_legacy_registry_to_platform_root(self):
        from xyequip.batch import BatchEquipmentService
        from xyequip.legacy import xy_equip_maker

        xy_equip_maker.SCRIPT_PROPS = Path("missing-script-properties.json")
        BatchEquipmentService(EQUIPMENT_ROOT.parent)
        self.assertEqual(xy_equip_maker.SCRIPT_PROPS, EQUIPMENT_ROOT / "profiles/script_properties.json")

    def test_master_workbook_is_inspected_without_translation(self):
        from xyequip.batch import BatchEquipmentService

        service = BatchEquipmentService(EQUIPMENT_ROOT.parent)
        inspection = service.inspect(EQUIPMENT_ROOT / "templates/XuanYuanItems.xlsx")
        self.assertGreater(len(inspection.headers), 20)
        self.assertEqual(inspection.row_count, 0)
        self.assertEqual(inspection.source, EQUIPMENT_ROOT / "templates/XuanYuanItems.xlsx")

    def test_compile_preserves_source_row_and_generates_internal_txt(self):
        from xyequip.batch import BatchEquipmentService

        service = BatchEquipmentService(EQUIPMENT_ROOT.parent)
        with tempfile.TemporaryDirectory() as directory:
            fixture = Path(directory) / "09_装备批量生成.xlsx"
            workbook = load_workbook(EQUIPMENT_ROOT / "templates/XuanYuanItems.xlsx")
            sheet = workbook["装备导入表"]
            sheet["A2"] = "XLSX契约头盔"
            sheet["B2"] = "头盔"
            sheet["C2"] = 1
            workbook.save(fixture)
            workbook.close()
            result = service.compile(fixture)
            self.assertEqual(result.rows[0].source_row, 2)
            self.assertIn("[装备]", result.rows[0].equipment_text)
            self.assertIn("DuraMax=60000", result.rows[0].equipment_text)
            self.assertTrue(result.workbook_hash)

    def test_unknown_green_property_is_rejected(self):
        from xyequip.batch import BatchEquipmentService, BatchSourceError

        service = BatchEquipmentService(EQUIPMENT_ROOT.parent)
        row = {"名称": "坏属性装备", "部位": "头盔", "绿字1类型": "从未登记的属性", "绿字1数值": "10"}
        with self.assertRaisesRegex(BatchSourceError, "未识别绿字属性"):
            service.compile_rows([row], Path("memory.xlsx"))

    def test_outputs_are_under_platform_equipment_root(self):
        from xyequip.batch import BatchEquipmentService

        service = BatchEquipmentService(EQUIPMENT_ROOT.parent)
        self.assertEqual(service.output_root, EQUIPMENT_ROOT / "outputs")
        self.assertNotIn("MirServer", str(service.output_root))

    def test_master_workbook_contains_extended_slot_guide(self):
        workbook = EQUIPMENT_ROOT / "templates/XuanYuanItems.xlsx"
        book = load_workbook(workbook, read_only=False, data_only=False)
        values = {
            str(cell.value)
            for sheet in book.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        }

        self.assertIn("部位说明", book.sheetnames)
        self.assertIn("时装武器", values)
        self.assertIn("时装首饰盒6", values)
        self.assertIn("时装生肖12", values)
        self.assertIn("背包神器", values)
        self.assertIn("称号卷", values)
        validations = list(book["装备导入表"].data_validations.dataValidation)
        self.assertTrue(any("部位说明!$A$2:$A$148" in str(rule.formula1) for rule in validations))


if __name__ == "__main__":
    unittest.main()
