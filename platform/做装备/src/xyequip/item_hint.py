from __future__ import annotations

import hashlib
import re
import shutil
import sqlite3
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Iterable

from .legacy import xy_equip_maker as legacy_core

if TYPE_CHECKING:
    from .batch import BatchCompileResult
    from .paths import EquipmentPaths


RESOURCE_LIBRARY = "XY_ItemHint.wzl"
RESOURCE_SLOT = 10
RESOURCE_STEM = "XY_ItemHint"
FINAL_SUFFIX = r"\-\<Img:27:10:0:0>\-\<PlayImg:10:19:8:120:0:-82:1>"
CATEGORY_EFFECT = {
    "制式装备": 3,
    "稀有专属": 4,
    "追梦神器": 5,
}
EFFECT_DEFINITIONS = {
    1: "1,10,0,1,200,0,0,-80,0,1,0,0,1",
    2: "2,10,1,1,200,0,0,-80,0,1,1,1,0",
    3: "3,10,5,1,200,-32,0,-80,0,0,1,1,0",
    4: "4,10,7,6,160,-32,0,0,0,0,1,1,0",
    5: "5,10,13,6,120,-32,0,0,0,0,1,1,0",
}
EXPECTED_ASSET_HASHES = {
    "XY_ItemHint.wzl": "68A33C07949CF81562D1DE91D17F9F449414FB73D4878050C74340944F2CE3A7",
    "XY_ItemHint.wzx": "EB9075711E346FDACEDCEE667A4CEAEDE16F0B233DBE765142A9E96498295B63",
}


class ItemHintError(RuntimeError):
    pass


@dataclass(frozen=True)
class ItemHintSelection:
    source_row: int
    name: str
    category: str

    @property
    def effects(self) -> tuple[int, int, int]:
        return (1, 2, CATEGORY_EFFECT[self.category])


def item_hint_asset_paths(platform_root: Path) -> tuple[Path, Path]:
    root = Path(platform_root).resolve()
    asset_root = root / "做装备" / "assets" / "item_hint" / "candidate.13"
    if not asset_root.is_dir() and getattr(sys, "_MEIPASS", None):
        asset_root = Path(sys._MEIPASS) / "做装备" / "assets" / "item_hint" / "candidate.13"
    return asset_root / "XY_ItemHint.wzl", asset_root / "XY_ItemHint.wzx"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def validate_item_hint_assets(platform_root: Path) -> list[str]:
    blockers: list[str] = []
    for path in item_hint_asset_paths(platform_root):
        if not path.is_file():
            blockers.append(f"装备悬浮资源缺失：{path}")
            continue
        expected = EXPECTED_ASSET_HASHES[path.name]
        actual = _sha256(path)
        if actual != expected:
            blockers.append(f"装备悬浮资源哈希不匹配：{path.name}，期望 {expected}，实际 {actual}")
    return blockers


def selections_from_compiled(compiled: "BatchCompileResult") -> tuple[ItemHintSelection, ...]:
    selected: list[ItemHintSelection] = []
    for row in compiled.rows:
        values = dict(row.source_values)
        category = values.get("悬浮分类", "").strip()
        if not category or category == "不处理":
            continue
        if category not in CATEGORY_EFFECT:
            raise ItemHintError(
                f"第{row.source_row}行 {row.name}：悬浮分类必须是制式装备、稀有专属、追梦神器或不处理"
            )
        selected.append(ItemHintSelection(row.source_row, row.name, category))
    return tuple(selected)


def _read_text(path: Path) -> str:
    return legacy_core.read_text_auto(path) if path.is_file() else ""


def _newline(text: str) -> str:
    return "\r\n" if "\r\n" in text else "\n"


def _write_lines(path: Path, lines: list[str], original: str = "") -> None:
    newline = _newline(original)
    trailing = bool(original.endswith(("\n", "\r"))) or not original
    text = newline.join(lines) + (newline if trailing and lines else "")
    path.parent.mkdir(parents=True, exist_ok=True)
    legacy_core.write_mir_text(path, text)


def _parse_effect_definitions(text: str) -> dict[int, str]:
    result: dict[int, str] = {}
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith((";", "#")):
            continue
        try:
            index = int(line.split(",", 1)[0].strip())
        except (ValueError, IndexError):
            continue
        if index in result:
            raise ItemHintError(f"EffectHintBG.txt 存在重复特效编号：{index}")
        result[index] = line
    return result


def _parse_bindings(text: str) -> tuple[list[str], dict[str, list[tuple[int, int, int]]]]:
    lines = text.splitlines()
    bindings: dict[str, list[tuple[int, int, int]]] = {}
    for raw in lines:
        stripped = raw.strip()
        if not stripped or stripped.startswith((";", "#")):
            continue
        parts = stripped.rsplit(None, 3)
        if len(parts) != 4:
            continue
        name = parts[0].strip()
        try:
            values = (int(parts[1]), int(parts[2]), int(parts[3]))
        except ValueError:
            continue
        bindings.setdefault(name, []).append(values)
    return lines, bindings


def _strip_managed_suffix(value: str) -> str:
    value = value.replace(FINAL_SUFFIX, "")
    value = re.sub(r"(?:\\-)?\\<Img:27:10:[^>]*>", "", value)
    value = re.sub(r"(?:\\-)?\\<PlayImg:10:(?:19|20|21|22|23|24|25|26):[^>]*>", "", value)
    return value.rstrip("\\-")


def _description_matches(text: str, name: str) -> list[int]:
    prefix = f"{name}="
    return [index for index, line in enumerate(text.splitlines()) if line.startswith(prefix)]


def preflight_item_hints(
    paths: "EquipmentPaths",
    selections: Iterable[ItemHintSelection],
    *,
    descriptions_will_be_created: bool = False,
) -> tuple[list[str], list[str]]:
    selections = tuple(selections)
    if not selections:
        return [], []
    blockers = validate_item_hint_assets(paths.platform_root)
    warnings: list[str] = []
    if paths.client_data is None:
        blockers.append("装备悬浮分类需要选择客户端 data 目录，以事务同步 XY_ItemHint.wzl/wzx")
        return blockers, warnings
    if not paths.client_data.is_dir():
        blockers.append(f"客户端 data 目录不存在：{paths.client_data}")
        return blockers, warnings

    for asset in item_hint_asset_paths(paths.platform_root):
        target = paths.client_data / asset.name
        if target.is_file() and _sha256(target) != EXPECTED_ASSET_HASHES[asset.name]:
            blockers.append(f"客户端已存在不同内容的 {asset.name}，禁止覆盖")
        elif not target.exists():
            warnings.append(f"将事务安装客户端悬浮资源：{target}")

    image_list = _read_text(paths.effect_image_list)
    image_lines = image_list.splitlines()
    if len(image_lines) <= RESOURCE_SLOT:
        blockers.append(f"EffectImageList.txt 未登记到资源编号 {RESOURCE_SLOT}，禁止猜测补齐")
    elif image_lines[RESOURCE_SLOT].strip().casefold() != RESOURCE_LIBRARY.casefold():
        blockers.append(
            f"EffectImageList.txt 编号 {RESOURCE_SLOT} 已被占用：{image_lines[RESOURCE_SLOT].strip()}"
        )

    try:
        definitions = _parse_effect_definitions(_read_text(paths.effect_hint_definitions))
        for index, expected in EFFECT_DEFINITIONS.items():
            actual = definitions.get(index)
            if actual is not None and actual.replace(" ", "") != expected:
                blockers.append(f"悬浮特效编号 {index} 与平台最终合同冲突：{actual}")
            elif actual is None:
                warnings.append(f"将创建悬浮特效编号 {index}")
    except ItemHintError as exc:
        blockers.append(str(exc))

    try:
        _lines, bindings = _parse_bindings(_read_text(paths.effect_hint_items))
        desc_text = _read_text(paths.item_desc)
        for selection in selections:
            if len(bindings.get(selection.name, [])) > 1:
                blockers.append(f"重复悬浮绑定：{selection.name}")
            if not descriptions_will_be_created:
                matches = _description_matches(desc_text, selection.name)
                if len(matches) != 1:
                    blockers.append(f"ItemDescList.txt 中 {selection.name} 必须且只能存在一条，实际 {len(matches)} 条")
                else:
                    value = desc_text.splitlines()[matches[0]].split("=", 1)[1]
                    stripped = _strip_managed_suffix(value)
                    if "\\<Img:" in stripped or "\\<PlayImg:" in stripped:
                        blockers.append(f"{selection.name} 已有非平台管理的图片悬浮标签，禁止覆盖")
    except ItemHintError as exc:
        blockers.append(str(exc))
    return blockers, warnings


def apply_item_hints(paths: "EquipmentPaths", selections: Iterable[ItemHintSelection]) -> None:
    selections = tuple(selections)
    if not selections:
        return
    blockers, _warnings = preflight_item_hints(paths, selections)
    if blockers:
        raise ItemHintError("；".join(blockers))

    assert paths.client_data is not None
    for asset in item_hint_asset_paths(paths.platform_root):
        destination = paths.client_data / asset.name
        if not destination.exists():
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(asset, destination)

    definitions_text = _read_text(paths.effect_hint_definitions)
    definition_lines = definitions_text.splitlines()
    definitions = _parse_effect_definitions(definitions_text)
    for index, expected in EFFECT_DEFINITIONS.items():
        if index not in definitions:
            definition_lines.append(expected)
    _write_lines(paths.effect_hint_definitions, definition_lines, definitions_text)

    binding_text = _read_text(paths.effect_hint_items)
    binding_lines, bindings = _parse_bindings(binding_text)
    names = {selection.name for selection in selections}
    kept: list[str] = []
    for raw in binding_lines:
        stripped = raw.strip()
        parts = stripped.rsplit(None, 3)
        if len(parts) == 4 and parts[0].strip() in names:
            continue
        kept.append(raw)
    for selection in selections:
        if len(bindings.get(selection.name, [])) > 1:
            raise ItemHintError(f"重复悬浮绑定：{selection.name}")
        kept.append(f"{selection.name}\t1\t2\t{CATEGORY_EFFECT[selection.category]}")
    _write_lines(paths.effect_hint_items, kept, binding_text)

    desc_text = _read_text(paths.item_desc)
    desc_lines = desc_text.splitlines()
    for selection in selections:
        matches = _description_matches(desc_text, selection.name)
        if len(matches) != 1:
            raise ItemHintError(f"ItemDescList.txt 中 {selection.name} 必须且只能存在一条，实际 {len(matches)} 条")
        index = matches[0]
        key, value = desc_lines[index].split("=", 1)
        desc_lines[index] = f"{key}={_strip_managed_suffix(value)}{FINAL_SUFFIX}"
        desc_text = "\n".join(desc_lines)
    _write_lines(paths.item_desc, desc_lines, _read_text(paths.item_desc))


def export_hint_workbook(paths: "EquipmentPaths", output: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    if not paths.db_path.is_file():
        raise ItemHintError(f"目标服装备数据库不存在：{paths.db_path}")
    connection = sqlite3.connect(f"file:{paths.db_path.as_posix()}?mode=ro", uri=True)
    try:
        names = [str(row[0]) for row in connection.execute("SELECT DISTINCT Name FROM StdItems ORDER BY Name")]
    finally:
        connection.close()
    _lines, bindings = _parse_bindings(_read_text(paths.effect_hint_items))
    reverse = {(1, 2, effect): category for category, effect in CATEGORY_EFFECT.items()}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "装备悬浮分类"
    sheet.append(["名称", "悬浮分类", "备注"])
    for name in names:
        values = bindings.get(name, [])
        category = reverse.get(values[0], "不处理") if len(values) == 1 else "不处理"
        sheet.append([name, category, ""])
    validation = DataValidation(
        type="list", formula1='"制式装备,稀有专属,追梦神器,不处理"', allow_blank=False
    )
    sheet.add_data_validation(validation)
    if sheet.max_row >= 2:
        validation.add(f"B2:B{sheet.max_row}")
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 42
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
