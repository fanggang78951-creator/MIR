from __future__ import annotations

import hashlib
import importlib.util
import io
import os
import re
import shutil
import sqlite3
import stat
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Callable, Iterable

import openpyxl

from .sqlitepatch import SqlitePatchError, apply_sqlite_upsert
from .installer import InstallPlan, Installer, PlannedChange
from .repository import PackageRepository


class NpcBundleError(ValueError):
    pass


def load_necklace_luck_bundle(
    platform_root: Path,
    workbook: Path,
    *,
    x: int,
    y: int,
):
    """Load the one existing necklace-luck generator for bundle compilation.

    The bundle adapter deliberately contains no reinforcement business logic.
    It only supplies the coordinates assigned by the batch preflight and asks
    the established lab generator to parse and compile the workbook.
    """

    tool = Path(platform_root).resolve() / "labs" / "necklace_luck" / "tools" / "deploy_candidate.py"
    if not tool.is_file():
        raise NpcBundleError(f"平台缺少项链幸运唯一生成器：{tool}")
    name = "_xydp_necklace_luck_generator_" + hashlib.sha256(str(tool).encode("utf-8")).hexdigest()[:12]
    module = sys.modules.get(name)
    if module is None:
        spec = importlib.util.spec_from_file_location(name, tool)
        if spec is None or spec.loader is None:
            raise NpcBundleError(f"无法加载项链幸运生成器：{tool}")
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        try:
            spec.loader.exec_module(module)
        except Exception:
            sys.modules.pop(name, None)
            raise
    try:
        base, tiers = module.parse_workbook(
            Path(workbook).resolve(),
            {"坐标X": int(x), "坐标Y": int(y)},
        )
    except Exception as exc:
        raise NpcBundleError(f"项链幸运配置无法由现有生成器读取：{exc}") from exc
    return module, base, tiers


def stage_necklace_luck(
    shadow: "BundleShadowWorkspace",
    platform_root: Path,
    workbook: Path,
    *,
    x: int,
    y: int,
) -> dict[str, object]:
    """Compile the paid formal NPC into the shared sparse shadow target."""

    module, _base, _tiers = load_necklace_luck_bundle(
        platform_root, workbook, x=x, y=y,
    )
    try:
        outputs, result, _already = module.build_outputs(
            shadow.server,
            Path(workbook).resolve(),
            base_overrides={"坐标X": int(x), "坐标Y": int(y)},
            formal_release=True,
        )
    except Exception as exc:
        raise NpcBundleError(f"项链幸运正式候选编译失败：{exc}") from exc

    changes: list[PlannedChange] = []
    for target, after in outputs.items():
        try:
            relative = target.resolve().relative_to(shadow.server.resolve()).as_posix()
        except ValueError as exc:
            raise NpcBundleError(f"项链幸运生成器输出越过服务端：{target}") from exc
        before = target.read_bytes() if target.is_file() else None
        changes.append(PlannedChange(
            relative,
            before,
            after,
            "bundle-necklace-luck",
            "xy.optional.necklace-luck",
            "server",
        ))
    shadow.apply_changes(changes)

    npc_path = shadow.server / module.NPC_REL
    npc = npc_path.read_bytes().decode("gb18030")
    for forbidden in ("测试人物清零", "XY_VERIFY", "XY-TEST-MONITOR"):
        if forbidden in npc:
            raise NpcBundleError(f"项链幸运正式NPC仍含测试内容：{forbidden}")
    if "黄金树芽×1，元宝×2000" not in npc:
        raise NpcBundleError("项链幸运NPC未编译本次付费资源")
    return dict(result)


@dataclass
class BundleShadowWorkspace:
    """Sparse writable overlay used to compose specialist platform plans.

    Existing platform compilers intentionally read the target filesystem.
    The bundle compiler therefore lets every specialist preflight see the
    previous candidate result without ever writing the live server.  Map files
    are linked read-only for coordinate checks and are explicitly forbidden as
    transaction outputs because maps are outside this bundle's scope.
    """

    live_server: Path
    live_client: Path | None
    live_launcher: Path | None
    server: Path
    client: Path | None
    launcher: Path | None
    _originals: dict[tuple[str, str], bytes | None]
    _operations: dict[tuple[str, str], tuple[str, str, str]]

    @classmethod
    def create(
        cls,
        *,
        server: Path,
        client: Path | None,
        launcher: Path | None,
        working_root: Path,
    ) -> "BundleShadowWorkspace":
        live_server = Path(server).resolve()
        live_client = Path(client).resolve() if client is not None else None
        live_launcher = Path(launcher).resolve() if launcher is not None else None
        root = Path(working_root).resolve()
        if root.exists():
            raise NpcBundleError(f"影子目标目录已存在，拒绝覆盖：{root}")
        shadow_server = root / "server"
        shadow_client = root / "client" if live_client is not None else None
        shadow_launcher = root / "launcher" if live_launcher is not None else None
        try:
            envir = live_server / "Mir200" / "Envir"
            database = live_server / "Mud2" / "DB"
            maps = live_server / "Mir200" / "Map"
            required = (live_server / "Mir200" / "M2Server.exe", envir, envir / "MapInfo.txt")
            missing = [str(path) for path in required if not path.exists()]
            if missing:
                raise NpcBundleError("影子目标缺少服务端基础文件：" + "、".join(missing))
            shutil.copytree(envir, shadow_server / "Mir200" / "Envir")
            if database.is_dir():
                shutil.copytree(database, shadow_server / "Mud2" / "DB")
            else:
                (shadow_server / "Mud2" / "DB").mkdir(parents=True)
            shadow_map = shadow_server / "Mir200" / "Map"
            shadow_map.mkdir(parents=True)
            if maps.is_dir():
                for source in maps.rglob("*"):
                    relative = source.relative_to(maps)
                    target = shadow_map / relative
                    if source.is_dir():
                        target.mkdir(parents=True, exist_ok=True)
                    elif source.is_file():
                        target.parent.mkdir(parents=True, exist_ok=True)
                        try:
                            os.link(source, target)
                        except OSError:
                            shutil.copy2(source, target)
            (shadow_server / "Mir200" / "M2Server.exe").write_bytes(b"")
            setup = live_server / "Mir200" / "!setup.txt"
            if setup.is_file():
                shutil.copy2(setup, shadow_server / "Mir200" / "!setup.txt")

            if shadow_client is not None:
                (shadow_client / "data").mkdir(parents=True)
                for name in ("fenghao.dat", "XY_EquipmentCollection.wzl", "XY_EquipmentCollection.wzx"):
                    source = live_client / "data" / name
                    if source.is_file():
                        shutil.copy2(source, shadow_client / "data" / name)
            if shadow_launcher is not None:
                patch_data = shadow_launcher / "补丁文件夹" / "Data"
                patch_data.mkdir(parents=True)
                for name in ("fenghao.dat", "XY_EquipmentCollection.wzl", "XY_EquipmentCollection.wzx"):
                    source = live_launcher / "补丁文件夹" / "Data" / name
                    if source.is_file():
                        shutil.copy2(source, patch_data / name)
        except Exception:
            shutil.rmtree(root, ignore_errors=True)
            raise
        return cls(
            live_server, live_client, live_launcher,
            shadow_server, shadow_client, shadow_launcher,
            {}, {},
        )

    def _roots(self, scope: str, *, live: bool) -> Path:
        roots = {
            "server": self.live_server if live else self.server,
            "client": self.live_client if live else self.client,
            "launcher": self.live_launcher if live else self.launcher,
        }
        root = roots.get(scope)
        if root is None:
            raise NpcBundleError(f"影子目标缺少{scope}作用域")
        return root

    @staticmethod
    def _safe(root: Path, relative: str) -> Path:
        normalized = relative.replace("\\", "/").lstrip("/")
        path = (root / Path(normalized.replace("/", os.sep))).resolve()
        try:
            path.relative_to(root.resolve())
        except ValueError as exc:
            raise NpcBundleError(f"影子目标路径越界：{relative}") from exc
        return path

    def apply_changes(self, changes: Iterable[PlannedChange]) -> None:
        for change in changes:
            relative = change.relative_path.replace("\\", "/")
            if change.scope == "server" and relative.casefold().startswith("mir200/map/"):
                raise NpcBundleError(f"地图文件不在本轮范围：{relative}")
            shadow_path = self._safe(self._roots(change.scope, live=False), relative)
            current = shadow_path.read_bytes() if shadow_path.is_file() else None
            if current != change.before:
                raise NpcBundleError(
                    f"影子候选前置指纹不一致：{change.scope}:{relative}"
                )
            key = (change.scope, relative.casefold())
            if key not in self._originals:
                live_path = self._safe(self._roots(change.scope, live=True), relative)
                self._originals[key] = live_path.read_bytes() if live_path.is_file() else None
            shadow_path.parent.mkdir(parents=True, exist_ok=True)
            shadow_path.write_bytes(change.after)
            self._operations[key] = (change.operation, change.package_id, relative)

    def final_changes(self) -> list[PlannedChange]:
        result: list[PlannedChange] = []
        for (scope, relative_key), before in self._originals.items():
            operation, package_id, relative = self._operations[(scope, relative_key)]
            shadow_path = self._safe(self._roots(scope, live=False), relative)
            after = shadow_path.read_bytes() if shadow_path.is_file() else b""
            if before != after:
                result.append(PlannedChange(relative, before, after, operation, package_id, scope))
        return result


NUCLEUS_LOOKS = {
    "群星之核LV1": 6278,
    "群星之核LV2": 6567,
    "群星之核LV3": 6917,
    "群星之核LV4": 6847,
    "群星之核LV5": 6290,
}


@dataclass(frozen=True)
class TitleDefinition:
    """The business definition used when assigning native title Shape values.

    ``shape`` is the requested number.  It is deliberately excluded from the
    semantic comparison: an existing title with the same name and attributes
    is reused even when the workbook requested a different number.
    """

    name: str
    shape: int
    attack: int
    magic: int
    taoism: int
    hp: int
    mp: int
    critical_rate: int

    @property
    def attributes(self) -> tuple[int, ...]:
        return (
            self.attack,
            self.magic,
            self.taoism,
            self.hp,
            self.mp,
            self.critical_rate,
        )


def allocate_title_shapes(
    existing: Iterable[TitleDefinition],
    requested: Iterable[TitleDefinition],
    *,
    fallback_start: int = 212,
) -> dict[str, int]:
    """Assign title numbers without overwriting unrelated native titles."""

    existing_rows = tuple(existing)
    requested_rows = tuple(requested)
    if not 0 <= fallback_start <= 255:
        raise NpcBundleError("称号候补编号起点必须在0到255之间")
    existing_by_name: dict[str, TitleDefinition] = {}
    occupied: dict[int, str] = {}
    for row in existing_rows:
        if not row.name.strip() or not 0 <= row.shape <= 255:
            raise NpcBundleError("现有称号名称或编号无效")
        if row.name in existing_by_name:
            raise NpcBundleError(f"现有同名称号不唯一：{row.name}")
        if row.shape in occupied:
            raise NpcBundleError(
                f"现有称号编号重复：{row.shape}（{occupied[row.shape]}、{row.name}）"
            )
        existing_by_name[row.name] = row
        occupied[row.shape] = row.name

    names = [row.name for row in requested_rows]
    duplicates = sorted({name for name in names if names.count(name) > 1})
    if duplicates:
        raise NpcBundleError("待安装称号名称重复：" + "、".join(duplicates))

    result: dict[str, int] = {}
    for row in requested_rows:
        if not row.name.strip() or not 0 <= row.shape <= 255:
            raise NpcBundleError(f"待安装称号名称或编号无效：{row.name}")
        current = existing_by_name.get(row.name)
        if current is not None:
            if current.attributes != row.attributes:
                raise NpcBundleError(f"同名称号定义不一致：{row.name}")
            result[row.name] = current.shape
            continue

        if row.shape not in occupied:
            selected = row.shape
        else:
            selected = next(
                (candidate for candidate in range(fallback_start, 256) if candidate not in occupied),
                -1,
            )
            if selected < 0:
                raise NpcBundleError(f"没有可用称号编号：{row.name}")
        occupied[selected] = row.name
        result[row.name] = selected
    return result


def _nucleus_values(name: str, looks: int) -> dict[str, object]:
    values: dict[str, object] = {
        "Name": name,
        "StdMode": 48,
        "Shape": 0,
        "Weight": 0,
        "Anicount": 0,
        "Source": 0,
        "Reserved": 0,
        "Looks": looks,
        "DuraMax": 1000,
        "Need": 0,
        "NeedLevel": 0,
        "Price": 0,
        "Stock": 0,
        "Color": 0,
        "OverLap": 4,
        "Light": 0,
        "Horse": 0,
        "Element": 0,
        "Expand1": 9,
        "Expand2": 0,
    }
    for column in (
        "Ac", "Ac2", "Mac", "Mac2", "Dc", "Dc2", "Mc", "Mc2", "Sc", "Sc2",
        "HP", "MP",
    ):
        values[column] = 0
    return values


def stage_nuclei(database_bytes: bytes) -> tuple[bytes, tuple[str, ...]]:
    """Create the five configured nuclei in a staged SQLite image.

    Existing exact rows are a byte-for-byte no-op; an existing name with a
    different definition is a blocker through the platform SQLite patcher.
    """

    staged = database_bytes
    created: list[str] = []
    for name, looks in NUCLEUS_LOOKS.items():
        before = staged
        try:
            staged = apply_sqlite_upsert(staged, {
                "type": "sqlite_upsert",
                "table": "StdItems",
                "unique_key": ["Name"],
                "conflict_keys": [],
                "allocate": {"Idx": "max_plus_one"},
                "on_conflict": "error",
                "values": _nucleus_values(name, looks),
            })
        except SqlitePatchError as exc:
            raise NpcBundleError(f"群星之核创建预检失败：{exc}") from exc
        if staged != before:
            created.append(name)
    return staged, tuple(created)


def choose_npc_positions(
    *,
    width: int,
    height: int,
    count: int,
    occupied: set[tuple[int, int]],
    is_walkable: Callable[[int, int], bool],
    minimum_distance: int = 3,
    center: tuple[int, int] | None = None,
) -> list[tuple[int, int]]:
    """Deterministically select walkable positions from the map centre out."""

    if width < 3 or height < 3 or count < 0:
        raise NpcBundleError("NPC坐标分配参数无效")
    if minimum_distance < 1:
        raise NpcBundleError("NPC安全间距必须至少为1")
    origin = center or (width // 2, height // 2)
    candidates = [
        (x, y)
        for y in range(1, height - 1)
        for x in range(1, width - 1)
    ]
    candidates.sort(key=lambda point: (
        max(abs(point[0] - origin[0]), abs(point[1] - origin[1])),
        (point[0] - origin[0]) ** 2 + (point[1] - origin[1]) ** 2,
        point[1],
        point[0],
    ))

    chosen: list[tuple[int, int]] = []
    forbidden = set(occupied)
    for point in candidates:
        if not is_walkable(*point):
            continue
        if any(
            max(abs(point[0] - other[0]), abs(point[1] - other[1])) < minimum_distance
            for other in (*forbidden, *chosen)
        ):
            continue
        chosen.append(point)
        if len(chosen) == count:
            return chosen
    raise NpcBundleError(
        f"地图{width}×{height}找不到{count}个满足间距{minimum_distance}的可走空位"
    )


@dataclass(frozen=True)
class BundleDocument:
    relative_path: str
    filename: str
    category: str
    sha256: str
    size: int


@dataclass(frozen=True)
class ExcludedDocument:
    relative_path: str
    reason: str


@dataclass(frozen=True)
class NpcBundleManifest:
    source: str
    source_kind: str
    source_sha256: str
    file_count: int
    xlsx_count: int
    documents: tuple[BundleDocument, ...]
    excluded: tuple[ExcludedDocument, ...]
    category_counts: dict[str, int]
    blockers: tuple[str, ...]

    @property
    def active_xlsx_count(self) -> int:
        return len(self.documents)

    def by_category(self, category: str) -> tuple[BundleDocument, ...]:
        return tuple(item for item in self.documents if item.category == category)


@dataclass(frozen=True)
class NpcBundleContentAudit:
    seal_levels: int
    rebirth_levels: int
    main_title_levels: int
    independent_title_levels: int
    synthesis_recipes: int
    skill_chains: int
    skill_levels: int
    collection_items: int
    collection_groups: int
    referenced_item_names: tuple[str, ...]
    missing_item_names: tuple[str, ...]
    duplicate_item_names: tuple[str, ...]
    blockers: tuple[str, ...]


@dataclass
class NpcBundleBatchPlan:
    source: str
    source_sha256: str
    server: str
    client: str
    launcher: str
    blockers: list[str]
    warnings: list[str]
    changes: list[PlannedChange]
    child_reports: list[dict[str, object]]
    manifest: NpcBundleManifest | None = None
    audit: NpcBundleContentAudit | None = None
    install_plan: InstallPlan | None = None


EXPECTED_CATEGORY_COUNTS = {
    "audit": 1,
    "initial_camp": 6,
    "skill_upgrade": 1,
    "weapon_enchant": 1,
    "global": 8,
    "seal_catchup": 9,
    "independent_title": 6,
    "synthesis": 16,
    "rebirth": 4,
    "main_title": 10,
}


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def workbook_has_usable_dimension(path: Path) -> bool:
    """Return whether every worksheet has explicit dimension metadata.

    A number of supplied V2.1 workbooks contain valid cells but omit the
    optional ``dimension`` node.  openpyxl's read-only reader then reports an
    empty sheet.  Bundle materialisation rewrites candidate copies so legacy
    readers can keep their existing implementation.
    """

    try:
        with zipfile.ZipFile(path) as archive:
            sheet_names = [name for name in archive.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)]
            if not sheet_names:
                return False
            for name in sheet_names:
                data = archive.read(name)
                if b"<dimension" not in data:
                    return False
    except (OSError, zipfile.BadZipFile, KeyError):
        return False
    return True


def _enabled(value: object) -> bool:
    return str(value or "").strip().casefold() in {"是", "1", "true", "yes", "启用", "可安装", "已填"}


def _sheet_rows(path: Path, sheet_name: str, *, header_search: tuple[str, ...]) -> list[dict[str, object]]:
    book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    try:
        if sheet_name not in book.sheetnames:
            raise NpcBundleError(f"{path.name}缺少工作表：{sheet_name}")
        sheet = book[sheet_name]
        values = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    header_index = -1
    headers: list[str] = []
    for index, row in enumerate(values[:20]):
        candidate = [str(value or "").strip() for value in row]
        if all(name in candidate for name in header_search):
            header_index = index
            headers = candidate
            break
    if header_index < 0:
        raise NpcBundleError(f"{path.name}/{sheet_name}未找到表头：{'、'.join(header_search)}")
    result: list[dict[str, object]] = []
    for row in values[header_index + 1:]:
        record = {header: row[column] if column < len(row) else None for column, header in enumerate(headers) if header}
        if any(value not in (None, "") for value in record.values()):
            result.append(record)
    return result


_CURRENCIES = {"金币", "元宝", "灵符", "金刚石"}
_ITEM_NAME_HEADERS = {
    "材料A", "材料B", "材料C", "材料D",
    "材料A名称", "材料B名称", "材料C名称", "材料D名称",
    "装备名称", "当前装备", "下一装备", "下一档装备", "目标装备",
    "产出名称", "物品名称", "奖励物品", "奖励物品名称", "名称",
}
_PACKED_ITEM_HEADERS = {"材料", "单次材料", "所需材料", "材料清单"}


def _split_packed_items(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    result: list[str] = []
    for token in re.split(r"[；;、,，]+", text):
        name = re.split(r"[×xX*]\s*\d", token.strip(), maxsplit=1)[0].strip()
        if name and name not in _CURRENCIES:
            result.append(name)
    return result


def _collect_candidate_item_names(path: Path) -> set[str]:
    """Collect item references from the bundle's user-facing table columns."""

    book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    names: set[str] = set()
    try:
        for sheet in book.worksheets:
            if sheet.title in {"当前服装备", "可选数据", "星辰原版参考", "填写说明", "使用说明"}:
                continue
            rows = list(sheet.iter_rows(values_only=True))
            header_index = None
            headers: list[str] = []
            for index, row in enumerate(rows[:20]):
                candidate = [str(value or "").strip() for value in row]
                if any(value in _ITEM_NAME_HEADERS or value in _PACKED_ITEM_HEADERS for value in candidate):
                    header_index = index
                    headers = candidate
                    break
            if header_index is None:
                continue
            status_column = headers.index("状态") if "状态" in headers else None
            input_type_column = headers.index("输入类型") if "输入类型" in headers else None
            for row in rows[header_index + 1:]:
                if status_column is not None and status_column < len(row):
                    status = str(row[status_column] or "").strip()
                    if status and not _enabled(status):
                        continue
                for column, header in enumerate(headers):
                    if column >= len(row):
                        continue
                    value = str(row[column] or "").strip()
                    if not value:
                        continue
                    if header in _PACKED_ITEM_HEADERS:
                        names.update(_split_packed_items(value))
                    elif header in _ITEM_NAME_HEADERS:
                        if header == "名称" and input_type_column is not None:
                            input_type = str(row[input_type_column] or "").strip()
                            if input_type != "物品":
                                continue
                        if value not in _CURRENCIES:
                            names.add(value)
    finally:
        book.close()
    return names


def _safe_relative(value: str) -> str:
    normalized = value.replace("\\", "/").lstrip("/")
    path = PurePosixPath(normalized)
    if not normalized or path.is_absolute() or any(part in ("", ".", "..") for part in path.parts):
        raise NpcBundleError(f"压缩包包含不安全路径：{value}")
    return path.as_posix()


def _strip_single_root(paths: list[str]) -> list[str]:
    roots = {PurePosixPath(path).parts[0] for path in paths}
    if len(roots) != 1:
        return paths
    return [PurePosixPath(*PurePosixPath(path).parts[1:]).as_posix() for path in paths]


def _classify(relative: str) -> tuple[str | None, str | None]:
    parts = PurePosixPath(relative).parts
    if not parts:
        return None, None
    filename = parts[-1]
    joined = "/".join(parts[:-1])
    if filename == "40_命格NPC规则.xlsx":
        return None, "命格按用户决定暂缓"
    if joined == "00_初始营地" and filename == "04_转生.xlsx":
        return None, "旧营地转生入口保持停用"
    if joined == "00_总控与审计" and filename == "00_NPC经济总账_V2.1.xlsx":
        return "audit", None
    if joined == "00_初始营地":
        return "initial_camp", None
    if joined == "C01_漂流群岛_技能":
        return "skill_upgrade", None
    if filename == "42_武器附魔.xlsx" and joined.startswith("平台暂缺字段_"):
        return "weapon_enchant", None
    if joined == "00_全局系统":
        return "global", None
    if joined.startswith("21_神印_每大陆追赶NPC/"):
        return "seal_catchup", None
    if joined.startswith("22_独立称号/"):
        return "independent_title", None
    if joined.startswith("34_通用物品合成_按NPC分表/"):
        return "synthesis", None
    if joined.startswith("04_转生_按大陆分表/"):
        return "rebirth", None
    if joined.startswith("22_贯穿称号_每大陆一个NPC/"):
        return "main_title", None
    return None, None


class NpcBundleService:
    def __init__(self, platform_root: Path):
        self.root = Path(platform_root).resolve()

    @staticmethod
    def _zip_files(path: Path) -> list[tuple[str, bytes]]:
        result: list[tuple[str, bytes]] = []
        with zipfile.ZipFile(path) as archive:
            raw_names: list[str] = []
            records: list[tuple[zipfile.ZipInfo, str]] = []
            for info in archive.infolist():
                if info.is_dir():
                    continue
                mode = info.external_attr >> 16
                if mode and stat.S_ISLNK(mode):
                    raise NpcBundleError(f"压缩包禁止包含路径链接：{info.filename}")
                safe = _safe_relative(info.filename)
                raw_names.append(safe)
                records.append((info, safe))
            stripped = _strip_single_root(raw_names)
            for (info, _), relative in zip(records, stripped):
                result.append((relative, archive.read(info)))
        return result

    @staticmethod
    def _directory_files(path: Path) -> list[tuple[str, bytes]]:
        result: list[tuple[str, bytes]] = []
        for item in sorted(path.rglob("*"), key=lambda value: value.as_posix().casefold()):
            if item.is_symlink():
                raise NpcBundleError(f"目录禁止包含路径链接：{item}")
            if item.is_file():
                result.append((item.relative_to(path).as_posix(), item.read_bytes()))
        return result

    def inspect_bundle(self, source: Path) -> NpcBundleManifest:
        path = Path(source).resolve()
        if path.is_file() and path.suffix.casefold() == ".zip":
            source_kind = "zip"
            source_hash = _digest(path.read_bytes())
            files = self._zip_files(path)
        elif path.is_dir():
            source_kind = "directory"
            files = self._directory_files(path)
            source_hash = _digest(
                "\n".join(f"{relative}\0{_digest(data)}" for relative, data in files).encode("utf-8")
            )
        else:
            raise NpcBundleError(f"NPC批量包不存在或类型不受支持：{path}")

        seen: set[str] = set()
        documents: list[BundleDocument] = []
        excluded: list[ExcludedDocument] = []
        blockers: list[str] = []
        xlsx_count = 0
        for raw_relative, data in files:
            relative = _safe_relative(raw_relative)
            key = relative.casefold()
            if key in seen:
                blockers.append(f"相对路径重复（忽略大小写）：{relative}")
                continue
            seen.add(key)
            if Path(relative).suffix.casefold() != ".xlsx":
                continue
            xlsx_count += 1
            category, reason = _classify(relative)
            if reason:
                excluded.append(ExcludedDocument(relative, reason))
            elif category:
                documents.append(BundleDocument(
                    relative, PurePosixPath(relative).name, category, _digest(data), len(data)
                ))
            else:
                blockers.append(f"未识别的XLSX：{relative}")

        documents.sort(key=lambda item: item.relative_path.casefold())
        excluded.sort(key=lambda item: item.relative_path.casefold())
        counts = {
            category: sum(1 for item in documents if item.category == category)
            for category in EXPECTED_CATEGORY_COUNTS
        }
        for category, expected in EXPECTED_CATEGORY_COUNTS.items():
            actual = counts[category]
            if actual != expected:
                blockers.append(f"{category}数量应为{expected}，实际{actual}")
        return NpcBundleManifest(
            str(path), source_kind, source_hash, len(files), xlsx_count,
            tuple(documents), tuple(excluded), counts, tuple(blockers),
        )

    def materialize_bundle(self, source: Path, destination: Path) -> Path:
        """Extract a read-only source into a normalized candidate directory."""

        source_path = Path(source).resolve()
        destination_path = Path(destination).resolve()
        if source_path.is_file() and source_path.suffix.casefold() == ".zip":
            files = self._zip_files(source_path)
            stem = source_path.stem
        elif source_path.is_dir():
            files = self._directory_files(source_path)
            stem = source_path.name
        else:
            raise NpcBundleError(f"NPC批量包不存在或类型不受支持：{source_path}")
        root = destination_path / f"{stem}-normalized"
        if root.exists():
            raise NpcBundleError(f"候选规范化目录已存在，拒绝覆盖：{root}")
        root.mkdir(parents=True)
        try:
            for relative, data in files:
                safe = _safe_relative(relative)
                target = root.joinpath(*PurePosixPath(safe).parts)
                target.parent.mkdir(parents=True, exist_ok=True)
                if target.suffix.casefold() != ".xlsx":
                    target.write_bytes(data)
                    continue
                try:
                    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=False)
                    if workbook.calculation is not None:
                        workbook.calculation.fullCalcOnLoad = True
                        workbook.calculation.forceFullCalc = True
                    workbook.save(target)
                    workbook.close()
                except Exception as exc:  # openpyxl raises several format-specific errors
                    raise NpcBundleError(f"规范化工作簿失败：{safe}：{exc}") from exc
                if not workbook_has_usable_dimension(target):
                    raise NpcBundleError(f"规范化后工作簿仍缺少dimension：{safe}")
        except Exception:
            shutil.rmtree(root, ignore_errors=True)
            raise
        return root

    @staticmethod
    def _database(server: Path) -> Path:
        for candidate in (
            Path(server) / "Mud2" / "DB" / "ApexM2.DB",
            Path(server) / "Mud2" / "DB" / "StdItems.DB",
        ):
            if candidate.is_file():
                return candidate
        raise NpcBundleError("目标服缺少Mud2/DB/ApexM2.DB（或兼容StdItems.DB）")

    def audit_content(self, source: Path, server: Path) -> NpcBundleContentAudit:
        """Audit all filled business chains without writing source or target."""

        import tempfile

        manifest = self.inspect_bundle(source)
        blockers = list(manifest.blockers)
        with tempfile.TemporaryDirectory(prefix="xy-npc-bundle-audit-") as temp:
            root = self.materialize_bundle(source, Path(temp))
            by_relative = {item.relative_path: root.joinpath(*PurePosixPath(item.relative_path).parts) for item in manifest.documents}

            global_seal = next(
                path for relative, path in by_relative.items()
                if relative == "00_全局系统/21_神印基础属性.xlsx"
            )
            seal_rows = _sheet_rows(global_seal, "配置", header_search=("档位/等级", "状态"))
            seal_levels = sum(1 for row in seal_rows if _enabled(row.get("状态")))

            rebirth_levels = 0
            for document in manifest.by_category("rebirth"):
                rows = _sheet_rows(by_relative[document.relative_path], "配置", header_search=("档位/等级", "状态"))
                rebirth_levels += sum(1 for row in rows if row.get("档位/等级") not in (None, ""))

            main_title_levels = 0
            for document in manifest.by_category("main_title"):
                rows = _sheet_rows(by_relative[document.relative_path], "配置", header_search=("档位/等级", "状态"))
                main_title_levels += sum(1 for row in rows if row.get("档位/等级") not in (None, ""))

            independent_title_levels = 0
            for document in manifest.by_category("independent_title"):
                rows = _sheet_rows(by_relative[document.relative_path], "配置", header_search=("档位/等级", "状态"))
                independent_title_levels += sum(1 for row in rows if row.get("档位/等级") not in (None, ""))

            synthesis_recipes = 0
            for document in manifest.by_category("synthesis"):
                rows = _sheet_rows(by_relative[document.relative_path], "合成配方", header_search=("配方ID", "状态"))
                synthesis_recipes += sum(1 for row in rows if _enabled(row.get("状态")))

            skill_document = manifest.by_category("skill_upgrade")[0]
            skill_path = by_relative[skill_document.relative_path]
            skill_rules = _sheet_rows(skill_path, "强化规则", header_search=("启用", "原技能名称"))
            skill_costs = _sheet_rows(skill_path, "每级消耗", header_search=("原技能名称", "目标强化等级"))
            active_skills = {str(row.get("原技能名称") or "").strip() for row in skill_rules if _enabled(row.get("启用"))}
            skill_levels = sum(1 for row in skill_costs if str(row.get("原技能名称") or "").strip() in active_skills)

            collection_path = next(
                path for relative, path in by_relative.items()
                if relative == "00_全局系统/33_装备收集图鉴.xlsx"
            )
            collection_rows = _sheet_rows(collection_path, "收集明细", header_search=("状态", "收集ID", "装备名称"))
            group_rows = _sheet_rows(collection_path, "分类奖励", header_search=("状态", "分类ID"))
            collection_items = sum(1 for row in collection_rows if _enabled(row.get("状态")))
            collection_groups = sum(1 for row in group_rows if _enabled(row.get("状态")))

            referenced_names: set[str] = set()
            for path in by_relative.values():
                referenced_names.update(_collect_candidate_item_names(path))
            referenced_names -= _CURRENCIES

        try:
            connection = sqlite3.connect(f"file:{self._database(server).as_posix()}?mode=ro", uri=True)
            try:
                rows = connection.execute("SELECT Name, COUNT(*) FROM StdItems GROUP BY Name").fetchall()
            finally:
                connection.close()
        except sqlite3.Error as exc:
            raise NpcBundleError(f"无法只读检查目标服StdItems：{exc}") from exc
        counts = {str(name): int(count) for name, count in rows}
        missing = tuple(sorted(name for name in referenced_names if counts.get(name, 0) == 0))
        duplicates = tuple(sorted(name for name in referenced_names if counts.get(name, 0) > 1))
        expected = {
            "seal_levels": (seal_levels, 255),
            "rebirth_levels": (rebirth_levels, 20),
            "main_title_levels": (main_title_levels, 100),
            "independent_title_levels": (independent_title_levels, 12),
            "synthesis_recipes": (synthesis_recipes, 38),
            "skill_chains": (len(active_skills), 5),
            "skill_levels": (skill_levels, 45),
            "collection_items": (collection_items, 619),
            "collection_groups": (collection_groups, 46),
        }
        for name, (actual, wanted) in expected.items():
            if actual != wanted:
                blockers.append(f"{name}应为{wanted}，实际{actual}")
        if duplicates:
            blockers.append("目标服物品名称不唯一：" + "、".join(duplicates))
        allowed_missing = {f"群星之核LV{level}" for level in range(1, 6)}
        unexpected_missing = sorted(set(missing) - allowed_missing)
        if unexpected_missing:
            blockers.append("目标服缺少批量包引用物品：" + "、".join(unexpected_missing))
        return NpcBundleContentAudit(
            seal_levels=seal_levels,
            rebirth_levels=rebirth_levels,
            main_title_levels=main_title_levels,
            independent_title_levels=independent_title_levels,
            synthesis_recipes=synthesis_recipes,
            skill_chains=len(active_skills),
            skill_levels=skill_levels,
            collection_items=collection_items,
            collection_groups=collection_groups,
            referenced_item_names=tuple(sorted(referenced_names)),
            missing_item_names=missing,
            duplicate_item_names=duplicates,
            blockers=tuple(blockers),
        )

    @staticmethod
    def _preserved_rage_npc(server: Path) -> dict[str, object]:
        merchant = Path(server) / "Mir200" / "Envir" / "MerChant.txt"
        text = merchant.read_bytes().decode("gb18030")
        rows = [
            line.split("\t") for line in text.splitlines()
            if line.split("\t", 1)[0].casefold() == "玄渊运营/狂暴之力".casefold()
        ]
        if len(rows) != 1 or len(rows[0]) < 7:
            raise NpcBundleError(f"当前已验收狂暴NPC注册应唯一，实际{len(rows)}条")
        fields = rows[0]
        script = (
            Path(server) / "Mir200" / "Envir" / "Market_Def" /
            f"{fields[0]}-{fields[1]}.txt"
        )
        if not script.is_file():
            raise NpcBundleError(f"当前已验收狂暴NPC脚本不存在：{script}")
        body = script.read_bytes().decode("gb18030")
        call = "#CALL [\\玄渊功能\\狂暴\\狂暴NPC接口.txt] @XY_RAGE_NPC_MAIN"
        if body.count(call) != 1:
            raise NpcBundleError("当前狂暴NPC未唯一调用已验收狂暴接口，禁止用旧表覆盖")
        return {
            "state": "preserved-game-accepted",
            "script": fields[0], "map": fields[1],
            "x": int(fields[2]), "y": int(fields[3]), "name": fields[4],
        }

    def preflight_bundle(
        self,
        source: Path,
        server: Path,
        *,
        client: Path,
        launcher: Path,
    ) -> NpcBundleBatchPlan:
        """Compile every in-scope V2.1 workbook into one target transaction."""

        source = Path(source).resolve()
        server = Path(server).resolve()
        client = Path(client).resolve()
        launcher = Path(launcher).resolve()
        plan = NpcBundleBatchPlan(
            str(source), "", str(server), str(client), str(launcher), [], [], [], [],
        )
        work_root = self.root / "work"
        work_root.mkdir(parents=True, exist_ok=True)
        try:
            manifest = self.inspect_bundle(source)
            plan.manifest = manifest
            plan.source_sha256 = manifest.source_sha256
            plan.blockers.extend(manifest.blockers)
            audit = self.audit_content(source, server)
            plan.audit = audit
            plan.blockers.extend(audit.blockers)
            if plan.blockers:
                return plan

            package_ids: list[str] = ["xy.bundle.elden-npc-v21"]
            package_versions: dict[str, str] = {"xy.bundle.elden-npc-v21": "2.1.0-candidate.1"}
            warnings: list[str] = []

            with tempfile.TemporaryDirectory(prefix="elden-npc-v21-", dir=work_root) as temp_text:
                temp = Path(temp_text)
                bundle_root = self.materialize_bundle(source, temp / "input")
                shadow = BundleShadowWorkspace.create(
                    server=server, client=client, launcher=launcher,
                    working_root=temp / "shadow",
                )

                def apply_child(name: str, child) -> None:
                    blockers = list(getattr(child, "blockers", []))
                    report = {
                        "name": name,
                        "blockers": blockers,
                        "warnings": list(getattr(child, "warnings", [])),
                        "changes": len(getattr(child, "changes", [])),
                    }
                    for field_name in (
                        "npcs", "title_mapping", "synthesis_recipes",
                        "service_costs", "soul_tasks",
                    ):
                        records = getattr(child, field_name, None)
                        if not records:
                            continue
                        report[field_name] = [
                            dict(item) if isinstance(item, dict) else dict(vars(item))
                            for item in records
                        ]
                    plan.child_reports.append(report)
                    if blockers:
                        raise NpcBundleError(name + "：" + "；".join(blockers))
                    shadow.apply_changes(child.changes)
                    warnings.extend(report["warnings"])
                    inner = getattr(child, "install_plan", None)
                    if inner is not None:
                        for package_id in inner.package_ids:
                            if package_id not in package_ids:
                                package_ids.append(package_id)
                        package_versions.update(inner.package_versions)

                database_relative = "Mud2/DB/ApexM2.DB"
                database_path = shadow.server / Path(database_relative.replace("/", os.sep))
                database_before = database_path.read_bytes()
                database_after, created_nuclei = stage_nuclei(database_before)
                shadow.apply_changes([
                    PlannedChange(
                        database_relative, database_before, database_after,
                        "create-bundle-nuclei", "xy.bundle.elden-npc-v21", "server",
                    )
                ])
                plan.child_reports.append({
                    "name": "five-nuclei", "blockers": [], "warnings": [],
                    "changes": int(database_before != database_after),
                    "created": list(created_nuclei),
                })

                # The accepted rage core contains a later event-hook repair than
                # the supplied historical workbook. Preserve and audit it rather
                # than using the old workbook to overwrite that accepted logic.
                plan.child_reports.append({"name": "rage", **self._preserved_rage_npc(shadow.server)})

                from .initial_camp import InitialCampService
                camp = bundle_root / "00_初始营地"
                camp_overrides = {path.name: path for path in camp.glob("*.xlsx")}
                apply_child(
                    "initial-camp",
                    InitialCampService(self.root).preflight_selected(
                        shadow.server, camp,
                        {"donate", "sponsor", "horse", "relic", "gift"},
                        shadow.client,
                        operation="npc-bundle-v21",
                        material_overrides=camp_overrides,
                    ),
                )

                from .skill_upgrade import SkillUpgradeService
                skill = next(bundle_root.rglob("12_战士技能强化.xlsx"))
                apply_child(
                    "skill-upgrade",
                    SkillUpgradeService(self.root).preflight(
                        shadow.server, skill, client=shadow.client, launcher=shadow.launcher,
                    ),
                )

                from .growth_stage import GrowthStageService
                growth = [
                    bundle_root / "00_全局系统" / "23_宁姆格福_圣律之剑.xlsx",
                    bundle_root / "00_全局系统" / "24_宁姆格福_黄金圣物.xlsx",
                ]
                apply_child("growth-stage", GrowthStageService(self.root).preflight(shadow.server, growth))

                from .elden_bundle_features import EldenBundleFeatureService
                feature_plan = EldenBundleFeatureService(self.root).preflight_progression(
                    shadow.server, bundle_root, client=shadow.client, launcher=shadow.launcher,
                )
                apply_child("elden-progression-and-npcs", feature_plan)

                merchant_text = (
                    shadow.server / "Mir200" / "Envir" / "MerChant.txt"
                ).read_bytes().decode("gb18030")
                necklace_rows = [
                    row.split("\t") for row in merchant_text.splitlines()
                    if row.split("\t", 1)[0].casefold() == "玄渊实验室/项链幸运".casefold()
                ]
                if len(necklace_rows) == 1 and len(necklace_rows[0]) >= 4:
                    neck_x, neck_y = int(necklace_rows[0][2]), int(necklace_rows[0][3])
                elif not necklace_rows:
                    neck_x, neck_y = 102, 75
                else:
                    raise NpcBundleError(f"项链幸运NPC注册应至多一条，实际{len(necklace_rows)}条")
                neck_result = stage_necklace_luck(
                    shadow, self.root,
                    bundle_root / "00_全局系统" / "17_项链幸运强化.xlsx",
                    x=neck_x, y=neck_y,
                )
                package_ids.append("xy.optional.necklace-luck")
                package_versions["xy.optional.necklace-luck"] = "1.0.0-candidate.25-bundle"
                plan.child_reports.append({
                    "name": "necklace-luck", "blockers": [],
                    "warnings": list(neck_result.get("warnings", [])),
                    "changes": len(neck_result.get("changes", [])),
                    "npc": {"map": "XY_NMGF_MAIN", "x": neck_x, "y": neck_y},
                })

                from .equipment_collection import EquipmentCollectionService
                apply_child(
                    "equipment-collection",
                    EquipmentCollectionService(self.root).preflight(
                        bundle_root / "00_全局系统" / "33_装备收集图鉴.xlsx",
                        shadow.server, shadow.client,
                    ),
                )
                # The specialist module owns the client resources. The bundle
                # additionally mirrors them to the selected launcher patch root.
                collection_mirrors: list[PlannedChange] = []
                for name in ("XY_EquipmentCollection.wzl", "XY_EquipmentCollection.wzx"):
                    source_file = shadow.client / "data" / name
                    if not source_file.is_file():
                        raise NpcBundleError(f"图鉴客户端资源未生成：{source_file}")
                    relative = f"补丁文件夹/Data/{name}"
                    target = shadow.launcher / Path(relative.replace("/", os.sep))
                    collection_mirrors.append(PlannedChange(
                        relative, target.read_bytes() if target.is_file() else None,
                        source_file.read_bytes(), "mirror-collection-resource",
                        "xy.optional.equipment-collection", "launcher",
                    ))
                shadow.apply_changes(collection_mirrors)

                from .recycle_config import RecycleConfigService
                apply_child(
                    "equipment-recycle",
                    RecycleConfigService(self.root).preflight(
                        bundle_root / "00_全局系统" / "19_装备回收配置.xlsx",
                        shadow.server,
                    ),
                )

                changes = shadow.final_changes()
                if any(
                    item.scope == "server" and item.relative_path.replace("\\", "/").casefold().startswith("mir200/map/")
                    for item in changes
                ):
                    raise NpcBundleError("批量NPC候选意外包含地图文件，已阻止")
                qfunction = (
                    shadow.server / "Mir200" / "Envir" / "Market_Def" / "QFunction-0.txt"
                ).read_bytes().decode("gb18030")
                qmanage = (
                    shadow.server / "Mir200" / "Envir" / "MapQuest_Def" / "QManage.txt"
                ).read_bytes().decode("gb18030")
                from .textpatch import scan_labels
                for label_name, text in (("QFunction", qfunction), ("QManage", qmanage)):
                    duplicates = scan_labels(text).duplicates
                    if duplicates:
                        raise NpcBundleError(f"{label_name}最终候选存在重复标签：{duplicates}")
                for change in changes:
                    relative_key = change.relative_path.replace("\\", "/").casefold()
                    if (
                        change.scope != "server"
                        or not relative_key.startswith("mir200/envir/")
                        or not relative_key.endswith(".txt")
                    ):
                        continue
                    if b"\n" in change.after.replace(b"\r\n", b""):
                        raise NpcBundleError(f"最终候选出现非CRLF脚本：{change.relative_path}")
                    change.after.decode("gb18030")

                plan.changes = changes
                warnings.extend([
                    "狂暴保留当前已验收接口脚本，未用压缩包中的历史表覆盖。",
                    "命格、独立材料回收、地图、怪物和地图素材均未纳入本事务。",
                    "安装只代表已部署待游戏实测；不会更新current-release.json。",
                ])
                plan.warnings = list(dict.fromkeys(warnings))
                plan.install_plan = InstallPlan(
                    target_root=str(server), client_root=str(client), launcher_root=str(launcher),
                    package_ids=list(dict.fromkeys(package_ids)),
                    package_versions=package_versions,
                    parameters={
                        "source": str(source), "source_sha256": manifest.source_sha256,
                        "active_workbooks": manifest.active_xlsx_count,
                        "excluded": [item.__dict__ for item in manifest.excluded],
                        "child_reports": plan.child_reports,
                    },
                    changes=changes,
                    warnings=plan.warnings,
                    operation_type="npc-bundle-v21",
                    candidate_packages=list(dict.fromkeys(package_ids)),
                )
            current = self.inspect_bundle(source)
            if current.source_sha256 != plan.source_sha256:
                raise NpcBundleError("批量包在预检期间发生变化，请重新预检")
        except (OSError, UnicodeError, ValueError, RuntimeError, sqlite3.Error, zipfile.BadZipFile) as exc:
            plan.blockers.append(str(exc))
            plan.install_plan = None
            plan.changes = []
        return plan

    def install_bundle(self, plan: NpcBundleBatchPlan):
        if plan.blockers or plan.install_plan is None:
            raise NpcBundleError("NPC批量包安装被阻止：\n" + "\n".join(plan.blockers))
        current = self.inspect_bundle(Path(plan.source))
        if current.source_sha256 != plan.source_sha256:
            raise NpcBundleError("NPC批量包在预检后发生变化，请重新预检")
        repository = PackageRepository(self.root / "packages")
        repository.refresh()
        return Installer(repository, self.root / "backups").install(plan.install_plan)

    def rollback_bundle(self, server: Path, transaction_id: str) -> None:
        repository = PackageRepository(self.root / "packages")
        repository.refresh()
        Installer(repository, self.root / "backups").rollback(Path(server), transaction_id)


__all__ = [
    "BundleDocument", "BundleShadowWorkspace", "ExcludedDocument", "NpcBundleError",
    "NUCLEUS_LOOKS", "NpcBundleManifest", "NpcBundleContentAudit", "NpcBundleBatchPlan", "NpcBundleService",
    "TitleDefinition", "allocate_title_shapes", "choose_npc_positions", "stage_nuclei",
    "load_necklace_luck_bundle", "stage_necklace_luck", "workbook_has_usable_dimension",
]
