from __future__ import annotations

import hashlib
import re
import sqlite3
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from .encoding import TextDocument, encode_text_document, read_text_document
from .installer import InstallPlan, Installer, PlannedChange
from .repository import PackageRepository
from .sqlitepatch import SqlitePatchError, apply_sqlite_magic_skill_update
from .textpatch import TextPatchError, ensure_event_label, install_event_hook, install_managed_block


PACKAGE_ID = "xy.optional.skill-upgrade.warrior.bundle"
PACKAGE_VERSION = "2.1.0-candidate.1"
QFUNCTION = "Mir200/Envir/Market_Def/QFunction-0.txt"
QMANAGE = "Mir200/Envir/MapQuest_Def/QManage.txt"
BACKFILL = "Mir200/Envir/QuestDiary/玄渊实验室/技能强化/登录补写.txt"
DATABASE = "Mud2/DB/ApexM2.DB"
CLIENT_DESC = "data/SkillUpgradeDesc.Dat"
LAUNCHER_DESC = "补丁文件夹/Data/SkillUpgradeDesc.Dat"
LEGACY_PACKAGE = "xy.lab.warrior-skill-upgrade"
YES = {"是", "yes", "true", "1", "启用"}


class SkillUpgradeError(ValueError):
    pass


@dataclass(frozen=True)
class Skill:
    mag_id: int
    name: str
    job: int
    max_train_level: int
    can_upgrade: int
    max_upgrade_level: int


@dataclass(frozen=True)
class Rule:
    source: str
    unlock_level: int
    target: str


@dataclass(frozen=True)
class Cost:
    source: str
    level: int
    material_a: str
    material_a_count: int
    material_b: str
    material_b_count: int
    currency_a: str
    currency_a_count: int
    currency_b: str
    currency_b_count: int


@dataclass
class SkillUpgradePlan:
    server: str
    workbook: str
    client: str | None = None
    launcher: str | None = None
    rules: list[dict[str, object]] = field(default_factory=list)
    blockers: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    changes: list[PlannedChange] = field(default_factory=list)
    install_plan: InstallPlan | None = None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _positive(value: Any, label: str) -> int:
    try:
        result = int(float(value))
    except (TypeError, ValueError) as exc:
        raise SkillUpgradeError(f"{label}必须填写正整数") from exc
    if result <= 0:
        raise SkillUpgradeError(f"{label}必须大于0")
    return result


def _sheet_records(sheet: openpyxl.worksheet.worksheet.Worksheet, markers: set[str]) -> list[dict[str, object]]:
    rows = list(sheet.iter_rows(values_only=True))
    header_index = -1
    headers: list[str] = []
    for index, row in enumerate(rows[:30]):
        candidate = [_text(value) for value in row]
        if markers <= set(candidate):
            header_index = index
            headers = candidate
            break
    if header_index < 0:
        raise SkillUpgradeError(f"{sheet.title}未找到表头：{'、'.join(sorted(markers))}")
    result: list[dict[str, object]] = []
    for row in rows[header_index + 1:]:
        record = {name: row[column] if column < len(row) else None for column, name in enumerate(headers) if name}
        if any(value not in (None, "") for value in record.values()):
            result.append(record)
    return result


def load_workbook(path: Path) -> tuple[list[Rule], list[Cost]]:
    source = Path(path).resolve()
    if not source.is_file():
        raise SkillUpgradeError(f"战士技能强化表不存在：{source}")
    book = openpyxl.load_workbook(source, read_only=False, data_only=True)
    try:
        for name in ("强化规则", "每级消耗"):
            if name not in book.sheetnames:
                raise SkillUpgradeError(f"战士技能强化表缺少工作表：{name}")
        rule_rows = _sheet_records(
            book["强化规则"], {"启用", "原技能名称", "解锁强化等级", "解锁技能名称"}
        )
        cost_rows = _sheet_records(
            book["每级消耗"], {"原技能名称", "目标强化等级", "材料A名称", "货币A类型"}
        )
    finally:
        book.close()

    rules: list[Rule] = []
    for number, row in enumerate(rule_rows, 1):
        if _text(row.get("启用")).casefold() not in YES:
            continue
        source_name = _text(row.get("原技能名称"))
        target_name = _text(row.get("解锁技能名称"))
        level = _positive(row.get("解锁强化等级"), f"强化规则第{number}行解锁强化等级")
        if level > 9 or not source_name or not target_name or source_name == target_name:
            raise SkillUpgradeError(f"强化规则第{number}行技能名称或解锁等级无效")
        rules.append(Rule(source_name, level, target_name))
    if len(rules) != 5 or len({rule.source.casefold() for rule in rules}) != len(rules):
        raise SkillUpgradeError("本批量包必须启用5条且原技能名称唯一的强化链")

    active = {rule.source for rule in rules}
    costs: list[Cost] = []
    for number, row in enumerate(cost_rows, 1):
        source_name = _text(row.get("原技能名称"))
        if source_name not in active:
            continue
        level = _positive(row.get("目标强化等级"), f"每级消耗第{number}行目标强化等级")
        if level > 9:
            raise SkillUpgradeError(f"每级消耗第{number}行目标强化等级只能为1至9")
        values = {
            "material_a": _text(row.get("材料A名称")),
            "material_b": _text(row.get("材料B名称")),
            "currency_a": _text(row.get("货币A类型")),
            "currency_b": _text(row.get("货币B类型")),
        }
        if not values["material_a"] and not values["material_b"] and not values["currency_a"] and not values["currency_b"]:
            raise SkillUpgradeError(f"每级消耗第{number}行至少填写一组材料或货币")
        if values["material_a"] and values["material_b"] and values["material_a"] == values["material_b"]:
            raise SkillUpgradeError(f"每级消耗第{number}行两种材料不能相同")
        currencies = {"金币", "元宝", "账户金刚石", "声望点"}
        if any(value and value not in currencies for value in (values["currency_a"], values["currency_b"])):
            raise SkillUpgradeError(f"每级消耗第{number}行包含不支持的原生货币")
        if values["currency_a"] and values["currency_b"] and values["currency_a"] == values["currency_b"]:
            raise SkillUpgradeError(f"每级消耗第{number}行两种货币不能相同")
        def optional_count(name: str, field: str) -> int:
            raw = row.get(field)
            if not name:
                if raw not in (None, "", 0, "0"):
                    raise SkillUpgradeError(f"每级消耗第{number}行{field}有数量但名称为空")
                return 0
            return _positive(raw, f"每级消耗第{number}行{field}")
        costs.append(Cost(
            source_name, level,
            values["material_a"], optional_count(values["material_a"], "材料A数量"),
            values["material_b"], optional_count(values["material_b"], "材料B数量"),
            values["currency_a"], optional_count(values["currency_a"], "货币A数量"),
            values["currency_b"], optional_count(values["currency_b"], "货币B数量"),
        ))
    grouped: dict[str, list[int]] = defaultdict(list)
    for cost in costs:
        grouped[cost.source].append(cost.level)
    for source_name in active:
        if sorted(grouped[source_name]) != list(range(1, 10)):
            raise SkillUpgradeError(f"{source_name}必须且只能填写强化1至9的九行消耗")
    return rules, costs


def _currency_check(name: str, count: int) -> str:
    return {
        "金币": f"CHECKGOLD {count}",
        "元宝": f"CHECKGAMEGOLD > {count - 1}",
        "账户金刚石": f"CHECKGAMEDIAMOND {count}",
        "声望点": f"CHECKCREDITPOINT > {count - 1}",
    }[name]


def _currency_take(name: str, count: int) -> str:
    return {
        "金币": f"GOLDCOUNT - {count}",
        "元宝": f"GAMEGOLD - {count}",
        "账户金刚石": f"GAMEDIAMOND - {count}",
        "声望点": f"CREDITPOINT - {count}",
    }[name]


def compile_qfunction(rules: list[Rule], costs: list[Cost], selected: dict[str, Skill]) -> str:
    by_cost = {(cost.source, cost.level): cost for cost in costs}
    lines: list[str] = []
    for rule in rules:
        mag_id = selected[rule.source].mag_id
        lines += [
            f"[@SkillLevelEx{mag_id}]", "#IF", f"CHECKSKILL {rule.source} = 3 0", "#ACT",
            f"GOTO @XY_WSU_{mag_id}_DISPATCH", "BREAK", "#ELSEACT",
            f"MESSAGEBOX {rule.source}必须先达到普通3级，才能进行强化。", "BREAK", "",
            f"[@XY_WSU_{mag_id}_DISPATCH]", "#IF", f"CHECKSKILL {rule.source} > 8 1", "#ACT",
            f"MESSAGEBOX {rule.source}已经达到强化9。", "BREAK", "",
        ]
        for level in range(1, 10):
            cost = by_cost[(rule.source, level)]
            summaries = [
                f"{name}×{count}" for name, count in (
                    (cost.material_a, cost.material_a_count), (cost.material_b, cost.material_b_count),
                    (cost.currency_a, cost.currency_a_count), (cost.currency_b, cost.currency_b_count),
                ) if name
            ]
            lines += [
                "#IF", f"CHECKSKILL {rule.source} = {level - 1} 1", "#ACT",
                f"MESSAGEBOX 强化{level}需要：{'、'.join(summaries)}。是否继续？ @XY_WSU_{mag_id}_L{level}_CONFIRM @取消",
                "BREAK", "",
            ]
        lines += ["MESSAGEBOX 当前强化等级异常，已停止操作。", "BREAK", ""]
        for level in range(1, 10):
            cost = by_cost[(rule.source, level)]
            checks = []
            actions = []
            for name, count in ((cost.material_a, cost.material_a_count), (cost.material_b, cost.material_b_count)):
                if name:
                    checks.append(f"CHECKITEM {name} {count}")
                    actions.append(f"TAKE {name} {count}")
            for name, count in ((cost.currency_a, cost.currency_a_count), (cost.currency_b, cost.currency_b_count)):
                if name:
                    checks.append(_currency_check(name, count))
                    actions.append(_currency_take(name, count))
            lines += [
                f"[@XY_WSU_{mag_id}_L{level}_CONFIRM]", "#IF",
                f"CHECKSKILL {rule.source} = 3 0", f"CHECKSKILL {rule.source} = {level - 1} 1",
                *checks, "#ACT", *actions,
                f"SKILLLEVEL {rule.source} = {level} 1",
            ]
            if level == rule.unlock_level:
                lines += [f"GOTO @XY_WSU_{mag_id}_UNLOCK", "BREAK"]
            else:
                lines += [f"SENDMSG 6 {rule.source}已强化至强化{level}。", "BREAK"]
            lines += ["#ELSEACT", "MESSAGEBOX 强化所需的物品或货币不足，请备齐后再来。", "BREAK", ""]
        lines += [
            f"[@XY_WSU_{mag_id}_UNLOCK]", "#IF", f"NOT CheckMagicName {rule.target}", "#ACT",
            f"ADDSKILL {rule.target} 3", f"SENDMSG 6 {rule.source}达到强化{rule.unlock_level}，已领悟3级{rule.target}。", "BREAK", "",
            "#IF", f"CheckMagicName {rule.target}", f"CHECKSKILL {rule.target} < 3 0", "#ACT",
            f"SKILLLEVEL {rule.target} = 3 0", f"SENDMSG 6 {rule.target}已补至普通3级。", "BREAK", "",
            f"SENDMSG 6 {rule.source}已强化，并保留已学会的{rule.target}。", "BREAK", "",
        ]
    return "\n".join(lines).rstrip() + "\n"


def compile_backfill(rules: list[Rule]) -> str:
    lines = ["[@XY_WSU_LOGIN_BACKFILL]", "{", "; 登录补写：只补目标技能，不改原技能强化等级", ""]
    for index, rule in enumerate(rules, 1):
        next_label = f"@XY_WSU_BACKFILL_NEXT_{index}"
        lines += [
            "#IF", f"CHECKSKILL {rule.source} > {rule.unlock_level - 1} 1", "#ACT", f"GOTO @XY_WSU_BACKFILL_{index}", "#ELSEACT", f"GOTO {next_label}", "",
            f"[@XY_WSU_BACKFILL_{index}]", "#IF", f"NOT CheckMagicName {rule.target}", "#ACT", f"ADDSKILL {rule.target} 3", f"GOTO {next_label}", "",
            "#IF", f"CheckMagicName {rule.target}", f"CHECKSKILL {rule.target} < 3 0", "#ACT", f"SKILLLEVEL {rule.target} = 3 0", f"GOTO {next_label}", "",
            f"GOTO {next_label}", "", f"[{next_label}]", "",
        ]
    lines += ["BREAK", "}", ""]
    return "\n".join(lines)


def compile_desc(rules: list[Rule]) -> str:
    lines: list[str] = []
    for rule in rules:
        for level in range(1, 10):
            suffix = f"；达到强化{rule.unlock_level}后领悟3级{rule.target}" if level == rule.unlock_level else ""
            lines.append(f"普通技能,{rule.source},{level},强化{level}{suffix}")
    return "\r\n".join(lines) + "\r\n"


def _strip_legacy(text: str, kind: str) -> str:
    if kind == "block":
        pattern = re.compile(
            rf"(?ms)^; XYDP-LAB-BEGIN {re.escape(LEGACY_PACKAGE)} SHA256=([0-9a-f]{{64}})\r?\n(.*?)^; XYDP-LAB-END {re.escape(LEGACY_PACKAGE)}\s*\r?\n?"
        )
    else:
        pattern = re.compile(
            rf"(?ms)^; XYDP-LAB-HOOK-BEGIN {re.escape(LEGACY_PACKAGE)} SHA256=([0-9a-f]{{64}})\r?\n(.*?)^; XYDP-LAB-HOOK-END {re.escape(LEGACY_PACKAGE)}\s*\r?\n?"
        )
    match = pattern.search(text)
    if match is None:
        begin = f"XYDP-LAB-{'HOOK-' if kind != 'block' else ''}BEGIN {LEGACY_PACKAGE}"
        end = f"XYDP-LAB-{'HOOK-' if kind != 'block' else ''}END {LEGACY_PACKAGE}"
        if begin in text or end in text:
            raise SkillUpgradeError("旧技能强化受管标记不完整")
        return text
    normalized = match.group(2).replace("\r\n", "\n")
    # The historical lab generator used two checksum conventions: the
    # QFunction block was hashed as GB18030 with one trailing LF, while the
    # QManage hook was hashed as UTF-8 without a trailing LF.  Verify against
    # the exact convention that produced each managed marker before replacing
    # it; do not silently accept a hand-edited legacy block.
    if kind == "block":
        body = normalized.rstrip("\n") + "\n"
        digest = hashlib.sha256(body.encode("gb18030")).hexdigest()
    else:
        body = normalized.rstrip("\n")
        digest = hashlib.sha256(body.encode("utf-8")).hexdigest()
    if digest != match.group(1):
        raise SkillUpgradeError("旧技能强化受管块已被手工修改")
    return text[:match.start()] + text[match.end():]


def _database_records(database: Path) -> tuple[list[Skill], Counter[str], Counter[str]]:
    connection = sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True)
    try:
        skills = [Skill(*map(int_or_text, row)) for row in connection.execute(
            "SELECT MagID,MagName,Job,MaxTrainLv,CanUpgrade,MaxUpgradeLv FROM Magic"
        )]
        skill_names = Counter(skill.name for skill in skills)
        item_names = Counter(str(row[0]) for row in connection.execute(
            "SELECT Name FROM StdItems WHERE Name IS NOT NULL AND trim(Name)<>''"
        ))
    finally:
        connection.close()
    return skills, skill_names, item_names


def int_or_text(value: object) -> object:
    if isinstance(value, str):
        return value
    return int(value)


class SkillUpgradeService:
    def __init__(self, platform_root: Path):
        self.root = Path(platform_root).resolve()
        repository = PackageRepository(self.root / "packages")
        repository.refresh()
        self.installer = Installer(repository, self.root / "backups")

    @staticmethod
    def _change(root: Path, relative: str, after: bytes, operation: str, *, scope: str = "server") -> PlannedChange | None:
        path = root / Path(relative.replace("/", "\\"))
        before = path.read_bytes() if path.is_file() else None
        if before == after:
            return None
        return PlannedChange(relative, before, after, operation, PACKAGE_ID, scope)

    def preflight(
        self,
        server: Path,
        workbook: Path,
        *,
        client: Path | None = None,
        launcher: Path | None = None,
    ) -> SkillUpgradePlan:
        plan = SkillUpgradePlan(
            str(Path(server).resolve()), str(Path(workbook).resolve()),
            str(Path(client).resolve()) if client else None,
            str(Path(launcher).resolve()) if launcher else None,
        )
        try:
            root = Path(server).resolve()
            client_root = Path(client).resolve() if client else None
            launcher_root = Path(launcher).resolve() if launcher else None
            rules, costs = load_workbook(Path(workbook))
            database_path = root / Path(DATABASE.replace("/", "\\"))
            if not database_path.is_file():
                raise SkillUpgradeError(f"目标服缺少{DATABASE}")
            skills, skill_names, item_names = _database_records(database_path)
            selected: dict[str, Skill] = {}
            warnings: list[str] = []
            by_name: dict[str, list[Skill]] = defaultdict(list)
            for skill in skills:
                by_name[skill.name].append(skill)
            for rule in rules:
                source_rows = by_name.get(rule.source, [])
                if len(source_rows) != 1:
                    raise SkillUpgradeError(f"原技能名称无法唯一识别：{rule.source}（{len(source_rows)}条）")
                source_skill = source_rows[0]
                if source_skill.job not in {0, 1, 2} or source_skill.max_train_level < 3:
                    raise SkillUpgradeError(f"原技能不是可强化的普通三职业技能：{rule.source}")
                selected[rule.source] = source_skill
                target_rows = by_name.get(rule.target, [])
                if not target_rows:
                    raise SkillUpgradeError(f"解锁技能不存在：{rule.target}")
                comparable = {(row.job, row.max_train_level) for row in target_rows}
                if len(comparable) != 1 or next(iter(comparable))[0] not in {0, 1, 2}:
                    raise SkillUpgradeError(f"解锁技能同名记录不等价：{rule.target}")
                if len(target_rows) > 1:
                    warnings.append(
                        f"解锁技能{rule.target}在Magic中有{len(target_rows)}条等价记录；按名称执行，保留目标库原状。"
                    )
            for cost in costs:
                for item in (cost.material_a, cost.material_b):
                    if not item:
                        continue
                    if item_names[item] == 0:
                        raise SkillUpgradeError(f"技能强化材料不存在：{item}")
                    if item_names[item] > 1:
                        raise SkillUpgradeError(f"技能强化材料名称不唯一：{item}")

            database = database_path.read_bytes()
            for rule in rules:
                skill = selected[rule.source]
                database = apply_sqlite_magic_skill_update(database, {
                    "type": "sqlite_magic_skill_update", "target": DATABASE,
                    "match": {"MagID": skill.mag_id, "MagName": skill.name, "Job": skill.job},
                    "values": {"CanUpgrade": 1, "MaxUpgradeLv": 9},
                    "allow_normal_professions": True,
                })

            qf_path = root / Path(QFUNCTION.replace("/", "\\"))
            qm_path = root / Path(QMANAGE.replace("/", "\\"))
            qf_doc = read_text_document(qf_path)
            qm_doc = read_text_document(qm_path)
            qf_text = _strip_legacy(qf_doc.text, "block")
            qf_text = install_managed_block(
                qf_text, PACKAGE_ID, compile_qfunction(rules, costs, selected), qf_doc.newline,
            ).text
            qm_text = _strip_legacy(qm_doc.text, "hook")
            qm_text = ensure_event_label(qm_text, PACKAGE_ID, "MAIN1", qm_doc.newline).text
            qm_text = install_event_hook(
                qm_text, PACKAGE_ID, "MAIN1",
                "#IF\n#ACT\n#CALL [\\玄渊实验室\\技能强化\\登录补写.txt] @XY_WSU_LOGIN_BACKFILL",
                qm_doc.newline,
            ).text
            backfill = compile_backfill(rules).replace("\n", "\r\n").encode("gb18030")
            description = compile_desc(rules).encode("gb18030")

            changes = [
                self._change(root, DATABASE, database, "enable-five-skill-upgrade"),
                self._change(root, QFUNCTION, encode_text_document(qf_doc, qf_text), "compile-five-skill-upgrade"),
                self._change(root, QMANAGE, encode_text_document(qm_doc, qm_text), "install-skill-login-backfill"),
                self._change(root, BACKFILL, backfill, "generate-skill-login-backfill"),
            ]
            if client_root is not None:
                changes.append(self._change(client_root, CLIENT_DESC, description, "generate-skill-description", scope="client"))
            if launcher_root is not None:
                changes.append(self._change(launcher_root, LAUNCHER_DESC, description, "mirror-skill-description", scope="launcher"))
            plan.rules = [rule.__dict__ for rule in rules]
            plan.warnings.extend(warnings)
            plan.warnings.append("五条技能强化链已静态编译；仍需重载后逐链完成游戏验收。")
            plan.changes = [change for change in changes if change is not None]
            plan.install_plan = InstallPlan(
                target_root=str(root), client_root=str(client_root) if client_root else None,
                launcher_root=str(launcher_root) if launcher_root else None,
                package_ids=[PACKAGE_ID], package_versions={PACKAGE_ID: PACKAGE_VERSION},
                parameters={"workbook": str(Path(workbook).resolve()), "rule_count": len(rules)},
                changes=plan.changes, warnings=plan.warnings,
                operation_type="warrior-skill-upgrade-bundle", candidate_packages=[PACKAGE_ID],
            )
        except (OSError, ValueError, RuntimeError, sqlite3.Error, SqlitePatchError, TextPatchError) as exc:
            plan.blockers.append(str(exc))
        return plan

    def install(self, plan: SkillUpgradePlan):
        if plan.blockers or plan.install_plan is None:
            raise SkillUpgradeError("技能强化安装被阻止：\n" + "\n".join(plan.blockers))
        return self.installer.install(plan.install_plan)


__all__ = [
    "Cost", "Rule", "Skill", "SkillUpgradeError", "SkillUpgradePlan", "SkillUpgradeService",
    "compile_backfill", "compile_desc", "compile_qfunction", "load_workbook",
]
