from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Any, Iterable
from uuid import uuid4

from openpyxl import load_workbook

from .mingge_native import (
    DIRECT_PROPERTY_SPECS,
    SCRIPT_PROPERTY_SPECS,
    STATIC_ONLY_PROPERTY_SPECS,
    TEXTLINE_RUNTIME_SPECS,
    TRUSTED_EQUIP_SLOTS,
)


NPC_PACKAGE_ID = "xy.optional.mingge-npc"
CONTENT_PACKAGE_ID = "xy.optional.mingge-content"
NPC_ROUTE = "mingge-npc"
CONTENT_ROUTE = "mingge-content"
API_VERSION = 1

BRIDGE_RELATIVE = Path("Mir200/Envir/QuestDiary/玄渊命格/命格公共接口.txt")
NPC_RULES_RELATIVE = Path("Mir200/Envir/QuestDiary/玄渊命格/命格NPC规则.txt")
CONTENT_PROVIDER_RELATIVE = Path("Mir200/Envir/QuestDiary/玄渊命格/命格内容提供者.txt")
QFUNCTION_RELATIVE = Path("Mir200/Envir/Market_Def/QFunction-0.txt")
TEXTVAR_RELATIVE = Path("Mir200/Envir/CustomItemPropertyTextVarList.txt")
ATTACK_SPEED_CORE_RELATIVE = Path("Mir200/Envir/QuestDiary/玄渊攻速突破/全身攻速阈值核心.txt")
ATTACK_SPEED_TEXT_LINE = 40
ATTACK_SPEED_TEXT = "{攻速突破∶|251}+$$2"
ATTACK_SPEED_CALL = "#CALL [\\玄渊攻速突破\\全身攻速阈值核心.txt] @XY_AS_CAP_RECALC"
CONTENT_TEXTVAR_START = 33
LEGACY_P3_TEXTVAR_SHA256 = {
    1001: (33, "7541D05D4ECA33106061C763572255D88A433D272D2436C6CBDED5F6BB9E9941"),
    1002: (34, "3F88F705ADF3886A05291F393C6B3BD9B04FE815277ED9943489B8620C94D9E0"),
    1003: (35, "C9573C38A91FA44D65530AB02F8F4FB4C1046F6CA0AC68B02D4146748EEACBD0"),
    1004: (36, "F7918D57DDBCE527D4039DECD0F1F150F9B7A3BEBA2D82AE505EDFABC09B4B3E"),
    1005: (37, "C3FA1CCD7E7018269CCC8456F90D708588AA632D467E2A30FF090396588F17FF"),
    1006: (38, "9B2178B2EE492F67A20EE19DE580B2B8B2BD0EB5887C8CF05C732853DAE981C7"),
}
CONTENT_PROPERTY_ROWS = (1, 2, 3, 4, 5, 6, 7, 8, 17, 18, 19)

NPC_MARKER = "XY-MG-NPC-V1"
CONTENT_MARKER = "XY-MG-CONTENT-V1"
CONTENT_QFUNCTION_MARKER = "XY-MG-CONTENT-QFUNCTION-V1-BEGIN"

NPC_SHEETS = ("系统设置", "槽位开放", "洗练规则", "品质保底", "命格羁绊", "NPC文案")
CONTENT_SHEETS = ("命格定义", "命格属性", "颜色方案", "数据字典")

_SAFE_SETTING_ID = re.compile(r"^[A-Za-z][A-Za-z0-9_-]{2,31}$")
_LABEL = re.compile(r"^[A-Za-z0-9_\u4e00-\u9fff-]{1,64}$")


class MinggeDualError(ValueError):
    pass


@dataclass(frozen=True)
class MinggeNpcSettings:
    system_id: str
    target_item: str
    equip_part: str
    equip_slot: int
    slot_count: int
    npc_script_relative: str
    npc_hook_label: str
    default_pool_id: int
    missing_content_message: str


@dataclass(frozen=True)
class MinggeSlotRule:
    slot: int
    enabled: bool
    display_name: str
    condition_type: str
    condition_operator: str
    condition_value: int
    cost_type: str
    cost_name: str
    cost_amount: int
    failure_message: str


@dataclass(frozen=True)
class MinggeWashRule:
    rule_id: int
    enabled: bool
    display_name: str
    batch_count: int
    cost_type: str
    cost_name: str
    cost_amount: int
    prerequisite_type: str
    prerequisite_value: int
    pool_id: int
    failure_message: str


@dataclass(frozen=True)
class MinggeQualityRule:
    quality_id: int
    display_name: str
    base_weight: int
    miss_pity: int
    fixed_cycle: int
    reset_on_hit: bool
    note: str


@dataclass(frozen=True)
class MinggeBondRule:
    bond_id: int
    enabled: bool
    display_name: str
    member_type: str
    member_ids: tuple[int, ...]
    required_count: int
    reward_property: str
    reward_value: int
    display_text: str


@dataclass(frozen=True)
class MinggeNpcWorkbook:
    path: Path
    settings: MinggeNpcSettings
    slot_rules: tuple[MinggeSlotRule, ...]
    wash_rules: tuple[MinggeWashRule, ...]
    quality_rules: tuple[MinggeQualityRule, ...]
    bond_rules: tuple[MinggeBondRule, ...]
    npc_text: dict[str, str]


@dataclass(frozen=True)
class MinggeColorScheme:
    scheme_id: int
    display_name: str
    kind: str
    name_colors: tuple[int, ...]
    attribute_colors: tuple[int, ...]
    separator_color: int
    note: str


@dataclass(frozen=True)
class MinggeCandidate:
    candidate_id: int
    state: str
    display_name: str
    quality_id: int
    pool_id: int
    equip_part: str
    color_scheme_id: int
    note: str

    @property
    def enabled(self) -> bool:
        return self.state == "active"

    @property
    def retained(self) -> bool:
        return self.state in {"active", "retired"}


@dataclass(frozen=True)
class MinggeAttribute:
    candidate_id: int
    order: int
    property_name: str
    value: int
    unit: str
    color_scheme_id: int | None
    route: str
    binding: int
    canonical_key: str
    effect_status: str
    note: str


@dataclass(frozen=True)
class MinggeContentWorkbook:
    path: Path
    candidates: tuple[MinggeCandidate, ...]
    attributes: dict[int, tuple[MinggeAttribute, ...]]
    color_schemes: dict[int, MinggeColorScheme]


@dataclass(frozen=True)
class DualFileChange:
    relative_path: str
    path: Path
    kind: str
    before: bytes | None
    after: bytes
    before_sha256: str
    after_sha256: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class DualInstallPlan:
    route: str
    package_id: str
    plan_id: str
    server_root: Path
    workbook: Path
    workbook_sha256: str
    changes: tuple[DualFileChange, ...]
    blockers: tuple[str, ...]
    warnings: tuple[str, ...]
    previews: tuple[tuple[str, tuple[tuple[str, int], ...]], ...] = ()


@dataclass(frozen=True)
class DualInstallReceipt:
    transaction_id: str
    status: str
    receipt_path: Path
    affected_paths: tuple[Path, ...]


def _sha256(data: bytes | None) -> str:
    return hashlib.sha256(data).hexdigest() if data is not None else ""


def _script_bytes(text: str) -> bytes:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").rstrip("\n")
    return (normalized.replace("\n", "\r\n") + "\r\n").encode("gb18030")


def _script_text(data: bytes | None) -> str:
    if not data:
        return ""
    try:
        return data.decode("gb18030").replace("\r\n", "\n").replace("\r", "\n")
    except UnicodeDecodeError as exc:
        raise MinggeDualError("目标脚本不是GB18030可解码文本") from exc


def _value(value: Any, label: str) -> str:
    text = "" if value is None else str(value).strip()
    if not text:
        raise MinggeDualError(f"{label}不能为空")
    return text


def _integer(value: Any, label: str, minimum: int = 0, maximum: int = 2_147_483_647) -> int:
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise MinggeDualError(f"{label}必须是整数") from exc
    if result < minimum or result > maximum:
        raise MinggeDualError(f"{label}必须在{minimum}至{maximum}之间")
    return result


def _enabled(value: Any, label: str = "启用") -> bool:
    text = str(value).strip().casefold()
    if text in {"是", "true", "1", "启用", "yes"}:
        return True
    if text in {"否", "false", "0", "停用", "no"}:
        return False
    raise MinggeDualError(f"{label}必须填写是或否")


def _candidate_state(value: Any, label: str = "启用") -> str:
    text = str(value).strip().casefold()
    if text in {"是", "true", "1", "启用", "yes"}:
        return "active"
    if text in {"退役", "retired"}:
        return "retired"
    if text in {"否", "false", "0", "停用", "no"}:
        return "disabled"
    raise MinggeDualError(f"{label}必须填写是、退役或否")


def _color_sequence(value: Any, label: str) -> tuple[int, ...]:
    raw = _value(value, label).replace("，", ",")
    result = tuple(_integer(item.strip(), label, 0, 255) for item in raw.split(",") if item.strip())
    if not result:
        raise MinggeDualError(f"{label}至少需要一个色号")
    return result


def _headers(sheet, expected: tuple[str, ...]) -> None:
    actual = tuple("" if cell.value is None else str(cell.value).strip() for cell in sheet[1])
    if actual != expected:
        raise MinggeDualError(f"工作表{sheet.title}表头必须为：{'、'.join(expected)}")


def _rows(sheet, expected: tuple[str, ...]) -> Iterable[tuple[int, tuple[Any, ...]]]:
    _headers(sheet, expected)
    for row_index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = tuple(values[: len(expected)])
        if all(value is None or str(value).strip() == "" for value in row):
            continue
        yield row_index, row


def _load_exact(path: Path, sheet_names: tuple[str, ...]):
    source = Path(path).resolve()
    if not source.is_file():
        raise MinggeDualError(f"配置表不存在：{source}")
    workbook = load_workbook(source, read_only=True, data_only=True)
    if tuple(workbook.sheetnames) != sheet_names:
        workbook.close()
        raise MinggeDualError(f"配置表必须依次包含：{'、'.join(sheet_names)}")
    return source, workbook


def load_mingge_npc_workbook(path: Path) -> MinggeNpcWorkbook:
    source, workbook = _load_exact(path, NPC_SHEETS)
    try:
        settings_sheet = workbook["系统设置"]
        settings_headers = ("配置项", "配置值", "说明")
        settings = {str(values[0]).strip(): values[1] for _, values in _rows(settings_sheet, settings_headers)}
        required = (
            "系统编号", "目标装备", "装备部位", "装备位置", "槽位数量",
            "NPC脚本相对路径", "NPC接入标签", "默认候选池ID", "内容包缺失提示",
        )
        missing = [key for key in required if key not in settings]
        if missing:
            raise MinggeDualError("系统设置缺少：" + "、".join(missing))
        system_id = _value(settings["系统编号"], "系统编号")
        if not _SAFE_SETTING_ID.fullmatch(system_id):
            raise MinggeDualError("系统编号必须以字母开头且只含字母、数字、下划线或短横线")
        equip_part = _value(settings["装备部位"], "装备部位")
        if equip_part not in TRUSTED_EQUIP_SLOTS:
            raise MinggeDualError(f"装备部位未登记：{equip_part}")
        equip_slot = _integer(settings["装备位置"], "装备位置", 0, 99)
        if TRUSTED_EQUIP_SLOTS[equip_part] != equip_slot:
            raise MinggeDualError(f"装备部位{equip_part}对应位置{TRUSTED_EQUIP_SLOTS[equip_part]}，与填写值{equip_slot}不一致")
        npc_relative = _value(settings["NPC脚本相对路径"], "NPC脚本相对路径").replace("\\", "/")
        pure = PurePosixPath(npc_relative)
        if pure.is_absolute() or ".." in pure.parts or not npc_relative.startswith("Mir200/Envir/Market_Def/"):
            raise MinggeDualError("NPC脚本相对路径必须位于Mir200/Envir/Market_Def内")
        hook = _value(settings["NPC接入标签"], "NPC接入标签").lstrip("@").strip("[]")
        if not _LABEL.fullmatch(hook):
            raise MinggeDualError("NPC接入标签格式非法")
        parsed_settings = MinggeNpcSettings(
            system_id=system_id,
            target_item=_value(settings["目标装备"], "目标装备"),
            equip_part=equip_part,
            equip_slot=equip_slot,
            slot_count=_integer(settings["槽位数量"], "槽位数量", 1, 8),
            npc_script_relative=npc_relative,
            npc_hook_label=hook,
            default_pool_id=_integer(settings["默认候选池ID"], "默认候选池ID", 1, 999999),
            missing_content_message=_value(settings["内容包缺失提示"], "内容包缺失提示"),
        )

        slot_headers = ("槽位", "启用", "显示名称", "条件类型", "条件运算", "条件值", "消耗类型", "消耗名称", "消耗数量", "失败提示")
        slot_rules = tuple(
            MinggeSlotRule(
                _integer(v[0], f"槽位开放第{row}行槽位", 1, 8), _enabled(v[1]),
                _value(v[2], f"槽位开放第{row}行显示名称"), _value(v[3], f"槽位开放第{row}行条件类型"),
                _value(v[4], f"槽位开放第{row}行条件运算"), _integer(v[5], f"槽位开放第{row}行条件值"),
                _value(v[6], f"槽位开放第{row}行消耗类型"), "" if v[7] is None else str(v[7]).strip(),
                _integer(v[8], f"槽位开放第{row}行消耗数量"), _value(v[9], f"槽位开放第{row}行失败提示"),
            )
            for row, v in _rows(workbook["槽位开放"], slot_headers)
        )
        for rule in slot_rules:
            if rule.condition_type not in {"等级", "无条件"}:
                raise MinggeDualError(f"槽位{rule.slot}条件类型首版只支持：等级、无条件")
            if rule.condition_type == "等级" and rule.condition_operator != ">=":
                raise MinggeDualError(f"槽位{rule.slot}等级条件运算首版只支持>=")
            if rule.cost_type not in {"无", "元宝", "物品"}:
                raise MinggeDualError(f"槽位{rule.slot}消耗类型只支持：无、元宝、物品")
        active_slots = sorted(rule.slot for rule in slot_rules if rule.enabled)
        if active_slots != list(range(1, parsed_settings.slot_count + 1)):
            raise MinggeDualError("启用槽位必须从1开始连续并与槽位数量一致")

        wash_headers = ("规则ID", "启用", "名称", "次数", "消耗类型", "消耗名称", "消耗数量", "前置条件类型", "前置条件值", "候选池ID", "失败提示")
        wash_rules = tuple(
            MinggeWashRule(
                _integer(v[0], f"洗练规则第{row}行规则ID", 1), _enabled(v[1]), _value(v[2], f"洗练规则第{row}行名称"),
                _integer(v[3], f"洗练规则第{row}行次数", 1, 100), _value(v[4], f"洗练规则第{row}行消耗类型"),
                _value(v[5], f"洗练规则第{row}行消耗名称"), _integer(v[6], f"洗练规则第{row}行消耗数量"),
                _value(v[7], f"洗练规则第{row}行前置条件类型"), _integer(v[8], f"洗练规则第{row}行前置条件值"),
                _integer(v[9], f"洗练规则第{row}行候选池ID", 1), _value(v[10], f"洗练规则第{row}行失败提示"),
            )
            for row, v in _rows(workbook["洗练规则"], wash_headers)
            if _enabled(v[1])
        )
        if not wash_rules:
            raise MinggeDualError("至少需要一条启用的洗练规则")
        for rule in wash_rules:
            if rule.cost_type not in {"无", "元宝", "物品"}:
                raise MinggeDualError(f"洗练规则{rule.rule_id}消耗类型只支持：无、元宝、物品")
            if rule.prerequisite_type not in {"已选择槽位", "等级", "无条件"}:
                raise MinggeDualError(f"洗练规则{rule.rule_id}前置条件只支持：已选择槽位、等级、无条件")

        quality_headers = ("品质ID", "品质名称", "基础权重", "连续未出保底次数", "固定周期次数", "命中后清零", "说明")
        quality_rules = tuple(
            MinggeQualityRule(
                _integer(v[0], f"品质保底第{row}行品质ID", 1), _value(v[1], f"品质保底第{row}行品质名称"),
                _integer(v[2], f"品质保底第{row}行基础权重", 0, 100000), _integer(v[3], f"品质保底第{row}行连续未出保底次数"),
                _integer(v[4], f"品质保底第{row}行固定周期次数"), _enabled(v[5], "命中后清零"),
                "" if v[6] is None else str(v[6]).strip(),
            )
            for row, v in _rows(workbook["品质保底"], quality_headers)
        )
        if not quality_rules or sum(item.base_weight for item in quality_rules) <= 0:
            raise MinggeDualError("品质基础权重合计必须大于0")
        if any(item.quality_id > 9 for item in quality_rules):
            raise MinggeDualError("品质ID首版必须在1至9之间，以便平台稳定分配人物持久计数")

        bond_headers = ("羁绊ID", "启用", "羁绊名称", "成员类型", "成员ID列表", "需要数量", "奖励属性", "奖励数值", "显示文字")
        bond_rules = []
        for row, v in _rows(workbook["命格羁绊"], bond_headers):
            members = tuple(_integer(item.strip(), f"命格羁绊第{row}行成员ID", 1) for item in str(v[4]).replace("，", ",").split(",") if item.strip())
            bond_rules.append(MinggeBondRule(
                _integer(v[0], f"命格羁绊第{row}行羁绊ID", 1), _enabled(v[1]), _value(v[2], f"命格羁绊第{row}行羁绊名称"),
                _value(v[3], f"命格羁绊第{row}行成员类型"), members, _integer(v[5], f"命格羁绊第{row}行需要数量", 1, 8),
                _value(v[6], f"命格羁绊第{row}行奖励属性"), _integer(v[7], f"命格羁绊第{row}行奖励数值"),
                _value(v[8], f"命格羁绊第{row}行显示文字"),
            ))

        text_headers = ("文案ID", "启用", "显示文字", "说明")
        npc_text = {
            _value(v[0], f"NPC文案第{row}行文案ID"): _value(v[2], f"NPC文案第{row}行显示文字")
            for row, v in _rows(workbook["NPC文案"], text_headers) if _enabled(v[1])
        }
        return MinggeNpcWorkbook(source, parsed_settings, slot_rules, wash_rules, quality_rules, tuple(bond_rules), npc_text)
    finally:
        workbook.close()


def _property_spec(name: str):
    return (
        TEXTLINE_RUNTIME_SPECS.get(name)
        or DIRECT_PROPERTY_SPECS.get(name)
        or SCRIPT_PROPERTY_SPECS.get(name)
        or STATIC_ONLY_PROPERTY_SPECS.get(name)
    )


def load_mingge_content_workbook(path: Path) -> MinggeContentWorkbook:
    source, workbook = _load_exact(path, CONTENT_SHEETS)
    try:
        color_headers = ("方案ID", "中文名称", "类型", "名称颜色序列", "属性颜色序列", "分隔符颜色", "说明")
        schemes: dict[int, MinggeColorScheme] = {}
        for row, v in _rows(workbook["颜色方案"], color_headers):
            scheme_id = _integer(v[0], f"颜色方案第{row}行方案ID", 1)
            if scheme_id in schemes:
                raise MinggeDualError(f"颜色方案ID重复：{scheme_id}")
            kind = _value(v[2], f"颜色方案第{row}行类型")
            if kind not in {"统一色", "原生多色"}:
                raise MinggeDualError(f"颜色方案第{row}行类型必须为统一色或原生多色")
            schemes[scheme_id] = MinggeColorScheme(
                scheme_id, _value(v[1], f"颜色方案第{row}行中文名称"), kind,
                _color_sequence(v[3], f"颜色方案第{row}行名称颜色序列"),
                _color_sequence(v[4], f"颜色方案第{row}行属性颜色序列"),
                _integer(v[5], f"颜色方案第{row}行分隔符颜色", 0, 255),
                "" if v[6] is None else str(v[6]).strip(),
            )

        candidate_headers = ("命格ID", "启用", "命格名称", "品质ID", "候选池ID", "适用部位", "颜色方案ID", "备注")
        candidates: list[MinggeCandidate] = []
        seen_ids: set[int] = set()
        for row, v in _rows(workbook["命格定义"], candidate_headers):
            candidate_id = _integer(v[0], f"命格定义第{row}行命格ID", 1, 999999999)
            if candidate_id in seen_ids:
                raise MinggeDualError(f"命格ID重复：{candidate_id}")
            seen_ids.add(candidate_id)
            scheme_id = _integer(v[6], f"命格定义第{row}行颜色方案ID", 1)
            if scheme_id not in schemes:
                raise MinggeDualError(f"命格定义第{row}行颜色方案ID不存在：{scheme_id}")
            part = _value(v[5], f"命格定义第{row}行适用部位")
            if part not in TRUSTED_EQUIP_SLOTS:
                raise MinggeDualError(f"命格定义第{row}行适用部位未登记：{part}")
            candidates.append(MinggeCandidate(
                candidate_id, _candidate_state(v[1]), _value(v[2], f"命格定义第{row}行命格名称"),
                _integer(v[3], f"命格定义第{row}行品质ID", 1), _integer(v[4], f"命格定义第{row}行候选池ID", 1),
                part, scheme_id, "" if v[7] is None else str(v[7]).strip(),
            ))
        all_candidates = tuple(candidates)
        candidates = [item for item in all_candidates if item.retained]
        if not candidates:
            raise MinggeDualError("命格定义至少需要一条启用或退役记录")

        attr_headers = ("命格ID", "顺序", "中文属性", "数值", "单位", "属性颜色方案ID", "备注")
        attrs: dict[int, list[MinggeAttribute]] = {item.candidate_id: [] for item in candidates}
        for row, v in _rows(workbook["命格属性"], attr_headers):
            candidate_id = _integer(v[0], f"命格属性第{row}行命格ID", 1)
            if candidate_id not in seen_ids:
                raise MinggeDualError(f"命格属性第{row}行引用的命格ID不存在：{candidate_id}")
            if candidate_id not in attrs:
                continue
            name = _value(v[2], f"命格属性第{row}行中文属性")
            spec = _property_spec(name)
            if spec is None:
                raise MinggeDualError(f"命格属性第{row}行中文属性未登记：{name}")
            override = None if v[5] is None or str(v[5]).strip() == "" else _integer(v[5], f"命格属性第{row}行属性颜色方案ID", 1)
            if override is not None and override not in schemes:
                raise MinggeDualError(f"命格属性第{row}行属性颜色方案ID不存在：{override}")
            attrs[candidate_id].append(MinggeAttribute(
                candidate_id, _integer(v[1], f"命格属性第{row}行顺序", 1, 99), name,
                _integer(v[3], f"命格属性第{row}行数值", -2_147_483_648, 2_147_483_647),
                "" if v[4] is None else str(v[4]).strip(), override, spec.route, spec.binding,
                spec.canonical_key, spec.effect_status, "" if v[6] is None else str(v[6]).strip(),
            ))
        frozen_attrs: dict[int, tuple[MinggeAttribute, ...]] = {}
        for candidate in candidates:
            rows = sorted(attrs[candidate.candidate_id], key=lambda item: item.order)
            if not rows:
                raise MinggeDualError(f"命格{candidate.candidate_id}至少需要一条属性")
            if len({item.order for item in rows}) != len(rows):
                raise MinggeDualError(f"命格{candidate.candidate_id}属性顺序重复")
            frozen_attrs[candidate.candidate_id] = tuple(rows)
        return MinggeContentWorkbook(source, tuple(sorted(candidates, key=lambda item: item.candidate_id)), frozen_attrs, schemes)
    finally:
        workbook.close()


BRIDGE_BASE = """{
; XY-MG-API-V1 PLATFORM SHARED BRIDGE
; 此文件是平台公共合同，任一命格子包均可幂等创建；子包只维护自己的受管块。
[@XY_MG_API_STATUS]
#IF
#ACT
MOV N$XY_MG_API_STATUS 0
; XY-MG-CONTENT-STATUS-ANCHOR
BREAK

[@XY_MG_API_DRAW]
#IF
#ACT
MOV N$XY_MG_API_STATUS 0
; XY-MG-CONTENT-DRAW-ANCHOR
BREAK

[@XY_MG_API_READ]
#IF
#ACT
MOV N$XY_MG_API_STATUS 0
; XY-MG-CONTENT-READ-ANCHOR
BREAK

[@XY_MG_API_RECALC]
#IF
#ACT
MOV N$XY_MG_API_STATUS 0
; XY-MG-CONTENT-RECALC-ANCHOR
BREAK
}
"""


def _managed_block(marker: str, lines: Iterable[str]) -> str:
    return "\n".join((f"; {marker}-BEGIN", *lines, f"; {marker}-END"))


def _find_block(text: str, marker: str) -> tuple[int, int] | None:
    begin = f"; {marker}-BEGIN"
    end = f"; {marker}-END"
    start = text.find(begin)
    finish = text.find(end)
    if start < 0 and finish < 0:
        return None
    if start < 0 or finish < start or text.find(begin, start + 1) >= 0 or text.find(end, finish + 1) >= 0:
        raise MinggeDualError(f"受管块{marker}重复或残缺")
    return start, finish + len(end)


def _upsert_after_anchor(text: str, anchor: str, marker: str, lines: Iterable[str]) -> tuple[str, str]:
    block = _managed_block(marker, lines)
    found = _find_block(text, marker)
    if found:
        start, end = found
        return text[:start] + block + text[end:], block
    if text.count(anchor) != 1:
        raise MinggeDualError(f"公共接口锚点必须恰好出现一次：{anchor}")
    at = text.index(anchor) + len(anchor)
    return text[:at] + "\n" + block + text[at:], block


def _upsert_after_label(text: str, label: str, marker: str, lines: Iterable[str]) -> tuple[str, str]:
    block = _managed_block(marker, lines)
    found = _find_block(text, marker)
    if found:
        start, end = found
        return text[:start] + block + text[end:], block
    anchor = f"[@{label}]"
    matches = list(re.finditer(rf"(?im)^\s*{re.escape(anchor)}\s*$", text))
    if len(matches) != 1:
        raise MinggeDualError(f"NPC接入标签必须恰好出现一次：{anchor}")
    at = matches[0].end()
    return text[:at] + "\n" + block + text[at:], block


def _remove_block(text: str, marker: str, expected: str | None = None) -> str:
    found = _find_block(text, marker)
    if not found:
        raise MinggeDualError(f"回滚时受管块不存在：{marker}")
    start, end = found
    actual = text[start:end]
    if expected is not None and actual != expected:
        raise MinggeDualError(f"回滚时受管块内容已变化：{marker}")
    before = text[:start].rstrip("\n")
    after = text[end:].lstrip("\n")
    return before + ("\n" if before and after else "") + after


def _patch_existing_npc_actions(root: Path, npc_text: str, book: MinggeNpcWorkbook) -> DualFileChange:
    calls = re.findall(r"(?im)^\s*#CALL\s+\[\\([^\]]+\.txt)\]\s+@(\w+)\s*$", npc_text)
    if len(calls) != 1:
        raise MinggeDualError("现有NPC必须恰好调用一个QuestDiary命格核心，平台才能只接管动作而保留布局")
    called_path, called_label = calls[0]
    if called_label != "XY_MG_MAIN":
        raise MinggeDualError("现有NPC命格核心入口必须为@XY_MG_MAIN")
    relative = Path("Mir200/Envir/QuestDiary") / Path(called_path.replace("\\", "/"))
    core_path = _target(root, relative)
    if not core_path.is_file():
        raise MinggeDualError(f"现有NPC调用的命格核心不存在：{relative}")
    before = core_path.read_bytes()
    text = _script_text(before)
    markers: list[str] = []
    blocks: list[str] = []
    action_specs: list[tuple[str, str, tuple[str, ...]]] = []
    for slot in range(1, book.settings.slot_count + 1):
        action_specs.append((
            f"XY_MG_SLOT{slot}", f"XY-MG-NPC-ACTION-SLOT-{slot}-V1",
            ("#IF", "#ACT", f"#CALL [\\玄渊命格\\命格NPC规则.txt] @XY_MG_NPC_SELECT_SLOT_{slot}", "GOTO @XY_MG_PANEL_ROUTE", "BREAK"),
        ))
    one = next((item for item in book.wash_rules if item.batch_count == 1), None)
    ten = next((item for item in book.wash_rules if item.batch_count == 10), None)
    if one is not None:
        action_specs.append(("XY_MG_WASH_ONE", "XY-MG-NPC-ACTION-WASH-ONE-V1", ("#IF", "#ACT", f"#CALL [\\玄渊命格\\命格NPC规则.txt] @XY_MG_NPC_WASH_{one.rule_id}", "GOTO @XY_MG_PANEL_ROUTE", "BREAK")))
    if ten is not None:
        action_specs.append(("XY_MG_WASH_TEN", "XY-MG-NPC-ACTION-WASH-TEN-V1", ("#IF", "#ACT", f"#CALL [\\玄渊命格\\命格NPC规则.txt] @XY_MG_NPC_WASH_{ten.rule_id}", "GOTO @XY_MG_PANEL_ROUTE", "BREAK")))
    if book.bond_rules:
        action_specs.append(("XY_MG_BOND_INFO", "XY-MG-NPC-ACTION-BOND-V1", ("#IF", "#ACT", "#CALL [\\玄渊命格\\命格NPC规则.txt] @XY_MG_NPC_BOND_INFO", "BREAK")))
    for label, marker, script_lines in action_specs:
        text, block = _upsert_after_label(text, label, marker, script_lines)
        markers.append(marker)
        blocks.append(block)
    return _change(root, relative, "managed_blocks", before, _script_bytes(text), {"markers": markers, "blocks": blocks})


def _bridge_before(path: Path) -> tuple[bytes | None, str]:
    before = path.read_bytes() if path.is_file() else None
    if before is None:
        return None, BRIDGE_BASE
    text = _script_text(before)
    if "; XY-MG-API-V1 PLATFORM SHARED BRIDGE" not in text:
        raise MinggeDualError("目标服已存在非平台命格公共接口，禁止覆盖")
    for anchor in (
        "; XY-MG-CONTENT-STATUS-ANCHOR", "; XY-MG-CONTENT-DRAW-ANCHOR",
        "; XY-MG-CONTENT-READ-ANCHOR", "; XY-MG-CONTENT-RECALC-ANCHOR",
    ):
        if text.count(anchor) != 1:
            raise MinggeDualError(f"命格公共接口合同损坏：{anchor}")
    normalized = text.rstrip("\n")
    lines = normalized.splitlines()
    if not (lines[0].strip() == "{" and lines[-1].strip() == "}"):
        if any(line.strip() in {"{", "}"} for line in lines):
            raise MinggeDualError("命格公共接口可调用边界不完整")
        normalized = "{\n" + normalized + "\n}"
    return before, normalized


def _candidate_segments(book: MinggeContentWorkbook, candidate: MinggeCandidate) -> tuple[tuple[str, int], ...]:
    scheme = book.color_schemes[candidate.color_scheme_id]
    segments: list[tuple[str, int]] = []
    if scheme.kind == "统一色":
        color = scheme.name_colors[0]
        for start in range(0, len(candidate.display_name), 2):
            segments.append((candidate.display_name[start:start + 2], color))
    else:
        for index, char in enumerate(candidate.display_name):
            segments.append((char, scheme.name_colors[index % len(scheme.name_colors)]))
    for index, attribute in enumerate(book.attributes[candidate.candidate_id]):
        segments.append(("\\", scheme.separator_color))
        attr_scheme = book.color_schemes.get(attribute.color_scheme_id or candidate.color_scheme_id, scheme)
        color = attr_scheme.attribute_colors[index % len(attr_scheme.attribute_colors)]
        unit = attribute.unit or ("" if attribute.route == "direct" else "%" if attribute.value else "")
        segments.append((f"{attribute.property_name}{attribute.value:+d}{unit}", color))
    return tuple(segments)


def _candidate_text(segments: tuple[tuple[str, int], ...]) -> str:
    return "".join(f"{{{text}|{color}}}" for text, color in segments)


def _render_npc_rules_legacy(book: MinggeNpcWorkbook) -> str:
    s = book.settings
    lines = [
        "; XY-MG-NPC-RULES-V1", f"; SYSTEM_ID={s.system_id}", f"; TARGET_ITEM={s.target_item}",
        f"; EQUIP_PART={s.equip_part}", f"; EQUIP_SLOT={s.equip_slot}", f"; SLOT_COUNT={s.slot_count}", "",
        "[@XY_MG_NPC_MAIN]", "#IF", "#ACT", f"MOV N$XY_MG_API_VERSION {API_VERSION}",
        f"MOV N$XY_MG_API_EQUIP_SLOT {s.equip_slot}", f"MOV N$XY_MG_API_POOL_ID {s.default_pool_id}",
        "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_STATUS", "BREAK", "",
    ]
    for slot in book.slot_rules:
        if not slot.enabled:
            continue
        lines.extend((
            f"[@XY_MG_NPC_SELECT_SLOT_{slot.slot}]", "#IF", f"CHECKLEVELEX > {slot.condition_value - 1}",
            "#ACT", f"MOV N$XY_MG_API_SLOT {slot.slot}", "BREAK", "#ELSEACT",
            f"MESSAGEBOX {slot.failure_message}", "BREAK", "",
        ))
    quality_total = sum(rule.base_weight for rule in book.quality_rules)
    for rule in book.wash_rules:
        lines.extend((
            f"[@XY_MG_NPC_WASH_{rule.rule_id}]", "#IF", "#ACT",
            "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_STATUS",
            "#IF", "EQUAL N$XY_MG_API_STATUS 0", "#ACT",
            f"MESSAGEBOX {s.missing_content_message}", "BREAK",
            "#IF", "EQUAL N$XY_MG_API_SLOT 0", "#ACT", f"MESSAGEBOX {rule.failure_message}", "BREAK",
        ))
        if rule.cost_type == "元宝":
            lines.extend(("#IF", f"CHECKGAMEGOLD > {max(rule.cost_amount - 1, 0)}", "#ACT", f"GAMEGOLD - {rule.cost_amount}"))
        elif rule.cost_type == "无":
            lines.extend(("#IF", "#ACT"))
        else:
            lines.extend(("#IF", f"CHECKITEM {rule.cost_name} {rule.cost_amount}", "#ACT", f"TAKE {rule.cost_name} {rule.cost_amount}"))
        lines.extend((f"MOV N$XY_MG_API_POOL_ID {rule.pool_id}", f"MOV N$XY_MG_API_BATCH {rule.batch_count}"))
        remaining = quality_total
        enabled_qualities = [item for item in book.quality_rules if item.base_weight > 0]
        for index, quality in enumerate(enabled_qualities):
            label = f"@XY_MG_NPC_QUALITY_{rule.rule_id}_{quality.quality_id}"
            if index == len(enabled_qualities) - 1:
                lines.extend(("#IF", "#ACT", f"GOTO {label}", "BREAK"))
            else:
                lines.extend(("#IF", f"RANDOMEX {quality.base_weight} {remaining}", "#ACT", f"GOTO {label}", "BREAK", "#IF", "#ACT"))
                remaining -= quality.base_weight
        lines.append("")
        for quality in enabled_qualities:
            lines.extend((
                f"[@XY_MG_NPC_QUALITY_{rule.rule_id}_{quality.quality_id}]", "#IF", "#ACT",
                f"MOV N$XY_MG_API_QUALITY_ID {quality.quality_id}",
                "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_DRAW", "BREAK", "",
            ))
    lines.extend(("[@XY_MG_NPC_BOND_INFO]", "#IF", "#ACT", "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_READ"))
    for bond in book.bond_rules:
        lines.append(f"; BOND {bond.bond_id} ENABLED={int(bond.enabled)} MEMBERS={','.join(map(str, bond.member_ids))} NEED={bond.required_count} REWARD={bond.reward_property}{bond.reward_value:+d}")
    lines.extend(("BREAK", ""))
    return "\n".join(lines)


def _unlock_var(slot: int) -> str:
    return f"U{600 + slot}"


def _miss_var(quality_id: int, slot: int) -> str:
    return f"U{620 + quality_id * 8 + slot}"


def _cycle_var(quality_id: int, slot: int) -> str:
    return f"U{820 + quality_id * 8 + slot}"


def _render_npc_rules(book: MinggeNpcWorkbook) -> str:
    s = book.settings
    qualities = tuple(item for item in book.quality_rules if item.base_weight > 0)
    total_weight = sum(item.base_weight for item in qualities)
    lines = [
        "; XY-MG-NPC-RULES-V1", f"; SYSTEM_ID={s.system_id}", f"; TARGET_ITEM={s.target_item}",
        f"; EQUIP_PART={s.equip_part}", f"; EQUIP_SLOT={s.equip_slot}", f"; SLOT_COUNT={s.slot_count}",
        "; 人物持久变量由平台稳定分配：U601-U608槽位，U629起品质未出，U829起固定周期。", "",
        "[@XY_MG_NPC_MAIN]", "#IF", "#ACT", f"MOV N$XY_MG_API_VERSION {API_VERSION}",
        f"MOV N$XY_MG_API_EQUIP_SLOT {s.equip_slot}", f"MOV N$XY_MG_API_POOL_ID {s.default_pool_id}",
        "MOV N$XY_MG_API_QUALITY_ID 0", "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_STATUS", "BREAK", "",
    ]
    for slot in book.slot_rules:
        if not slot.enabled:
            continue
        open_label = f"@XY_MG_NPC_OPEN_SLOT_{slot.slot}"
        selected_label = f"@XY_MG_NPC_SLOT_READY_{slot.slot}"
        lines.extend((
            f"[@XY_MG_NPC_SELECT_SLOT_{slot.slot}]", "#IF", f"EQUAL {_unlock_var(slot.slot)} 1", "#ACT", f"GOTO {selected_label}", "BREAK",
        ))
        if slot.condition_type == "无条件":
            lines.extend(("#IF", "#ACT", f"GOTO {open_label}", "BREAK"))
        else:
            lines.extend(("#IF", f"CHECKLEVELEX > {max(slot.condition_value - 1, 0)}", "#ACT", f"GOTO {open_label}", "BREAK", "#ELSEACT", f"MESSAGEBOX {slot.failure_message}", "BREAK"))
        lines.extend((f"[{open_label}]", "#IF"))
        if slot.cost_type == "元宝":
            lines.extend((f"CHECKGAMEGOLD > {max(slot.cost_amount - 1, 0)}", "#ACT", f"GAMEGOLD - {slot.cost_amount}", f"MOV {_unlock_var(slot.slot)} 1", f"GOTO {selected_label}", "BREAK", "#ELSEACT", f"MESSAGEBOX {slot.failure_message}", "BREAK"))
        elif slot.cost_type == "物品":
            lines.extend((f"CHECKITEM {slot.cost_name} {slot.cost_amount}", "#ACT", f"TAKE {slot.cost_name} {slot.cost_amount}", f"MOV {_unlock_var(slot.slot)} 1", f"GOTO {selected_label}", "BREAK", "#ELSEACT", f"MESSAGEBOX {slot.failure_message}", "BREAK"))
        else:
            lines.extend(("#ACT", f"MOV {_unlock_var(slot.slot)} 1", f"GOTO {selected_label}", "BREAK"))
        lines.extend((f"[{selected_label}]", "#IF", "#ACT", f"MOV N$XY_MG_API_SLOT {slot.slot}", "BREAK", ""))

    for rule in book.wash_rules:
        pay_label = f"@XY_MG_NPC_PAY_{rule.rule_id}"
        lines.extend((
            f"[@XY_MG_NPC_WASH_{rule.rule_id}]", "#IF", "#ACT",
            "MOV N$XY_MG_API_QUALITY_ID 0", "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_STATUS",
            "#IF", "EQUAL N$XY_MG_API_STATUS 0", "#ACT", f"MESSAGEBOX {s.missing_content_message}", "BREAK",
        ))
        if rule.prerequisite_type == "已选择槽位":
            lines.extend(("#IF", "EQUAL N$XY_MG_API_SLOT 0", "#ACT", f"MESSAGEBOX {rule.failure_message}", "BREAK"))
        elif rule.prerequisite_type == "等级":
            lines.extend(("#IF", f"CHECKLEVELEX > {max(rule.prerequisite_value - 1, 0)}", "#ACT", f"GOTO {pay_label}", "BREAK", "#ELSEACT", f"MESSAGEBOX {rule.failure_message}", "BREAK"))
        lines.extend((f"[{pay_label}]", "#IF"))
        if rule.cost_type == "元宝":
            lines.extend((f"CHECKGAMEGOLD > {max(rule.cost_amount - 1, 0)}", "#ACT", f"GAMEGOLD - {rule.cost_amount}", "#ELSEACT", f"MESSAGEBOX {rule.failure_message}", "BREAK", "#IF", "#ACT"))
        elif rule.cost_type == "物品":
            lines.extend((f"CHECKITEM {rule.cost_name} {rule.cost_amount}", "#ACT", f"TAKE {rule.cost_name} {rule.cost_amount}", "#ELSEACT", f"MESSAGEBOX {rule.failure_message}", "BREAK", "#IF", "#ACT"))
        else:
            lines.append("#ACT")
        lines.extend((f"MOV N$XY_MG_API_POOL_ID {rule.pool_id}", f"MOV N$XY_MG_NPC_BATCH_LEFT {rule.batch_count}", f"GOTO @XY_MG_NPC_BATCH_{rule.rule_id}", "BREAK", ""))

        lines.extend((f"[@XY_MG_NPC_BATCH_{rule.rule_id}]", "#IF", "LARGE N$XY_MG_NPC_BATCH_LEFT 0", "#ACT"))
        for slot in range(1, s.slot_count + 1):
            lines.extend(("#IF", f"EQUAL N$XY_MG_API_SLOT {slot}", "#ACT", f"GOTO @XY_MG_NPC_QUALITY_{rule.rule_id}_SLOT_{slot}", "BREAK"))
        lines.extend(("#IF", "#ACT", "BREAK", ""))

        for slot in range(1, s.slot_count + 1):
            lines.append(f"[@XY_MG_NPC_QUALITY_{rule.rule_id}_SLOT_{slot}]")
            for quality in sorted(qualities, key=lambda item: item.quality_id, reverse=True):
                if quality.fixed_cycle > 0:
                    lines.extend(("#IF", f"LARGE {_cycle_var(quality.quality_id, slot)} {max(quality.fixed_cycle - 2, 0)}", "#ACT", f"GOTO @XY_MG_NPC_DRAW_{rule.rule_id}_{slot}_{quality.quality_id}", "BREAK"))
                if quality.miss_pity > 0:
                    lines.extend(("#IF", f"LARGE {_miss_var(quality.quality_id, slot)} {max(quality.miss_pity - 2, 0)}", "#ACT", f"GOTO @XY_MG_NPC_DRAW_{rule.rule_id}_{slot}_{quality.quality_id}", "BREAK"))
            remaining = total_weight
            for index, quality in enumerate(qualities):
                label = f"@XY_MG_NPC_DRAW_{rule.rule_id}_{slot}_{quality.quality_id}"
                if index == len(qualities) - 1:
                    lines.extend(("#IF", "#ACT", f"GOTO {label}", "BREAK"))
                else:
                    lines.extend(("#IF", f"RANDOMEX {quality.base_weight} {remaining}", "#ACT", f"GOTO {label}", "BREAK", "#IF", "#ACT"))
                    remaining -= quality.base_weight
            lines.append("")
            for selected in qualities:
                lines.extend((f"[@XY_MG_NPC_DRAW_{rule.rule_id}_{slot}_{selected.quality_id}]", "#IF", "#ACT", f"MOV N$XY_MG_API_QUALITY_ID {selected.quality_id}", "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_DRAW"))
                for quality in qualities:
                    if quality.quality_id == selected.quality_id and quality.reset_on_hit:
                        lines.extend((f"MOV {_miss_var(quality.quality_id, slot)} 0", f"MOV {_cycle_var(quality.quality_id, slot)} 0"))
                    else:
                        lines.extend((f"INC {_miss_var(quality.quality_id, slot)} 1", f"INC {_cycle_var(quality.quality_id, slot)} 1"))
                lines.extend(("DEC N$XY_MG_NPC_BATCH_LEFT 1", f"GOTO @XY_MG_NPC_BATCH_{rule.rule_id}", "BREAK", ""))

    lines.extend(("[@XY_MG_NPC_BOND_INFO]", "#IF", "#ACT", "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_READ"))
    for bond in book.bond_rules:
        lines.append(f"; BOND {bond.bond_id} ENABLED={int(bond.enabled)} MEMBERS={','.join(map(str, bond.member_ids))} NEED={bond.required_count} REWARD={bond.reward_property}{bond.reward_value:+d} TEXT={bond.display_text}")
    lines.extend(("BREAK", ""))
    return "\n".join(lines)


def _content_aggregates(book: MinggeContentWorkbook):
    result = []
    seen: set[tuple[int, int, int | None, str]] = set()
    for candidate in book.candidates:
        for attribute in book.attributes[candidate.candidate_id]:
            spec = _property_spec(attribute.property_name)
            if spec.route not in {"direct", "script_bind"} and spec.text_line != ATTACK_SPEED_TEXT_LINE:
                continue
            key = (spec.binding, spec.type3, spec.type4, spec.value_command)
            if key not in seen:
                seen.add(key)
                result.append((spec, attribute.order))
    return tuple(result)


def _content_increment_lines(book: MinggeContentWorkbook, candidate: MinggeCandidate, aggregates: tuple[object, ...]) -> tuple[str, ...]:
    index_by_key = {
        (spec.binding, spec.type3, spec.type4, spec.value_command): index + 1
        for index, (spec, _position) in enumerate(aggregates)
    }
    lines: list[str] = []
    for attribute in book.attributes[candidate.candidate_id]:
        spec = _property_spec(attribute.property_name)
        if spec.route not in {"direct", "script_bind"} and spec.text_line != ATTACK_SPEED_TEXT_LINE:
            continue
        number = index_by_key[(spec.binding, spec.type3, spec.type4, spec.value_command)]
        lines.append(f"INC N$XY_MG_CONTENT_SUM{number}A {attribute.value}")
    return tuple(lines)


def _clear_content_row(equip: str, row: int) -> tuple[str, ...]:
    return (
        f"SetCustomItemAbil {equip} {row} 0 0", f"SetCustomItemAbil {equip} {row} 1 0",
        f"SetCustomItemAbil {equip} {row} 2 0", f"SetCustomItemAbil {equip} {row} 3 0",
        f"SetCustomItemAbil {equip} {row} 4 0", f"SetCustomItemValueEx {equip} {row} = 0 0 0",
    )


def _dynamic_content_property_lines(equip: str, index: int, spec: object, position: int) -> tuple[str, ...]:
    """Write one non-zero aggregate into the next free native instance row.

    Catalog size is deliberately unrelated to the eleven physical rows.  Rows are
    assigned from the properties actually present on the current item, so a large
    workbook is not rejected and no catalog entry is dropped by Python ``zip``.
    """
    sum_var = f"N$XY_MG_CONTENT_SUM{index}A"
    lines = [
        "#IF", f"NOT EQUAL {sum_var} 0", "#ACT",
        "INC N$XY_MG_CONTENT_PROP_INDEX 1",
        "MOV N$XY_MG_CONTENT_PROP_ROW 0",
    ]
    for ordinal, row in enumerate(CONTENT_PROPERTY_ROWS, start=1):
        lines.extend((
            "#IF", f"EQUAL N$XY_MG_CONTENT_PROP_INDEX {ordinal}", "#ACT",
            f"MOV N$XY_MG_CONTENT_PROP_ROW {row}",
        ))
    lines.extend((
        "#IF", "LARGE N$XY_MG_CONTENT_PROP_ROW 0", "#ACT",
        f"SetCustomItemAbil {equip} <$STR(N$XY_MG_CONTENT_PROP_ROW)> 0 250",
        f"SetCustomItemAbil {equip} <$STR(N$XY_MG_CONTENT_PROP_ROW)> 1 {spec.binding}",
        f"SetCustomItemAbil {equip} <$STR(N$XY_MG_CONTENT_PROP_ROW)> 2 "
        f"{spec.text_line if spec.text_line == ATTACK_SPEED_TEXT_LINE else position}",
        f"SetCustomItemAbil {equip} <$STR(N$XY_MG_CONTENT_PROP_ROW)> 3 {spec.type3}",
    ))
    if spec.type4 is not None:
        lines.append(f"SetCustomItemAbil {equip} <$STR(N$XY_MG_CONTENT_PROP_ROW)> 4 {spec.type4}")
    if spec.text_line == ATTACK_SPEED_TEXT_LINE:
        lines.append(
            f"SetCustomItemValueEx {equip} <$STR(N$XY_MG_CONTENT_PROP_ROW)> = "
            f"{ATTACK_SPEED_TEXT_LINE} <$STR({sum_var})> 0"
        )
    elif spec.value_command == "value_ex":
        lines.append(f"SetCustomItemValueEx {equip} <$STR(N$XY_MG_CONTENT_PROP_ROW)> = <$STR({sum_var})> 0 0")
    else:
        lines.append(f"SetCustomItemValue {equip} <$STR(N$XY_MG_CONTENT_PROP_ROW)> = <$STR({sum_var})>")
    lines.extend((
        "#IF", "EQUAL N$XY_MG_CONTENT_PROP_ROW 0", "#ACT",
        "MOV N$XY_MG_CONTENT_PROP_OVERFLOW 1",
    ))
    return tuple(lines)


_RUNTIME_VARIABLES = {
    33: "N$XY_MG_CONTENT_MONSTER_ABSORB",
    34: "N$XY_MG_CONTENT_TAO_PERCENT",
    35: "N$XY_MG_CONTENT_FATAL",
    36: "N$XY_MG_CONTENT_EXECUTION_BP",
    37: "N$XY_MG_CONTENT_CORPSE",
    38: "N$XY_MG_CONTENT_BLAST",
    39: "N$XY_MG_CONTENT_TAIL",
}
_SCRIPT_RUNTIME_VARIABLES = {
    50: "N$XY_MG_CONTENT_TOUGHNESS",
    51: "N$XY_MG_CONTENT_EXECUTION_BONUS",
    53: "N$XY_MG_CONTENT_DAMAGE_COEFF",
}


def _runtime_attributes(book: MinggeContentWorkbook, candidate: MinggeCandidate):
    result = []
    for attribute in book.attributes[candidate.candidate_id]:
        spec = _property_spec(attribute.property_name)
        if (
            spec.route == "textline_runtime" and spec.text_line != ATTACK_SPEED_TEXT_LINE
        ) or (spec.route == "script_bind" and spec.binding in _SCRIPT_RUNTIME_VARIABLES):
            result.append((attribute, spec))
    return tuple(result)


def _render_content_provider(book: MinggeContentWorkbook, line_by_id: dict[int, int]) -> str:
    aggregates = _content_aggregates(book)
    lines = ["{", "; XY-MG-CONTENT-PROVIDER-V1"]
    for candidate in book.candidates:
        lines.append(f"; CANDIDATE {candidate.candidate_id} TEXTVAR {line_by_id[candidate.candidate_id]}")
    lines.extend(("", "[@XY_MG_CONTENT_STATUS]", "#IF", "#ACT", "MOV N$XY_MG_API_STATUS 1", "BREAK", ""))
    lines.extend(("[@XY_MG_CONTENT_DRAW]", "#IF", "#ACT", "MOV N$XY_MG_API_STATUS 0"))
    groups: dict[tuple[int, int, int], list[MinggeCandidate]] = {}
    for candidate in book.candidates:
        if not candidate.enabled:
            continue
        groups.setdefault((candidate.pool_id, candidate.quality_id, TRUSTED_EQUIP_SLOTS[candidate.equip_part]), []).append(candidate)
    for (pool_id, quality_id, equip_slot), candidates in sorted(groups.items()):
        label = f"@XY_MG_CONTENT_POOL_{pool_id}_{quality_id}_{equip_slot}"
        lines.extend(("#IF", f"EQUAL N$XY_MG_API_POOL_ID {pool_id}", f"EQUAL N$XY_MG_API_QUALITY_ID {quality_id}", f"EQUAL N$XY_MG_API_EQUIP_SLOT {equip_slot}", "#ACT", f"GOTO {label}", "BREAK"))
    lines.extend(("#IF", "#ACT", "MESSAGEBOX 当前候选池或品质没有启用命格，未改变装备。", "BREAK", ""))
    for (pool_id, quality_id, equip_slot), candidates in sorted(groups.items()):
        pool_label = f"@XY_MG_CONTENT_POOL_{pool_id}_{quality_id}_{equip_slot}"
        lines.extend((f"[{pool_label}]", "#IF", "#ACT"))
        for index, candidate in enumerate(candidates):
            choose = f"@XY_MG_CONTENT_CHOOSE_{candidate.candidate_id}"
            remaining = len(candidates) - index
            if remaining == 1:
                lines.extend((f"GOTO {choose}", "BREAK"))
            else:
                lines.extend(("#IF", f"RANDOMEX 1 {remaining}", "#ACT", f"GOTO {choose}", "BREAK", "#IF", "#ACT"))
        lines.append("")
    for candidate in book.candidates:
        scheme = book.color_schemes[candidate.color_scheme_id]
        lines.extend((
            f"[@XY_MG_CONTENT_CHOOSE_{candidate.candidate_id}]", "#IF", "#ACT",
            f"MOV N$XY_MG_API_RESULT_ID {candidate.candidate_id}",
            f"MOV N$XY_MG_API_TEXT_LINE {line_by_id[candidate.candidate_id]}",
            f"MOV N$XY_MG_API_TEXT_COLOR {scheme.name_colors[0]}",
            f"MOV S$XY_MG_API_RESULT_TEXT {candidate.display_name}",
            "MOV S$XY_MG_PANEL_RESULT <$STR(S$XY_MG_API_RESULT_TEXT)>",
            "GOTO @XY_MG_CONTENT_WRITE", "BREAK", "",
        ))
    lines.extend(("[@XY_MG_CONTENT_WRITE]", "#IF", "#ACT"))
    for slot in range(1, 9):
        lines.extend(("#IF", f"EQUAL N$XY_MG_API_SLOT {slot}", "#ACT", f"GOTO @XY_MG_CONTENT_WRITE_SLOT_{slot}", "BREAK"))
    lines.extend(("#IF", "#ACT", "MESSAGEBOX 当前命格槽位无效，未改变装备。", "BREAK", ""))
    for slot in range(1, 9):
        row = 8 + slot
        lines.extend((
            f"[@XY_MG_CONTENT_WRITE_SLOT_{slot}]", "#IF", "#ACT",
            "LockUpdateItem <$STR(N$XY_MG_API_EQUIP_SLOT)>",
            f"SetCustomItemAbil <$STR(N$XY_MG_API_EQUIP_SLOT)> {row} 0 <$STR(N$XY_MG_API_TEXT_COLOR)>",
            f"SetCustomItemAbil <$STR(N$XY_MG_API_EQUIP_SLOT)> {row} 1 60",
            f"SetCustomItemAbil <$STR(N$XY_MG_API_EQUIP_SLOT)> {row} 2 {slot}",
            f"SetCustomItemAbil <$STR(N$XY_MG_API_EQUIP_SLOT)> {row} 3 0",
            f"SetCustomItemAbil <$STR(N$XY_MG_API_EQUIP_SLOT)> {row} 4 9",
            f"SetCustomItemValueEx <$STR(N$XY_MG_API_EQUIP_SLOT)> {row} = <$STR(N$XY_MG_API_TEXT_LINE)> <$STR(N$XY_MG_API_RESULT_ID)> 0",
            "GOTO @XY_MG_CONTENT_REBUILD", "BREAK", "",
        ))
    lines.extend(("[@XY_MG_CONTENT_READ]", "#IF", "#ACT"))
    for slot in range(1, 9):
        row = 8 + slot
        lines.append(f"GetCustomItemValueEx <$STR(N$XY_MG_API_EQUIP_SLOT)> {row} N$XY_MG_TMPP N$XY_MG_TEXT_{slot} N$XY_MG_ID_{slot} N$XY_MG_TMPC")
    lines.extend(("MOV N$XY_MG_API_STATUS 1", "BREAK", "", "[@XY_MG_CONTENT_REBUILD]", "#IF", "#ACT"))
    for index in range(1, len(aggregates) + 1):
        lines.append(f"MOV N$XY_MG_CONTENT_SUM{index}A 0")
    for slot in range(1, 9):
        row = 8 + slot
        lines.append(f"GetCustomItemValueEx <$STR(N$XY_MG_API_EQUIP_SLOT)> {row} N$XY_MG_CONTENT_TMPP N$XY_MG_CONTENT_TEXT_{slot} N$XY_MG_CONTENT_ID_{slot} N$XY_MG_CONTENT_TMPC")
    for slot in range(1, 9):
        for candidate in book.candidates:
            increments = _content_increment_lines(book, candidate, aggregates)
            if not increments:
                continue
            lines.extend(("#IF", f"EQUAL N$XY_MG_CONTENT_ID_{slot} {candidate.candidate_id}", "#ACT", *increments))
    equip = "<$STR(N$XY_MG_API_EQUIP_SLOT)>"
    for row in CONTENT_PROPERTY_ROWS:
        lines.extend(_clear_content_row(equip, row))
    lines.extend(("MOV N$XY_MG_CONTENT_PROP_INDEX 0", "MOV N$XY_MG_CONTENT_PROP_ROW 0", "MOV N$XY_MG_CONTENT_PROP_OVERFLOW 0"))
    for index, (spec, position) in enumerate(aggregates, start=1):
        lines.extend(_dynamic_content_property_lines(equip, index, spec, position))
    lines.extend((
        f"UpdateItem {equip}",
        "#IF", "EQUAL N$XY_MG_CONTENT_PROP_OVERFLOW 1", "#ACT",
        "MESSAGEBOX 当前灵玉实际同时存在的原生属性超过装备记录行容量；本次未静默截断，请调整候选组合。",
        "#IF", "#ACT",
        "GOTO @XY_MG_CONTENT_RECALC_ALL", "BREAK", "",
    ))

    lines.extend(("[@XY_MG_CONTENT_RECALC]", "#IF", "#ACT", "GOTO @XY_MG_CONTENT_RECALC_ALL", "BREAK", "", "[@XY_MG_CONTENT_RECALC_ALL]", "#IF", "#ACT"))
    for variable in (*_RUNTIME_VARIABLES.values(), *_SCRIPT_RUNTIME_VARIABLES.values()):
        lines.append(f"MOV {variable} 0")
    for equip_slot in sorted({TRUSTED_EQUIP_SLOTS[item.equip_part] for item in book.candidates}):
        for slot in range(1, 9):
            row = slot + 8
            lines.append(f"GetCustomItemValueEx {equip_slot} {row} N$XY_MG_CONTENT_RT_P N$XY_MG_CONTENT_RT_TEXT N$XY_MG_CONTENT_RT_ID_{equip_slot}_{slot} N$XY_MG_CONTENT_RT_C")
        candidates = [item for item in book.candidates if TRUSTED_EQUIP_SLOTS[item.equip_part] == equip_slot]
        for slot in range(1, 9):
            for candidate in candidates:
                runtime = _runtime_attributes(book, candidate)
                if not runtime:
                    continue
                lines.extend(("#IF", f"EQUAL N$XY_MG_CONTENT_RT_ID_{equip_slot}_{slot} {candidate.candidate_id}", "#ACT"))
                for attribute, spec in runtime:
                    variable = _RUNTIME_VARIABLES[spec.text_line] if spec.route == "textline_runtime" else _SCRIPT_RUNTIME_VARIABLES[spec.binding]
                    lines.append(f"INC {variable} {attribute.value}")
    lines.extend((
        "ChangeHumAbility 9 = 0", "ChangeHumAbility 10 = 0",
        "MOV N$XY_MG_CONTENT_TAO_LOW 0", "MOV N$XY_MG_CONTENT_TAO_HIGH 0",
        "CalcPercent <$SC> <$STR(N$XY_MG_CONTENT_TAO_PERCENT)> N$XY_MG_CONTENT_TAO_LOW",
        "CalcPercent <$MAXSC> <$STR(N$XY_MG_CONTENT_TAO_PERCENT)> N$XY_MG_CONTENT_TAO_HIGH",
        "ChangeHumAbility 9 + <$STR(N$XY_MG_CONTENT_TAO_LOW)>",
        "ChangeHumAbility 10 + <$STR(N$XY_MG_CONTENT_TAO_HIGH)>",
        "AddHumNewValue 22 = <$STR(N$XY_MG_CONTENT_FATAL)>",
    ))
    if any(
        _property_spec(attribute.property_name).text_line == ATTACK_SPEED_TEXT_LINE
        for values in book.attributes.values()
        for attribute in values
    ):
        lines.append(ATTACK_SPEED_CALL)
    lines.extend(("MOV N$XY_MG_API_STATUS 1", "BREAK", "}"))
    return "\n".join(lines)


def _attack_speed_dependency_blockers(root: Path, book: MinggeContentWorkbook) -> list[str]:
    if not any(
        _property_spec(attribute.property_name).text_line == ATTACK_SPEED_TEXT_LINE
        for values in book.attributes.values()
        for attribute in values
    ):
        return []
    blockers: list[str] = []
    core_path = _target(root, ATTACK_SPEED_CORE_RELATIVE)
    if not core_path.is_file():
        blockers.append("攻速突破依赖核心缺失")
    else:
        core = _script_text(core_path.read_bytes())
        if core.count("[@XY_AS_CAP_RECALC]") != 1 or "GetAllCustomItemValueByTextLine 60 -1 40 " not in core:
            blockers.append("攻速突破依赖核心缺少TextLine40统一重算接口")
    textvar_path = _target(root, TEXTVAR_RELATIVE)
    if not textvar_path.is_file():
        blockers.append("攻速突破依赖TextVar40文件缺失")
    else:
        text_lines = _script_text(textvar_path.read_bytes()).splitlines()
        if len(text_lines) < ATTACK_SPEED_TEXT_LINE or text_lines[ATTACK_SPEED_TEXT_LINE - 1].strip() != ATTACK_SPEED_TEXT:
            blockers.append("攻速突破依赖TextVar40合同缺失")
    qfunction_path = _target(root, QFUNCTION_RELATIVE)
    if not qfunction_path.is_file():
        blockers.append("攻速突破依赖QFunction缺失")
    else:
        qfunction = _script_text(qfunction_path.read_bytes())
        if qfunction.count("; XY-AS-CAP-V1-BEGIN") != 1 or qfunction.count("; XY-AS-CAP-V1-END") != 1:
            blockers.append("攻速突破依赖QFunction受管接口缺失或重复")
    return blockers


def _content_bridge(text: str) -> tuple[str, list[str], list[str]]:
    specs = (
        ("; XY-MG-CONTENT-STATUS-ANCHOR", CONTENT_MARKER, ("#CALL [\\玄渊命格\\命格内容提供者.txt] @XY_MG_CONTENT_STATUS",)),
        ("; XY-MG-CONTENT-DRAW-ANCHOR", CONTENT_MARKER + "-DRAW", ("#CALL [\\玄渊命格\\命格内容提供者.txt] @XY_MG_CONTENT_DRAW",)),
        ("; XY-MG-CONTENT-READ-ANCHOR", CONTENT_MARKER + "-READ", ("#CALL [\\玄渊命格\\命格内容提供者.txt] @XY_MG_CONTENT_READ",)),
        ("; XY-MG-CONTENT-RECALC-ANCHOR", CONTENT_MARKER + "-RECALC", ("#CALL [\\玄渊命格\\命格内容提供者.txt] @XY_MG_CONTENT_RECALC",)),
    )
    markers: list[str] = []
    blocks: list[str] = []
    for anchor, marker, lines in specs:
        text, block = _upsert_after_anchor(text, anchor, marker, lines)
        markers.append(marker)
        blocks.append(block)
    return text, markers, blocks


def _qfunction_content(text: str, book: MinggeContentWorkbook) -> tuple[str, list[str], list[str]]:
    marker = "XY-MG-CONTENT-QFUNCTION-V1"
    block = _managed_block(marker, (
        "[@XY_MG_CONTENT_RECALC_API]", "#IF", "#ACT",
        "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_RECALC", "BREAK",
    ))
    found = _find_block(text, marker)
    if found:
        start, end = found
        text = text[:start] + block + text[end:]
    else:
        text = text.rstrip("\n") + ("\n\n" if text.strip() else "") + block + "\n"
    markers = [marker]
    blocks = [block]
    runtime = [
        (attribute, _property_spec(attribute.property_name))
        for values in book.attributes.values() for attribute in values
        if _property_spec(attribute.property_name).route == "textline_runtime"
        or (_property_spec(attribute.property_name).route == "script_bind" and _property_spec(attribute.property_name).binding in _SCRIPT_RUNTIME_VARIABLES)
    ]
    if not runtime:
        return text, markers, blocks
    for event in ("PlayLogin", "TakeOnEx", "TakeOffEx"):
        event_marker = f"XY-MG-CONTENT-RECALC-{event.upper()}-V1"
        text, event_block = _upsert_after_label(
            text, event, event_marker,
            ("#IF", "#ACT", "#CALL [\\玄渊命格\\命格公共接口.txt] @XY_MG_API_RECALC"),
        )
        markers.append(event_marker)
        blocks.append(event_block)
    text_lines = {spec.text_line for _attribute, spec in runtime if spec.route == "textline_runtime"}
    bindings = {spec.binding for _attribute, spec in runtime if spec.route == "script_bind"}
    anchor_specs = (
        (33, None, "; XY_EQUIP_MAKER_MONSTER_ABSORB_ANCHOR", "MONSTER_ABSORB", ("#IF", "#ACT", "INC N$XY_MDA_Raw <$STR(N$XY_MG_CONTENT_MONSTER_ABSORB)>")),
        (36, None, "; XY_EXECUTION_LAB_CHANCE_ANCHOR", "EXECUTION", ("#IF", "#ACT", "MOV N$XY_MG_CONTENT_EVT <$STR(N$XY_MG_CONTENT_EXECUTION_BP)>", "MUL N$XY_MG_CONTENT_EVT 100", "INC N$XY_EXEC_ChanceBP <$STR(N$XY_MG_CONTENT_EVT)>")),
        (37, None, "; XY_EQUIP_MAKER_CORPSE_ANCHOR", "CORPSE", ("#IF", "#ACT", "INC N$XY_CorpseRate <$STR(N$XY_MG_CONTENT_CORPSE)>")),
        (38, None, "; XY_EQUIP_MAKER_RUNTIME_BLAST_ANCHOR", "BLAST", ("#IF", "#ACT", "INC N$XY_RT_Blast <$STR(N$XY_MG_CONTENT_BLAST)>")),
        (39, None, "; XY_EQUIP_MAKER_KILL_RESET", "TAIL", ("#IF", "#ACT", "INC N$XY_TailKillRate <$STR(N$XY_MG_CONTENT_TAIL)>")),
        (None, 50, "; XY_EXECUTION_LAB_TOUGHNESS_ANCHOR", "TOUGHNESS", ("#IF", "#ACT", "INC N$XY_EXEC_Toughness <$STR(N$XY_MG_CONTENT_TOUGHNESS)>")),
        (None, 51, "; XY_EXECUTION_LAB_PVE_BONUS_ANCHOR", "EXECUTION_BONUS", ("#IF", "#ACT", "INC N$XY_EXEC_PVEEquipBonusPercent <$STR(N$XY_MG_CONTENT_EXECUTION_BONUS)>")),
        (None, 53, "; XY_EQUIP_MAKER_RUNTIME_DAMAGE_COEFFICIENT_ANCHOR", "DAMAGE_COEFFICIENT", ("#IF", "#ACT", "INC N$XY_RT_DamageCoeff <$STR(N$XY_MG_CONTENT_DAMAGE_COEFF)>")),
    )
    for text_line, binding, anchor, suffix, script_lines in anchor_specs:
        if (text_line is not None and text_line not in text_lines) or (binding is not None and binding not in bindings):
            continue
        runtime_marker = f"XY-MG-CONTENT-RUNTIME-{suffix}-V1"
        text, runtime_block = _upsert_after_anchor(text, anchor, runtime_marker, script_lines)
        markers.append(runtime_marker)
        blocks.append(runtime_block)
    return text, markers, blocks


def _target(server: Path, relative: Path) -> Path:
    root = Path(server).resolve()
    if root.is_symlink():
        raise MinggeDualError("服务端根目录不能是路径链接")
    target = (root / relative).resolve()
    try:
        target.relative_to(root)
    except ValueError as exc:
        raise MinggeDualError(f"目标路径越界：{relative}") from exc
    return target


def _change(server: Path, relative: Path, kind: str, before: bytes | None, after: bytes, metadata: dict[str, Any] | None = None) -> DualFileChange:
    return DualFileChange(str(relative).replace("\\", "/"), _target(server, relative), kind, before, after, _sha256(before), _sha256(after), metadata or {})


def _plan_id(route: str, workbook_hash: str, changes: Iterable[DualFileChange]) -> str:
    payload = route + workbook_hash + "".join(item.after_sha256 for item in changes)
    return hashlib.sha256(payload.encode("ascii")).hexdigest()[:24]


def _is_exact_legacy_p3_textvar(candidate_id: int, line_no: int, text: str) -> bool:
    expected = LEGACY_P3_TEXTVAR_SHA256.get(candidate_id)
    if expected is None or expected[0] != line_no:
        return False
    return hashlib.sha256(text.encode("gb18030")).hexdigest().upper() == expected[1]


class _DualService:
    route: str
    package_id: str

    def __init__(self, platform_root: Path):
        self.platform_root = Path(platform_root).resolve()

    @property
    def backup_root(self) -> Path:
        return self.platform_root / "backups" / self.route

    def _validate_install_inputs(self, plan: DualInstallPlan) -> None:
        """Specialized services may bind additional read-only dependencies to a plan."""

    @staticmethod
    def _atomic_write(path: Path, data: bytes, expected: bytes | None) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temp = path.with_name(path.name + f".{uuid4().hex}.xydp-new")
        try:
            with temp.open("xb") as stream:
                stream.write(data)
                stream.flush()
                os.fsync(stream.fileno())
            current = path.read_bytes() if path.is_file() else None
            if current != expected:
                raise MinggeDualError(f"原子提交前目标文件已变化：{path}")
            os.replace(temp, path)
        finally:
            if temp.exists():
                temp.unlink()

    @classmethod
    def _recover_committed(cls, committed) -> list[str]:
        failures: list[str] = []
        for path, before, after in reversed(committed):
            try:
                current = path.read_bytes() if path.is_file() else None
                if current != after:
                    failures.append(f"保留其他写者的新字节：{path}")
                    continue
                if before is None:
                    path.unlink()
                else:
                    cls._atomic_write(path, before, after)
            except (OSError, MinggeDualError) as exc:
                failures.append(f"恢复失败：{path}: {exc}")
        return failures

    def install(self, plan: DualInstallPlan) -> DualInstallReceipt:
        from .target_lock import target_lock
        with target_lock(plan.server_root):
            return self._install_locked(plan)

    def _install_locked(self, plan: DualInstallPlan) -> DualInstallReceipt:
        if plan.route != self.route or plan.package_id != self.package_id:
            raise MinggeDualError("安装计划路由与服务不匹配")
        if plan.blockers:
            raise MinggeDualError("安装被阻止：" + "；".join(plan.blockers))
        self._validate_install_inputs(plan)
        if _sha256(plan.workbook.read_bytes()) != plan.workbook_sha256:
            raise MinggeDualError("预检后配置表已变化，请重新预检")
        seen: set[Path] = set()
        for change in plan.changes:
            if _target(plan.server_root, Path(change.relative_path)) != change.path.resolve() or change.path in seen:
                raise MinggeDualError("安装计划目标越界或重复")
            seen.add(change.path)
            if _sha256(change.before) != change.before_sha256 or _sha256(change.after) != change.after_sha256:
                raise MinggeDualError("安装计划内容哈希不一致")
            current = change.path.read_bytes() if change.path.is_file() else None
            if _sha256(current) != change.before_sha256:
                raise MinggeDualError(f"预检后目标文件已变化：{change.relative_path}")
        transaction_id = datetime.now().strftime("%Y%m%d_%H%M%S_") + uuid4().hex[:8]
        transaction = self.backup_root / transaction_id
        transaction.mkdir(parents=True, exist_ok=False)
        receipt_changes: list[dict[str, Any]] = []
        committed = []
        try:
            # Back up every original before the first target mutation.
            for index, change in enumerate(plan.changes):
                backup_file = None
                if change.before is not None:
                    backup_file = transaction / "files" / f"{index:03d}.bin"
                    backup_file.parent.mkdir(parents=True, exist_ok=True)
                    backup_file.write_bytes(change.before)
                receipt_changes.append({
                    "relative_path": change.relative_path, "kind": change.kind,
                    "before_sha256": change.before_sha256, "after_sha256": change.after_sha256,
                    "backup_file": str(backup_file.relative_to(transaction)) if backup_file else None,
                    "metadata": change.metadata,
                })
            for change in plan.changes:
                current = change.path.read_bytes() if change.path.is_file() else None
                if _sha256(current) != change.before_sha256:
                    raise MinggeDualError(f"提交前目标文件已变化：{change.relative_path}")
                self._atomic_write(change.path, change.after, change.before)
                committed.append((change.path, change.before, change.after))
                if _sha256(change.path.read_bytes()) != change.after_sha256:
                    raise MinggeDualError(f"安装后哈希不匹配：{change.relative_path}")
            # Detect non-cooperating edits to a previously committed target too.
            for path, _, after in committed:
                if not path.is_file() or path.read_bytes() != after:
                    raise MinggeDualError(f"提交期间目标文件已变化：{path}")
            payload = {
                "schema_version": 1, "operation": self.route, "package_id": self.package_id,
                "transaction_id": transaction_id, "status": "installed-pending-game-verification",
                "server": str(plan.server_root), "workbook": str(plan.workbook),
                "workbook_sha256": plan.workbook_sha256, "plan_id": plan.plan_id,
                "changes": receipt_changes,
            }
            receipt_path = transaction / "receipt.json"
            receipt_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
            return DualInstallReceipt(transaction_id, payload["status"], receipt_path, tuple(item.path for item in plan.changes))
        except Exception as exc:
            failures = self._recover_committed(committed)
            if failures:
                (transaction / "recovery_conflicts.json").write_text(json.dumps(failures, ensure_ascii=False, indent=2), encoding="utf-8")
                raise MinggeDualError("安装失败且恢复需要人工处理：" + "；".join(failures)) from exc
            raise

    def rollback(self, server: Path, transaction_id: str) -> DualInstallReceipt:
        from .target_lock import target_lock
        with target_lock(server):
            return self._rollback_locked(server, transaction_id)

    def _rollback_locked(self, server: Path, transaction_id: str) -> DualInstallReceipt:
        if not re.fullmatch(r"\d{8}_\d{6}_[0-9a-f]{8}", transaction_id):
            raise MinggeDualError("事务号格式无效")
        transaction = self.backup_root / transaction_id
        receipt_path = transaction / "receipt.json"
        if not receipt_path.is_file():
            raise MinggeDualError(f"事务收据不存在：{transaction_id}")
        payload = json.loads(receipt_path.read_text(encoding="utf-8"))
        if payload.get("operation") != self.route or payload.get("package_id") != self.package_id:
            raise MinggeDualError("事务不属于当前子包")
        if Path(payload.get("server", "")).resolve() != Path(server).resolve():
            raise MinggeDualError("事务目标服务端不匹配")
        prepared = []
        seen: set[Path] = set()
        # Compute and validate ALL rollback outputs before changing any target.
        for item in reversed(payload.get("changes", [])):
            relative = Path(item["relative_path"])
            path = _target(server, relative)
            if path in seen:
                raise MinggeDualError("回滚事务包含重复目标")
            seen.add(path)
            kind = item["kind"]
            metadata = item.get("metadata") or {}
            if kind == "shared_base":
                continue
            if not path.is_file():
                raise MinggeDualError(f"回滚目标缺失：{relative}")
            current = path.read_bytes()
            before = None
            if item.get("backup_file"):
                backup = _target(transaction, Path(item["backup_file"]))
                before = backup.read_bytes()
                if _sha256(before) != item["before_sha256"]:
                    raise MinggeDualError(f"回滚备份哈希不匹配：{relative}")
            elif item.get("before_sha256"):
                raise MinggeDualError(f"回滚备份缺失：{relative}")
            if kind == "exclusive":
                if _sha256(current) != item["after_sha256"]:
                    raise MinggeDualError(f"回滚前文件发生未知变化：{relative}")
                restored = before
            elif _sha256(current) == item["after_sha256"] and before is not None:
                restored = before
            elif kind == "managed_blocks":
                text = _script_text(current)
                if len(metadata["markers"]) != len(metadata["blocks"]):
                    raise MinggeDualError(f"回滚受管块元数据不一致：{relative}")
                for marker, block in reversed(list(zip(metadata["markers"], metadata["blocks"]))):
                    text = _remove_block(text, marker, block)
                restored = _script_bytes(text)
            elif kind == "indexed_lines":
                lines = _script_text(current).splitlines()
                start = int(metadata["start"])
                after_lines = list(metadata["after_lines"])
                if start < 1 or lines[start - 1:start - 1 + len(after_lines)] != after_lines:
                    raise MinggeDualError(f"回滚前TextVar受管行发生变化：{relative}")
                lines[start - 1:start - 1 + len(after_lines)] = list(metadata["before_lines"])
                restored = _script_bytes("\n".join(lines))
            else:
                raise MinggeDualError(f"未知回滚类型：{kind}")
            prepared.append((path, current, restored))
        committed = []
        try:
            for path, current, restored in prepared:
                if not path.is_file() or path.read_bytes() != current:
                    raise MinggeDualError(f"回滚提交前文件已变化：{path}")
                if restored is None:
                    path.unlink()
                else:
                    self._atomic_write(path, restored, current)
                committed.append((path, current, restored))
            for path, _, restored in committed:
                if (path.read_bytes() if path.is_file() else None) != restored:
                    raise MinggeDualError(f"回滚提交期间目标已变化：{path}")
            rollback_path = transaction / "rollback.json"
            rollback_path.write_text(json.dumps({
                "schema_version": 1, "operation": self.route + "-rollback", "transaction_id": transaction_id,
                "status": "rolled-back", "affected_paths": [str(path) for path, _, _ in prepared],
            }, ensure_ascii=False, indent=2), encoding="utf-8")
            return DualInstallReceipt(transaction_id, "rolled-back", rollback_path, tuple(path for path, _, _ in prepared))
        except Exception as exc:
            failures = self._recover_committed(committed)
            if failures:
                (transaction / "rollback_recovery_conflicts.json").write_text(json.dumps(failures, ensure_ascii=False, indent=2), encoding="utf-8")
                raise MinggeDualError("回滚失败且恢复需要人工处理：" + "；".join(failures)) from exc
            raise


class MinggeNpcService(_DualService):
    route = NPC_ROUTE
    package_id = NPC_PACKAGE_ID

    def preflight(self, server: Path, workbook: Path) -> DualInstallPlan:
        blockers: list[str] = []
        warnings = ["NPC规则包独立安装；命格内容包缺失时洗练会在扣费前停止。", "当前服仅作金样本，本计划不操作引擎。"]
        changes: list[DualFileChange] = []
        source = Path(workbook).resolve()
        workbook_hash = _sha256(source.read_bytes()) if source.is_file() else ""
        try:
            root = Path(server).resolve()
            if not (root / "Mir200/Envir").is_dir():
                raise MinggeDualError("目标目录不是可识别的翎风/LFM2服务端")
            book = load_mingge_npc_workbook(source)
            bridge_path = _target(root, BRIDGE_RELATIVE)
            bridge_before, bridge_text = _bridge_before(bridge_path)
            bridge_after = _script_bytes(bridge_text)
            if bridge_before != bridge_after:
                changes.append(_change(root, BRIDGE_RELATIVE, "shared_base", bridge_before, bridge_after))
            rules_path = _target(root, NPC_RULES_RELATIVE)
            rules_before = rules_path.read_bytes() if rules_path.is_file() else None
            rules_after = _script_bytes(_render_npc_rules(book))
            changes.append(_change(root, NPC_RULES_RELATIVE, "exclusive", rules_before, rules_after))
            npc_relative = Path(book.settings.npc_script_relative)
            npc_path = _target(root, npc_relative)
            if not npc_path.is_file():
                raise MinggeDualError(f"现有NPC脚本不存在：{book.settings.npc_script_relative}")
            npc_before = npc_path.read_bytes()
            npc_text = _script_text(npc_before)
            npc_after_text, block = _upsert_after_label(
                npc_text, book.settings.npc_hook_label, NPC_MARKER,
                ("#CALL [\\玄渊命格\\命格NPC规则.txt] @XY_MG_NPC_MAIN",),
            )
            npc_after = _script_bytes(npc_after_text)
            changes.append(_change(root, npc_relative, "managed_blocks", npc_before, npc_after, {"markers": [NPC_MARKER], "blocks": [block]}))
            changes.append(_patch_existing_npc_actions(root, npc_text, book))
        except (OSError, MinggeDualError, ValueError) as exc:
            blockers.append(str(exc))
        plan_id = _plan_id(self.route, workbook_hash, changes)
        return DualInstallPlan(self.route, self.package_id, plan_id, Path(server).resolve(), source, workbook_hash, tuple(changes), tuple(blockers), tuple(warnings))


class MinggeContentService(_DualService):
    route = CONTENT_ROUTE
    package_id = CONTENT_PACKAGE_ID

    def preflight(self, server: Path, workbook: Path) -> DualInstallPlan:
        blockers: list[str] = []
        warnings = ["内容包独立安装；未安装NPC规则包时只提供内容接口和属性重算入口。", "真实游戏验收前状态保持candidate。"]
        changes: list[DualFileChange] = []
        previews: list[tuple[str, tuple[tuple[str, int], ...]]] = []
        source = Path(workbook).resolve()
        workbook_hash = _sha256(source.read_bytes()) if source.is_file() else ""
        try:
            root = Path(server).resolve()
            if not (root / "Mir200/Envir").is_dir():
                raise MinggeDualError("目标目录不是可识别的翎风/LFM2服务端")
            book = load_mingge_content_workbook(source)
            static = [attr for values in book.attributes.values() for attr in values if attr.route == "static_only"]
            if static:
                names = sorted({item.property_name for item in static})
                raise MinggeDualError("以下属性只有静态目录、没有实例实效出口，禁止伪装生效：" + "、".join(names))
            speed_blockers = _attack_speed_dependency_blockers(root, book)
            if speed_blockers:
                raise MinggeDualError("；".join(speed_blockers))
            bridge_path = _target(root, BRIDGE_RELATIVE)
            bridge_before, bridge_text = _bridge_before(bridge_path)
            bridge_with_content, markers, blocks = _content_bridge(bridge_text)
            bridge_after = _script_bytes(bridge_with_content)
            changes.append(_change(root, BRIDGE_RELATIVE, "managed_blocks", bridge_before, bridge_after, {"markers": markers, "blocks": blocks}))

            textvar_path = _target(root, TEXTVAR_RELATIVE)
            if not textvar_path.is_file():
                raise MinggeDualError("CustomItemPropertyTextVarList.txt不存在")
            textvar_before = textvar_path.read_bytes()
            text_lines = _script_text(textvar_before).splitlines()

            provider_path = _target(root, CONTENT_PROVIDER_RELATIVE)
            provider_before = provider_path.read_bytes() if provider_path.is_file() else None
            line_by_id: dict[int, int] = {}
            previously_owned_lines: set[int] = set()
            if provider_before:
                for candidate_id, line in re.findall(r"(?m)^; CANDIDATE (\d+) TEXTVAR (\d+)\s*$", _script_text(provider_before)):
                    line_by_id[int(candidate_id)] = int(line)
                    previously_owned_lines.add(int(line))
            next_line = CONTENT_TEXTVAR_START
            skipped_occupied_lines: list[int] = []
            for candidate in book.candidates:
                if candidate.candidate_id in line_by_id:
                    continue
                while True:
                    current = text_lines[next_line - 1] if next_line <= len(text_lines) else ""
                    if next_line not in line_by_id.values() and (
                        not current.strip()
                        or _is_exact_legacy_p3_textvar(candidate.candidate_id, next_line, current)
                    ):
                        line_by_id[candidate.candidate_id] = next_line
                        next_line += 1
                        break
                    if current.strip() and next_line not in skipped_occupied_lines:
                        skipped_occupied_lines.append(next_line)
                    next_line += 1
            provider_after = _script_bytes(_render_content_provider(book, line_by_id))
            changes.append(_change(root, CONTENT_PROVIDER_RELATIVE, "exclusive", provider_before, provider_after))
            if skipped_occupied_lines:
                warnings.append(
                    "TextVar已自动跳过其他系统占用行："
                    + "、".join(str(line) for line in skipped_occupied_lines)
                )
            highest = max(line_by_id.values())
            while len(text_lines) < highest:
                text_lines.append("")
            before_lines: list[str] = []
            after_lines: list[str] = []
            adopted_legacy_lines: list[int] = []
            ordered = sorted(book.candidates, key=lambda item: line_by_id[item.candidate_id])
            managed_lines = previously_owned_lines | {line_by_id[item.candidate_id] for item in ordered}
            start = min(managed_lines)
            end = max(managed_lines)
            by_line = {line_by_id[item.candidate_id]: item for item in ordered}
            for line_no in range(start, end + 1):
                before_lines.append(text_lines[line_no - 1])
                candidate = by_line.get(line_no)
                if candidate is None:
                    after_lines.append("" if line_no in previously_owned_lines else text_lines[line_no - 1])
                    continue
                if line_no not in previously_owned_lines and text_lines[line_no - 1].strip():
                    if _is_exact_legacy_p3_textvar(candidate.candidate_id, line_no, text_lines[line_no - 1]):
                        adopted_legacy_lines.append(line_no)
                    else:
                        raise MinggeDualError(f"TextVar第{line_no}行已被其他系统占用，禁止覆盖")
                segments = _candidate_segments(book, candidate)
                rendered = _candidate_text(segments)
                rendered_bytes = len(rendered.encode("gb18030"))
                if rendered_bytes > 128:
                    warnings.append(
                        f"命格{candidate.candidate_id}的TextVar多色文本为{rendered_bytes}字节；"
                        "128字符限制缺少当前引擎实证，已改为显示验收警告。"
                    )
                after_lines.append(rendered)
                previews.append((candidate.display_name, segments))
            text_lines[start - 1:end] = after_lines
            textvar_after = _script_bytes("\n".join(text_lines))
            changes.append(_change(root, TEXTVAR_RELATIVE, "indexed_lines", textvar_before, textvar_after, {"start": start, "before_lines": before_lines, "after_lines": after_lines, "adopted_legacy_lines": adopted_legacy_lines}))
            if adopted_legacy_lines:
                warnings.append("已按当前服精确旧内容接管旧P3命格TextVar第33至38行；其他占用仍会阻止覆盖。")

            qfunction_path = _target(root, QFUNCTION_RELATIVE)
            if not qfunction_path.is_file():
                raise MinggeDualError("QFunction-0.txt不存在")
            q_before = qfunction_path.read_bytes()
            q_after_text, q_markers, q_blocks = _qfunction_content(_script_text(q_before), book)
            q_after = _script_bytes(q_after_text)
            changes.append(_change(root, QFUNCTION_RELATIVE, "managed_blocks", q_before, q_after, {"markers": q_markers, "blocks": q_blocks}))
        except (OSError, MinggeDualError, ValueError) as exc:
            blockers.append(str(exc))
        plan_id = _plan_id(self.route, workbook_hash, changes)
        return DualInstallPlan(self.route, self.package_id, plan_id, Path(server).resolve(), source, workbook_hash, tuple(changes), tuple(blockers), tuple(warnings), tuple(previews))


__all__ = [
    "API_VERSION", "BRIDGE_RELATIVE", "CONTENT_PACKAGE_ID", "CONTENT_PROVIDER_RELATIVE",
    "CONTENT_QFUNCTION_MARKER", "CONTENT_ROUTE", "CONTENT_TEXTVAR_START", "DualInstallPlan",
    "DualInstallReceipt", "MinggeContentService", "MinggeDualError", "MinggeNpcService",
    "NPC_PACKAGE_ID", "NPC_ROUTE", "NPC_RULES_RELATIVE", "load_mingge_content_workbook",
    "load_mingge_npc_workbook",
]
