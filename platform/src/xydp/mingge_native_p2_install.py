"""命格38号候选的P2三属性隔离测试安装链。

P2只允许固定的鞭尸灵玉测试候选：防御+5、生命值+100、魔法值+100。
安装器只写三个脚本配置文件；玩家执行命令后，M2会修改当前测试物品实例。
"""

from __future__ import annotations

import hashlib
import hmac
import json
import os
import stat
import subprocess
import tempfile
import uuid
from base64 import b64decode, b64encode
from contextlib import contextmanager
from ctypes import POINTER, Structure, WinError, byref, c_char, c_longlong, c_void_p, cast, create_string_buffer, sizeof, windll
from ctypes.wintypes import BOOL, DWORD, HANDLE, LPCWSTR
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path

from .mingge_native import NativeCandidate, compile_native_candidate, read_native_candidate_workbook


COMMAND_NAME = "命格平台三属性验收"
COMMAND_NUMBER = 98
MANAGED_BEGIN = "; XY-MG-NATIVE-P2-BEGIN"
MANAGED_END = "; XY-MG-NATIVE-P2-END"
TEST_SCRIPT_RELATIVE = "Mir200/Envir/QuestDiary/玄渊验收/命格平台三属性验收.txt"
USERCMD_RELATIVE = "Mir200/Envir/UserCmd.txt"
QFUNCTION_RELATIVE = "Mir200/Envir/Market_Def/QFunction-0.txt"
TARGET_RELATIVES = (USERCMD_RELATIVE, QFUNCTION_RELATIVE, TEST_SCRIPT_RELATIVE)
EXPECTED_CANDIDATE_ID = "tiebi-vitality-p2"
EXPECTED_TARGET_ITEM = "鞭尸灵玉"
EXPECTED_MINGGE_NAME = "【三元命格·测试】"
ENGINE_VERSION_BASELINE = "LFM2-20260707"
EXPECTED_M2SERVER_VERSION = "4.0.0.0"
EXPECTED_M2SERVER_SHA256 = "6b23758a704d29ab57da0f9f673a64854dcbd531ce9d6c36741702bc678a8f27"
EXPECTED_CANDIDATE_SPEC_SHA256 = "fc80d75fb151974cc1155acd11cb4b90a662d864092609bc19c0e1aa32e85c52"
EXPECTED_COMPILED_SCRIPT_SHA256 = "a4cd8376bb5a341ab1ac254c5926246a085084e34118fcc72949db9157ac260f"
MANIFEST_DOMAIN = b"XYDP-MG-P2-MANIFEST\0"
JOURNAL_DOMAIN = b"XYDP-MG-P2-JOURNAL\0"
ACTIVE_DOMAIN = b"XYDP-MG-P2-ACTIVE\0"
COORDINATION_DOMAIN = b"XYDP-MG-P2-COORDINATION\0"
CRYPTPROTECT_LOCAL_MACHINE = 0x4
WAIT_OBJECT_0 = 0
WAIT_ABANDONED = 0x80
WAIT_TIMEOUT = 0x102
GENERIC_READ = 0x80000000
GENERIC_WRITE = 0x40000000
DELETE_ACCESS = 0x00010000
FILE_SHARE_READ = 0x00000001
FILE_SHARE_WRITE = 0x00000002
CREATE_NEW = 1
OPEN_EXISTING = 3
FILE_FLAG_OPEN_REPARSE_POINT = 0x00200000
FILE_FLAG_BACKUP_SEMANTICS = 0x02000000
FILE_ATTRIBUTE_REPARSE_POINT = 0x00000400
FILE_ATTRIBUTE_DIRECTORY = 0x00000010
FILE_ATTRIBUTE_TAG_INFO_CLASS = 9
FILE_DISPOSITION_INFO_CLASS = 4
FILE_BEGIN = 0
INVALID_HANDLE_VALUE = HANDLE(-1).value


class NativeP2TestInstallError(ValueError):
    """P2隔离测试安装契约不满足。"""


@dataclass(frozen=True)
class NativeP2Evidence:
    property_name: str
    evidence_status: str
    source_id: str
    source_type: str
    source_locator: str
    snapshot_path: str
    captured_at: str
    source_sha256: str


P2_EVIDENCE = (
    NativeP2Evidence(
        "防御", "已实测确认", "LF-ITEM-002/P1-GAME-ACCEPTANCE", "game_acceptance",
        r"D:\codex交班记录\翎风引擎知识库\装备与物品\LF-ITEM-002_命格原生重建能力边界与阶段路线.txt",
        r"D:\codex交班记录\翎风引擎知识库\证据原件\网上资料\20260826_P2三属性\P1_LF-ITEM-002_验收快照_20260826.txt",
        "2026-08-26", "801A986C34D43A9D76586D07F25F462CB3D96FFDAFAA879B606FE0C155CB3500",
    ),
    NativeP2Evidence(
        "生命值", "网上资料待实测", "LFM2-OFFICIAL-CUSTOM-ABIL-BIND-6", "official_web",
        "https://www.5cq.com/help/翎风/topics/游戏引擎反外挂系统/游戏功能详解/自定义属性说明书.htm",
        r"D:\codex交班记录\翎风引擎知识库\证据原件\网上资料\20260826_P2三属性\官方_自定义属性说明书_20260826.htm",
        "2026-08-26", "42D317DE781D322A78885F75CA00944B12EE81786FF119C6815E0B78C8E4D87E",
    ),
    NativeP2Evidence(
        "魔法值", "网上资料待实测", "LFM2-OFFICIAL-CUSTOM-ABIL-BIND-7", "official_web",
        "https://www.5cq.com/help/翎风/topics/游戏引擎反外挂系统/游戏功能详解/自定义属性说明书.htm",
        r"D:\codex交班记录\翎风引擎知识库\证据原件\网上资料\20260826_P2三属性\官方_自定义属性说明书_20260826.htm",
        "2026-08-26", "42D317DE781D322A78885F75CA00944B12EE81786FF119C6815E0B78C8E4D87E",
    ),
)


@dataclass(frozen=True)
class NativeP2PlannedFile:
    relative_path: str
    path: Path
    existed_before: bool
    before_sha256: str | None
    after_sha256: str
    after: bytes


@dataclass(frozen=True)
class NativeP2InstallPlan:
    plan_id: str
    candidate_id: str
    candidate_spec_sha256: str
    compiled_script_sha256: str
    workbook: Path
    workbook_sha256: str
    server_root: Path
    engine_version: str
    engine_sha256: str
    files: tuple[NativeP2PlannedFile, ...]
    blockers: tuple[str, ...]
    warnings: tuple[str, ...]
    evidence_status: tuple[str, ...]
    evidence: tuple[NativeP2Evidence, ...]


@dataclass(frozen=True)
class NativeP2InstallReceipt:
    transaction_id: str
    status: str
    candidate_id: str
    receipt_path: Path
    files: tuple[str, ...]


@dataclass(frozen=True)
class NativeP2RollbackReceipt:
    transaction_id: str
    status: str
    receipt_path: Path


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _read(path: Path) -> bytes | None:
    return path.read_bytes() if path.is_file() else None


def _atomic_write(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    temporary_path = Path(temporary)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_path, path)
    except BaseException:
        temporary_path.unlink(missing_ok=True)
        raise


def _canonical_json_bytes(payload: object) -> bytes:
    return (json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode("utf-8")


class _DataBlob(Structure):
    _fields_ = (("cbData", DWORD), ("pbData", POINTER(c_char)))


def _blob_from_bytes(payload: bytes) -> tuple[_DataBlob, object]:
    buffer = create_string_buffer(payload)
    return _DataBlob(len(payload), cast(buffer, POINTER(c_char))), buffer


def _dpapi_protect(payload: bytes) -> bytes:
    source, source_buffer = _blob_from_bytes(payload)
    output = _DataBlob()
    ok = windll.crypt32.CryptProtectData(
        byref(source), "XYDP Mingge P2 HMAC", None, None, None,
        CRYPTPROTECT_LOCAL_MACHINE, byref(output),
    )
    _ = source_buffer
    if not ok:
        raise NativeP2TestInstallError("无法用Windows DPAPI保护P2事务密钥")
    try:
        return bytes((c_char * output.cbData).from_address(cast(output.pbData, c_void_p).value))
    finally:
        windll.kernel32.LocalFree(output.pbData)


def _dpapi_unprotect(payload: bytes) -> bytes:
    source, source_buffer = _blob_from_bytes(payload)
    output = _DataBlob()
    ok = windll.crypt32.CryptUnprotectData(
        byref(source), None, None, None, None, 0, byref(output),
    )
    _ = source_buffer
    if not ok:
        raise NativeP2TestInstallError("无法解封P2事务完整性密钥")
    try:
        return bytes((c_char * output.cbData).from_address(cast(output.pbData, c_void_p).value))
    finally:
        windll.kernel32.LocalFree(output.pbData)


def _secret_directory() -> Path:
    program_data = os.environ.get("PROGRAMDATA")
    if not program_data:
        raise NativeP2TestInstallError("系统没有PROGRAMDATA，无法建立P2事务密钥边界")
    return Path(program_data) / "XuanYuanDevPlatform" / "secrets"


def _acl_snapshot(path: Path) -> dict[str, object]:
    environment = os.environ.copy()
    environment["XYDP_P2_ACL_TARGET"] = str(path)
    windows_directory = Path(os.environ.get("WINDIR", r"C:\Windows"))
    program_files = Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
    environment["PSModulePath"] = os.pathsep.join((
        str(program_files / "WindowsPowerShell" / "Modules"),
        str(windows_directory / "System32" / "WindowsPowerShell" / "v1.0" / "Modules"),
    ))
    script = (
        "$ErrorActionPreference='Stop';"
        "$a=Get-Acl -LiteralPath $env:XYDP_P2_ACL_TARGET;"
        "$r=@($a.Access|ForEach-Object{[pscustomobject]@{Identity=$_.IdentityReference.Translate([System.Security.Principal.SecurityIdentifier]).Value;"
        "Type=$_.AccessControlType.ToString();Rights=$_.FileSystemRights.ToString();Inherited=$_.IsInherited}});"
        "[pscustomobject]@{Protected=$a.AreAccessRulesProtected;Rules=$r}|ConvertTo-Json -Compress -Depth 4"
    )
    completed = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True, text=True, encoding="utf-8-sig", errors="strict", env=environment, check=False,
    )
    if completed.returncode != 0:
        raise NativeP2TestInstallError(f"读取P2密钥ACL失败：{completed.stderr.strip()}")
    try:
        return json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        raise NativeP2TestInstallError("读取P2密钥ACL得到无效结果") from exc


def _verify_secret_acl(path: Path) -> None:
    snapshot = _acl_snapshot(path)
    rules = snapshot.get("Rules", [])
    if isinstance(rules, dict):
        rules = [rules]
    allowed = {"S-1-5-18", "S-1-5-32-544"}
    if snapshot.get("Protected") is not True or not isinstance(rules, list) or len(rules) != 2:
        raise NativeP2TestInstallError(f"P2密钥ACL不是仅SYSTEM和Administrators的受保护DACL：{path}")
    identities: set[str] = set()
    for rule in rules:
        if not isinstance(rule, dict) or rule.get("Type") != "Allow" or rule.get("Inherited") is not False:
            raise NativeP2TestInstallError(f"P2密钥ACL含继承、拒绝或无效ACE：{path}")
        identity = str(rule.get("Identity", ""))
        if identity not in allowed or str(rule.get("Rights", "")) != "FullControl":
            raise NativeP2TestInstallError(f"P2密钥ACL含额外主体或非完全控制ACE：{path}")
        identities.add(identity)
    if len(identities) != 2:
        raise NativeP2TestInstallError(f"P2密钥ACL主体不唯一：{path}")


def _restrict_secret_acl(path: Path) -> None:
    environment = os.environ.copy()
    environment["XYDP_P2_ACL_TARGET"] = str(path)
    windows_directory = Path(os.environ.get("WINDIR", r"C:\Windows"))
    program_files = Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
    environment["PSModulePath"] = os.pathsep.join((
        str(program_files / "WindowsPowerShell" / "Modules"),
        str(windows_directory / "System32" / "WindowsPowerShell" / "v1.0" / "Modules"),
    ))
    script = (
        "$ErrorActionPreference='Stop';"
        "$isDir=Test-Path -LiteralPath $env:XYDP_P2_ACL_TARGET -PathType Container;"
        "$acl=if($isDir){New-Object System.Security.AccessControl.DirectorySecurity}else{New-Object System.Security.AccessControl.FileSecurity};"
        "$acl.SetAccessRuleProtection($true,$false);"
        "$inherit=if($isDir){[System.Security.AccessControl.InheritanceFlags]'ContainerInherit,ObjectInherit'}else{[System.Security.AccessControl.InheritanceFlags]::None};"
        "foreach($sidText in @('S-1-5-18','S-1-5-32-544')){"
        "$sid=[System.Security.Principal.SecurityIdentifier]::new($sidText);"
        "$rule=[System.Security.AccessControl.FileSystemAccessRule]::new($sid,[System.Security.AccessControl.FileSystemRights]::FullControl,$inherit,[System.Security.AccessControl.PropagationFlags]::None,[System.Security.AccessControl.AccessControlType]::Allow);"
        "$acl.AddAccessRule($rule)|Out-Null};"
        "Set-Acl -LiteralPath $env:XYDP_P2_ACL_TARGET -AclObject $acl"
    )
    completed = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", script],
        capture_output=True, text=True, encoding="utf-8-sig", errors="strict", env=environment, check=False,
    )
    if completed.returncode != 0:
        raise NativeP2TestInstallError(f"P2事务密钥ACL收紧失败：{completed.stderr.strip()}")
    _verify_secret_acl(path)


def _assert_secret_path_no_reparse(path: Path) -> None:
    program_data = Path(os.environ.get("PROGRAMDATA", "")).absolute()
    if not str(program_data) or not path.absolute().is_relative_to(program_data):
        raise NativeP2TestInstallError("P2事务密钥路径不在PROGRAMDATA内")
    current = program_data
    if current.exists() and _is_reparse_point(current):
        raise NativeP2TestInstallError(f"P2密钥路径包含重解析点：{current}")
    for part in path.absolute().relative_to(program_data).parts:
        current = current / part
        if current.exists() and _is_reparse_point(current):
            raise NativeP2TestInstallError(f"P2密钥路径包含重解析点：{current}")


def _integrity_key_path() -> Path:
    return _secret_directory() / "mingge-p2-hmac.key.dpapi"


def _load_or_create_integrity_key_held(base: Path) -> tuple[str, bytes]:
    path = _integrity_key_path()
    if path.is_file():
        _assert_secret_path_no_reparse(path)
        _verify_secret_acl(path.parent)
        _verify_secret_acl(path)
        try:
            envelope = json.loads(path.read_text(encoding="utf-8"))
            key = _dpapi_unprotect(b64decode(envelope["protected_key"]))
        except (KeyError, ValueError, json.JSONDecodeError) as exc:
            raise NativeP2TestInstallError("P2事务密钥文件损坏") from exc
        key_id = hashlib.sha256(key).hexdigest()[:16]
        if envelope.get("key_id") != key_id or len(key) != 32:
            raise NativeP2TestInstallError("P2事务密钥ID或长度不匹配")
        return key_id, key
    open_transactions = [item for item in base.iterdir() if item.is_dir() and (item / "manifest.json").exists()] if base.exists() else []
    if open_transactions:
        raise NativeP2TestInstallError("P2事务密钥缺失且已有历史事务，禁止自动重建")
    directory_existed = path.parent.exists()
    path.parent.mkdir(parents=True, exist_ok=True)
    created_key = False
    try:
        _assert_secret_path_no_reparse(path.parent)
        _restrict_secret_acl(path.parent)
        key = os.urandom(32)
        key_id = hashlib.sha256(key).hexdigest()[:16]
        envelope = {"schema_version": 1, "key_id": key_id, "protected_key": b64encode(_dpapi_protect(key)).decode("ascii")}
        _atomic_write(path, _canonical_json_bytes(envelope))
        created_key = True
        _restrict_secret_acl(path)
        _verify_secret_acl(path.parent)
        _verify_secret_acl(path)
        return key_id, key
    except BaseException:
        if created_key or path.exists():
            path.unlink(missing_ok=True)
        if not directory_existed:
            try:
                path.parent.rmdir()
            except OSError:
                pass
        raise


def _load_or_create_integrity_key(base: Path) -> tuple[str, bytes]:
    with _hold_secret_tree():
        return _load_or_create_integrity_key_held(base)


def _seal(payload: object, key: bytes, domain: bytes) -> str:
    return hmac.new(key, domain + _canonical_json_bytes(payload), hashlib.sha256).hexdigest()


def _verify_seal(payload: object, signature: str, key: bytes, domain: bytes, label: str) -> None:
    if not isinstance(signature, str) or not hmac.compare_digest(_seal(payload, key, domain), signature.lower()):
        raise NativeP2TestInstallError(f"{label}完整性封印不匹配")


def _decode_script(data: bytes, label: str) -> str:
    try:
        return data.decode("gb18030")
    except UnicodeDecodeError as exc:
        raise NativeP2TestInstallError(f"{label}不是GB18030/GBK可读文本") from exc


def _encode_script(text: str) -> bytes:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    return (normalized.rstrip("\n") + "\n").replace("\n", "\r\n").encode("gb18030")


def _is_reparse_point(path: Path) -> bool:
    if path.is_symlink() or getattr(os.path, "isjunction", lambda _: False)(path):
        return True
    try:
        attributes = path.lstat().st_file_attributes
    except (AttributeError, FileNotFoundError, OSError):
        return False
    return bool(attributes & getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400))


class _FileAttributeTagInfo(Structure):
    _fields_ = (("FileAttributes", DWORD), ("ReparseTag", DWORD))


class _FileDispositionInfo(Structure):
    _fields_ = (("DeleteFile", BOOL),)


@dataclass
class _HeldFileLeaf:
    path: Path
    handle: HANDLE | None
    payload: bytes | None
    created: bool = False
    delete_pending: bool = False


def _open_validated_directory_handle(path: Path, label: str) -> HANDLE:
    """Open a directory itself, reject reparse metadata, and deny delete sharing."""
    candidate = Path(path).absolute()
    create_file = windll.kernel32.CreateFileW
    create_file.argtypes = (LPCWSTR, DWORD, DWORD, c_void_p, DWORD, DWORD, HANDLE)
    create_file.restype = HANDLE
    handle = create_file(
        str(candidate),
        GENERIC_READ,
        FILE_SHARE_READ | FILE_SHARE_WRITE,
        None,
        OPEN_EXISTING,
        FILE_FLAG_BACKUP_SEMANTICS | FILE_FLAG_OPEN_REPARSE_POINT,
        None,
    )
    if handle == INVALID_HANDLE_VALUE:
        raise NativeP2TestInstallError(f"无法安全打开{label}目录句柄：{candidate}；{WinError()}")
    info = _FileAttributeTagInfo()
    get_info = windll.kernel32.GetFileInformationByHandleEx
    get_info.argtypes = (HANDLE, DWORD, c_void_p, DWORD)
    get_info.restype = BOOL
    try:
        if not get_info(handle, FILE_ATTRIBUTE_TAG_INFO_CLASS, byref(info), sizeof(info)):
            raise NativeP2TestInstallError(f"无法读取{label}目录句柄属性：{candidate}；{WinError()}")
        if not info.FileAttributes & FILE_ATTRIBUTE_DIRECTORY:
            raise NativeP2TestInstallError(f"{label}路径不是目录：{candidate}")
        if info.FileAttributes & FILE_ATTRIBUTE_REPARSE_POINT:
            raise NativeP2TestInstallError(f"{label}目录不能是符号链接、联接或重解析点：{candidate}")
        return handle
    except BaseException:
        windll.kernel32.CloseHandle(handle)
        raise


def _open_validated_file_handle(path: Path, label: str, *, writable: bool, create_new: bool = False) -> HANDLE:
    """Open the leaf itself without following reparse metadata and deny WRITE/DELETE sharing."""
    candidate = Path(path).absolute()
    create_file = windll.kernel32.CreateFileW
    create_file.argtypes = (LPCWSTR, DWORD, DWORD, c_void_p, DWORD, DWORD, HANDLE)
    create_file.restype = HANDLE
    desired_access = GENERIC_READ | (GENERIC_WRITE | DELETE_ACCESS if writable else 0)
    handle = create_file(
        str(candidate),
        desired_access,
        FILE_SHARE_READ,
        None,
        CREATE_NEW if create_new else OPEN_EXISTING,
        FILE_FLAG_OPEN_REPARSE_POINT,
        None,
    )
    if handle == INVALID_HANDLE_VALUE:
        raise NativeP2TestInstallError(f"无法安全打开{label}叶子句柄：{candidate}；{WinError()}")
    info = _FileAttributeTagInfo()
    get_info = windll.kernel32.GetFileInformationByHandleEx
    get_info.argtypes = (HANDLE, DWORD, c_void_p, DWORD)
    get_info.restype = BOOL
    try:
        if not get_info(handle, FILE_ATTRIBUTE_TAG_INFO_CLASS, byref(info), sizeof(info)):
            raise NativeP2TestInstallError(f"无法读取{label}叶子句柄属性：{candidate}；{WinError()}")
        if info.FileAttributes & FILE_ATTRIBUTE_DIRECTORY:
            raise NativeP2TestInstallError(f"{label}叶子不能是目录：{candidate}")
        if info.FileAttributes & FILE_ATTRIBUTE_REPARSE_POINT:
            raise NativeP2TestInstallError(f"{label}叶子不能是符号链接、联接或重解析点：{candidate}")
        return handle
    except BaseException:
        windll.kernel32.CloseHandle(handle)
        raise


def _seek_file_handle(handle: HANDLE, offset: int = 0) -> None:
    new_position = c_longlong()
    set_pointer = windll.kernel32.SetFilePointerEx
    set_pointer.argtypes = (HANDLE, c_longlong, POINTER(c_longlong), DWORD)
    set_pointer.restype = BOOL
    if not set_pointer(handle, c_longlong(offset), byref(new_position), FILE_BEGIN):
        raise NativeP2TestInstallError(f"无法移动P2叶子句柄位置：{WinError()}")


def _read_file_handle(handle: HANDLE) -> bytes:
    size = c_longlong()
    get_size = windll.kernel32.GetFileSizeEx
    get_size.argtypes = (HANDLE, POINTER(c_longlong))
    get_size.restype = BOOL
    if not get_size(handle, byref(size)) or size.value < 0:
        raise NativeP2TestInstallError(f"无法读取P2叶子句柄长度：{WinError()}")
    _seek_file_handle(handle)
    remaining = size.value
    chunks: list[bytes] = []
    read_file = windll.kernel32.ReadFile
    read_file.argtypes = (HANDLE, c_void_p, DWORD, POINTER(DWORD), c_void_p)
    read_file.restype = BOOL
    while remaining:
        amount = min(remaining, 1024 * 1024)
        buffer = create_string_buffer(amount)
        read = DWORD()
        if not read_file(handle, buffer, amount, byref(read), None):
            raise NativeP2TestInstallError(f"无法从P2叶子句柄读取内容：{WinError()}")
        if read.value == 0:
            raise NativeP2TestInstallError("P2叶子句柄在固定长度内提前读到EOF")
        chunks.append(buffer.raw[:read.value])
        remaining -= read.value
    return b"".join(chunks)


def _write_file_handle_exact(leaf: _HeldFileLeaf, payload: bytes, fault_name: str) -> None:
    if leaf.handle is None:
        raise NativeP2TestInstallError(f"P2目标叶子没有排他写句柄：{leaf.path}")
    _seek_file_handle(leaf.handle)
    set_end = windll.kernel32.SetEndOfFile
    set_end.argtypes = (HANDLE,)
    set_end.restype = BOOL
    if not set_end(leaf.handle):
        raise NativeP2TestInstallError(f"无法截断P2目标叶子：{leaf.path}；{WinError()}")
    leaf.payload = b""
    _transaction_fault(fault_name)
    write_file = windll.kernel32.WriteFile
    write_file.argtypes = (HANDLE, c_void_p, DWORD, POINTER(DWORD), c_void_p)
    write_file.restype = BOOL
    offset = 0
    while offset < len(payload):
        chunk = payload[offset:offset + 1024 * 1024]
        buffer = create_string_buffer(chunk)
        written = DWORD()
        if not write_file(leaf.handle, buffer, len(chunk), byref(written), None):
            raise NativeP2TestInstallError(f"无法写入P2目标叶子：{leaf.path}；{WinError()}")
        if written.value == 0:
            raise NativeP2TestInstallError(f"P2目标叶子写入没有取得进展：{leaf.path}")
        offset += written.value
    flush = windll.kernel32.FlushFileBuffers
    flush.argtypes = (HANDLE,)
    flush.restype = BOOL
    if not flush(leaf.handle):
        raise NativeP2TestInstallError(f"无法持久化P2目标叶子：{leaf.path}；{WinError()}")
    observed = _read_file_handle(leaf.handle)
    if observed != payload:
        raise NativeP2TestInstallError(f"P2目标叶子同句柄写后读回不匹配：{leaf.path}")
    leaf.payload = observed


def _delete_file_handle_exact(leaf: _HeldFileLeaf, fault_name: str) -> None:
    if leaf.handle is None:
        raise NativeP2TestInstallError(f"P2删除目标没有排他叶子句柄：{leaf.path}")
    _transaction_fault(fault_name)
    disposition = _FileDispositionInfo(True)
    set_info = windll.kernel32.SetFileInformationByHandle
    set_info.argtypes = (HANDLE, DWORD, c_void_p, DWORD)
    set_info.restype = BOOL
    if not set_info(leaf.handle, FILE_DISPOSITION_INFO_CLASS, byref(disposition), sizeof(disposition)):
        raise NativeP2TestInstallError(f"无法用同一排他句柄删除P2目标叶子：{leaf.path}；{WinError()}")
    leaf.delete_pending = True
    leaf.payload = None


def _close_file_leaves(leaves: list[_HeldFileLeaf]) -> None:
    for leaf in reversed(leaves):
        if leaf.handle is not None:
            windll.kernel32.CloseHandle(leaf.handle)
            leaf.handle = None


@contextmanager
def _hold_target_leaves(server: Path) -> list[_HeldFileLeaf]:
    leaves: list[_HeldFileLeaf] = []
    try:
        for relative in TARGET_RELATIVES:
            path = Path(server) / Path(relative)
            if os.path.lexists(path):
                handle = _open_validated_file_handle(path, "业务目标", writable=True)
                leaves.append(_HeldFileLeaf(path, handle, _read_file_handle(handle)))
            else:
                leaves.append(_HeldFileLeaf(path, None, None))
        yield leaves
    finally:
        _close_file_leaves(leaves)


def _create_missing_target_handle(leaf: _HeldFileLeaf, index: int) -> None:
    if leaf.handle is not None:
        return
    _transaction_fault(f"before_create_new_target:{index}")
    handle = _open_validated_file_handle(leaf.path, "新业务目标", writable=True, create_new=True)
    leaf.handle = handle
    leaf.payload = _read_file_handle(handle)
    leaf.created = True


@contextmanager
def _hold_directory_tree(root: Path, relative_directories: tuple[Path, ...], *, create: bool) -> Path:
    """Hold raw root and every selected descendant without FILE_SHARE_DELETE."""
    raw_root = Path(root).absolute()
    handles: list[HANDLE] = []
    opened: set[str] = set()

    def hold(path: Path, label: str) -> None:
        key = os.path.normcase(str(Path(path).absolute()))
        if key in opened:
            return
        handle = _open_validated_directory_handle(path, label)
        handles.append(handle)
        opened.add(key)

    try:
        hold(raw_root, "原始入口")
        for relative in sorted(relative_directories, key=lambda item: (len(item.parts), item.as_posix())):
            if relative.is_absolute() or any(part in {"", ".", ".."} for part in relative.parts):
                raise NativeP2TestInstallError(f"受保护目录链含非法相对路径：{relative}")
            current = raw_root
            for part in relative.parts:
                current = current / part
                if create:
                    try:
                        current.mkdir()
                    except FileExistsError:
                        pass
                if not current.is_dir():
                    raise NativeP2TestInstallError(f"受保护目录不存在或不是目录：{current}")
                hold(current, "事务关键")
        yield raw_root.resolve()
    finally:
        for handle in reversed(handles):
            windll.kernel32.CloseHandle(handle)


@contextmanager
def _hold_secret_tree() -> Path:
    secret = _secret_directory().absolute()
    secret_existed = secret.exists()
    secret.parent.mkdir(parents=True, exist_ok=True)
    _assert_secret_path_no_reparse(secret.parent)
    with _hold_directory_tree(secret.parent, (Path(secret.name),), create=True):
        _assert_secret_path_no_reparse(secret)
        if secret_existed:
            _verify_secret_acl(secret)
        else:
            _restrict_secret_acl(secret)
        yield secret


def _assert_no_reparse_path(path: Path, root: Path) -> None:
    root = root.resolve()
    candidate = path.absolute()
    if not candidate.resolve(strict=False).is_relative_to(root):
        raise NativeP2TestInstallError(f"目标路径越出服务端目录：{path}")
    current = root
    if _is_reparse_point(current):
        raise NativeP2TestInstallError(f"服务端目录不能是重解析点：{current}")
    for part in candidate.relative_to(root).parts:
        current = current / part
        if current.exists() and _is_reparse_point(current):
            raise NativeP2TestInstallError(f"目标路径包含符号链接、联接或重解析点：{current}")


def _planned_file(server_root: Path, relative_path: str, after: bytes) -> NativeP2PlannedFile:
    path = server_root / Path(relative_path)
    _assert_no_reparse_path(path, server_root)
    before = _read(path)
    return NativeP2PlannedFile(
        relative_path=relative_path,
        path=path,
        existed_before=before is not None,
        before_sha256=_sha256_bytes(before) if before is not None else None,
        after_sha256=_sha256_bytes(after),
        after=after,
    )


def _current_hash(path: Path) -> str | None:
    data = _read(path)
    return _sha256_bytes(data) if data is not None else None


def _engine_identity(server: Path) -> tuple[str, str, list[str]]:
    executable = server / "Mir200" / "M2Server.exe"
    if not executable.is_file():
        return "missing", "", [f"目标端缺少M2Server.exe，无法确认当前引擎身份：{executable}"]
    digest = _sha256_bytes(executable.read_bytes())
    blockers: list[str] = []
    if digest.lower() != EXPECTED_M2SERVER_SHA256:
        blockers.append("目标M2Server.exe不是当前冻结版本，P2官网工作假设禁止跨版本使用")
    return EXPECTED_M2SERVER_VERSION, digest, blockers


def _workbook_evidence_blockers(path: Path, candidate_id: str) -> list[str]:
    from openpyxl import load_workbook

    blockers: list[str] = []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        properties = workbook["自定义属性"]
        header = {str(cell.value).strip(): index for index, cell in enumerate(next(properties.iter_rows()))}
        claims: dict[str, str] = {}
        for row in properties.iter_rows(min_row=2, values_only=True):
            if str(row[header["候选ID"]] or "").strip() == candidate_id and str(row[header["启用"]] or "").strip().lower() not in {"false", "0", "否", "no"}:
                claims[str(row[header["中文属性"]] or "").strip()] = str(row[header["备注"]] or "").strip()
        required = {"防御": "已实测确认", "生命值": "网上资料待实测", "魔法值": "网上资料待实测"}
        if set(claims) != set(required) or "已实测确认" not in claims.get("防御", "") or any(
            "网上资料待实测" not in claims.get(name, "") or "已实测确认" in claims.get(name, "")
            for name in ("生命值", "魔法值")
        ):
            blockers.append("P2候选表逐属性证据状态必须精确为：防御已实测确认；生命值、魔法值网上资料待实测")

        help_sheet = workbook["属性说明"]
        help_header = {str(cell.value).strip(): index for index, cell in enumerate(next(help_sheet.iter_rows()))}
        help_claims: dict[str, str] = {}
        for row in help_sheet.iter_rows(min_row=2, values_only=True):
            name = str(row[help_header["中文属性"]] or "").strip()
            if name in required:
                help_claims[name] = str(row[help_header["实效状态"]] or "").strip()
        for name in ("生命值", "魔法值"):
            claim = help_claims.get(name, "")
            if "已实测" in claim or "待实测" not in claim:
                blockers.append(f"属性说明不得把{name}升级为已实测，必须明确待实测")
    except (KeyError, StopIteration, IndexError) as exc:
        blockers.append(f"P2候选表缺少证据状态所需工作表或列：{exc}")
    finally:
        workbook.close()
    return blockers


def _evidence_blockers(evidence_items: tuple[NativeP2Evidence, ...]) -> list[str]:
    blockers: list[str] = []
    if any(not item.source_locator or not item.snapshot_path or not item.captured_at or not item.source_sha256 for item in evidence_items):
        blockers.append("P2证据来源、日期或快照哈希不完整")
    if any(item.source_type == "official_web" and item.evidence_status != "网上资料待实测" for item in evidence_items):
        blockers.append("官网资料不得自动升级为已实测确认")
    for evidence in evidence_items:
        snapshot = Path(evidence.snapshot_path)
        if not snapshot.is_file():
            blockers.append(f"P2证据快照不存在：{snapshot}")
        elif _sha256_bytes(snapshot.read_bytes()).upper() != evidence.source_sha256.upper():
            blockers.append(f"P2证据快照哈希漂移：{snapshot}")
    return blockers


def _candidate_spec(candidate: NativeCandidate) -> dict[str, object]:
    compiled = compile_native_candidate(candidate)
    return {
        "candidate_id": candidate.candidate_id,
        "target_item_name": candidate.target_item_name,
        "equip_slot": candidate.equip_slot,
        "mingge_name": candidate.mingge_name,
        "default_text_color": candidate.default_text_color,
        "properties": [
            {
                "order": item.order, "display_name": item.display_name, "value": item.value,
                "value2": item.value2, "value3": item.value3, "property_row": item.property_row,
                "position": item.position, "binding": item.binding, "color": item.color,
                "type3": item.type3, "type4": item.type4,
                "value_command": item.value_command, "dependency": item.dependency,
            }
            for item in candidate.properties
        ],
        "segments": [
            {
                "order": item.order, "role": item.role, "text": item.text,
                "color": item.color, "source_property": item.source_property,
            }
            for item in candidate.segments
        ],
        "compiled_script_sha256": _sha256_bytes(compiled.payload),
    }


def _candidate_spec_hash(candidate: NativeCandidate) -> str:
    payload = json.dumps(_candidate_spec(candidate), ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return _sha256_bytes(payload)


def _candidate_blockers(candidate: NativeCandidate) -> list[str]:
    blockers: list[str] = []
    if candidate.candidate_id != EXPECTED_CANDIDATE_ID:
        blockers.append(f"P2只允许固定候选ID：{EXPECTED_CANDIDATE_ID}")
    if candidate.target_item_name != EXPECTED_TARGET_ITEM or candidate.equip_slot != 17:
        blockers.append("P2只允许可废弃测试装备鞭尸灵玉及灵玉位置17")
    if candidate.mingge_name != EXPECTED_MINGGE_NAME or candidate.default_text_color != 255:
        blockers.append("P2命格名称或默认文字颜色不符合固定黄金测试夹具")
    expected = (
        (1, "防御", 5, 17, 17, 1, 250, 0, 9, "value_ex"),
        (2, "生命值", 100, 18, 18, 6, 250, 0, 9, "value_ex"),
        (3, "魔法值", 100, 19, 19, 7, 250, 0, 9, "value_ex"),
    )
    actual = tuple(
        (item.order, item.display_name, item.value, item.property_row, item.position,
         item.binding, item.color, item.type3, item.type4, item.value_command)
        for item in candidate.properties
    )
    if actual != expected or any(item.value2 != 0 or item.value3 != 0 or item.dependency is not None for item in candidate.properties):
        blockers.append("P2候选必须精确三属性：防御+5、生命值+100、魔法值+100；禁止脚本依赖属性")
    expected_segments = (
        (1, "命格名称", "", "", 249),
        (2, "属性名称", "防御", "", 250), (3, "正负号", "防御", "", 69), (4, "属性值", "防御", "", 69),
        (5, "固定文字", "", "·", 255),
        (6, "属性名称", "生命值", "", 251), (7, "正负号", "生命值", "", 69), (8, "属性值", "生命值", "", 69),
        (9, "固定文字", "", "·", 255),
        (10, "属性名称", "魔法值", "", 250), (11, "正负号", "魔法值", "", 69), (12, "属性值", "魔法值", "", 69),
    )
    actual_segments = tuple((item.order, item.role, item.source_property, item.text, item.color) for item in candidate.segments)
    if actual_segments != expected_segments:
        blockers.append("P2显示片段、来源或颜色不符合固定测试夹具")
    compiled = compile_native_candidate(candidate)
    compiled_hash = _sha256_bytes(compiled.payload)
    spec_hash = _candidate_spec_hash(candidate)
    if compiled_hash != EXPECTED_COMPILED_SCRIPT_SHA256 or spec_hash != EXPECTED_CANDIDATE_SPEC_SHA256:
        blockers.append("P2候选编译结果或规范哈希不等于冻结黄金夹具")
    forbidden = ("GAMEGOLD", "LINKGIVEITEM", "SQL", "@RECALCABILITYS", "RANDOM", "DELAY")
    upper = compiled.script.upper()
    if any(token in upper for token in forbidden):
        blockers.append("P2编译脚本包含数据库、扣费、随机、延迟或重算禁用命令")
    lines = compiled.script.splitlines()
    if lines.count("LockUpdateItem 17") != 1 or lines.count("UpdateItem 17") != 1:
        blockers.append("P2编译脚本必须恰好包含一对LockUpdateItem/UpdateItem")
    else:
        start = lines.index("LockUpdateItem 17")
        end = lines.index("UpdateItem 17")
        allowed_prefixes = (
            "SetCustomItemAbil 17 ", "SetCustomItemValueEx 17 ",
            "SetCustomItemText 17 ", "SetCustomItemTextColor 17 ",
        )
        if start >= end or any(not line.startswith(allowed_prefixes) for line in lines[start + 1:end]):
            blockers.append("P2的LockUpdateItem至UpdateItem之间必须是连续固定写段，禁止条件、跳转、调用或延迟")
    blockers.extend(_evidence_blockers(P2_EVIDENCE))
    return blockers


def _read_lines(prefix: str) -> list[str]:
    lines: list[str] = []
    for row in (17, 18, 19):
        stem = f"N$XY_MG_P2_{prefix}{row}"
        lines.extend((
            f"GetCustomItemAbil 17 {row} 0 {stem}C",
            f"GetCustomItemAbil 17 {row} 1 {stem}B",
            f"GetCustomItemAbil 17 {row} 2 {stem}P",
            f"GetCustomItemAbil 17 {row} 3 {stem}M",
            f"GetCustomItemValueEx 17 {row} {stem}T {stem}V1 {stem}V2 {stem}V3",
        ))
    return lines


def _row_conditions(prefix: str, row: int, values: tuple[int, ...]) -> list[str]:
    suffixes = ("C", "B", "P", "M", "T", "V1", "V2", "V3")
    return [f"EQUAL N$XY_MG_P2_{prefix}{row}{suffix} {value}" for suffix, value in zip(suffixes, values)]


BASELINE = {
    17: (250, 1, 17, 0, 0, 5, 0, 0),
    18: (0, 0, 0, 0, 0, 0, 0, 0),
    19: (0, 0, 0, 0, 0, 0, 0, 0),
}
EXPECTED_AFTER = {
    17: (250, 1, 17, 0, 0, 5, 0, 0),
    18: (250, 6, 18, 0, 0, 100, 0, 0),
    19: (250, 7, 19, 0, 0, 100, 0, 0),
}


def _all_conditions(prefix: str, expected: dict[int, tuple[int, ...]]) -> list[str]:
    result: list[str] = []
    for row in (17, 18, 19):
        result.extend(_row_conditions(prefix, row, expected[row]))
    return result


def _render_test_script(candidate: NativeCandidate) -> bytes:
    command_lines = compile_native_candidate(candidate).script.rstrip("\r\n").split("\r\n")
    body = [
        "[@XY_MG_NATIVE_P2_MAIN]", "{", "#IF", "ISADMIN", "CHECKUSEITEM 17", "#ACT",
        "GOTO @XY_MG_NATIVE_P2_CHECK_NAME", "BREAK", "#ELSEACT",
        "MESSAGEBOX 仅允许管理员在灵玉位置执行P2隔离验收，未写入。", "BREAK", "",
        "[@XY_MG_NATIVE_P2_CHECK_NAME]", "#IF", f"EQUAL <$JADE> {candidate.target_item_name}", "#ACT",
        "GOTO @XY_MG_NATIVE_P2_READ_A", "BREAK", "#ELSEACT",
        f"MESSAGEBOX 当前灵玉不是“{candidate.target_item_name}”，已阻止。", "BREAK", "",
        "[@XY_MG_NATIVE_P2_READ_A]", "#IF", "#ACT", "GetCustomItemText 17 <$STR(S1)>",
        "GetCustomItemTextColor 17 <$STR(N1)>", *_read_lines("A"),
        "GOTO @XY_MG_NATIVE_P2_CHECK_ALREADY", "BREAK", "",
        "[@XY_MG_NATIVE_P2_CHECK_ALREADY]", "#IF", *_all_conditions("A", EXPECTED_AFTER), "#ACT",
        "GOTO @XY_MG_NATIVE_P2_ALREADY", "BREAK", "#ELSEACT", "GOTO @XY_MG_NATIVE_P2_CHECK_BASELINE", "BREAK", "",
        "[@XY_MG_NATIVE_P2_CHECK_BASELINE]", "#IF", *_all_conditions("A", BASELINE), "#ACT",
        "GOTO @XY_MG_NATIVE_P2_READY", "BREAK", "#ELSEACT", "GOTO @XY_MG_NATIVE_P2_BLOCKED", "BREAK", "",
        "[@XY_MG_NATIVE_P2_READY]", "#SAY",
        f"平台候选：{candidate.candidate_id}\\",
        "防御行17已是P1黄金值；属性行18、19必须为空且已读回为全0。\\",
        "本次会把同一件可废弃测试灵玉更新为防御+5、生命值+100、魔法值+100，并整体替换命格显示文字。\\",
        "生命值和魔法值依据官网规则进入测试，当前仍是网上资料待实测；不扣费、不抽取，安装器不直写数据库。\\",
        "<写入P2三属性候选/@XY_MG_NATIVE_P2_READ_B>  <退出/@EXIT>", "",
        "[@XY_MG_NATIVE_P2_ALREADY]", "#SAY",
        "当前三行已精确等于P2候选，未重复写入。\\",
        "请核对人物防御、最大生命、最大魔法和悬浮多色文字。\\", "<退出/@EXIT>", "",
        "[@XY_MG_NATIVE_P2_BLOCKED]", "#SAY",
        "已阻止P2写入：要求行17为250/1/17/0/0/5/0/0，属性行18、19必须为空。\\",
        "当前物品不是P1黄金基线，也不是已完成的P2成品；请保留截图并退出。\\", "<退出/@EXIT>", "",
        "[@XY_MG_NATIVE_P2_READ_B]", "#IF", "ISADMIN", "CHECKUSEITEM 17",
        f"EQUAL <$JADE> {candidate.target_item_name}", "#ACT", *_read_lines("B"),
        "GOTO @XY_MG_NATIVE_P2_CHECK_B", "BREAK", "#ELSEACT",
        "MESSAGEBOX 点击后管理员、位置或装备名称已变化，未写入。", "BREAK", "",
        "[@XY_MG_NATIVE_P2_CHECK_B]", "#IF", *_all_conditions("B", BASELINE), "#ACT",
        "GOTO @XY_MG_NATIVE_P2_READ_C", "BREAK", "#ELSEACT",
        "MESSAGEBOX 点击后属性行17至19已变化，未写入；请重新输入@命格平台三属性验收。", "BREAK", "",
        "[@XY_MG_NATIVE_P2_READ_C]", "#IF", "ISADMIN", "CHECKUSEITEM 17",
        f"EQUAL <$JADE> {candidate.target_item_name}", "#ACT", *_read_lines("C"),
        "GOTO @XY_MG_NATIVE_P2_CHECK_C", "BREAK", "#ELSEACT",
        "MESSAGEBOX 最终门禁的管理员、位置或装备名称已变化，未写入。", "BREAK", "",
        "[@XY_MG_NATIVE_P2_CHECK_C]", "#IF", *_all_conditions("C", BASELINE), "#ACT",
        "GOTO @XY_MG_NATIVE_P2_APPLY", "BREAK", "#ELSEACT",
        "MESSAGEBOX 最终实际读回与P1基线不符，未写入；请重新输入@命格平台三属性验收。", "BREAK", "",
        "[@XY_MG_NATIVE_P2_APPLY]", "#IF", "#ACT", *command_lines,
        "GOTO @XY_MG_NATIVE_P2_READ_AFTER", "BREAK", "",
        "[@XY_MG_NATIVE_P2_READ_AFTER]", "#IF", "#ACT", "GetCustomItemText 17 <$STR(S1)>",
        "GetCustomItemTextColor 17 <$STR(N1)>", *_read_lines("D"),
        "GOTO @XY_MG_NATIVE_P2_CHECK_AFTER", "BREAK", "",
        "[@XY_MG_NATIVE_P2_CHECK_AFTER]", "#IF", *_all_conditions("D", EXPECTED_AFTER), "#ACT",
        "GOTO @XY_MG_NATIVE_P2_SUCCESS", "BREAK", "#ELSEACT", "GOTO @XY_MG_NATIVE_P2_AFTER_MISMATCH", "BREAK", "",
        "[@XY_MG_NATIVE_P2_SUCCESS]", "#SAY",
        "P2三属性写后读回通过：\\",
        "行17=250/1/17/0/0/5/0/0\\行18=250/6/18/0/0/100/0/0\\行19=250/7/19/0/0/100/0/0\\",
        "请截图本面板；再核对人物防御、最大生命、最大魔法与悬浮框多色文字，最后小退重登复核。\\",
        "<退出/@EXIT>", "",
        "[@XY_MG_NATIVE_P2_AFTER_MISMATCH]", "#SAY",
        "P2写后读回异常，请立即停止使用并隔离这件测试物品。文件回滚不能恢复已经修改的物品实例。\\",
        "请截图本面板后退出。\\", "<退出/@EXIT>", "}",
    ]
    rendered = _encode_script("\n".join(body))
    decoded = rendered.decode("gb18030").upper()
    if "@RECALCABILITYS" in decoded:
        raise NativeP2TestInstallError("P2验收脚本禁止进入@RecalcAbilitys")
    return rendered


def _render_usercmd(current: bytes) -> tuple[bytes, list[str]]:
    text = _decode_script(current, "UserCmd.txt")
    blockers: list[str] = []
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    exact = f"{COMMAND_NAME}\t{COMMAND_NUMBER}"
    parsed = [tuple(line.split("\t")) for line in lines if len(line.split("\t")) == 2]
    exact_count = sum(1 for name, number in parsed if name == COMMAND_NAME and number == str(COMMAND_NUMBER))
    name_count = sum(1 for name, _ in parsed if name == COMMAND_NAME)
    number_count = sum(1 for _, number in parsed if number == str(COMMAND_NUMBER))
    if exact_count > 1 or name_count > 1 or number_count > 1:
        blockers.append(f"UserCmd的{COMMAND_NAME}/编号{COMMAND_NUMBER}映射必须全文件唯一")
    for line in lines:
        parts = line.split("\t")
        if len(parts) != 2:
            continue
        if parts[0] == COMMAND_NAME and parts[1] != str(COMMAND_NUMBER):
            blockers.append(f"UserCmd命令名{COMMAND_NAME}已占用其他编号")
        if parts[1] == str(COMMAND_NUMBER) and parts[0] != COMMAND_NAME:
            blockers.append(f"UserCmd编号{COMMAND_NUMBER}已被{parts[0]}占用")
    if exact_count == 1:
        return current, blockers
    if blockers:
        return current, blockers
    suffix = b"" if current.endswith((b"\r\n", b"\n", b"\r")) or not current else b"\r\n"
    return current + suffix + exact.encode("gb18030") + b"\r\n", blockers


def _qfunction_block() -> str:
    return "\n".join((
        MANAGED_BEGIN, f"[@UserCmd{COMMAND_NUMBER}]", "#IF", "ISADMIN", "#ACT",
        "#CALL [\\玄渊验收\\命格平台三属性验收.txt] @XY_MG_NATIVE_P2_MAIN", "BREAK", "#ELSEACT",
        "SENDMSG 6 [命格平台三属性验收] 仅管理员测试角色可使用。", "BREAK", MANAGED_END,
    ))


def _render_qfunction(current: bytes) -> tuple[bytes, list[str]]:
    text = _decode_script(current, "QFunction-0.txt")
    blockers: list[str] = []
    block = _qfunction_block()
    if text.count(MANAGED_BEGIN) or text.count(MANAGED_END):
        if text.count(MANAGED_BEGIN) != 1 or text.count(MANAGED_END) != 1:
            blockers.append("QFunction的P2受管标记不唯一")
        if block.replace("\n", "\r\n") not in text and block not in text:
            blockers.append("QFunction已存在不匹配的P2受管块")
        if text.count(f"[@UserCmd{COMMAND_NUMBER}]") != 1:
            blockers.append(f"QFunction的@UserCmd{COMMAND_NUMBER}必须在全文件恰好出现一次")
        return current, blockers
    if text.count(f"[@UserCmd{COMMAND_NUMBER}]"):
        blockers.append(f"QFunction的@UserCmd{COMMAND_NUMBER}已被其他逻辑占用")
        return current, blockers
    suffix = "" if text.endswith(("\r\n", "\n", "\r")) or not text else "\r\n"
    return (text + suffix + "\r\n" + block.replace("\n", "\r\n") + "\r\n").encode("gb18030"), blockers


def plan_native_p2_test_install(workbook_path: Path, candidate_id: str, server_root: Path) -> NativeP2InstallPlan:
    """只读预检固定P2候选，生成精确三文件计划。"""
    book = read_native_candidate_workbook(Path(workbook_path))
    matches = [item for item in book.candidates if item.candidate_id == candidate_id]
    if len(matches) != 1:
        raise NativeP2TestInstallError(f"启用候选必须唯一存在：{candidate_id}")
    candidate = matches[0]
    server_entry = Path(server_root).absolute()
    server = server_entry.resolve()
    blockers = _candidate_blockers(candidate)
    blockers.extend(_workbook_evidence_blockers(book.source, candidate_id))
    if _is_reparse_point(server_entry):
        blockers.append(f"服务端入口不能是符号链接、联接或重解析点：{server_entry}")
    engine_version, engine_sha256, engine_blockers = _engine_identity(server)
    blockers.extend(engine_blockers)
    usercmd_path = server / Path(USERCMD_RELATIVE)
    qfunction_path = server / Path(QFUNCTION_RELATIVE)
    try:
        _assert_no_reparse_path(usercmd_path, server)
        _assert_no_reparse_path(qfunction_path, server)
    except NativeP2TestInstallError as exc:
        blockers.append(str(exc))
    if not usercmd_path.is_file():
        blockers.append(f"UserCmd不存在：{usercmd_path}")
        usercmd_after = b""
    else:
        usercmd_after, found = _render_usercmd(usercmd_path.read_bytes())
        blockers.extend(found)
    if not qfunction_path.is_file():
        blockers.append(f"QFunction不存在：{qfunction_path}")
        qfunction_after = b""
    else:
        qfunction_after, found = _render_qfunction(qfunction_path.read_bytes())
        blockers.extend(found)
    test_script = _render_test_script(candidate)
    test_path = server / Path(TEST_SCRIPT_RELATIVE)
    existing_test = _read(test_path)
    if existing_test is not None and existing_test != test_script:
        blockers.append("P2隔离测试脚本已存在且内容不匹配，禁止覆盖")
    files = (
        _planned_file(server, USERCMD_RELATIVE, usercmd_after),
        _planned_file(server, QFUNCTION_RELATIVE, qfunction_after),
        _planned_file(server, TEST_SCRIPT_RELATIVE, test_script),
    )
    if all(item.before_sha256 == item.after_sha256 for item in files):
        blockers.append("P2三个固定目标已经全部是安装后内容，禁止建立无变化的新事务")
    if tuple(item.relative_path for item in files) != TARGET_RELATIVES:
        blockers.append("P2文件目标集合不符合固定三文件合同")
    workbook_hash = _sha256_bytes(book.source.read_bytes())
    spec_hash = _candidate_spec_hash(candidate)
    compiled_hash = _sha256_bytes(compile_native_candidate(candidate).payload)
    evidence_digest = _sha256_bytes(_canonical_json_bytes([asdict(item) for item in P2_EVIDENCE]))
    identity = "|".join((candidate_id, spec_hash, workbook_hash, engine_sha256, evidence_digest,
                         *(f"{i.relative_path}:{i.before_sha256}:{i.after_sha256}" for i in files)))
    return NativeP2InstallPlan(
        plan_id=_sha256_bytes(identity.encode("utf-8"))[:24], candidate_id=candidate_id,
        candidate_spec_sha256=spec_hash, compiled_script_sha256=compiled_hash,
        workbook=book.source, workbook_sha256=workbook_hash, server_root=server,
        engine_version=engine_version, engine_sha256=engine_sha256, files=files,
        blockers=tuple(dict.fromkeys(blockers)),
        warnings=(
            "P1入口与脚本保持原样；P2只增加独立管理员验收入口",
            "安装器不直写数据库；玩家执行命令后M2会修改当前测试物品的持久实例数据",
            "生命值与魔法值为官网规则工作假设，游戏验收前不得标为已实测",
            "文件回滚不能恢复已经修改的物品实例，失败时应隔离或废弃测试物品",
        ),
        evidence_status=tuple(item.evidence_status for item in P2_EVIDENCE), evidence=P2_EVIDENCE,
    )


def _assert_plan_is_current(
    plan: NativeP2InstallPlan,
    target_leaves: list[_HeldFileLeaf] | None = None,
) -> None:
    if plan.blockers:
        raise NativeP2TestInstallError("安装计划存在阻断：" + "；".join(plan.blockers))
    if plan.candidate_id != EXPECTED_CANDIDATE_ID:
        raise NativeP2TestInstallError("P2计划候选ID不匹配")
    if plan.candidate_spec_sha256 != EXPECTED_CANDIDATE_SPEC_SHA256 or plan.compiled_script_sha256 != EXPECTED_COMPILED_SCRIPT_SHA256:
        raise NativeP2TestInstallError("P2计划不是冻结黄金候选")
    if len(plan.files) != len(TARGET_RELATIVES):
        raise NativeP2TestInstallError("P2计划文件数量不符合固定三文件合同")
    for index, (item, relative) in enumerate(zip(plan.files, TARGET_RELATIVES)):
        expected_path = plan.server_root / Path(relative)
        if item.relative_path != relative:
            raise NativeP2TestInstallError(f"P2计划文件顺序或相对路径不符合固定三文件合同：{index}")
        if item.path.absolute() != expected_path.absolute():
            raise NativeP2TestInstallError(f"P2计划文件绝对目标不符合固定三文件合同：{index}")
        if item.after_sha256 != _sha256_bytes(item.after):
            raise NativeP2TestInstallError(f"P2计划文件安装后内容与哈希不匹配：{index}")
    if all(item.before_sha256 == item.after_sha256 for item in plan.files):
        raise NativeP2TestInstallError("P2三个固定目标已经全部是安装后内容，禁止建立无变化的新事务")
    if _sha256_bytes(plan.workbook.read_bytes()) != plan.workbook_sha256:
        raise NativeP2TestInstallError("38号候选表在预检后发生变化，请重新预检")
    fresh = read_native_candidate_workbook(plan.workbook)
    matches = [item for item in fresh.candidates if item.candidate_id == plan.candidate_id]
    if len(matches) != 1 or _candidate_spec_hash(matches[0]) != plan.candidate_spec_sha256:
        raise NativeP2TestInstallError("P2候选规范在预检后发生变化，请重新预检")
    evidence_errors = _evidence_blockers(plan.evidence)
    evidence_errors.extend(_workbook_evidence_blockers(plan.workbook, plan.candidate_id))
    engine_version, engine_sha256, engine_errors = _engine_identity(plan.server_root)
    if evidence_errors:
        raise NativeP2TestInstallError("P2证据在预检后发生变化：" + "；".join(evidence_errors))
    if engine_errors or engine_version != plan.engine_version or engine_sha256 != plan.engine_sha256:
        raise NativeP2TestInstallError("目标引擎身份在预检后发生变化：" + "；".join(engine_errors))
    for index, item in enumerate(plan.files):
        _assert_no_reparse_path(item.path, plan.server_root)
        if target_leaves is None:
            current_hash = _current_hash(item.path)
        else:
            current = _refresh_target_leaf(target_leaves[index])
            current_hash = _sha256_bytes(current) if current is not None else None
        if current_hash != item.before_sha256:
            raise NativeP2TestInstallError(f"目标文件在预检后发生变化，请重新预检：{item.path}")


def _transaction_base(platform_root: Path) -> Path:
    raw_root = Path(platform_root).absolute()
    if _is_reparse_point(raw_root):
        raise NativeP2TestInstallError("平台原始入口不能是符号链接、联接或重解析点")
    if not raw_root.is_dir():
        raise NativeP2TestInstallError(f"平台根目录不存在或不是目录：{raw_root}")
    root = raw_root.resolve()
    base = root / "backups" / "mingge-native-p2-test"
    if not base.resolve(strict=False).is_relative_to(root):
        raise NativeP2TestInstallError("命格P2事务目录越界")
    if _is_reparse_point(root):
        raise NativeP2TestInstallError("平台根目录不能是重解析点")
    _assert_no_reparse_path(base, root)
    return base


@contextmanager
def _transaction_lock(platform_root: Path, server_root: Path):
    raw_platform = Path(platform_root).absolute()
    raw_server = Path(server_root).absolute()
    if _is_reparse_point(raw_platform):
        raise NativeP2TestInstallError("平台原始入口不能是符号链接、联接或重解析点")
    raw_platform.mkdir(parents=True, exist_ok=True)
    if _is_reparse_point(raw_server):
        raise NativeP2TestInstallError("服务端原始入口不能是符号链接、联接或重解析点")
    protected = (
        Path("backups"),
        Path("backups") / "mingge-native-p2-test",
        Path("backups") / "mingge-native-p2-test" / "active",
    )
    with _hold_directory_tree(raw_platform, protected, create=True):
        base = _transaction_base(raw_platform)
        mutex_name = "Global\\XYDP_SERVER_SCRIPT_" + _sha256_bytes(str(raw_server.resolve()).lower().encode("utf-8"))[:24]
        create_mutex = windll.kernel32.CreateMutexW
        create_mutex.argtypes = (c_void_p, BOOL, LPCWSTR)
        create_mutex.restype = HANDLE
        handle = create_mutex(None, False, mutex_name)
        if not handle:
            raise NativeP2TestInstallError("无法创建服务端脚本事务全局互斥锁")
        wait_for_single = windll.kernel32.WaitForSingleObject
        wait_for_single.argtypes = (HANDLE, DWORD)
        wait_for_single.restype = DWORD
        wait_result = wait_for_single(handle, 0)
        if wait_result not in (WAIT_OBJECT_0, WAIT_ABANDONED):
            windll.kernel32.CloseHandle(handle)
            if wait_result == WAIT_TIMEOUT:
                raise NativeP2TestInstallError("另一个服务端脚本安装或回滚事务正在执行")
            raise NativeP2TestInstallError("无法取得服务端脚本事务全局互斥锁")
        try:
            yield base
        finally:
            windll.kernel32.ReleaseMutex(handle)
            windll.kernel32.CloseHandle(handle)


def _receipt_payload(plan: NativeP2InstallPlan, transaction_id: str, entries: list[dict[str, object]], status: str) -> dict[str, object]:
    return {
        "schema_version": 2, "operation": "mingge-native-p2-test", "status": status,
        "transaction_id": transaction_id, "candidate_id": plan.candidate_id,
        "plan_id": plan.plan_id, "candidate_spec_sha256": plan.candidate_spec_sha256,
        "compiled_script_sha256": plan.compiled_script_sha256,
        "engine_version_baseline": ENGINE_VERSION_BASELINE,
        "server_root": str(plan.server_root), "workbook": str(plan.workbook),
        "workbook_sha256": plan.workbook_sha256, "files": entries,
        "evidence": [asdict(item) for item in plan.evidence],
        "evidence_status": "working_hypothesis_pending_game_validation",
        "game_validation_status": "pending", "equip_slots": [17], "payment": False,
        "writes_server": True, "writes_client": False, "writes_database_directly": False,
        "runtime_item_instance_may_change": True, "runtime_item_instance_restored_by_file_rollback": False,
        "p1_route_untouched": True, "old_mingge_route_untouched": True,
        "updated_at_utc": datetime.now(timezone.utc).isoformat(),
    }


def _manifest_payload(
    plan: NativeP2InstallPlan,
    transaction_id: str,
    platform_root: Path,
    key_id: str,
    generation: str,
    entries: list[dict[str, object]],
) -> dict[str, object]:
    module_path = Path(__file__).resolve()
    cli_path = module_path.with_name("cli.py")
    return {
        "schema_version": 4,
        "operation": "mingge-native-p2-test",
        "transaction_id": transaction_id,
        "generation": generation,
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "key_id": key_id,
        "tool": {
            "module_sha256": _sha256_bytes(module_path.read_bytes()),
            "cli_sha256": _sha256_bytes(cli_path.read_bytes()) if cli_path.is_file() else None,
        },
        "platform_root": str(Path(platform_root).resolve()),
        "server_root": str(plan.server_root),
        "engine": {
            "required_baseline": ENGINE_VERSION_BASELINE,
            "detected_version": plan.engine_version,
            "binary_sha256": plan.engine_sha256,
        },
        "candidate": {
            "candidate_id": plan.candidate_id,
            "candidate_spec_sha256": plan.candidate_spec_sha256,
            "compiled_script_sha256": plan.compiled_script_sha256,
            "workbook_path": str(plan.workbook),
            "workbook_sha256": plan.workbook_sha256,
            "equip_slots": [17],
            "payment": False,
            "writes_database_directly": False,
            "runtime_item_instance_restored_by_file_rollback": False,
        },
        "evidence": [asdict(item) for item in plan.evidence],
        "runtime_contract": {
            "expected_before": {str(key): list(value) for key, value in BASELINE.items()},
            "expected_after": {str(key): list(value) for key, value in EXPECTED_AFTER.items()},
            "item_policy": "disposable_test_instance_only",
        },
        "files": entries,
    }


def _write_manifest(
    transaction: Path,
    manifest: dict[str, object],
    key: bytes,
    *,
    server_root: Path,
    transaction_id: str,
    generation: str,
) -> None:
    manifest_path = transaction / "manifest.json"
    signature_path = transaction / "manifest.hmac"
    _atomic_write(manifest_path, _canonical_json_bytes(manifest))
    _set_coordination_setup_state(
        server_root, transaction_id, generation, key, "MANIFEST_JSON_WRITTEN",
    )
    _transaction_fault("after_manifest_json")
    _atomic_write(signature_path, (_seal(manifest, key, MANIFEST_DOMAIN) + "\n").encode("ascii"))
    _set_coordination_setup_state(
        server_root, transaction_id, generation, key, "MANIFEST_HMAC_WRITTEN",
    )
    _transaction_fault("after_manifest_hmac")
    loaded = json.loads(manifest_path.read_text(encoding="utf-8"))
    _verify_seal(loaded, signature_path.read_text(encoding="ascii").strip(), key, MANIFEST_DOMAIN, "P2 manifest")
    _set_coordination_setup_state(
        server_root, transaction_id, generation, key, "MANIFEST_VERIFIED",
    )


def _coordination_path(server_root: Path) -> Path:
    server_id = _sha256_bytes(str(Path(server_root).resolve()).lower().encode("utf-8"))[:24]
    path = _secret_directory() / f"mingge-p2-coordination-{server_id}.json"
    if path.parent != _secret_directory():
        raise NativeP2TestInstallError("P2受保护协调文件越界")
    return path


def _read_coordination(server_root: Path, key: bytes, *, required: bool) -> dict[str, object] | None:
    with _hold_secret_tree():
        path = _coordination_path(server_root)
        if not path.exists():
            if required:
                raise NativeP2TestInstallError("P2受保护协调状态缺失，拒绝仅信任普通备份树")
            return None
        if _is_reparse_point(path):
            raise NativeP2TestInstallError("P2受保护协调文件不能是重解析点")
        _verify_secret_acl(path.parent)
        _verify_secret_acl(path)
        try:
            envelope = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise NativeP2TestInstallError("P2受保护协调状态不可读或损坏") from exc
        signature = envelope.pop("hmac", None)
        _verify_seal(envelope, signature, key, COORDINATION_DOMAIN, "P2 coordination")
        reservation_files = envelope.get("reservation_files")
        predecessor_owner = envelope.get("predecessor_owner")
        if (
            envelope.get("schema_version") != 3
            or Path(str(envelope.get("server_root", ""))).resolve() != Path(server_root).resolve()
            or envelope.get("target_relatives") != list(TARGET_RELATIVES)
            or envelope.get("owner_state") not in {"none", "pending_active", "active", "pending_closed", "closed"}
            or envelope.get("setup_state") not in {
                "RESERVED", "MANIFEST_JSON_WRITTEN", "MANIFEST_HMAC_WRITTEN",
                "MANIFEST_VERIFIED", "PREPARED", "BUSINESS_WRITES", "ABORTED",
            }
            or not isinstance(envelope.get("business_write_started"), bool)
            or not isinstance(reservation_files, list)
            or len(reservation_files) != len(TARGET_RELATIVES)
            or tuple(
                item.get("relative_path") for item in reservation_files if isinstance(item, dict)
            ) != TARGET_RELATIVES
            or any(
                not isinstance(item, dict)
                or item.get("index") != index
                or item.get("before_sha256") is not None and not isinstance(item.get("before_sha256"), str)
                or not isinstance(item.get("after_sha256"), str)
                for index, item in enumerate(reservation_files)
            )
            or not isinstance(envelope.get("journal_sequence"), int)
            or int(envelope.get("journal_sequence", -1)) < 0
            or not isinstance(envelope.get("journal_head_hmac"), str)
            or predecessor_owner is not None and (
                not isinstance(predecessor_owner, dict)
                or set(predecessor_owner) != {"transaction_id", "generation", "state"}
                or not isinstance(predecessor_owner.get("transaction_id"), str)
                or not predecessor_owner.get("transaction_id")
                or not isinstance(predecessor_owner.get("generation"), str)
                or not predecessor_owner.get("generation")
                or predecessor_owner.get("state") != "closed"
            )
        ):
            raise NativeP2TestInstallError("P2受保护协调状态基础合同不匹配")
        envelope["hmac"] = signature
        return envelope


def _write_coordination(payload: dict[str, object], key: bytes) -> dict[str, object]:
    clean = {name: value for name, value in payload.items() if name != "hmac"}
    clean["updated_at_utc"] = datetime.now(timezone.utc).isoformat()
    clean["hmac"] = _seal(clean, key, COORDINATION_DOMAIN)
    with _hold_secret_tree():
        path = _coordination_path(Path(str(clean["server_root"])))
        _atomic_write(path, _canonical_json_bytes(clean))
        _restrict_secret_acl(path)
        _verify_secret_acl(path.parent)
        _verify_secret_acl(path)
    loaded = _read_coordination(Path(str(clean["server_root"])), key, required=True)
    if not loaded or any(loaded.get(name) != value for name, value in clean.items()):
        raise NativeP2TestInstallError("P2受保护协调状态写后验证失败")
    return loaded


def _begin_coordination(
    server_root: Path,
    transaction_id: str,
    generation: str,
    key: bytes,
    reservation_files: list[dict[str, object]],
    predecessor_owner: dict[str, object] | None,
) -> dict[str, object]:
    return _write_coordination({
        "schema_version": 3,
        "server_root": str(Path(server_root).resolve()),
        "transaction_id": transaction_id,
        "generation": generation,
        "phase": "MANIFEST_ONLY",
        "setup_state": "RESERVED",
        "business_write_started": False,
        "reservation_files": reservation_files,
        "journal_sequence": 0,
        "journal_head_hmac": "",
        "pending_event": None,
        "owner_state": "none",
        "predecessor_owner": predecessor_owner,
        "target_relatives": list(TARGET_RELATIVES),
    }, key)


def _set_coordination_setup_state(
    server_root: Path,
    transaction_id: str,
    generation: str,
    key: bytes,
    setup_state: str,
    *,
    business_write_started: bool | None = None,
) -> dict[str, object]:
    allowed = {
        "RESERVED", "MANIFEST_JSON_WRITTEN", "MANIFEST_HMAC_WRITTEN",
        "MANIFEST_VERIFIED", "PREPARED", "BUSINESS_WRITES", "ABORTED",
    }
    if setup_state not in allowed:
        raise NativeP2TestInstallError("P2受保护准备阶段非法")
    coordination = _require_coordination(server_root, transaction_id, generation, key)
    already_started = bool(coordination.get("business_write_started"))
    if already_started and business_write_started is False:
        raise NativeP2TestInstallError("P2受保护准备状态禁止回退业务写入标记")
    if setup_state == "ABORTED" and already_started:
        raise NativeP2TestInstallError("P2业务写入已开始，禁止按零写入准备事务中止")
    coordination["setup_state"] = setup_state
    if business_write_started is not None:
        coordination["business_write_started"] = business_write_started
    return _write_coordination(coordination, key)


def _require_coordination(
    server_root: Path,
    transaction_id: str,
    generation: str,
    key: bytes,
) -> dict[str, object]:
    coordination = _read_coordination(server_root, key, required=True)
    if (
        not coordination
        or coordination.get("transaction_id") != transaction_id
        or coordination.get("generation") != generation
    ):
        raise NativeP2TestInstallError("P2事务不匹配受保护协调安装代，拒绝旧事务复活或回滚")
    return coordination


def _set_coordination_owner_state(
    server_root: Path,
    transaction_id: str,
    generation: str,
    key: bytes,
    owner_state: str,
) -> dict[str, object]:
    if owner_state not in {"none", "pending_active", "active", "pending_closed", "closed"}:
        raise NativeP2TestInstallError("P2受保护协调归属状态非法")
    coordination = _require_coordination(server_root, transaction_id, generation, key)
    coordination["owner_state"] = owner_state
    return _write_coordination(coordination, key)


def _reservation_files(plan: NativeP2InstallPlan) -> list[dict[str, object]]:
    return [
        {
            "index": index,
            "relative_path": item.relative_path,
            "before_sha256": item.before_sha256,
            "after_sha256": item.after_sha256,
        }
        for index, item in enumerate(plan.files)
    ]


def _abort_preparation_reservation(
    base: Path,
    transaction: Path,
    server: Path,
    transaction_id: str,
    key: bytes,
    coordination: dict[str, object],
    targets: list[_HeldFileLeaf],
) -> NativeP2RollbackReceipt:
    """Idempotently close a protected current-generation reservation proven to have zero business writes."""
    generation = str(coordination.get("generation", ""))
    if (
        coordination.get("transaction_id") != transaction_id
        or coordination.get("business_write_started") is not False
        or coordination.get("setup_state") not in {
            "RESERVED", "MANIFEST_JSON_WRITTEN", "MANIFEST_HMAC_WRITTEN",
            "MANIFEST_VERIFIED", "ABORTED",
        }
        or coordination.get("journal_sequence") != 0
        or coordination.get("journal_head_hmac") != ""
        or coordination.get("pending_event") is not None
    ):
        raise NativeP2TestInstallError("P2准备事务没有可证明的当前代零业务写入中止状态")
    reservation = coordination.get("reservation_files")
    if not isinstance(reservation, list) or len(reservation) != len(TARGET_RELATIVES):
        raise NativeP2TestInstallError("P2准备事务缺少受保护的三目标预留绑定")
    for index, (entry, leaf) in enumerate(zip(reservation, targets)):
        if not isinstance(entry, dict) or entry.get("index") != index or entry.get("relative_path") != TARGET_RELATIVES[index]:
            raise NativeP2TestInstallError("P2准备事务受保护目标预留合同不匹配")
        current = _refresh_target_leaf(leaf)
        current_hash = _sha256_bytes(current) if current is not None else None
        if current_hash != entry.get("before_sha256"):
            raise NativeP2TestInstallError("P2准备事务目标不再是受保护预留的零写入状态，失败关闭")
    if coordination.get("setup_state") != "ABORTED" or coordination.get("phase") != "ABORTED" or coordination.get("owner_state") != "closed":
        coordination["setup_state"] = "ABORTED"
        coordination["phase"] = "ABORTED"
        coordination["owner_state"] = "closed"
        coordination = _write_coordination(coordination, key)
    owner = _read_active_owner(base, server, key)
    if not owner or owner.get("transaction_id") != transaction_id or owner.get("generation") != generation or owner.get("state") != "closed":
        _write_active_owner(base, server, transaction_id, generation, key, "closed")
    receipt_path = transaction / "abort-receipt.json"
    _atomic_write(receipt_path, _canonical_json_bytes({
        "schema_version": 1,
        "operation": "mingge-native-p2-test-preparation-abort",
        "status": "rolled-back",
        "transaction_id": transaction_id,
        "generation": generation,
        "business_files_written": False,
    }))
    return NativeP2RollbackReceipt(transaction_id, "rolled-back", receipt_path)


def _validate_journal_event_contract(
    event: dict[str, object],
    generation: str,
) -> None:
    """Validate fields shared by every ordinary or pending journal event."""
    file_index = event.get("file_index")
    relative_path = event.get("relative_path")
    observed_before_length = event.get("observed_before_length")
    if (
        event.get("generation") != generation
        or file_index is None and relative_path is not None
        or file_index is not None and (
            isinstance(file_index, bool)
            or not isinstance(file_index, int)
            or not 0 <= file_index < len(TARGET_RELATIVES)
            or relative_path != TARGET_RELATIVES[file_index]
        )
        or observed_before_length is not None and (
            isinstance(observed_before_length, bool)
            or not isinstance(observed_before_length, int)
            or observed_before_length < 0
        )
    ):
        raise NativeP2TestInstallError("P2事务日志安装代、固定目标或源长度合同不匹配")


def _read_journal_files(
    transaction: Path,
    key: bytes,
    transaction_id: str,
    generation: str,
) -> list[dict[str, object]]:
    journal_dir = transaction / "journal"
    if not journal_dir.exists():
        return []
    events: list[dict[str, object]] = []
    previous = ""
    paths = sorted(journal_dir.glob("*.json"))
    for expected_sequence, path in enumerate(paths, start=1):
        envelope = json.loads(path.read_text(encoding="utf-8"))
        signature = envelope.pop("hmac", None)
        if envelope.get("sequence") != expected_sequence or path.stem != f"{expected_sequence:06d}":
            raise NativeP2TestInstallError("P2事务日志序号不连续")
        if envelope.get("transaction_id") != transaction_id or envelope.get("previous_event_hmac") != previous:
            raise NativeP2TestInstallError("P2事务日志链断裂")
        _verify_seal(envelope, signature, key, JOURNAL_DOMAIN, "P2 journal")
        _validate_journal_event_contract(envelope, generation)
        previous = str(signature)
        envelope["hmac"] = signature
        events.append(envelope)
    return events


def _read_journal(
    transaction: Path,
    key: bytes,
    transaction_id: str,
    server_root: Path,
    generation: str,
    *,
    reconcile_pending: bool = False,
) -> list[dict[str, object]]:
    coordination = _require_coordination(server_root, transaction_id, generation, key)
    events = _read_journal_files(transaction, key, transaction_id, generation)
    pending = coordination.get("pending_event")
    confirmed_sequence = int(coordination["journal_sequence"])
    confirmed_head = str(coordination["journal_head_hmac"])
    if pending is not None:
        if not reconcile_pending or not isinstance(pending, dict):
            raise NativeP2TestInstallError("P2 journal存在受保护的未决事件，必须显式恢复")
        if pending.get("sequence") != confirmed_sequence + 1 or pending.get("previous_event_hmac") != confirmed_head:
            raise NativeP2TestInstallError("P2 journal未决事件与受保护尾锚不匹配")
        _verify_seal(
            {name: value for name, value in pending.items() if name != "hmac"},
            pending.get("hmac"), key, JOURNAL_DOMAIN, "P2 pending journal",
        )
        _validate_journal_event_contract(pending, generation)
        if len(events) == confirmed_sequence:
            journal_dir = transaction / "journal"
            journal_dir.mkdir(parents=True, exist_ok=True)
            _atomic_write(journal_dir / f"{confirmed_sequence + 1:06d}.json", _canonical_json_bytes(pending))
        elif len(events) != confirmed_sequence + 1 or events[-1] != pending:
            raise NativeP2TestInstallError("P2 journal未决事件与平台日志不一致")
        coordination["journal_sequence"] = confirmed_sequence + 1
        coordination["journal_head_hmac"] = str(pending["hmac"])
        coordination["phase"] = str(pending["state"])
        coordination["pending_event"] = None
        coordination = _write_coordination(coordination, key)
        events = _read_journal_files(transaction, key, transaction_id, generation)
        confirmed_sequence = int(coordination["journal_sequence"])
        confirmed_head = str(coordination["journal_head_hmac"])
    if len(events) != confirmed_sequence:
        raise NativeP2TestInstallError("P2 journal缺失、截短或含未锚定尾事件")
    actual_head = str(events[-1]["hmac"]) if events else ""
    if actual_head != confirmed_head:
        raise NativeP2TestInstallError("P2 journal尾部与受保护尾锚不匹配")
    expected_phase = str(events[-1]["state"]) if events else "MANIFEST_ONLY"
    if coordination.get("phase") != expected_phase:
        raise NativeP2TestInstallError("P2 journal状态与受保护协调阶段不匹配")
    return events


def _append_journal(
    transaction: Path,
    key: bytes,
    transaction_id: str,
    state: str,
    action: str,
    file_index: int | None = None,
    observed_before_sha256: str | None = None,
    observed_after_sha256: str | None = None,
    observed_before_length: int | None = None,
    *,
    server_root: Path,
    generation: str,
) -> dict[str, object]:
    events = _read_journal(
        transaction, key, transaction_id, server_root, generation, reconcile_pending=True,
    )
    coordination = _require_coordination(server_root, transaction_id, generation, key)
    sequence = int(coordination["journal_sequence"]) + 1
    previous = str(coordination["journal_head_hmac"])
    event: dict[str, object] = {
        "sequence": sequence,
        "transaction_id": transaction_id,
        "generation": generation,
        "state": state,
        "action": action,
        "file_index": file_index,
        "relative_path": TARGET_RELATIVES[file_index] if isinstance(file_index, int) and not isinstance(file_index, bool) and 0 <= file_index < len(TARGET_RELATIVES) else None,
        "observed_before_sha256": observed_before_sha256,
        "observed_after_sha256": observed_after_sha256,
        "observed_before_length": observed_before_length,
        "previous_event_hmac": previous,
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
    }
    event["hmac"] = _seal(event, key, JOURNAL_DOMAIN)
    coordination["pending_event"] = event
    _write_coordination(coordination, key)
    journal_dir = transaction / "journal"
    journal_dir.mkdir(parents=True, exist_ok=True)
    _atomic_write(journal_dir / f"{sequence:06d}.json", _canonical_json_bytes(event))
    coordination["journal_sequence"] = sequence
    coordination["journal_head_hmac"] = str(event["hmac"])
    coordination["phase"] = state
    coordination["pending_event"] = None
    _write_coordination(coordination, key)
    _read_journal(transaction, key, transaction_id, server_root, generation)
    return event


def _active_owner_path(base: Path, server_root: Path) -> Path:
    server_id = _sha256_bytes(str(Path(server_root).resolve()).lower().encode("utf-8"))[:24]
    path = base / "active" / f"{server_id}.json"
    if not path.resolve(strict=False).is_relative_to(base.resolve(strict=False)):
        raise NativeP2TestInstallError("P2活动事务归属文件越界")
    return path


def _read_active_owner(base: Path, server_root: Path, key: bytes) -> dict[str, object] | None:
    path = _active_owner_path(base, server_root)
    if not path.exists():
        return None
    if _is_reparse_point(path) or _is_reparse_point(path.parent):
        raise NativeP2TestInstallError("P2活动事务归属路径不能是重解析点")
    envelope = json.loads(path.read_text(encoding="utf-8"))
    signature = envelope.pop("hmac", None)
    _verify_seal(envelope, signature, key, ACTIVE_DOMAIN, "P2 active owner")
    if envelope.get("schema_version") != 1 or Path(str(envelope.get("server_root", ""))).resolve() != Path(server_root).resolve():
        raise NativeP2TestInstallError("P2活动事务归属合同不匹配")
    envelope["hmac"] = signature
    return envelope


def _write_active_owner(
    base: Path,
    server_root: Path,
    transaction_id: str,
    generation: str,
    key: bytes,
    state: str,
) -> None:
    if state not in {"active", "closed"}:
        raise NativeP2TestInstallError("P2活动事务归属状态非法")
    path = _active_owner_path(base, server_root)
    path.parent.mkdir(parents=True, exist_ok=True)
    if _is_reparse_point(path.parent):
        raise NativeP2TestInstallError("P2活动事务归属目录不能是重解析点")
    payload: dict[str, object] = {
        "schema_version": 1,
        "server_root": str(Path(server_root).resolve()),
        "transaction_id": transaction_id,
        "generation": generation,
        "state": state,
        "target_relatives": list(TARGET_RELATIVES),
        "updated_at_utc": datetime.now(timezone.utc).isoformat(),
    }
    payload["hmac"] = _seal(payload, key, ACTIVE_DOMAIN)
    _atomic_write(path, _canonical_json_bytes(payload))
    loaded = _read_active_owner(base, server_root, key)
    if not loaded or loaded.get("transaction_id") != transaction_id or loaded.get("generation") != generation or loaded.get("state") != state:
        raise NativeP2TestInstallError("P2活动事务归属写后验证失败")


def _require_active_owner(
    base: Path,
    server_root: Path,
    transaction_id: str,
    generation: str,
    key: bytes,
) -> dict[str, object]:
    owner = _read_active_owner(base, server_root, key)
    if not owner or owner.get("state") != "active" or owner.get("transaction_id") != transaction_id or owner.get("generation") != generation:
        raise NativeP2TestInstallError("P2事务不再拥有当前服务端安装代，拒绝旧事务覆盖")
    return owner


def _load_sealed_manifest(
    platform_root: Path,
    transaction_id: str,
    server_root: Path,
    *,
    reconcile_pending: bool = False,
) -> tuple[Path, dict[str, object], bytes, list[dict[str, object]]]:
    base = _transaction_base(platform_root)
    transaction = base / transaction_id
    if transaction.name != transaction_id or not transaction.resolve(strict=False).is_relative_to(base.resolve(strict=False)):
        raise NativeP2TestInstallError("P2事务ID非法或目录越界")
    manifest_path = transaction / "manifest.json"
    signature_path = transaction / "manifest.hmac"
    if not manifest_path.is_file() or not signature_path.is_file():
        raise NativeP2TestInstallError("P2事务缺少独立manifest或HMAC封印")
    key_id, key = _load_or_create_integrity_key(base)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("key_id") != key_id:
        raise NativeP2TestInstallError("P2事务manifest的密钥ID不匹配")
    _verify_seal(manifest, signature_path.read_text(encoding="ascii").strip(), key, MANIFEST_DOMAIN, "P2 manifest")
    if manifest.get("schema_version") != 4 or manifest.get("operation") != "mingge-native-p2-test" or manifest.get("transaction_id") != transaction_id:
        raise NativeP2TestInstallError("P2事务manifest基础合同不匹配")
    if not isinstance(manifest.get("generation"), str) or len(str(manifest.get("generation"))) < 16:
        raise NativeP2TestInstallError("P2事务manifest缺少有效安装代")
    server_entry = Path(server_root).absolute()
    if _is_reparse_point(server_entry):
        raise NativeP2TestInstallError("回滚服务端入口不能是重解析点")
    server = server_entry.resolve()
    if Path(str(manifest.get("server_root", ""))).resolve() != server:
        raise NativeP2TestInstallError("回滚目标服务端与P2 manifest不一致")
    candidate = manifest.get("candidate", {})
    if not isinstance(candidate, dict) or candidate != {
        "candidate_id": EXPECTED_CANDIDATE_ID,
        "candidate_spec_sha256": EXPECTED_CANDIDATE_SPEC_SHA256,
        "compiled_script_sha256": EXPECTED_COMPILED_SCRIPT_SHA256,
        "workbook_path": candidate.get("workbook_path") if isinstance(candidate, dict) else None,
        "workbook_sha256": candidate.get("workbook_sha256") if isinstance(candidate, dict) else None,
        "equip_slots": [17], "payment": False, "writes_database_directly": False,
        "runtime_item_instance_restored_by_file_rollback": False,
    }:
        raise NativeP2TestInstallError("P2 manifest候选或安全字段不符合冻结合同")
    engine = manifest.get("engine", {})
    if not isinstance(engine, dict) or engine.get("required_baseline") != ENGINE_VERSION_BASELINE or engine.get("detected_version") != EXPECTED_M2SERVER_VERSION or engine.get("binary_sha256") != EXPECTED_M2SERVER_SHA256:
        raise NativeP2TestInstallError("P2 manifest引擎身份不符合冻结合同")
    if manifest.get("evidence") != [asdict(item) for item in P2_EVIDENCE]:
        raise NativeP2TestInstallError("P2 manifest证据合同不匹配")
    entries = manifest.get("files")
    if not isinstance(entries, list) or len(entries) != len(TARGET_RELATIVES) or tuple(item.get("relative_path") for item in entries if isinstance(item, dict)) != TARGET_RELATIVES:
        raise NativeP2TestInstallError("P2 manifest目标集合、顺序或数量不符合固定合同")
    for index, entry in enumerate(entries):
        if not isinstance(entry, dict) or entry.get("index") != index:
            raise NativeP2TestInstallError("P2 manifest目标索引不符合固定合同")
        _entry_after_bytes(entry)
    events = _read_journal(
        transaction, key, transaction_id, server, str(manifest["generation"]),
        reconcile_pending=reconcile_pending,
    )
    return transaction, manifest, key, events


def _entry_after_bytes(entry: dict[str, object]) -> bytes:
    encoded = entry.get("after_payload_b64")
    if not isinstance(encoded, str):
        raise NativeP2TestInstallError("P2 manifest缺少封印的安装后目标内容")
    try:
        payload = b64decode(encoded, validate=True)
    except ValueError as exc:
        raise NativeP2TestInstallError("P2 manifest安装后目标内容不是有效Base64") from exc
    if _sha256_bytes(payload) != entry.get("after_sha256") or len(payload) != entry.get("after_length"):
        raise NativeP2TestInstallError("P2 manifest安装后目标内容与封印哈希或长度不匹配")
    return payload


def _classify_transaction_files(server: Path, entries: list[dict[str, object]]) -> list[str]:
    classifications: list[str] = []
    for relative, entry in zip(TARGET_RELATIVES, entries):
        path = server / Path(relative)
        _assert_no_reparse_path(path, server)
        current = _current_hash(path)
        if current == entry.get("before_sha256"):
            classifications.append("before")
        elif current == entry.get("after_sha256"):
            classifications.append("after")
        else:
            classifications.append("unknown")
    return classifications


@contextmanager
def _hold_verified_manifest_backups(
    transaction: Path,
    entries: list[dict[str, object]],
) -> list[_HeldFileLeaf]:
    """Validate backup bytes from no-write/no-delete-share handles and keep those handles held."""
    leaves: list[_HeldFileLeaf] = []
    try:
        for index, entry in enumerate(entries):
            if entry.get("index") != index or entry.get("backup_relative_path") not in (f"files/{index:02d}.bin", None):
                raise NativeP2TestInstallError("P2 manifest备份索引或固定相对路径不匹配")
            backup = transaction / "files" / f"{index:02d}.bin"
            if entry.get("existed_before"):
                if not os.path.lexists(backup):
                    raise NativeP2TestInstallError(f"P2回滚备份缺失或哈希不匹配：{entry.get('relative_path')}")
                handle = _open_validated_file_handle(backup, "回滚备份", writable=False)
                payload = _read_file_handle(handle)
                leaf = _HeldFileLeaf(backup, handle, payload)
                leaves.append(leaf)
                if (
                    _sha256_bytes(payload) != entry.get("backup_sha256")
                    or entry.get("backup_sha256") != entry.get("before_sha256")
                    or len(payload) != entry.get("before_length")
                ):
                    raise NativeP2TestInstallError(f"P2回滚备份缺失或哈希不匹配：{entry.get('relative_path')}")
            else:
                if entry.get("backup_relative_path") is not None or entry.get("backup_sha256") is not None or os.path.lexists(backup):
                    raise NativeP2TestInstallError("原本不存在的P2目标不应带备份")
                leaves.append(_HeldFileLeaf(backup, None, None))
        yield leaves
    finally:
        _close_file_leaves(leaves)


def _verify_manifest_backups(transaction: Path, entries: list[dict[str, object]]) -> None:
    with _hold_verified_manifest_backups(transaction, entries):
        pass


def _refresh_target_leaf(leaf: _HeldFileLeaf) -> bytes | None:
    if leaf.delete_pending:
        return None
    if leaf.handle is None and os.path.lexists(leaf.path):
        leaf.handle = _open_validated_file_handle(leaf.path, "业务目标", writable=True)
        leaf.payload = _read_file_handle(leaf.handle)
    elif leaf.handle is not None:
        leaf.payload = _read_file_handle(leaf.handle)
    return leaf.payload


def _classify_held_targets(
    leaves: list[_HeldFileLeaf],
    entries: list[dict[str, object]],
    backups: list[_HeldFileLeaf],
    events: list[dict[str, object]],
    *,
    allow_current_intent_prefix: bool,
) -> list[str]:
    classifications: list[str] = []
    pending_index: int | None = None
    pending_payload: bytes | None = None
    if allow_current_intent_prefix and events:
        event = events[-1]
        action = event.get("action")
        index = event.get("file_index")
        if isinstance(index, int) and 0 <= index < len(entries):
            if action == "WRITE_INTENT" and event.get("state") == "INSTALLING":
                candidate = _entry_after_bytes(entries[index])
                if event.get("observed_after_sha256") == entries[index].get("after_sha256"):
                    pending_index, pending_payload = index, candidate
            elif action == "ROLLBACK_INTENT" and event.get("state") == "ROLLING_BACK":
                candidate = backups[index].payload
                if event.get("observed_after_sha256") == entries[index].get("before_sha256"):
                    pending_index, pending_payload = index, candidate
    for index, (leaf, entry) in enumerate(zip(leaves, entries)):
        current = _refresh_target_leaf(leaf)
        current_hash = _sha256_bytes(current) if current is not None else None
        if current_hash == entry.get("before_sha256"):
            classifications.append("before")
        elif current_hash == entry.get("after_sha256"):
            classifications.append("after")
        elif (
            index == pending_index
            and current is not None
            and pending_payload is not None
            and len(current) <= len(pending_payload)
            and pending_payload.startswith(current)
        ):
            classifications.append("intent-prefix")
        else:
            classifications.append("unknown")
    return classifications


def _terminal_owner_is_bound(
    coordination: dict[str, object],
    active_owner: dict[str, object] | None,
    transaction_id: str,
    generation: str,
) -> bool:
    if active_owner is None:
        return True
    if (
        active_owner.get("transaction_id") == transaction_id
        and active_owner.get("generation") == generation
        and active_owner.get("state") in {"active", "closed"}
    ):
        return True
    predecessor = coordination.get("predecessor_owner")
    return bool(
        isinstance(predecessor, dict)
        and predecessor.get("state") == "closed"
        and active_owner.get("state") == "closed"
        and active_owner.get("transaction_id") == predecessor.get("transaction_id")
        and active_owner.get("generation") == predecessor.get("generation")
    )


def _nonactive_terminal_owner_is_bound(
    coordination: dict[str, object],
    active_owner: dict[str, object] | None,
    transaction_id: str,
    generation: str,
) -> bool:
    if active_owner is None:
        return True
    if (
        active_owner.get("transaction_id") == transaction_id
        and active_owner.get("generation") == generation
        and active_owner.get("state") == "closed"
    ):
        return True
    predecessor = coordination.get("predecessor_owner")
    return bool(
        isinstance(predecessor, dict)
        and predecessor.get("state") == "closed"
        and active_owner.get("state") == "closed"
        and active_owner.get("transaction_id") == predecessor.get("transaction_id")
        and active_owner.get("generation") == predecessor.get("generation")
    )


def _rollback_event_index_matches(
    event: dict[str, object],
    entries: list[dict[str, object]],
    generation: str,
) -> bool:
    index = event.get("file_index")
    return bool(
        isinstance(index, int)
        and not isinstance(index, bool)
        and 0 <= index < len(entries)
        and event.get("generation") == generation
        and event.get("relative_path") == TARGET_RELATIVES[index]
        and entries[index].get("index") == index
        and entries[index].get("relative_path") == TARGET_RELATIVES[index]
    )


def _rollback_source_matches_entry(event: dict[str, object], entry: dict[str, object]) -> bool:
    length = event.get("observed_before_length")
    if isinstance(length, bool) or not isinstance(length, int) or length < 0:
        return False
    source_hash = event.get("observed_before_sha256")
    before_length = entry.get("before_length")
    if source_hash == entry.get("before_sha256") and length == before_length:
        return True
    after_payload = _entry_after_bytes(entry)
    return bool(
        length <= len(after_payload)
        and source_hash == _sha256_bytes(after_payload[:length])
    )


def _paired_rollback_done(
    events: list[dict[str, object]],
    entries: list[dict[str, object]],
    generation: str,
) -> bool:
    if len(events) < 2:
        return False
    intent, done = events[-2:]
    if (
        done.get("state") != "ROLLING_BACK"
        or done.get("action") != "ROLLBACK_DONE"
        or not _rollback_event_index_matches(done, entries, generation)
    ):
        return False
    index = int(done["file_index"])
    entry = entries[index]
    return bool(
        done.get("observed_after_sha256") == entry.get("before_sha256")
        and _rollback_source_matches_entry(done, entry)
        and intent.get("state") == "ROLLING_BACK"
        and intent.get("action") == "ROLLBACK_INTENT"
        and _rollback_event_index_matches(intent, entries, generation)
        and intent.get("file_index") == index
        and intent.get("observed_before_sha256") == done.get("observed_before_sha256")
        and intent.get("observed_before_length") == done.get("observed_before_length")
        and intent.get("observed_after_sha256") == entry.get("before_sha256")
    )


def _recovery_required_action_matches(
    event: dict[str, object],
    entries: list[dict[str, object]],
    generation: str,
) -> bool:
    if event.get("state") != "RECOVERY_REQUIRED" or event.get("generation") != generation:
        return False
    action = event.get("action")
    if action in {"UNKNOWN_TARGET_DRIFT", "AUTO_RECOVERY_FAILED"}:
        return bool(
            event.get("file_index") is None
            and event.get("relative_path") is None
            and event.get("observed_before_sha256") is None
            and event.get("observed_after_sha256") is None
            and event.get("observed_before_length") is None
        )
    if action == "ROLLBACK_VERIFY_FAILED":
        return bool(
            _rollback_event_index_matches(event, entries, generation)
            and event.get("observed_before_sha256") is None
            and event.get("observed_after_sha256") is None
            and event.get("observed_before_length") is None
        )
    if action == "ROLLBACK_DRIFT" and _rollback_event_index_matches(event, entries, generation):
        source_hash = event.get("observed_before_sha256")
        source_length = event.get("observed_before_length")
        valid_source = bool(
            source_hash is None
            and isinstance(source_length, int)
            and not isinstance(source_length, bool)
            and source_length == 0
            or isinstance(source_hash, str)
            and len(source_hash) == 64
            and all(character in "0123456789abcdef" for character in source_hash)
            and isinstance(source_length, int)
            and not isinstance(source_length, bool)
            and source_length >= 0
        )
        return valid_source and event.get("observed_after_sha256") is None
    return False


def _require_recovery_required_terminal(
    coordination: dict[str, object],
    active_owner: dict[str, object] | None,
    events: list[dict[str, object]],
    entries: list[dict[str, object]],
    classifications: list[str],
    transaction_id: str,
    generation: str,
) -> None:
    if (
        coordination.get("transaction_id") != transaction_id
        or coordination.get("generation") != generation
        or coordination.get("phase") != "RECOVERY_REQUIRED"
        or coordination.get("setup_state") != "BUSINESS_WRITES"
        or coordination.get("business_write_started") is not True
        or not events
        or not _recovery_required_action_matches(events[-1], entries, generation)
        or any(value != "before" for value in classifications)
    ):
        raise NativeP2TestInstallError(
            "RECOVERY_REQUIRED只允许当前受保护代的全before terminal-only收尾"
        )
    owner_state = coordination.get("owner_state")
    if owner_state == "active":
        owner_matches = bool(
            active_owner
            and active_owner.get("transaction_id") == transaction_id
            and active_owner.get("generation") == generation
            and active_owner.get("state") == "active"
        )
    elif owner_state == "none":
        owner_matches = _nonactive_terminal_owner_is_bound(
            coordination, active_owner, transaction_id, generation,
        )
    else:
        owner_matches = False
    if not owner_matches:
        raise NativeP2TestInstallError("RECOVERY_REQUIRED受保护归属来源不匹配")


def _require_pending_closed_terminal(
    coordination: dict[str, object],
    active_owner: dict[str, object] | None,
    events: list[dict[str, object]],
    entries: list[dict[str, object]],
    classifications: list[str],
    transaction_id: str,
    generation: str,
) -> None:
    if (
        coordination.get("transaction_id") != transaction_id
        or coordination.get("generation") != generation
        or coordination.get("owner_state") != "pending_closed"
        or coordination.get("business_write_started") is not True
        or coordination.get("setup_state") != "BUSINESS_WRITES"
        or not events
        or coordination.get("phase") != events[-1].get("state")
        or any(value != "before" for value in classifications)
        or not _terminal_owner_is_bound(
            coordination, active_owner, transaction_id, generation,
        )
    ):
        raise NativeP2TestInstallError("待关闭回滚不满足受保护terminal-only状态表")
    tail = events[-1]
    if _recovery_required_action_matches(tail, entries, generation):
        return
    generic_tail = bool(
        tail.get("file_index") is None
        and tail.get("relative_path") is None
        and tail.get("observed_before_sha256") is None
        and tail.get("observed_after_sha256") is None
        and tail.get("observed_before_length") is None
    )
    allowed_generic = (
        ("ROLLING_BACK", "ROLLBACK_START"),
        ("PREPARED", "TRANSACTION_PREPARED"),
        ("ROLLED_BACK", "ROLLBACK_COMPLETE"),
        ("ROLLED_BACK", "AUTO_RECOVERY_COMPLETE"),
    )
    if generic_tail and (tail.get("state"), tail.get("action")) in allowed_generic:
        return
    if _paired_rollback_done(events, entries, generation):
        return
    raise NativeP2TestInstallError("待关闭回滚缺少合法的受保护终结来源")


def _nonactive_rollback_resume_mode(
    coordination: dict[str, object],
    active_owner: dict[str, object] | None,
    events: list[dict[str, object]],
    entries: list[dict[str, object]],
    classifications: list[str],
    transaction_id: str,
    generation: str,
) -> str:
    """Return the protected non-active ROLLING_BACK state-table transition."""
    if (
        coordination.get("transaction_id") != transaction_id
        or coordination.get("generation") != generation
        or coordination.get("phase") != "ROLLING_BACK"
        or coordination.get("business_write_started") is not True
        or coordination.get("setup_state") != "BUSINESS_WRITES"
        or not events
        or events[-1].get("state") != "ROLLING_BACK"
    ):
        raise NativeP2TestInstallError("非活动回滚不满足当前受保护业务写入状态表")

    tail = events[-1]
    tail_index = tail.get("file_index")
    tail_index_matches = _rollback_event_index_matches(tail, entries, generation)

    if coordination.get("owner_state") == "none":
        if active_owner is not None and active_owner.get("state") != "closed":
            raise NativeP2TestInstallError("非活动回滚发现普通活动归属，失败关闭")
        if tail.get("action") == "ROLLBACK_INTENT" and tail_index_matches:
            entry = entries[tail_index]
            if (
                tail.get("observed_after_sha256") != entry.get("before_sha256")
                or not _rollback_source_matches_entry(tail, entry)
            ):
                raise NativeP2TestInstallError("非活动回滚逐文件意图的目标before哈希不匹配")
            if any(value != "before" for value in classifications[tail_index + 1:]):
                raise NativeP2TestInstallError("非活动回滚逐文件意图不符合固定逆序")
            classification = classifications[tail_index]
            if classification in {"before", "intent-prefix"}:
                return "business"
            if (
                classification == "after"
                and tail.get("observed_before_sha256") == entry.get("after_sha256")
            ):
                return "business"
        elif _paired_rollback_done(events, entries, generation):
            if (
                classifications[tail_index] == "before"
                and not any(value != "before" for value in classifications[tail_index + 1:])
            ):
                return "business"
        elif (
            tail.get("action") == "ROLLBACK_START"
            and tail.get("file_index") is None
            and tail.get("relative_path") is None
            and tail.get("observed_before_sha256") is None
            and tail.get("observed_after_sha256") is None
            and tail.get("observed_before_length") is None
            and all(value in {"before", "after"} for value in classifications)
        ):
            return "business"
        raise NativeP2TestInstallError("非活动回滚缺少匹配的逐文件意图或通用开始锚")

    raise NativeP2TestInstallError("非活动回滚的受保护归属状态不在允许状态表")


def _transaction_fault(name: str) -> None:
    """Test-only fault boundary; production execution is a no-op."""
    _ = name


def install_native_p2_test_candidate(plan: NativeP2InstallPlan, platform_root: Path) -> NativeP2InstallReceipt:
    """事务安装P2固定候选；manifest与日志封印成功后才写三个固定目标。"""
    _assert_plan_is_current(plan)
    with _transaction_lock(platform_root, plan.server_root) as base:
        target_parents = tuple(dict.fromkeys(Path(relative).parent for relative in TARGET_RELATIVES))
        with (
            _hold_directory_tree(plan.server_root, target_parents, create=True),
            _hold_target_leaves(plan.server_root) as target_leaves,
        ):
            _assert_plan_is_current(plan, target_leaves)
            for item, leaf in zip(plan.files, target_leaves):
                current = _refresh_target_leaf(leaf)
                current_hash = _sha256_bytes(current) if current is not None else None
                if current_hash != item.before_sha256:
                    raise NativeP2TestInstallError(f"目标文件在排他叶子句柄内发生变化：{item.path}")
            key_id, key = _load_or_create_integrity_key(base)
            with _hold_secret_tree():
                coordination = _read_coordination(plan.server_root, key, required=False)
                active_owner = _read_active_owner(base, plan.server_root, key)
                predecessor_owner: dict[str, object] | None = None
                historical = [
                    item for item in base.iterdir()
                    if item.is_dir() and item.name != "active" and (item / "manifest.json").exists()
                ]
                if coordination is None:
                    if active_owner or historical:
                        raise NativeP2TestInstallError("P2受保护协调状态缺失但普通备份树已有归属或历史，失败关闭")
                else:
                    if coordination.get("owner_state") != "closed" or coordination.get("phase") not in {"ROLLED_BACK", "ABORTED"}:
                        raise NativeP2TestInstallError(
                            f"当前服务端已有未关闭P2协调安装代：{coordination.get('transaction_id')}"
                        )
                    if (
                        not active_owner
                        or active_owner.get("state") != "closed"
                        or active_owner.get("transaction_id") != coordination.get("transaction_id")
                        or active_owner.get("generation") != coordination.get("generation")
                    ):
                        raise NativeP2TestInstallError("P2普通归属文件缺失或不匹配受保护协调历史，失败关闭")
                    predecessor_owner = {
                        "transaction_id": str(active_owner["transaction_id"]),
                        "generation": str(active_owner["generation"]),
                        "state": "closed",
                    }

                transaction_id = datetime.now().strftime("%Y%m%d_%H%M%S_") + uuid.uuid4().hex[:8]
                generation = uuid.uuid4().hex
                transaction = base / transaction_id
                if transaction.exists():
                    raise NativeP2TestInstallError(f"事务目录已存在：{transaction}")
                transaction.mkdir()
                control_dirs = (
                    Path(transaction_id),
                    Path(transaction_id) / "files",
                    Path(transaction_id) / "journal",
                )
                with _hold_directory_tree(base, control_dirs, create=True):
                    files_dir = transaction / "files"
                    _begin_coordination(
                        plan.server_root, transaction_id, generation, key, _reservation_files(plan),
                        predecessor_owner,
                    )
                    _transaction_fault("after_coordination_reservation")
                    entries: list[dict[str, object]] = []
                    for index, (item, leaf) in enumerate(zip(plan.files, target_leaves)):
                        before = _refresh_target_leaf(leaf)
                        current_hash = _sha256_bytes(before) if before is not None else None
                        if current_hash != item.before_sha256:
                            raise NativeP2TestInstallError(f"目标文件在排他叶子句柄内发生变化：{item.path}")
                        backup = files_dir / f"{index:02d}.bin"
                        if before is not None:
                            _atomic_write(backup, before)
                        entries.append({
                            "index": index, "relative_path": item.relative_path, "existed_before": before is not None,
                            "before_sha256": item.before_sha256, "after_sha256": item.after_sha256,
                            "before_length": len(before) if before is not None else 0, "after_length": len(item.after),
                            "after_payload_b64": b64encode(item.after).decode("ascii"),
                            "backup_relative_path": f"files/{index:02d}.bin" if before is not None else None,
                            "backup_sha256": item.before_sha256 if before is not None else None,
                        })
                    with _hold_verified_manifest_backups(transaction, entries) as backup_leaves:
                        receipt_path = transaction / "receipt.json"
                        manifest = _manifest_payload(plan, transaction_id, platform_root, key_id, generation, entries)
                        _write_manifest(
                            transaction, manifest, key,
                            server_root=plan.server_root,
                            transaction_id=transaction_id,
                            generation=generation,
                        )
                        _append_journal(
                            transaction, key, transaction_id, "PREPARED", "TRANSACTION_PREPARED",
                            server_root=plan.server_root, generation=generation,
                        )
                        _set_coordination_setup_state(
                            plan.server_root, transaction_id, generation, key, "PREPARED",
                        )
                        _atomic_write(receipt_path, (json.dumps(_receipt_payload(plan, transaction_id, entries, "PREPARED"), ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
                        written: list[int] = []
                        try:
                            _set_coordination_setup_state(
                                plan.server_root, transaction_id, generation, key, "BUSINESS_WRITES",
                                business_write_started=True,
                            )
                            for item, leaf in zip(plan.files, target_leaves):
                                current = _refresh_target_leaf(leaf)
                                current_hash = _sha256_bytes(current) if current is not None else None
                                if current_hash != item.before_sha256:
                                    raise NativeP2TestInstallError(f"目标文件最终排他分类发生变化：{item.path}")
                            _transaction_fault("after_target_leaves_validated_for_install")
                            for index, (item, leaf) in enumerate(zip(plan.files, target_leaves)):
                                current = _refresh_target_leaf(leaf)
                                current_hash = _sha256_bytes(current) if current is not None else None
                                if current_hash != item.before_sha256:
                                    raise NativeP2TestInstallError(f"目标文件写入前发生变化：{item.path}")
                                _append_journal(
                                    transaction, key, transaction_id, "INSTALLING", "WRITE_INTENT", index,
                                    current_hash, item.after_sha256,
                                    len(current) if current is not None else 0,
                                    server_root=plan.server_root, generation=generation,
                                )
                                _create_missing_target_handle(leaf, index)
                                written.append(index)
                                _write_file_handle_exact(
                                    leaf, item.after, f"install_after_target_handle_truncate:{index}",
                                )
                                if _sha256_bytes(leaf.payload or b"") != item.after_sha256:
                                    raise NativeP2TestInstallError(f"安装后同句柄哈希不匹配：{item.path}")
                                _append_journal(
                                    transaction, key, transaction_id, "INSTALLING", "WRITE_DONE", index,
                                    item.before_sha256, item.after_sha256,
                                    len(current) if current is not None else 0,
                                    server_root=plan.server_root, generation=generation,
                                )
                                state = _receipt_payload(plan, transaction_id, entries, "INSTALLING")
                                state["written_count"] = index + 1
                                _atomic_write(receipt_path, (json.dumps(state, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
                            for item, leaf in zip(plan.files, target_leaves):
                                current = _refresh_target_leaf(leaf)
                                if current is None or _sha256_bytes(current) != item.after_sha256:
                                    raise NativeP2TestInstallError(f"安装后同句柄哈希不匹配：{item.path}")
                            _set_coordination_owner_state(
                                plan.server_root, transaction_id, generation, key, "pending_active",
                            )
                            _append_journal(
                                transaction, key, transaction_id, "INSTALLED", "INSTALL_COMPLETE",
                                server_root=plan.server_root, generation=generation,
                            )
                        except Exception as original_error:
                            recovery_error: Exception | None = None
                            for index in reversed(written):
                                entry = entries[index]
                                leaf = target_leaves[index]
                                try:
                                    current = _refresh_target_leaf(leaf)
                                    current_hash = _sha256_bytes(current) if current is not None else None
                                    current_length = len(current) if current is not None else 0
                                    _append_journal(
                                        transaction, key, transaction_id, "ROLLING_BACK", "ROLLBACK_INTENT", index,
                                        current_hash, entry.get("before_sha256"), current_length,
                                        server_root=plan.server_root, generation=generation,
                                    )
                                    before_payload = backup_leaves[index].payload
                                    if before_payload is None:
                                        if leaf.handle is not None:
                                            _delete_file_handle_exact(
                                                leaf, f"auto_recovery_before_delete_target_handle:{index}",
                                            )
                                    else:
                                        if leaf.handle is None:
                                            leaf.handle = _open_validated_file_handle(
                                                leaf.path, "自动恢复目标", writable=True, create_new=True,
                                            )
                                            leaf.created = True
                                        _write_file_handle_exact(
                                            leaf, before_payload,
                                            f"auto_recovery_after_target_handle_truncate:{index}",
                                        )
                                    _append_journal(
                                        transaction, key, transaction_id, "ROLLING_BACK", "ROLLBACK_DONE", index,
                                        current_hash, entry.get("before_sha256"), current_length,
                                        server_root=plan.server_root, generation=generation,
                                    )
                                except Exception as exc:  # pragma: no cover - extreme recovery path
                                    recovery_error = exc
                            for index, (entry, leaf) in enumerate(zip(entries, target_leaves)):
                                try:
                                    current = _refresh_target_leaf(leaf)
                                    current_hash = _sha256_bytes(current) if current is not None else None
                                    if current_hash != entry.get("before_sha256"):
                                        recovery_error = NativeP2TestInstallError(
                                            f"自动恢复后目标仍有未知漂移：{TARGET_RELATIVES[index]}"
                                        )
                                except Exception as exc:
                                    recovery_error = exc
                            status = "RECOVERY_REQUIRED" if recovery_error else "RECOVERED_AFTER_INSTALL_FAILURE"
                            try:
                                if recovery_error:
                                    _set_coordination_owner_state(
                                        plan.server_root, transaction_id, generation, key, "none",
                                    )
                                    _append_journal(
                                        transaction, key, transaction_id,
                                        "RECOVERY_REQUIRED", "AUTO_RECOVERY_FAILED",
                                        server_root=plan.server_root, generation=generation,
                                    )
                                else:
                                    _set_coordination_owner_state(
                                        plan.server_root, transaction_id, generation, key, "pending_closed",
                                    )
                                    _append_journal(
                                        transaction, key, transaction_id,
                                        "ROLLED_BACK", "AUTO_RECOVERY_COMPLETE",
                                        server_root=plan.server_root, generation=generation,
                                    )
                                    _write_active_owner(base, plan.server_root, transaction_id, generation, key, "closed")
                                    _set_coordination_owner_state(
                                        plan.server_root, transaction_id, generation, key, "closed",
                                    )
                            except Exception:
                                pass
                            try:
                                _atomic_write(receipt_path, (json.dumps(_receipt_payload(plan, transaction_id, entries, status), ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
                            except Exception:
                                pass
                            if recovery_error:
                                raise NativeP2TestInstallError(f"安装失败且自动恢复不完整：{recovery_error}") from recovery_error
                            raise original_error

                        _transaction_fault("after_installed_journal")
                        _write_active_owner(base, plan.server_root, transaction_id, generation, key, "active")
                        _set_coordination_owner_state(plan.server_root, transaction_id, generation, key, "active")
                        try:
                            _atomic_write(receipt_path, (json.dumps(_receipt_payload(plan, transaction_id, entries, "INSTALLED"), ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
                        except Exception:
                            # receipt只用于展示；可信状态由受保护协调状态和完整HMAC日志共同封印。
                            pass
                        return NativeP2InstallReceipt(transaction_id, "installed-test-candidate", plan.candidate_id, receipt_path, TARGET_RELATIVES)


def inspect_native_p2_test_transaction(platform_root: Path, transaction_id: str, server_root: Path) -> dict[str, object]:
    transaction, manifest, key, events = _load_sealed_manifest(platform_root, transaction_id, server_root)
    entries = list(manifest["files"])
    _verify_manifest_backups(transaction, entries)
    server = Path(server_root).resolve()
    active_owner = _read_active_owner(_transaction_base(platform_root), server, key)
    return {
        "transaction_id": transaction_id,
        "state": events[-1]["state"] if events else "MANIFEST_ONLY",
        "file_classification": dict(zip(TARGET_RELATIVES, _classify_transaction_files(server, entries))),
        "manifest_verified": True,
        "journal_events": len(events),
        "owns_active_generation": bool(
            active_owner and active_owner.get("state") == "active"
            and active_owner.get("transaction_id") == transaction_id
            and active_owner.get("generation") == manifest.get("generation")
        ),
    }


def _rollback_from_manifest(
    transaction: Path,
    manifest: dict[str, object],
    key: bytes,
    transaction_id: str,
    server: Path,
    targets: list[_HeldFileLeaf],
    backups: list[_HeldFileLeaf],
    events: list[dict[str, object]],
) -> None:
    generation = str(manifest["generation"])
    entries = list(manifest["files"])
    if events:
        tail = events[-1]
        if tail.get("state") == "ROLLING_BACK" and tail.get("action") == "ROLLBACK_INTENT":
            if not _rollback_event_index_matches(tail, entries, generation):
                raise NativeP2TestInstallError("逐文件回滚意图未绑定当前代固定目标")
            tail_entry = entries[int(tail["file_index"])]
            if (
                tail.get("observed_after_sha256") != tail_entry.get("before_sha256")
                or not _rollback_source_matches_entry(tail, tail_entry)
            ):
                raise NativeP2TestInstallError("逐文件回滚意图的source hash/length合同不匹配")
        elif tail.get("state") == "ROLLING_BACK" and tail.get("action") == "ROLLBACK_DONE":
            if not _paired_rollback_done(events, entries, generation):
                raise NativeP2TestInstallError("逐文件回滚完成锚未与相邻意图精确配对")
    classifications = _classify_held_targets(
        targets, entries, backups, events, allow_current_intent_prefix=True,
    )
    if "unknown" in classifications:
        _append_journal(
            transaction, key, transaction_id, "RECOVERY_REQUIRED", "UNKNOWN_TARGET_DRIFT",
            server_root=server, generation=generation,
        )
        raise NativeP2TestInstallError("P2事务存在既非安装前也非安装后的未知文件漂移，零覆盖停止")
    _transaction_fault("after_target_leaves_validated_for_rollback")

    def restore_index(index: int, *, reauthorizing_prefix: bool) -> None:
        entry = entries[index]
        path = server / Path(TARGET_RELATIVES[index])
        leaf = targets[index]
        current_payload = _refresh_target_leaf(leaf)
        current = _sha256_bytes(current_payload) if current_payload is not None else None
        current_length = len(current_payload) if current_payload is not None else 0
        if current == entry.get("before_sha256"):
            classifications[index] = "before"
            return
        if classifications[index] not in {"after", "intent-prefix"}:
            _append_journal(
                transaction, key, transaction_id, "RECOVERY_REQUIRED", "ROLLBACK_DRIFT", index, current, None,
                current_length,
                server_root=server, generation=generation,
            )
            raise NativeP2TestInstallError(f"P2回滚前目标发生未知漂移：{path}")
        _append_journal(
            transaction, key, transaction_id, "ROLLING_BACK", "ROLLBACK_INTENT", index, current,
            entry.get("before_sha256"), current_length,
            server_root=server, generation=generation,
        )
        if reauthorizing_prefix:
            _transaction_fault(f"rollback_prefix_reauthorized_before_write:{index}")
        before_payload = backups[index].payload
        if before_payload is not None:
            if leaf.handle is None:
                raise NativeP2TestInstallError(f"P2回滚目标缺少排他叶子句柄：{path}")
            _write_file_handle_exact(
                leaf, before_payload, f"rollback_after_target_handle_truncate:{index}",
            )
        else:
            if leaf.handle is None:
                raise NativeP2TestInstallError(f"P2回滚删除目标缺少排他叶子句柄：{path}")
            _delete_file_handle_exact(leaf, f"before_delete_target_handle:{index}")
        verified_payload = _refresh_target_leaf(leaf)
        verified_hash = _sha256_bytes(verified_payload) if verified_payload is not None else None
        if verified_hash != entry.get("before_sha256"):
            _append_journal(
                transaction, key, transaction_id, "RECOVERY_REQUIRED", "ROLLBACK_VERIFY_FAILED", index,
                server_root=server, generation=generation,
            )
            raise NativeP2TestInstallError(f"P2回滚后哈希不匹配：{path}")
        _append_journal(
            transaction, key, transaction_id, "ROLLING_BACK", "ROLLBACK_DONE", index,
            current, entry.get("before_sha256"), current_length,
            server_root=server, generation=generation,
        )
        classifications[index] = "before"

    # A current exact-prefix is authorized by only the final per-file intent.  Re-seal that
    # same generation/index/intended hash and finish the file before a generic state event
    # can replace the journal tail anchor.
    for index in reversed(range(len(entries))):
        if classifications[index] == "intent-prefix":
            restore_index(index, reauthorizing_prefix=True)

    _append_journal(
        transaction, key, transaction_id, "ROLLING_BACK", "ROLLBACK_START",
        server_root=server, generation=generation,
    )
    for index in reversed(range(len(entries))):
        if classifications[index] != "before":
            restore_index(index, reauthorizing_prefix=False)
    _set_coordination_owner_state(server, transaction_id, generation, key, "pending_closed")
    _append_journal(
        transaction, key, transaction_id, "ROLLED_BACK", "ROLLBACK_COMPLETE",
        server_root=server, generation=generation,
    )


def rollback_native_p2_test_install(platform_root: Path, transaction_id: str, server_root: Path) -> NativeP2RollbackReceipt:
    """只信任HMAC manifest/journal，逐字节回滚已完整安装的P2事务。"""
    with _transaction_lock(platform_root, server_root) as base:
        target_parents = tuple(dict.fromkeys(Path(relative).parent for relative in TARGET_RELATIVES))
        control_dirs = (Path(transaction_id), Path(transaction_id) / "files", Path(transaction_id) / "journal")
        with (
            _hold_directory_tree(Path(server_root).absolute(), target_parents, create=False),
            _hold_directory_tree(base, control_dirs, create=False),
            _hold_secret_tree(),
        ):
            transaction, manifest, key, events = _load_sealed_manifest(platform_root, transaction_id, server_root)
            if not events or events[-1]["state"] not in {"INSTALLED", "ROLLING_BACK"}:
                raise NativeP2TestInstallError("P2事务不是已安装或可重入回滚状态；请先执行事务恢复预检")
            server = Path(server_root).resolve()
            generation = str(manifest["generation"])
            coordination = _require_coordination(server, transaction_id, generation, key)
            if coordination.get("owner_state") != "active":
                raise NativeP2TestInstallError("P2受保护协调状态不是活动安装代；请先执行事务恢复")
            _require_active_owner(base, server, transaction_id, generation, key)
            entries = list(manifest["files"])
            with (
                _hold_verified_manifest_backups(transaction, entries) as backups,
                _hold_target_leaves(server) as targets,
            ):
                _rollback_from_manifest(
                    transaction, manifest, key, transaction_id, server, targets, backups, events,
                )
                _transaction_fault("after_rolled_back_journal")
                _write_active_owner(base, server, transaction_id, generation, key, "closed")
                _set_coordination_owner_state(server, transaction_id, generation, key, "closed")
                receipt_path = transaction / "receipt.json"
                receipt = json.loads(receipt_path.read_text(encoding="utf-8")) if receipt_path.is_file() else {}
                receipt.update({"status": "ROLLED_BACK", "updated_at_utc": datetime.now(timezone.utc).isoformat()})
                try:
                    _atomic_write(receipt_path, (json.dumps(receipt, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
                except Exception:
                    pass
                rollback_path = transaction / "rollback-receipt.json"
                _atomic_write(rollback_path, (json.dumps({
                    "schema_version": 2, "operation": "mingge-native-p2-test-rollback", "status": "rolled-back",
                    "transaction_id": transaction_id, "server_root": str(server),
                    "runtime_item_instance_restored": False,
                }, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
                return NativeP2RollbackReceipt(transaction_id, "rolled-back", rollback_path)


def recover_native_p2_test_transaction(
    platform_root: Path,
    transaction_id: str,
    server_root: Path,
    strategy: str,
) -> NativeP2RollbackReceipt:
    """显式恢复中断事务；只允许回滚到before或在全部after时补记INSTALLED。"""
    if strategy not in {"rollback", "finalize-install"}:
        raise NativeP2TestInstallError("P2恢复策略只能是rollback或finalize-install")
    with _transaction_lock(platform_root, server_root) as base:
        target_parents = tuple(dict.fromkeys(Path(relative).parent for relative in TARGET_RELATIVES))
        control_dirs = (Path(transaction_id), Path(transaction_id) / "files", Path(transaction_id) / "journal")
        with (
            _hold_directory_tree(Path(server_root).absolute(), target_parents, create=False),
            _hold_directory_tree(base, control_dirs, create=False),
            _hold_secret_tree(),
        ):
            transaction = base / transaction_id
            if transaction.name != transaction_id or not transaction.resolve(strict=False).is_relative_to(base.resolve(strict=False)):
                raise NativeP2TestInstallError("P2事务ID非法或目录越界")
            _, preparation_key = _load_or_create_integrity_key_held(base)
            preparation_coordination = _read_coordination(Path(server_root).resolve(), preparation_key, required=False)
            preparation_states = {
                "RESERVED", "MANIFEST_JSON_WRITTEN", "MANIFEST_HMAC_WRITTEN",
                "MANIFEST_VERIFIED", "ABORTED",
            }
            if (
                preparation_coordination
                and preparation_coordination.get("transaction_id") == transaction_id
                and preparation_coordination.get("business_write_started") is False
                and preparation_coordination.get("setup_state") in preparation_states
                and preparation_coordination.get("journal_sequence") == 0
                and preparation_coordination.get("journal_head_hmac") == ""
                and preparation_coordination.get("pending_event") is None
            ):
                if strategy != "rollback":
                    raise NativeP2TestInstallError("P2零写入准备事务只能显式中止，不能补记安装完成")
                server = Path(server_root).resolve()
                with _hold_target_leaves(server) as preparation_targets:
                    return _abort_preparation_reservation(
                        base, transaction, server, transaction_id, preparation_key,
                        preparation_coordination, preparation_targets,
                    )

            transaction, manifest, key, events = _load_sealed_manifest(
                platform_root, transaction_id, server_root, reconcile_pending=True,
            )
            server = Path(server_root).resolve()
            generation = str(manifest["generation"])
            coordination = _require_coordination(server, transaction_id, generation, key)
            entries = list(manifest["files"])
            with (
                _hold_verified_manifest_backups(transaction, entries) as backups,
                _hold_target_leaves(server) as targets,
            ):
                classifications = _classify_held_targets(
                    targets, entries, backups, events, allow_current_intent_prefix=True,
                )
                if "unknown" in classifications:
                    if not events or events[-1]["state"] != "RECOVERY_REQUIRED":
                        _append_journal(
                            transaction, key, transaction_id, "RECOVERY_REQUIRED", "UNKNOWN_TARGET_DRIFT",
                            server_root=server, generation=generation,
                        )
                    raise NativeP2TestInstallError("P2事务存在未知文件漂移，恢复保持零覆盖")
                last_state = str(events[-1]["state"]) if events else "MANIFEST_ONLY"
                active_owner = _read_active_owner(base, server, key)

                if coordination.get("owner_state") == "pending_closed":
                    if strategy != "rollback":
                        raise NativeP2TestInstallError("待关闭回滚只能执行terminal-only恢复")
                    _require_pending_closed_terminal(
                        coordination, active_owner, events, entries, classifications,
                        transaction_id, generation,
                    )
                    if last_state != "ROLLED_BACK":
                        _append_journal(
                            transaction, key, transaction_id, "ROLLED_BACK", "ROLLBACK_COMPLETE",
                            server_root=server, generation=generation,
                        )
                    _write_active_owner(base, server, transaction_id, generation, key, "closed")
                    _set_coordination_owner_state(server, transaction_id, generation, key, "closed")
                    rollback_path = transaction / "rollback-receipt.json"
                    _atomic_write(rollback_path, (json.dumps({
                        "schema_version": 3, "operation": "mingge-native-p2-test-recovery",
                        "status": "rolled-back", "transaction_id": transaction_id,
                        "runtime_item_instance_restored": False,
                    }, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
                    return NativeP2RollbackReceipt(transaction_id, "rolled-back", rollback_path)

                if strategy == "finalize-install":
                    if last_state == "ROLLED_BACK":
                        raise NativeP2TestInstallError("P2事务已回滚，不允许补记INSTALLED")
                    if any(value != "after" for value in classifications):
                        raise NativeP2TestInstallError("只有三目标全部为安装后哈希时才能补记或协调INSTALLED")
                    if last_state == "INSTALLED":
                        if coordination.get("owner_state") == "active":
                            _require_active_owner(base, server, transaction_id, generation, key)
                        elif coordination.get("owner_state") == "pending_active":
                            current_active = bool(
                                active_owner
                                and active_owner.get("transaction_id") == transaction_id
                                and active_owner.get("generation") == generation
                                and active_owner.get("state") == "active"
                            )
                            previous_closed = bool(active_owner and active_owner.get("state") == "closed")
                            if active_owner is not None and not current_active and not previous_closed:
                                raise NativeP2TestInstallError("INSTALLED协调恢复发现其他活动归属，失败关闭")
                            if not current_active:
                                _write_active_owner(base, server, transaction_id, generation, key, "active")
                            _set_coordination_owner_state(server, transaction_id, generation, key, "active")
                        else:
                            raise NativeP2TestInstallError("INSTALLED状态没有可恢复的受保护归属转换")
                        return NativeP2RollbackReceipt(transaction_id, "installed-finalized", transaction / "manifest.json")
                    if last_state not in {"MANIFEST_ONLY", "PREPARED", "INSTALLING", "RECOVERY_REQUIRED"}:
                        raise NativeP2TestInstallError(f"P2事务状态{last_state}不允许补记INSTALLED")
                    if coordination.get("owner_state") not in {"none", "pending_active"}:
                        raise NativeP2TestInstallError("P2受保护协调归属状态不允许补记INSTALLED")
                    _set_coordination_owner_state(server, transaction_id, generation, key, "pending_active")
                    _append_journal(
                        transaction, key, transaction_id, "INSTALLED", "RECOVERY_FINALIZE_INSTALL",
                        server_root=server, generation=generation,
                    )
                    _write_active_owner(base, server, transaction_id, generation, key, "active")
                    _set_coordination_owner_state(server, transaction_id, generation, key, "active")
                    return NativeP2RollbackReceipt(transaction_id, "installed-finalized", transaction / "manifest.json")

                if last_state == "RECOVERY_REQUIRED":
                    _require_recovery_required_terminal(
                        coordination, active_owner, events, entries, classifications,
                        transaction_id, generation,
                    )
                    _set_coordination_owner_state(
                        server, transaction_id, generation, key, "pending_closed",
                    )
                    _append_journal(
                        transaction, key, transaction_id, "ROLLED_BACK", "ROLLBACK_COMPLETE",
                        server_root=server, generation=generation,
                    )
                    _write_active_owner(base, server, transaction_id, generation, key, "closed")
                    _set_coordination_owner_state(server, transaction_id, generation, key, "closed")
                    rollback_path = transaction / "rollback-receipt.json"
                    _atomic_write(rollback_path, (json.dumps({
                        "schema_version": 3, "operation": "mingge-native-p2-test-recovery",
                        "status": "rolled-back", "transaction_id": transaction_id,
                        "runtime_item_instance_restored": False,
                    }, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
                    return NativeP2RollbackReceipt(transaction_id, "rolled-back", rollback_path)

                if last_state == "ROLLED_BACK":
                    if any(value != "before" for value in classifications):
                        raise NativeP2TestInstallError("ROLLED_BACK协调恢复要求三个目标全部精确为安装前内容")
                    if coordination.get("owner_state") == "closed":
                        owner = _read_active_owner(base, server, key)
                        if not owner or owner.get("state") != "closed" or owner.get("transaction_id") != transaction_id or owner.get("generation") != generation:
                            raise NativeP2TestInstallError("已关闭P2协调状态的普通归属缺失或不匹配")
                    else:
                        raise NativeP2TestInstallError("ROLLED_BACK状态没有可恢复的受保护关闭转换")
                else:
                    resume_mode = "business"
                    if last_state == "INSTALLED":
                        if coordination.get("owner_state") != "active":
                            raise NativeP2TestInstallError("活动回滚缺少受保护协调归属")
                        _require_active_owner(base, server, transaction_id, generation, key)
                    elif last_state == "ROLLING_BACK":
                        if coordination.get("owner_state") == "active":
                            _require_active_owner(base, server, transaction_id, generation, key)
                        else:
                            resume_mode = _nonactive_rollback_resume_mode(
                                coordination, active_owner, events, entries, classifications,
                                transaction_id, generation,
                            )
                    elif active_owner and active_owner.get("state") == "active":
                        _require_active_owner(base, server, transaction_id, generation, key)
                    if resume_mode == "terminal":
                        _append_journal(
                            transaction, key, transaction_id, "ROLLED_BACK", "ROLLBACK_COMPLETE",
                            server_root=server, generation=generation,
                        )
                    else:
                        _rollback_from_manifest(
                            transaction, manifest, key, transaction_id, server,
                            targets, backups, events,
                        )
                    _write_active_owner(base, server, transaction_id, generation, key, "closed")
                    _set_coordination_owner_state(server, transaction_id, generation, key, "closed")

                rollback_path = transaction / "rollback-receipt.json"
                _atomic_write(rollback_path, (json.dumps({
                    "schema_version": 3, "operation": "mingge-native-p2-test-recovery",
                    "status": "rolled-back", "transaction_id": transaction_id,
                    "runtime_item_instance_restored": False,
                }, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
                return NativeP2RollbackReceipt(transaction_id, "rolled-back", rollback_path)
