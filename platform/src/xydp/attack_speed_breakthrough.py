from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .mingge_dual import (
    DualInstallPlan,
    MinggeDualError,
    _DualService,
    _change,
    _find_block,
    _managed_block,
    _plan_id,
    _script_bytes,
    _script_text,
    _sha256,
    _target,
    _upsert_after_label,
)


PACKAGE_ID = "xy.optional.attack-speed-breakthrough"
ROUTE = "attack-speed-breakthrough"
CORE_RELATIVE = Path("Mir200/Envir/QuestDiary/玄渊攻速突破/全身攻速阈值核心.txt")
QFUNCTION_RELATIVE = Path("Mir200/Envir/Market_Def/QFunction-0.txt")
QMANAGE_RELATIVE = Path("Mir200/Envir/MapQuest_Def/QManage.txt")
TEXTVAR_RELATIVE = Path("Mir200/Envir/CustomItemPropertyTextVarList.txt")
TEXTVAR_LINE = 40
TEXTVAR_VALUE = "{攻速突破∶|251}+$$2"
CORE_LABEL = "XY_AS_CAP_RECALC"
RECALC_LABEL = "XY_AS_FORMAL_RECALC_DELAY"
LOGIN_BRIDGE_LABEL = "XY_AS_LOGIN_RECALC_BRIDGE"
LOGIN_HOOK_MARKER = "XY-AS-LOGIN-BRIDGE-MAIN1"
LOGIN_LABEL_MARKER = "XY-AS-LOGIN-BRIDGE-LABEL"
PACKAGE_MARKER = "XY-AS-CAP-V1"
EQUIPMENT_ANCHOR = "; XY_EQUIP_MAKER_ATTACK_SPEED_BREAK_ANCHOR"
WASH_SPEED_MARKER = "XY-EQUIPMENT-WASH-IMPORT-V1-NORMAL-SPEED"
LEGACY_CORE_SHA256 = "DEAC16FFFC84C75398F1897F0043138E88F1E2FDFD76F03A7AAB2C6D55FA1216"
SHEETS = ("基础设置", "填写说明")
VARIABLES = (
    "N$XY_AS_BUSY",
    "N$XY_AS_RAW",
    "N$XY_AS_TMP",
    "N$XY_AS_BREAK_RAW",
    "N$XY_AS_BREAK",
    "N$XY_AS_FIXED_BREAK",
    "N$XY_AS_CAP",
    "N$XY_AS_EFFECTIVE",
)


class AttackSpeedBreakthroughError(MinggeDualError):
    pass


@dataclass(frozen=True)
class AttackSpeedWorkbook:
    path: Path
    schema_version: int
    system_id: str
    base_cap: int
    item_maximum: int
    total_maximum: int
    absolute_maximum: int
    textvar_line: int
    login_delay_ms: int
    equip_delay_ms: int


def _int(value, name: str) -> int:
    if value is None or isinstance(value, bool):
        raise AttackSpeedBreakthroughError(f"基础设置缺少整数：{name}")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise AttackSpeedBreakthroughError(f"基础设置{name}必须是整数") from exc
    return number


def load_attack_speed_workbook(path: Path) -> AttackSpeedWorkbook:
    source = Path(path).resolve()
    if not source.is_file():
        raise AttackSpeedBreakthroughError(f"攻速突破配置表不存在：{source}")
    workbook = load_workbook(source, data_only=False, read_only=True)
    try:
        missing = [name for name in SHEETS if name not in workbook.sheetnames]
        if missing:
            raise AttackSpeedBreakthroughError("攻速突破配置表缺少工作表：" + "、".join(missing))
        sheet = workbook["基础设置"]
        values: dict[str, object] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            key = str(row[0] or "").strip()
            if key:
                if key in values:
                    raise AttackSpeedBreakthroughError(f"基础设置配置项重复：{key}")
                values[key] = row[1]
    finally:
        workbook.close()
    book = AttackSpeedWorkbook(
        source,
        _int(values.get("配置版本"), "配置版本"),
        str(values.get("系统编号") or "").strip(),
        _int(values.get("基础攻速阈值"), "基础攻速阈值"),
        _int(values.get("单件突破上限"), "单件突破上限"),
        _int(values.get("全身突破上限"), "全身突破上限"),
        _int(values.get("绝对攻速上限"), "绝对攻速上限"),
        _int(values.get("TextVar行"), "TextVar行"),
        _int(values.get("登录延迟毫秒"), "登录延迟毫秒"),
        _int(values.get("穿脱延迟毫秒"), "穿脱延迟毫秒"),
    )
    expected = {
        "配置版本": (book.schema_version, 1),
        "系统编号": (book.system_id, "attack_speed_breakthrough"),
        "基础攻速阈值": (book.base_cap, 20),
        "单件突破上限": (book.item_maximum, 30),
        "全身突破上限": (book.total_maximum, 30),
        "绝对攻速上限": (book.absolute_maximum, 50),
        "TextVar行": (book.textvar_line, 40),
        "登录延迟毫秒": (book.login_delay_ms, 1000),
        "穿脱延迟毫秒": (book.equip_delay_ms, 100),
    }
    mismatches = [f"{key}必须为{want}，当前为{actual}" for key, (actual, want) in expected.items() if actual != want]
    if mismatches:
        raise AttackSpeedBreakthroughError("；".join(mismatches))
    return book


def _render_core(book: AttackSpeedWorkbook) -> str:
    return f"""{{
[@{CORE_LABEL}]
#IF
EQUAL N$XY_AS_BUSY 1
#ACT
BREAK

#IF
#ACT
MOV N$XY_AS_BUSY 1
ChangeSpeed 2 0
MOV N$XY_AS_RAW <$HITSPD>
MOV N$XY_AS_TMP 0
MOV N$XY_AS_BREAK_RAW 0
GetAllCustomItemValueByTextLine 60 -1 {book.textvar_line} N$XY_AS_TMP N$XY_AS_BREAK_RAW N$XY_AS_TMP
INC N$XY_AS_BREAK_RAW <$STR(N$XY_AS_FIXED_BREAK)>

#IF
SMALL N$XY_AS_RAW 0
#ACT
MOV N$XY_AS_RAW 0

#IF
#ACT
MOV N$XY_AS_BREAK <$STR(N$XY_AS_BREAK_RAW)>

#IF
SMALL N$XY_AS_BREAK 0
#ACT
MOV N$XY_AS_BREAK 0

#IF
LARGE N$XY_AS_BREAK {book.total_maximum}
#ACT
MOV N$XY_AS_BREAK {book.total_maximum}

#IF
#ACT
MOV N$XY_AS_CAP {book.base_cap}
INC N$XY_AS_CAP <$STR(N$XY_AS_BREAK)>

#IF
LARGE N$XY_AS_CAP {book.absolute_maximum}
#ACT
MOV N$XY_AS_CAP {book.absolute_maximum}

#IF
#ACT
MOV N$XY_AS_EFFECTIVE <$STR(N$XY_AS_RAW)>

#IF
LARGE N$XY_AS_EFFECTIVE <$STR(N$XY_AS_CAP)>
#ACT
MOV N$XY_AS_EFFECTIVE <$STR(N$XY_AS_CAP)>

#IF
LARGE N$XY_AS_EFFECTIVE {book.absolute_maximum}
#ACT
MOV N$XY_AS_EFFECTIVE {book.absolute_maximum}

#IF
SMALL N$XY_AS_EFFECTIVE 0
#ACT
MOV N$XY_AS_EFFECTIVE 0

#IF
#ACT
DEC N$XY_AS_EFFECTIVE <$STR(N$XY_AS_RAW)>
ChangeSpeed 2 <$STR(N$XY_AS_EFFECTIVE)>
MOV N$XY_AS_BUSY 0
BREAK
}}"""


def _wash_speed_block() -> str:
    return _managed_block(WASH_SPEED_MARKER, (
        "MOV N$XY_WASH_AS_TMP 0",
        "MOV N$XY_WASH_AS_NORMAL 0",
        "GetAllCustomItemValueByTextLine 60 -1 46 N$XY_WASH_AS_TMP N$XY_WASH_AS_NORMAL N$XY_WASH_AS_TMP",
        "INC N$XY_AS_RAW <$STR(N$XY_WASH_AS_NORMAL)>",
    ))


def _is_compatible_core(text: str, book: AttackSpeedWorkbook) -> bool:
    current = text.rstrip("\n")
    baseline = _render_core(book).rstrip("\n")
    if current == baseline:
        return True
    block = _wash_speed_block()
    if current.count(block) != 1:
        return False
    without_wash = current.replace(block + "\n", "", 1)
    return without_wash == baseline


def _event_block(event: str, delay: int) -> tuple[str, str]:
    marker = f"XY-AS-CAP-HOOK-{event}"
    return marker, _managed_block(marker, ("#IF", "#ACT", f"DELAYGOTO {delay} @{RECALC_LABEL}"))


def _legacy_event_block(event: str, delay: int) -> str:
    return "\n".join((
        f"; XY-AS-FORMAL-CAP-BEGIN {event}",
        "#IF",
        "#ACT",
        f"DELAYGOTO {delay} @{RECALC_LABEL}",
        f"; XY-AS-FORMAL-CAP-END {event}",
    ))


def _block_text(text: str, marker: str) -> str | None:
    found = _find_block(text, marker)
    return None if not found else text[found[0]:found[1]]


def _ensure_event_hook(text: str, event: str, delay: int) -> tuple[str, str | None, str | None]:
    legacy = _legacy_event_block(event, delay)
    legacy_begin = f"; XY-AS-FORMAL-CAP-BEGIN {event}"
    legacy_end = f"; XY-AS-FORMAL-CAP-END {event}"
    if legacy_begin in text or legacy_end in text:
        if text.count(legacy_begin) != 1 or text.count(legacy_end) != 1:
            raise AttackSpeedBreakthroughError(f"{event}旧攻速挂钩重复、残缺或内容漂移")
        if legacy in text:
            return text, None, None
        raise AttackSpeedBreakthroughError(f"{event}旧攻速挂钩重复、残缺或内容漂移")
    marker, expected = _event_block(event, delay)
    actual = _block_text(text, marker)
    if actual is not None:
        if actual != expected:
            raise AttackSpeedBreakthroughError(f"{event}攻速挂钩受管块已被手工修改")
        return text, None, None
    updated, block = _upsert_after_label(text, event, marker, ("#IF", "#ACT", f"DELAYGOTO {delay} @{RECALC_LABEL}"))
    return updated, marker, block


def _tail_blocks() -> tuple[str, str]:
    anchor_only = _managed_block(PACKAGE_MARKER, (
        "#IF",
        "#ACT",
        "MOV N$XY_AS_FIXED_BREAK 0",
        EQUIPMENT_ANCHOR,
    ))
    full = _managed_block(PACKAGE_MARKER, (
        f"[@{RECALC_LABEL}]",
        "#IF",
        "#ACT",
        "MOV N$XY_AS_FIXED_BREAK 0",
        EQUIPMENT_ANCHOR,
        f"#CALL [\\玄渊攻速突破\\全身攻速阈值核心.txt] @{CORE_LABEL}",
        "BREAK",
    ))
    return anchor_only, full


def _valid_fixed_equipment_extensions(text: str) -> bool:
    if not text.strip():
        return True
    chunks = [chunk for chunk in text.split("; XY-EQUIP-MAKER ") if chunk.strip()]
    for chunk in chunks:
        lines = [line.strip() for line in chunk.splitlines() if line.strip()]
        if len(lines) != 5:
            return False
        match = re.fullmatch(r"(.+): 攻速突破\+(\d+)", lines[0])
        if not match:
            return False
        name, value = match.groups()
        if lines[1:] != [
            "#IF",
            f"CHECKITEMW {name} 1",
            "#ACT",
            f"INC N$XY_AS_FIXED_BREAK {value}",
        ]:
            return False
    return True


def _ensure_recalc_tail(text: str) -> tuple[str, str | None, str | None]:
    anchor_only, full = _tail_blocks()
    actual = _block_text(text, PACKAGE_MARKER)
    if actual is not None:
        if actual not in {anchor_only, full}:
            prefixes = (
                anchor_only[:anchor_only.index(EQUIPMENT_ANCHOR) + len(EQUIPMENT_ANCHOR)],
                full[:full.index(EQUIPMENT_ANCHOR) + len(EQUIPMENT_ANCHOR)],
            )
            prefix = next((value for value in prefixes if actual.startswith(value)), "")
            end_marker = f"; {PACKAGE_MARKER}-END"
            full_suffix = "\n".join((
                f"#CALL [\\玄渊攻速突破\\全身攻速阈值核心.txt] @{CORE_LABEL}",
                "BREAK",
                end_marker,
            ))
            suffix = full_suffix if actual.endswith(full_suffix) else end_marker
            if not prefix or not actual.endswith(suffix):
                raise AttackSpeedBreakthroughError("攻速突破统一重算受管块已被手工修改")
            extensions = actual[len(prefix):-len(suffix)]
            if not _valid_fixed_equipment_extensions(extensions):
                raise AttackSpeedBreakthroughError("攻速突破固定装备扩展块格式异常")
        return text, None, None
    label = f"[@{RECALC_LABEL}]"
    matches = list(re.finditer(rf"(?im)^\s*{re.escape(label)}\s*$", text))
    if len(matches) > 1:
        raise AttackSpeedBreakthroughError(f"统一重算标签重复：{label}")
    if matches:
        start = matches[0].end()
        next_label = re.search(r"(?im)^\s*\[@", text[start:])
        end = len(text) if next_label is None else start + next_label.start()
        body = text[start:end].strip()
        expected = "\n".join(("#IF", "#ACT", f"#CALL [\\玄渊攻速突破\\全身攻速阈值核心.txt] @{CORE_LABEL}", "BREAK"))
        if body != expected:
            raise AttackSpeedBreakthroughError("现有统一重算标签正文不是已验收金样本，禁止猜测接管")
        updated, block = _upsert_after_label(text, RECALC_LABEL, PACKAGE_MARKER, (
            "#IF", "#ACT", "MOV N$XY_AS_FIXED_BREAK 0", EQUIPMENT_ANCHOR,
        ))
        return updated, PACKAGE_MARKER, block
    suffix = "" if not text or text.endswith("\n") else "\n"
    return text + suffix + full + "\n", PACKAGE_MARKER, full


def _remove_retired_login_bridge(text: str, delay: int) -> str:
    hook_block = _managed_block(LOGIN_HOOK_MARKER, (
        "#IF",
        "#ACT",
        f"DELAYGOTO {delay} @{LOGIN_BRIDGE_LABEL}",
    ))
    label_block = _managed_block(LOGIN_LABEL_MARKER, (
        f"[@{LOGIN_BRIDGE_LABEL}]",
        "#IF",
        "#ACT",
        f"GOTOLABEL 8 @{RECALC_LABEL} <$X> <$Y> 0 0",
        "BREAK",
    ))
    actual_hook = _block_text(text, LOGIN_HOOK_MARKER)
    actual_label = _block_text(text, LOGIN_LABEL_MARKER)
    if actual_hook is None and actual_label is None:
        return text
    if actual_hook != hook_block or actual_label != label_block:
        raise AttackSpeedBreakthroughError("退役QManage攻速登录桥重复、残缺或内容漂移")
    for block in (hook_block, label_block):
        text = text.replace(block + "\n", "", 1) if block + "\n" in text else text.replace(block, "", 1)
    return text


def _active_change_speed_files(root: Path, core_path: Path) -> list[Path]:
    conflicts: list[Path] = []
    envir = root / "Mir200/Envir"
    for path in envir.rglob("*.txt"):
        if path.resolve() == core_path.resolve() or not path.is_file():
            continue
        text = _scan_business_script(root, path)
        if re.search(r"(?im)^\s*ChangeSpeed\s+2(?:\s|$)", text):
            conflicts.append(path)
    return conflicts


def _assert_namespace(root: Path, core_path: Path, qfunction_path: Path) -> None:
    allowed = {core_path.resolve(), qfunction_path.resolve()}
    envir = root / "Mir200/Envir"
    for path in envir.rglob("*.txt"):
        if path.resolve() in allowed or not path.is_file():
            continue
        text = _scan_business_script(root, path)
        used = [name for name in VARIABLES if re.search(rf"(?<![A-Za-z0-9_]){re.escape(name)}(?![A-Za-z0-9_])", text)]
        if used:
            raise AttackSpeedBreakthroughError(f"变量命名空间冲突：{path.relative_to(root)} 使用 {'、'.join(used)}")


def _scan_business_script(root: Path, path: Path) -> str:
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise AttackSpeedBreakthroughError(f"无法读取业务脚本：{path.relative_to(root)}") from exc
    if data.startswith(b"\xef\xbb\xbf"):
        encodings = ("utf-8-sig", "gb18030", "utf-16-le")
    elif data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ("utf-16", "gb18030", "utf-8")
    else:
        encodings = ("gb18030", "utf-8-sig", "utf-8", "utf-16-le")
    for encoding in encodings:
        try:
            text = data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
        if "\x00" in text:
            continue
        return text.replace("\r\n", "\n").replace("\r", "\n")
    raise AttackSpeedBreakthroughError(
        f"业务脚本编码不可判定：{path.relative_to(root)}（已严格尝试GB18030、UTF-8和UTF-16LE）"
    )


class AttackSpeedBreakthroughService(_DualService):
    route = ROUTE
    package_id = PACKAGE_ID

    def preflight(self, server: Path, workbook: Path) -> DualInstallPlan:
        blockers: list[str] = []
        warnings = ["当前服已游戏验收；本平台包在其他翎风/LFM2目标服仍需逐服验收。"]
        changes = []
        source = Path(workbook).resolve()
        workbook_hash = _sha256(source.read_bytes()) if source.is_file() else ""
        try:
            root = Path(server).resolve()
            if not (root / "Mir200/Envir").is_dir():
                raise AttackSpeedBreakthroughError("目标目录不是可识别的翎风/LFM2服务端")
            book = load_attack_speed_workbook(source)
            core_path = _target(root, CORE_RELATIVE)
            qfunction_path = _target(root, QFUNCTION_RELATIVE)
            qmanage_path = _target(root, QMANAGE_RELATIVE)
            textvar_path = _target(root, TEXTVAR_RELATIVE)
            if not qfunction_path.is_file():
                raise AttackSpeedBreakthroughError("QFunction-0.txt不存在")
            if not qmanage_path.is_file():
                raise AttackSpeedBreakthroughError("QManage.txt不存在")
            if not textvar_path.is_file():
                raise AttackSpeedBreakthroughError("CustomItemPropertyTextVarList.txt不存在")
            conflicts = _active_change_speed_files(root, core_path)
            if conflicts:
                names = "、".join(str(path.relative_to(root)) for path in conflicts[:8])
                raise AttackSpeedBreakthroughError("发现本包核心之外的活动ChangeSpeed 2出口：" + names)
            _assert_namespace(root, core_path, qfunction_path)

            core_before = core_path.read_bytes() if core_path.is_file() else None
            core_after = _script_bytes(_render_core(book))
            if core_before is not None and core_before != core_after:
                current_text = _script_text(core_before)
                if _is_compatible_core(current_text, book):
                    core_after = core_before
                else:
                    current_hash = hashlib.sha256(core_before).hexdigest().upper()
                    if current_hash != LEGACY_CORE_SHA256:
                        raise AttackSpeedBreakthroughError("现有攻速突破核心不是已验收金样本或本包版本，禁止覆盖")
            if core_before != core_after:
                changes.append(_change(root, CORE_RELATIVE, "exclusive", core_before, core_after))

            textvar_before = textvar_path.read_bytes()
            lines = _script_text(textvar_before).splitlines()
            existing = lines[TEXTVAR_LINE - 1] if len(lines) >= TEXTVAR_LINE else ""
            if existing.strip() and existing != TEXTVAR_VALUE:
                raise AttackSpeedBreakthroughError(f"TextVar第{TEXTVAR_LINE}行已被其他系统占用，禁止覆盖")
            if existing != TEXTVAR_VALUE:
                while len(lines) < TEXTVAR_LINE - 1:
                    lines.append("")
                before_lines = [existing] if len(lines) >= TEXTVAR_LINE else []
                if len(lines) == TEXTVAR_LINE - 1:
                    lines.append(TEXTVAR_VALUE)
                else:
                    lines[TEXTVAR_LINE - 1] = TEXTVAR_VALUE
                textvar_after = _script_bytes("\n".join(lines))
                changes.append(_change(root, TEXTVAR_RELATIVE, "indexed_lines", textvar_before, textvar_after, {
                    "start": TEXTVAR_LINE, "before_lines": before_lines, "after_lines": [TEXTVAR_VALUE],
                }))

            q_before = qfunction_path.read_bytes()
            q_text = _script_text(q_before)
            markers: list[str] = []
            blocks: list[str] = []
            for event, delay in (("PlayLogin", book.login_delay_ms), ("TakeOnEx", book.equip_delay_ms), ("TakeOffEx", book.equip_delay_ms)):
                q_text, marker, block = _ensure_event_hook(q_text, event, delay)
                if marker and block:
                    markers.append(marker)
                    blocks.append(block)
            q_text, marker, block = _ensure_recalc_tail(q_text)
            if marker and block:
                markers.append(marker)
                blocks.append(block)
            q_after = _script_bytes(q_text)
            if q_before != q_after:
                changes.append(_change(root, QFUNCTION_RELATIVE, "managed_blocks", q_before, q_after, {
                    "markers": markers, "blocks": blocks,
                }))

            qmanage_before = qmanage_path.read_bytes()
            qmanage_text = _script_text(qmanage_before)
            qmanage_text = _remove_retired_login_bridge(qmanage_text, book.login_delay_ms)
            qmanage_after = _script_bytes(qmanage_text)
            if qmanage_before != qmanage_after:
                changes.append(_change(root, QMANAGE_RELATIVE, "managed_blocks", qmanage_before, qmanage_after, {
                    "markers": [LOGIN_HOOK_MARKER, LOGIN_LABEL_MARKER], "blocks": [],
                }))
        except (OSError, UnicodeError, ValueError, MinggeDualError) as exc:
            blockers.append(str(exc))
            changes = []
        return DualInstallPlan(
            ROUTE, PACKAGE_ID, _plan_id(ROUTE, workbook_hash, changes), Path(server).resolve(), source,
            workbook_hash, tuple(changes), tuple(blockers), tuple(warnings),
        )


__all__ = [
    "AttackSpeedBreakthroughError", "AttackSpeedBreakthroughService", "AttackSpeedWorkbook",
    "CORE_RELATIVE", "PACKAGE_ID", "QFUNCTION_RELATIVE", "QMANAGE_RELATIVE", "ROUTE", "TEXTVAR_RELATIVE",
    "load_attack_speed_workbook",
]
