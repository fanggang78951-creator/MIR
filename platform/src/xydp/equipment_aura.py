"""翎风/LFM2 装备范围光环候选生成与事务部署工具。

候选阶段只支持一件测试装备和一套特效。所有目标文件均先在内存中
生成，预检无阻止项后才允许事务写入；收据可逐字节回滚。
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import tempfile
import uuid
from dataclasses import asdict, dataclass, replace
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image


ENCODING = "gb18030"
NEWLINE = "\r\n"
KNOWN_PRIOR_CLIENT_RESOURCE_HASHES = {
    "xy_equipmentaura.wzl": {
        # candidate.1：切线刀轮、Y=-120。
        "0FD9782FC3978D409D82A4AE9CEAF2C7FA15A1300311580B0848277B49ABFE30",
        # candidate.3：径向刀轮、192x160、内圈4、Y=-75。
        "2C6FDB67120669BE76A8A28EFA8EF4536CCE34BD5AA4D8D9E2B76D39DCB12B38",
        # candidate.4：224x224、内圈28、坐标-72,-92。
        "82DDFD94E5B667336FCA727C878B98478940235CBF722E825EDA422AEAAD9814",
        # candidate.5：脚点锚定，坐标-90,-92。
        "7CE8786CEA6C3A058F108207E45A7E2B77D170F8031F329B995A3C082BC23698",
        # candidate.6：供体圆月斩，小尺寸8帧、坐标-282,-267。
        "EF3ED6C80BFA216E0B1EC7B520499535287EE6F148F5958A4512B018D41D354E",
    },
    "xy_equipmentaura.wzx": {
        # candidate.1～3共用的48帧索引。
        "1B3D93B8479821F0861D56EC6EBB84981C960773C3FAA0A4013A3B751493A077",
        # candidate.4：48帧、224x224、坐标-72,-92。
        "875EF2614AAEFA788EDD666A7893A96940B78F16F388A8F0E78B754F8F4A10C2",
        # candidate.6：供体圆月斩8帧索引。
        "B7C70C0A0C6A6D53BF85ED278E3798EE135D28BE5DFF29E50BC572874F7657FB",
    },
}


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def _sha256_file(path: Path) -> str:
    return _sha256_bytes(path.read_bytes())


def _is_known_prior_client_resource(name: str, digest: str) -> bool:
    return digest.upper() in KNOWN_PRIOR_CLIENT_RESOURCE_HASHES.get(name.lower(), set())


@dataclass(frozen=True)
class AuraSpec:
    item_name: str = "灵玉"
    timer_id: int = 97
    damage_multiplier: int = 2
    damage_range: int = 2
    interval_seconds: int = 1
    frame_count: int = 24
    frame_speed: int = 4
    resource_name: str = "XY_EquipmentAura.wzl"
    canvas_width: int = 224
    canvas_height: int = 224
    effect_x_offset: int = -90
    effect_y_offset: int = -92
    inner_radius: int = 28
    blade_count: int = 10
    source_rotation_degrees: float = -35.0
    effect_layer_count: int = 2
    visual_name: str = "金焰刀轮"

    @classmethod
    def round_moon(cls) -> "AuraSpec":
        """返回供体 E:\\JN28-圆月斩 的小尺寸单层候选规格。"""
        return cls(
            frame_count=8,
            frame_speed=30,
            canvas_width=559,
            canvas_height=429,
            effect_x_offset=-282,
            effect_y_offset=-267,
            effect_layer_count=1,
            visual_name="紫电圆月斩",
        )


@dataclass(frozen=True)
class AuraStyle:
    """一件装备对应的一段独立人物特效。"""

    item_name: str
    resource_index: int
    start_frame: int
    frame_count: int
    frame_speed: int
    layer: int
    style_id: int


class EquipmentAuraError(ValueError):
    """装备光环配置、目标或事务不满足安全契约。"""


@dataclass(frozen=True)
class AuraBinding:
    enabled: bool
    equipment_name: str
    style_id: str
    damage_multiplier: int
    interval_seconds: int
    priority: int
    note: str = ""


@dataclass(frozen=True)
class AuraWorkbook:
    bindings: tuple[AuraBinding, ...]

    @property
    def enabled_bindings(self) -> tuple[AuraBinding, ...]:
        return tuple(binding for binding in self.bindings if binding.enabled)


WORKBOOK_HEADERS = (
    "启用状态",
    "装备名称",
    "特效样式ID",
    "伤害倍率",
    "攻击间隔秒",
    "优先级",
    "备注",
)


@dataclass(frozen=True)
class RuntimeAuraStyleDefinition:
    style_id: str
    display_name: str
    start_frame: int
    frame_count: int = 8
    frame_speed_ms: int = 30
    damage_range: int = 3


RUNTIME_STYLE_CATALOG = tuple(
    RuntimeAuraStyleDefinition(style_id, display_name, start_frame)
    for style_id, display_name, start_frame in (
        ("style01", "青霜圆月", 0),
        ("style02", "紫金圆月", 8),
        ("style03", "赤焰圆月", 16),
        ("style04", "翡翠圆月", 24),
        ("style05", "圣辉圆月", 32),
        ("style06", "血月天轮", 40),
        ("style07", "冰魄天轮", 48),
        ("style08", "幽冥天轮", 56),
        ("style09", "雷霆天轮", 64),
        ("style10", "混沌天轮", 72),
    )
)


def _style_definition(style_id: str) -> RuntimeAuraStyleDefinition:
    for style in RUNTIME_STYLE_CATALOG:
        if style.style_id == style_id:
            return style
    raise EquipmentAuraError(f"未知特效样式: {style_id}")


def _strict_integer(value: object, field: str, minimum: int, maximum: int, row: int) -> int:
    if isinstance(value, bool):
        raise EquipmentAuraError(f"第{row}行{field}必须是{minimum}～{maximum}的整数")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if not isinstance(value, int) or not minimum <= value <= maximum:
        raise EquipmentAuraError(f"第{row}行{field}必须是{minimum}～{maximum}的整数")
    return value


def _binding_from_row(values: Sequence[object], row: int) -> AuraBinding:
    enabled_text = "" if values[0] is None else str(values[0]).strip()
    if enabled_text not in {"是", "否"}:
        raise EquipmentAuraError(f"第{row}行启用状态只允许填“是”或“否”")
    equipment_name = "" if values[1] is None else str(values[1]).strip()
    if not equipment_name:
        raise EquipmentAuraError(f"第{row}行装备名称不能为空")
    style_id = "" if values[2] is None else str(values[2]).strip()
    _style_definition(style_id)
    multiplier = _strict_integer(values[3], "伤害倍率", 1, 999, row)
    interval = _strict_integer(values[4], "攻击间隔", 1, 60, row)
    priority = _strict_integer(values[5], "优先级", 1, 999, row)
    note = "" if values[6] is None else str(values[6]).strip()
    return AuraBinding(
        enabled=enabled_text == "是",
        equipment_name=equipment_name,
        style_id=style_id,
        damage_multiplier=multiplier,
        interval_seconds=interval,
        priority=priority,
        note=note,
    )


def validate_binding_priorities(workbook: AuraWorkbook) -> None:
    enabled = workbook.enabled_bindings
    if not enabled:
        raise EquipmentAuraError("配置表没有启用任何装备光环")
    maximum = max(binding.priority for binding in enabled)
    owners = [binding.equipment_name for binding in enabled if binding.priority == maximum]
    if len(owners) > 1:
        raise EquipmentAuraError(
            f"最高优先级重复: {maximum} -> {', '.join(owners)}"
        )


def read_aura_workbook(path: Path) -> AuraWorkbook:
    """严格读取唯一母表；不从装备位置或服务端推测样式。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 正式平台构建会打包依赖。
        raise EquipmentAuraError("当前平台缺少XLSX读取组件 openpyxl") from exc

    path = Path(path)
    if not path.is_file():
        raise EquipmentAuraError(f"装备光环配置表不存在: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "装备绑定" not in workbook.sheetnames:
            raise EquipmentAuraError("配置表缺少“装备绑定”工作表")
        sheet = workbook["装备绑定"]
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers[: len(WORKBOOK_HEADERS)] != WORKBOOK_HEADERS:
            raise EquipmentAuraError(
                f"装备绑定表头不匹配，必须为: {', '.join(WORKBOOK_HEADERS)}"
            )
        bindings: list[AuraBinding] = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(WORKBOOK_HEADERS), values_only=True),
            start=2,
        ):
            if all(value is None or str(value).strip() == "" for value in row):
                continue
            bindings.append(_binding_from_row(row, row_number))
    finally:
        workbook.close()

    enabled_keys: dict[str, str] = {}
    for binding in bindings:
        if not binding.enabled:
            continue
        key = binding.equipment_name.casefold()
        previous = enabled_keys.get(key)
        if previous is not None:
            raise EquipmentAuraError(
                f"启用装备名称重复: {previous} / {binding.equipment_name}"
            )
        enabled_keys[key] = binding.equipment_name
    result = AuraWorkbook(tuple(bindings))
    validate_binding_priorities(result)
    return result


def resolve_target_equipment(database: Path, equipment_names: Sequence[str]) -> dict[str, int]:
    """以只读方式解析目标服 StdItems.Name，每个名称必须唯一。"""
    database = Path(database).resolve()
    if not database.is_file():
        raise EquipmentAuraError(f"目标装备数据库不存在: {database}")
    connection = sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True)
    try:
        rows = connection.execute("SELECT Idx, Name FROM StdItems").fetchall()
    except sqlite3.Error as exc:
        raise EquipmentAuraError(f"无法读取目标StdItems: {exc}") from exc
    finally:
        connection.close()
    by_name: dict[str, list[tuple[str, int]]] = {}
    for idx, raw_name in rows:
        name = "" if raw_name is None else str(raw_name).strip()
        by_name.setdefault(name.casefold(), []).append((name, int(idx)))
    resolved: dict[str, int] = {}
    for requested in equipment_names:
        matches = by_name.get(requested.casefold(), [])
        if not matches:
            raise EquipmentAuraError(f"装备不存在: {requested}")
        if len(matches) != 1:
            details = ", ".join(f"{name}(Idx={idx})" for name, idx in matches)
            raise EquipmentAuraError(f"装备名称不唯一: {requested} -> {details}")
        resolved[requested] = matches[0][1]
    return resolved


def _managed_style_commands(resource_index: int) -> tuple[tuple[str, str], ...]:
    commands: list[tuple[str, str]] = []
    for style_number in range(1, 11):
        style = _style_definition(f"style{style_number:02d}")
        start = (
            f"PLAYEFFECT {resource_index} {style.start_frame} "
            f"{style.frame_count} 0 {style.frame_speed_ms} 1"
        )
        commands.append((start, start.replace("PLAYEFFECT", "StopPlayEffect", 1)))
    return tuple(commands)


def compile_runtime_core(
    bindings: Sequence[AuraBinding],
    resource_index: int,
    *,
    timer_id: int = 97,
) -> str:
    """将配置编译为一个定时器、一条伤害出口的隔离运行时。"""
    enabled = tuple(binding for binding in bindings if binding.enabled)
    workbook = AuraWorkbook(enabled)
    validate_binding_priorities(workbook)
    if not 0 <= resource_index <= 9999:
        raise EquipmentAuraError("特效资源序号必须为0～9999")
    if not 1 <= timer_id <= 99:
        raise EquipmentAuraError("光环定时器必须为1～99")
    sorted_bindings = tuple(sorted(enabled, key=lambda item: item.priority, reverse=True))
    seen_names: set[str] = set()
    for binding in sorted_bindings:
        key = binding.equipment_name.casefold()
        if key in seen_names:
            raise EquipmentAuraError(f"启用装备名称重复: {binding.equipment_name}")
        seen_names.add(key)
        _style_definition(binding.style_id)
        _strict_integer(binding.damage_multiplier, "伤害倍率", 1, 999, 0)
        _strict_integer(binding.interval_seconds, "攻击间隔", 1, 60, 0)
        _strict_integer(binding.priority, "优先级", 1, 999, 0)

    style_commands = _managed_style_commands(resource_index)
    stop_all = tuple(stop for _, stop in style_commands)
    lines: list[str] = [
        "; XY-EQUIPMENT-AURA-10-CORE-BEGIN",
        "; 平台受管：装备名称选择、唯一优先级、安全区停止、仅伤害怪物。",
        "[@XY_EQUIP_AURA_COMMAND]",
        "{",
        "#IF",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_REFRESH",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_REFRESH]",
    ]
    for binding_number, binding in enumerate(sorted_bindings, start=1):
        lines.extend((
            "#IF",
            f"CHECKITEMW {binding.equipment_name}",
            "#ACT",
            f"MOV N$XY_AURA_ACTIVE {binding_number}",
            f"MOV N$XY_AURA_MULTIPLIER {binding.damage_multiplier}",
            f"SetOnTimer {timer_id} {binding.interval_seconds}",
            f"GOTO @XY_EQUIP_AURA_SYNC_{binding_number}",
            "BREAK",
        ))
    lines.extend(("#IF", "#ACT", "GOTO @XY_EQUIP_AURA_STOP", "BREAK", ""))

    for binding_number, binding in enumerate(sorted_bindings, start=1):
        style = _style_definition(binding.style_id)
        start_command = style_commands[int(binding.style_id[-2:]) - 1][0]
        lines.extend((
            f"[@XY_EQUIP_AURA_SYNC_{binding_number}]",
            "#IF",
            "NOT INSAFEZONE",
            f"EQUAL N$XY_AURA_VISUAL {style.start_frame + 1}",
            "#ACT",
            "BREAK",
            "#IF",
            "NOT INSAFEZONE",
            "#ACT",
            *stop_all,
            start_command,
            f"MOV N$XY_AURA_VISUAL {style.start_frame + 1}",
            "BREAK",
            "#IF",
            "#ACT",
            "GOTO @XY_EQUIP_AURA_VISUAL_STOP",
            "BREAK",
            "",
        ))

    lines.append("[@XY_EQUIP_AURA_TICK]")
    for binding_number, binding in enumerate(sorted_bindings, start=1):
        lines.extend((
            "#IF",
            f"CHECKITEMW {binding.equipment_name}",
            "NOT INSAFEZONE",
            "#ACT",
            f"GOTO @XY_EQUIP_AURA_TICK_{binding_number}",
            "BREAK",
        ))
    for binding in sorted_bindings:
        lines.extend((
            "#IF",
            f"CHECKITEMW {binding.equipment_name}",
            "#ACT",
            "GOTO @XY_EQUIP_AURA_TICK_SAFE",
            "BREAK",
        ))
    lines.extend(("#IF", "#ACT", "GOTO @XY_EQUIP_AURA_STOP", "BREAK", ""))

    for binding_number, binding in enumerate(sorted_bindings, start=1):
        style = _style_definition(binding.style_id)
        start_command = style_commands[int(binding.style_id[-2:]) - 1][0]
        lines.extend((
            f"[@XY_EQUIP_AURA_TICK_{binding_number}]",
            "#IF",
            f"EQUAL N$XY_AURA_ACTIVE {binding_number}",
            f"EQUAL N$XY_AURA_VISUAL {style.start_frame + 1}",
            "#ACT",
            "GOTO @XY_EQUIP_AURA_DAMAGE",
            "BREAK",
            "#IF",
            "#ACT",
            *stop_all,
            start_command,
            f"MOV N$XY_AURA_ACTIVE {binding_number}",
            f"MOV N$XY_AURA_MULTIPLIER {binding.damage_multiplier}",
            f"MOV N$XY_AURA_VISUAL {style.start_frame + 1}",
            "GOTO @XY_EQUIP_AURA_DAMAGE",
            "BREAK",
            "",
        ))

    lines.extend((
        "[@XY_EQUIP_AURA_DAMAGE]",
        "#IF",
        "#ACT",
        "MOV N$XY_AURA_DAMAGE <$MAXDC>",
        "MUL N$XY_AURA_DAMAGE <$STR(N$XY_AURA_MULTIPLIER)>",
        "RangeHarm <$X> <$Y> 3 1 6 <$STR(N$XY_AURA_DAMAGE)> 0 2",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_TICK_SAFE]",
        "#IF",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_VISUAL_STOP",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_VISUAL_STOP]",
        "#IF",
        "#ACT",
        *stop_all,
        "MOV N$XY_AURA_VISUAL 0",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_STOP]",
        "#IF",
        "#ACT",
        *stop_all,
        f"SetOffTimer {timer_id}",
        "MOV N$XY_AURA_ACTIVE 0",
        "MOV N$XY_AURA_VISUAL 0",
        "MOV N$XY_AURA_MULTIPLIER 0",
        "BREAK",
        "; XY-EQUIPMENT-AURA-10-CORE-END",
        "}",
        "",
    ))
    return NEWLINE.join(lines)


def _style_start(style: AuraStyle) -> str:
    return (
        f"PLAYEFFECT {style.resource_index} {style.start_frame} "
        f"{style.frame_count} 0 {style.frame_speed} {style.layer}"
    )


def _style_stop(style: AuraStyle) -> str:
    return _style_start(style).replace("PLAYEFFECT", "StopPlayEffect", 1)


def render_multi_style_core(styles: tuple[AuraStyle, ...], spec: AuraSpec) -> str:
    """生成三种灵玉共用的唯一光环核心脚本。"""
    if not styles:
        raise ValueError("至少需要一种光环样式")
    style_ids = [style.style_id for style in styles]
    if len(style_ids) != len(set(style_ids)) or any(style_id <= 0 for style_id in style_ids):
        raise ValueError("光环样式ID必须是互不重复的正整数")
    if len({style.item_name.lower() for style in styles}) != len(styles):
        raise ValueError("光环装备名称必须唯一")
    stop_all = tuple(_style_stop(style) for style in styles)

    lines: list[str] = [
        "; XY-EQUIPMENT-AURA-CORE-BEGIN",
        "; 候选：灵玉、灵玉1、灵玉2共用范围光环。安全区无特效、无伤害。",
        "[@XY_EQUIP_AURA_COMMAND]",
        "{",
        "#IF",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_REFRESH",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_REFRESH]",
    ]
    for style in styles:
        lines.extend(
            (
                "#IF",
                f"CHECKITEMW {style.item_name}",
                "#ACT",
                "MOV N$XY_AURA_ACTIVE 1",
                f"SetOnTimer {spec.timer_id} {spec.interval_seconds}",
                f"GOTO @XY_EQUIP_AURA_SYNC_STYLE_{style.style_id}",
                "BREAK",
            )
        )
    lines.extend(("#IF", "#ACT", "GOTO @XY_EQUIP_AURA_STOP", "BREAK", ""))

    for style in styles:
        lines.extend(
            (
                f"[@XY_EQUIP_AURA_SYNC_STYLE_{style.style_id}]",
                "#IF",
                "NOT INSAFEZONE",
                f"EQUAL N$XY_AURA_VISUAL {style.style_id}",
                "#ACT",
                "BREAK",
                "#IF",
                "NOT INSAFEZONE",
                "#ACT",
                *stop_all,
                _style_start(style),
                f"MOV N$XY_AURA_VISUAL {style.style_id}",
                "BREAK",
                "#IF",
                "#ACT",
                "GOTO @XY_EQUIP_AURA_VISUAL_STOP",
                "BREAK",
                "",
            )
        )

    lines.append("[@XY_EQUIP_AURA_TICK]")
    for style in styles:
        lines.extend(
            (
                "#IF",
                f"CHECKITEMW {style.item_name}",
                "NOT INSAFEZONE",
                "#ACT",
                f"GOTO @XY_EQUIP_AURA_TICK_STYLE_{style.style_id}",
                "BREAK",
            )
        )
    for style in styles:
        lines.extend(
            (
                "#IF",
                f"CHECKITEMW {style.item_name}",
                "#ACT",
                "GOTO @XY_EQUIP_AURA_TICK_SAFE",
                "BREAK",
            )
        )
    lines.extend(("#IF", "#ACT", "GOTO @XY_EQUIP_AURA_STOP", "BREAK", ""))

    for style in styles:
        lines.extend(
            (
                f"[@XY_EQUIP_AURA_TICK_STYLE_{style.style_id}]",
                "#IF",
                f"EQUAL N$XY_AURA_VISUAL {style.style_id}",
                "#ACT",
                "GOTO @XY_EQUIP_AURA_DAMAGE",
                "BREAK",
                "#IF",
                "#ACT",
                *stop_all,
                _style_start(style),
                f"MOV N$XY_AURA_VISUAL {style.style_id}",
                "GOTO @XY_EQUIP_AURA_DAMAGE",
                "BREAK",
                "",
            )
        )

    lines.extend(
        (
            "[@XY_EQUIP_AURA_DAMAGE]",
            "#IF",
            "#ACT",
            "MOV N$XY_AURA_DAMAGE <$MAXDC>",
            f"MUL N$XY_AURA_DAMAGE {spec.damage_multiplier}",
            (
                f"RangeHarm <$X> <$Y> {spec.damage_range} 1 6 "
                "<$STR(N$XY_AURA_DAMAGE)> 0 2"
            ),
            "BREAK",
            "",
            "[@XY_EQUIP_AURA_TICK_SAFE]",
            "#IF",
            "#ACT",
            "GOTO @XY_EQUIP_AURA_VISUAL_STOP",
            "BREAK",
            "",
            "[@XY_EQUIP_AURA_VISUAL_STOP]",
            "#IF",
            "#ACT",
            *stop_all,
            "MOV N$XY_AURA_VISUAL 0",
            "BREAK",
            "",
            "[@XY_EQUIP_AURA_STOP]",
            "#IF",
            "#ACT",
            *stop_all,
            f"SetOffTimer {spec.timer_id}",
            "MOV N$XY_AURA_ACTIVE 0",
            "MOV N$XY_AURA_VISUAL 0",
            "BREAK",
            "; XY-EQUIPMENT-AURA-CORE-END",
            "}",
            "",
        )
    )
    return NEWLINE.join(lines)


@dataclass(frozen=True)
class CandidatePaths:
    server_root: Path
    client_root: Path
    asset_wzl: Path
    asset_wzx: Path
    qfunction: Path
    qmanage: Path
    effect_list: Path
    database: Path
    core_script: Path
    client_wzl: Path
    client_wzx: Path

    @classmethod
    def from_roots(
        cls, server: Path, client: Path, asset_wzl: Path, asset_wzx: Path
    ) -> "CandidatePaths":
        server = Path(server).resolve()
        client = Path(client).resolve()
        envir = server / "Mir200" / "Envir"
        database_candidates = (
            server / "Mud2" / "DB" / "ApexM2.DB",
            server / "DBServer" / "ApexM2.DB",
            envir / "ApexM2.DB",
        )
        database = next((path for path in database_candidates if path.is_file()), database_candidates[0])
        return cls(
            server_root=server,
            client_root=client,
            asset_wzl=Path(asset_wzl).resolve(),
            asset_wzx=Path(asset_wzx).resolve(),
            qfunction=envir / "Market_Def" / "QFunction-0.txt",
            qmanage=envir / "MapQuest_Def" / "QManage.txt",
            effect_list=envir / "EffectImageList.txt",
            database=database,
            core_script=envir / "QuestDiary" / "玄渊实验室" / "装备范围光环" / "装备范围光环核心.txt",
            client_wzl=client / "Data" / "XY_EquipmentAura.wzl",
            client_wzx=client / "Data" / "XY_EquipmentAura.wzx",
        )


@dataclass(frozen=True)
class PlannedFile:
    path: Path
    before_hash: str | None
    after_bytes: bytes

    @property
    def after_hash(self) -> str:
        return _sha256_bytes(self.after_bytes)


@dataclass(frozen=True)
class CandidatePlan:
    paths: CandidatePaths
    spec: AuraSpec
    resource_index: int
    blockers: tuple[str, ...]
    warnings: tuple[str, ...]
    files: tuple[PlannedFile, ...]


def _effect_start(resource_index: int, spec: AuraSpec) -> tuple[str, ...]:
    back = (
        f"PLAYEFFECT {resource_index} 0 {spec.frame_count} 0 "
        f"{spec.frame_speed} 1"
    )
    if spec.effect_layer_count == 1:
        return (back,)
    if spec.effect_layer_count != 2:
        raise ValueError(f"不支持的特效层数: {spec.effect_layer_count}")
    front = (
        f"PLAYEFFECT {resource_index} {spec.frame_count} {spec.frame_count} 0 "
        f"{spec.frame_speed} 0"
    )
    return back, front


def _effect_stop(resource_index: int, spec: AuraSpec) -> tuple[str, ...]:
    return tuple(
        command.replace("PLAYEFFECT", "StopPlayEffect", 1)
        for command in _effect_start(resource_index, spec)
    )


def render_core_script(spec: AuraSpec, resource_index: int) -> str:
    """生成独立核心脚本；伤害只走 RangeHarm 怪物路线。"""
    start_effects = _effect_start(resource_index, spec)
    stop_effects = _effect_stop(resource_index, spec)
    lines = [
        "; XY-EQUIPMENT-AURA-CORE-BEGIN",
        f"; 候选：灵玉{spec.visual_name}。安全区无特效、无伤害。",
        "[@XY_EQUIP_AURA_COMMAND]",
        "{",
        "#IF",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_REFRESH",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_REFRESH]",
        "#IF",
        f"CHECKITEMW {spec.item_name}",
        "#ACT",
        "MOV N$XY_AURA_ACTIVE 1",
        f"SetOnTimer {spec.timer_id} {spec.interval_seconds}",
        "GOTO @XY_EQUIP_AURA_VISUAL_SYNC",
        "BREAK",
        "#IF",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_STOP",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_VISUAL_SYNC]",
        "#IF",
        "NOT INSAFEZONE",
        "EQUAL N$XY_AURA_VISUAL 0",
        "#ACT",
        *start_effects,
        "MOV N$XY_AURA_VISUAL 1",
        "BREAK",
        "#IF",
        "NOT INSAFEZONE",
        "#ACT",
        "BREAK",
        "#IF",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_VISUAL_STOP",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_VISUAL_STOP]",
        "#IF",
        "EQUAL N$XY_AURA_VISUAL 1",
        "#ACT",
        *stop_effects,
        "MOV N$XY_AURA_VISUAL 0",
        "BREAK",
        "#IF",
        "#ACT",
        "MOV N$XY_AURA_VISUAL 0",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_TICK]",
        "#IF",
        f"NOT CHECKITEMW {spec.item_name}",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_STOP",
        "BREAK",
        "#IF",
        "NOT INSAFEZONE",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_TICK_OUTSIDE",
        "BREAK",
        "#IF",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_TICK_SAFE",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_TICK_OUTSIDE]",
        "#IF",
        "EQUAL N$XY_AURA_VISUAL 0",
        "#ACT",
        *start_effects,
        "MOV N$XY_AURA_VISUAL 1",
        "#IF",
        "#ACT",
        "MOV N$XY_AURA_DAMAGE <$MAXDC>",
        f"MUL N$XY_AURA_DAMAGE {spec.damage_multiplier}",
        (
            f"RangeHarm <$X> <$Y> {spec.damage_range} 1 6 "
            "<$STR(N$XY_AURA_DAMAGE)> 0 2"
        ),
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_TICK_SAFE]",
        "#IF",
        "#ACT",
        "GOTO @XY_EQUIP_AURA_VISUAL_STOP",
        "BREAK",
        "",
        "[@XY_EQUIP_AURA_STOP]",
        "#IF",
        "#ACT",
        *stop_effects,
        f"SetOffTimer {spec.timer_id}",
        "MOV N$XY_AURA_ACTIVE 0",
        "MOV N$XY_AURA_VISUAL 0",
        "BREAK",
        "; XY-EQUIPMENT-AURA-CORE-END",
        "}",
        "",
    ]
    return NEWLINE.join(lines)


def _legacy_unwrapped_core(expected_core: bytes) -> bytes:
    """重建 candidate.1 唯一已知错误格式，供事务升级时精确比对。"""
    text = _decode_script(expected_core)
    text = text.replace(
        "[@XY_EQUIP_AURA_COMMAND]\r\n{\r\n",
        "[@XY_EQUIP_AURA_COMMAND]\r\n",
        1,
    )
    if not text.endswith("}\r\n"):
        raise ValueError("预期核心脚本缺少外层结束花括号")
    return _encode_script(text[:-3])


def _normalize_blade(source: Path, spec: AuraSpec) -> Image.Image:
    image = Image.open(source).convert("RGBA")
    bbox = image.getchannel("A").getbbox()
    if bbox is None:
        raise ValueError("刀剑源图没有可见像素")
    image = image.crop(bbox)
    if spec.source_rotation_degrees:
        image = image.rotate(
            spec.source_rotation_degrees,
            resample=Image.Resampling.BICUBIC,
            expand=True,
        )
        rotated_bbox = image.getchannel("A").getbbox()
        if rotated_bbox is None:
            raise ValueError("刀剑源图旋正后没有可见像素")
        image = image.crop(rotated_bbox)
    max_width, max_height = 70, 28
    scale = min(max_width / image.width, max_height / image.height)
    size = (max(1, round(image.width * scale)), max(1, round(image.height * scale)))
    return image.resize(size, Image.Resampling.LANCZOS)


def _paste_center(canvas: Image.Image, image: Image.Image, x: float, y: float) -> None:
    left = round(x - image.width / 2)
    top = round(y - image.height / 2)
    canvas.alpha_composite(image, (left, top))


def render_orbit_frames(source: Path, output_dir: Path, spec: AuraSpec) -> tuple[Path, ...]:
    """由单把透明刀剑生成24帧人物前后分层径向刀轮。"""
    blade = _normalize_blade(Path(source), spec)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    center_x, center_y = spec.canvas_width / 2, spec.canvas_height / 2
    # 刀柄从内圈外沿开始，刀身沿半径向外；不是沿圆周切线摆放。
    blade_center_radius = float(spec.inner_radius) + blade.width / 2
    generated: list[Image.Image] = []
    for frame_index in range(spec.frame_count):
        back = Image.new("RGBA", (spec.canvas_width, spec.canvas_height), (0, 0, 0, 0))
        front = Image.new("RGBA", (spec.canvas_width, spec.canvas_height), (0, 0, 0, 0))
        # 十把相同刀剑每转过一把刀的角距，画面即进入下一完整循环。
        # 因此24帧只覆盖36度，避免一轮内重复播放十个相同周期。
        phase = 2 * math.pi * frame_index / (spec.frame_count * spec.blade_count)
        for blade_index in range(spec.blade_count):
            angle = phase + 2 * math.pi * blade_index / spec.blade_count
            x = center_x + blade_center_radius * math.cos(angle)
            y = center_y + blade_center_radius * math.sin(angle)
            radial_degrees = math.degrees(angle)
            rotated = blade.rotate(-radial_degrees, resample=Image.Resampling.BICUBIC, expand=True)
            target = back if math.sin(angle) < 0 else front
            _paste_center(target, rotated, x, y)
        generated.extend((back, front))
    ordered = generated[0::2] + generated[1::2]
    for index, frame in enumerate(ordered):
        path = output_dir / f"{index:05d}.png"
        frame.save(path, format="PNG", optimize=False, compress_level=9)
        paths.append(path)
    return tuple(paths)


def _decode_script(data: bytes) -> str:
    return data.decode(ENCODING)


def _encode_script(text: str) -> bytes:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").replace("\n", NEWLINE)
    return normalized.encode(ENCODING)


def _managed_block(kind: str, name: str, body_lines: Iterable[str]) -> str:
    body = NEWLINE.join(body_lines) + NEWLINE
    digest = _sha256_bytes(body.encode(ENCODING))
    return (
        f"; XY-EQUIPMENT-AURA-{kind}-BEGIN {name} SHA256={digest}{NEWLINE}"
        f"{body}"
        f"; XY-EQUIPMENT-AURA-{kind}-END {name}{NEWLINE}"
    )


def _validate_existing_managed_blocks(text: str, kind: str) -> list[str]:
    blockers: list[str] = []
    pattern = re.compile(
        rf"; XY-EQUIPMENT-AURA-{kind}-BEGIN ([^\r\n]+?) SHA256=([0-9A-F]{{64}})\r?\n"
        rf"(.*?)"
        rf"; XY-EQUIPMENT-AURA-{kind}-END \1\r?\n?",
        re.IGNORECASE | re.DOTALL,
    )
    begin_count = len(re.findall(rf"; XY-EQUIPMENT-AURA-{kind}-BEGIN", text, re.IGNORECASE))
    matches = list(pattern.finditer(text))
    if begin_count != len(matches):
        blockers.append(f"受管块结构损坏: {kind}")
        return blockers
    for match in matches:
        body = (
            match.group(3)
            .replace("\r\n", "\n")
            .replace("\r", "\n")
            .replace("\n", NEWLINE)
        )
        actual = _sha256_bytes(body.encode(ENCODING))
        if actual != match.group(2).upper():
            blockers.append(f"受管块被手工修改: {kind}/{match.group(1)}")
    return blockers


def _inject_event_hook(
    text: str, event_name: str, body_lines: Iterable[str]
) -> tuple[str, str | None]:
    block = _managed_block(
        "HOOK",
        event_name,
        body_lines,
    )
    existing_marker = f"XY-EQUIPMENT-AURA-HOOK-BEGIN {event_name} "
    if existing_marker.lower() in text.lower():
        pattern = re.compile(
            rf"; XY-EQUIPMENT-AURA-HOOK-BEGIN {re.escape(event_name)} SHA256=[0-9A-F]{{64}}\r?\n"
            rf".*?"
            rf"; XY-EQUIPMENT-AURA-HOOK-END {re.escape(event_name)}\r?\n?",
            re.IGNORECASE | re.DOTALL,
        )
        matches = list(pattern.finditer(text))
        if len(matches) != 1:
            return text, f"受管事件块数量异常: {event_name}={len(matches)}"
        match = matches[0]
        return text[: match.start()] + block + text[match.end() :], None
    labels = list(re.finditer(rf"(?im)^\[@{re.escape(event_name)}\]\s*$", text))
    if len(labels) != 1:
        return text, f"事件标签数量异常: [@{event_name}]={len(labels)}"
    start = labels[0].start()
    next_label = re.search(r"(?im)^\[@[^\]]+\]\s*$", text[labels[0].end() :])
    end = labels[0].end() + next_label.start() if next_label else len(text)
    section = text[start:end]
    break_matches = list(re.finditer(r"(?im)^BREAK\s*$", section))
    if not break_matches:
        return text, f"事件标签缺少安全插入位置: [@{event_name}]"
    insert_at = start + break_matches[-1].start()
    prefix = "" if text[:insert_at].endswith(("\n", "\r")) else NEWLINE
    return text[:insert_at] + prefix + block + text[insert_at:], None


def _upsert_qfunction_login_refresh(text: str) -> tuple[str, str | None]:
    label = "XY_EQUIP_AURA_LOGIN_REFRESH"
    block = _managed_block(
        "EVENT",
        label,
        (
            f"[@{label}]",
            "#IF",
            "#ACT",
            "#CALL [\\玄渊实验室\\装备范围光环\\装备范围光环核心.txt] @XY_EQUIP_AURA_REFRESH",
            "BREAK",
        ),
    )
    marker = f"XY-EQUIPMENT-AURA-EVENT-BEGIN {label} "
    label_count = len(re.findall(rf"(?im)^\[@{re.escape(label)}\]\s*$", text))
    if marker.lower() in text.lower():
        pattern = re.compile(
            rf"; XY-EQUIPMENT-AURA-EVENT-BEGIN {re.escape(label)} SHA256=[0-9A-F]{{64}}\r?\n"
            rf".*?"
            rf"; XY-EQUIPMENT-AURA-EVENT-END {re.escape(label)}\r?\n?",
            re.IGNORECASE | re.DOTALL,
        )
        matches = list(pattern.finditer(text))
        if len(matches) != 1 or label_count != 1:
            return text, f"登录延迟刷新标签数量异常: [@{label}]={label_count}"
        match = matches[0]
        return text[: match.start()] + block + text[match.end() :], None
    if label_count:
        return text, f"登录延迟刷新标签已被占用: [@{label}]"
    if text and not text.endswith(("\n", "\r")):
        text += NEWLINE
    return text + NEWLINE + block, None


def _render_qfunction(original: bytes) -> tuple[bytes, list[str]]:
    text = _decode_script(original)
    blockers = _validate_existing_managed_blocks(text, "HOOK")
    blockers.extend(_validate_existing_managed_blocks(text, "EVENT"))
    if blockers:
        return original, blockers
    hooks = {
        "PlayLogin": (
            "#IF",
            "#ACT",
            "DELAYGOTO 2 @XY_EQUIP_AURA_LOGIN_REFRESH",
        ),
        "TakeOnEx": (
            "#IF",
            "#ACT",
            "#CALL [\\玄渊实验室\\装备范围光环\\装备范围光环核心.txt] @XY_EQUIP_AURA_REFRESH",
        ),
        "TakeOffEx": (
            "#IF",
            "#ACT",
            "#CALL [\\玄渊实验室\\装备范围光环\\装备范围光环核心.txt] @XY_EQUIP_AURA_REFRESH",
        ),
        "PlayDie": (
            "#IF",
            "#ACT",
            "#CALL [\\玄渊实验室\\装备范围光环\\装备范围光环核心.txt] @XY_EQUIP_AURA_STOP",
        ),
        "NpcRevival": (
            "#IF",
            "#ACT",
            "#CALL [\\玄渊实验室\\装备范围光环\\装备范围光环核心.txt] @XY_EQUIP_AURA_REFRESH",
        ),
        "EnterMap": (
            "#IF",
            "#ACT",
            "#CALL [\\玄渊实验室\\装备范围光环\\装备范围光环核心.txt] @XY_EQUIP_AURA_REFRESH",
        ),
    }
    for event_name, body_lines in hooks.items():
        if (
            event_name == "EnterMap"
            and f"XY-EQUIPMENT-AURA-HOOK-BEGIN {event_name} ".lower() not in text.lower()
            and not re.search(rf"(?im)^\[@{re.escape(event_name)}\]\s*$", text)
        ):
            # 旧服和部分空白夹具没有换图事件；定时器仍会在1秒内自愈。
            # 目标若存在 EnterMap，则必须受管接入以实现即时刷新。
            continue
        text, blocker = _inject_event_hook(text, event_name, body_lines)
        if blocker:
            blockers.append(blocker)
    if not blockers:
        text, blocker = _upsert_qfunction_login_refresh(text)
        if blocker:
            blockers.append(blocker)
    return _encode_script(text), blockers


def _render_qmanage(original: bytes, spec: AuraSpec) -> tuple[bytes, list[str]]:
    text = _decode_script(original)
    blockers = _validate_existing_managed_blocks(text, "EVENT")
    event_label = f"OnTimer{spec.timer_id}"
    marker = f"XY-EQUIPMENT-AURA-EVENT-BEGIN {event_label} "
    label_count = len(re.findall(rf"(?im)^\[@{re.escape(event_label)}\]\s*$", text))
    if marker.lower() in text.lower():
        block = _managed_block(
            "EVENT",
            event_label,
            (
                f"[@{event_label}]",
                "#IF",
                "#ACT",
                "#CALL [\\玄渊实验室\\装备范围光环\\装备范围光环核心.txt] @XY_EQUIP_AURA_TICK",
                "BREAK",
            ),
        )
        pattern = re.compile(
            rf"; XY-EQUIPMENT-AURA-EVENT-BEGIN {re.escape(event_label)} SHA256=[0-9A-F]{{64}}\r?\n"
            rf".*?"
            rf"; XY-EQUIPMENT-AURA-EVENT-END {re.escape(event_label)}\r?\n?",
            re.IGNORECASE | re.DOTALL,
        )
        matches = list(pattern.finditer(text))
        if len(matches) != 1:
            blockers.append(f"受管定时器块数量异常: {event_label}={len(matches)}")
            return original, blockers
        match = matches[0]
        updated = text[: match.start()] + block + text[match.end() :]
        return _encode_script(updated), blockers
    if label_count:
        blockers.append(f"定时器事件已被占用: [@{event_label}]")
        return original, blockers
    block = _managed_block(
        "EVENT",
        event_label,
        (
            f"[@{event_label}]",
            "#IF",
            "#ACT",
            "#CALL [\\玄渊实验室\\装备范围光环\\装备范围光环核心.txt] @XY_EQUIP_AURA_TICK",
            "BREAK",
        ),
    )
    if text and not text.endswith(("\n", "\r")):
        text += NEWLINE
    return _encode_script(text + NEWLINE + block), blockers


def _effect_list_plan(original: bytes, resource_name: str) -> tuple[bytes, int]:
    text = _decode_script(original)
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    matches = [index for index, line in enumerate(lines) if line.strip().lower() == resource_name.lower()]
    if matches:
        return original, matches[0]
    index = len(lines)
    lines.append(resource_name)
    return _encode_script(NEWLINE.join(lines) + NEWLINE), index


def _changed_file(path: Path, after: bytes) -> PlannedFile | None:
    before_hash = _sha256_file(path) if path.is_file() else None
    if path.is_file() and path.read_bytes() == after:
        return None
    return PlannedFile(path=path, before_hash=before_hash, after_bytes=after)


def _check_item(paths: CandidatePaths, spec: AuraSpec) -> list[str]:
    if not paths.database.is_file():
        return [f"装备数据库不存在: {paths.database}"]
    connection = sqlite3.connect(f"file:{paths.database.as_posix()}?mode=ro", uri=True)
    try:
        rows = connection.execute(
            "SELECT Idx, Name, StdMode, Looks FROM StdItems WHERE Name = ?", (spec.item_name,)
        ).fetchall()
    except sqlite3.Error as error:
        return [f"装备数据库读取失败: {error}"]
    finally:
        connection.close()
    if not rows:
        return [f"装备名称不存在: {spec.item_name}"]
    if len(rows) != 1:
        details = ", ".join(str(row[0]) for row in rows)
        return [f"装备名称不唯一: {spec.item_name} (Idx={details})"]
    return []


def _check_external_timer_owners(paths: CandidatePaths, spec: AuraSpec) -> list[str]:
    blockers: list[str] = []
    pattern = re.compile(
        rb"(?i)(?:SetOnTimer\s+" + str(spec.timer_id).encode("ascii") + rb"\b|\[@OnTimer" + str(spec.timer_id).encode("ascii") + rb"\])"
    )
    envir = paths.server_root / "Mir200" / "Envir"
    for path in envir.rglob("*.txt"):
        resolved = path.resolve()
        if resolved in {paths.qmanage.resolve(), paths.core_script.resolve()}:
            continue
        try:
            if pattern.search(path.read_bytes()):
                blockers.append(f"定时器{spec.timer_id}被其他脚本占用: {path}")
        except OSError as error:
            blockers.append(f"定时器占用扫描失败: {path} ({error})")
    return blockers


def plan_candidate(paths: CandidatePaths, spec: AuraSpec) -> CandidatePlan:
    blockers: list[str] = []
    warnings = ["M2运行时允许写入；新脚本和资源需由用户重载或重启后生效。"]
    required = (paths.qfunction, paths.qmanage, paths.effect_list, paths.asset_wzl, paths.asset_wzx)
    for path in required:
        if not path.is_file():
            blockers.append(f"必需文件不存在: {path}")
    missing_required = bool(blockers)
    blockers.extend(_check_item(paths, spec))
    if missing_required:
        return CandidatePlan(paths, spec, -1, tuple(blockers), tuple(warnings), ())

    blockers.extend(_check_external_timer_owners(paths, spec))

    qfunction_after, qfunction_blockers = _render_qfunction(paths.qfunction.read_bytes())
    qmanage_after, qmanage_blockers = _render_qmanage(paths.qmanage.read_bytes(), spec)
    blockers.extend(qfunction_blockers)
    blockers.extend(qmanage_blockers)
    effects_after, resource_index = _effect_list_plan(paths.effect_list.read_bytes(), spec.resource_name)
    effect_names = [
        line.strip().lower()
        for line in _decode_script(paths.effect_list.read_bytes()).replace("\r\n", "\n").split("\n")
        if line.strip()
    ]
    if effect_names.count(spec.resource_name.lower()) > 1:
        blockers.append(f"EffectImageList资源重复: {spec.resource_name}")

    core_after = _encode_script(render_core_script(spec, resource_index))
    if paths.core_script.is_file() and paths.core_script.read_bytes() != core_after:
        accepted_prior_cores = {_legacy_unwrapped_core(core_after)}
        legacy_base = replace(
            spec,
            frame_count=24,
            effect_layer_count=2,
            visual_name="金焰刀轮",
        )
        for speed in (4, 10, 30, 80):
            prior_core = _encode_script(
                render_core_script(replace(legacy_base, frame_speed=speed), resource_index)
            )
            accepted_prior_cores.add(prior_core)
            accepted_prior_cores.add(_legacy_unwrapped_core(prior_core))
        if paths.core_script.read_bytes() not in accepted_prior_cores:
            blockers.append(f"核心脚本已存在且内容不一致: {paths.core_script}")

    source_pairs = ((paths.asset_wzl, paths.client_wzl), (paths.asset_wzx, paths.client_wzx))
    for source, target in source_pairs:
        if target.is_file() and target.read_bytes() != source.read_bytes():
            if not _is_known_prior_client_resource(target.name, _sha256_file(target)):
                blockers.append(f"客户端同名资源冲突: {target}")
    if blockers:
        return CandidatePlan(paths, spec, resource_index, tuple(blockers), tuple(warnings), ())
    proposed = (
        (paths.qfunction, qfunction_after),
        (paths.qmanage, qmanage_after),
        (paths.effect_list, effects_after),
        (paths.core_script, core_after),
        (paths.client_wzl, paths.asset_wzl.read_bytes()),
        (paths.client_wzx, paths.asset_wzx.read_bytes()),
    )
    files = tuple(change for path, after in proposed if (change := _changed_file(path, after)))
    return CandidatePlan(paths, spec, resource_index, (), tuple(warnings), files)


def _atomic_write(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise


def apply_candidate(plan: CandidatePlan, backup_root: Path) -> Path:
    if plan.blockers:
        raise RuntimeError("预检存在阻止项，禁止安装: " + "；".join(plan.blockers))
    if not plan.files:
        raise RuntimeError("目标已是候选状态，无需重复安装")
    for change in plan.files:
        current_hash = _sha256_file(change.path) if change.path.is_file() else None
        if current_hash != change.before_hash:
            raise RuntimeError(f"预检后文件已变化: {change.path}")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_dir = Path(backup_root) / f"equipment_aura_{timestamp}"
    files_dir = backup_dir / "files"
    files_dir.mkdir(parents=True, exist_ok=False)
    entries: list[dict[str, object]] = []
    try:
        for index, change in enumerate(plan.files):
            existed = change.path.is_file()
            backup_path = files_dir / f"{index:02d}.bin"
            if existed:
                shutil.copyfile(change.path, backup_path)
            entries.append(
                {
                    "path": str(change.path),
                    "existed_before": existed,
                    "before_hash": change.before_hash,
                    "after_hash": change.after_hash,
                    "backup": str(backup_path) if existed else None,
                }
            )
        for change in plan.files:
            _atomic_write(change.path, change.after_bytes)
            if _sha256_file(change.path) != change.after_hash:
                raise RuntimeError(f"写入后哈希不符: {change.path}")
    except Exception as install_error:
        rollback_errors: list[str] = []
        for change, entry in reversed(list(zip(plan.files, entries))):
            path = change.path
            try:
                current_hash = _sha256_file(path) if path.is_file() else None
            except Exception as inspect_error:
                rollback_errors.append(f"{path}: 无法读取当前文件 ({inspect_error})")
                continue
            if current_hash == entry["before_hash"]:
                continue
            try:
                if entry["existed_before"]:
                    _atomic_write(path, Path(str(entry["backup"])).read_bytes())
                    restored_hash = _sha256_file(path)
                    if restored_hash != entry["before_hash"]:
                        raise RuntimeError(f"恢复后哈希不符: {restored_hash}")
                elif path.exists():
                    path.unlink()
            except Exception as rollback_error:
                rollback_errors.append(f"{path}: {rollback_error}")
        if rollback_errors:
            raise RuntimeError(
                "安装失败且部分目标无法验证或回滚；其余可恢复文件已继续处理: "
                + "；".join(rollback_errors)
            ) from install_error
        raise

    receipt = {
        "schema_version": 1,
        "operation": "equipment-aura-candidate",
        "created_at": datetime.now().astimezone().isoformat(),
        "resource_index": plan.resource_index,
        "spec": asdict(plan.spec),
        "files": entries,
        "status": "installed-pending-game-verification",
    }
    receipt_path = backup_dir / "receipt.json"
    _atomic_write(receipt_path, json.dumps(receipt, ensure_ascii=False, indent=2).encode("utf-8"))
    return receipt_path


def rollback_candidate(receipt_path: Path) -> None:
    receipt_path = Path(receipt_path)
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    if receipt.get("operation") not in {
        "equipment-aura-candidate",
        "equipment-aura-jades-candidate",
    }:
        raise RuntimeError("不是装备范围光环安装收据")
    entries = receipt["files"]
    for entry in entries:
        path = Path(entry["path"])
        current_hash = _sha256_file(path) if path.is_file() else None
        if current_hash != entry["after_hash"]:
            raise RuntimeError(f"安装后文件已被修改，禁止盲目回滚: {path}")
    for entry in reversed(entries):
        path = Path(entry["path"])
        if entry["existed_before"]:
            _atomic_write(path, Path(entry["backup"]).read_bytes())
            if _sha256_file(path) != entry["before_hash"]:
                raise RuntimeError(f"回滚哈希不符: {path}")
        elif path.exists():
            path.unlink()
    receipt["status"] = "rolled-back"
    receipt["rolled_back_at"] = datetime.now().astimezone().isoformat()
    _atomic_write(receipt_path, json.dumps(receipt, ensure_ascii=False, indent=2).encode("utf-8"))


# ---------------------------------------------------------------------------
# 10套光环平台服务：一份XLSX、一个核心、一个定时器、一条伤害出口。


def sha256(path: Path) -> str:
    return _sha256_file(Path(path))


@dataclass(frozen=True)
class AuraTargetPaths:
    server_root: Path
    client_root: Path
    login_root: Path
    qfunction: Path
    qmanage: Path
    effect_list: Path
    database: Path
    core_script: Path
    client_wzl: Path
    client_wzx: Path
    login_wzl: Path
    login_wzx: Path

    @classmethod
    def from_roots(cls, server: Path, client: Path, login: Path) -> "AuraTargetPaths":
        server = Path(server).resolve()
        client = Path(client).resolve()
        login = Path(login).resolve()
        envir = server / "Mir200" / "Envir"
        database_candidates = (
            server / "Mud2" / "DB" / "ApexM2.DB",
            server / "DBServer" / "ApexM2.DB",
            envir / "ApexM2.DB",
        )
        database = next((path for path in database_candidates if path.is_file()), database_candidates[0])
        return cls(
            server_root=server,
            client_root=client,
            login_root=login,
            qfunction=envir / "Market_Def" / "QFunction-0.txt",
            qmanage=envir / "MapQuest_Def" / "QManage.txt",
            effect_list=envir / "EffectImageList.txt",
            database=database,
            core_script=(
                envir
                / "QuestDiary"
                / "玄渊实验室"
                / "装备范围光环"
                / "装备范围光环核心.txt"
            ),
            client_wzl=client / "data" / "XY_EquipmentAura_10.wzl",
            client_wzx=client / "data" / "XY_EquipmentAura_10.wzx",
            login_wzl=login / "XY_EquipmentAura_10.wzl",
            login_wzx=login / "XY_EquipmentAura_10.wzx",
        )


@dataclass(frozen=True)
class AuraInstallPlan:
    plan_id: str
    platform_root: Path
    paths: AuraTargetPaths
    workbook_path: Path
    workbook_hash: str
    workbook: AuraWorkbook
    resource_index: int
    static_resource_action: str
    blockers: tuple[str, ...]
    warnings: tuple[str, ...]
    files: tuple[PlannedFile, ...]

    @property
    def affected_paths(self) -> tuple[Path, ...]:
        return tuple(change.path for change in self.files)


@dataclass(frozen=True)
class AuraInstallReceipt:
    transaction_id: str
    receipt_path: Path
    status: str
    changed_files: tuple[Path, ...]


def _wrap_runtime_core(runtime: str) -> bytes:
    lines = runtime.replace("\r\n", "\n").replace("\r", "\n").rstrip("\n").split("\n")
    return _encode_script(_managed_block("CORE10", "runtime", lines))


def _known_legacy_multi_style_core(effect_list: bytes) -> set[bytes]:
    """只允许平台历史上已生成的三种灵玉核心升级，不猜测手写脚本。"""
    names = [
        line.strip().lower()
        for line in _decode_script(effect_list).replace("\r\n", "\n").replace("\r", "\n").split("\n")
        if line.strip()
    ]
    try:
        original_index = names.index("xy_equipmentaura.wzl")
        jades_index = names.index("xy_equipmentaura_jades.wzl")
    except ValueError:
        return set()
    styles = (
        AuraStyle("灵玉2", jades_index, 8, 8, 30, 1, 3),
        AuraStyle("灵玉1", jades_index, 0, 8, 30, 1, 2),
        AuraStyle("灵玉", original_index, 0, 8, 30, 1, 1),
    )
    return {_encode_script(render_multi_style_core(styles, AuraSpec.round_moon()))}


def _scan_external_timer_97(paths: AuraTargetPaths) -> list[str]:
    blockers: list[str] = []
    pattern = re.compile(rb"(?i)(?:SetOnTimer\s+97\b|\[@OnTimer97\])")
    envir = paths.server_root / "Mir200" / "Envir"
    if not envir.is_dir():
        return blockers
    excluded = {paths.qmanage.resolve(), paths.core_script.resolve()}
    for path in envir.rglob("*.txt"):
        if path.resolve() in excluded:
            continue
        try:
            if pattern.search(path.read_bytes()):
                blockers.append(f"定时器97被其他脚本占用: {path}")
        except OSError as exc:
            blockers.append(f"定时器97占用扫描失败: {path} ({exc})")
    return blockers


def _plan_static_pair(
    source_wzl: Path,
    source_wzx: Path,
    target_wzl: Path,
    target_wzx: Path,
    label: str,
) -> tuple[str, list[str], list[tuple[Path, bytes]]]:
    blockers: list[str] = []
    proposed: list[tuple[Path, bytes]] = []
    if not source_wzl.is_file() or not source_wzx.is_file():
        return "blocked", [f"平台静态资源缺失: {source_wzl} / {source_wzx}"], []
    exists = (target_wzl.is_file(), target_wzx.is_file())
    if exists[0] != exists[1]:
        return "blocked", [f"{label}静态资源不完整: {target_wzl} / {target_wzx}"], []
    if all(exists):
        for source, target in ((source_wzl, target_wzl), (source_wzx, target_wzx)):
            if _sha256_file(source) != _sha256_file(target):
                blockers.append(f"{label}静态资源冲突: {target}")
        return ("blocked" if blockers else "reuse"), blockers, []
    proposed.extend(((target_wzl, source_wzl.read_bytes()), (target_wzx, source_wzx.read_bytes())))
    return "install", blockers, proposed


class EquipmentAuraService:
    """装备光环的唯一平台核心；GUI和CLI只调用本类。"""

    def __init__(self, platform_root: Path) -> None:
        self.platform_root = Path(platform_root).resolve()
        self.library_root = self.platform_root / "assets" / "equipment_aura" / "library"
        self.asset_wzl = self.library_root / "XY_EquipmentAura_10.wzl"
        self.asset_wzx = self.library_root / "XY_EquipmentAura_10.wzx"

    def list_styles(self) -> tuple[dict[str, object], ...]:
        return tuple(
            {
                "style_id": style.style_id,
                "display_name": style.display_name,
                "frame_start": style.start_frame,
                "frame_count": style.frame_count,
                "frame_speed_ms": style.frame_speed_ms,
                "range": 3,
                "size": "100%",
            }
            for style in RUNTIME_STYLE_CATALOG
        )

    def search_target_equipment(
        self,
        server: Path,
        keyword: str,
        *,
        limit: int = 100,
    ) -> tuple[dict[str, object], ...]:
        """按名称搜索目标服真实装备；结果只来自目标StdItems。"""
        server = Path(server).resolve()
        candidates = (
            server / "Mud2" / "DB" / "ApexM2.DB",
            server / "DBServer" / "ApexM2.DB",
            server / "Mir200" / "Envir" / "ApexM2.DB",
        )
        database = next((path for path in candidates if path.is_file()), candidates[0])
        if not database.is_file():
            raise EquipmentAuraError(f"目标装备数据库不存在: {database}")
        needle = str(keyword).strip().casefold()
        connection = sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True)
        try:
            rows = connection.execute("SELECT Idx, Name FROM StdItems ORDER BY Idx").fetchall()
        except sqlite3.Error as exc:
            raise EquipmentAuraError(f"无法读取目标StdItems: {exc}") from exc
        finally:
            connection.close()
        matches = []
        for idx, raw_name in rows:
            name = "" if raw_name is None else str(raw_name).strip()
            if not name or (needle and needle not in name.casefold()):
                continue
            matches.append({"idx": int(idx), "name": name})
            if len(matches) >= limit:
                break
        return tuple(matches)

    def upsert_binding(
        self,
        workbook_path: Path,
        equipment_name: str,
        style_id: str,
        damage_multiplier: int,
        interval_seconds: int,
        priority: int,
        *,
        enabled: bool = True,
        note: str = "由装备光环页面更新",
    ) -> int:
        """原子更新或追加一条绑定，保留工作簿的其余格式和工作表。"""
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - 正式平台构建会打包依赖。
            raise EquipmentAuraError("当前平台缺少XLSX读取组件 openpyxl") from exc

        workbook_path = Path(workbook_path).resolve()
        if not workbook_path.is_file():
            raise EquipmentAuraError(f"装备光环配置表不存在: {workbook_path}")
        equipment_name = str(equipment_name).strip()
        if not equipment_name:
            raise EquipmentAuraError("装备名称不能为空")
        _style_definition(str(style_id).strip())
        multiplier = _strict_integer(damage_multiplier, "伤害倍率", 1, 999, 0)
        interval = _strict_integer(interval_seconds, "攻击间隔", 1, 60, 0)
        binding_priority = _strict_integer(priority, "优先级", 1, 999, 0)

        workbook = load_workbook(workbook_path)
        temporary_path: Path | None = None
        try:
            if "装备绑定" not in workbook.sheetnames:
                raise EquipmentAuraError("配置表缺少“装备绑定”工作表")
            sheet = workbook["装备绑定"]
            headers = tuple(cell.value for cell in sheet[1])
            if headers[: len(WORKBOOK_HEADERS)] != WORKBOOK_HEADERS:
                raise EquipmentAuraError(
                    f"装备绑定表头不匹配，必须为: {', '.join(WORKBOOK_HEADERS)}"
                )
            matching_rows = [
                row
                for row in range(2, sheet.max_row + 1)
                if str(sheet.cell(row, 2).value or "").strip().casefold() == equipment_name.casefold()
            ]
            if len(matching_rows) > 1:
                raise EquipmentAuraError(f"配置表装备名称不唯一: {equipment_name}")
            row_number = matching_rows[0] if matching_rows else sheet.max_row + 1
            values = (
                "是" if enabled else "否",
                equipment_name,
                str(style_id).strip(),
                multiplier,
                interval,
                binding_priority,
                str(note).strip(),
            )
            for column, value in enumerate(values, start=1):
                sheet.cell(row_number, column).value = value

            backup_root = self.platform_root / "backups" / "equipment-aura-workbook"
            backup_root.mkdir(parents=True, exist_ok=True)
            backup_name = (
                datetime.now().strftime("%Y%m%d_%H%M%S_")
                + _sha256_file(workbook_path)[:12]
                + "_"
                + uuid.uuid4().hex[:6]
                + ".xlsx"
            )
            shutil.copyfile(workbook_path, backup_root / backup_name)
            temporary_path = workbook_path.with_name(
                f".{workbook_path.name}.{uuid.uuid4().hex}.tmp.xlsx"
            )
            workbook.save(temporary_path)
        finally:
            workbook.close()
        if temporary_path is None or not temporary_path.is_file():
            raise EquipmentAuraError("配置表临时副本生成失败")
        try:
            # 保存后再用正式严格读取器回读，防止写入无效行。
            read_aura_workbook(temporary_path)
            os.replace(temporary_path, workbook_path)
        finally:
            if temporary_path.exists():
                temporary_path.unlink()
        return row_number

    def preflight(
        self,
        server: Path,
        client: Path,
        login: Path,
        workbook: Path,
    ) -> AuraInstallPlan:
        paths = AuraTargetPaths.from_roots(server, client, login)
        workbook_path = Path(workbook).resolve()
        parsed = read_aura_workbook(workbook_path)
        blockers: list[str] = []
        warnings = [
            "M2运行时仍可预检与部署；脚本和静态资源需由用户重载或重启后生效。"
        ]
        if not paths.server_root.is_dir():
            blockers.append(f"服务端根目录不存在: {paths.server_root}")
        if not paths.client_root.is_dir():
            blockers.append(f"客户端根目录不存在: {paths.client_root}")
        if not paths.login_root.is_dir():
            blockers.append(f"登录器补丁目录不存在: {paths.login_root}")
        for path in (paths.qfunction, paths.qmanage, paths.effect_list, paths.database):
            if not path.is_file():
                blockers.append(f"必需目标文件不存在: {path}")
        if blockers:
            return AuraInstallPlan(
                uuid.uuid4().hex,
                self.platform_root,
                paths,
                workbook_path,
                _sha256_file(workbook_path),
                parsed,
                -1,
                "blocked",
                tuple(blockers),
                tuple(warnings),
                (),
            )

        try:
            resolve_target_equipment(
                paths.database,
                tuple(binding.equipment_name for binding in parsed.enabled_bindings),
            )
        except EquipmentAuraError as exc:
            blockers.append(str(exc))

        qfunction_after, qfunction_blockers = _render_qfunction(paths.qfunction.read_bytes())
        qmanage_after, qmanage_blockers = _render_qmanage(paths.qmanage.read_bytes(), AuraSpec.round_moon())
        blockers.extend(qfunction_blockers)
        blockers.extend(qmanage_blockers)
        blockers.extend(_scan_external_timer_97(paths))

        effect_original = paths.effect_list.read_bytes()
        effect_after, resource_index = _effect_list_plan(effect_original, "XY_EquipmentAura_10.wzl")
        effect_names = [
            line.strip().lower()
            for line in _decode_script(effect_original).replace("\r\n", "\n").replace("\r", "\n").split("\n")
            if line.strip()
        ]
        if effect_names.count("xy_equipmentaura_10.wzl") > 1:
            blockers.append("EffectImageList资源重复: XY_EquipmentAura_10.wzl")

        runtime = compile_runtime_core(parsed.enabled_bindings, resource_index)
        core_after = _wrap_runtime_core(runtime)
        if paths.core_script.is_file() and paths.core_script.read_bytes() != core_after:
            current = paths.core_script.read_bytes()
            current_text = _decode_script(current)
            if "XY-EQUIPMENT-AURA-CORE10-BEGIN" in current_text:
                blockers.extend(_validate_existing_managed_blocks(current_text, "CORE10"))
            elif current not in _known_legacy_multi_style_core(effect_original):
                blockers.append(f"核心脚本不是可安全升级的平台版本: {paths.core_script}")

        client_action, client_blockers, client_files = _plan_static_pair(
            self.asset_wzl,
            self.asset_wzx,
            paths.client_wzl,
            paths.client_wzx,
            "客户端",
        )
        login_action, login_blockers, login_files = _plan_static_pair(
            self.asset_wzl,
            self.asset_wzx,
            paths.login_wzl,
            paths.login_wzx,
            "登录器",
        )
        blockers.extend(client_blockers)
        blockers.extend(login_blockers)
        static_action = "blocked" if blockers else (
            "install" if "install" in {client_action, login_action} else "reuse"
        )
        if blockers:
            return AuraInstallPlan(
                uuid.uuid4().hex,
                self.platform_root,
                paths,
                workbook_path,
                _sha256_file(workbook_path),
                parsed,
                resource_index,
                static_action,
                tuple(blockers),
                tuple(warnings),
                (),
            )

        proposed: list[tuple[Path, bytes]] = [
            (paths.qfunction, qfunction_after),
            (paths.qmanage, qmanage_after),
            (paths.effect_list, effect_after),
            (paths.core_script, core_after),
        ]
        proposed.extend(client_files)
        proposed.extend(login_files)
        files = tuple(
            change
            for path, after in proposed
            if (change := _changed_file(path, after)) is not None
        )
        return AuraInstallPlan(
            uuid.uuid4().hex,
            self.platform_root,
            paths,
            workbook_path,
            _sha256_file(workbook_path),
            parsed,
            resource_index,
            static_action,
            (),
            tuple(warnings),
            files,
        )

    def _atomic_replace(self, path: Path, data: bytes) -> None:
        _atomic_write(path, data)

    def install(self, plan: AuraInstallPlan) -> AuraInstallReceipt:
        if plan.blockers:
            raise EquipmentAuraError("预检存在阻止项: " + "；".join(plan.blockers))
        if _sha256_file(plan.workbook_path) != plan.workbook_hash:
            raise EquipmentAuraError(f"预检后文件已变化: {plan.workbook_path}")
        for change in plan.files:
            current_hash = _sha256_file(change.path) if change.path.is_file() else None
            if current_hash != change.before_hash:
                raise EquipmentAuraError(f"预检后文件已变化: {change.path}")
        if not plan.files:
            return AuraInstallReceipt(plan.plan_id, Path(), "already-current", ())

        transaction_id = datetime.now().strftime("%Y%m%d_%H%M%S_") + uuid.uuid4().hex[:8]
        backup_dir = self.platform_root / "backups" / "equipment-aura" / transaction_id
        files_dir = backup_dir / "files"
        staging_dir = backup_dir / ".staging"
        files_dir.mkdir(parents=True, exist_ok=False)
        staging_dir.mkdir()
        entries: list[dict[str, object]] = []
        try:
            for index, change in enumerate(plan.files):
                existed = change.path.is_file()
                backup = files_dir / f"{index:02d}.bin"
                staged = staging_dir / f"{index:02d}.bin"
                if existed:
                    shutil.copyfile(change.path, backup)
                _atomic_write(staged, change.after_bytes)
                if _sha256_file(staged) != change.after_hash:
                    raise EquipmentAuraError(f"临时副本验证失败: {change.path}")
                entries.append({
                    "path": str(change.path),
                    "existed_before": existed,
                    "before_hash": change.before_hash,
                    "after_hash": change.after_hash,
                    "backup": str(backup) if existed else None,
                })
            for change in plan.files:
                self._atomic_replace(change.path, change.after_bytes)
                if _sha256_file(change.path) != change.after_hash:
                    raise EquipmentAuraError(f"写入后哈希不符: {change.path}")
        except BaseException:
            for change, entry in reversed(list(zip(plan.files, entries))):
                if bool(entry["existed_before"]):
                    _atomic_write(change.path, Path(str(entry["backup"])).read_bytes())
                elif change.path.exists():
                    change.path.unlink()
            raise
        finally:
            if staging_dir.exists():
                shutil.rmtree(staging_dir)

        receipt_payload = {
            "schema_version": 1,
            "operation": "equipment-aura",
            "transaction_id": transaction_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "status": "installed-pending-game-verification",
            "server": str(plan.paths.server_root),
            "client": str(plan.paths.client_root),
            "login": str(plan.paths.login_root),
            "workbook": str(plan.workbook_path),
            "workbook_hash": plan.workbook_hash,
            "resource_index": plan.resource_index,
            "static_resource_action": plan.static_resource_action,
            "bindings": [asdict(binding) for binding in plan.workbook.enabled_bindings],
            "warnings": list(plan.warnings),
            "files": entries,
        }
        receipt_path = backup_dir / "receipt.json"
        _atomic_write(
            receipt_path,
            (json.dumps(receipt_payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8"),
        )
        log_path = self.platform_root / "logs" / "equipment-aura" / f"{transaction_id}.json"
        _atomic_write(log_path, receipt_path.read_bytes())
        return AuraInstallReceipt(
            transaction_id,
            receipt_path,
            "installed-pending-game-verification",
            plan.affected_paths,
        )

    def rollback(self, server: Path, transaction_id: str) -> str:
        receipt_path = (
            self.platform_root / "backups" / "equipment-aura" / transaction_id / "receipt.json"
        )
        if not receipt_path.is_file():
            raise EquipmentAuraError(f"光环事务收据不存在: {transaction_id}")
        payload = json.loads(receipt_path.read_text(encoding="utf-8"))
        if payload.get("operation") != "equipment-aura":
            raise EquipmentAuraError("收据不属于装备光环")
        if Path(str(payload.get("server", ""))).resolve() != Path(server).resolve():
            raise EquipmentAuraError("回滚目标服务端与收据不一致")
        entries = payload["files"]
        for entry in entries:
            path = Path(entry["path"])
            current = _sha256_file(path) if path.is_file() else None
            if current != entry["after_hash"]:
                raise EquipmentAuraError(f"安装后文件已被修改，禁止盲目回滚: {path}")
        for entry in reversed(entries):
            path = Path(entry["path"])
            if entry["existed_before"]:
                _atomic_write(path, Path(entry["backup"]).read_bytes())
                if _sha256_file(path) != entry["before_hash"]:
                    raise EquipmentAuraError(f"回滚后哈希不符: {path}")
            elif path.exists():
                path.unlink()
        payload["status"] = "rolled-back"
        payload["rolled_back_at"] = datetime.now().astimezone().isoformat()
        _atomic_write(
            receipt_path,
            (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8"),
        )
        return "rolled-back"
