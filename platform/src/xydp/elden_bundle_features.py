from __future__ import annotations

"""Filled V2.1 multi-instance NPC compilers.

The ordinary platform modules intentionally model one workbook per feature.
The Elden V2.1 delivery contains one shared 255-level seal chain, ten capped
seal NPCs, a 100-level title chain split over ten NPCs, six independent title
chains and a native 20-level rebirth chain split over four NPCs.  This module
adds only that multi-instance compilation layer; it reuses the platform's
validated cost and native-currency helpers instead of introducing a second
business data format.
"""

from dataclasses import asdict, dataclass, field
from decimal import Decimal
import hashlib
from pathlib import Path
from typing import Iterable
import re
import sqlite3
import tempfile

import openpyxl

from .encoding import TextDocument, encode_text_document, read_text_document
from .initial_camp import (
    _cost_lines,
    _cost_summary,
    _enable_setup_flags,
    _integer,
    _merge_named_description_lines,
    _script_token,
    _text,
    _title_item_values as _rebirth_title_item_values,
)
from .installer import InstallPlan, Installer, PlannedChange
from .item_synthesis import (
    SynthesisWorkbook,
    _compile_npc_script,
    read_item_synthesis_workbook,
)
from .repository import PackageRepository
from .seal_title import (
    ABILITY_IDS,
    SEAL_COLUMNS,
    SEAL_STATE_SCRIPT_PATH,
    SEAL_VARIABLE,
    TITLE_COLUMNS,
    _cost_group_summary,
    _payment_mode,
    _summary,
    _title_item_values,
    _title_payment_routes,
    _geometry,
    _map_file_from_info,
    _walkable,
)
from .sqlitepatch import SqlitePatchError, apply_sqlite_upsert
from .target import TargetInspector
from .textpatch import (
    TextPatchError,
    ensure_event_label,
    install_event_hook,
    install_managed_anchor_hook,
    remove_event_hook,
    remove_managed_anchor_hook,
    scan_labels,
)
from .weapon_enchant import (
    CORE_RELATIVE as WEAPON_ENCHANT_CORE_RELATIVE,
    RUNTIME_RELATIVE as WEAPON_ENCHANT_RUNTIME_RELATIVE,
    WarAshCost,
    load_weapon_enchant_workbook,
    render_elden_war_ash_core,
    render_elden_war_ash_runtime,
    upsert_weapon_enchant_qfunction_hook,
)


class EldenBundleFeatureError(ValueError):
    pass


PACKAGE_ID = "xy.bundle.elden-npc-v21"
PACKAGE_VERSION = "2.1.0-candidate.1"

# This exact self-contained title NPC came from the verified
# xy.lab.seal-title-v2 rollout and was later copied to the continent trees.
# That historical rollout left a stale embedded body digest.  Accept only its
# full-file SHA-256 so any subsequent byte change still blocks takeover.
LEGACY_ACCEPTED_NPC_SHA256 = {
    "8786037867d5c0a3bf58009bb35a99f9cef97b8bbc349c65d27155fc97909859",
    "d1c39a75056fe9fe6d6f654467bca90ae49f39840da207a4afb7d606bd25d35c",
    "8bb71e8b6fa85562c9612c7ba5d64352c4d21978293ba455fa0696b5383dcaa2",
}
LEGACY_SEAL_LOGIN_HOOK_HASHES = (
    "808f8e1bd9498bc4c4ac6e44a5cbddeff4258c831892b4dfbdff78d46998d6ab",
)
LEGACY_WEAPON_ENCHANT_HASHES = {
    WEAPON_ENCHANT_CORE_RELATIVE.as_posix(): {
        "34914bb225baf066b32bd2456ed72f6d71c1c6ff7f2fce1fe5e268727faadbc8",
    },
    WEAPON_ENCHANT_RUNTIME_RELATIVE.as_posix(): {
        "0cc11723fbbc4351004b230f61bbeaef0917c8dbe323957b9c3baa68e6e9e8ed",
    },
}
DB_RELATIVE = "Mud2/DB/ApexM2.DB"
MERCHANT_RELATIVE = "Mir200/Envir/MerChant.txt"
QFUNCTION_RELATIVE = "Mir200/Envir/Market_Def/QFunction-0.txt"
QMANAGE_RELATIVE = "Mir200/Envir/MapQuest_Def/QManage.txt"
ITEMDESC_RELATIVE = "Mir200/Envir/ItemDescList.txt"
SETUP_RELATIVE = "Mir200/!setup.txt"
SEAL_STATE_RELATIVE = "Mir200/Envir/QuestDiary/XY_System/XuanYuanHumanVar.txt"
ATTRIBUTE_PANEL_RELATIVE = (
    "Mir200/Envir/QuestDiary/玄渊功能/非常驻/属性总览/玄渊三属性按钮.txt"
)


@dataclass(frozen=True)
class NpcAssignment:
    feature_id: str
    name: str
    script_path: str
    map_code: str
    x: int
    y: int
    appearance: int
    reused: bool


@dataclass
class EldenBundleFeaturePlan:
    server: str
    bundle_root: str
    client: str | None = None
    launcher: str | None = None
    npcs: list[NpcAssignment] = field(default_factory=list)
    title_mapping: list[dict[str, object]] = field(default_factory=list)
    synthesis_recipes: list[dict[str, object]] = field(default_factory=list)
    service_costs: list[dict[str, object]] = field(default_factory=list)
    soul_tasks: list[dict[str, object]] = field(default_factory=list)
    changes: list[PlannedChange] = field(default_factory=list)
    blockers: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    install_plan: InstallPlan | None = None


def load_synthesis_workbooks(bundle_root: Path) -> tuple[tuple[str, SynthesisWorkbook], ...]:
    root = Path(bundle_root) / "34_通用物品合成_按NPC分表"
    paths = sorted(root.rglob("34_通用物品合成.xlsx"), key=lambda path: path.as_posix().casefold())
    result: list[tuple[str, SynthesisWorkbook]] = []
    seen_ids: set[str] = set()
    for order, path in enumerate(paths, start=1):
        workbook = read_item_synthesis_workbook(path, allow_auto_coordinates=True)
        folder = path.parent.name
        match = re.fullmatch(r"(C\d{2})_(.+)", folder)
        if not match:
            raise EldenBundleFeatureError(f"合成NPC目录名称无效：{folder}")
        npc_id = f"SYN_{match.group(1)}_{order:02d}"
        if npc_id.casefold() in seen_ids:
            raise EldenBundleFeatureError(f"合成NPC业务ID重复：{npc_id}")
        seen_ids.add(npc_id.casefold())
        result.append((npc_id, workbook))
    if len(result) != 16:
        raise EldenBundleFeatureError(f"合成NPC应为16个，实际{len(result)}个")
    recipe_count = sum(len(workbook.recipes) for _, workbook in result)
    if recipe_count != 38:
        raise EldenBundleFeatureError(f"合成配方应为38条，实际{recipe_count}条")
    return tuple(result)


def compile_combined_title_synthesis(title_body: str, synthesis_body: str) -> str:
    title = title_body.replace("[@Main]", "[@EL_IND_C01_TITLE]", 1).strip()
    synthesis = synthesis_body.replace("[@Main]", "[@EL_IND_C01_SYNTH]", 1).strip()
    return (
        "[@Main]\n#IF\n#SAY\n"
        "我可以为您见证称号，也能将收集到的物品合成为新的器物。\\\n"
        "<称号晋升/@EL_IND_C01_TITLE>　<物品合成/@EL_IND_C01_SYNTH>　<关闭/@exit>\n\n"
        f"{title}\n\n{synthesis}\n"
    )


def _resolve_synthesis_indices(
    database: bytes,
    workbooks: Iterable[tuple[str, SynthesisWorkbook]],
) -> dict[str, int]:
    names = sorted({
        name
        for _, workbook in workbooks
        for recipe in workbook.recipes
        for name in (
            recipe.output_name,
            *(item.name for item in recipe.inputs if item.input_type == "物品"),
        )
    })
    marks = ",".join("?" for _ in names)
    rows = _database_rows(
        database,
        f"SELECT Idx,Name FROM StdItems WHERE Name IN ({marks}) ORDER BY Idx",
        tuple(names),
    )
    by_name: dict[str, list[int]] = {}
    for idx, name in rows:
        by_name.setdefault(str(name), []).append(int(idx))
    errors: list[str] = []
    resolved: dict[str, int] = {}
    for name in names:
        matches = by_name.get(name, [])
        if not matches:
            errors.append(f"目标服物品不存在：{name}")
        elif len(matches) > 1:
            errors.append(f"目标服物品名称不唯一：{name}")
        else:
            resolved[name] = matches[0]
    if errors:
        raise EldenBundleFeatureError("\n".join(errors))
    return resolved


@dataclass(frozen=True)
class ServiceCost:
    operation_id: str
    continent_id: str
    map_code: str
    material_amount: int
    yuanbao: int


@dataclass(frozen=True)
class SoulTask:
    task_id: str
    continent_id: str
    map_group: str
    task_name: str
    map_code: str
    materials: tuple[tuple[str, int], ...]
    attack: int
    magic: int
    taoism: int
    toughness: int
    execution_percent: int
    variable: str
    existing_npc_name: str = ""
    existing_script_path: str = ""
    existing_title: str = ""


def _ledger_rows(bundle_root: Path, sheet_name: str) -> tuple[dict[str, object], ...]:
    path = Path(bundle_root) / "00_总控与审计" / "00_NPC经济总账_V2.1.xlsx"
    if not path.is_file():
        raise EldenBundleFeatureError(f"经济总账不存在：{path}")
    book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    try:
        if sheet_name not in book.sheetnames:
            raise EldenBundleFeatureError(f"经济总账缺少工作表：{sheet_name}")
        matrix = list(book[sheet_name].iter_rows(values_only=True))
    finally:
        book.close()
    if not matrix:
        raise EldenBundleFeatureError(f"经济总账工作表为空：{sheet_name}")
    headers = [_text(value) for value in matrix[0]]
    return tuple({
        header: values[index] if index < len(values) else None
        for index, header in enumerate(headers) if header
    } for values in matrix[1:] if any(value not in (None, "") for value in values))


def load_wash_open_costs(
    bundle_root: Path,
) -> tuple[tuple[ServiceCost, ...], tuple[ServiceCost, ...]]:
    groups: dict[str, list[ServiceCost]] = {"普通洗练": [], "开光": []}
    for row in _ledger_rows(bundle_root, "121操作映射"):
        category = _text(row.get("类别"))
        if category not in groups:
            continue
        operation_id = _script_token(row.get("ID"), "操作ID")
        match = re.fullmatch(r"(C\d{2})-(?:WASH|OPEN)", operation_id)
        if not match:
            raise EldenBundleFeatureError(f"洗练/开光操作ID无效：{operation_id}")
        packed = _text(row.get("单次材料"))
        material = re.fullmatch(r"失色锻造石[×xX*](\d+)", packed)
        if not material:
            raise EldenBundleFeatureError(f"{operation_id}材料必须且只能填写失色锻造石：{packed}")
        gold = int(_integer(row.get("单次金币"), 0) or 0)
        yuanbao = int(_integer(row.get("单次元宝"), 0) or 0)
        if gold != 0 or yuanbao <= 0:
            raise EldenBundleFeatureError(f"{operation_id}必须使用失色锻造石+原生元宝")
        groups[category].append(ServiceCost(
            operation_id,
            match.group(1),
            _script_token(row.get("地图"), "地图"),
            int(material.group(1)),
            yuanbao,
        ))
    for category, rows in groups.items():
        rows.sort(key=lambda row: row.continent_id)
        expected = [f"C{number:02d}" for number in range(2, 11)]
        actual = [row.continent_id for row in rows]
        if actual != expected:
            raise EldenBundleFeatureError(f"{category}大陆链应为{expected}，实际{actual}")
    return tuple(groups["普通洗练"]), tuple(groups["开光"])


_SOUL_TASK_LOCATIONS = {
    "C01-QUEST-1": ("XX113", "拾骸女·梅芙", "玄渊任务/漂流群岛/拾骸女_梅芙", "海岸拾魂者"),
    "C01-QUEST-2": ("XX329", "断誓骑士·罗恩", "玄渊任务/漂流群岛/断誓骑士_罗恩", "风暴断誓者"),
    "C01-QUEST-3": ("XX104", "无冠史官·塞蕾", "玄渊任务/漂流群岛/无冠史官_塞蕾", "沉王终结者"),
    "C02-QUEST-1": ("XX346", "", "", ""),
    "C02-QUEST-2": ("XX125", "", "", ""),
    "C02-QUEST-3": ("XX171", "", "", ""),
    "C02-QUEST-4": ("XX248", "", "", ""),
    "C02-QUEST-5": ("XX144", "", "", ""),
}


def _packed_amounts(value: object, field_name: str) -> tuple[tuple[str, int], ...]:
    text = _text(value)
    result: list[tuple[str, int]] = []
    for token in re.split(r"[；;]", text):
        token = token.strip()
        if not token:
            continue
        match = re.fullmatch(r"(.+?)[×xX*](\d+)", token)
        if not match:
            raise EldenBundleFeatureError(f"{field_name}格式无效：{token}")
        result.append((_script_token(match.group(1), field_name), int(match.group(2))))
    if not result:
        raise EldenBundleFeatureError(f"{field_name}不能为空")
    return tuple(result)


def load_soul_tasks(bundle_root: Path) -> tuple[SoulTask, ...]:
    result: list[SoulTask] = []
    for row in _ledger_rows(bundle_root, "任务总账"):
        task_id = _script_token(row.get("任务ID"), "任务ID")
        location = _SOUL_TASK_LOCATIONS.get(task_id)
        if location is None:
            raise EldenBundleFeatureError(f"任务总账出现未登记任务：{task_id}")
        continent = f"C{int(_integer(row.get('大陆')) or 0):02d}"
        reward = _text(row.get("永久奖励"))
        values = {"攻击": 0, "魔法": 0, "道术": 0, "韧性": 0, "处决概率": 0}
        for name, amount in _packed_amounts(reward, f"{task_id}永久奖励"):
            if name not in values:
                raise EldenBundleFeatureError(f"{task_id}不支持的永久奖励：{name}")
            values[name] += amount
        map_code, existing_name, existing_path, existing_title = location
        sequence = int(task_id.rsplit("-", 1)[-1])
        variable = f"XY_SOUL_{continent}_{sequence:02d}" if continent == "C02" else ""
        result.append(SoulTask(
            task_id,
            continent,
            _script_token(row.get("地图组"), "地图组"),
            _script_token(row.get("任务名称"), "任务名称"),
            map_code,
            _packed_amounts(row.get("材料"), f"{task_id}材料"),
            values["攻击"], values["魔法"], values["道术"], values["韧性"], values["处决概率"],
            variable, existing_name, existing_path, existing_title,
        ))
    result.sort(key=lambda task: task.task_id)
    if [task.task_id for task in result] != [
        "C01-QUEST-1", "C01-QUEST-2", "C01-QUEST-3",
        "C02-QUEST-1", "C02-QUEST-2", "C02-QUEST-3", "C02-QUEST-4", "C02-QUEST-5",
    ]:
        raise EldenBundleFeatureError("任务总账必须完整包含C01三条与C02五条魂魄任务")
    return tuple(result)


def load_war_ash_costs(bundle_root: Path) -> tuple[WarAshCost, WarAshCost, WarAshCost]:
    expected = ("C02-ASH-NORMAL", "C02-ASH-HIGH", "C02-ASH-OVERWRITE")
    found: dict[str, WarAshCost] = {}
    for row in _ledger_rows(bundle_root, "121操作映射"):
        operation_id = _text(row.get("ID"))
        if operation_id not in expected:
            continue
        if _text(row.get("地图")).casefold() != "xy_nmgf_main":
            raise EldenBundleFeatureError(f"{operation_id}必须位于XY_NMGF_MAIN")
        if _text(row.get("NPC名称")) != "战灰刻印师·盖尔":
            raise EldenBundleFeatureError(f"{operation_id}NPC名称与经济总账不一致")
        found[operation_id] = WarAshCost(
            operation_id,
            _packed_amounts(row.get("单次材料"), f"{operation_id}单次材料"),
            int(_integer(row.get("单次金币"), 0) or 0),
            int(_integer(row.get("单次元宝"), 0) or 0),
        )
    if tuple(found) != expected:
        raise EldenBundleFeatureError(f"战灰操作应为{list(expected)}，实际{list(found)}")
    normal, rare, clear = (found[item] for item in expected)
    if (normal.gold, normal.yuanbao) != (0, 15000):
        raise EldenBundleFeatureError("普通战灰必须使用原生元宝15000")
    if (rare.gold, rare.yuanbao) != (0, 50000):
        raise EldenBundleFeatureError("稀有战灰必须使用原生元宝50000")
    if (clear.gold, clear.yuanbao) != (100000, 0):
        raise EldenBundleFeatureError("清除战灰必须使用原生金币100000")
    return normal, rare, clear


def compile_soul_task(task: SoulTask) -> str:
    if not task.variable:
        raise EldenBundleFeatureError(f"{task.task_id}由既有已验收任务承载，不应重新编译")
    prefix = task.task_id.replace("-", "_")
    material_summary = "、".join(f"{name}×{amount}" for name, amount in task.materials)
    rewards: list[str] = []
    if task.toughness:
        rewards.append(f"韧性+{task.toughness}")
    if task.execution_percent:
        rewards.append(f"处决概率+{task.execution_percent}%")
    reward_summary = "、".join(rewards)
    lines = [
        "[@Main]", "#IF", f"CHECKVAR HUMAN {task.variable} = 1", "#ACT",
        f"GOTO @{prefix}_DONE", "BREAK", "#ELSEACT", f"GOTO @{prefix}_QUEST", "BREAK", "",
        f"[@{prefix}_QUEST]", "#IF", "#SAY", f"【{task.task_name}】\\",
        f"请带回：{material_summary}\\", f"永久奖励：{reward_summary}\\",
        f"<提交魂魄信物/@{prefix}_SUBMIT>　<关闭/@exit>", "",
        f"[@{prefix}_SUBMIT]", "#IF", f"CHECKVAR HUMAN {task.variable} = 0",
    ]
    lines.extend(f"CHECKITEM {name} {amount}" for name, amount in task.materials)
    lines.append("#ACT")
    lines.extend(f"TAKE {name} {amount}" for name, amount in task.materials)
    lines.extend([
        f"CALCVAR HUMAN {task.variable} = 1",
        f"SAVEVAR HUMAN {task.variable} {SEAL_STATE_SCRIPT_PATH}",
    ])
    if task.toughness:
        lines.append(f"INC N$XY_EXEC_Toughness {task.toughness}")
    lines.extend([
        "#CALL [\\玄渊功能\\非常驻\\属性总览\\玄渊三属性按钮.txt] @XY_UI_STATUS_REFRESH",
        f"SENDMSG 6 [{task.task_name}] 任务完成，永久获得{reward_summary}。",
        f"GOTO @{prefix}_DONE", "BREAK", "#ELSEACT",
        "MESSAGEBOX 所需魂魄材料不足，或该任务已经完成。", "BREAK", "",
        f"[@{prefix}_DONE]", "#IF", "#SAY", f"【{task.task_name}】\\",
        f"您已经完成本任务，永久奖励：{reward_summary}。\\", "<关闭/@exit>",
    ])
    return "\n".join(lines).rstrip() + "\n"


def compile_task_toughness_hook(tasks: Iterable[SoulTask]) -> str:
    lines: list[str] = []
    for task in tasks:
        if task.variable and task.toughness:
            lines += [
                "#IF", f"CHECKVAR HUMAN {task.variable} = 1", "#ACT",
                f"INC N$XY_EXEC_Toughness {task.toughness}",
            ]
    return "\n".join(lines)


def compile_task_execution_hook(tasks: Iterable[SoulTask], variable: str) -> str:
    lines: list[str] = []
    for task in tasks:
        if not task.variable or not task.execution_percent:
            continue
        amount = task.execution_percent * 100 if variable == "N$XY_EXEC_ChanceBP" else task.execution_percent
        lines += [
            "#IF", f"CHECKVAR HUMAN {task.variable} = 1", "#ACT", f"INC {variable} {amount}",
        ]
    return "\n".join(lines)


def _resolve_unique_items(database: bytes, names: Iterable[str]) -> dict[str, int]:
    unique = sorted(set(names))
    if not unique:
        return {}
    marks = ",".join("?" for _ in unique)
    rows = _database_rows(
        database,
        f"SELECT Idx,Name FROM StdItems WHERE Name IN ({marks}) ORDER BY Idx",
        tuple(unique),
    )
    grouped: dict[str, list[int]] = {}
    for idx, name in rows:
        grouped.setdefault(str(name), []).append(int(idx))
    errors: list[str] = []
    result: dict[str, int] = {}
    for name in unique:
        matches = grouped.get(name, [])
        if not matches:
            errors.append(f"目标服物品不存在：{name}")
        elif len(matches) > 1:
            errors.append(f"目标服物品名称不唯一：{name}")
        else:
            result[name] = matches[0]
    if errors:
        raise EldenBundleFeatureError("\n".join(errors))
    return result


def _validate_existing_c1_tasks(server: Path, database: bytes, qfunction_text: str, tasks: Iterable[SoulTask]) -> None:
    merchant = read_text_document(Path(server) / Path(MERCHANT_RELATIVE.replace("/", "\\"))).text
    rows = [line.split("\t") for line in merchant.splitlines()]
    for task in tasks:
        if task.continent_id != "C01":
            continue
        matches = [
            fields for fields in rows
            if len(fields) >= 7
            and fields[0].casefold() == task.existing_script_path.casefold()
            and fields[1].casefold() == task.map_code.casefold()
            and fields[4] == task.existing_npc_name
        ]
        if len(matches) != 1:
            raise EldenBundleFeatureError(f"既有魂魄任务NPC注册应唯一：{task.task_id}")
        script_path = Path(server) / "Mir200" / "Envir" / "Market_Def" / (
            f"{task.existing_script_path}-{task.map_code}.txt".replace("/", "\\")
        )
        if not script_path.is_file():
            raise EldenBundleFeatureError(f"既有魂魄任务脚本不存在：{script_path}")
        script = script_path.read_bytes().decode("gb18030")
        for name, amount in task.materials:
            if f"CHECKITEM {name} {amount}" not in script or f"TAKE {name} {amount}" not in script:
                raise EldenBundleFeatureError(f"既有魂魄任务材料与总账不一致：{task.task_id}/{name}")
        if f"CHECKFENGHAO {task.existing_title}" not in script or f"GIVEFENGHAO {task.existing_title} 1" not in script:
            raise EldenBundleFeatureError(f"既有魂魄任务完成状态与总账不一致：{task.task_id}")
        title_rows = _database_rows(
            database,
            "SELECT Dc,Dc2,Mc,Mc2,Sc,Sc2 FROM StdItems WHERE StdMode=70 AND Name=?",
            (task.existing_title,),
        )
        expected = (task.attack, task.attack, task.magic, task.magic, task.taoism, task.taoism)
        if len(title_rows) != 1 or tuple(int(value or 0) for value in title_rows[0]) != expected:
            raise EldenBundleFeatureError(f"既有魂魄任务称号属性与总账不一致：{task.task_id}")
        hook_pattern = re.compile(
            rf"CHECKFENGHAO {re.escape(task.existing_title)}\s+#ACT\s+INC N\$XY_EXEC_Toughness {task.toughness}",
            re.IGNORECASE,
        )
        if not hook_pattern.search(qfunction_text):
            raise EldenBundleFeatureError(f"既有魂魄任务韧性钩子与总账不一致：{task.task_id}")


def _managed_body_or_text(template: str) -> str:
    pattern = re.compile(
        rf"^; XYDP-FILE-BEGIN {re.escape(PACKAGE_ID)} ([A-Za-z0-9_.-]+) "
        r"SHA256=([0-9a-f]{64})\r?\n(.*?)"
        rf"^; XYDP-FILE-END {re.escape(PACKAGE_ID)} \1\s*$",
        re.MULTILINE | re.DOTALL,
    )
    match = pattern.search(template)
    if match is None:
        return template
    body = match.group(3).replace("\r\n", "\n").replace("\r", "\n").rstrip("\n") + "\n"
    if hashlib.sha256(body.encode("gb18030")).hexdigest() != match.group(2):
        raise EldenBundleFeatureError("已安装洗练/开光NPC受管正文被修改")
    return body


def compile_wash_variant(template: str, material_amount: int, yuanbao: int) -> str:
    text = _managed_body_or_text(template).replace("\r\n", "\n").replace("\r", "\n")
    summary = f"失色锻造石{material_amount}个和{yuanbao}元宝"
    text, display_count = re.subn(
        r"<&Text:每次消耗：[^:\r\n]+:382:108\{FCOLOR=243\}>",
        f"<&Text:每次消耗：{summary}:382:108{{FCOLOR=243}}>",
        text,
    )
    if display_count != 1:
        raise EldenBundleFeatureError("装备洗练底座的费用显示结构已变化")
    guard = (
        f"NOT CHECKITEM 失色锻造石 {material_amount}\n#ACT\n"
        f"MESSAGEBOX 洗练装备需要{summary}。\nBREAK\n#IF\n"
        f"NOT CHECKGAMEGOLD > {yuanbao - 1}"
    )
    text, guard_count = re.subn(
        r"(?:NOT CHECKITEM 失色锻造石 \d+\n#ACT\nMESSAGEBOX [^\n]+\nBREAK\n#IF\n)?"
        r"NOT CHECKGAMEGOLD > \d+",
        guard,
        text,
    )
    if guard_count != 2:
        raise EldenBundleFeatureError("装备洗练底座的两次资源复核结构已变化")
    text = re.sub(
        r"MESSAGEBOX 洗练装备需要[^\n，。]*(，本次未修改装备。|。)?",
        lambda match: f"MESSAGEBOX 洗练装备需要{summary}" + (match.group(1) or "。"),
        text,
    )
    text, action_count = re.subn(
        r"(?:TAKE 失色锻造石 \d+\n)?GAMEGOLD - \d+",
        f"TAKE 失色锻造石 {material_amount}\nGAMEGOLD - {yuanbao}",
        text,
    )
    if action_count != 1:
        raise EldenBundleFeatureError("装备洗练底座的最终扣费结构已变化")
    return text.rstrip("\n") + "\n"


def compile_opening_variant(template: str, material_amount: int, yuanbao: int) -> str:
    text = _managed_body_or_text(template).replace("\r\n", "\n").replace("\r", "\n")
    summary = f"失色锻造石{material_amount}个和{yuanbao}元宝"
    text, display_count = re.subn(
        r"<&Text:重复开光：[^:\r\n]+:382:152\{FCOLOR=161\}>",
        f"<&Text:重复开光：{summary}:382:152{{FCOLOR=161}}>",
        text,
    )
    if display_count != 1:
        raise EldenBundleFeatureError("装备开光底座的费用显示结构已变化")
    new_block = (
        "[@XYEW_REOPEN_CHECK]\n#if\n"
        f"CHECKITEM 失色锻造石 {material_amount}\n"
        f"CHECKGAMEGOLD > {yuanbao - 1}\n#act\n"
        "MOV N$XYEW_O_开光锁 1\n"
        f"TAKE 失色锻造石 {material_amount}\nGAMEGOLD - {yuanbao}\n"
        "goto @XYEW_REOPEN_CORE\nbreak\n#ELSEACT\n"
        f"messagebox 重复开光需要{summary}。\nbreak"
    )
    pattern = re.compile(
        r"\[@XYEW_REOPEN_CHECK\]\n.*?(?=\n\[@XYEW_REOPEN_CORE\])",
        re.DOTALL | re.IGNORECASE,
    )
    text, block_count = pattern.subn(new_block, text)
    if block_count != 1:
        raise EldenBundleFeatureError("装备开光底座的重复开光结算结构已变化")
    return text.rstrip("\n") + "\n"


def _title_token(value: object, field_name: str) -> str:
    """Validate a native title name while retaining the common ``[称号]`` prefix."""

    token = _text(value)
    if not token:
        raise EldenBundleFeatureError(f"{field_name}不能为空")
    if re.search(r"[\s<>\\/{};]", token):
        raise EldenBundleFeatureError(f"{field_name}含脚本不支持的字符：{token}")
    return token


@dataclass(frozen=True)
class ChainNpc:
    chain_id: str
    npc_name: str
    map_code: str
    appearance: int
    rows: tuple[dict[str, object], ...]


@dataclass(frozen=True)
class V21Progression:
    seal_rows: tuple[dict[str, object], ...]
    seal_npcs: tuple[ChainNpc, ...]
    main_title_rows: tuple[dict[str, object], ...]
    main_titles: tuple[ChainNpc, ...]
    independent_titles: tuple[ChainNpc, ...]
    rebirth_rows: tuple[dict[str, object], ...]
    rebirths: tuple[ChainNpc, ...]


def _enabled_state(value: object) -> bool:
    return _text(value).casefold() in {"可安装", "待配置", "是", "启用", "1", "true", "yes"}


def load_config_rows(path: Path, sheet_name: str = "配置") -> tuple[dict[str, object], ...]:
    """Read a filled candidate workbook without mutating its pending status.

    In this delivery ``待配置`` means that target-dependent title numbers and
    coordinates still need preflight allocation.  Rows must still be fully
    populated and are compiled only after that allocation succeeds.
    """

    source = Path(path)
    if not source.is_file():
        raise EldenBundleFeatureError(f"工作簿不存在：{source}")
    book = openpyxl.load_workbook(source, read_only=False, data_only=True)
    try:
        if sheet_name not in book.sheetnames:
            raise EldenBundleFeatureError(f"{source.name}缺少工作表：{sheet_name}")
        sheet = book[sheet_name]
        matrix = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    header_index = -1
    headers: list[str] = []
    for index, row in enumerate(matrix[:20]):
        candidate = [_text(value) for value in row]
        if "档位/等级" in candidate and "状态" in candidate:
            header_index = index
            headers = candidate
            break
    if header_index < 0:
        raise EldenBundleFeatureError(f"{source.name}/{sheet_name}未找到配置表头")
    result: list[dict[str, object]] = []
    for values in matrix[header_index + 1:]:
        row = {
            header: values[column] if column < len(values) else None
            for column, header in enumerate(headers) if header
        }
        if not any(value not in (None, "") for value in row.values()):
            continue
        if not _enabled_state(row.get("状态")):
            continue
        # The economic ledger keeps the native currency name on several rows
        # whose amount is deliberately zero.  Zero means that route is absent,
        # not a malformed name/count pair.  Normalize only the in-memory
        # candidate row; the user's read-only workbook remains untouched.
        for suffix in ("A", "B"):
            for name_prefix, count_prefix in (("材料", "材料"), ("货币", "货币")):
                name_key = f"{name_prefix}{suffix}"
                count_key = f"{count_prefix}{suffix}数量"
                amount = _integer(row.get(count_key), 0) or 0
                if amount == 0:
                    row[name_key] = None
                    row[count_key] = None
        result.append(row)
    if not result:
        raise EldenBundleFeatureError(f"{source.name}/{sheet_name}没有可编译的已填写行")
    return tuple(result)


def _one_metadata(rows: Iterable[dict[str, object]], key: str, source: str) -> object:
    values = {_text(row.get(key)) for row in rows if _text(row.get(key))}
    if len(values) != 1:
        raise EldenBundleFeatureError(f"{source}的{key}必须全部一致，当前：{sorted(values)}")
    return next(iter(values))


def _chain(path: Path, chain_id: str) -> ChainNpc:
    rows = load_config_rows(path)
    npc_name = _script_token(_one_metadata(rows, "NPC名称", path.name), "NPC名称")
    map_code = _script_token(_one_metadata(rows, "地图", path.name), "地图")
    appearances = {_integer(row.get("外观")) for row in rows if row.get("外观") not in (None, "")}
    if len(appearances) != 1 or next(iter(appearances)) is None or int(next(iter(appearances))) < 0:
        raise EldenBundleFeatureError(f"{path.name}的外观必须全部一致且为非负整数")
    return ChainNpc(chain_id, npc_name, map_code, int(next(iter(appearances))), rows)


def _levels(rows: Iterable[dict[str, object]], source: str) -> list[int]:
    result: list[int] = []
    for row in rows:
        value = _integer(row.get("档位/等级"))
        if value is None:
            raise EldenBundleFeatureError(f"{source}存在空档位")
        result.append(value)
    if result != list(range(result[0], result[0] + len(result))):
        raise EldenBundleFeatureError(f"{source}档位不连续：{result[:3]}…{result[-3:]}")
    return result


def _validate_seal_rows(rows: tuple[dict[str, object], ...], expected: int, source: str) -> None:
    levels = _levels(rows, source)
    if levels != list(range(1, expected + 1)):
        raise EldenBundleFeatureError(f"{source}必须严格为1至{expected}档")
    previous = {column: 0 for column in SEAL_COLUMNS}
    for level, row in enumerate(rows, 1):
        for column in SEAL_COLUMNS:
            value = _integer(row.get(column))
            if value is None or value < previous[column]:
                raise EldenBundleFeatureError(f"{source}第{level}档{column}不是非递减累计整数")
            previous[column] = value
        checks, _actions = _cost_lines(row)
        if not checks:
            raise EldenBundleFeatureError(f"{source}第{level}档没有材料或货币")


def _validate_title_rows(rows: tuple[dict[str, object], ...], source: str) -> None:
    _levels(rows, source)
    names: list[str] = []
    previous = {column: 0 for column in TITLE_COLUMNS}
    for index, row in enumerate(rows, 1):
        name = _title_token(row.get("称号名称"), f"{source}第{index}档称号名称")
        names.append(name)
        for column in TITLE_COLUMNS:
            value = _integer(row.get(column))
            if value is None or value < 0:
                raise EldenBundleFeatureError(f"{source}第{index}档{column}必须为非负整数")
            # Main-title totals span files; independent chains are checked by
            # their own caller.  Within any supplied chain totals may not drop.
            if value < previous[column]:
                raise EldenBundleFeatureError(f"{source}第{index}档{column}低于本链上一档")
            previous[column] = value
        if not _title_payment_routes(row):
            raise EldenBundleFeatureError(f"{source}第{index}档没有支付路线")
    if len(names) != len(set(names)):
        raise EldenBundleFeatureError(f"{source}称号名称重复")


def load_v21_progression(root: Path) -> V21Progression:
    base = Path(root)
    global_seal = load_config_rows(base / "00_全局系统" / "21_神印基础属性.xlsx")
    _validate_seal_rows(global_seal, 255, "全局神印")
    global_spec = _chain(base / "00_全局系统" / "21_神印基础属性.xlsx", "SEAL_GLOBAL")
    catchups = tuple(
        _chain(path, f"SEAL_{path.parent.name.split('_', 1)[0]}")
        for path in sorted((base / "21_神印_每大陆追赶NPC").rglob("21_神印基础属性.xlsx"))
    )
    expected_caps = [50, 80, 110, 140, 170, 195, 220, 240, 255]
    if [len(item.rows) for item in catchups] != expected_caps:
        raise EldenBundleFeatureError(
            f"神印追赶NPC上限应为{expected_caps}，实际{[len(item.rows) for item in catchups]}"
        )
    comparable = (
        "档位/等级", "攻击", "魔法", "道术", "HP", "材料A", "材料A数量",
        "材料B", "材料B数量", "货币A", "货币A数量", "货币B", "货币B数量",
    )
    for item in catchups:
        _validate_seal_rows(item.rows, len(item.rows), item.chain_id)
        for expected, actual in zip(global_seal, item.rows):
            if tuple(expected.get(key) for key in comparable) != tuple(actual.get(key) for key in comparable):
                raise EldenBundleFeatureError(f"{item.chain_id}不是全局255档神印的完整前缀")

    main_chains = tuple(
        _chain(path, f"MAIN_{path.parent.name.split('_', 1)[0]}")
        for path in sorted((base / "22_贯穿称号_每大陆一个NPC").rglob("22_称号晋升.xlsx"))
    )
    main_rows = tuple(row for chain in main_chains for row in chain.rows)
    if len(main_chains) != 10 or _levels(main_rows, "贯穿称号") != list(range(1, 101)):
        raise EldenBundleFeatureError("贯穿称号必须是10个NPC、全局1至100级")
    # Validate the 100-level chain as one chain, not as ten isolated sheets.
    _validate_title_rows(main_rows, "贯穿称号")

    independent = tuple(
        _chain(path, f"INDEPENDENT_{path.parent.name.split('_', 1)[0]}")
        for path in sorted((base / "22_独立称号").rglob("22_称号晋升.xlsx"))
    )
    if len(independent) != 6 or any(len(item.rows) != 2 for item in independent):
        raise EldenBundleFeatureError("独立称号必须是6条二级链")
    for item in independent:
        _validate_title_rows(item.rows, item.chain_id)

    rebirth_chains = tuple(
        _chain(path, f"REBIRTH_{path.parent.name.split('_', 1)[0]}")
        for path in sorted((base / "04_转生_按大陆分表").rglob("04_转生.xlsx"))
    )
    rebirth_rows = tuple(row for chain in rebirth_chains for row in chain.rows)
    if len(rebirth_chains) != 4 or _levels(rebirth_rows, "转生") != list(range(1, 21)):
        raise EldenBundleFeatureError("转生必须是4个NPC、全局1至20转")
    for index, row in enumerate(rebirth_rows, 1):
        checks, _actions = _cost_lines(row)
        if not checks:
            raise EldenBundleFeatureError(f"转生第{index}档没有材料或货币")
        raw = row.get("神力增加")
        value = Decimal(str(raw)) * Decimal("100") if raw not in (None, "") else Decimal("0")
        if value <= 0 or value != value.to_integral_value():
            raise EldenBundleFeatureError(f"转生第{index}档神力增加无法换算为整数百分比")

    return V21Progression(
        global_seal,
        (global_spec, *catchups),
        main_rows,
        main_chains,
        independent,
        rebirth_rows,
        rebirth_chains,
    )


def _seal_totals(rows: tuple[dict[str, object], ...]) -> list[dict[str, int]]:
    return [
        {column: int(_integer(row.get(column)) or 0) for column in SEAL_COLUMNS}
        for row in rows
    ]


def compile_seal_npc(
    all_rows: tuple[dict[str, object], ...],
    cap: int,
    npc_name: str,
) -> str:
    if not 1 <= cap <= len(all_rows):
        raise EldenBundleFeatureError(f"{npc_name}神印上限无效：{cap}")
    values = _seal_totals(all_rows)
    lines = ["[@Main]"]
    lines += [
        "#IF", f"CHECKVAR HUMAN {SEAL_VARIABLE} < 0", "#ACT", "GOTO @EL_SEAL_INVALID", "BREAK", "",
        "#IF", f"CHECKVAR HUMAN {SEAL_VARIABLE} > {cap - 1}", "#ACT", "GOTO @EL_SEAL_CAP", "BREAK", "",
    ]
    route_size = 20
    for start in range(0, cap, route_size):
        end = min(cap - 1, start + route_size - 1)
        lines += [
            "#IF", f"CHECKVAR HUMAN {SEAL_VARIABLE} > {start - 1}",
            f"CHECKVAR HUMAN {SEAL_VARIABLE} < {end + 1}", "#ACT",
            f"GOTO @EL_SEAL_ROUTE_{start}_{end}", "BREAK", "",
        ]
    lines += ["#IF", "#ACT", "GOTO @EL_SEAL_INVALID", "BREAK", ""]

    for start in range(0, cap, route_size):
        end = min(cap - 1, start + route_size - 1)
        lines.append(f"[@EL_SEAL_ROUTE_{start}_{end}]")
        for current in range(start, end + 1):
            lines += [
                "#IF", f"CHECKVAR HUMAN {SEAL_VARIABLE} = {current}", "#ACT",
                f"GOTO @EL_SEAL_VIEW_{current + 1}", "BREAK", "",
            ]
        lines += ["GOTO @EL_SEAL_INVALID", "BREAK", ""]

    previous = {column: 0 for column in SEAL_COLUMNS}
    for index in range(1, cap + 1):
        row = all_rows[index - 1]
        total = values[index - 1]
        checks, actions = _cost_lines(row)
        deltas = {column: total[column] - previous[column] for column in SEAL_COLUMNS}
        current_summary = "未修行" if index == 1 else _summary(previous)
        lines += [
            f"[@EL_SEAL_VIEW_{index}]", "#IF", "#SAY", "< /FCOLOR=250>\\",
            f"<> <{npc_name}：/FCOLOR=158> <只显示当前可进行的下一档/FCOLOR=218>\\",
            "<> <---------------------------------------------------------------/FCOLOR=10>\\",
            f"<> <当前阶段:/FCOLOR=161>{{{index - 1}档/SCOLOR=249}}\\",
            f"<> <当前总值:/FCOLOR=161>{{{current_summary}/SCOLOR=249}}\\",
            f"<> <下一阶段:/FCOLOR=161>{{{index}档/SCOLOR=253}}\\",
            f"<> <下一总值:/FCOLOR=161>{{{_summary(total)}/SCOLOR=147}}\\",
            f"<> <所需资源:/FCOLOR=161>{{{_cost_summary(row)}/SCOLOR=251}}\\",
            "<> <---------------------------------------------------------------/FCOLOR=10>\\",
            f"<> 【<确认修行/@EL_SEAL_APPLY_{index}>】　<关闭/@exit>\\", "",
            f"[@EL_SEAL_APPLY_{index}]", "#IF", f"CHECKVAR HUMAN {SEAL_VARIABLE} = {index - 1}", *checks,
            "#ACT", f"CALCVAR HUMAN {SEAL_VARIABLE} = {index}",
            f"SAVEVAR HUMAN {SEAL_VARIABLE} {SEAL_STATE_SCRIPT_PATH}",
        ]
        for column in SEAL_COLUMNS:
            for ability_id in ABILITY_IDS[column]:
                if deltas[column]:
                    lines.append(f"ChangeHumAbilityEX {ability_id} + {deltas[column]}")
        lines += [
            *actions,
            f"SENDMSG 6 [神印修行] 已提升至{index}档。",
            "GOTO @Main", "BREAK", "#ELSEACT",
            "MESSAGEBOX 修行所需的物品或货币不足，请备齐后再来。", "BREAK", "",
        ]
        previous = total
    lines += [
        "[@EL_SEAL_CAP]", "#IF", "#SAY", "< /FCOLOR=250>\\",
        f"<> <{npc_name}：/FCOLOR=158> <本处可修行的最高阶段为{cap}档/FCOLOR=218>\\",
        f"<> <当前已达到或超过{cap}档，请前往后续大陆继续修行。/FCOLOR=251>\\",
        "<> <关闭/@exit>\\", "",
        "[@EL_SEAL_INVALID]", "#IF", "#SAY",
        "修行记录异常，请联系游戏管理员。\\", "<关闭/@exit>",
    ]
    return "\n".join(lines).rstrip() + "\n"


def compile_title_chain(
    rows: tuple[dict[str, object], ...],
    all_chain_titles: list[str],
    prefix: str,
    npc_name: str,
) -> str:
    levels = _levels(rows, npc_name)
    start, end = levels[0], levels[-1]
    if end > len(all_chain_titles):
        raise EldenBundleFeatureError(f"{npc_name}称号范围超出全链")
    local_titles = [_title_token(row.get("称号名称"), "称号名称") for row in rows]
    if local_titles != all_chain_titles[start - 1:end]:
        raise EldenBundleFeatureError(f"{npc_name}称号不是全链{start}至{end}级")
    lines = ["[@Main]"]
    for global_level in range(len(all_chain_titles), end, -1):
        lines += [
            "#IF", f"CHECKFENGHAO {all_chain_titles[global_level - 1]}", "#ACT",
            f"GOTO @{prefix}_PASSED", "BREAK", "",
        ]
    for global_level in range(end, start - 1, -1):
        target = f"@{prefix}_FULL" if global_level == end else f"@{prefix}_VIEW_{global_level + 1}"
        lines += [
            "#IF", f"CHECKFENGHAO {all_chain_titles[global_level - 1]}", "#ACT",
            f"GOTO {target}", "BREAK", "",
        ]
    if start == 1:
        lines += ["#IF", "#ACT", f"GOTO @{prefix}_VIEW_1", "BREAK", ""]
    else:
        lines += [
            "#IF", f"CHECKFENGHAO {all_chain_titles[start - 2]}", "#ACT",
            f"GOTO @{prefix}_VIEW_{start}", "BREAK", "",
            "#IF", "#ACT", f"GOTO @{prefix}_PREREQUISITE", "BREAK", "",
        ]

    for row, level, title in zip(rows, levels, local_titles):
        previous = all_chain_titles[level - 2] if level > 1 else ""
        payment_mode = _payment_mode(row)
        payment_routes = _title_payment_routes(row)
        total = {column: int(_integer(row.get(column)) or 0) for column in TITLE_COLUMNS}
        material_summary = _cost_group_summary(row, (("材料A", "材料A数量"), ("材料B", "材料B数量")))
        currency_summary = _cost_group_summary(row, (("货币A", "货币A数量"), ("货币B", "货币B数量")))
        currency_a = _cost_group_summary(row, (("货币A", "货币A数量"),))
        currency_b = _cost_group_summary(row, (("货币B", "货币B数量"),))
        lines += [
            f"[@{prefix}_VIEW_{level}]", "#IF", "#SAY", "< /FCOLOR=250>\\",
            f"<> <{npc_name}：/FCOLOR=158> <只显示当前可晋升的下一级/FCOLOR=218>\\",
            f"<> <当前称号:/FCOLOR=161>{{{previous or '无'}/SCOLOR=249}}\\",
            f"<> <下一级称号:/FCOLOR=161>{{{level}级（{title}）/SCOLOR=253}}\\",
            f"<> <需要材料:/FCOLOR=161>{{{material_summary}/SCOLOR=251}}\\",
        ]
        if payment_mode in {"二选一", "二选一B免材料"}:
            mode_text = "货币A需材料；货币B免材料" if payment_mode == "二选一B免材料" else "二选一，任选一种"
            lines += [
                f"<> <支付方式:/FCOLOR=161>{{{mode_text}/SCOLOR=253}}\\",
                f"<> <货币A:/FCOLOR=161>{{{currency_a}/SCOLOR=251}}　<货币B:/FCOLOR=161>{{{currency_b}/SCOLOR=251}}\\",
            ]
        else:
            lines.append(f"<> <需要货币:/FCOLOR=161>{{{currency_summary}/SCOLOR=251}}\\")
        lines += [
            f"<> <完整属性:/FCOLOR=161>{{{_summary(total, include_mp=True, include_drop=True)}/SCOLOR=147}}\\",
        ]
        buttons = "　".join(
            f"【<{label}/@{prefix}_APPLY_{level}{'_' + code if code else ''}>】"
            for code, label, _checks, _actions in payment_routes
        )
        lines += [f"<> {buttons}　<关闭/@exit>\\", ""]
        for route_code, _label, checks, actions in payment_routes:
            suffix = f"_{route_code}" if route_code else ""
            lines += [f"[@{prefix}_APPLY_{level}{suffix}]", "#IF"]
            if previous:
                lines.append(f"CHECKFENGHAO {previous}")
            for current_or_higher in all_chain_titles[level - 1:]:
                lines.append(f"NOT CHECKFENGHAO {current_or_higher}")
            lines += [
                *checks, "#ACT", f"GIVEFENGHAO {title} 1",
                f"GOTO @{prefix}_CONFIRM_{level}{suffix}", "BREAK", "#ELSEACT",
                "MESSAGEBOX 晋升所需的物品或货币不足，请备齐后再来。", "BREAK", "",
                f"[@{prefix}_CONFIRM_{level}{suffix}]", "#IF", f"CHECKFENGHAO {title}",
            ]
            if previous:
                lines.append(f"CHECKFENGHAO {previous}")
            lines += [*checks, "#ACT", *actions]
            if previous:
                lines.append(f"RECYCFENGHAO {previous}")
            lines += [
                f"SENDMSG 6 [称号晋升] 已晋升为{title}。", "GOTO @Main", "BREAK",
                "#ELSEACT", f"RECYCFENGHAO {title}",
                "MESSAGEBOX 称号晋升未能完成，请稍后再试；本次未消耗物品和货币。", "BREAK", "",
            ]
    final_values = {column: int(_integer(rows[-1].get(column)) or 0) for column in TITLE_COLUMNS}
    lines += [
        f"[@{prefix}_FULL]", "#IF", "#SAY", "< /FCOLOR=250>\\",
        f"<> <{npc_name}：/FCOLOR=158> <当前大陆的称号已经全部完成/FCOLOR=218>\\",
        f"<> <当前称号:/FCOLOR=161>{{{local_titles[-1]}/SCOLOR=253}}\\",
        f"<> <完整属性:/FCOLOR=161>{{{_summary(final_values, include_mp=True, include_drop=True)}/SCOLOR=147}}\\",
        "<> <关闭/@exit>\\", "",
        f"[@{prefix}_PASSED]", "#IF", "#SAY",
        f"您的贯穿称号已经超过本大陆{end}级，无需在此再次晋升。\\", "<关闭/@exit>", "",
        f"[@{prefix}_PREREQUISITE]", "#IF", "#SAY",
        f"请先在上一大陆完成{start - 1}级称号“{all_chain_titles[start - 2]}”，再来继续晋升。\\", "<关闭/@exit>",
    ]
    return "\n".join(lines).rstrip() + "\n"


def _rebirth_cumulative(rows: tuple[dict[str, object], ...]) -> list[int]:
    result: list[int] = []
    total = Decimal("0")
    for index, row in enumerate(rows, 1):
        raw = row.get("神力增加")
        increment = Decimal(str(raw)) * Decimal("100") if raw not in (None, "") else Decimal("0")
        if increment <= 0 or increment != increment.to_integral_value():
            raise EldenBundleFeatureError(f"转生第{index}档神力增加无效")
        total += increment
        result.append(int(total))
    return result


def compile_rebirth_power_hook(rows: tuple[dict[str, object], ...], variable: str) -> str:
    values = _rebirth_cumulative(rows)
    lines = [
        "; 转生神力：以翎风引擎原生转生等级为唯一权威，称号只作显示镜像",
        "#IF", "#ACT", "MOV N$XY_REBIRTH_PowerBonus 0",
    ]
    for level, value in enumerate(values, 1):
        lines += ["#IF"]
        lines += [f"CHECKRENEWLEVEL > {level - 1}"] if level == len(values) else [f"CHECKRENEWLEVEL = {level}"]
        lines += ["#ACT", f"MOV N$XY_REBIRTH_PowerBonus {value}"]
    lines += [
        "#IF", "LARGE N$XY_REBIRTH_PowerBonus 0", "#ACT",
        f"INC {variable} <$STR(N$XY_REBIRTH_PowerBonus)>",
    ]
    return "\n".join(lines)


def compile_rebirth_segment(
    rows: tuple[dict[str, object], ...],
    all_titles: list[str],
    prefix: str,
    npc_name: str,
) -> str:
    levels = _levels(rows, npc_name)
    start, end = levels[0], levels[-1]
    if end > len(all_titles) or len(rows) != 5:
        raise EldenBundleFeatureError(f"{npc_name}转生范围无效")
    local_titles = [_title_token(row.get("称号名称"), "转生称号") for row in rows]
    lines = ["[@Main]"]
    lines += [
        "#IF", f"CHECKRENEWLEVEL < {start - 1}", "#ACT", f"GOTO @{prefix}_PREREQUISITE", "BREAK", "",
        "#IF", f"CHECKRENEWLEVEL > {end}", "#ACT", f"GOTO @{prefix}_PASSED", "BREAK", "",
        "#IF", f"CHECKRENEWLEVEL = {end}", "#ACT", f"GOTO @{prefix}_FULL", "BREAK", "",
    ]
    for current in range(start - 1, end):
        lines += [
            "#IF", f"CHECKRENEWLEVEL = {current}", "#ACT",
            f"GOTO @{prefix}_VIEW_{current + 1}", "BREAK", "",
        ]
    lines += ["MESSAGEBOX 当前转生状态异常，请联系管理员处理。", "BREAK", ""]

    for row, level, title in zip(rows, levels, local_titles):
        checks, actions = _cost_lines(row)
        previous_level = level - 1
        summary = _text(row.get("属性与数值"))
        lines += [
            f"[@{prefix}_VIEW_{level}]", "#IF", "#SAY", "< /FCOLOR=250>\\",
            f"<> <{npc_name}：/FCOLOR=158> <以引擎原生转生等级为准/FCOLOR=218>\\",
            f"<> <当前转生:/FCOLOR=161>{{{previous_level}转/SCOLOR=249}}\\",
            f"<> <下一转生:/FCOLOR=161>{{{level}转（{title}）/SCOLOR=253}}\\",
            f"<> <所需资源:/FCOLOR=161>{{{_cost_summary(row)}/SCOLOR=251}}\\",
            f"<> <转生效果:/FCOLOR=161>{{{summary}/SCOLOR=147}}\\",
            "<> <转生后人物等级不变，当前经验将清空。/FCOLOR=245>\\",
            f"<> 【<确认转生/@{prefix}_APPLY_{level}>】　<关闭/@exit>\\", "",
            f"[@{prefix}_APPLY_{level}]", "#IF", f"CHECKRENEWLEVEL = {previous_level}", *checks,
            "#ACT", f"GIVEFENGHAO {title} 1", f"GOTO @{prefix}_TITLE_CONFIRM_{level}", "BREAK",
            "#ELSEACT", "MESSAGEBOX 当前转生阶段不符，或所需物品、货币不足。", "BREAK", "",
            f"[@{prefix}_TITLE_CONFIRM_{level}]", "#IF", f"CHECKRENEWLEVEL = {previous_level}", f"CHECKFENGHAO {title}",
            *checks, "#ACT", "RENEWLEVEL 1", f"GOTO @{prefix}_ENGINE_CONFIRM_{level}", "BREAK",
            "#ELSEACT", f"RECYCFENGHAO {title}", "MESSAGEBOX 转生未能完成，本次未消耗物品和货币。", "BREAK", "",
            f"[@{prefix}_ENGINE_CONFIRM_{level}]", "#IF", f"CHECKRENEWLEVEL = {level}", f"CHECKFENGHAO {title}",
            *checks, "#ACT", *actions,
        ]
        for other in all_titles:
            if other != title:
                lines.append(f"RECYCFENGHAO {other}")
        lines += [
            f"SENDMSG 6 [转生] 恭喜您晋升至{level}转，获得称号{title}。", "GOTO @Main", "BREAK",
            "#ELSEACT", f"RECYCFENGHAO {title}", "MESSAGEBOX 转生未能完成，本次未扣除物品和货币。", "BREAK", "",
        ]
    lines += [
        f"[@{prefix}_FULL]", "#IF", "#SAY", f"您已完成本处开放的{start}至{end}转。\\", "<关闭/@exit>", "",
        f"[@{prefix}_PREREQUISITE]", "#IF", "#SAY", f"请先在上一大陆完成{start - 1}转，再来继续转生。\\", "<关闭/@exit>", "",
        f"[@{prefix}_PASSED]", "#IF", "#SAY", f"您的转生等级已经超过本处开放的{end}转，请前往后续大陆。\\", "<关闭/@exit>",
    ]
    return "\n".join(lines).rstrip() + "\n"


def _managed_npc(kind: str, body: str) -> bytes:
    normalized = body.replace("\r\n", "\n").rstrip("\n") + "\n"
    digest = hashlib.sha256(normalized.encode("gb18030")).hexdigest()
    text = (
        f"; XYDP-FILE-BEGIN {PACKAGE_ID} {kind} SHA256={digest}\n"
        f"{normalized}"
        f"; XYDP-FILE-END {PACKAGE_ID} {kind}\n"
    )
    return text.replace("\n", "\r\n").encode("gb18030")


def _exclusive_script_bytes(text: str) -> bytes:
    return (text.replace("\r\n", "\n").replace("\r", "\n").rstrip("\n") + "\n").replace(
        "\n", "\r\n"
    ).encode("gb18030")


def _validate_legacy_exclusive_script(relative: str, before: bytes | None, desired: bytes) -> None:
    if before is None or before == desired:
        return
    accepted = LEGACY_WEAPON_ENCHANT_HASHES.get(relative, set())
    if hashlib.sha256(before).hexdigest() not in accepted:
        raise EldenBundleFeatureError(f"已有武器附魔底座发生未授权漂移：{relative}")


def _validate_existing_npc(data: bytes | None, accepted: set[tuple[str, str]]) -> None:
    if data is None:
        return
    if hashlib.sha256(data).hexdigest() in LEGACY_ACCEPTED_NPC_SHA256:
        return
    try:
        text = data.decode("gb18030")
    except UnicodeError as exc:
        raise EldenBundleFeatureError("现有NPC脚本不是GB18030，拒绝接管") from exc
    pattern = re.compile(
        r"^; XYDP-FILE-BEGIN ([A-Za-z0-9_.-]+) ([A-Za-z0-9_.-]+) SHA256=([0-9a-f]{64})\r?\n"
        r"(.*?)^; XYDP-FILE-END \1 \2\s*$",
        re.MULTILINE | re.DOTALL,
    )
    match = pattern.search(text)
    if match is None:
        raise EldenBundleFeatureError("目标NPC脚本不是可验证的平台受管文件，拒绝覆盖")
    identity = (match.group(1), match.group(2))
    if identity not in accepted and identity[0] != PACKAGE_ID:
        raise EldenBundleFeatureError(f"目标NPC脚本由其他成果包管理：{identity[0]}/{identity[1]}")
    body = match.group(4).replace("\r\n", "\n").rstrip("\n") + "\n"
    if hashlib.sha256(body.encode("gb18030")).hexdigest() != match.group(3):
        raise EldenBundleFeatureError("目标NPC受管块已被手工修改，拒绝覆盖")


def _database_rows(database: bytes, sql: str, parameters: tuple[object, ...] = ()) -> list[tuple[object, ...]]:
    path: Path | None = None
    connection: sqlite3.Connection | None = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".db") as handle:
            path = Path(handle.name)
            handle.write(database)
        connection = sqlite3.connect(path)
        integrity = connection.execute("PRAGMA integrity_check").fetchone()
        if not integrity or integrity[0] != "ok":
            raise EldenBundleFeatureError(f"SQLite完整性检查失败：{integrity}")
        return connection.execute(sql, parameters).fetchall()
    except sqlite3.Error as exc:
        raise EldenBundleFeatureError(f"读取目标数据库失败：{exc}") from exc
    finally:
        if connection is not None:
            connection.close()
        if path is not None:
            path.unlink(missing_ok=True)


def _title_totals(row: dict[str, object]) -> dict[str, int]:
    return {column: int(_integer(row.get(column), 0) or 0) for column in TITLE_COLUMNS}


def _allocate_and_upsert_titles(
    database: bytes,
    progression: V21Progression,
) -> tuple[bytes, dict[str, int], list[dict[str, object]]]:
    existing_rows = _database_rows(
        database,
        'SELECT Name,Shape,Dc,Dc2,Mc,Mc2,Sc,Sc2,HP,MP FROM StdItems WHERE StdMode=70',
    )
    existing_by_name = {str(row[0]): row for row in existing_rows}
    occupied = {int(row[1]): str(row[0]) for row in existing_rows}
    requested: list[tuple[str, dict[str, object], str]] = []
    requested.extend(("main", row, "title") for row in progression.main_title_rows)
    for chain in progression.independent_titles:
        requested.extend((chain.chain_id, row, "title") for row in chain.rows)
    requested.extend(("rebirth", row, "rebirth") for row in progression.rebirth_rows)

    names = [_title_token(row.get("称号名称"), "称号名称") for _chain, row, _kind in requested]
    duplicate_names = sorted({name for name in names if names.count(name) > 1})
    if duplicate_names:
        raise EldenBundleFeatureError("待安装称号名称重复：" + "、".join(duplicate_names))

    mapping: dict[str, int] = {}
    report: list[dict[str, object]] = []
    staged = database
    for chain_id, row, kind in requested:
        name = _title_token(row.get("称号名称"), "称号名称")
        requested_shape = _integer(row.get("称号编号"))
        if requested_shape is None or not 0 <= requested_shape <= 255:
            raise EldenBundleFeatureError(f"称号{name}的规划编号必须为0至255")
        existing = existing_by_name.get(name)
        if existing is not None:
            actual_shape = int(existing[1])
            expected = (0, 0, 0, 0, 0, 0, 0, 0) if kind == "rebirth" else (
                _title_totals(row)["攻击"], _title_totals(row)["攻击"],
                _title_totals(row)["魔法"], _title_totals(row)["魔法"],
                _title_totals(row)["道术"], _title_totals(row)["道术"],
                _title_totals(row)["HP"], _title_totals(row)["MP"],
            )
            actual = tuple(int(value or 0) for value in existing[2:])
            if actual != expected:
                raise EldenBundleFeatureError(f"同名称号静态属性不一致：{name}")
        elif requested_shape not in occupied:
            actual_shape = requested_shape
        else:
            actual_shape = next((value for value in range(212, 256) if value not in occupied), -1)
            if actual_shape < 0:
                raise EldenBundleFeatureError(f"没有可用称号编号：{name}")
        if existing is None:
            occupied[actual_shape] = name
        mapping[name] = actual_shape
        report.append({
            "chain": chain_id,
            "name": name,
            "requested_shape": requested_shape,
            "actual_shape": actual_shape,
            "reused": existing is not None,
        })
        values = (
            _rebirth_title_item_values(name, actual_shape, 248)
            if kind == "rebirth"
            else _title_item_values(name, actual_shape, _title_totals(row))
        )
        staged = apply_sqlite_upsert(staged, {
            "type": "sqlite_upsert",
            "table": "StdItems",
            "unique_key": ["Name"],
            "conflict_keys": [["StdMode", "Shape"]],
            "allocate": {"Idx": "max_plus_one"},
            "on_conflict": "error",
            "values": values,
        })
    return staged, mapping, report


class _NpcRegistry:
    def __init__(self, server: Path):
        self.server = Path(server)
        self.merchant_path = self.server / Path(MERCHANT_RELATIVE.replace("/", "\\"))
        self.document = read_text_document(self.merchant_path)
        self.rows: list[list[str]] = [line.split("\t") for line in self.document.text.splitlines()]
        self.mapinfo = read_text_document(self.server / "Mir200" / "Envir" / "MapInfo.txt").text
        self._maps: dict[str, tuple[bytes, int, int, int]] = {}
        self.assignments: list[NpcAssignment] = []

    def remove(self, script_path: str, map_code: str) -> None:
        matches = [
            index for index, fields in enumerate(self.rows)
            if len(fields) >= 2 and fields[0].casefold() == script_path.casefold()
            and fields[1].casefold() == map_code.casefold()
        ]
        if len(matches) > 1:
            raise EldenBundleFeatureError(f"MerChant注册重复：{script_path}/{map_code}")
        if matches:
            self.rows.pop(matches[0])

    def _map(self, map_code: str) -> tuple[bytes, int, int, int]:
        key = map_code.casefold()
        if key not in self._maps:
            filename = _map_file_from_info(self.mapinfo, map_code)
            path = self.server / "Mir200" / "Map" / filename
            if not path.is_file():
                raise EldenBundleFeatureError(f"目标服缺少地图文件：{filename}")
            data = path.read_bytes()
            width, height, cell_size = _geometry(data, map_code)
            self._maps[key] = (data, width, height, cell_size)
        return self._maps[key]

    def _occupied(self, map_code: str, excluding: set[tuple[str, str]]) -> set[tuple[int, int]]:
        result: set[tuple[int, int]] = set()
        for fields in self.rows:
            if len(fields) < 4 or fields[1].casefold() != map_code.casefold():
                continue
            if (fields[0].casefold(), fields[1].casefold()) in excluding:
                continue
            try:
                result.add((int(fields[2]), int(fields[3])))
            except ValueError:
                continue
        return result

    def assign(
        self,
        feature_id: str,
        name: str,
        map_code: str,
        appearance: int,
        script_path: str,
        *,
        reuse_path: str | None = None,
    ) -> NpcAssignment:
        candidates = [
            (index, fields) for index, fields in enumerate(self.rows)
            if len(fields) >= 7 and fields[1].casefold() == map_code.casefold()
            and fields[0].casefold() in {script_path.casefold(), (reuse_path or script_path).casefold()}
        ]
        if len(candidates) > 1:
            raise EldenBundleFeatureError(f"NPC复用目标不唯一：{feature_id}/{map_code}")
        reused = bool(candidates)
        if candidates:
            index, fields = candidates[0]
            try:
                x, y = int(fields[2]), int(fields[3])
            except ValueError as exc:
                raise EldenBundleFeatureError(f"NPC坐标不是整数：{fields}") from exc
            data, width, height, cell_size = self._map(map_code)
            if not _walkable(data, width, height, cell_size, x, y):
                raise EldenBundleFeatureError(f"已有NPC坐标不可走：{map_code} {x},{y}")
            self.rows[index] = [script_path, map_code, str(x), str(y), name, "0", str(appearance), "0"]
        else:
            name_conflicts = [
                fields for fields in self.rows
                if len(fields) >= 5 and fields[1].casefold() == map_code.casefold()
                and fields[4].casefold() == name.casefold()
            ]
            if name_conflicts:
                raise EldenBundleFeatureError(f"NPC名称已被其他脚本占用：{map_code}/{name}")
            data, width, height, cell_size = self._map(map_code)
            occupied = self._occupied(map_code, set())
            center = (100, 75) if map_code.casefold() == "xy_nmgf_main" else (width // 2, height // 2)
            candidates_points = [
                (x, y)
                for y in range(1, height - 1)
                for x in range(1, width - 1)
            ]
            candidates_points.sort(key=lambda point: (
                max(abs(point[0] - center[0]), abs(point[1] - center[1])),
                (point[0] - center[0]) ** 2 + (point[1] - center[1]) ** 2,
                point[1], point[0],
            ))
            point = next((
                candidate for candidate in candidates_points
                if _walkable(data, width, height, cell_size, *candidate)
                and all(max(abs(candidate[0] - other[0]), abs(candidate[1] - other[1])) >= 3 for other in occupied)
            ), None)
            if point is None:
                raise EldenBundleFeatureError(f"地图{map_code}找不到未占用可走NPC坐标")
            x, y = point
            self.rows.append([script_path, map_code, str(x), str(y), name, "0", str(appearance), "0"])
        assignment = NpcAssignment(feature_id, name, script_path, map_code, x, y, appearance, reused)
        self.assignments.append(assignment)
        return assignment

    def bytes(self) -> bytes:
        lines = ["\t".join(fields) for fields in self.rows]
        text = self.document.newline.join(lines)
        if lines:
            text += self.document.newline
        return encode_text_document(self.document, text)


def _title_description_entries(
    progression: V21Progression,
    shape_mapping: dict[str, int],
) -> tuple[list[str], list[str]]:
    del shape_mapping  # Shape is already represented in StdItems; descriptions are name keyed.
    itemdesc: list[str] = []
    fenghao: list[str] = []
    rows = [*progression.main_title_rows]
    for chain in progression.independent_titles:
        rows.extend(chain.rows)
    for row in rows:
        name = _title_token(row.get("称号名称"), "称号名称")
        values = _title_totals(row)
        summary = _summary(values, include_mp=True, include_drop=True)
        fenghao.append(f"{name}={summary}")
        itemdesc.append(
            f"{name}=\\242/　艾尔登法环晋升称号\\-\\146/　[称号属性]："
            f"\\251/　攻击+{values['攻击']}-{values['攻击']}"
            f"\\251/　魔法+{values['魔法']}-{values['魔法']}"
            f"\\251/　道术+{values['道术']}-{values['道术']}"
            f"\\251/　HP+{values['HP']}\\251/　MP+{values['MP']}"
            f"\\251/　基础爆率+{values['基础爆率']}%"
        )
    cumulative = _rebirth_cumulative(progression.rebirth_rows)
    for row, power in zip(progression.rebirth_rows, cumulative):
        name = _title_token(row.get("称号名称"), "转生称号")
        level = int(_integer(row.get("档位/等级")) or 0)
        text = f"原生转生{level}转显示称号；累计神力+{power}%"
        fenghao.append(f"{name}={text}")
        itemdesc.append(f"{name}=\\242/　艾尔登法环转生称号\\-\\251/　{text}")
    return itemdesc, fenghao


def _title_drop_hook(progression: V21Progression, variable: str) -> str:
    rows = [*progression.main_title_rows]
    for chain in progression.independent_titles:
        rows.extend(chain.rows)
    lines: list[str] = []
    for row in rows:
        amount = int(_integer(row.get("基础爆率"), 0) or 0)
        if amount:
            name = _title_token(row.get("称号名称"), "称号名称")
            lines += ["#IF", f"CHECKFENGHAO {name}", "#ACT", f"INC {variable} {amount}"]
    return "\n".join(lines)


class EldenBundleFeatureService:
    """Compile the target-dependent progression part of the filled NPC bundle."""

    def __init__(self, platform_root: Path):
        self.root = Path(platform_root).resolve()
        repository = PackageRepository(self.root / "packages")
        repository.refresh()
        self.installer = Installer(repository, self.root / "backups")

    @staticmethod
    def _change(
        root: Path,
        relative: str,
        after: bytes,
        operation: str,
        *,
        scope: str = "server",
    ) -> PlannedChange | None:
        path = root / Path(relative.replace("/", "\\"))
        before = path.read_bytes() if path.is_file() else None
        if before == after:
            return None
        return PlannedChange(relative, before, after, operation, PACKAGE_ID, scope)

    def preflight_progression(
        self,
        server: Path,
        bundle_root: Path,
        *,
        client: Path,
        launcher: Path,
    ) -> EldenBundleFeaturePlan:
        plan = EldenBundleFeaturePlan(
            str(Path(server).resolve()), str(Path(bundle_root).resolve()),
            str(Path(client).resolve()), str(Path(launcher).resolve()),
        )
        try:
            progression = load_v21_progression(Path(bundle_root))
            wash_costs, opening_costs = load_wash_open_costs(Path(bundle_root))
            soul_tasks = load_soul_tasks(Path(bundle_root))
            war_ash_costs = load_war_ash_costs(Path(bundle_root))
            target = TargetInspector.inspect(Path(server))
            client_root = Path(client).resolve()
            launcher_root = Path(launcher).resolve()
            if not (client_root / "data").is_dir():
                raise EldenBundleFeatureError(f"客户端缺少data目录：{client_root}")
            if not launcher_root.is_dir():
                raise EldenBundleFeatureError(f"登录器生成器目录不存在：{launcher_root}")

            database_path = target.root / Path(DB_RELATIVE.replace("/", "\\"))
            database = database_path.read_bytes()
            database, shape_mapping, mapping_report = _allocate_and_upsert_titles(database, progression)
            plan.title_mapping = mapping_report
            synthesis_workbooks = load_synthesis_workbooks(Path(bundle_root))
            synthesis_indices = _resolve_synthesis_indices(database, synthesis_workbooks)
            _resolve_unique_items(
                database,
                (name for task in soul_tasks for name, _amount in task.materials),
            )
            _resolve_unique_items(
                database,
                (name for cost in war_ash_costs for name, _amount in cost.materials),
            )
            qf_path = target.root / Path(QFUNCTION_RELATIVE.replace("/", "\\"))
            qf_doc = read_text_document(qf_path)
            _validate_existing_c1_tasks(target.root, database, qf_doc.text, soul_tasks)

            registry = _NpcRegistry(target.root)
            # Old generic entries would bypass continent-specific progression.
            registry.remove("初始营地NPC/神树赐福", "xycamp")
            registry.remove("玄渊成长/转生", "xycamp")
            registry.remove("玄渊武器附魔/武器附魔", "3")

            scripts: list[tuple[NpcAssignment, str, bytes, set[tuple[str, str]]]] = []
            seal_paths = ["玄渊成长/神印修行"] + [
                f"玄渊法环/神印/{item.chain_id.split('_')[-1]}" for item in progression.seal_npcs[1:]
            ]
            for spec, script_path in zip(progression.seal_npcs, seal_paths):
                assignment = registry.assign(
                    spec.chain_id, spec.npc_name, spec.map_code, spec.appearance, script_path,
                )
                body = compile_seal_npc(progression.seal_rows, len(spec.rows), spec.npc_name)
                scripts.append((
                    assignment,
                    f"seal-{spec.chain_id.casefold()}",
                    _managed_npc(f"seal-{spec.chain_id.casefold()}", body),
                    {("xy.lab.seal-title-v2", "seal-self-contained")},
                ))

            tree_reuse = {
                "MAIN_C02": "初始营地NPC/神树赐福",
                "MAIN_C04": "初始营地NPC/海域神树",
                "MAIN_C05": "初始营地NPC/利耶尼亚神树",
                "MAIN_C06": "初始营地NPC/盖利德神树",
                "MAIN_C07": "初始营地NPC/雪原神树",
                "MAIN_C08": "初始营地NPC/巨人神树",
                "MAIN_C09": "初始营地NPC/亚坛神树",
                "MAIN_C10": "初始营地NPC/罗德尔神树",
            }
            all_main_titles = [_title_token(row.get("称号名称"), "称号名称") for row in progression.main_title_rows]
            for spec in progression.main_titles:
                suffix = spec.chain_id.split("_")[-1]
                desired = tree_reuse.get(spec.chain_id, f"玄渊法环/称号/贯穿/{suffix}")
                assignment = registry.assign(
                    spec.chain_id, spec.npc_name, spec.map_code, spec.appearance, desired,
                    reuse_path=tree_reuse.get(spec.chain_id),
                )
                body = compile_title_chain(spec.rows, all_main_titles, f"EL_MAIN_{suffix}", spec.npc_name)
                scripts.append((
                    assignment,
                    f"main-title-{suffix.casefold()}",
                    _managed_npc(f"main-title-{suffix.casefold()}", body),
                    {("xy.lab.seal-title-v2", "title-self-contained")},
                ))

            shared_independent_index: int | None = None
            shared_independent_assignment: NpcAssignment | None = None
            shared_independent_body = ""
            for spec in progression.independent_titles:
                suffix = spec.chain_id.split("_")[-1]
                path = f"玄渊法环/称号/独立/{suffix}"
                assignment = registry.assign(
                    spec.chain_id, spec.npc_name, spec.map_code, spec.appearance, path,
                )
                titles = [_title_token(row.get("称号名称"), "称号名称") for row in spec.rows]
                body = compile_title_chain(spec.rows, titles, f"EL_IND_{suffix}", spec.npc_name)
                if spec.chain_id == "INDEPENDENT_C01":
                    shared_independent_index = len(scripts)
                    shared_independent_assignment = assignment
                    shared_independent_body = body
                scripts.append((
                    assignment,
                    f"independent-title-{suffix.casefold()}",
                    _managed_npc(f"independent-title-{suffix.casefold()}", body),
                    set(),
                ))

            rebirth_titles = [_title_token(row.get("称号名称"), "转生称号") for row in progression.rebirth_rows]
            for spec in progression.rebirths:
                suffix = spec.chain_id.split("_")[-1]
                path = f"玄渊法环/转生/{suffix}"
                assignment = registry.assign(
                    spec.chain_id, spec.npc_name, spec.map_code, spec.appearance, path,
                )
                body = compile_rebirth_segment(spec.rows, rebirth_titles, f"EL_REBIRTH_{suffix}", spec.npc_name)
                scripts.append((
                    assignment,
                    f"rebirth-{suffix.casefold()}",
                    _managed_npc(f"rebirth-{suffix.casefold()}", body),
                    set(),
                ))

            shared_synthesis_count = 0
            for npc_id, workbook in synthesis_workbooks:
                settings = workbook.settings
                synthesis_body = _compile_npc_script(settings, workbook.recipes, synthesis_indices)
                folder = Path(workbook.path).parent.name
                if (
                    settings.npc_name == "拾荒匠·阿砾"
                    and settings.map_code.casefold() == "xy_drifting_islands"
                ):
                    if shared_independent_index is None or shared_independent_assignment is None:
                        raise EldenBundleFeatureError("漂流群岛拾荒匠缺少独立称号NPC，无法合并合成功能")
                    if settings.appearance != shared_independent_assignment.appearance:
                        raise EldenBundleFeatureError("漂流群岛拾荒匠的称号与合成外观不一致")
                    combined = compile_combined_title_synthesis(shared_independent_body, synthesis_body)
                    scripts[shared_independent_index] = (
                        shared_independent_assignment,
                        "independent-title-synthesis-c01",
                        _managed_npc("independent-title-synthesis-c01", combined),
                        set(),
                    )
                    shared_synthesis_count += 1
                    assignment = shared_independent_assignment
                else:
                    path = f"玄渊法环/合成/{folder}"
                    assignment = registry.assign(
                        npc_id,
                        settings.npc_name,
                        settings.map_code,
                        settings.appearance,
                        path,
                    )
                    scripts.append((
                        assignment,
                        f"synthesis-{npc_id.casefold()}",
                        _managed_npc(f"synthesis-{npc_id.casefold()}", synthesis_body + "\n"),
                        set(),
                    ))
                for recipe in workbook.recipes:
                    plan.synthesis_recipes.append({
                        "npc_id": npc_id,
                        "npc_name": assignment.name,
                        "map_code": assignment.map_code,
                        **asdict(recipe),
                    })
            if shared_synthesis_count != 1:
                raise EldenBundleFeatureError(
                    f"漂流群岛拾荒匠合成功能应合并1次，实际{shared_synthesis_count}次"
                )

            wash_template_path = target.root / Path(
                "Mir200/Envir/Market_Def/玄渊实验室/装备洗练-XY_NMGF_MAIN.txt".replace("/", "\\")
            )
            opening_template_path = target.root / Path(
                "Mir200/Envir/Market_Def/玄渊实验室/装备开光-XY_NMGF_MAIN.txt".replace("/", "\\")
            )
            if not wash_template_path.is_file() or not opening_template_path.is_file():
                raise EldenBundleFeatureError("目标服缺少已验收的装备洗练或装备开光底座")
            wash_template = wash_template_path.read_bytes().decode("gb18030")
            opening_template = opening_template_path.read_bytes().decode("gb18030")

            war_ash_assignment = registry.assign(
                "WAR_ASH_C02",
                "战灰刻印师·盖尔",
                "XY_NMGF_MAIN",
                244,
                "玄渊法环/战灰刻印师/战灰刻印师",
            )
            war_ash_workbook = (
                Path(bundle_root) / "平台暂缺字段_不要直接安装" / "42_武器附魔.xlsx"
            )
            war_ash_book = load_weapon_enchant_workbook(
                war_ash_workbook,
                setting_overrides={
                    "NPC脚本相对路径": (
                        "Mir200/Envir/Market_Def/"
                        f"{war_ash_assignment.script_path}-{war_ash_assignment.map_code}.txt"
                    ),
                    "NPC坐标X": war_ash_assignment.x,
                    "NPC坐标Y": war_ash_assignment.y,
                    "消耗类型": "元宝",
                    "消耗数量": 0,
                    "复用母版包ID": "xy.optional.equipment-wash-opening",
                },
            )
            war_ash_core = _exclusive_script_bytes(render_elden_war_ash_core(
                _managed_body_or_text(wash_template),
                war_ash_book,
                *war_ash_costs,
            ))
            war_ash_runtime = _exclusive_script_bytes(render_elden_war_ash_runtime(war_ash_book))
            scripts.append((
                war_ash_assignment,
                "war-ash-c02",
                _managed_npc(
                    "war-ash-c02",
                    "[@main]\n#IF\n#ACT\n"
                    "#CALL [\\玄渊武器附魔\\武器附魔核心.txt] @main\nBREAK\n",
                ),
                set(),
            ))
            for cost in war_ash_costs:
                plan.service_costs.append({"service": "war_ash", **asdict(cost)})

            for cost in wash_costs:
                script_path = (
                    "玄渊实验室/装备洗练"
                    if cost.continent_id == "C02"
                    else f"玄渊法环/洗练/{cost.continent_id}"
                )
                assignment = registry.assign(
                    f"WASH_{cost.continent_id}", "装备洗练", cost.map_code, 220, script_path,
                )
                body = compile_wash_variant(wash_template, cost.material_amount, cost.yuanbao)
                scripts.append((
                    assignment,
                    f"wash-{cost.continent_id.casefold()}",
                    _managed_npc(f"wash-{cost.continent_id.casefold()}", body),
                    set(),
                ))
                plan.service_costs.append({"service": "wash", **asdict(cost)})
            for cost in opening_costs:
                script_path = (
                    "玄渊实验室/装备开光"
                    if cost.continent_id == "C02"
                    else f"玄渊法环/开光/{cost.continent_id}"
                )
                assignment = registry.assign(
                    f"OPEN_{cost.continent_id}", "装备开光", cost.map_code, 221, script_path,
                )
                body = compile_opening_variant(opening_template, cost.material_amount, cost.yuanbao)
                scripts.append((
                    assignment,
                    f"opening-{cost.continent_id.casefold()}",
                    _managed_npc(f"opening-{cost.continent_id.casefold()}", body),
                    set(),
                ))
                plan.service_costs.append({"service": "opening", **asdict(cost)})

            for task in soul_tasks:
                plan.soul_tasks.append(asdict(task))
                if task.continent_id == "C01":
                    continue
                path = f"玄渊法环/任务/{task.task_id.replace('-', '_')}"
                assignment = registry.assign(
                    task.task_id,
                    task.task_name,
                    task.map_code,
                    226,
                    path,
                )
                scripts.append((
                    assignment,
                    f"soul-{task.task_id.casefold()}",
                    _managed_npc(f"soul-{task.task_id.casefold()}", compile_soul_task(task)),
                    set(),
                ))

            changes: list[PlannedChange | None] = []
            for assignment, _kind, desired, accepted in scripts:
                npc_labels = scan_labels(desired.decode("gb18030"))
                if npc_labels.duplicates:
                    raise EldenBundleFeatureError(
                        f"NPC脚本存在重复标签：{assignment.name}/{npc_labels.duplicates}"
                    )
                relative = (
                    f"Mir200/Envir/Market_Def/{assignment.script_path}-{assignment.map_code}.txt"
                )
                path = target.root / Path(relative.replace("/", "\\"))
                before = path.read_bytes() if path.is_file() else None
                try:
                    _validate_existing_npc(before, accepted)
                except EldenBundleFeatureError as exc:
                    raise EldenBundleFeatureError(f"{relative}：{exc}") from exc
                changes.append(self._change(target.root, relative, desired, "compile-elden-progression-npc"))
            for relative_path, desired, operation in (
                (WEAPON_ENCHANT_CORE_RELATIVE, war_ash_core, "compile-elden-war-ash-core"),
                (WEAPON_ENCHANT_RUNTIME_RELATIVE, war_ash_runtime, "compile-elden-war-ash-runtime"),
            ):
                relative = relative_path.as_posix()
                path = target.root / Path(relative.replace("/", "\\"))
                before = path.read_bytes() if path.is_file() else None
                _validate_legacy_exclusive_script(relative, before, desired)
                changes.append(self._change(target.root, relative, desired, operation))
            changes.append(self._change(target.root, MERCHANT_RELATIVE, registry.bytes(), "register-elden-progression-npcs"))
            changes.append(self._change(target.root, DB_RELATIVE, database, "upsert-elden-progression-titles"))

            qf_text = qf_doc.text
            for anchor in ("XY_EQUIP_MAKER_DROP_ANCHOR", "XY_EQUIP_MAKER_RUNTIME_DROP_ANCHOR"):
                qf_text = remove_managed_anchor_hook(qf_text, "xy.lab.seal-title-v2", anchor).text
            qf_text = install_managed_anchor_hook(
                qf_text, PACKAGE_ID, "XY_EQUIP_MAKER_DROP_ANCHOR",
                _title_drop_hook(progression, "N$XY_最终爆率"), qf_doc.newline,
                canonical_before_existing_hooks=True,
            ).text
            qf_text = install_managed_anchor_hook(
                qf_text, PACKAGE_ID, "XY_EQUIP_MAKER_RUNTIME_DROP_ANCHOR",
                _title_drop_hook(progression, "N$XY_RT_Drop"), qf_doc.newline,
                canonical_before_existing_hooks=True,
            ).text
            for anchor in ("XY_EQUIP_MAKER_POWER_ANCHOR", "XY_EQUIP_MAKER_RUNTIME_POWER_ANCHOR"):
                qf_text = remove_managed_anchor_hook(qf_text, "xy.initial-camp.seven-npcs", anchor).text
            qf_text = install_managed_anchor_hook(
                qf_text, PACKAGE_ID, "XY_EQUIP_MAKER_POWER_ANCHOR",
                compile_rebirth_power_hook(progression.rebirth_rows, "N$倍攻"), qf_doc.newline,
                canonical_before_existing_hooks=True,
            ).text
            qf_text = install_managed_anchor_hook(
                qf_text, PACKAGE_ID, "XY_EQUIP_MAKER_RUNTIME_POWER_ANCHOR",
                compile_rebirth_power_hook(progression.rebirth_rows, "N$XY_RT_Power"), qf_doc.newline,
                canonical_before_existing_hooks=True,
            ).text
            qf_text = install_managed_anchor_hook(
                qf_text, PACKAGE_ID, "XY_EXECUTION_LAB_TOUGHNESS_ANCHOR",
                compile_task_toughness_hook(soul_tasks), qf_doc.newline,
                canonical_before_existing_hooks=True,
            ).text
            qf_text = install_managed_anchor_hook(
                qf_text, PACKAGE_ID, "XY_EXECUTION_LAB_CHANCE_ANCHOR",
                compile_task_execution_hook(soul_tasks, "N$XY_EXEC_ChanceBP"), qf_doc.newline,
                canonical_before_existing_hooks=True,
            ).text
            qf_text = upsert_weapon_enchant_qfunction_hook(qf_text)
            labels = scan_labels(qf_text)
            if labels.duplicates:
                raise EldenBundleFeatureError(f"QFunction存在重复标签：{labels.duplicates}")
            changes.append(self._change(
                target.root, QFUNCTION_RELATIVE,
                encode_text_document(qf_doc, qf_text), "install-elden-progression-hooks",
            ))

            panel_path = target.root / Path(ATTRIBUTE_PANEL_RELATIVE.replace("/", "\\"))
            panel_doc = read_text_document(panel_path)
            panel_text = install_managed_anchor_hook(
                panel_doc.text, PACKAGE_ID, "XY_EQUIP_MAKER_PANEL_EXEC_CHANCE_ANCHOR",
                compile_task_execution_hook(soul_tasks, "N$XY_UI_C_Value02"), panel_doc.newline,
                canonical_before_existing_hooks=True,
            ).text
            changes.append(self._change(
                target.root, ATTRIBUTE_PANEL_RELATIVE,
                encode_text_document(panel_doc, panel_text), "install-elden-soul-task-panel-hook",
            ))

            qm_path = target.root / Path(QMANAGE_RELATIVE.replace("/", "\\"))
            qm_doc = read_text_document(qm_path)
            qm_text = remove_event_hook(
                qm_doc.text,
                "xy.lab.seal-title-v2",
                "Login",
                qm_doc.newline,
                LEGACY_SEAL_LOGIN_HOOK_HASHES,
            ).text
            qm_text = ensure_event_label(qm_text, PACKAGE_ID, "Login", qm_doc.newline).text
            login_lines = [
                "#IF", "#ACT",
                f"VAR Integer HUMAN {SEAL_VARIABLE}",
                f"LOADVAR HUMAN {SEAL_VARIABLE} {SEAL_STATE_SCRIPT_PATH}",
            ]
            for task in soul_tasks:
                if task.variable:
                    login_lines += [
                        f"VAR Integer HUMAN {task.variable}",
                        f"LOADVAR HUMAN {task.variable} {SEAL_STATE_SCRIPT_PATH}",
                    ]
            qm_text = install_event_hook(
                qm_text, PACKAGE_ID, "Login",
                "\n".join(login_lines),
                qm_doc.newline,
            ).text
            changes.append(self._change(
                target.root, QMANAGE_RELATIVE,
                encode_text_document(qm_doc, qm_text), "install-elden-seal-login-state",
            ))
            state_path = target.root / Path(SEAL_STATE_RELATIVE.replace("/", "\\"))
            if not state_path.is_file():
                changes.append(self._change(
                    target.root, SEAL_STATE_RELATIVE, b"\xef\xbb\xbf", "initialize-elden-seal-state",
                ))

            itemdesc_entries, fenghao_entries = _title_description_entries(progression, shape_mapping)
            itemdesc_path = target.root / Path(ITEMDESC_RELATIVE.replace("/", "\\"))
            itemdesc_doc = read_text_document(itemdesc_path) if itemdesc_path.is_file() else TextDocument("", "gb18030", "\r\n")
            changes.append(self._change(
                target.root, ITEMDESC_RELATIVE,
                _merge_named_description_lines(itemdesc_doc, itemdesc_entries),
                "merge-elden-title-itemdesc",
            ))
            setup_path = target.root / Path(SETUP_RELATIVE.replace("/", "\\"))
            setup_doc = read_text_document(setup_path)
            changes.append(self._change(
                target.root, SETUP_RELATIVE,
                _enable_setup_flags(setup_doc, ("SendItemDescList", "SendTzItemDescList")),
                "enable-elden-title-descriptions",
            ))

            client_candidates = ["data/fenghao.dat", "Resources/fenghao.dat"]
            client_targets = [
                relative for relative in client_candidates
                if (client_root / Path(relative.replace("/", "\\"))).is_file()
            ] or [client_candidates[0]]
            final_fenghao: bytes | None = None
            for relative in client_targets:
                path = client_root / Path(relative.replace("/", "\\"))
                document = read_text_document(path) if path.is_file() else TextDocument("", "gb18030", "\r\n")
                after = _merge_named_description_lines(document, fenghao_entries)
                final_fenghao = after if relative.casefold() == "data/fenghao.dat" else final_fenghao
                changes.append(self._change(client_root, relative, after, "merge-elden-title-fenghao", scope="client"))
            if final_fenghao is None:
                source_path = client_root / "data" / "fenghao.dat"
                source_doc = read_text_document(source_path) if source_path.is_file() else TextDocument("", "gb18030", "\r\n")
                final_fenghao = _merge_named_description_lines(source_doc, fenghao_entries)
            launcher_relative = "补丁文件夹/Data/fenghao.dat"
            changes.append(self._change(
                launcher_root, launcher_relative, final_fenghao,
                "mirror-elden-title-fenghao", scope="launcher",
            ))

            plan.npcs = registry.assignments
            plan.changes = [change for change in changes if change is not None]
            plan.warnings.extend([
                "神印以HUMAN变量XY_SEAL_BASE_LEVEL作为唯一进度；10个NPC只限制本处可达上限。",
                "100级贯穿称号严格校验上一大陆第10级称号；高级成功后回收前一级。",
                "20转进度只读写翎风引擎原生转生等级；称号只是显示镜像。",
                "称号编号已按目标数据库实时分配；映射已同步脚本、StdItems、ItemDescList、客户端与登录器fenghao.dat。",
                "16组合成功能共38条配方；漂流群岛拾荒匠的称号与合成入口合并为同一个NPC。",
                "C2至C10普通洗练、装备开光费用直接读取经济总账；首次开光继续免费，重复开光原子扣除材料与原生元宝。",
                "C1三条魂魄任务保留已验收NPC、称号和完成状态；C2五条任务使用HUMAN变量防重并随登录载入。",
                "魂魄任务的韧性与处决概率同步接入实效重算链和属性图标显示链。",
                "战灰刻印师在原武器附魔模块内扩展普通、稀有、清除三种操作，所有材料与原生货币二次复核后一次扣除。",
                "旧地图3的武器附魔NPC注册已停用，战灰刻印师改由宁姆格福承载；旧脚本文件仅作追溯不再被注册。",
            ])
            plan.install_plan = InstallPlan(
                target_root=str(target.root), client_root=str(client_root), launcher_root=str(launcher_root),
                package_ids=[PACKAGE_ID], package_versions={PACKAGE_ID: PACKAGE_VERSION},
                parameters={
                    "bundle_root": str(Path(bundle_root).resolve()),
                    "seal_levels": 255,
                    "main_title_levels": 100,
                    "independent_title_levels": 12,
                    "rebirth_levels": 20,
                    "synthesis_npcs": 16,
                    "synthesis_recipes": 38,
                    "wash_npcs": 9,
                    "opening_npcs": 9,
                    "soul_tasks": 8,
                    "war_ash_operations": 3,
                    "title_mapping": mapping_report,
                    "npcs": [asdict(item) for item in plan.npcs],
                },
                changes=plan.changes,
                warnings=plan.warnings,
                operation_type="elden-npc-v21-progression",
                candidate_packages=[PACKAGE_ID],
                superseded_package_ids=["xy.lab.seal-title-v2"],
            )
        except (OSError, UnicodeError, ValueError, sqlite3.Error, SqlitePatchError, TextPatchError) as exc:
            plan.blockers.append(str(exc))
        return plan


    def install(self, plan: EldenBundleFeaturePlan):
        if plan.blockers or plan.install_plan is None:
            raise EldenBundleFeatureError("艾尔登法环多实例成长NPC安装被阻止：\n" + "\n".join(plan.blockers))
        return self.installer.install(plan.install_plan)


__all__ = [
    "ChainNpc", "EldenBundleFeatureError", "EldenBundleFeaturePlan",
    "EldenBundleFeatureService", "NpcAssignment", "V21Progression",
    "compile_rebirth_power_hook", "compile_rebirth_segment", "compile_seal_npc",
    "compile_title_chain", "load_config_rows", "load_v21_progression",
]
