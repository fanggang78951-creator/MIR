from __future__ import annotations

import hashlib
import math
import re
from dataclasses import dataclass, replace
from pathlib import Path

from .encoding import encode_text_document, read_text_document
from .monster_library import (
    MonsterLibraryChange,
    MonsterLibraryError,
    MonsterLibraryPlan,
    MonsterLibraryReceipt,
    MonsterLibraryService,
    MonsterRecord,
    _monster_rows,
    _monster_hash,
    _sha256_bytes,
    _sha256_file,
    appr_to_library,
)
from .user_documents import preferred_document


class MonsterWorkbookError(MonsterLibraryError):
    pass


SHEET_NAME = "怪物生成"
WORKBOOK_NAME = "20_怪物批量生成.xlsx"
ALLOWED_STATUSES = {"可安装", "待配置", "忽略"}
LEGACY_HEADERS = (
    "状态", "怪物名称", "等级", "经验", "血量", "防御", "魔防",
    "最小攻击", "最大攻击", "命中", "模型库编号", "备注",
)
HEADERS = (
    "状态", "怪物名称", "等级", "经验", "血量", "防御", "魔防",
    "最小攻击", "最大攻击", "命中", "颜色", "模型库编号", "备注",
)
OPTIONAL_HEADERS = (
    "模型Appr", "魔法值", "魔法攻击", "道术攻击", "行走速度", "攻击速度", "不死系",
)
SUPPORTED_HEADERS = HEADERS + OPTIONAL_HEADERS
REQUIRED_VALUE_COLUMNS = (
    "怪物名称", "等级", "经验", "血量", "防御", "魔防", "最小攻击", "最大攻击", "命中",
)
FIELD_BY_HEADER = {
    "等级": "Lvl",
    "经验": "Exp",
    "血量": "HP",
    "魔法值": "MP",
    "防御": "AC",
    "魔防": "MAC",
    "最小攻击": "DC",
    "最大攻击": "DCMAX",
    "魔法攻击": "MC",
    "道术攻击": "SC",
    "命中": "HIT",
    "行走速度": "WALK_SPD",
    "攻击速度": "ATTACK_SPD",
    "不死系": "Undead",
}
POSITIVE_FIELDS = {"HP", "WALK_SPD", "ATTACK_SPD"}
MODEL_FIELD_ALLOWLIST = frozenset({
    "Race", "RaceImg", "Appr", "SPEED", "WalkStep", "WalkWait",
    "AttackState", "AttackSource", "DisableSimpleActor",
})
NEUTRAL_TEMPLATE_NAME = "稻草人"
DEFAULT_NAME_COLOR = 151
NAME_COLOR_LABELS = {
    151: "黄色",
    154: "蓝色",
    242: "紫色",
    249: "红色",
    70: "橙色",
    250: "绿色",
    254: "青色",
    245: "粉色",
}
NAME_COLOR_ALIASES = {
    "黄": 151, "黄色": 151, "普通": 151, "普通怪": 151, "151": 151,
    "蓝": 154, "蓝色": 154, "精英": 154, "精英怪": 154, "154": 154,
    "紫": 242, "紫色": 242, "首领": 242, "首领怪": 242, "242": 242,
    "红": 249, "红色": 249, "boss": 249, "boss怪": 249, "249": 249,
    "橙": 70, "橙色": 70, "70": 70,
    "绿": 250, "绿色": 250, "250": 250,
    "青": 254, "青色": 254, "254": 254,
    "粉": 245, "粉色": 245, "245": 245,
}
MISSING_COLOR_HEADER = object()


@dataclass(frozen=True)
class MonsterWorkbookInspection:
    source: str
    workbook_hash: str
    sheet: str
    data_rows: int
    install_rows: int
    pending_rows: int
    ignored_rows: int


@dataclass(frozen=True)
class MonsterSpec:
    source_row: int
    name: str
    model_id: int | None
    model_appr: int | None
    name_color: int | None
    overrides: tuple[tuple[str, int], ...]
    note: str


@dataclass(frozen=True)
class MonsterWorkbookCompileResult:
    source: str
    workbook_hash: str
    specs: tuple[MonsterSpec, ...]
    blockers: tuple[str, ...]
    warnings: tuple[str, ...]
    skipped: tuple[str, ...]
    inspection: MonsterWorkbookInspection


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _integer(value: object, *, row: int, column: str, required: bool = False) -> int | None:
    text = _cell_text(value)
    if not text:
        if required:
            raise MonsterWorkbookError(f"第{row}行：{column}不能为空")
        return None
    if text.startswith("="):
        raise MonsterWorkbookError(f"第{row}行：{column}不允许使用公式")
    if isinstance(value, bool):
        raise MonsterWorkbookError(f"第{row}行：{column}必须是整数")
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            raise MonsterWorkbookError(f"第{row}行：{column}必须是整数")
        result = int(value)
    else:
        try:
            result = int(text)
        except (TypeError, ValueError) as exc:
            raise MonsterWorkbookError(f"第{row}行：{column}必须是整数") from exc
        if str(result) != text and text not in {f"+{result}", f"-{abs(result)}"}:
            raise MonsterWorkbookError(f"第{row}行：{column}必须是整数")
    if result < 0:
        raise MonsterWorkbookError(f"第{row}行：{column}不能为负数")
    return result


def _name_color(value: object, *, row: int) -> int | None:
    if value is MISSING_COLOR_HEADER:
        return None
    text = _cell_text(value)
    if not text:
        return DEFAULT_NAME_COLOR
    result = NAME_COLOR_ALIASES.get(text.casefold())
    if result is None:
        allowed = "、".join(NAME_COLOR_LABELS.values())
        raise MonsterWorkbookError(f"第{row}行：颜色只能填写{allowed}，留空默认为黄色")
    return result


class MonsterWorkbookService:
    """Compile a user XLSX into custom Monster rows backed by the local V2 model library."""

    def __init__(self, platform_root: Path):
        self.platform_root = Path(platform_root).resolve()
        self.default_workbook = preferred_document(
            self.platform_root,
            WORKBOOK_NAME,
        )
        self.library = MonsterLibraryService(self.platform_root)

    @staticmethod
    def _rows(workbook_path: Path) -> tuple[str, list[tuple[int, dict[str, object]]]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise MonsterWorkbookError("平台缺少XLSX读取组件 openpyxl") from exc

        path = Path(workbook_path).resolve()
        if not path.is_file():
            raise MonsterWorkbookError(f"怪物生成表不存在：{path}")
        if path.suffix.casefold() != ".xlsx":
            raise MonsterWorkbookError("怪物生成表只接受 .xlsx 文件")
        try:
            book = load_workbook(path, read_only=True, data_only=False)
        except Exception as exc:
            raise MonsterWorkbookError(f"怪物生成表无法读取，请先保存并关闭Excel/WPS：{exc}") from exc
        try:
            if SHEET_NAME not in book.sheetnames:
                raise MonsterWorkbookError(f"怪物生成表缺少工作表：{SHEET_NAME}")
            sheet = book[SHEET_NAME]
            raw_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = tuple(_cell_text(item) for item in raw_headers)
            duplicates = sorted({item for item in headers if item and headers.count(item) > 1})
            if duplicates:
                raise MonsterWorkbookError(f"怪物生成表表头重复：{', '.join(duplicates)}")
            missing = [item for item in LEGACY_HEADERS if item not in headers]
            if missing:
                raise MonsterWorkbookError(f"怪物生成表缺少列：{', '.join(missing)}")
            indexes = {name: headers.index(name) for name in SUPPORTED_HEADERS if name in headers}
            rows: list[tuple[int, dict[str, object]]] = []
            for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                row = {
                    name: (
                        MISSING_COLOR_HEADER if name == "颜色" and name not in indexes
                        else values[indexes[name]] if name in indexes and indexes[name] < len(values)
                        else None
                    )
                    for name in SUPPORTED_HEADERS
                }
                if not any(
                    _cell_text(value) for value in row.values()
                    if value is not MISSING_COLOR_HEADER
                ):
                    continue
                rows.append((row_number, row))
            return str(path), rows
        finally:
            book.close()

    def compile(self, workbook_path: Path) -> MonsterWorkbookCompileResult:
        path = Path(workbook_path).resolve()
        source, rows = self._rows(path)
        workbook_hash = hashlib.sha256(path.read_bytes()).hexdigest().upper()
        specs: list[MonsterSpec] = []
        blockers: list[str] = []
        warnings: list[str] = []
        skipped: list[str] = []
        seen_names: set[str] = set()
        install_rows = pending_rows = ignored_rows = 0

        for row_number, row in rows:
            status = _cell_text(row["状态"])
            name = _cell_text(row["怪物名称"])
            if status not in ALLOWED_STATUSES:
                blockers.append(f"第{row_number}行 {name or '未命名怪物'}：状态必须是可安装、待配置或忽略")
                continue
            if status == "忽略":
                ignored_rows += 1
                skipped.append(f"第{row_number}行 {name or '未命名怪物'}：状态为忽略")
                continue
            if status == "待配置":
                pending_rows += 1
                blockers.append(f"第{row_number}行 {name or '未命名怪物'}：状态为待配置")
                continue
            install_rows += 1
            try:
                for column in REQUIRED_VALUE_COLUMNS:
                    if not _cell_text(row[column]):
                        raise MonsterWorkbookError(f"第{row_number}行：{column}不能为空")
                if len(name) > 60:
                    raise MonsterWorkbookError(f"第{row_number}行：怪物名称不能超过60个字符")
                name_key = name.casefold()
                if name_key in seen_names:
                    raise MonsterWorkbookError(f"第{row_number}行：怪物名称重复：{name}")
                model_id = _integer(row["模型库编号"], row=row_number, column="模型库编号")
                model_appr = _integer(row["模型Appr"], row=row_number, column="模型Appr")
                name_color = _name_color(row["颜色"], row=row_number)
                overrides: dict[str, int] = {}
                for header, field in FIELD_BY_HEADER.items():
                    required = header in REQUIRED_VALUE_COLUMNS
                    parsed = _integer(row[header], row=row_number, column=header, required=required)
                    if parsed is not None:
                        if field in POSITIVE_FIELDS and parsed == 0:
                            raise MonsterWorkbookError(f"第{row_number}行：{header}必须大于0")
                        overrides[field] = parsed
                if overrides["DCMAX"] < overrides["DC"]:
                    raise MonsterWorkbookError(f"第{row_number}行：最大攻击不能小于最小攻击")
                seen_names.add(name_key)
                specs.append(MonsterSpec(
                    source_row=row_number,
                    name=name,
                    model_id=model_id,
                    model_appr=model_appr,
                    name_color=name_color,
                    overrides=tuple(overrides.items()),
                    note=_cell_text(row["备注"]),
                ))
            except MonsterWorkbookError as exc:
                blockers.append(str(exc))

        if not specs and not blockers:
            blockers.append("怪物生成表没有可安装的数据行")
        inspection = MonsterWorkbookInspection(
            source=source,
            workbook_hash=workbook_hash,
            sheet=SHEET_NAME,
            data_rows=len(rows),
            install_rows=install_rows,
            pending_rows=pending_rows,
            ignored_rows=ignored_rows,
        )
        return MonsterWorkbookCompileResult(
            source=source,
            workbook_hash=workbook_hash,
            specs=tuple(specs),
            blockers=tuple(blockers),
            warnings=tuple(warnings),
            skipped=tuple(skipped),
            inspection=inspection,
        )

    def inspect(self, workbook_path: Path) -> MonsterWorkbookInspection:
        return self.compile(workbook_path).inspection

    @staticmethod
    def _stable_pick(candidates: list[MonsterRecord], token: str) -> MonsterRecord:
        ordered = sorted(candidates, key=lambda item: item.monster_id)
        digest = hashlib.sha256(token.encode("utf-8")).digest()
        return ordered[int.from_bytes(digest[:8], "big") % len(ordered)]

    def _resolve_model(
        self,
        spec: MonsterSpec,
        records: list[MonsterRecord],
        workbook_hash: str,
    ) -> tuple[MonsterRecord, str]:
        ready_verified = [
            record for record in records
            if record.status == "ready" and record.closure_status == "ready_verified"
        ]
        if spec.model_id is not None:
            registered = next((record for record in records if record.monster_id == spec.model_id), None)
            if registered is None:
                raise MonsterWorkbookError(f"第{spec.source_row}行 {spec.name}：模型库编号{spec.model_id}不存在")
            if registered.status != "ready" or registered.closure_status not in {"ready_verified", "ready_opaque"}:
                raise MonsterWorkbookError(
                    f"第{spec.source_row}行 {spec.name}：模型库编号{spec.model_id}已收录，"
                    f"但资料未闭环（{registered.closure_status}）"
                )
            model = registered
            if spec.model_appr is not None and model.appearance_id != spec.model_appr:
                raise MonsterWorkbookError(
                    f"第{spec.source_row}行 {spec.name}：模型库编号{spec.model_id}的Appr是{model.appearance_id}，与填写的{spec.model_appr}不一致"
                )
            return model, "指定库编号"
        candidates = ready_verified
        mode = "自动随机"
        if spec.model_appr is not None:
            candidates = [record for record in ready_verified if record.appearance_id == spec.model_appr]
            mode = "指定Appr"
            if not candidates:
                raise MonsterWorkbookError(f"第{spec.source_row}行 {spec.name}：模型Appr {spec.model_appr}没有完整可用补丁")
        if not candidates:
            raise MonsterWorkbookError("默认怪物库没有可用的完整模型")
        token = f"{workbook_hash}|{spec.source_row}|{spec.name}|{mode}"
        return self._stable_pick(candidates, token), mode

    @staticmethod
    def _compose_new_monster_values(
        spec: MonsterSpec,
        neutral_values: dict[str, object],
        model_values: dict[str, object],
    ) -> dict[str, object]:
        values = dict(neutral_values)
        values["Name"] = spec.name
        for field in MODEL_FIELD_ALLOWLIST:
            values[field] = model_values[field]
        values.update(dict(spec.overrides))
        return values

    @classmethod
    def _custom_record(
        cls,
        spec: MonsterSpec,
        base: MonsterRecord,
        neutral_values: dict[str, object],
        *,
        source_rowid: int | None = None,
    ) -> MonsterRecord:
        values = cls._compose_new_monster_values(spec, neutral_values, base.monster_values)
        monster_json, monster_hash = _monster_hash(values)
        return replace(
            base,
            monster_id=spec.source_row,
            source_rowid=spec.source_row if source_rowid is None else source_rowid,
            monster_name=spec.name,
            level=int(values["Lvl"]),
            exp=int(values["Exp"]),
            hp=int(values["HP"]),
            hit=int(values["HIT"]),
            monster_json=monster_json,
            monster_hash=monster_hash,
            status="ready",
            skip_reason=None,
        )

    @staticmethod
    def _existing_record(
        spec: MonsterSpec,
        rowid: int,
        current_values: dict[str, object],
    ) -> MonsterRecord:
        values = dict(current_values)
        values["Name"] = spec.name
        values.update(dict(spec.overrides))
        monster_json, monster_hash = _monster_hash(values)
        appearance_id = int(values["Appr"])
        library_no, slot_no = appr_to_library(appearance_id)
        return MonsterRecord(
            monster_id=spec.source_row,
            source_rowid=rowid,
            monster_name=spec.name,
            appearance_id=appearance_id,
            library_no=library_no,
            slot_no=slot_no,
            race=int(values["Race"]),
            race_img=int(values["RaceImg"]),
            level=int(values["Lvl"]),
            exp=int(values["Exp"]),
            hp=int(values["HP"]),
            hit=int(values["HIT"]),
            monster_json=monster_json,
            monster_hash=monster_hash,
            source_kind="existing",
            resource_name="",
            source_path="",
            companion_path=None,
            source_hash=None,
            companion_hash=None,
            source_size=0,
            preview_path=None,
            preview_state="existing",
            status="ready",
            skip_reason=None,
            pak_password=None,
        )

    def preflight(
        self,
        workbook_path: Path,
        server_root: Path,
        client_data: Path,
        *,
        generator_login_dir: Path | None = None,
    ) -> MonsterLibraryPlan:
        compiled = self.compile(workbook_path)
        records_by_status = self.library.list_monsters("all")
        ready_verified = [
            record for record in records_by_status
            if record.status == "ready" and record.closure_status == "ready_verified"
        ]
        server_root = Path(server_root).resolve()
        client_data = Path(client_data).resolve()
        target_database = server_root / "Mud2" / "DB" / "ApexM2.DB"
        target_rows: dict[str, list[tuple[int, dict[str, object]]]] = {}
        if target_database.is_file():
            try:
                for rowid, values in _monster_rows(target_database):
                    target_rows.setdefault(str(values["Name"]).casefold(), []).append((rowid, values))
            except MonsterLibraryError as exc:
                target_rows = {}
                compiled_blockers = [str(exc)]
            else:
                compiled_blockers = []
        else:
            compiled_blockers = []

        try:
            compatible_ready, compatibility_conflicts = self.library.compatible_models(
                ready_verified, server_root, client_data,
            )
        except MonsterLibraryError as exc:
            compatible_ready = []
            compatibility_conflicts = {}
            compiled_blockers.append(str(exc))
        compatible_ids = {record.monster_id for record in compatible_ready}
        records: list[MonsterRecord] = []
        blockers = list(compiled.blockers) + compiled_blockers
        assignments: list[dict[str, object]] = []
        resource_required_names: set[str] = set()
        appearance_changed_names: list[str] = []
        auto_reselected: list[dict[str, object]] = []
        selection_warnings: list[str] = []
        for spec in compiled.specs:
            try:
                existing = target_rows.get(spec.name.casefold(), [])
                if len(existing) > 1:
                    raise MonsterWorkbookError(f"目标服存在多个同名怪物，无法唯一同步：{spec.name}")
                explicit_model = spec.model_id is not None or spec.model_appr is not None
                if existing and not explicit_model:
                    rowid, current_values = existing[0]
                    record = self._existing_record(spec, rowid, current_values)
                    records.append(record)
                    assignments.append({
                        "source_row": spec.source_row,
                        "monster_name": spec.name,
                        "action": "更新属性",
                        "mode": "保留现有模型",
                        "model_id": None,
                        "model_name": spec.name,
                        "appr": record.appearance_id,
                        "library_no": record.library_no,
                        "slot_no": record.slot_no,
                    })
                    continue

                base, mode = self._resolve_model(spec, records_by_status, compiled.workbook_hash)
                original = base
                if not explicit_model and base.monster_id not in compatible_ids:
                    if not compatible_ready:
                        raise MonsterWorkbookError(
                            f"第{spec.source_row}行 {spec.name}：目标客户端没有可兼容的完整怪物模型"
                        )
                    base, _ = self._resolve_model(spec, compatible_ready, compiled.workbook_hash)
                    mode = "自动兼容改选"
                    auto_reselected.append({
                        "source_row": spec.source_row,
                        "monster_name": spec.name,
                        "original_model_id": original.monster_id,
                        "original_model_name": original.monster_name,
                        "original_library_no": original.library_no,
                        "reason": compatibility_conflicts.get(original.monster_id, "目标补丁不兼容"),
                        "selected_model_id": base.monster_id,
                        "selected_model_name": base.monster_name,
                        "selected_library_no": base.library_no,
                    })
                if existing:
                    rowid, current_values = existing[0]
                    records.append(self._custom_record(
                        spec,
                        base,
                        current_values,
                        source_rowid=rowid,
                    ))
                else:
                    neutral = target_rows.get(NEUTRAL_TEMPLATE_NAME.casefold(), [])
                    if not neutral:
                        raise MonsterWorkbookError(
                            f"目标服缺少唯一中性模板：{NEUTRAL_TEMPLATE_NAME}；无法新增怪物"
                        )
                    if len(neutral) != 1:
                        raise MonsterWorkbookError(
                            f"目标服中性模板重复：{NEUTRAL_TEMPLATE_NAME}；无法新增怪物"
                        )
                    records.append(self._custom_record(spec, base, neutral[0][1]))
                resource_required_names.add(spec.name.casefold())
                action = "更新属性并换模" if existing else "新增"
                if existing:
                    appearance_changed_names.append(spec.name)
                opaque = base.closure_status == "ready_opaque"
                if opaque:
                    selection_warnings.append(
                        f"第{spec.source_row}行 {spec.name}：模型库编号{base.monster_id}处于ready_opaque，"
                        "仅可显式选择，部署后必须完成单怪验收"
                    )
                assignments.append({
                    "source_row": spec.source_row,
                    "monster_name": spec.name,
                    "action": action,
                    "mode": mode,
                    "model_id": base.monster_id,
                    "model_name": base.monster_name,
                    "appr": base.appearance_id,
                    "library_no": base.library_no,
                    "slot_no": base.slot_no,
                    "single_monster_acceptance_required": opaque,
                })
            except MonsterWorkbookError as exc:
                blockers.append(str(exc))

        if records:
            plan = self.library.preflight_records(
                records,
                server_root,
                client_data,
                allow_updates=True,
                resource_required_names=resource_required_names,
                appearance_changed_names=appearance_changed_names,
                conflicts_are_blockers=True,
                generator_login_dir=generator_login_dir,
            )
        else:
            plan = MonsterLibraryPlan(
                server_root=server_root,
                client_data=client_data,
                selected_ids=(),
                monster_names=(),
                library_numbers=(),
                changes=[],
                blockers=["怪物生成表没有可生成的有效怪物"],
                warnings=[],
                skipped=[],
            )
        plan.operation = "monster-workbook-install"
        plan.source_workbook_path = compiled.source
        plan.source_workbook_hash = compiled.workbook_hash
        plan.model_assignments = tuple(assignments)
        plan.auto_reselected = tuple(auto_reselected)
        plan.blockers.extend(blockers)
        plan.warnings.extend(compiled.warnings)
        plan.warnings.extend(selection_warnings)
        plan.skipped.extend(compiled.skipped)
        self._add_name_color_changes(plan, compiled.specs)
        return plan

    @staticmethod
    def _add_name_color_changes(plan: MonsterLibraryPlan, specs: tuple[MonsterSpec, ...]) -> None:
        eligible = {name.casefold() for name in plan.monster_names}
        selected = [
            spec for spec in specs
            if spec.name.casefold() in eligible and spec.name_color is not None
        ]
        if not selected:
            plan.name_color_assignments = ()
            return

        mon_gen = plan.server_root / "Mir200" / "Envir" / "MonGen.txt"
        if not mon_gen.is_file():
            plan.name_color_assignments = tuple({
                "source_row": spec.source_row,
                "monster_name": spec.name,
                "color": spec.name_color,
                "color_name": NAME_COLOR_LABELS[spec.name_color],
                "matched_rows": 0,
                "changed_rows": 0,
            } for spec in selected)
            plan.warnings.append("目标服尚无MonGen.txt；颜色保留在XLSX中，增加刷怪后重新同步本表即可应用")
            return

        try:
            document = read_text_document(mon_gen)
        except (OSError, UnicodeError) as exc:
            plan.blockers.append(f"MonGen.txt无法安全读取：{exc}")
            return

        spec_by_name = {spec.name.casefold(): spec for spec in selected}
        matched = {key: 0 for key in spec_by_name}
        changed = {key: 0 for key in spec_by_name}
        trailing_newline = document.text.endswith(("\r\n", "\n", "\r"))
        lines = document.text.splitlines()

        for index, line in enumerate(lines):
            stripped = line.strip()
            if not stripped or stripped.startswith(";"):
                continue
            parts = re.split(r"\s+", stripped)
            if len(parts) < 4:
                continue
            key = parts[3].casefold()
            spec = spec_by_name.get(key)
            if spec is None:
                continue
            matched[key] += 1
            if len(parts) not in {7, 8, 9}:
                plan.blockers.append(
                    f"MonGen第{index + 1}行 {spec.name} 为{len(parts)}列，无法安全写入名字颜色"
                )
                continue
            color = str(spec.name_color)
            if len(parts) == 7:
                base = line.rstrip()
                lines[index] = base + "\t0\t" + color + line[len(base):]
            elif len(parts) == 8:
                base = line.rstrip()
                lines[index] = base + "\t" + color + line[len(base):]
            elif parts[8] != color:
                match = re.search(r"\S+(\s*)$", line)
                if match is None:
                    plan.blockers.append(f"MonGen第{index + 1}行 {spec.name} 无法定位颜色字段")
                    continue
                lines[index] = line[:match.start()] + color + match.group(1)
            else:
                continue
            changed[key] += 1

        new_text = document.newline.join(lines)
        if trailing_newline:
            new_text += document.newline
        try:
            after = encode_text_document(document, new_text)
        except UnicodeEncodeError as exc:
            plan.blockers.append(f"MonGen.txt颜色写入无法保持原编码：{exc}")
            return

        plan.name_color_assignments = tuple({
            "source_row": spec.source_row,
            "monster_name": spec.name,
            "color": spec.name_color,
            "color_name": NAME_COLOR_LABELS[spec.name_color],
            "matched_rows": matched[spec.name.casefold()],
            "changed_rows": changed[spec.name.casefold()],
        } for spec in selected)
        missing = [spec.name for spec in selected if matched[spec.name.casefold()] == 0]
        if missing:
            display = "、".join(missing[:10]) + ("等" if len(missing) > 10 else "")
            plan.warnings.append(
                f"{len(missing)}只怪物当前没有MonGen刷新行，颜色暂未写入：{display}；以后增加刷怪后重新同步本表"
            )
        before = mon_gen.read_bytes()
        if after != before and not any("MonGen第" in item for item in plan.blockers):
            plan.mon_gen_after = after
            plan.changes.append(MonsterLibraryChange(
                scope="server",
                target_path=str(mon_gen),
                source_path=None,
                before_hash=_sha256_file(mon_gen),
                after_hash=_sha256_bytes(after),
                kind="monster-name-colors",
            ))

    def install(self, plan: MonsterLibraryPlan) -> MonsterLibraryReceipt:
        if plan.operation != "monster-workbook-install":
            raise MonsterWorkbookError("当前计划不是怪物表格生成计划")
        return self.library.install(plan)

    def rollback(self, transaction_id: str) -> None:
        self.library.rollback(transaction_id)

    @staticmethod
    def plan_summary(plan: MonsterLibraryPlan) -> dict[str, object]:
        return MonsterLibraryService.plan_summary(plan)
