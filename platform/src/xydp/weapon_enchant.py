from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any, Mapping

from openpyxl import load_workbook

from .mingge_dual import (
    DualInstallPlan,
    MinggeDualError,
    _DualService,
    _change,
    _managed_block,
    _plan_id,
    _script_bytes,
    _script_text,
    _sha256,
    _target,
    _upsert_after_label,
)


PACKAGE_ID = "xy.optional.weapon-enchant"
ROUTE = "weapon-enchant"
SHEETS = ("系统设置", "附魔槽位", "洗练规则", "品质规则", "词条候选")
SHEETS_V2 = ("基础设置", "结果候选", "普通洗练设置", "词条路由库", "填写说明")
PLATFORM_ROOT = Path(__file__).resolve().parents[2]
LEGACY_WASH_TEMPLATE = PLATFORM_ROOT / "packages/verified/xy.optional.equipment-wash-opening/payload/server/npc/装备洗练.txt"
LEGACY_WASH_SHA256 = "2775B933E0EE97319E31357FB855DBCF1FC96406F874A03B612D78A4255FED2B"
LEGACY_WASH_EFFECT = Path("Mir200/Envir/EffectImageList.txt")
LEGACY_WASH_EFFECT_LINE = "XY_EquipmentWorkbench.wz"
LEGACY_WASH_SERVER_DEPENDENCIES = (
    (Path("Mir200/Envir/QuestDiary/玄渊实验室/装备洗练/仙级直出.txt"), "04CA49BB6CD51BA3C06EF3F846B81BC21A2DBE3D2183ABB886EB80DA92428F9E"),
    (Path("Mir200/Envir/QuestDiary/玄渊实验室/装备洗练/圣级直出.txt"), "608598D30E668194724EB321B05CBD57839A722AED71D3E48F3310A97FB12CEE"),
)
CORE_RELATIVE = Path("Mir200/Envir/QuestDiary/玄渊武器附魔/武器附魔核心.txt")
RUNTIME_RELATIVE = Path("Mir200/Envir/QuestDiary/玄渊武器附魔/武器附魔运行时.txt")
QFUNCTION_RELATIVE = Path("Mir200/Envir/Market_Def/QFunction-0.txt")
MERCHANT_RELATIVE = Path("Mir200/Envir/MerChant.txt")
QFUNCTION_MARKER = "XY-WEAPON-ENCHANT-V1-ATTACK"
SKILL_IDS = {"烈火剑法": 26, "逐日剑法": 56, "开天斩": 66}
EFFECT_CODES = {"必定暴击": 1, "致命一击": 2}


class WeaponEnchantError(MinggeDualError):
    pass


@dataclass(frozen=True)
class WeaponEnchantSettings:
    system_id: str
    equip_part: str
    slot_count: int
    row_start: int
    npc_relative: str
    npc_map: str
    npc_x: int
    npc_y: int
    npc_name: str
    npc_appearance: int
    item_box: int
    schema_version: int = 1
    allowed_std_modes: tuple[int, ...] = (5, 6)
    clear_rows: tuple[int, ...] = (9, 10, 11)
    cost_type: str = "元宝"
    cost_name: str = ""
    cost_amount: int = 100
    display_prefix: str = "附魔属性："
    template_package_id: str = "xy.optional.equipment-wash-opening"


@dataclass(frozen=True)
class EnchantSlot:
    slot: int
    enabled: bool
    display_name: str
    condition: str
    value: int


@dataclass(frozen=True)
class WashRule:
    rule_id: int
    name: str
    count: int
    cost_type: str
    cost_name: str
    cost_amount: int
    pool: str


@dataclass(frozen=True)
class QualityRule:
    name: str
    weight: int
    pity: int
    color: int


@dataclass(frozen=True)
class EnchantAffix:
    stable_id: int
    name: str
    skill: str
    effect_type: str
    effect_value: int
    quality: str
    weight: int
    quality_weight: int
    color: int
    route: str
    effect_status: str
    skill_id: int
    effect_code: int
    code: str = ""
    source: str = ""
    property_name: str = ""
    native_bind: int = 0
    percent: int = 0
    new_item_index: int = -1


@dataclass(frozen=True)
class LegacyWashRule:
    name: str
    enabled: bool
    denominator: int
    color: int
    note: str


@dataclass(frozen=True)
class WarAshCost:
    operation_id: str
    materials: tuple[tuple[str, int], ...]
    gold: int
    yuanbao: int


@dataclass(frozen=True)
class WeaponEnchantWorkbook:
    path: Path
    settings: WeaponEnchantSettings
    slots: tuple[EnchantSlot, ...]
    wash_rules: tuple[WashRule, ...]
    qualities: tuple[QualityRule, ...]
    affixes: tuple[EnchantAffix, ...]
    schema_version: int = 1
    legacy_rules: tuple[LegacyWashRule, ...] = ()


def _text(value: Any, label: str) -> str:
    result = "" if value is None else str(value).strip()
    if not result:
        raise WeaponEnchantError(f"{label}不能为空")
    return result


def _int(value: Any, label: str, low: int = 0, high: int = 2_147_483_647) -> int:
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise WeaponEnchantError(f"{label}必须是整数") from exc
    if not low <= result <= high:
        raise WeaponEnchantError(f"{label}必须在{low}至{high}之间")
    return result


def _enabled(value: Any) -> bool:
    text = str(value).strip().casefold()
    if text in {"是", "1", "true", "启用", "yes"}:
        return True
    if text in {"否", "0", "false", "停用", "no"}:
        return False
    raise WeaponEnchantError("是否启用只能填写是或否")


def _rows(sheet, headers: tuple[str, ...]):
    actual = tuple("" if cell.value is None else str(cell.value).strip() for cell in sheet[1])
    if actual != headers:
        raise WeaponEnchantError(f"工作表{sheet.title}表头必须为：{'、'.join(headers)}")
    for index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        row = tuple(values[: len(headers)])
        if any(value is not None and str(value).strip() for value in row):
            yield index, row


def _stable_id(name: str, skill: str, effect: str) -> int:
    key = "|".join(part.strip().casefold() for part in (name, skill, effect))
    return 100_000_000 + int(hashlib.sha256(key.encode("utf-8")).hexdigest()[:7], 16)


def _stable_id_from_code(code: str) -> int:
    key = code.strip().casefold()
    return 100_000_000 + int(hashlib.sha256(key.encode("utf-8")).hexdigest()[:7], 16)


def _load_v1_weapon_enchant_workbook(path: Path) -> WeaponEnchantWorkbook:
    source = Path(path).resolve()
    if not source.is_file():
        raise WeaponEnchantError(f"配置表不存在：{source}")
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if tuple(wb.sheetnames) != SHEETS:
            raise WeaponEnchantError("工作表必须依次为：" + "、".join(SHEETS))
        settings = {str(v[0]).strip(): v[1] for _, v in _rows(wb["系统设置"], ("配置项", "配置值"))}
        required = ("系统编号", "目标装备部位", "附魔槽位数", "受管属性起始行", "NPC脚本相对路径", "NPC地图", "NPC坐标X", "NPC坐标Y", "NPC名称", "NPC外观", "物品框编号")
        missing = [name for name in required if name not in settings]
        if missing:
            raise WeaponEnchantError("系统设置缺少：" + "、".join(missing))
        system_id = _text(settings["系统编号"], "系统编号")
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{2,31}", system_id):
            raise WeaponEnchantError("系统编号格式非法")
        if _text(settings["目标装备部位"], "目标装备部位") != "武器":
            raise WeaponEnchantError("首版目标装备部位只允许武器")
        npc_relative = _text(settings["NPC脚本相对路径"], "NPC脚本相对路径").replace("\\", "/")
        pure = PurePosixPath(npc_relative)
        if pure.is_absolute() or ".." in pure.parts or not npc_relative.startswith("Mir200/Envir/Market_Def/"):
            raise WeaponEnchantError("NPC脚本相对路径必须位于Mir200/Envir/Market_Def内")
        parsed_settings = WeaponEnchantSettings(
            system_id, "武器", _int(settings["附魔槽位数"], "附魔槽位数", 1, 3),
            _int(settings["受管属性起始行"], "受管属性起始行", 9, 17), npc_relative,
            _text(settings["NPC地图"], "NPC地图"), _int(settings["NPC坐标X"], "NPC坐标X", 0, 999),
            _int(settings["NPC坐标Y"], "NPC坐标Y", 0, 999), _text(settings["NPC名称"], "NPC名称"),
            _int(settings["NPC外观"], "NPC外观", 0, 9999), _int(settings["物品框编号"], "物品框编号", 1, 99),
        )
        if parsed_settings.row_start + parsed_settings.slot_count - 1 > 19:
            raise WeaponEnchantError("受管属性行不得超过19")
        slots = tuple(EnchantSlot(_int(v[0], f"附魔槽位第{r}行槽位", 1, 3), _enabled(v[1]), _text(v[2], "显示名称"), _text(v[3], "开放条件"), _int(v[4], "条件值")) for r, v in _rows(wb["附魔槽位"], ("槽位", "启用", "显示名称", "开放条件", "条件值")))
        active = sorted(item.slot for item in slots if item.enabled)
        if active != list(range(1, parsed_settings.slot_count + 1)):
            raise WeaponEnchantError("启用槽位必须从1开始连续")
        washes = tuple(WashRule(_int(v[0], "规则ID", 1), _text(v[2], "名称"), _int(v[3], "次数", 1, 100), _text(v[4], "消耗类型"), "" if v[5] is None else str(v[5]).strip(), _int(v[6], "消耗数量"), _text(v[7], "候选池")) for _, v in _rows(wb["洗练规则"], ("规则ID", "启用", "名称", "次数", "消耗类型", "消耗名称", "消耗数量", "候选池")) if _enabled(v[1]))
        qualities = tuple(QualityRule(_text(v[0], "品质"), _int(v[1], "权重"), _int(v[2], "保底次数"), _int(v[3], "颜色", 0, 255)) for _, v in _rows(wb["品质规则"], ("品质", "权重", "保底次数", "颜色")))
        quality_rules = {item.name: item for item in qualities}
        affixes = []
        seen = set()
        for row, v in _rows(wb["词条候选"], ("词条名称", "适用技能", "效果类型", "效果值", "品质", "权重", "颜色", "是否启用")):
            if not _enabled(v[7]):
                continue
            name, skill, effect = _text(v[0], f"第{row}行词条名称"), _text(v[1], f"第{row}行适用技能"), _text(v[2], f"第{row}行效果类型")
            quality = _text(v[4], "品质")
            if quality not in quality_rules:
                raise WeaponEnchantError(f"第{row}行品质未定义")
            stable = _stable_id(name, skill, effect)
            if stable in seen:
                raise WeaponEnchantError(f"第{row}行稳定词条ID冲突")
            seen.add(stable)
            runtime = effect in EFFECT_CODES and skill in SKILL_IDS
            affixes.append(EnchantAffix(stable, name, skill, effect, _int(v[3], "效果值"), quality, _int(v[5], "权重", 1), quality_rules[quality].weight, _int(v[6], "颜色", 0, 255), "textline_runtime" if runtime else "static_only", "candidate", SKILL_IDS.get(skill, 0), EFFECT_CODES.get(effect, 0)))
        if not affixes:
            raise WeaponEnchantError("至少需要一条启用词条")
        return WeaponEnchantWorkbook(source, parsed_settings, slots, washes, qualities, tuple(affixes))
    finally:
        wb.close()


_NATIVE_BINDINGS = {
    "自身防御": (1, 0, -1),
    "防御": (1, 0, -1),
    "自定魔防": (2, 0, -1),
    "魔御": (2, 0, -1),
    "自身攻击": (3, 0, -1),
    "攻击": (3, 0, -1),
    "自身魔法": (4, 0, -1),
    "魔法": (4, 0, -1),
    "自身道术": (5, 0, -1),
    "道术": (5, 0, -1),
    "生命值": (6, 0, -1),
    "魔法值": (7, 0, -1),
    "暴击伤害": (8, 1, -1),
    "致命一击": (9, 1, 0),
    "攻击伤害": (10, 1, 1),
    "神圣一击": (11, 0, -1),
    "真实一击": (12, 0, -1),
    "麻痹一击": (13, 0, -1),
    "冰冻一击": (14, 0, -1),
    "攻击速度": (15, 0, -1),
}


def _optional_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _parse_std_modes(value: Any) -> tuple[int, ...]:
    text = _text(value, "允许StdMode")
    try:
        result = tuple(int(part.strip()) for part in re.split(r"[,，|]", text) if part.strip())
    except ValueError as exc:
        raise WeaponEnchantError("允许StdMode必须用逗号填写整数") from exc
    if not result or any(item < 0 or item > 255 for item in result):
        raise WeaponEnchantError("允许StdMode必须在0至255之间")
    return result


def _parse_clear_rows(value: Any) -> tuple[int, ...]:
    text = _text(value, "兼容清理行")
    match = re.fullmatch(r"(\d+)-(\d+)", text)
    if match:
        start, end = (int(match.group(1)), int(match.group(2)))
        if start > end:
            raise WeaponEnchantError("兼容清理行起始值不得大于结束值")
        result = tuple(range(start, end + 1))
    else:
        result = tuple(int(part.strip()) for part in re.split(r"[,，]", text) if part.strip())
    if not result or min(result) < 0 or max(result) > 19:
        raise WeaponEnchantError("兼容清理行必须位于0至19")
    return result


def _load_v2_weapon_enchant_workbook(
    path: Path,
    setting_overrides: Mapping[str, Any] | None = None,
) -> WeaponEnchantWorkbook:
    source = Path(path).resolve()
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if tuple(wb.sheetnames) != SHEETS_V2:
            raise WeaponEnchantError("V2工作表必须依次为：" + "、".join(SHEETS_V2))
        values = {str(v[0]).strip(): v[1] for _, v in _rows(wb["基础设置"], ("配置项", "配置值"))}
        if setting_overrides:
            unknown = sorted(set(setting_overrides) - set(values))
            if unknown:
                raise WeaponEnchantError("批量包基础设置覆盖出现未知字段：" + "、".join(unknown))
            values.update(setting_overrides)
        required = (
            "配置版本", "系统编号", "目标装备部位", "允许StdMode", "受管身份行", "兼容清理行",
            "NPC脚本相对路径", "NPC地图", "NPC坐标X", "NPC坐标Y", "NPC名称", "NPC外观",
            "物品框编号", "消耗类型", "消耗数量", "显示前缀", "复用母版包ID",
        )
        missing = [name for name in required if name not in values]
        if missing:
            raise WeaponEnchantError("基础设置缺少：" + "、".join(missing))
        if _int(values["配置版本"], "配置版本", 2, 2) != 2:
            raise WeaponEnchantError("配置版本只允许2")
        system_id = _text(values["系统编号"], "系统编号")
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{2,31}", system_id):
            raise WeaponEnchantError("系统编号格式非法")
        if _text(values["目标装备部位"], "目标装备部位") != "武器":
            raise WeaponEnchantError("V2目标装备部位只允许武器")
        npc_relative = _text(values["NPC脚本相对路径"], "NPC脚本相对路径").replace("\\", "/")
        pure = PurePosixPath(npc_relative)
        if pure.is_absolute() or ".." in pure.parts or not npc_relative.startswith("Mir200/Envir/Market_Def/"):
            raise WeaponEnchantError("NPC脚本相对路径必须位于Mir200/Envir/Market_Def内")
        allowed_std_modes = _parse_std_modes(values["允许StdMode"])
        if allowed_std_modes != (5, 6):
            raise WeaponEnchantError("当前已验收武器附魔只允许StdMode 5,6")
        managed_row = _int(values["受管身份行"], "受管身份行", 0, 19)
        clear_rows = _parse_clear_rows(values["兼容清理行"])
        if managed_row not in clear_rows:
            raise WeaponEnchantError("受管身份行必须包含在兼容清理行中")
        cost_type = _text(values["消耗类型"], "消耗类型")
        if cost_type != "元宝":
            raise WeaponEnchantError("V2当前只支持元宝消耗")
        settings = WeaponEnchantSettings(
            system_id=system_id,
            equip_part="武器",
            slot_count=1,
            row_start=managed_row,
            npc_relative=npc_relative,
            npc_map=_text(values["NPC地图"], "NPC地图"),
            npc_x=_int(values["NPC坐标X"], "NPC坐标X", 0, 999),
            npc_y=_int(values["NPC坐标Y"], "NPC坐标Y", 0, 999),
            npc_name=_text(values["NPC名称"], "NPC名称"),
            npc_appearance=_int(values["NPC外观"], "NPC外观", 0, 9999),
            item_box=_int(values["物品框编号"], "物品框编号", 1, 99),
            schema_version=2,
            allowed_std_modes=allowed_std_modes,
            clear_rows=clear_rows,
            cost_type=cost_type,
            cost_name=_optional_text(values.get("消耗名称")),
            cost_amount=_int(values["消耗数量"], "消耗数量", 0),
            display_prefix=_text(values["显示前缀"], "显示前缀"),
            template_package_id=_text(values["复用母版包ID"], "复用母版包ID"),
        )
        if settings.template_package_id != "xy.optional.equipment-wash-opening":
            raise WeaponEnchantError("复用母版包ID必须为xy.optional.equipment-wash-opening")

        route_specs: dict[tuple[str, str], tuple[str, str]] = {}
        for row, v in _rows(wb["词条路由库"], ("属性来源", "属性名称", "路由", "状态", "单位", "适用范围", "说明")):
            source_name = _text(v[0], f"词条路由库第{row}行属性来源")
            property_name = _optional_text(v[1])
            key = (source_name, property_name)
            if key in route_specs:
                raise WeaponEnchantError(f"词条路由库第{row}行重复：{source_name}/{property_name}")
            route_specs[key] = (_text(v[2], "路由"), _text(v[3], "状态"))

        affixes: list[EnchantAffix] = []
        seen_codes: set[str] = set()
        seen_ids: set[int] = set()
        headers = ("词条编码", "显示名称", "属性来源", "属性名称", "适用技能", "效果类型", "效果值", "概率权重", "实际概率", "颜色", "是否启用")
        for row, v in _rows(wb["结果候选"], headers):
            if not _enabled(v[10]):
                continue
            code = _text(v[0], f"第{row}行词条编码")
            normalized_code = code.casefold()
            if normalized_code in seen_codes:
                raise WeaponEnchantError(f"第{row}行词条编码重复")
            seen_codes.add(normalized_code)
            stable = _stable_id_from_code(code)
            if stable in seen_ids:
                raise WeaponEnchantError(f"第{row}行稳定词条ID冲突")
            seen_ids.add(stable)
            name = _text(v[1], f"第{row}行显示名称")
            source_name = _text(v[2], f"第{row}行属性来源")
            property_name = _optional_text(v[3])
            spec = route_specs.get((source_name, property_name))
            if spec is None:
                raise WeaponEnchantError(f"第{row}行没有对应词条路由：{source_name}/{property_name}")
            route, status = spec
            skill = _optional_text(v[4])
            effect_type = _optional_text(v[5])
            effect_value = _int(v[6], f"第{row}行效果值", 0)
            native_bind, percent, new_item_index = _NATIVE_BINDINGS.get(property_name, (0, 0, -1))
            if route == "skill_damage_runtime" and (skill not in SKILL_IDS or effect_type not in EFFECT_CODES):
                raise WeaponEnchantError(f"第{row}行技能或效果类型没有当前端已验证路由")
            if route == "native_bind" and native_bind == 0:
                raise WeaponEnchantError(f"第{row}行已有洗练词条没有原生绑定映射：{property_name}")
            affixes.append(EnchantAffix(
                stable_id=stable,
                name=name,
                skill=skill,
                effect_type=effect_type,
                effect_value=effect_value,
                quality="",
                weight=_int(v[7], f"第{row}行概率权重", 1),
                quality_weight=1,
                color=_int(v[9], f"第{row}行颜色", 0, 255),
                route=route,
                effect_status=status,
                skill_id=SKILL_IDS.get(skill, 0),
                effect_code=EFFECT_CODES.get(effect_type, 0),
                code=code,
                source=source_name,
                property_name=property_name,
                native_bind=native_bind,
                percent=percent,
                new_item_index=new_item_index,
            ))
        if not affixes:
            raise WeaponEnchantError("至少需要一条启用结果候选")

        legacy_rules = tuple(
            LegacyWashRule(
                _text(v[0], f"普通洗练设置第{row}行结果名称"),
                _enabled(v[1]),
                _int(v[2], f"普通洗练设置第{row}行触发分母", 1),
                _int(v[3], f"普通洗练设置第{row}行颜色", 0, 255),
                _optional_text(v[4]),
            )
            for row, v in _rows(wb["普通洗练设置"], ("结果名称", "是否启用", "触发分母", "颜色", "说明"))
        )
        required_legacy = {"神佑", "天赐", "圣级", "仙级", "传说", "上古", "灵级"}
        if {item.name for item in legacy_rules} != required_legacy:
            raise WeaponEnchantError("普通洗练设置必须且只能包含：" + "、".join(sorted(required_legacy)))
        return WeaponEnchantWorkbook(
            source,
            settings,
            (EnchantSlot(1, True, "武器附魔", "无条件", 0),),
            (WashRule(1, "洗练", 1, settings.cost_type, settings.cost_name, settings.cost_amount, "V2结果池"),),
            (),
            tuple(affixes),
            2,
            legacy_rules,
        )
    finally:
        wb.close()


def load_weapon_enchant_workbook(
    path: Path,
    *,
    setting_overrides: Mapping[str, Any] | None = None,
) -> WeaponEnchantWorkbook:
    source = Path(path).resolve()
    if not source.is_file():
        raise WeaponEnchantError(f"配置表不存在：{source}")
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        sheetnames = tuple(wb.sheetnames)
    finally:
        wb.close()
    if sheetnames == SHEETS:
        if setting_overrides:
            raise WeaponEnchantError("批量包基础设置覆盖只支持V2五表")
        return _load_v1_weapon_enchant_workbook(source)
    if sheetnames == SHEETS_V2:
        return _load_v2_weapon_enchant_workbook(source, setting_overrides)
    raise WeaponEnchantError("工作表结构无法识别；仅支持旧五表或武器附魔V2五表")


def _render_runtime_v1(book: WeaponEnchantWorkbook) -> str:
    lines = ["{", "; XY-WEAPON-ENCHANT-V1", "[@XY_WE_ATTACK_DAMAGE]", "#IF", "#ACT"]
    for index in range(book.settings.slot_count):
        row = book.settings.row_start + index
        lines.append(f"GetCustomItemValueEx 1 {row} N$XY_WE_T{index+1} N$XY_WE_ID{index+1} N$XY_WE_VALUE{index+1} N$XY_WE_TYPE{index+1}")
    for index in range(book.settings.slot_count):
        for affix in book.affixes:
            if affix.route != "textline_runtime":
                continue
            lines.extend(("#IF", f"EQUAL N$XY_WE_ID{index+1} {affix.stable_id}", f"EQUAL <$CURRRUSEMAGICID> {affix.skill_id}", "#ACT", f"ChangeDamageValue 1 + <$STR(N$XY_WE_VALUE{index+1})>"))
    lines.extend(("BREAK", "}"))
    return "\n".join(lines)


def _render_core_v1(book: WeaponEnchantWorkbook) -> str:
    box = book.settings.item_box
    rows = [book.settings.row_start + i for i in range(book.settings.slot_count)]
    labels = "|".join(f"<$STR(S$XY_WE_LABEL{i+1})>" for i in range(book.settings.slot_count))
    lines = [
        "{", "; XY-WEAPON-ENCHANT-V1", "[@main]", "#SAY",
        f"<ITEMBOX:{box}:12:1:30:70:76:76:请放入武器>",
        "请选择附魔槽位：" + "  ".join(f"<{slot.display_name}/@XY_WE_SLOT_{slot.slot}>" for slot in book.slots if slot.enabled),
        "  ".join(f"<{rule.name}/@XY_WE_WASH_{rule.rule_id}>" for rule in book.wash_rules) + "    <关闭/@exit>", "",
        f"[@ItemIntoBox{box}]", "#IF", "#ACT", "GOTO @main", "BREAK", "",
    ]
    for slot in book.slots:
        if not slot.enabled:
            continue
        lines.extend((f"[@XY_WE_SLOT_{slot.slot}]", "#IF", "#ACT", f"MOV N$XY_WE_SELECTED_SLOT {slot.slot}", f"SENDMSG 6 已选择{slot.display_name}。", "GOTO @main", ""))
    for rule in book.wash_rules:
        lines.extend((f"[@XY_WE_WASH_{rule.rule_id}]", "#IF"))
        if rule.cost_type == "元宝":
            lines.append(f"CHECKGAMEGOLD > {max(0, rule.cost_amount - 1)}")
        elif rule.cost_type == "物品":
            lines.append(f"CHECKITEM {rule.cost_name} {rule.cost_amount}")
        elif rule.cost_type != "无":
            raise WeaponEnchantError(f"洗练规则{rule.rule_id}消耗类型只支持无、元宝、物品")
        lines.extend(("#ACT",))
        if rule.cost_type == "元宝":
            lines.append(f"GAMEGOLD - {rule.cost_amount}")
        elif rule.cost_type == "物品":
            lines.append(f"TAKE {rule.cost_name} {rule.cost_amount}")
        lines.extend((f"MOV N$XY_WE_DRAW_COUNT {rule.count}", "GOTO @XY_WE_PICK", "#ELSEACT", "MESSAGEBOX 附魔所需材料或货币不足。", "BREAK", ""))
    effective_weights = {item.stable_id: item.weight * item.quality_weight for item in book.affixes}
    total_weight = sum(effective_weights.values())
    lines.extend(("[@XY_WE_PICK]", "#IF", "SMALL N$XY_WE_SELECTED_SLOT 1", "#ACT", "MOV N$XY_WE_SELECTED_SLOT 1", ""))
    for affix in book.affixes[:-1]:
        denominator = max(2, round(total_weight / effective_weights[affix.stable_id]))
        lines.extend(("#IF", f"RANDOM {denominator}", "#ACT", f"MOV N$XY_WE_PICK_ID {affix.stable_id}", f"MOV N$XY_WE_PICK_VALUE {affix.effect_value}", f"MOV N$XY_WE_PICK_TYPE {affix.effect_code}", f"MOV N$XY_WE_PICK_COLOR {affix.color}", "GOTO @XY_WE_PREPARE", ""))
    fallback = book.affixes[-1]
    lines.extend(("#IF", "#ACT", f"MOV N$XY_WE_PICK_ID {fallback.stable_id}", f"MOV N$XY_WE_PICK_VALUE {fallback.effect_value}", f"MOV N$XY_WE_PICK_TYPE {fallback.effect_code}", f"MOV N$XY_WE_PICK_COLOR {fallback.color}", "GOTO @XY_WE_PREPARE", "", "[@XY_WE_PREPARE]"))
    for i, row in enumerate(rows):
        lines.extend((f"GetCustomItemValueEx boxitem{box} {row} N$XY_WE_T{i+1} N$XY_WE_ID{i+1} N$XY_WE_VALUE{i+1} N$XY_WE_TYPE{i+1}", f"MOV S$XY_WE_LABEL{i+1} 无", f"MOV N$XY_WE_COLOR{i+1} 255"))
    for i in range(book.settings.slot_count):
        lines.extend(("#IF", f"EQUAL N$XY_WE_SELECTED_SLOT {i+1}", "#ACT", f"MOV N$XY_WE_ID{i+1} N$XY_WE_PICK_ID", f"MOV N$XY_WE_VALUE{i+1} N$XY_WE_PICK_VALUE", f"MOV N$XY_WE_TYPE{i+1} N$XY_WE_PICK_TYPE", f"MOV N$XY_WE_COLOR{i+1} N$XY_WE_PICK_COLOR"))
    for i in range(book.settings.slot_count):
        for affix in book.affixes:
            lines.extend(("#IF", f"EQUAL N$XY_WE_ID{i+1} {affix.stable_id}", "#ACT", f"MOV S$XY_WE_LABEL{i+1} {affix.name}", f"MOV N$XY_WE_COLOR{i+1} {affix.color}"))
    lines.extend(("#IF", "#ACT", "MOV N$XY_WE_DISPLAY_COLOR N$XY_WE_PICK_COLOR"))
    lines.extend(("[@XY_WE_WRITE]",))
    lines.extend(("#IF", "#ACT", f"LockUpdateItem boxitem{box}"))
    for row in rows:
        lines.extend((f"SetCustomItemAbil boxitem{box} {row} 4 0", f"SetCustomItemValueEx boxitem{box} {row} = 0 0 0"))
    for i, row in enumerate(rows):
        lines.extend((f"SetCustomItemAbil boxitem{box} {row} 60 0", f"SetCustomItemValueEx boxitem{box} {row} = <$STR(N$XY_WE_ID{i+1})> <$STR(N$XY_WE_VALUE{i+1})> <$STR(N$XY_WE_TYPE{i+1})>"))
    lines.extend((f"SetCustomItemText boxitem{box} [武器附魔:{labels}]", f"SetCustomItemTextColor boxitem{box} <$STR(N$XY_WE_DISPLAY_COLOR)>", f"UpdateItem boxitem{box}", "MESSAGEBOX 附魔写入完成；技能暴击与致命一击仍需游戏验收。", "BREAK", "}"))
    return "\n".join(lines)


def _legacy_template_text(platform_root: Path | None = None) -> str:
    template = (
        Path(platform_root).resolve() / "packages/verified/xy.optional.equipment-wash-opening/payload/server/npc/装备洗练.txt"
        if platform_root is not None
        else LEGACY_WASH_TEMPLATE
    )
    if not template.is_file():
        raise WeaponEnchantError("verified装备洗练母版不存在")
    raw = template.read_bytes()
    if hashlib.sha256(raw).hexdigest().upper() != LEGACY_WASH_SHA256:
        raise WeaponEnchantError("verified装备洗练母版哈希漂移，禁止派生")
    return raw.decode("utf-8").replace("\r\n", "\n").replace("\r", "\n")


def _clear_identity_lines(box: int, rows: tuple[int, ...]) -> list[str]:
    lines: list[str] = []
    for row in rows:
        for kind in range(5):
            lines.append(f"SetCustomItemAbil boxitem{box} {row} {kind} 0")
        lines.append(f"SetCustomItemValueEx boxitem{box} {row} = 0 0 0")
    return lines


def _war_ash_cost_summary(cost: WarAshCost) -> str:
    parts = [f"{name}×{amount}" for name, amount in cost.materials]
    if cost.gold:
        parts.append(f"金币×{cost.gold}")
    if cost.yuanbao:
        parts.append(f"元宝×{cost.yuanbao}")
    if not parts:
        raise WeaponEnchantError(f"{cost.operation_id}消耗不能为空")
    return "、".join(parts)


def _war_ash_validate_cost(cost: WarAshCost) -> None:
    if not re.fullmatch(r"[A-Za-z0-9_-]{3,64}", cost.operation_id):
        raise WeaponEnchantError("战灰操作ID格式无效")
    seen: set[str] = set()
    for name, amount in cost.materials:
        if not name.strip() or any(char.isspace() for char in name.strip()):
            raise WeaponEnchantError(f"{cost.operation_id}材料名称无效：{name}")
        if name.casefold() in seen:
            raise WeaponEnchantError(f"{cost.operation_id}材料重复：{name}")
        seen.add(name.casefold())
        if amount <= 0:
            raise WeaponEnchantError(f"{cost.operation_id}材料数量必须大于0")
    if cost.gold < 0 or cost.yuanbao < 0:
        raise WeaponEnchantError(f"{cost.operation_id}货币数量不能为负数")
    _war_ash_cost_summary(cost)


def _war_ash_positive_checks(cost: WarAshCost) -> list[str]:
    lines = [f"CHECKITEM {name} {amount}" for name, amount in cost.materials]
    if cost.gold:
        lines.append(f"CHECKGOLD {cost.gold}")
    if cost.yuanbao:
        lines.append(f"CHECKGAMEGOLD > {max(0, cost.yuanbao - 1)}")
    return lines


def _war_ash_deductions(cost: WarAshCost) -> list[str]:
    lines = [f"TAKE {name} {amount}" for name, amount in cost.materials]
    if cost.gold:
        lines.append(f"TAKE 金币 {cost.gold}")
    if cost.yuanbao:
        lines.append(f"GAMEGOLD - {cost.yuanbao}")
    return lines


def _war_ash_negative_guard(cost: WarAshCost, label: str) -> str:
    checks: list[str] = []
    for name, amount in cost.materials:
        checks.append(f"NOT CHECKITEM {name} {amount}")
    if cost.gold:
        checks.append(f"NOT CHECKGOLD {cost.gold}")
    if cost.yuanbao:
        checks.append(f"NOT CHECKGAMEGOLD > {max(0, cost.yuanbao - 1)}")
    lines: list[str] = []
    for index, check in enumerate(checks):
        lines.append(check)
        if index != len(checks) - 1:
            lines.extend((
                "#ACT",
                f"MESSAGEBOX {label}所需战灰材料或货币不足，本次未扣除资源。",
                "BREAK",
                "#IF",
            ))
    return "\n".join(lines)


def _war_ash_pick_lines(affixes: tuple[EnchantAffix, ...]) -> list[str]:
    if not affixes:
        raise WeaponEnchantError("稀有战灰至少需要一条已验证技能词条")
    remaining = sum(item.weight for item in affixes)
    lines: list[str] = []
    for index, affix in enumerate(affixes):
        lines.append("#IF")
        if index != len(affixes) - 1:
            lines.append(f"RANDOMEX {affix.weight} {remaining}")
        lines.extend((
            "#ACT",
            f"MOV N$XY_WE_PICK_ID {affix.stable_id}",
            f"MOV N$XY_WE_PICK_VALUE {affix.effect_value}",
            f"MOV N$XY_WE_PICK_TYPE {affix.effect_code}",
            f"MOV N$XY_WE_PICK_COLOR {affix.color}",
            f"MOV S$XY_WE_PICK_LABEL {affix.name}",
            "GOTO @XY_ASH_RARE_WRITE",
            "BREAK",
        ))
        remaining -= affix.weight
    return lines


def render_elden_war_ash_core(
    wash_template: str,
    book: WeaponEnchantWorkbook,
    normal_cost: WarAshCost,
    rare_cost: WarAshCost,
    clear_cost: WarAshCost,
) -> str:
    """Derive the three-operation Elden war-ash NPC from accepted wash logic.

    The normal path retains the target's accepted native wash pipeline.  The
    rare path reuses the verified skill-affix identity/runtime contract.  The
    clear path only removes rows owned by weapon enchant.  Every path checks
    all resources before making a single atomic deduction.
    """

    if book.schema_version != 2:
        raise WeaponEnchantError("战灰刻印只支持武器附魔V2词条表")
    for cost in (normal_cost, rare_cost, clear_cost):
        _war_ash_validate_cost(cost)
    if normal_cost.gold or normal_cost.yuanbao <= 0:
        raise WeaponEnchantError("普通战灰必须使用多材料+原生元宝")
    skill_affixes = tuple(
        item for item in book.affixes
        if item.route == "skill_damage_runtime" and item.effect_status == "verified"
    )
    if len(skill_affixes) != 6:
        raise WeaponEnchantError("稀有战灰必须且只能包含6条已验证技能词条")

    text = wash_template.replace("\r\n", "\n").replace("\r", "\n").strip()
    if text.startswith("{") or text.endswith("}"):
        raise WeaponEnchantError("装备洗练底座不应包含外部CALL花括号")
    text = _replace_once(
        text,
        "<ITEMBOX:27:12:1:158:140:76:76:15,19,22,23,24,26,64,62,53,63,51,28:放入需要洗练的装备>",
        "<ITEMBOX:27:12:1:158:140:76:76:5,6:放入需要刻印战灰的武器>",
        "战灰武器框",
    )
    text = _replace_once(text, "<&Text:装备洗练:24:10", "<&Text:战灰刻印:24:10", "战灰标题")
    text = _replace_once(text, "请将需要洗练的装备放入左侧法阵。", "请将需要刻印战灰的武器放入左侧法阵。", "战灰说明")
    text, display_count = re.subn(
        r"每次消耗：[^:\r\n]+(?=:382:108\{FCOLOR=243\})",
        f"普通：{_war_ash_cost_summary(normal_cost)}",
        text,
    )
    if display_count != 1:
        raise WeaponEnchantError("verified洗练母版锚点异常：普通战灰费用")
    text = _replace_once(
        text,
        "洗练会覆盖装备原有洗练词条。",
        f"稀有：{_war_ash_cost_summary(rare_cost)}",
        "稀有战灰费用",
    )
    text = _replace_once(
        text,
        "天赐、神佑装备洗练后需要开光。",
        f"清除：{_war_ash_cost_summary(clear_cost)}",
        "清除战灰费用",
    )
    text = _replace_once(text, "<&Text:放入装备:167:266", "<&Text:放入武器:167:266", "放入武器提示")
    text = _replace_once(
        text,
        "<&Text:确认洗练:420:315{FCOLOR=243}/@XYEW_WASH_EXEC>",
        "<&Text:普通战灰:350:300{FCOLOR=243}/@XY_ASH_NORMAL>\n"
        "<&Text:稀有战灰:445:300{FCOLOR=253}/@XY_ASH_RARE>\n"
        "<&Text:清除战灰:540:300{FCOLOR=161}/@XY_ASH_CLEAR>",
        "战灰三操作按钮",
    )
    text = _replace_once(text, "[@XYEW_WASH_EXEC]", "[@XY_ASH_NORMAL]", "普通战灰入口")

    guard = _war_ash_negative_guard(normal_cost, "普通战灰")
    text, guard_count = re.subn(
        r"(?:(?:NOT CHECKITEM [^\r\n]+|NOT CHECKGOLD \d+)\n#ACT\n"
        r"MESSAGEBOX 洗练装备需要[^\r\n]+\nBREAK\n#IF\n)*"
        r"NOT CHECKGAMEGOLD > \d+",
        guard,
        text,
        flags=re.IGNORECASE,
    )
    if guard_count != 2:
        raise WeaponEnchantError("装备洗练底座的两次资源复核结构已变化")
    text, message_count = re.subn(
        r"MESSAGEBOX 洗练装备需要[^\r\n]*",
        "MESSAGEBOX 普通战灰所需战灰材料或货币不足，本次未扣除资源。",
        text,
        flags=re.IGNORECASE,
    )
    if message_count != 2:
        raise WeaponEnchantError("装备洗练底座的资源不足提示结构已变化")
    deduction = "\n".join(_war_ash_deductions(normal_cost))
    text, deduction_count = re.subn(
        r"(?:TAKE [^\r\n]+\n)*GAMEGOLD - \d+\b",
        deduction,
        text,
        count=1,
        flags=re.IGNORECASE,
    )
    if deduction_count != 1:
        raise WeaponEnchantError("装备洗练底座的最终扣费结构已变化")

    rare_checks = _war_ash_positive_checks(rare_cost)
    rare_deductions = _war_ash_deductions(rare_cost)
    clear_checks = _war_ash_positive_checks(clear_cost)
    clear_deductions = _war_ash_deductions(clear_cost)
    box = 27
    identity = book.settings.row_start
    extra = [
        "",
        "[@XY_ASH_RARE]", "#IF", f"SMALL <$BOXITEM[{box}].MAKEINDEX> 1", "#ACT",
        "MESSAGEBOX 请先放入需要刻印的武器。", "BREAK",
        "#IF", "EQUAL N$XY_ASH_BUSY 1", "#ACT", "BREAK",
        "#IF", *rare_checks, "#ACT", "MOV N$XY_ASH_BUSY 1", "GOTO @XY_ASH_RARE_PICK", "BREAK",
        "#ELSEACT", "MESSAGEBOX 所需战灰材料或货币不足，本次未扣除资源。", "BREAK", "",
        "[@XY_ASH_RARE_PICK]", *_war_ash_pick_lines(skill_affixes), "",
        "[@XY_ASH_RARE_WRITE]", "#IF", *rare_checks, "#ACT", *rare_deductions,
        f"LockUpdateItem boxitem{box}", *_clear_identity_lines(box, book.settings.clear_rows),
        f"SetCustomItemAbil boxitem{box} {identity} 0 <$STR(N$XY_WE_PICK_COLOR)>",
        f"SetCustomItemAbil boxitem{box} {identity} 1 60",
        f"SetCustomItemAbil boxitem{box} {identity} 2 {identity}",
        f"SetCustomItemAbil boxitem{box} {identity} 3 0",
        f"SetCustomItemAbil boxitem{box} {identity} 4 9",
        f"SetCustomItemValueEx boxitem{box} {identity} = <$STR(N$XY_WE_PICK_ID)> <$STR(N$XY_WE_PICK_VALUE)> <$STR(N$XY_WE_PICK_TYPE)>",
        f"SetCustomItemText boxitem{box} [战灰属性：<$STR(S$XY_WE_PICK_LABEL)>]",
        f"SetCustomItemTextColor boxitem{box} <$STR(N$XY_WE_PICK_COLOR)>",
        f"UpdateItem boxitem{box}", "MOV N$XY_ASH_BUSY 0",
        "MESSAGEBOX 稀有战灰刻印完成：<$STR(S$XY_WE_PICK_LABEL)>", "BREAK",
        "#ELSEACT", "MOV N$XY_ASH_BUSY 0", "MESSAGEBOX 所需战灰材料或货币已变化，本次未扣除资源。", "BREAK", "",
        "[@XY_ASH_CLEAR]", "#IF", f"SMALL <$BOXITEM[{box}].MAKEINDEX> 1", "#ACT",
        "MESSAGEBOX 请先放入需要清除战灰的武器。", "BREAK",
        "#IF", *clear_checks, "#ACT", *clear_deductions,
        f"LockUpdateItem boxitem{box}", *_clear_identity_lines(box, book.settings.clear_rows),
        f"SetCustomItemText boxitem{box} [战灰属性：无]", f"SetCustomItemTextColor boxitem{box} 255",
        f"UpdateItem boxitem{box}", "MESSAGEBOX 战灰已清除。", "BREAK",
        "#ELSEACT", "MESSAGEBOX 所需战灰材料或货币不足，本次未扣除资源。", "BREAK",
    ]
    return "\n".join(("{", "; XY-ELDEN-WAR-ASH-V1", text.rstrip(), *extra, "}"))


def render_elden_war_ash_runtime(book: WeaponEnchantWorkbook) -> str:
    return _render_runtime_v2(book)


def upsert_weapon_enchant_qfunction_hook(text: str) -> str:
    updated, _block = _upsert_after_label(
        text,
        "AttackDamage",
        QFUNCTION_MARKER,
        ("#IF", "#ACT", "#CALL [\\玄渊武器附魔\\武器附魔运行时.txt] @XY_WE_ATTACK_DAMAGE"),
    )
    return updated


def _selection_lines(book: WeaponEnchantWorkbook) -> list[str]:
    lines: list[str] = []
    total = sum(item.weight for item in book.affixes)
    remaining = total
    for index, affix in enumerate(book.affixes):
        last = index == len(book.affixes) - 1
        lines.append("#IF")
        if not last:
            lines.append(f"RANDOMEX {affix.weight} {remaining}")
        lines.append("#ACT")
        if affix.route == "legacy_pool":
            lines.extend(("GOTO @XY_WE_LEGACY_BEGIN", "BREAK"))
        else:
            lines.extend((
                f"MOV N$XY_WE_PICK_ID {affix.stable_id}",
                f"MOV N$XY_WE_PICK_VALUE {affix.effect_value}",
                f"MOV N$XY_WE_PICK_TYPE {affix.effect_code}",
                f"MOV N$XY_WE_PICK_COLOR {affix.color}",
                f"MOV S$XY_WE_PICK_LABEL {affix.name}",
            ))
            if affix.route == "native_bind":
                new0 = affix.effect_value if affix.new_item_index == 0 else 0
                new1 = affix.effect_value if affix.new_item_index == 1 else 0
                lines.extend((
                    f"MOV N$XY_WE_PICK_BIND {affix.native_bind}",
                    f"MOV N$XY_WE_PICK_PERCENT {affix.percent}",
                    f"MOV N$XY_WE_PICK_NEW0 {new0}",
                    f"MOV N$XY_WE_PICK_NEW1 {new1}",
                    "GOTO @XY_WE_NATIVE_WRITE",
                    "BREAK",
                ))
            else:
                lines.extend(("GOTO @XY_WE_SPECIAL_WRITE", "BREAK"))
        remaining -= affix.weight
    return lines


def _replace_once(text: str, old: str, new: str, label: str) -> str:
    if text.count(old) != 1:
        raise WeaponEnchantError(f"verified洗练母版锚点异常：{label}")
    return text.replace(old, new, 1)


def _render_core_v2(book: WeaponEnchantWorkbook, platform_root: Path | None = None) -> str:
    box = book.settings.item_box
    if box != 27:
        raise WeaponEnchantError("复用verified洗练母版时物品框编号必须为27")
    text = _legacy_template_text(platform_root)
    text = _replace_once(
        text,
        "<ITEMBOX:27:12:1:158:140:76:76:15,19,22,23,24,26,64,62,53,63,51,28:放入需要洗练的装备>",
        "<ITEMBOX:27:12:1:158:140:76:76:5,6:放入需要附魔的武器>",
        "武器物品框",
    )
    text = _replace_once(text, "<&Text:装备洗练:24:10", "<&Text:武器附魔:24:10", "标题")
    text = _replace_once(text, "请将需要洗练的装备放入左侧法阵。", "请将需要附魔的武器放入左侧法阵。", "说明")
    text = _replace_once(text, "每次消耗：100元宝", f"每次消耗：{book.settings.cost_amount}元宝", "费用显示")
    text = _replace_once(text, "not checkgamegold > 99", f"not checkgamegold > {max(0, book.settings.cost_amount - 1)}", "余额检查")
    text = _replace_once(text, "洗练装备需要100元宝", f"武器附魔需要{book.settings.cost_amount}元宝", "余额提示")
    text = _replace_once(text, "gamegold - 100", f"gamegold - {book.settings.cost_amount}", "扣费")

    legacy_rules = {item.name: item for item in book.legacy_rules}
    for name, old, command in (("神佑", 350, "RANDOM"), ("天赐", 300, "RANDOM"), ("圣级", 250, "RANDOM"), ("仙级", 150, "RANDOM"), ("传说", 10, "random"), ("上古", 5, "random"), ("灵级", 3, "random")):
        rule = legacy_rules[name]
        replacement = f"RANDOM {rule.denominator}" if rule.enabled else "RANDOMEX 0 1"
        text = _replace_once(text, f"\n{command} {old}\n", f"\n{replacement}\n", f"{name}概率")

    initial_lock = "LockUpdateItem boxitem27\nSetUpgradeItem 27\nSetCustomItemValue boxitem27 0 = 0"
    text = _replace_once(text, "SetUpgradeItem 27\nSetCustomItemValue boxitem27 0 = 0", initial_lock, "初始连续锁")
    clear_lines = _clear_identity_lines(box, book.settings.clear_rows)
    text = _replace_once(
        text,
        "SetCustomItemValue boxitem27 9 = 0\nSetNewItemValue boxitem27 0 = 0",
        "SetCustomItemValue boxitem27 9 = 0\n" + "\n".join(clear_lines) + "\nSetNewItemValue boxitem27 0 = 0",
        "身份行清理",
    )
    text = _replace_once(text, "[@区分级别洗练]", "[@XY_WE_LEGACY_BEGIN]", "普通洗练入口")
    router = "\n".join(("[@区分级别洗练]", *_selection_lines(book), "", "[@XY_WE_LEGACY_BEGIN]"))
    text = _replace_once(text, "[@XY_WE_LEGACY_BEGIN]", router, "结果路由插入")

    identity_write = [
        f"SetCustomItemAbil boxitem{box} {book.settings.row_start} 0 <$STR(N$XY_WE_PICK_COLOR)>",
        f"SetCustomItemAbil boxitem{box} {book.settings.row_start} 1 60",
        f"SetCustomItemAbil boxitem{box} {book.settings.row_start} 2 {book.settings.row_start}",
        f"SetCustomItemAbil boxitem{box} {book.settings.row_start} 3 0",
        f"SetCustomItemAbil boxitem{box} {book.settings.row_start} 4 9",
        f"SetCustomItemValueEx boxitem{box} {book.settings.row_start} = <$STR(N$XY_WE_PICK_ID)> <$STR(N$XY_WE_PICK_VALUE)> <$STR(N$XY_WE_PICK_TYPE)>",
    ]
    special = [
        "[@XY_WE_SPECIAL_WRITE]", "#IF", "#ACT", f"LockUpdateItem boxitem{box}",
        *_clear_identity_lines(box, book.settings.clear_rows),
        *identity_write,
        f"SetCustomItemText boxitem{box} [{book.settings.display_prefix}<$STR(S$XY_WE_PICK_LABEL)>]",
        f"SetCustomItemTextColor boxitem{box} <$STR(N$XY_WE_PICK_COLOR)>",
        f"UpdateItem boxitem{box}",
        "MOV N$XYEW_W_洗练CD 0",
        "MESSAGEBOX 武器附魔完成：<$STR(S$XY_WE_PICK_LABEL)>",
        "BREAK", "",
        "[@XY_WE_NATIVE_WRITE]", "#IF", "#ACT", f"LockUpdateItem boxitem{box}",
        *_clear_identity_lines(box, book.settings.clear_rows),
        f"SetCustomItemAbil boxitem{box} 0 0 <$STR(N$XY_WE_PICK_COLOR)>",
        f"SetCustomItemAbil boxitem{box} 0 1 <$STR(N$XY_WE_PICK_BIND)>",
        f"SetCustomItemAbil boxitem{box} 0 2 0",
        f"SetCustomItemAbil boxitem{box} 0 3 <$STR(N$XY_WE_PICK_PERCENT)>",
        f"SetCustomItemValue boxitem{box} 0 = <$STR(N$XY_WE_PICK_VALUE)>",
        f"SetNewItemValue boxitem{box} 0 = <$STR(N$XY_WE_PICK_NEW0)>",
        f"SetNewItemValue boxitem{box} 1 = <$STR(N$XY_WE_PICK_NEW1)>",
        *identity_write,
        f"SetCustomItemText boxitem{box} [{book.settings.display_prefix}<$STR(S$XY_WE_PICK_LABEL)>]",
        f"SetCustomItemTextColor boxitem{box} <$STR(N$XY_WE_PICK_COLOR)>",
        f"UpdateItem boxitem{box}",
        "MOV N$XYEW_W_洗练CD 0",
        "MESSAGEBOX 武器附魔完成：<$STR(S$XY_WE_PICK_LABEL)>",
        "BREAK",
    ]
    return "\n".join(("{", "; XY-WEAPON-ENCHANT-V2-DERIVED", text.rstrip(), *special, "}"))


def _render_runtime_v2(book: WeaponEnchantWorkbook) -> str:
    row = book.settings.row_start
    lines = [
        "{", "; XY-WEAPON-ENCHANT-V2-DERIVED", "[@XY_WE_ATTACK_DAMAGE]", "#IF", "#ACT",
        f"GetCustomItemValueEx 1 {row} N$XY_WE_T N$XY_WE_ID N$XY_WE_VALUE N$XY_WE_TYPE",
    ]
    for affix in book.affixes:
        if affix.route != "skill_damage_runtime":
            continue
        lines.extend((
            "#IF",
            f"EQUAL N$XY_WE_ID {affix.stable_id}",
            f"EQUAL <$CURRRUSEMAGICID> {affix.skill_id}",
            "#ACT",
            "ChangeDamageValue 1 + <$STR(N$XY_WE_VALUE)>",
        ))
    lines.extend(("BREAK", "}"))
    return "\n".join(lines)


def _render_runtime(book: WeaponEnchantWorkbook) -> str:
    return _render_runtime_v2(book) if book.schema_version == 2 else _render_runtime_v1(book)


def _render_core(book: WeaponEnchantWorkbook, platform_root: Path | None = None) -> str:
    return _render_core_v2(book, platform_root) if book.schema_version == 2 else _render_core_v1(book)


def _merchant_line(book: WeaponEnchantWorkbook) -> str:
    s = book.settings
    script = str(PurePosixPath(s.npc_relative).relative_to("Mir200/Envir/Market_Def")).removesuffix(f"-{s.npc_map}.txt")
    return f"{script}\t{s.npc_map}\t{s.npc_x}\t{s.npc_y}\t{s.npc_name}\t0\t{s.npc_appearance}\t0"


def _assert_verified_wash_dependency(root: Path) -> None:
    effect_path = _target(root, LEGACY_WASH_EFFECT)
    if not effect_path.is_file() or LEGACY_WASH_EFFECT_LINE not in _script_text(effect_path.read_bytes()).splitlines():
        raise WeaponEnchantError("未检测到已验证装备洗练底座的EffectImageList资源注册，请先安装xy.optional.equipment-wash-opening")
    for relative, expected_hash in LEGACY_WASH_SERVER_DEPENDENCIES:
        path = _target(root, relative)
        if not path.is_file():
            raise WeaponEnchantError(f"未检测到已验证装备洗练底座文件：{relative.as_posix()}")
        actual_hash = _sha256(path.read_bytes()).upper()
        if actual_hash != expected_hash:
            raise WeaponEnchantError(f"装备洗练底座文件已漂移，禁止派生覆盖：{relative.as_posix()}")


class WeaponEnchantService(_DualService):
    route = ROUTE
    package_id = PACKAGE_ID

    def preflight(self, server: Path, workbook: Path) -> DualInstallPlan:
        source = Path(workbook).resolve()
        workbook_hash = _sha256(source.read_bytes()) if source.is_file() else ""
        blockers, warnings, changes = [], ["文件回滚不会恢复已写入玩家武器实例；实例恢复需另行洗回或使用备份装备。"], []
        try:
            root = Path(server).resolve()
            if not (root / "Mir200/Envir").is_dir():
                raise WeaponEnchantError("目标目录不是可识别的翎风/LFM2服务端")
            book = load_weapon_enchant_workbook(source)
            if book.schema_version == 1:
                warnings.extend(("旧五表按V1兼容生成；建议迁移到正式42号V2中文表。", "技能专属必定暴击/致命一击沿用旧candidate状态。", "品质权重和颜色已进入候选脚本；持久保底尚无当前端已确认的人物持久变量出口。"))
                enabled_pity = [rule.name for rule in book.qualities if rule.pity > 0]
                if enabled_pity:
                    raise WeaponEnchantError("保底持久化路线待验证；请将以下品质的保底次数设为0：" + "、".join(enabled_pity))
                unsupported = [a.name for a in book.affixes if a.route == "static_only"]
                if unsupported:
                    raise WeaponEnchantError("以下词条没有真实属性出口，禁止仅生成显示：" + "、".join(unsupported))
            else:
                warnings.append("武器附魔V2复用已验收装备洗练母版；新增组合仍需在目标服完成一次游戏验收。")
                blocked_affixes = [a.property_name or a.name for a in book.affixes if a.route == "candidate" or a.effect_status != "verified"]
                if blocked_affixes:
                    raise WeaponEnchantError("以下启用词条实效路线待验证，禁止安装：" + "、".join(blocked_affixes))
                _assert_verified_wash_dependency(root)
            for relative, rendered in ((CORE_RELATIVE, _render_core(book, self.platform_root)), (RUNTIME_RELATIVE, _render_runtime(book))):
                path = _target(root, relative)
                before = path.read_bytes() if path.is_file() else None
                after = _script_bytes(rendered)
                if before != after:
                    changes.append(_change(root, relative, "exclusive", before, after))
            npc_relative = Path(book.settings.npc_relative)
            npc_path = _target(root, npc_relative)
            npc_before = npc_path.read_bytes() if npc_path.is_file() else None
            npc_after = _script_bytes("[@main]\n#IF\n#ACT\n#CALL [\\玄渊武器附魔\\武器附魔核心.txt] @main\nBREAK")
            if npc_before != npc_after:
                changes.append(_change(root, npc_relative, "exclusive", npc_before, npc_after))
            qf_path = _target(root, QFUNCTION_RELATIVE)
            if not qf_path.is_file():
                raise WeaponEnchantError("QFunction-0.txt不存在")
            qf_before = qf_path.read_bytes()
            qf_before_text = _script_text(qf_before)
            qf_after_text, block = _upsert_after_label(qf_before_text, "AttackDamage", QFUNCTION_MARKER, ("#IF", "#ACT", "#CALL [\\玄渊武器附魔\\武器附魔运行时.txt] @XY_WE_ATTACK_DAMAGE"))
            qf_after = qf_before if qf_before_text.rstrip("\n") == qf_after_text.rstrip("\n") else _script_bytes(qf_after_text)
            if qf_before != qf_after:
                changes.append(_change(root, QFUNCTION_RELATIVE, "managed_blocks", qf_before, qf_after, {"markers": [QFUNCTION_MARKER], "blocks": [block]}))
            merchant_path = _target(root, MERCHANT_RELATIVE)
            if not merchant_path.is_file():
                raise WeaponEnchantError("MerChant.txt不存在")
            merchant_before = merchant_path.read_bytes()
            merchant_text = _script_text(merchant_before)
            line = _merchant_line(book)
            conflicting = [item for item in merchant_text.splitlines() if item.strip() and item.split()[0].replace("\\", "/") == line.split()[0] and item.strip() != line]
            if conflicting:
                raise WeaponEnchantError("MerChant已存在同脚本不同配置，禁止覆盖")
            if line not in merchant_text.splitlines():
                merchant_after = _script_bytes(merchant_text.rstrip("\n") + ("\n" if merchant_text.strip() else "") + line)
                changes.append(_change(root, MERCHANT_RELATIVE, "exclusive", merchant_before, merchant_after))
        except (OSError, ValueError, WeaponEnchantError) as exc:
            blockers.append(str(exc))
        return DualInstallPlan(ROUTE, PACKAGE_ID, _plan_id(ROUTE, workbook_hash, changes), Path(server).resolve(), source, workbook_hash, tuple(changes), tuple(blockers), tuple(warnings))
